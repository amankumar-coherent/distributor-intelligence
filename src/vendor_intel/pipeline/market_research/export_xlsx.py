"""Excel export for a market's full accumulated company list.

Re-checks the WHOLE state file (every batch ever returned for this market,
not just the latest one) for duplicates and incomplete rows before writing —
schema.finalize_row() already prevents an incomplete row from ever entering
state, and state.add_batch() already dedupes by company_dedupe_key, but this
export-time pass is a second, independent guard against either of those
being bypassed (e.g. a state file hand-edited or written by an older
pipeline version) rather than trusting that invariant blindly.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from vendor_intel.pipeline.market_research import schema
from vendor_intel.pipeline.market_research.state import MarketResearchState, market_slug
from vendor_intel.utils.domains import company_dedupe_key

OUTPUT_DIR = Path(__file__).resolve().parents[4] / "output" / "market_research"


def _one_line(value: object) -> str:
    """Collapse whitespace so a value can never force a row taller.

    A single embedded newline makes Excel render that row multi-line even
    with wrap_text off, which is exactly the "expanded cell" look this
    export avoids. AI Mode answers occasionally carry stray newlines
    inside a value, so normalise here rather than trusting upstream.
    """
    return " ".join(str(value or "").split())


def _dedupe_and_validate(
    companies: list[dict[str, str]], *, market: str = ""
) -> tuple[list[dict[str, str]], list[str]]:
    """Returns (clean_rows, dropped_reasons). A row is dropped if its
    company_dedupe_key repeats an earlier row, or if any required column
    (per schema.REQUIRED_FIELDS) is missing/placeholder — never exported
    with a blank/N/A cell."""
    seen: set[str] = set()
    clean: list[dict[str, str]] = []
    dropped: list[str] = []

    for row in companies:
        name = row.get("Company", "")
        if not name:
            dropped.append("(unnamed row) — missing Company")
            continue
        key = company_dedupe_key(name)
        if key in seen:
            dropped.append(f"{name} — duplicate of an earlier row")
            continue

        # Only the firmographic columns are required. The six contact
        # fields are legitimately blank when the real value could not be
        # verified (schema.CONTACT_FIELDS), so they must not drop a row.
        # Market-aware: a column this market does not export (the
        # physical-channel column, for software markets) must not be able
        # to drop a row for being blank.
        required = (schema.required_fields_for(market) if market
                    else schema.REQUIRED_FIELDS)
        missing = [
            col for col in required if schema.is_placeholder(row.get(col))
        ]
        if missing:
            dropped.append(f"{name} — incomplete: {', '.join(missing)}")
            continue

        seen.add(key)
        clean.append(row)

    return clean, dropped



def _definition_banner(market: str) -> list[tuple[str, bool]]:
    """Lines to print above the table: the market and the inclusion rule.

    Returns (text, is_title) pairs. The rule text is derived from the SAME
    function the discovery prompt uses, so the banner can never drift from
    the criteria that actually selected the rows.
    """
    from vendor_intel.pipeline.market_research.discovery import (
        _DEF_SOFTWARE_SERVICES,
        distributor_definition_for,
    )

    # Which rule did this market actually use? The banner has to match
    # the collection rule, or it describes data that was never gathered
    # that way -- a training market was carrying the physical-goods
    # wording above a list of pilot academies.
    from vendor_intel.pipeline.market_research.discovery import (
        _DEF_HYBRID_TECH,
        _DEF_SERVICE_PROVIDER,
    )

    definition = distributor_definition_for(market)

    if definition is _DEF_SERVICE_PROVIDER:
        rule = (
            "SERVICE PROVIDER DEFINITION (services market): a company that "
            "PERFORMS this service commercially for third-party customers - "
            "it is contracted to do the work, or sells the course, "
            "inspection, survey or certification as its own offering. "
            "Operators, contractors, consultancies, accredited training "
            "organisations and certification bodies all qualify when this "
            "work is what customers pay them for. EXCLUDED: manufacturers "
            "whose only service is support or familiarisation for their own "
            "products; equipment sellers with no service delivery; "
            "marketplaces and directories that merely list providers."
        )
    elif definition is _DEF_HYBRID_TECH:
        rule = (
            "DISTRIBUTOR / CHANNEL PARTNER DEFINITION (hardware + software "
            "market): this market sells both hardware and software, so two "
            "shapes qualify - a stocking or authorized DISTRIBUTOR buying "
            "hardware in volume and reselling it, and a CHANNEL PARTNER, "
            "value-added reseller or SYSTEM INTEGRATOR that resells, bundles "
            "or implements another vendor's platform, software or modules. "
            "EXCLUDED: manufacturers, OEMs and their sales subsidiaries; "
            "vendors selling their own platform; consultancies with nothing "
            "to resell; and consumer-only retailers."
        )
    elif definition is _DEF_SOFTWARE_SERVICES:
        rule = (
            "DISTRIBUTOR DEFINITION (software / cloud market): a company whose "
            "business is RESELLING, distributing, white-labelling or aggregating "
            "another vendor's platform to its own customers - authorized "
            "reseller, master agent / master distributor, value-added reseller, "
            "white-label or wholesale aggregator, or a telecom aggregator "
            "reselling messaging/voice capacity. In this market the CHANNEL "
            "PARTNERS ARE THE DISTRIBUTORS, because there is no physical stock "
            "to warehouse. EXCLUDED: the platform vendors themselves and their "
            "own sales subsidiaries, system integrators and consultancies "
            "building bespoke solutions, and referral-only affiliates."
        )
    else:
        rule = (
            "DISTRIBUTOR DEFINITION (physical goods market): a company whose "
            "PRIMARY business is buying in volume from manufacturers and "
            "reselling/distributing onward - wholesale, authorized, "
            "value-added or stocking distributor. EXCLUDED: channel partners, "
            "resellers, VARs, dealers, agents and brokers that do not stock or "
            "wholesale; manufacturers and OEMs (including a manufacturer's own "
            "sales subsidiary); system integrators; and pure service providers "
            "or consultancies."
        )
    note = (
        "Continent / Geography = countries with a physical site (HQ, office, "
        "warehouse, plant). Operational Presence = every country the company "
        "does business in, which is WIDER. Blank or NPD in a contact column "
        "means no published evidence was found - it is never a guess."
    )
    # The product scope, when the market has one. Without it the sheet
    # states which CHANNEL ROLE was required but not which PRODUCT, and
    # that was exactly the ambiguity that let ophthalmic equipment
    # distributors into a dry-eye drug market unnoticed.
    from vendor_intel.pipeline.market_research.discovery import product_scope_for

    scope = product_scope_for(market)
    rows = [(market, True), (rule, False)]
    if scope:
        rows.append((" ".join(scope.split()), False))
    rows.append((note, False))
    return rows

def export_xlsx(market: str, *, path: Path | None = None) -> Path:
    """Writes every verified company accumulated so far for `market` to an
    .xlsx file, after a final dedup + completeness pass across the whole
    list. Returns the path written."""
    state = MarketResearchState.load(market)
    clean_rows, dropped = _dedupe_and_validate(state.companies, market=market)

    # Not schema.COLUMNS: a software/cloud market drops the
    # "Online / Offline / Both" channel column, which has no meaning when
    # the product is delivered over the internet. See schema.columns_for.
    columns = schema.columns_for(market)

    out_path = path or (OUTPUT_DIR / f"{market_slug(market)}.xlsx")
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    cell_font = Font(name="Calibri", size=9)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    # No wrap_text: wrapping is what forced every row to be multiple lines
    # tall. Long values (the per-country lists especially) now stay on one
    # line and are read either by widening the column or from the formula
    # bar / the cell's own tooltip, which keeps the sheet scannable — 493
    # rows at ~6 lines each is unusable for scrolling or comparing rows.
    plain = Alignment(wrap_text=False, vertical="center", horizontal="left")
    plain_c = Alignment(wrap_text=False, vertical="center", horizontal="center")
    center_cols = {"#", "Founded", "Online / Offline / Both", "Country Code / Region Code", "Office No."}

    # Per-column widths, tuned for single-line rows. Kept moderate rather
    # than sized to the longest value: the geography columns can hold a
    # 300-character country list, and a 300-wide column would push every
    # other field off-screen. They are wide enough to show the leading
    # region plus the first countries, and the full value is available in
    # the formula bar (or by widening that one column).
    _WIDTHS = {
        "#": 5,
        "Company": 32,
        "Website": 30,
        "Founded": 9,
        "Headquarters": 28,
        "Continent / Geography": 46,
        "Operational Presence": 46,
        "Ownership": 18,
        "Employees": 15,
        "Core Categories": 34,
        "Specialty Focus": 38,
        "Key Products / Technologies": 38,
        "Online / Offline / Both": 13,
        "Contact Person": 22,
        "Role": 26,
        "Email": 30,
        "LinkedIn": 36,
        "Office No.": 18,
        "Office Website": 30,
        "Office Email": 30,
        "Country Code / Region Code": 34,
    }
    # Single-line rows. Excel's default is ~15px; 16 gives the 9pt font a
    # little breathing room without making the sheet tall. Explicitly set
    # (rather than left to "auto") so the height is deterministic and does
    # not grow back if a cell ever contains a newline.
    _ROW_HEIGHT = 16

    # Definition banner above the table.
    #
    # Requested because "distributor" does not mean the same thing in every
    # market: in CPaaS and other software markets the channel partners ARE
    # the distributors, while in a physical-goods market they are explicitly
    # excluded. Anyone opening the file needs to know which rule produced
    # these rows, or the inclusion criteria are invisible and the sheet
    # cannot be audited.
    banner_rows = _definition_banner(market)
    for text, is_title in banner_rows:
        ws.append([text])
        r = ws.max_row
        cell = ws.cell(r, 1)
        cell.font = Font(bold=is_title, size=11 if is_title else 9,
                         color="1F4E79" if is_title else "333333")
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        ws.merge_cells(start_row=r, start_column=1,
                       end_row=r, end_column=len(columns))
        ws.row_dimensions[r].height = 18 if is_title else 30
    ws.append([])                       # blank spacer row

    ws.append(columns)
    # AFTER appending, so this is the row the header actually landed on.
    # Computing it beforehand was off by one, which anchored the autofilter
    # on the spacer row instead of the header.
    header_row = ws.max_row
    for i, h in enumerate(columns, 1):
        cell = ws.cell(header_row, i)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=False, vertical="center", horizontal="center")
        cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = _WIDTHS.get(h, 22)
    ws.row_dimensions[header_row].height = 20
    # Freeze below the banner AND the header, so both stay visible.
    ws.freeze_panes = f"B{header_row + 1}"

    for idx, row in enumerate(clean_rows, 1):
        values = [
            idx if col == "#" else _one_line(row.get(col, "")) for col in columns
        ]
        ws.append(values)
        r = ws.max_row
        for i, h in enumerate(columns, 1):
            cell = ws.cell(r, i)
            cell.font = cell_font
            cell.border = thin
            cell.alignment = plain_c if h in center_cols else plain
        ws.row_dimensions[r].height = _ROW_HEIGHT

    last = get_column_letter(len(columns))
    # Anchor the filter on the REAL header row, not A1 — the banner now
    # occupies the first rows, and filtering from A1 would treat the
    # definition text as the header and break sorting.
    ws.auto_filter.ref = f"A{header_row}:{last}{ws.max_row}"

    if dropped:
        audit_ws = wb.create_sheet("Export Audit")
        audit_ws.append(["Dropped at export time (duplicate or incomplete)"])
        for reason in dropped:
            audit_ws.append([reason])
        audit_ws.column_dimensions["A"].width = 90

    wb.save(out_path)
    return out_path
