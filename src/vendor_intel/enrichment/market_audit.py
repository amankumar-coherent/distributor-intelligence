"""Web-search + LLM market audit (final_audit_v2 logic), batch-friendly."""
from __future__ import annotations

import asyncio
import csv
import io
import json
import os
import time
from pathlib import Path
from typing import Any


def load_companies_from_csv(csv_path: Path) -> list[dict[str, str]]:
    if not csv_path.is_file():
        return []
    lines = csv_path.read_text(encoding="utf-8", errors="replace").splitlines(keepends=True)
    if not lines:
        return []
    header = lines[0]
    data_lines = [line for line in lines[1:] if not line.strip().startswith("===")]
    reader = csv.DictReader(io.StringIO(header + "".join(data_lines)))
    result: list[dict[str, str]] = []
    for row in reader:
        name = (row.get("Company") or "").strip()
        if not name:
            continue
        result.append(
            {
                "name": name,
                "website": (row.get("Website") or "").strip(),
                "functionality": (row.get("Functionality") or row.get("functionality") or "").strip(),
                "key_products": (row.get("Key_Products") or row.get("key_products") or "").strip(),
                "summary": (row.get("Summary") or row.get("summary") or "").strip()[:300],
                "role": (row.get("Role") or row.get("role") or "").strip(),
                "confidence": (row.get("Confidence") or row.get("confidence") or "").strip(),
                "headquarters": (row.get("Headquarters") or row.get("headquarters") or "").strip(),
            }
        )
    return result


async def _web_search(query: str, num: int = 5) -> list[dict]:
    try:
        from ddgs import DDGS

        def _search() -> list:
            with DDGS() as ddgs:
                return list(ddgs.text(query, max_results=num))

        return await asyncio.to_thread(_search)
    except Exception:
        return []


async def gather_company_evidence(company: dict[str, str]) -> str:
    name = company["name"]
    website = company.get("website", "")
    queries = [
        f'"{name}" company what do they do products services',
        f'"{name}" {website} official',
    ]
    snippets: list[str] = []
    for q in queries:
        results = await _web_search(q, num=4)
        for r in results:
            title = r.get("title", "")
            body = r.get("body", "")
            href = r.get("href", "")
            snippets.append(f"[{title}] {href}\n{body[:200]}")
        await asyncio.sleep(0.3)
    return "\n\n".join(snippets)[:3000]


async def audit_one_company(
    company: dict[str, str],
    *,
    market_query: str,
    market_desc: str,
    client: Any,
    model: str,
    sem: asyncio.Semaphore,
) -> dict[str, Any]:
    async with sem:
        name = company["name"]
        profile = f"""Company: {name}
Website: {company.get('website', '')}
Role: {company.get('role', '')}
Functionality: {company.get('functionality', '')}
Key Products: {company.get('key_products', '')}
Headquarters: {company.get('headquarters', '')}
Summary: {company.get('summary', '')}"""
        evidence = await gather_company_evidence(company)
        prompt = f"""You are a strict market research auditor. Verify if this company truly operates in the stated market.

MARKET: {market_query}
MARKET SCOPE: {market_desc}

COMPANY PROFILE (from our database):
{profile}

WEB EVIDENCE (from live search):
{evidence}

Based on ALL evidence above, decide:
1. Does this company GENUINELY operate in this market?
2. Are their products/services a real match for the market scope?
3. Is this a legitimate, identifiable business (not a generic name or fake entry)?

STRICT RULES:
- "keep"   = clear evidence company operates in this market
- "review" = partial match or insufficient evidence — needs manual check
- "remove" = wrong industry, not a real company, or clearly irrelevant
- Use web evidence to override the profile if they contradict

Return JSON:
{{
  "company": "{name}",
  "action": "keep|review|remove",
  "market_relevant": "yes|partial|no|unclear",
  "what_they_do": "1-2 sentences from web evidence",
  "red_flags": "none or specific issue",
  "confidence": "high|medium|low",
  "reason": "brief justification using evidence"
}}"""
        for attempt in range(3):
            try:
                resp = await client.chat.completions.create(
                    model=model,
                    messages=[{"role": "user", "content": prompt}],
                    temperature=0,
                    max_tokens=350,
                    response_format={"type": "json_object"},
                )
                result = json.loads(resp.choices[0].message.content or "{}")
                result["website"] = company.get("website", "")
                result["original_confidence"] = company.get("confidence", "")
                return result
            except Exception:
                if attempt < 2:
                    await asyncio.sleep(2**attempt)
        return {
            "company": name,
            "action": "review",
            "market_relevant": "unclear",
            "what_they_do": "",
            "red_flags": "API error",
            "confidence": "low",
            "reason": "Could not audit",
            "website": company.get("website", ""),
        }


def _find_header_row(ws) -> int:
    for i, row in enumerate(ws.rows, 1):
        if row[0].value == "Company":
            return i
    raise ValueError("No 'Company' header row")


def remove_from_excel(xlsx_path: Path, remove_set: set[str]) -> int:
    import openpyxl

    if not xlsx_path.is_file():
        return 0
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    header_row = _find_header_row(ws)
    rows_to_delete: list[int] = []
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        val = row[0].value
        if val and str(val).strip().lower() in remove_set:
            rows_to_delete.append(row[0].row)
    for rn in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(rn)
    wb.save(xlsx_path)
    return len(rows_to_delete)


def remove_from_csv(csv_path: Path, remove_set: set[str]) -> int:
    if not csv_path.is_file():
        return 0
    lines = csv_path.read_text(encoding="utf-8", errors="replace").splitlines(keepends=True)
    if not lines:
        return 0
    header = lines[0]
    data_lines = [line for line in lines[1:] if not line.strip().startswith("===")]
    reader = csv.DictReader(io.StringIO(header + "".join(data_lines)))
    fieldnames = reader.fieldnames or []
    all_rows = list(reader)
    kept = [r for r in all_rows if r.get("Company", "").strip().lower() not in remove_set]
    removed = len(all_rows) - len(kept)
    with csv_path.open("w", encoding="utf-8", newline="", errors="replace") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(kept)
    return removed


def apply_removals_to_result(result: dict[str, Any], remove_set: set[str]) -> int:
    rows = result.get("relevant_companies") or []
    if not rows or not remove_set:
        return 0
    kept = [
        r
        for r in rows
        if str(r.get("company") or r.get("brand") or "").strip().lower() not in remove_set
    ]
    removed = len(rows) - len(kept)
    result["relevant_companies"] = kept
    return removed


def write_market_report(
    report_path: Path,
    *,
    market_query: str,
    market_desc: str,
    results: list[dict],
) -> None:
    keep = [r for r in results if r.get("action") == "keep"]
    review = [r for r in results if r.get("action") == "review"]
    remove_list = [r for r in results if r.get("action") == "remove"]

    lines = [
        f"# Audit Report — {market_query}\n\n",
        f"**Market scope:** {market_desc}\n\n",
        "| Status | Count |\n|--------|-------|\n",
        f"| Keep | {len(keep)} |\n",
        f"| Review | {len(review)} |\n",
        f"| Remove | {len(remove_list)} |\n",
        f"| **Total** | **{len(results)}** |\n\n",
        "## All Companies\n\n",
        "| # | Company | Action | Market Relevant | What They Do | Red Flags |\n",
        "|---|---------|--------|-----------------|--------------|----------|\n",
    ]
    for i, r in enumerate(results, 1):
        lines.append(
            f"| {i} | **{r.get('company', '?')}** | {r.get('action', '?')} | "
            f"{r.get('market_relevant', '?')} | {str(r.get('what_they_do', ''))[:70]} | "
            f"{str(r.get('red_flags', 'none'))[:60]} |\n"
        )
    if remove_list:
        lines.append("\n## Companies to Remove\n\n")
        lines.append("| Company | Website | Reason |\n|---------|---------|--------|\n")
        for r in remove_list:
            lines.append(
                f"| {r.get('company', '')} | {r.get('website', '')} | "
                f"{str(r.get('reason', ''))[:80]} |\n"
            )
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("".join(lines), encoding="utf-8")


async def audit_market_exports(
    *,
    output_dir: Path,
    market_query: str,
    csv_path: Path | None = None,
    xlsx_path: Path | None = None,
    json_path: Path | None = None,
    result: dict[str, Any] | None = None,
    market_type: str = "general",
    market_desc: str | None = None,
    concurrent: int = 8,
    fmcg_xlsx: bool = True,
) -> dict[str, Any]:
    """Audit exported companies; remove 'remove' verdicts from CSV/XLSX/JSON."""
    from backend.config import get_settings, sync_opencode_to_openai_env
    from openai import AsyncOpenAI

    sync_opencode_to_openai_env()
    cfg = get_settings()
    if not cfg.openai_api_key:
        print("  [audit] OPENAI_API_KEY not set — skipping audit.", flush=True)
        return {"skipped": True, "reason": "no_api_key"}

    output_dir = Path(output_dir)
    csv_path = Path(csv_path) if csv_path else None
    xlsx_path = Path(xlsx_path) if xlsx_path else None
    json_path = Path(json_path) if json_path else None

    if csv_path is None or not csv_path.is_file():
        candidates = sorted(output_dir.glob("*.csv"))
        csv_path = candidates[0] if candidates else None
    if csv_path is None or not csv_path.is_file():
        print(f"  [audit] No CSV found in {output_dir} — skipping.", flush=True)
        return {"skipped": True, "reason": "no_csv"}

    if xlsx_path is None or not xlsx_path.is_file():
        xlsx_candidates = sorted(output_dir.glob("*.xlsx"))
        xlsx_path = xlsx_candidates[0] if xlsx_candidates else None

    if json_path is None or not json_path.is_file():
        json_candidates = sorted(output_dir.glob("*.json"))
        json_path = json_candidates[0] if json_candidates else None

    desc = (market_desc or market_query or "this market").strip()
    cache_path = output_dir / "audit_cache.json"
    report_path = output_dir / "audit_report.md"

    cached: dict[str, dict] = {}
    if cache_path.is_file():
        try:
            cached_list = json.loads(cache_path.read_text(encoding="utf-8"))
            cached = {r["company"].lower(): r for r in cached_list if r.get("company")}
            print(f"  [audit] Loaded {len(cached)} cached results", flush=True)
        except Exception:
            cached = {}

    companies = load_companies_from_csv(csv_path)
    print(f"  [audit] {market_query}: {len(companies)} companies to audit", flush=True)
    if not companies:
        return {"skipped": True, "reason": "empty_csv", "total": 0}

    to_audit = [c for c in companies if c["name"].lower() not in cached]
    if to_audit:
        client = AsyncOpenAI(api_key=cfg.openai_api_key)
        model = (os.getenv("OPENAI_MODEL") or cfg.model_extract or "gpt-4o-mini").strip()
        sem = asyncio.Semaphore(max(1, concurrent))
        t0 = time.perf_counter()
        new_results = await asyncio.gather(
            *[
                audit_one_company(
                    c,
                    market_query=market_query,
                    market_desc=desc,
                    client=client,
                    model=model,
                    sem=sem,
                )
                for c in to_audit
            ]
        )
        print(
            f"  [audit] Audited {len(new_results)} companies in {time.perf_counter() - t0:.1f}s",
            flush=True,
        )
        for r in new_results:
            if r and r.get("company"):
                cached[r["company"].lower()] = r
        cache_path.write_text(json.dumps(list(cached.values()), indent=2, ensure_ascii=False), encoding="utf-8")

    all_results = [cached[c["name"].lower()] for c in companies if c["name"].lower() in cached]
    keep = [r for r in all_results if r.get("action") == "keep"]
    review = [r for r in all_results if r.get("action") == "review"]
    remove_list = [r for r in all_results if r.get("action") == "remove"]
    print(
        f"  [audit] Results: {len(keep)} keep | {len(review)} review | {len(remove_list)} remove",
        flush=True,
    )

    xl_removed = csv_removed = json_removed = 0
    if remove_list:
        remove_set = {r["company"].strip().lower() for r in remove_list}
        csv_removed = remove_from_csv(csv_path, remove_set)
        if xlsx_path and xlsx_path.is_file():
            xl_removed = remove_from_excel(xlsx_path, remove_set)
        if result is not None:
            json_removed = apply_removals_to_result(result, remove_set)
            if json_path:
                json_path.write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
        print(
            f"  [audit] Removed rows — CSV: {csv_removed}, Excel: {xl_removed}, JSON: {json_removed}",
            flush=True,
        )

    write_market_report(report_path, market_query=market_query, market_desc=desc, results=all_results)

    stats = {
        "total": len(all_results),
        "keep": len(keep),
        "review": len(review),
        "remove": len(remove_list),
        "removed_names": [r["company"] for r in remove_list],
        "audit_cache_path": str(cache_path),
        "audit_report_path": str(report_path),
    }
    if result is not None:
        result["market_audit"] = stats
        result["market_audit_completed_at"] = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    return stats
