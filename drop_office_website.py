#!/usr/bin/env python3
"""Remove the "Office Website" column from finished workbooks.

The pipeline still exports it, but it duplicates "Website" for almost
every row and was dropped from the delivered files by request. Copying a
fresh export over a cleaned Desktop workbook silently reintroduces it,
which is what this repairs.

Edits in place, preserving the banner rows above the header, the column
widths, the freeze pane and the autofilter -- reformatting the sheet
would undo the layout work already done on these files.

Usage:
    python drop_office_website.py <file.xlsx> [more.xlsx ...]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TARGET = "office website"


def header_row(ws, limit: int = 12) -> int | None:
    """Row index of the real header, past the banner.

    Identified by width: the banner rows are single merged cells, the
    header is the only wide row near the top.
    """
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit,
                                         values_only=True), 1):
        filled = sum(1 for c in row if c is not None and str(c).strip())
        if filled > 3:
            return i
    return None


def drop(path: Path) -> str:
    wb = load_workbook(path)
    ws = wb.active
    hdr = header_row(ws)
    if hdr is None:
        return "no header found - skipped"

    target_col = None
    for cell in ws[hdr]:
        if cell.value and str(cell.value).strip().lower() == TARGET:
            target_col = cell.column
            break
    if target_col is None:
        return "already removed"

    # Widths are keyed by letter and do NOT shift with the delete, so they
    # must be re-mapped by hand or every column past the deleted one keeps
    # the width of its left neighbour.
    widths = {letter: dim.width
              for letter, dim in ws.column_dimensions.items() if dim.width}
    ws.delete_cols(target_col)

    ncols = ws.max_column
    for letter in list(ws.column_dimensions):
        del ws.column_dimensions[letter]
    for idx in range(1, ncols + 1):
        src = idx if idx < target_col else idx + 1
        w = widths.get(get_column_letter(src))
        if w:
            ws.column_dimensions[get_column_letter(idx)].width = w

    # Autofilter spans a range that no longer exists after the delete.
    if ws.auto_filter.ref:
        ws.auto_filter.ref = (
            f"A{hdr}:{get_column_letter(ncols)}{ws.max_row}")

    wb.save(path)
    return f"removed (now {ncols} cols)"


def main() -> None:
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        raise SystemExit(2)
    for raw in args:
        p = Path(raw)
        if not p.is_file():
            print(f"  {p.name}: not found")
            continue
        print(f"  {p.name}: {drop(p)}")


if __name__ == "__main__":
    main()
