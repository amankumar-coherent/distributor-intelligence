# -*- coding: utf-8 -*-
"""Agent web-search fills for smartwatch + ADC empty Contact Person / Role / LinkedIn."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

NPD = "Not publicly disclosed"
EMPTY = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}

SW_PATH = Path(
    r"D:\vendor-intel\output\run4\global_smartwatch_market_global"
    r"\global_smartwatch_market.xlsx"
)
ADC_PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global"
    r"\global_antibody_drug_conjugates_market.xlsx"
)

# Only write if Contact Person currently empty. Sales/procurement titles only.
SW = {
    "Dicker Data": {
        "person": "Yasser Elgammal",
        "role": "Head of Sales, ANZ",
        "li": "https://www.linkedin.com/in/yasserelgammal",
    },
    "CDW": {
        "person": "Liz Connelly",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/liz-h-connelly",
    },
    "D&H Distributing": {
        "person": "Marty Bauerlein",
        "role": "Chief Consumer and Commercial Officer",
        "li": "https://www.linkedin.com/in/marty-bauerlein",
    },
    "Westcoast": {
        "person": "Phil Bell",
        "role": "Chief Revenue Officer",
        "li": "https://www.linkedin.com/in/phil-bell-57547018",
    },
    "ASBIS": {
        "person": "Yuri Ulasovich",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/yuri-ulasovich-0658494",
    },
    "ScanSource": {
        "person": "Mark David Jones",
        "role": "Vice President of Sales",
        "li": "https://www.linkedin.com/in/mark-david-jones-9a883716",
    },
    "Softchoice": {
        "person": "Jeff Knowlton",
        "role": "Senior Vice President of Sales",
        "li": "https://www.linkedin.com/in/knowltonjeff",
    },
    "Esprinet": {
        "person": "Sergio Grassi",
        "role": "Chief Revenue Officer, Value Green and Services",
        "li": "https://www.linkedin.com/in/sergiograssi",
    },
    "TD SYNNEX": {
        "person": "Phil Brown",
        "role": "VP Sales UK and Ireland",
        "li": "https://www.linkedin.com/in/phil-brown-16657714",
    },
    "SHI International": {
        "person": "Brian McGrath",
        "role": "SVP, US Commercial UK and Europe",
        "li": "https://www.linkedin.com/in/bcmcgrath",
    },
    "Insight Enterprises": {
        "person": "Bob Kane",
        "role": "North America Chief Revenue Officer",
        "li": "https://www.linkedin.com/in/robertfkane",
    },
    "Computacenter": {
        "person": "Lieven Bergmans",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/lieven-bergmans",
    },
    "Zones": {
        "person": "Dominic Camden",
        "role": "SVP, Alliance Sales",
        "li": "https://www.linkedin.com/in/dominiccamden",
    },
    "Intcomex": {
        "person": "Joseph Bouhadana",
        "role": "VP, Value and Cloud Channel Sales",
        "li": "https://www.linkedin.com/in/josephbouhadana",
    },
    "Connection": {
        "person": "Rick Gilligan",
        "role": "President of Sales",
        "li": "https://www.linkedin.com/in/rick-gilligan-723872",
    },
    "ALSO Holding": {
        "person": "Holger Dölle",
        "role": "Head of Sales, ALSO Deutschland",
        "li": "https://www.linkedin.com/in/holger-dölle-b106731a2",
    },
    "Petra Industries": {
        "person": "Jay Wehba",
        "role": "Vice President of National Accounts",
        "li": "https://www.linkedin.com/in/jay-wehba-8a556526",
    },
    "Almo Corporation": {
        "person": "John Riley",
        "role": "SVP of Sales East",
        "li": "https://www.linkedin.com/in/johnriley111",
    },
    "Softcat": {
        "person": "Anthony Cowen",
        "role": "Chief Revenue Officer",
        "li": "https://www.linkedin.com/in/anthony-cowen-9b74bb12",
    },
    "ELKO Group": {
        "person": "Martins Ozolins",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/martinsozolins",
    },
    "Savex Computers": {
        "person": "Nitin Ganjoo",
        "role": "Sales Head",
        "li": "https://www.linkedin.com/in/nitin-ganjoo-b551a814",
    },
    "AB SA": {
        "person": "Patrycja Gawarecka",
        "role": "Sales and Marketing Director",
        "li": "https://www.linkedin.com/in/patrycja-gawarecka",
    },
    "Cancom": {
        "person": "Marc Schmid",
        "role": "Director Sales, Direct",
        "li": "https://www.linkedin.com/in/marc-schmid-513a371b",
    },
    "Komplett Group": {
        "person": "Trygve Hillesland",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/trygve-hillesland",
    },
    "Komputronik": {
        "person": "Marcin Mordak",
        "role": "Omnichannel Sales Director",
        "li": "https://www.linkedin.com/in/marcin-mordak-7a68b782",
    },
    "PB Tech": {
        "person": "Dave Moynahan",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/dave-moynahan-2b84b26",
    },
    "Verkkokauppa.com": {
        "person": "Pekka Litmanen",
        "role": "Chief Experience and Sales Officer",
        "li": "https://www.linkedin.com/in/pekkalitmanen",
    },
    "Fnac Darty": {
        "person": "Florence Lemetais",
        "role": "Chief Commercial and Marketing Officer",
        "li": "https://www.linkedin.com/in/florence-lemetais-lajouanie",
    },
    "Alza.cz": {
        "person": "Juraj Matúšek",
        "role": "Commercial Division Director",
        "li": "https://www.linkedin.com/in/jurajmatusek",
    },
    "eMAG": {
        "person": "Stejara Pircan",
        "role": "Senior VP Commercial",
        "li": "https://www.linkedin.com/in/stejara-pircan-67132aa",
    },
    "MediaExpert": {
        "person": "Rafał Tomaszewski",
        "role": "Purchasing Director",
        "li": "https://www.linkedin.com/in/rafałtomaszewski",
    },
}
ADC = {
    "Euroclone S.p.A.": {
        "person": "Franco Ranzani",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/francoranzani",
    },
    "Mawdsleys": {
        "person": "Richard Day",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/richard-day-9131433",
    },
    "Mawdsley-Brooks": {
        "person": "Richard Day",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/richard-day-9131433",
    },
    "Nordic Biosite AB": {
        "person": "Aaron Venables",
        "role": "Chief Commercial Officer, Europa Biosite",
        "li": "https://www.linkedin.com/in/aaronvenables",
    },
    "Nordic BioSite ApS": {
        "person": "Aaron Venables",
        "role": "Chief Commercial Officer, Europa Biosite",
        "li": "https://www.linkedin.com/in/aaronvenables",
    },
    "Nordic Biosite AS": {
        "person": "Aaron Venables",
        "role": "Chief Commercial Officer, Europa Biosite",
        "li": "https://www.linkedin.com/in/aaronvenables",
    },
    "Nordic BioSite Oy": {
        "person": "Aaron Venables",
        "role": "Chief Commercial Officer, Europa Biosite",
        "li": "https://www.linkedin.com/in/aaronvenables",
    },
    "Labclinics": {
        "person": "Montse Gallardo",
        "role": "Director of Sales",
        "li": "https://www.linkedin.com/in/montse-gallardo-b8714a50",
    },
    "LubioScience GmbH": {
        "person": "Vladimir Zlateski",
        "role": "Head of Sales",
        "li": "https://www.linkedin.com/in/zlateskivladimir",
    },
    "SteinCares": {
        "person": "Carlos Abarca Calvo",
        "role": "Head of Procurement",
        "li": "https://www.linkedin.com/in/carlos-abarca-calvo-09051157",
    },
    "Difare Ecuador": {
        "person": "Daniel Moya",
        "role": "National Commercial Head",
        "li": "https://www.linkedin.com/in/daniel-moya-b2b12332",
    },
    "Assay Matrix Pty Ltd": {
        "person": "Sandhya Nambiar",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/sandhya-nambiar-1b0a7916",
    },
    "Drogueria del Sud Argentina": {
        "person": "Eric Perez Ochoteco",
        "role": "Director Comercial",
        "li": "https://www.linkedin.com/in/eric-perez-ochoteco-34710325",
    },
    "Herba Chemosan": {
        "person": "Gerhard Feytl",
        "role": "Verkaufsdirektor",
        "li": "https://www.linkedin.com/in/gerhard-feytl-14801b8b",
    },
    "Ibn Sina Pharma Egypt": {
        "person": "Mohamed Abbas",
        "role": "Head of Sales",
        "li": "https://www.linkedin.com/in/mohamed-abbas-898aab16b",
    },
    "Medical Supply Company Ltd": {
        "person": "Lisa Tang",
        "role": "Head of Sales and Marketing",
        "li": "https://www.linkedin.com/in/lisa-tang-4039185b",
    },
    "Gene X-Press": {
        "person": "Gonzalo Rojas",
        "role": "Head of Sales",
        "li": "https://www.linkedin.com/in/gonzalo-rojas-84493829",
    },
    "Izinta Trading Co.": {
        "person": "Szilvia Pipei",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/pipei-szilvia-740a40167",
    },
    "Farmasierra Distribución": {
        "person": "Carlos Lopez Lallana",
        "role": "Director Comercial",
        "li": "https://www.linkedin.com/in/carlos-lopez-lallana-8b250024",
    },
    "Farmaenlace Ecuador": {
        "person": "Miguel Angel Perez",
        "role": "Gerente Comercial",
        "li": "https://www.linkedin.com/in/miguel-angel-perez-883b6462",
    },
    "Farmaenlace Colombia": {
        "person": "Miguel Angel Perez",
        "role": "Gerente Comercial",
        "li": "https://www.linkedin.com/in/miguel-angel-perez-883b6462",
    },
    "Max Jenne Arzneimittel-Großhandlung": {
        "person": "Manuel Wolf",
        "role": "Vertriebsleiter",
        "li": "https://www.linkedin.com/in/manuel-wolf-6a9b60215",
    },
    "Drogueria Betances": {
        "person": "Carmen Millie Jusino",
        "role": "VP Sales and Marketing",
        "li": "https://www.linkedin.com/in/carmen-millie-jusino-53a86a15",
    },
    "Iwai Chemicals Company": {
        "person": "Hideki Maruyama",
        "role": "Managing Director, Strategic Sales and Marketing",
        "li": "https://www.linkedin.com/in/hideki-maruyama-a2450823",
    },
    "Shoppers Drug Mart Specialty Health Network": {
        "person": "Matt Simioni",
        "role": "Senior Director, Strategic Pharmaceutical Partnerships and Business Development",
        "li": "https://www.linkedin.com/in/matthew-simioni",
    },
}


def _headers(ws):
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            return r, {h: i + 1 for i, h in enumerate(vals) if h}
    raise RuntimeError("headers not found")


def _empty(v):
    return (str(v or "").strip().lower() in EMPTY) or (v is None)


def save_atomic(wb, path: Path):
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def apply(path: Path, contacts: dict) -> int:
    if not contacts:
        print("NO CONTACTS", path.name)
        return 0
    wb = load_workbook(path)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    hr, col = _headers(ws)
    c_co = col["Company"]
    c_cp = col["Contact Person"]
    c_role = col["Role"]
    c_li = col["LinkedIn"]
    n = 0
    unmatched = set(contacts)
    for row in range(hr + 1, ws.max_row + 1):
        co = str(ws.cell(row, c_co).value or "").strip()
        if co not in contacts:
            continue
        unmatched.discard(co)
        if not _empty(ws.cell(row, c_cp).value):
            print("SKIP already filled", co)
            continue
        rec = contacts[co]
        ws.cell(row, c_cp).value = rec["person"]
        ws.cell(row, c_role).value = rec["role"]
        ws.cell(row, c_li).value = rec["li"]
        n += 1
        print(f"FILL {co} | {rec['person']} | {rec['role']}")
    save_atomic(wb, path)
    wb.close()
    print("filled", n, path.name)
    if unmatched:
        print("UNMATCHED", unmatched)
    return n


def main():
    apply(SW_PATH, SW)
    apply(ADC_PATH, ADC)


if __name__ == "__main__":
    main()
