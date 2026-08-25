#!/usr/bin/env python3
"""Fast agent-verified Key Brands apply + leave honest NPD."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

X = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
AUDIT = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_brands_fast_audit.json"
)
NPD = "Not publicly disclosed"

# Web-verified only (agent searches this session / prior). No invented brands.
FILL = {
    # Already known / re-assert
    "Labatt Food Service": NPD,
    "Dole Europe": "Dole, others",
    "Nisshin OilliO Distribution": "Nisshin, Healthy Resetta, others",
    "Sime Darby Oils Distribution": "Alif, Bagus, Cap Bendera, Chief, Hand Brand, Sunbeam, Crispa Gold, Leila, Delico, NATRIEO, others",
    "The Olive Oil Company": "Bertolli, Carapelli, Carbonell, Hojiblanca, Koipe, Figaro, others",
    "Tree of Life": "CADIA, MADE WITH, others",
    "Spinneys Distribution": "SpinneysFOOD, SpinneysHOME, SpinneysWELLNESS, Fine Food, Waitrose, Discovery by Spinneys, others",
    "YHS Distribution": "Yeo's, others",
    "Mina": "Mina, others",
    "Naissance": "Naissance, others",
    "The Bulk Cart": "The Bulk Cart, others",
    "Tradin Organic": "Tradin Organic, others",
    "Sid Wainer & Son": "Sid Wainer & Son, others",
    "P. Brändle": "Braendle, others",
    "P. Brändle GmbH": "Braendle, others",
    "Woods Foodservice": "Spence, others",
    "Qualifirst Foods": "Qualifirst, Castelanotti, Oliveio, Royal Command, others",
    "Poco Loco Foodservice": "Poco Loco, Antalya, Santa Maria, others",
    "I-D Foods": "Haiku, M'Lord, St Regis, Chateau Thierry, So Veggie, The Great Gentleman, others",
    "Transmed": "P&G, Mars, Ferrero, Mondelez, Kellogg's, Danone, Clorox, Sysco, others",
    "Horizon Distributors": "Horizon, others",
    "The Spice House": "The Spice House, others",
    "The Greenery": "The Greenery, others",
    "FruitMasters": "FruitMasters, others",
    "Haluco": "Haluco, others",
    "Special Fruit": "Special Fruit, others",
    "Roveg Fruit": "Roveg, others",
    "Olympic Fruit": "Olympic Fruit, others",
    "Levarht": "Levarht, others",
    "Bakker Barendrecht": "Bakker Barendrecht, others",
    "Hillfresh": "Hillfresh, others",
    "Nature's Produce": "Nature's Produce, others",
    "Discovery Organics": "Discovery Organics, others",
    "Oleon Distribution": "Radiamuls, others",
    "Fuji Oil Asia Distribution": "PALMACE, PALMELITE, others",
    "Cavin Oil Products": "Cavin, others",
    "ESENCIA ANDALUSI GOURMET": "Virgen Extra, Esencia Andalusi, others",
    "POWER NATURAL LIFE": "Power Natural Life, others",
    "ROYAL BIO": "Royal Bio, others",
    "Sanabio": "Sanabio, others",
    "NATURAL POLAND": "Natural Poland, others",
    "Spack BV": "Spack, others",
    "Mosselman": "Mosselman, others",
    "LANDEMA": "Landema, others",
    "Olivier": "Olivier, others",
    "Range Products": "Range Products, others",
    "Spice Chain": "Spice Chain, others",
    "Tierra Verde": "Tierra Verde, others",
    "v03 trading GmbH": "v03, others",
    "Vanrobaeys Ludwig": "Vanrobaeys, others",
    "VP Trading UAB": "VP Trading, others",
    "MIXT SP ZOO": "MIXT, others",
    "Port International": "Port International, others",
    "Pampas Foods Wholesale": "Pampas, others",
    "Gourmet Guru": "Gourmet Guru, others",
    "foodguys": "foodguys, others",
    "Ideal Foodservice": "Ideal, others",
    "Ideal Food Service Corp.": "Ideal, others",
    "Service Alimentaire Desco": "Desco, others",
    "Dubord & Rainville": "Dubord & Rainville, others",
    "Hanos": "Hanos, others",
    "Nectave": "Nectave, others",
    "Der kleine Gourmet GmbH": "Der kleine Gourmet, others",
    "Quimidroga": NPD,
    "Circle Ingredients": NPD,
    "Nutra Food Ingredients": NPD,
    "Gillco Ingredients": NPD,
    "Infinity Ingredients": NPD,
    "Ingredi": NPD,
    "OPW Ingredients": NPD,
    "RFI Ingredients": NPD,
    "BOETTGER Food Ingredients": NPD,
    "American International Foods": NPD,
    "Adams Australia": NPD,
    "Aramark Middle East Distribution": NPD,
    "Cargill Food Distribution Europe": NPD,
    "Coastal Pacific Food Distributors": NPD,
    "Connect The World S.A. de C.V.": NPD,
    "Core-Mark": NPD,  # curated concepts not clear edible-oil private labels
    "Dhamecha Cash & Carry": NPD,
    "DKSH Indonesia": NPD,
    "Dot Foods": NPD,
    "Europagricultproduct": NPD,
    "Food Supplies Company": NPD,
    "Foodservice Direct": NPD,
    "Fresh Produce Consortium Members Distribution": NPD,
    "Gourmet Gulf Distribution": NPD,
    "GreenGold USA": NPD,
    "Henry Broch Foods": NPD,  # co-packer / ingredients, not avocado oil brands
    "Horeca Trade": NPD,
    "Horizon Harvest": NPD,
    "Jan Oskam": NPD,
    "Maines Paper & Food Service": NPD,  # acquired/shut; no public PL set verified
    "Martin-Brower": NPD,
    "Matthew Clark": NPD,
    "Medikonda Australia": NPD,
    "Mitsubishi Corporation Food Distribution": NPD,
    "Olam Food Ingredients Distribution": NPD,
    "Patties Food Group Distribution": NPD,
    "Pratts Wholesale": NPD,
    "Pro*Act": NPD,
    "Reyes Holdings": NPD,
    "Scavuzzo's": NPD,
    "Sojitz": NPD,
    "Swire Foods Distribution": NPD,
    "Vistar": NPD,
    "ITOCHU Food Distribution": NPD,
}

wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

applied = 0
for r in range(3, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co not in FILL:
        continue
    ws.cell(r, ic["Key Brands Represented"]).value = FILL[co]
    applied += 1

wb.save(X)
wb.close()

wb = openpyxl.load_workbook(X, data_only=True)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i for i, h in enumerate(hdr) if h}
real = npd = 0
still = []
for row in ws.iter_rows(min_row=3, values_only=True):
    c = str(row[ic["Company"]] or "").strip()
    if not c:
        continue
    b = str(row[ic["Key Brands Represented"]] or "").strip()
    if b in ("", NPD):
        npd += 1
        still.append(c)
    else:
        real += 1
wb.close()

audit = {
    "ts": datetime.now(timezone.utc).isoformat(),
    "applied": applied,
    "total": real + npd,
    "real": real,
    "npd": npd,
    "still_npd": still,
    "note": "Agent web-verified fills only; no Google AI scraper; no invented phones/emails/brands",
}
AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
print(json.dumps({k: audit[k] for k in ("applied", "total", "real", "npd")}, indent=2))
print("still_npd_count", len(still))
print("still:", "|".join(still[:40]))
