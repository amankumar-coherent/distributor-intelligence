# -*- coding: utf-8 -*-
"""Fill remaining flexpack Employees NPD via Google AI scraper + DeepSeek (evidence only)."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")

from vendor_intel.clients.google_ai_scraper import (  # noqa: E402
    google_ai_ask,
    google_ai_ping,
)

OUT = ROOT / "output/web_expand/global_flexible_packaging_market_global"
XLSX = OUT / "global_flexible_packaging_market_global_FINAL.xlsx"
AUDIT = OUT / "flexpack_employees_google_ai_audit.json"
SUMMARY = OUT / "flexpack_employees_google_ai_summary.json"

NPD = "Not publicly disclosed"
NPD_SET = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "null",
}

# Use en-dash to match existing FINAL sheet style
BANDS = [
    (1, 10, "1–10"),
    (11, 50, "11–50"),
    (51, 200, "51–200"),
    (201, 500, "201–500"),
    (501, 1_000, "501–1,000"),
    (1_001, 5_000, "1,001–5,000"),
    (5_001, 10_000, "5,001–10,000"),
    (10_001, 50_000, "10,001–50,000"),
    (50_001, None, "50,001+"),
]
VALID = {lab for _, _, lab in BANDS}

ALLOW_MEGA = {
    "bunzl",
    "grainger",
    "fastenal",
    "uline",
    "hd supply",
    "office depot",
    "officeworks",
    "lyreco",
    "askul",
    "msc",
    "veritiv",
    "imperial dade",
    "bradyplus",
    "ahlsell",
    "mcmaster",
    "global industrial",
    "raja",
    "antalis",
    "macfarlane",
}

# Optional extras verified in prior passes (safe)
HAND: dict[str, str] = {}

# Limit for one run (0 = all). Override with EMP_GAI_LIMIT.
LIMIT = int(os.getenv("EMP_GAI_LIMIT") or "0")


def emptyish(v: object) -> bool:
    return str(v or "").strip().lower() in NPD_SET


def number_to_range(n: int) -> str:
    if n <= 0:
        return ""
    for lo, hi, lab in BANDS:
        if hi is None:
            if n >= lo:
                return lab
        elif lo <= n <= hi:
            return lab
    return "50,001+"


def normalize_to_range(raw: str) -> str:
    if emptyish(raw):
        return ""
    s = str(raw).strip().replace("–", "-").replace("—", "-").replace("�", "-")
    for _, _, lab in BANDS:
        a = s.replace(",", "").replace(" ", "").lower()
        b = lab.replace("–", "-").replace(",", "").replace(" ", "").lower()
        if a == b:
            return lab
    m = re.search(r"(\d[\d,]*)\s*[-–—]\s*(\d[\d,]*)", s)
    if m:
        a = int(m.group(1).replace(",", ""))
        b = int(m.group(2).replace(",", ""))
        return number_to_range((a + b) // 2)
    m = re.search(r"(\d+(?:\.\d+)?)\s*[kK]\b", s)
    if m:
        return number_to_range(int(float(m.group(1)) * 1000))
    m = re.search(r"\b(\d{1,3}(?:,\d{3})+|\d{2,7})\+?\b", s)
    if m:
        return number_to_range(int(m.group(1).replace(",", "")))
    return ""


def company_tokens(name: str) -> list[str]:
    stop = {
        "inc",
        "ltd",
        "llc",
        "co",
        "corp",
        "company",
        "packaging",
        "pack",
        "supply",
        "supplies",
        "the",
        "and",
        "of",
        "group",
        "gmbh",
        "pty",
    }
    toks = re.findall(r"[A-Za-z0-9]+", name.lower())
    return [t for t in toks if len(t) >= 3 and t not in stop][:4]


def _absolute_headcounts(evidence: str) -> list[int]:
    nums: list[int] = []
    pats = [
        r"(?:approximately|about|around|nearly|over|more than|~)\s*([\d,]{1,7})\+?\s*"
        r"(?:employees?|staff|associates|people|team members|workers)\b",
        r"(?:employees?|headcount|associates|staff|workforce|team members)"
        r"[^\d]{0,24}(?:of|:)?\s*(?:approximately|about|around|nearly|over|~)?\s*([\d,]{1,7})\+?\b",
        r"\b([\d,]{1,7})\+?\s*(?:employees?|staff|associates)\b",
        r"employs\s+(?:approximately|about|around|nearly|over|~)?\s*([\d,]{1,7})\+?\b",
        r"workforce\s+of\s+(?:approximately|about|around|nearly|over|~)?\s*([\d,]{1,7})\+?\b",
        r"\b([\d,]{2,7})\s+(?:full[- ]time\s+)?employees?\b",
        r"company size[:\s]+([\d,]{1,7})\s*[-–]\s*([\d,]{1,7})",
    ]
    for pat in pats:
        for m in re.finditer(pat, evidence, re.I):
            if m.lastindex and m.lastindex >= 2 and m.group(2):
                try:
                    lo = int(m.group(1).replace(",", ""))
                    hi = int(m.group(2).replace(",", ""))
                    if 1 <= lo <= hi <= 5_000_000:
                        nums.append((lo + hi) // 2)
                    continue
                except ValueError:
                    pass
            n = int(m.group(1).replace(",", ""))
            if 1 <= n <= 5_000_000 and n not in range(1900, 2031):
                nums.append(n)
    return nums


def evidence_supports(company: str, rng: str, evidence: str) -> bool:
    if not rng or not evidence or rng not in VALID:
        return False
    ev = evidence.lower()
    if not re.search(
        r"employee|headcount|associates|staff|workforce|team members|people|employs|company size",
        ev,
    ):
        return False
    toks = company_tokens(company)
    if toks and not any(t in ev for t in toks):
        return False
    abs_nums = _absolute_headcounts(evidence)
    if not abs_nums:
        return False
    if not any(number_to_range(n) == rng for n in abs_nums):
        return False
    if rng in {"10,001–50,000", "50,001+"}:
        if not any(g in company.lower() for g in ALLOW_MEGA):
            return False
    if rng == "5,001–10,000" and not any(g in company.lower() for g in ALLOW_MEGA):
        if not any(n >= 5001 for n in abs_nums):
            return False
    return True


def extract_regex_range(evidence: str) -> str:
    nums = _absolute_headcounts(evidence)
    if not nums:
        return ""
    return number_to_range(nums[0])


def ds_client() -> tuple[OpenAI | None, str]:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    if not key:
        return None, model
    return OpenAI(api_key=key, base_url=base, timeout=60.0), model


def parse_json(text: str) -> dict:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    try:
        return json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


def deepseek_from_evidence(client: OpenAI, model: str, company: str, evidence: str) -> dict:
    labs = " | ".join(sorted(VALID))
    system = (
        "Extract THIS company's employee headcount from EVIDENCE ONLY. Never invent.\n"
        "Ignore parent conglomerate / wrong namesake headcounts.\n"
        "JSON: {found:bool, employee_count:int|null, range:string, confidence:high|medium|low|none, reason:string}\n"
        f"range must be one of: {labs} or empty.\n"
        "If no clear number/band for this company: found=false."
    )
    user = f"Company: {company}\n\nEVIDENCE:\n{(evidence or '(none)')[:7500]}"
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=220,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        return parse_json(resp.choices[0].message.content or "{}")
    except Exception as e:  # noqa: BLE001
        return {"found": False, "range": "", "confidence": "none", "reason": str(e)[:120]}


async def ask_employees(company: str, website: str, hq: str) -> str:
    geo = f" ({hq})" if hq else ""
    host = ""
    if website:
        host = re.sub(r"^https?://", "", website.strip(), flags=re.I).split("/")[0]
    site = f" {host}" if host else ""
    queries = [
        f"{company}{geo} number of employees headcount LinkedIn company size",
        f"How many employees does {company} have?{site}",
        f"{company}{site} employees staff headcount",
    ]
    chunks: list[str] = []
    for q in queries:
        data = await google_ai_ask(q, llm_clean=True)
        md = str(data.get("markdown") or "").strip()
        err = str(data.get("error") or "")
        if md and len(md) >= 40:
            chunks.append(f"Q: {q}\n{md}")
        elif err:
            chunks.append(f"Q: {q}\n[error:{err}]")
        joined = "\n\n".join(chunks)
        if extract_regex_range(joined):
            break
        await asyncio.sleep(0.7)
    return "\n\n----\n\n".join(chunks)[:12000]


async def main_async() -> None:
    client, model = ds_client()
    ok = await google_ai_ping()
    print(f"GoogleAI ping={ok} DeepSeek={'yes' if client else 'no'} model={model}", flush=True)
    if not ok:
        print("ERROR: Google AI scraper not reachable on :15551", flush=True)
        return

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_EMPLOYEES_GAI_{stamp}.xlsx")
    wb0 = openpyxl.load_workbook(XLSX)
    wb0.save(bak)
    wb0.close()
    print("backup", bak.name, flush=True)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    gaps: list[int] = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        cur = str(ws.cell(r, ic["Employees"]).value or "").strip()
        if co and emptyish(cur):
            gaps.append(r)

    if LIMIT > 0:
        gaps = gaps[:LIMIT]
    print(f"npd gaps={len(gaps)} (LIMIT={LIMIT or 'all'})", flush=True)

    hand_n = gai_n = ds_n = left = 0
    audit: list[dict] = []

    for i, r in enumerate(gaps, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic.get("Website", 0)).value or "").strip() if "Website" in ic else ""
        hq = ""
        for key in ("Headquarters", "Continent / Geography"):
            if key in ic:
                hq = str(ws.cell(r, ic[key]).value or "").strip()
                if hq and hq.lower() not in NPD_SET:
                    break

        if co in HAND:
            ws.cell(r, ic["Employees"]).value = HAND[co]
            hand_n += 1
            audit.append({"company": co, "employees": HAND[co], "source": "hand"})
            print(f"[{i}/{len(gaps)}] HAND {co} => {HAND[co]}", flush=True)
            continue

        print(f"[{i}/{len(gaps)}] GoogleAI {co} ...", flush=True)
        evidence = await ask_employees(co, web, hq)
        rng = ""
        source = "still_npd"

        cand = extract_regex_range(evidence)
        if cand and evidence_supports(co, cand, evidence):
            rng = cand
            source = "google_ai_regex"
            gai_n += 1
        elif client and evidence.strip():
            data = deepseek_from_evidence(client, model, co, evidence)
            conf = str(data.get("confidence") or "").lower()
            raw = normalize_to_range(str(data.get("range") or ""))
            if not raw and data.get("employee_count"):
                try:
                    raw = number_to_range(int(data["employee_count"]))
                except Exception:
                    raw = ""
            if data.get("found") and conf in {"high", "medium"} and raw in VALID:
                if evidence_supports(co, raw, evidence):
                    rng = raw
                    source = "google_ai_deepseek"
                    ds_n += 1
            time.sleep(0.2)

        if rng:
            ws.cell(r, ic["Employees"]).value = rng
            audit.append(
                {
                    "company": co,
                    "employees": rng,
                    "source": source,
                    "evidence_chars": len(evidence),
                    "evidence_preview": evidence[:500],
                }
            )
            print(f"  FILL {co} => {rng} [{source}]", flush=True)
        else:
            left += 1
            ws.cell(r, ic["Employees"]).value = NPD
            audit.append(
                {
                    "company": co,
                    "employees": NPD,
                    "source": "still_npd",
                    "evidence_chars": len(evidence),
                    "evidence_preview": evidence[:400],
                }
            )
            print(f"  NPD {co}", flush=True)

        if i % 5 == 0:
            wb.save(XLSX)
            AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
            print(f"  checkpoint saved ({i})", flush=True)

    wb.save(XLSX)
    wb.close()

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ei = headers.index("Employees")
    filled = npd = 0
    ctr: Counter[str] = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = str(row[ei] or "").strip()
        ctr[v] += 1
        if emptyish(v):
            npd += 1
        else:
            filled += 1
    wb.close()

    stats = {
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "backup": str(bak),
        "gaps_at_start": len(gaps),
        "hand": hand_n,
        "google_ai_regex": gai_n,
        "google_ai_deepseek": ds_n,
        "still_npd": left,
        "employees_ranged": filled,
        "employees_npd": npd,
        "distribution": dict(ctr.most_common()),
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    SUMMARY.write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(stats, indent=2), flush=True)


def main() -> None:
    asyncio.run(main_async())


if __name__ == "__main__":
    main()
