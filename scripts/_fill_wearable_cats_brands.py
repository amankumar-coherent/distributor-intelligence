# -*- coding: utf-8 -*-
"""Verify/fill Core Categories, Specialty Focus, Key Brands Represented on wearable FINAL.

Definitions (user):
1) Core Categories — main product categories distributed
2) Specialty Focus — market/tech/segment expertise
3) Key Brands Represented — manufacturer brands officially distributed/represented

Uses website scrape + DeepSeek. Never invent brands — NPD if not evidenced.
"""
from __future__ import annotations

import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import httpx
import openpyxl
from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

OUT = ROOT / "output/web_expand/global_wearable_medical_devices_market_global"
XLSX = OUT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
AUDIT = OUT / "wearable_cats_brands_audit.json"
SUMMARY = OUT / "wearable_cats_brands_summary.json"

NPD = "Not publicly disclosed"
NPD_SET = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
}

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124.0.0.0 Safari/537.36"

# Agent web-verified for previously empty Core/Specialty (+ brands when known)
HAND: dict[str, dict[str, str]] = {
    "Africure Pharmaceuticals": {
        "Core Categories": "Generic Pharmaceuticals; Antiretrovirals; Antibiotics; Antimalarials; Essential Medicines",
        "Specialty Focus": "Manufactures and distributes essential generic medicines across Sub-Saharan Africa, including ARVs, antibiotics, and antimalarials from regional plants.",
        "Key Brands Represented": "Africure, others",
    },
    "Bidvest Medical": {
        "Core Categories": "Pharmaceuticals; OTC Medicines; Critical Care Products; Consumer Healthcare",
        "Specialty Focus": "Bidvest Branded Products pharmaceutical and healthcare distribution in South Africa spanning Rx, OTC, and consumer health.",
        "Key Brands Represented": "Panado, Compral, Epimax, Bioplus, others",
    },
    "Daehwa Pharmaceutical": {
        "Core Categories": "Prescription Pharmaceuticals; OTC Pharmaceuticals; Health Functional Foods",
        "Specialty Focus": "Korean pharmaceutical company focused on cardiovascular, hepatic, CNS, and specialty medicines plus health functional foods.",
        "Key Brands Represented": "Daehwa, D.H. Carbamazepine, Atorva, Topstatin, Losartin, others",
    },
    "MediSystem Pharmacy": {
        "Core Categories": "Long-Term Care Pharmacy; Medication Packaging; Clinical Pharmacy Services; Medication Delivery",
        "Specialty Focus": "Specialty long-term care pharmacy (Shoppers Drug Mart) providing dispensing, packaging, and clinical consulting to LTC, retirement, and assisted-living facilities across Canada.",
        "Key Brands Represented": NPD,  # dispenses multi-brand Rx; principals not listed as agency brands
    },
    "Enerspect Medical Solutions": {
        "Core Categories": "AEDs; ALS Monitor/Defibrillators; Emergency Cardiac Equipment; First Response Supplies",
        "Specialty Focus": "Distributes and supports AEDs and ALS defibrillator/monitors for workplaces, EMS, and emergency cardiac response.",
        "Key Brands Represented": "ZOLL, Philips, Defibtech, HeartSine, LIFEPAK, others",
    },
    "Stevens Company": {
        "Core Categories": "Surgical Supplies; Hospital Furniture; Operating Room Equipment; Clinic Supplies; Medical Consumables",
        "Specialty Focus": "Canada’s largest family-owned full-line medical supply and equipment distributor serving hospitals, physicians, and long-term care.",
        "Key Brands Represented": NPD,  # 500+ manufacturing partners; specific list not public on About
    },
    "Globalmed Malaysia": {
        "Core Categories": "Medical Devices; Critical Care Equipment; Respiratory & Sleep Therapy; Patient Monitoring; Hospital Equipment",
        "Specialty Focus": "Malaysian medical device distributor focused on hospital projects, tenders, and critical-care / respiratory equipment supply.",
        "Key Brands Represented": NPD,  # globalmed.com.my site not usable; do not invent principals
    },
    "Primary Care Supplies": {
        "Core Categories": "Primary Care Medical Equipment; Diagnostic Devices; Clinical Consumables; Pharmacy Fridges; Clinic Furniture",
        "Specialty Focus": "UK distributor of medical equipment and supplies for primary care / frontline NHS and clinic settings for 20+ years.",
        "Key Brands Represented": "seca, Cardioline, ndd, Labcold, Marsden, Clinell, CoaguChek, NObreath, others",
    },
}


def emptyish(v: object) -> bool:
    return str(v or "").strip().lower() in NPD_SET or str(v or "").strip().lower().startswith("verify:")


def weak_brands(v: object) -> bool:
    s = str(v or "").strip()
    if emptyish(s):
        return True
    low = s.lower()
    if low.startswith("multi-vendor") or low.startswith("multi-brand pharmacy") or low.startswith("multi-principal"):
        parts = [p.strip() for p in re.split(r"[,;]", s) if p.strip()]
        real = [
            p
            for p in parts
            if p.lower() not in {"others", "etc", "etc."}
            and "multi-" not in p.lower()
            and "assortment" not in p.lower()
            and "principals" not in p.lower()
            and "subsidiary" not in p.lower()
            and "cardinal health medical products" not in p.lower()
        ]
        # Cardinal-style with only weak + others
        if len(real) <= 1 and any("cardinal health medical" in p.lower() for p in parts):
            return True
        return len(real) == 0
    return False


def strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def fetch(url: str) -> str:
    try:
        with httpx.Client(
            headers={"User-Agent": UA},
            verify=False,
            timeout=httpx.Timeout(16.0, connect=5.0),
            follow_redirects=True,
        ) as c:
            r = c.get(url)
            if r.status_code >= 400:
                return ""
            return strip_html(r.text or "")[:7000]
    except Exception:
        return ""


def gather_evidence(company: str, website: str) -> str:
    chunks: list[str] = [f"Company: {company}", f"Website: {website}"]
    base = (website or "").strip()
    if base and not re.match(r"^https?://", base, re.I):
        base = "https://" + base
    if base:
        paths = [
            "",
            "/about",
            "/about-us",
            "/company",
            "/products",
            "/brands",
            "/partners",
            "/our-brands",
            "/en/about",
            "/en/products",
            "/portfolio",
            "/solutions",
        ]
        for path in paths:
            u = base if not path else urljoin(base.rstrip("/") + "/", path.lstrip("/"))
            t = fetch(u)
            if len(t) > 250:
                chunks.append(f"URL:{u}\n{t[:2800]}")
            if len(chunks) >= 5:
                break
            time.sleep(0.15)
    return "\n\n---\n\n".join(chunks)[:12000]


def ds_client() -> tuple[OpenAI, str]:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    if not key:
        raise SystemExit("OPENAI_API_KEY missing")
    return OpenAI(api_key=key, base_url=base, timeout=90.0), model


def parse_json(text: str) -> dict:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    try:
        return json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


def ensure_others(brands: str) -> str:
    t = (brands or "").strip().rstrip(",")
    if not t or emptyish(t):
        return NPD
    if re.search(r",\s*others\s*$", t, re.I):
        return re.sub(r",\s*others\s*$", ", others", t, flags=re.I)
    return t + ", others"


def clean_brands(raw: str) -> str:
    if emptyish(raw) or str(raw).strip().lower() in {"none", "null", "unknown"}:
        return NPD
    parts = re.split(r"[,;/|]+", str(raw))
    out: list[str] = []
    seen: set[str] = set()
    junk = {
        "others",
        "etc",
        "various",
        "multiple",
        "n/a",
        "none",
        "unknown",
        "brands",
        "brand",
        "principals",
        "suppliers",
        "portfolio",
        "multi-vendor",
        "multi-brand",
        "assortment",
    }
    for p in parts:
        b = re.sub(r"\s+", " ", p).strip(" .")
        if not b or len(b) < 2 or b.lower() in junk:
            continue
        if re.search(r"multi-vendor|multi-brand|assortment|not publicly|principals|private label", b, re.I):
            continue
        if len(b) > 55:
            continue
        key = b.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(b)
    if not out:
        return NPD
    return ensure_others(", ".join(out[:14]))


def clean_core(raw: str) -> str:
    s = re.sub(r"\s+", " ", str(raw or "")).strip(" ;,")
    if emptyish(s) or len(s) < 3:
        return ""
    # semicolon-separated categories
    parts = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"[;|]", s) if p.strip()]
    parts = [p for p in parts if len(p) >= 2][:10]
    return "; ".join(parts) if parts else ""


def clean_specialty(raw: str) -> str:
    s = re.sub(r"\s+", " ", str(raw or "")).strip()
    if emptyish(s) or len(s) < 20:
        return ""
    return s[:420]


def deepseek_extract(client: OpenAI, model: str, company: str, website: str, hq: str, evidence: str) -> dict:
    system = (
        "You extract distributor landscape fields from EVIDENCE ONLY. Never invent.\n"
        "Return JSON:\n"
        "{\n"
        '  "core_categories": "semicolon-separated main product categories distributed, or empty",\n'
        '  "specialty_focus": "1-2 sentences on what the distributor is particularly known for, or empty",\n'
        '  "key_brands": "comma-separated manufacturer/brand names they distribute/represent, or empty",\n'
        '  "confidence": "high|medium|low|none"\n'
        "}\n"
        "Rules:\n"
        "- Core Categories = types of products sold/distributed (not a company slogan).\n"
        "- Specialty Focus = specific market/tech/customer focus expertise.\n"
        "- Key Brands = only named manufacturer brands evidenced as distributed/represented/partnered.\n"
        "- If brands are not named in evidence, key_brands must be empty (do not write multi-vendor).\n"
        "- Do not invent emails, phones, or brand names.\n"
    )
    user = (
        f"Company: {company}\nWebsite: {website}\nHeadquarters: {hq}\n\n"
        f"EVIDENCE:\n{(evidence or '(none)')[:9000]}"
    )
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=450,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        return parse_json(resp.choices[0].message.content or "{}")
    except Exception as e:  # noqa: BLE001
        return {"confidence": "none", "error": str(e)[:160]}


def brands_supported(brands: str, evidence: str) -> str:
    """Keep only brand tokens that appear in evidence (case-insensitive)."""
    if emptyish(brands) or brands == NPD:
        return NPD
    ev = (evidence or "").lower()
    kept: list[str] = []
    for b in [x.strip() for x in brands.replace(";", ",").split(",")]:
        if not b or b.lower() == "others":
            continue
        # require brand (or first significant token) in evidence
        token = re.sub(r"[^A-Za-z0-9&.\\- ]", "", b).strip()
        if len(token) < 2:
            continue
        # allow slight spacing variants
        needle = token.lower()
        if needle in ev or needle.replace(" ", "") in ev.replace(" ", ""):
            kept.append(b)
            continue
        # first word >= 4 chars
        first = needle.split()[0] if needle.split() else ""
        if len(first) >= 4 and first in ev:
            kept.append(b)
    if not kept:
        return NPD
    return ensure_others(", ".join(kept[:14]))


def main() -> None:
    client, model = ds_client()
    print(f"DeepSeek=yes model={model}", flush=True)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_CATS_BRANDS_{stamp}.xlsx")
    wb0 = openpyxl.load_workbook(XLSX)
    wb0.save(bak)
    wb0.close()
    print("backup", bak.name, flush=True)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    # determine worklist
    work: list[int] = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        core = ws.cell(r, ic["Core Categories"]).value
        spec = ws.cell(r, ic["Specialty Focus"]).value
        brands = ws.cell(r, ic["Key Brands Represented"]).value
        need_core = emptyish(core)
        need_spec = emptyish(spec)
        need_brands = weak_brands(brands)
        if co in HAND or need_core or need_spec or need_brands:
            work.append(r)
    print(f"rows to process={len(work)}", flush=True)

    # apply HAND first
    hand_n = 0
    audit: list[dict] = []
    for r in work:
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co not in HAND:
            continue
        hmap = HAND[co]
        for col, val in hmap.items():
            if col in ic and val:
                cur = ws.cell(r, ic[col]).value
                if emptyish(cur) or (col == "Key Brands Represented" and weak_brands(cur)) or True:
                    # always apply hand overrides for these keys
                    ws.cell(r, ic[col]).value = val
        hand_n += 1
        audit.append({"company": co, "source": "hand", **hmap})
        print("HAND", co, flush=True)

    # gather evidence for remaining needing scrape
    todo = []
    for r in work:
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co in HAND:
            # still may want brands/core refresh? HAND already set — skip DS
            continue
        core = ws.cell(r, ic["Core Categories"]).value
        spec = ws.cell(r, ic["Specialty Focus"]).value
        brands = ws.cell(r, ic["Key Brands Represented"]).value
        if emptyish(core) or emptyish(spec) or weak_brands(brands):
            todo.append(r)

    print(f"gathering evidence for {len(todo)} ...", flush=True)
    evidence_map: dict[int, str] = {}

    def job(r: int) -> tuple[int, str]:
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        return r, gather_evidence(co, web)

    with ThreadPoolExecutor(max_workers=8) as ex:
        futs = [ex.submit(job, r) for r in todo]
        done = 0
        for fut in as_completed(futs):
            r, ev = fut.result()
            evidence_map[r] = ev
            done += 1
            if done % 25 == 0 or done == len(todo):
                print(f"  evidence {done}/{len(todo)}", flush=True)

    filled_core = filled_spec = filled_brands = npd_brands = 0
    for i, r in enumerate(todo, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        hq = str(ws.cell(r, ic.get("Headquarters", 0)).value or "").strip() if "Headquarters" in ic else ""
        evidence = evidence_map.get(r) or ""

        need_core = emptyish(ws.cell(r, ic["Core Categories"]).value)
        need_spec = emptyish(ws.cell(r, ic["Specialty Focus"]).value)
        need_brands = weak_brands(ws.cell(r, ic["Key Brands Represented"]).value)

        data = deepseek_extract(client, model, co, web, hq, evidence)
        conf = str(data.get("confidence") or "").lower()
        core = clean_core(str(data.get("core_categories") or ""))
        spec = clean_specialty(str(data.get("specialty_focus") or ""))
        brands = clean_brands(str(data.get("key_brands") or ""))
        if brands != NPD:
            brands = brands_supported(brands, evidence)

        row_audit: dict = {
            "company": co,
            "source": "deepseek_web",
            "confidence": conf,
            "evidence_chars": len(evidence),
        }

        if need_core and core and conf in {"high", "medium", "low"}:
            ws.cell(r, ic["Core Categories"]).value = core
            filled_core += 1
            row_audit["Core Categories"] = core
        elif need_core:
            row_audit["Core Categories"] = "(still empty)"

        if need_spec and spec and conf in {"high", "medium", "low"}:
            ws.cell(r, ic["Specialty Focus"]).value = spec
            filled_spec += 1
            row_audit["Specialty Focus"] = spec
        elif need_spec:
            row_audit["Specialty Focus"] = "(still empty)"

        if need_brands:
            if brands != NPD and conf in {"high", "medium"}:
                ws.cell(r, ic["Key Brands Represented"]).value = brands
                filled_brands += 1
                row_audit["Key Brands Represented"] = brands
            else:
                ws.cell(r, ic["Key Brands Represented"]).value = NPD
                npd_brands += 1
                row_audit["Key Brands Represented"] = NPD

        audit.append(row_audit)
        print(
            f"[{i}/{len(todo)}] {co} core={'Y' if 'Core Categories' in row_audit and not str(row_audit.get('Core Categories')).startswith('(') else 'n'} "
            f"spec={'Y' if 'Specialty Focus' in row_audit and not str(row_audit.get('Specialty Focus','')).startswith('(') else 'n'} "
            f"brands={row_audit.get('Key Brands Represented','-')[:50] if need_brands else 'skip'}",
            flush=True,
        )
        time.sleep(0.2)
        if i % 40 == 0:
            wb.save(XLSX)
            print("  checkpoint", flush=True)

    wb.save(XLSX)
    wb.close()

    # recount
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic2 = {str(h): i + 1 for i, h in enumerate(headers) if h}

    def count_col(name: str) -> tuple[int, int, int]:
        filled = empty = weak = 0
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, ic2[name]).value
            if name == "Key Brands Represented":
                if emptyish(v):
                    empty += 1
                elif weak_brands(v):
                    weak += 1
                else:
                    filled += 1
            else:
                if emptyish(v):
                    empty += 1
                else:
                    filled += 1
        return filled, empty, weak

    stats = {
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "backup": str(bak),
        "hand": hand_n,
        "filled_core": filled_core,
        "filled_spec": filled_spec,
        "filled_brands": filled_brands,
        "brands_set_npd": npd_brands,
        "Core Categories": dict(zip(["filled", "empty", "weak"], count_col("Core Categories"))),
        "Specialty Focus": dict(zip(["filled", "empty", "weak"], count_col("Specialty Focus"))),
        "Key Brands Represented": dict(
            zip(["filled", "empty_or_npd", "weak_stub"], count_col("Key Brands Represented"))
        ),
    }
    # NPD counts as emptyish for brands
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    SUMMARY.write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(stats, indent=2), flush=True)


if __name__ == "__main__":
    main()
