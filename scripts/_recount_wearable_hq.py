# -*- coding: utf-8 -*-
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"


def has_city(hq: str) -> bool:
    s = (hq or "").strip()
    if not s or "," not in s:
        return False
    parts = [p.strip() for p in s.split(",") if p.strip()]
    if len(parts) < 2:
        return False
    city = parts[0]
    if len(city) < 2:
        return False
    low = s.lower()
    if low in {
        "singapore, singapore",
        "hong kong sar, china",
        "hong kong, china",
        "hong kong, hong kong",
    }:
        return True
    countries_only = {
        "united states",
        "usa",
        "united kingdom",
        "uk",
        "germany",
        "france",
        "japan",
        "china",
        "india",
        "brazil",
        "canada",
        "australia",
        "singapore",
        "malaysia",
        "south africa",
        "sweden",
        "spain",
        "italy",
        "netherlands",
        "belgium",
        "switzerland",
        "denmark",
        "finland",
        "norway",
        "mexico",
        "korea",
        "south korea",
        "uae",
        "united arab emirates",
        "hong kong",
        "new zealand",
        "ireland",
        "argentina",
        "kenya",
        "qatar",
        "kuwait",
        "oman",
        "saudi arabia",
        "europe",
        "iraq",
        "honduras",
        "el salvador",
        "bolivia",
        "puerto rico",
    }
    if city.lower() in countries_only:
        return False
    return True


def main() -> None:
    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers) if h}
    co_i, hq_i = ic["Company"], ic["Headquarters"]
    city = 0
    still = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, co_i).value or "").strip()
        hq = str(ws.cell(r, hq_i).value or "").strip()
        if has_city(hq):
            city += 1
        else:
            still.append(
                {
                    "company": co,
                    "hq": hq,
                    "website": str(ws.cell(r, ic["Website"]).value or "").strip(),
                }
            )
    audit = {
        "city_country_total": city,
        "country_only_remaining": len(still),
        "rows": ws.max_row - 1,
        "still_all": still,
    }
    (ROOT / "_wearable_hq_fill_audit.json").write_text(
        json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print("rows", ws.max_row - 1)
    print("city+country", city)
    print("country-only remaining", len(still))
    print("--- still ---")
    for s in still:
        print(s["hq"], "|", s["company"], "|", s["website"])


if __name__ == "__main__":
    main()
