#!/usr/bin/env python3
"""Fill Contact Person where empty/NPD using OpenAI Responses web_search (live web)."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

load_dotenv(ROOT / ".env")

from vendor_intel.pipeline.chatgpt_expand import _web_search_json  # noqa: E402

NPD = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd", "leadership"}
BAD = {"ceo", "cto", "inc", "ltd", "llc", "gmbh", "team", "leadership", "executive", "appoints"}

FILES = [
    ("smartwatch", ROOT / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL.xlsx", "Smartwatch Companies", "Global Smartwatch Market"),
    ("green", ROOT / "output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL.xlsx", "Green Chemical Companies", "Global Green Chemical Market"),
    ("firewall", ROOT / "output/run4/global_firewall_market_global/global_firewall_market_FINAL.xlsx", "Firewall Companies", "Global Firewall Market"),
    ("smartphone", ROOT / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx", "Companies", "Global Smartphone Market"),
    ("avocado", ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL.xlsx", "Avocado Oil Distributors", "Global Avocado Oil Market"),
]


def empty(v: str) -> bool:
    return str(v or "").strip().lower() in NPD


def looks_person(name: str, company: str = "") -> bool:
    n = (name or "").strip()
    if len(n) < 4 or len(n) > 55 or "\n" in n:
        return False
    parts = n.split()
    if not (2 <= len(parts) <= 4):
        return False
    if any(p.lower().strip(".,") in BAD for p in parts):
        return False
    if not all(re.match(r"^[A-Za-z][A-Za-z.'\-]*$", p) for p in parts):
        return False
    if company and name.strip().lower() == company.strip().lower():
        return False
    return True


def load(path: Path, sheet: str):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hi = None
    hdr = None
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            hi, hdr = i, vals
            break
    data = []
    for r in rows[hi + 1 :]:
        if not r or not r[hdr.index("Company")]:
            continue
        data.append({hdr[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip()) for j in range(len(hdr)) if hdr[j]})
    return hdr, data, hi + 1


def save(path: Path, sheet: str, headers, header_row, rows):
    wb = load_workbook(path)
    ws = wb[sheet]
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    font = Font(name="Calibri", size=9)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row.get(h, "") if h else "" for h in headers])
        r = ws.max_row
        for i in range(1, len(headers) + 1):
            c = ws.cell(r, i)
            c.font = font
            c.border = thin
            c.alignment = wrap
    wb.save(path)


async def fill_one(client, model, company, website, market_q):
    data = await asyncio.to_thread(
        _web_search_json,
        client,
        model,
        (
            "Find the public CEO/Founder/Managing Director/Sales Director for this "
            "distributor or channel company using live web search. Return JSON only: "
            '{"contact_person":"Full Name or empty","role":"Title or empty",'
            '"email":"public email or empty","confidence":"high|medium|low"} '
            "Do NOT invent names."
        ),
        (
            f"Market: {market_q}\nCompany: {company}\nWebsite: {website}\n"
            f'Search LinkedIn and official site for leadership of "{company}".'
        ),
        label="final-contact",
        max_tokens=400,
    )
    if not isinstance(data, dict):
        return {}
    person = str(data.get("contact_person") or "").strip()
    role = str(data.get("role") or "").strip()
    email = str(data.get("email") or "").strip()
    if not looks_person(person, company):
        return {}
    out = {"Contact Person": person}
    if role and role.lower() not in NPD and len(role) < 80:
        out["Role"] = role
    if email and "@" in email and email.lower() not in NPD:
        out["Email"] = email
    return out


async def process(key, path, sheet, market_q, client, model):
    headers, rows, header_row = load(path, sheet)
    ckpt_path = path.parent / f"final_oai_contact_{key}.json"
    ckpt = json.loads(ckpt_path.read_text(encoding="utf-8")) if ckpt_path.exists() else {}
    for row in rows:
        c = row.get("Company") or ""
        if c in ckpt and empty(row.get("Contact Person")) and looks_person(ckpt[c].get("Contact Person", ""), c):
            row["Contact Person"] = ckpt[c]["Contact Person"]
            if ckpt[c].get("Role") and empty(row.get("Role")):
                row["Role"] = ckpt[c]["Role"]
            if ckpt[c].get("Email") and empty(row.get("Email")):
                row["Email"] = ckpt[c]["Email"]

    todo = [i for i, r in enumerate(rows) if empty(r.get("Contact Person")) or not looks_person(r.get("Contact Person") or "", r.get("Company") or "")]
    # only fill empties / NPD, not overwrite good names
    todo = [i for i, r in enumerate(rows) if empty(r.get("Contact Person"))]
    limit = os.getenv("FINAL_OAI_LIMIT")
    if limit and limit.isdigit():
        todo = todo[: int(limit)]
    print(f"\n===== {key} empty_contact={len(todo)} / total={len(rows)} =====", flush=True)
    sem = asyncio.Semaphore(int(os.getenv("FINAL_OAI_CONCURRENT") or "3"))
    filled = 0

    async def one(i):
        nonlocal filled
        row = rows[i]
        company = row.get("Company") or ""
        async with sem:
            print(f"  [oai] {company[:55].encode('ascii','replace').decode()}", flush=True)
            hit = await fill_one(client, model, company, row.get("Website") or "", market_q)
            if hit.get("Contact Person"):
                row["Contact Person"] = hit["Contact Person"]
                if hit.get("Role") and empty(row.get("Role")):
                    row["Role"] = hit["Role"]
                if hit.get("Email") and empty(row.get("Email")):
                    row["Email"] = hit["Email"]
                src = row.get("Data Sources") or ""
                if "final_oai_web_contact" not in src:
                    row["Data Sources"] = (src + "; final_oai_web_contact").strip("; ")
                ckpt[company] = {
                    "Contact Person": row["Contact Person"],
                    "Role": row.get("Role") or "",
                    "Email": row.get("Email") or "",
                }
                filled += 1
                ckpt_path.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")

    await asyncio.gather(*[one(i) for i in todo])
    save(path, sheet, headers, header_row, rows)
    named = sum(1 for r in rows if looks_person(r.get("Contact Person") or "", r.get("Company") or ""))
    still = sum(1 for r in rows if empty(r.get("Contact Person")))
    print(f"  filled=+{filled} named={named} still_empty/npd={still}", flush=True)
    return {"market": key, "total": len(rows), "filled": filled, "named": named, "still_empty": still}


async def main():
    client = OpenAI(
        api_key=(os.getenv("OPENAI_API_KEY") or "").strip(),
        base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.openai.com/v1").rstrip("/"),
    )
    model = (os.getenv("OPENAI_MODEL") or "gpt-4o-mini").strip()
    only = os.getenv("FINAL_ONLY", "").strip()
    files = FILES
    if only:
        files = [f for f in FILES if f[0] in only.split(",")]
    results = []
    for item in files:
        results.append(await process(*item, client=client, model=model))
    out = ROOT / "output/five_markets_oai_contact_summary.json"
    out.write_text(json.dumps(results, indent=2), encoding="utf-8")
    print("\nTOTAL companies", sum(r["total"] for r in results))
    print(out)


if __name__ == "__main__":
    asyncio.run(main())
