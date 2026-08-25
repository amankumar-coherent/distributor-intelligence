# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pathlib import Path
import re

script = Path(r"D:\vendor-intel\scripts\verify_fill_hq_flexpack.py").read_text(encoding="utf-8")
block = script.split("MANUAL_SEEDS", 1)[1].split("COUNTRY_ONLY", 1)[0]
seeds = dict(re.findall(r'"([^"]+)":\s*"([^"]+)"', block))

# Extra seeds to merge into script later
EXTRA = {
    "Allpack": "Cannock, United Kingdom",
    "Aropack": "Hythe, Kent, United Kingdom",
    "Bags Plus": "Malvern, Victoria, Australia",
    "Box Brothers": "Woodland Hills, California, United States",
    "Carolina Retail Packaging": "Lexington, South Carolina, United States",
    "Horna": "Grafenrheinfeld, Germany",
    "Chase Packaging": "Enfield, United Kingdom",
    "Complete Packaging Products": "Troy, Michigan, United States",
    "Delaware Valley Packaging": "Bensalem, Pennsylvania, United States",
    "Springpack": "Worcester, United Kingdom",
    "RAJA": "Tremblay-en-France, France",
    "Packaging Corporation of America": "Lake Forest, Illinois, United States",
}
seeds.update(EXTRA)

FP = Path(r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global\global_flexible_packaging_market_global_FINAL.xlsx")
wb = load_workbook(FP)
ws = wb.active
applied = 0
names = []
for row in ws.iter_rows(min_row=2):
    name = str(row[0].value or "").strip()
    names.append(name)
    if name in seeds:
        row[3].value = seeds[name]
        applied += 1
wb.save(FP)
print(f"applied_seeds={applied} total_seed_keys={len(seeds)}")

# fuzzy check missing extras
for k in EXTRA:
    if k not in names:
        hits = [n for n in names if k.lower() in n.lower() or n.lower() in k.lower()]
        print(f"MISSING exact '{k}' approx={hits[:5]}")
