import openpyxl

X = r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
wb = openpyxl.load_workbook(X, data_only=True)
ws = wb.active
hdr = [c.value for c in ws[2]]
cols = [h for h in hdr if h]
ic = {str(h): i for i, h in enumerate(hdr) if h}
NPD = "Not publicly disclosed"

rows = []
for row in ws.iter_rows(min_row=3, values_only=True):
    c = str(row[ic["Company"]] or "").strip()
    if not c:
        continue
    rows.append(row)
n = len(rows)
print("COMPANY_ROWS", n)
print()
print(f"{'Column':<35} {'Empty':>6} {'NPD':>6} {'Empty+NPD':>10} {'Filled':>7}")
print("-" * 70)
for h in cols:
    empty = npd = 0
    for row in rows:
        v = row[ic[h]]
        s = "" if v is None else str(v).strip()
        if not s:
            empty += 1
        elif s == NPD:
            npd += 1
    print(f"{h:<35} {empty:>6} {npd:>6} {empty + npd:>10} {n - empty - npd:>7}")
print("-" * 70)
wb.close()
