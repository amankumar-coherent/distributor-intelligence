# -*- coding: utf-8 -*-
"""Patch Founded column from verified dict + any good cache hits. Never invent."""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
CACHE = ROOT / "_wearable_founded_harvest_cache.json"
AUDIT = ROOT / "_wearable_founded_patch_audit.json"
REMAINING = ROOT / "_wearable_founded_still_needed.json"

VERIFIED: dict[str, str] = {
    "AAH Pharmaceuticals": "1892",
    "AddLife AB": "2014",
    "AddLife": "2014",
    "Adcock Ingram": "1890",
    "Adcock Ingram Critical Care Distribution": "1890",
    "180 Medical": "2002",
    "Advanced Diabetes Supply": "2002",
    "ActivStyle": "1997",
    "AdaptHealth": "2012",
    "Medline": "1966",
    "Medline Industries": "1966",
    "Henry Schein": "1932",
    "Cardinal Health": "1971",
    "McKesson": "1833",
    "McKesson Corporation": "1833",
    "McKesson Canada": "1833",
    "Suzuken": "1932",
    "Alfresa Holdings Corporation": "2003",
    "Alfresa Holdings": "2003",
    "Toho Holdings": "1948",
    "Bound Tree Medical": "1978",
    "Life-Assist": "1977",
    "Vitality Medical": "2000",
    "Concordance Healthcare Solutions": "2016",
    "Mediq": "1899",
    "Grifols": "1909",
    "Medipal Holdings": "1898",
    "MEDIPAL HOLDINGS": "1898",
    "Medipal": "1898",
    "Alliance Healthcare": "1997",
    "Alliance Healthcare UK": "1997",
    "Alliance Healthcare France": "1997",
    "Alliance Healthcare Deutschland": "1997",
    "Alliance Healthcare Netherlands": "1997",
    "Alliance Healthcare Spain": "1997",
    "Byram Healthcare": "1968",
    "Edgepark": "1928",
    "Edgepark Medical Supplies": "1928",
    "Owens & Minor": "1882",
    "Owens and Minor": "1882",
    "DKSH": "1865",
    "Zuellig Pharma": "1922",
    "Phoenix Pharmahandel": "1994",
    "PHOENIX Group": "1994",
    "Phoenix Group": "1994",
    "Numotion": "2001",
    "Canadian Hospital Specialties": "1967",
    "CHS Ltd": "1967",
    "Zahrawi Group": "1989",
    "Williams Medical Supplies": "1986",
    "Shields Health Solutions": "2012",
    "Walgreens": "1901",
    "Walgreens Boots Alliance": "2014",
    "Invacare": "1885",
    "Invacare Europe": "1885",
    "Invacare Corporation": "1885",
    "Performance Health": "1918",
    "Rotech Healthcare": "1981",
    "Liberator Medical Supply": "1999",
    "Surgipharm": "1985",
    "B. Braun": "1839",
    "Coloplast": "1957",
    "eBay Inc.": "1995",
    "Grainger": "1927",
    "Fastenal": "1967",
    "Arrow Electronics": "1935",
    "Anixter": "1957",
    "CDW Corporation": "1984",
    "Direct Supply": "1985",
    "Apotek Hjärtat": "2010",
    "Jean Coutu": "1969",
    "MatsukiyoCocokara & Co.": "2015",
    "As One Corporation": "1961",
    "Comifar": "1995",
    "Oriola": "1907",
    "Movianto": "1885",
    "Tri-anim Health Services": "1975",
    "Cirúrgica Fernandes": "1946",
    "Al-Bahar Medical": "2002",
    "Africa Health Supplies": "2015",
    "Bidvest Medical": "1988",
    "Guardian Pharmacy": "2004",
    "Health Products For You": "2002",
    "Healthcare 21": "2003",
    "Healthcare Direct": "1993",
    "Medi Assist India": "2000",
    "Medisource Supply (Asia)": "2003",
    "Medshop Australia": "2005",
    "Meditrade GmbH": "1988",
    "Device Technologies": "1992",
    "Defries Industries": "1980",
    "Buffalo Supply": "1983",
    "Buffalo Hospital Supply": "1989",
    "Jackson Allison Medical & Surgical": "1981",
    "US MED": "1996",
    "AED Superstore": "2002",
    "Mobility123": "2003",
    "Easmed": "2006",
    "LifeHealthcare": "2006",
    "Farmasierra Distribución": "1996",
    "Mediscor": "1989",
    "Pharma Dynamics": "2011",
    "Tata 1mg": "2015",
    "Junnimed Services": "2004",
    "Shenzhen Medlink Electronics": "2004",
    "Shenzhen Medke Technology": "2008",
    "Veloce Solutions": "2012",
    "MedSell": "2018",
    "Terumo Australia": "1973",
    "NDC Inc.": "1902",
}

REJECT = {
    "Health Care Australia",
    "MediLink",
    "Carex Health Brands",
    "Newegg Business",
    "Zoro",
    "Rakuten Medical",
    "Shoprite Holdings",
}

KW = (
    "founded",
    "established",
    "est.",
    "since",
    "incorporated",
    "started",
    "formed",
    "created",
    "establishment",
    "dating back",
    "origins",
)


def empty(v: str) -> bool:
    s = (v or "").strip().lower()
    return (not s) or s.startswith("not public") or s in {
        "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd", "null"
    }


def year_ok(y: str) -> bool:
    return y.isdigit() and 1800 <= int(y) <= datetime.now().year


def main() -> None:
    cache = {}
    if CACHE.exists():
        cache = json.loads(CACHE.read_text(encoding="utf-8"))

    # merge verified into lookup
    lookup = dict(VERIFIED)
    for name, res in cache.items():
        if name in REJECT or name in lookup:
            continue
        y = str(res.get("founded") or "").strip()
        if not year_ok(y):
            continue
        ev = (res.get("evidence") or "").lower()
        if any(bad in ev for bad in ("medical fair asia", "tens 7000", "customer since", "copyright")):
            continue
        if res.get("verified") or res.get("ok") and any(k in ev for k in KW):
            lookup[name] = y

    wb = load_workbook(XLSX)
    ws = wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(hdr) if h}

    filled = 0
    still = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        cur = str(ws.cell(r, ic["Founded"]).value or "").strip()
        if not empty(cur) and cur.isdigit():
            continue
        y = lookup.get(co, "")
        if y and year_ok(y):
            ws.cell(r, ic["Founded"]).value = y
            filled += 1
        else:
            ws.cell(r, ic["Founded"]).value = "Not publicly disclosed"
            still.append({"company": co, "website": web})

    wb.save(XLSX)
    wb.close()

    # recount
    wb2 = load_workbook(XLSX, data_only=True)
    ws2 = wb2.active
    hdr2 = [str(c.value or "").strip() for c in ws2[1]]
    fi = hdr2.index("Founded") + 1
    numeric = npd = 0
    for r in range(2, ws2.max_row + 1):
        f = str(ws2.cell(r, fi).value or "").strip()
        if f.isdigit():
            numeric += 1
        else:
            npd += 1
    wb2.close()

    REMAINING.write_text(json.dumps(still, indent=2, ensure_ascii=False), encoding="utf-8")
    audit = {
        "filled_this_patch": filled,
        "lookup_size": len(lookup),
        "numeric_total": numeric,
        "npd_total": npd,
        "still": len(still),
        "still_sample": still[:80],
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(audit, indent=2)[:2000])
    print("still_count", len(still), "numeric", numeric, "npd", npd)


if __name__ == "__main__":
    main()
