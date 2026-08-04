#!/usr/bin/env python3
"""Backfill empty Headquarters on newly added firewall rows (SearXNG/Wikipedia only)."""
from __future__ import annotations

import asyncio
import re
import sys
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import expand_sw_fw_100_web as m  # noqa: E402

PATH = ROOT / "output/run4/global_firewall_market_global/global_firewall_market_FINAL.xlsx"
SHEET = "Firewall Companies"
PRIOR = 278  # rows before expand


async def hq_for(client: httpx.AsyncClient, company: str) -> str:
    wiki = await m.wiki_summary(client, company)
    if wiki.get("Headquarters"):
        return wiki["Headquarters"]
    hits = await m.searx_search(client, f'"{company}" headquarters OR "based in" OR "head office"')
    blob = "\n".join(f"{h['title']} {h['snippet']}" for h in hits)
    mm = re.search(
        r"(?:headquartered|based|head office|HQ)\s+in\s+([A-Z][^,|.]{2,50})",
        blob,
        re.I,
    )
    if mm:
        return mm.group(1).strip()
    # city, country pattern near company
    mm2 = re.search(r"\b([A-Z][a-zA-Z\- ]{2,30},\s*(?:[A-Z][a-zA-Z\- ]{2,30}|[A-Z]{2,}))\b", blob)
    if mm2:
        return mm2.group(1).strip()
    return ""


async def main() -> None:
    wb = load_workbook(PATH)
    ws = wb[SHEET]
    # find header
    header_row = None
    headers = []
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            header_row = i
            headers = vals
            break
    assert header_row
    ci = headers.index("Company")
    hi = headers.index("Headquarters")
    ri = headers.index("Region Segment") if "Region Segment" in headers else None
    cri = headers.index("Cities / Regions") if "Cities / Regions" in headers else None
    rsi = headers.index("Regions Served") if "Regions Served" in headers else None

    # data rows start header_row+1; new rows are last 100
    max_r = ws.max_row
    start = max_r - 99  # last 100 rows
    targets = []
    for r in range(start, max_r + 1):
        company = str(ws.cell(r, ci + 1).value or "").strip()
        hq = str(ws.cell(r, hi + 1).value or "").strip()
        if company and (not hq or hq.lower() in m.NPD):
            targets.append((r, company))
    print(f"HQ backfill targets: {len(targets)}", flush=True)

    sem = asyncio.Semaphore(6)
    filled = 0

    async with httpx.AsyncClient() as client:

        async def one(r: int, company: str) -> None:
            nonlocal filled
            async with sem:
                print(f"  HQ {company[:50]}", flush=True)
                hq = await hq_for(client, company)
                if not hq:
                    # fallback from region segment
                    reg = str(ws.cell(r, (ri or 0) + 1).value or "").strip() if ri is not None else ""
                    hq = reg or "Not publicly disclosed"
                else:
                    filled += 1
                ws.cell(r, hi + 1).value = hq
                if cri is not None and not str(ws.cell(r, cri + 1).value or "").strip():
                    ws.cell(r, cri + 1).value = hq
                if rsi is not None and not str(ws.cell(r, rsi + 1).value or "").strip():
                    if ri is not None:
                        ws.cell(r, rsi + 1).value = str(ws.cell(r, ri + 1).value or hq)

        await asyncio.gather(*[one(r, c) for r, c in targets])

    wb.save(PATH)
    print(f"HQ filled from web: {filled}; remaining use region fallback; saved {PATH.name}", flush=True)


if __name__ == "__main__":
    asyncio.run(main())
