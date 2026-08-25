# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pathlib import Path

COUNTRY_ONLY = {
    "united states", "usa", "us", "united kingdom", "uk", "great britain", "england",
    "germany", "deutschland", "australia", "canada", "france", "india", "mexico",
    "spain", "ireland", "netherlands", "sweden", "japan", "italy", "singapore",
    "new zealand", "south africa", "switzerland", "austria", "belgium", "china",
    "brazil", "poland", "portugal", "denmark", "norway", "finland",
}

def needs(hq: str) -> bool:
    h = (hq or "").strip()
    if not h or "not publicly" in h.lower():
        return True
    if "," in h:
        left = h.split(",")[0].strip()
        if left and left.lower() not in COUNTRY_ONLY and len(left) >= 3:
            return False
    if h.lower() in COUNTRY_ONLY:
        return True
    return "," not in h and len(h.split()) <= 3

FP = Path(r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global\global_flexible_packaging_market_global_FINAL.xlsx")
wb = load_workbook(FP, data_only=True)
ws = wb.active
weak = []
for r in ws.iter_rows(min_row=2, values_only=True):
    name = str(r[0] or "").strip()
    site = str(r[1] or "").strip()
    hq = str(r[3] or "").strip()
    if needs(hq):
        weak.append(f"{name}\t{site}\t{hq}")
out = FP.parent / "hq_still_weak.txt"
out.write_text("\n".join(weak), encoding="utf-8")
print(f"weak={len(weak)} wrote {out}")
for line in weak[:60]:
    print(line)
