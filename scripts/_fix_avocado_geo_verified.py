#!/usr/bin/env python3
"""Apply web-verified Founded/HQ/Continent/Operational Presence fixes."""
from __future__ import annotations

import re
from collections import defaultdict

import openpyxl
from pathlib import Path

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
    r"\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)

# country -> region code used in this file
COUNTRY_REGION = {
    "united states": "NA",
    "usa": "NA",
    "u.s.": "NA",
    "puerto rico": "NA",
    "canada": "NA",
    "mexico": "NA",
    "costa rica": "LATAM",
    "guatemala": "LATAM",
    "honduras": "LATAM",
    "el salvador": "LATAM",
    "nicaragua": "LATAM",
    "panama": "LATAM",
    "colombia": "LATAM",
    "dominican republic": "LATAM",
    "chile": "LATAM",
    "brazil": "LATAM",
    "argentina": "LATAM",
    "peru": "LATAM",
    "ecuador": "LATAM",
    "uruguay": "LATAM",
    "venezuela": "LATAM",
    "curaçao": "LATAM",
    "curacao": "LATAM",
    "jamaica": "LATAM",
    "barbados": "LATAM",
    "aruba": "LATAM",
    "trinidad and tobago": "LATAM",
    "united states virgin islands": "LATAM",
    "u.s. virgin islands": "LATAM",
    "united kingdom": "EU",
    "england": "EU",
    "scotland": "EU",
    "ireland": "EU",
    "germany": "EU",
    "france": "EU",
    "netherlands": "EU",
    "belgium": "EU",
    "luxembourg": "EU",
    "spain": "EU",
    "italy": "EU",
    "portugal": "EU",
    "switzerland": "EU",
    "austria": "EU",
    "sweden": "EU",
    "norway": "EU",
    "denmark": "EU",
    "finland": "EU",
    "iceland": "EU",
    "poland": "EU",
    "czech republic": "EU",
    "slovakia": "EU",
    "hungary": "EU",
    "romania": "EU",
    "greece": "EU",
    "bulgaria": "EU",
    "croatia": "EU",
    "serbia": "EU",
    "slovenia": "EU",
    "estonia": "EU",
    "latvia": "EU",
    "lithuania": "EU",
    "ukraine": "EU",
    "belarus": "EU",
    "moldova": "EU",
    "turkey": "EU",
    "georgia": "EU",
    "russia": "EU",
    "russian federation": "EU",
    "kazakhstan": "EU",
    "uzbekistan": "EU",
    "united arab emirates": "ME",
    "saudi arabia": "ME",
    "qatar": "ME",
    "oman": "ME",
    "bahrain": "ME",
    "kuwait": "ME",
    "jordan": "ME",
    "israel": "ME",
    "pakistan": "APAC",
    "india": "APAC",
    "china": "APAC",
    "japan": "APAC",
    "korea": "APAC",
    "south korea": "APAC",
    "taiwan": "APAC",
    "hong kong": "APAC",
    "singapore": "APAC",
    "malaysia": "APAC",
    "indonesia": "APAC",
    "thailand": "APAC",
    "vietnam": "APAC",
    "philippines": "APAC",
    "australia": "APAC",
    "new zealand": "APAC",
    "south africa": "AF",
    "tunisia": "AF",
    "zimbabwe": "AF",
    "botswana": "AF",
    "namibia": "AF",
    "zambia": "AF",
    "malawi": "AF",
    "mozambique": "AF",
    "lesotho": "AF",
    "eswatini": "AF",
    "egypt": "AF",
    "ghana": "AF",
}


def continent_from_presence(presence: str) -> str:
    parts = [p.strip() for p in re.split(r"[;]", presence) if p.strip()]
    buckets: dict[str, list[str]] = defaultdict(list)
    for p in parts:
        key = p.lower()
        reg = None
        for cname, code in COUNTRY_REGION.items():
            if cname == key or cname in key:
                reg = code
                break
        if not reg:
            # keep unknown under APAC? skip unknown rather than invent region
            continue
        if p not in buckets[reg]:
            buckets[reg].append(p)
    order = ["NA", "EU", "ME", "APAC", "LATAM", "AF"]
    chunks = []
    for code in order:
        if code in buckets:
            chunks.append(f"{code}({', '.join(buckets[code])})")
    for code, countries in buckets.items():
        if code not in order:
            chunks.append(f"{code}({', '.join(countries)})")
    return "; ".join(chunks) if chunks else presence


# Web-verified overrides only (sources in comments)
OVERRIDES = {
    "Anchor Foods Wholesale": {
        # Wikipedia / anchorfoods.com.au — Fremantle AU manufacturer; presence AU (not UK)
        "founded": "1854",
        "hq": "Fremantle, Western Australia, Australia",
        "presence": "Australia",
    },
    "Barentz": {
        # barentz.com/where-we-are + contact HQ Hoofddorp; founded 1953
        "founded": "1953",
        "hq": "Hoofddorp, Netherlands",
        "presence": (
            "Netherlands; Belgium; Luxembourg; Germany; France; United Kingdom; Ireland; "
            "Spain; Portugal; Italy; Switzerland; Austria; Sweden; Norway; Denmark; Finland; "
            "Iceland; Poland; Czech Republic; Slovakia; Hungary; Romania; Bulgaria; Greece; "
            "Croatia; Serbia; Slovenia; Estonia; Latvia; Lithuania; Ukraine; Belarus; "
            "Turkey; Georgia; Kazakhstan; Uzbekistan; Russian Federation; United Arab Emirates; "
            "Tunisia; South Africa; Zimbabwe; United States; Canada; Mexico; Costa Rica; "
            "Guatemala; Honduras; El Salvador; Dominican Republic; Brazil; Argentina; Chile; "
            "Colombia; Peru; Ecuador; Uruguay; Curaçao; Singapore; China; India; Indonesia; "
            "Malaysia; Thailand; Vietnam; Philippines; Pakistan; Australia; New Zealand"
        ),
    },
    "Costco": {
        # Costco SEC / investor warehouse count (2026)
        "founded": "1983",
        "hq": "Issaquah, Washington, United States",
        "presence": (
            "United States; Puerto Rico; Canada; Mexico; Japan; United Kingdom; Korea; "
            "Australia; Taiwan; China; Spain; France; Sweden; Iceland; New Zealand"
        ),
    },
    "Walmart": {
        # corporate.walmart.com international markets / 10-K
        "founded": "1962",
        "hq": "Bentonville, Arkansas, United States",
        "presence": (
            "United States; Canada; Mexico; Chile; China; India; Costa Rica; El Salvador; "
            "Guatemala; Honduras; Nicaragua; Botswana; Eswatini; Lesotho; Malawi; "
            "Mozambique; Namibia; South Africa; Zambia"
        ),
    },
    "PriceSmart": {
        # PriceSmart FY25 10-K / investors
        "founded": "1994",
        "hq": "San Diego, California, United States",
        "presence": (
            "United States; Colombia; Costa Rica; Panama; Guatemala; Dominican Republic; "
            "Trinidad and Tobago; El Salvador; Honduras; Nicaragua; Jamaica; Aruba; "
            "Barbados; United States Virgin Islands"
        ),
    },
    "Makro": {
        # Makro South Africa (Massmart/Walmart) — Sandton; founded 1971; SA only for this entity
        "founded": "1971",
        "hq": "Sandton, Gauteng, South Africa",
        "presence": "South Africa",
    },
    "Scavuzzo's": {
        # scavuzzos.com / RocketReach — Kansas City foodservice; not Argentina
        "founded": "1906",
        "hq": "Kansas City, Kansas, United States",
        "presence": "United States",
    },
    "Aramark Middle East Distribution": {
        # ME ops listed; keep parent HQ US but continent from ME presence
        "hq": "Philadelphia, Pennsylvania, United States",
        "presence": "United Arab Emirates; Saudi Arabia; Qatar; Oman",
    },
    "Cargill Food Distribution Europe": {
        "hq": "Minnetonka, Minnesota, United States",
        # keep Europe presence already listed — only fix continent rebuild
    },
    "Fresh Del Monte Europe": {
        "hq": "Coral Gables, Florida, United States",
    },
}


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

    changed = 0
    for row in range(3, ws.max_row + 1):
        co = str(ws.cell(row, ic["Company"]).value or "").strip()
        if co not in OVERRIDES:
            continue
        hit = OVERRIDES[co]
        if hit.get("founded"):
            ws.cell(row, ic["Founded"]).value = hit["founded"]
        if hit.get("hq"):
            ws.cell(row, ic["Headquarters"]).value = hit["hq"]
        if hit.get("presence"):
            ws.cell(row, ic["Operational Presence"]).value = hit["presence"]
        # always rebuild continent from (updated) presence for these rows
        presence = str(ws.cell(row, ic["Operational Presence"]).value or "").strip()
        cont = continent_from_presence(presence)
        if cont:
            ws.cell(row, ic["Continent / Geography"]).value = cont
        changed += 1
        print(f"fixed {co}")
        print(f"  HQ={ws.cell(row, ic['Headquarters']).value}")
        print(f"  C={ws.cell(row, ic['Continent / Geography']).value}")
        print(f"  P={str(ws.cell(row, ic['Operational Presence']).value)[:120]}...")

    # also rebuild continent for any row where HQ country missing from presence was flagged
    # already handled via overrides

    wb.save(XLSX)
    wb.close()
    print(f"changed={changed} saved={XLSX}")


if __name__ == "__main__":
    main()
