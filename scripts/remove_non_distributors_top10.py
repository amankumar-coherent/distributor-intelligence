"""Remove non-distributor rows from top-10 *_distributors_expanded.xlsx files.

Drops government portals, marketplaces, airlines, associations, and
trade-facilitation listings that are not real distributors.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(r"D:\vendor-intel\output")

# Exact / substring match on Company name (case-insensitive)
DROP_NAME_SUBSTR = [
    "pdunippd",
    "gesund.bund.de",
    "goldsupplier",
    "koreanair",
    "korean air",
    "kotra",
    "indiamart",
    "alibaba textile",
    "alibaba group b2b",
    "mask sheets oem exporters",
    "textile fabric store (india exporters",
    "vestibular disorders association",
    "world instant noodles association",
    "government portal",
    "federal ministry",
    "marketplace directory",
    "b2b supplier directory",
]

# Exact company names (lowercase) that are associations/portals, not distributors
DROP_NAME_EXACT = {
    "instantnoodles",
    "instantnoodles.org",
}

# Also drop if Business Type or Data Sources clearly mark non-distributor entity
DROP_BLOB_SUBSTR = [
    "n/a — government",
    "n/a — marketplace",
    "government portal",
    "government institute",
    "not a distributor",
    "marketplace directory",
    "federal health portal",
    "industry association",
]


def should_drop(company: str, business_type: str, sources: str, summary: str) -> bool:
    name = (company or "").lower().strip()
    blob = " ".join([company or "", business_type or "", sources or "", summary or ""]).lower()
    if name in DROP_NAME_EXACT:
        return True
    if any(s in name for s in DROP_NAME_SUBSTR):
        return True
    if any(s in blob for s in DROP_BLOB_SUBSTR):
        return True
    return False


def process(path: Path) -> tuple[int, list[str]]:
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=False))
    hdr_i = next(
        i
        for i, r in enumerate(rows[:12])
        if r and r[0].value and str(r[0].value).strip() == "Company"
    )
    header = [str(c.value or "").strip() for c in rows[hdr_i]]
    ci = header.index("Company")
    bti = header.index("Business Type") if "Business Type" in header else None
    si = header.index("Data Sources") if "Data Sources" in header else None
    sui = header.index("Summary") if "Summary" in header else None

    # Collect data rows to keep (values), then rewrite sheet body
    keep_vals: list[tuple] = []
    removed: list[str] = []
    for r in rows[hdr_i + 1 :]:
        company = str(r[ci].value or "").strip()
        if not company:
            continue
        bt = str(r[bti].value or "") if bti is not None else ""
        sources = str(r[si].value or "") if si is not None else ""
        summary = str(r[sui].value or "") if sui is not None else ""
        if should_drop(company, bt, sources, summary):
            removed.append(company)
            continue
        keep_vals.append(tuple(c.value for c in r[: len(header)]))

    if not removed:
        wb.close()
        return 0, []

    # Excel rows are 1-indexed; header is at hdr_i+1; first data at hdr_i+2
    header_row = hdr_i + 1
    first_data_row = hdr_i + 2
    # Delete all data rows, keep title + blank + header
    while ws.max_row >= first_data_row:
        ws.delete_rows(first_data_row)

    for i, vals in enumerate(keep_vals):
        for col, val in enumerate(vals, 1):
            ws.cell(first_data_row + i, col).value = val

    if "Run Summary" in wb.sheetnames:
        ws2 = wb["Run Summary"]
        ws2["A5"] = f"Distributors after non-distributor removal: {len(keep_vals)}"
        ws2["A6"] = f"Removed: {', '.join(removed)}"

    wb.save(path)
    wb.close()
    return len(removed), removed


def main() -> None:
    total = 0
    for path in sorted(ROOT.rglob("*_distributors_expanded.xlsx")):
        n, removed = process(path)
        total += n
        if removed:
            print(f"{path.parent.name}: removed {n}")
            for name in removed:
                print(f"  - {name}")
        else:
            print(f"{path.parent.name}: none")
    print(f"\nTotal removed: {total}")


if __name__ == "__main__":
    main()
