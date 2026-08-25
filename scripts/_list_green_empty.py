# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pathlib import Path

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}

files = {
    "xlsx": Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market.xlsx"),
    "FINAL": Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL.xlsx"),
}

for key, path in files.items():
    wb = load_workbook(path, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and ("Contact Person" in vals or "Contact" in vals):
            headers = vals
            hr = r
            break
    print("===", key, "sheet", ws.title, "header_row", hr)
    print("headers:", headers)
    tot = cp = li_url = li_npd = li_empty = 0
    empties = []
    named_npd = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {
            headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
        }
        co = d.get("Company") or ""
        if not co:
            continue
        tot += 1
        person = (d.get("Contact Person") or d.get("Contact") or "").strip()
        role = (d.get("Role") or "").strip()
        li = (d.get("LinkedIn") or "").strip()
        web = (d.get("Website") or "").strip()
        hq = (d.get("HQ") or d.get("Headquarters") or d.get("Country") or "").strip()
        has_cp = bool(person) and person.lower() not in NPD
        if has_cp:
            cp += 1
            if "linkedin.com" in li.lower():
                li_url += 1
            else:
                li_npd += 1
                named_npd.append((co, person, role, li))
        else:
            li_empty += 1
            empties.append((co, web, hq, person, role, li))
    print(f"rows={tot} cp={cp} li_url={li_url} named_npd={li_npd} empty_cp={li_empty}")
    print("--- EMPTY (first 80) ---")
    for i, e in enumerate(empties[:80], 1):
        print(f"{i}\t{e[0]}\t{e[1]}\t{e[2]}")
    print("empty_total", len(empties))
    wb.close()
