# -*- coding: utf-8 -*-
from datetime import datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook

p = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
    r"\global_wearable_medical_devices_market_global_FINAL.xlsx"
)
stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
bak = p.with_name(f"{p.stem}_BEFORE_DROP_DISTTYPE_{stamp}.xlsx")
shutil.copy2(p, bak)

wb = load_workbook(p)
ws = wb.active
headers = [c.value for c in ws[1]]
print("before:", headers)
if "Distribution Type" not in headers:
    raise SystemExit("Distribution Type column not found")
idx = headers.index("Distribution Type") + 1
ws.delete_cols(idx)
wb.save(p)

wb2 = load_workbook(p, read_only=True)
h2 = [c.value for c in next(wb2.active.iter_rows(min_row=1, max_row=1))]
print("after:", h2)
print("backup:", bak.name)
print("rows:", wb2.active.max_row - 1)
