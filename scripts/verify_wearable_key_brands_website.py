#!/usr/bin/env python3
"""Follow-up pass: fill Key Brands Represented by crawling each company's own
website (homepage + likely brand/partner pages) — targets only rows still
showing 'Not publicly disclosed' (or empty) after the Google AI / ddgs pass.

No LLM anywhere: brand candidates come from three regex-only signals on the
raw page HTML/text —
  1. <img alt="..."> values (brand logos are very often alt-text only)
  2. <li> items inside a list that follows a "Brands/Partners/Principals" heading
  3. the same brand-list trigger phrase regex used in the search-based pass

No Claude subagents; plain httpx fetches run directly in this process.

Resumable: one JSON checkpoint per company under
output/web_expand/global_wearable_medical_devices_market_global/verify_checkpoints_brands/
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import urljoin, urlparse

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import httpx
import openpyxl
from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

XLSX = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/global_wearable_medical_devices_market_global_FINAL.xlsx"
SHEET = "Wearable Medical Companies"
CKPT_DIR = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/verify_checkpoints_brands"
CKPT_DIR.mkdir(parents=True, exist_ok=True)

CONCURRENCY = int(os.getenv("BRANDS_CONCURRENCY", "5"))
NOT_DISCLOSED = "Not publicly disclosed"
TIMEOUT = float(os.getenv("BRANDS_FETCH_TIMEOUT_SEC", "12"))

CANDIDATE_PATHS = ["", "/brands", "/our-brands", "/partners", "/our-partners", "/principals", "/manufacturers"]

_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36"
)
_HEADERS = {"User-Agent": _UA, "Accept-Language": "en-US,en;q=0.9"}

_IMG_ALT_RE = re.compile(r'<img[^>]+alt=["\']([^"\']{2,60})["\']', re.I)
_HEADING_RE = re.compile(
    r"<(?:h[1-6]|strong|b)[^>]*>\s*(?:our\s+)?(?:brands?|partners?|principals?|manufacturers?)\b[^<]{0,40}</",
    re.I,
)
_LI_RE = re.compile(r"<li[^>]*>(.*?)</li>", re.I | re.S)
_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"\s+")

_BRAND_TRIGGER_RE = re.compile(
    r"(?:brands?|principals?|partners?|manufacturers?)[\s—–\-:,]*"
    r"(?:such as|include[sd]?|including|represented|distributed|carried|sold|are|:)?[\s—–\-:,]*"
    r"((?:[A-Z][\w&.'\-]+(?:\s+[A-Z][\w&.'\-]+){0,3})"
    r"(?:\s*,\s*(?:and\s+)?[A-Z][\w&.'\-]+(?:\s+[A-Z][\w&.'\-]+){0,3}){1,14})",
)

_JUNK_ALT = {
    "logo", "icon", "banner", "arrow", "menu", "close", "search", "facebook", "twitter",
    "instagram", "linkedin", "youtube", "home", "header", "footer", "spacer", "placeholder",
    "avatar", "star", "flag", "none", "image", "photo", "background", "hero", "thumbnail",
    "play", "video", "loading", "spinner", "chevron", "caret", "check", "plus", "minus",
}


def _strip_tags(html: str) -> str:
    return _WS_RE.sub(" ", _TAG_RE.sub(" ", html)).strip()


def _clean_candidate(name: str, company_tokens: set[str]) -> str | None:
    n = _strip_tags(name).strip(" .,;:-|")
    n = re.sub(r"\.(jpe?g|png|gif|svg|webp|bmp|ico)$", "", n, flags=re.I)
    n = re.sub(r"[_\-]\d+$", "", n)
    n = n.strip(" .,;:-|_")
    if not n or len(n) < 2 or len(n) > 40:
        return None
    low = n.lower()
    if low in _JUNK_ALT:
        return None
    if any(t and t in low for t in company_tokens):
        return None
    if not re.search(r"[A-Za-z]", n):
        return None
    if not re.search(r"[A-Z]", n):  # require a proper-noun-like capital letter
        return None
    if re.match(r"^\d", n):  # filename-style residue ("1625 Ply4")
        return None
    if re.search(r"\d", n) and re.search(r"[_\-]", n):  # e.g. "SVL-BORD1212-34"
        return None
    if n.isupper() and len(n) > 12:
        return None
    return n


def _company_tokens(company: str) -> set[str]:
    toks = re.sub(r"[^a-z0-9\s]", " ", company.lower()).split()
    stop = {"inc", "ltd", "llc", "corp", "corporation", "company", "co", "group", "the", "and"}
    return {t for t in toks if len(t) >= 4 and t not in stop}


def extract_brand_candidates(html: str, text: str, company: str) -> list[str]:
    tokens = _company_tokens(company)
    found: list[str] = []

    for m in _IMG_ALT_RE.finditer(html):
        c = _clean_candidate(m.group(1), tokens)
        if c:
            found.append(c)

    for hm in _HEADING_RE.finditer(html):
        window = html[hm.end() : hm.end() + 3000]
        ul_match = re.search(r"<ul[^>]*>(.*?)</ul>", window, re.I | re.S)
        if ul_match:
            for lm in _LI_RE.finditer(ul_match.group(1)):
                c = _clean_candidate(lm.group(1), tokens)
                if c:
                    found.append(c)

    for m in _BRAND_TRIGGER_RE.finditer(text):
        raw = m.group(1)
        parts = [p.strip(" .,;:") for p in re.split(r",|\band\b", raw) if p.strip(" .,;:")]
        for p in parts:
            c = _clean_candidate(p, tokens)
            if c:
                found.append(c)

    seen: set[str] = set()
    out: list[str] = []
    for f in found:
        key = f.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(f)
    return out[:20]


def _fetch(url: str) -> tuple[str, str]:
    try:
        with httpx.Client(follow_redirects=True, timeout=TIMEOUT, headers=_HEADERS) as client:
            r = client.get(url)
            if r.status_code >= 400:
                return "", ""
            html = r.text or ""
            text = _strip_tags(html)
            return html[:200000], text[:20000]
    except Exception:
        return "", ""


async def research_company(row: dict) -> dict:
    company = row["company"]
    website = row["website"]
    ckpt_path = CKPT_DIR / f"{row['row']:04d}.json"
    if ckpt_path.exists():
        try:
            return json.loads(ckpt_path.read_text(encoding="utf-8"))
        except Exception:
            pass

    base = website.strip()
    if not base:
        result = {"row": row["row"], "company": company, "key_brands": "", "pages_checked": 0}
        ckpt_path.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
        return result
    if not base.startswith("http"):
        base = "https://" + base.lstrip("/")
    parsed = urlparse(base)
    root = f"{parsed.scheme}://{parsed.netloc}"

    all_candidates: list[str] = []
    pages_checked = 0
    for path in CANDIDATE_PATHS:
        url = urljoin(root, path)
        html, text = await asyncio.to_thread(_fetch, url)
        if not html:
            continue
        pages_checked += 1
        cands = extract_brand_candidates(html, text, company)
        if cands:
            all_candidates.extend(cands)
            if len(all_candidates) >= 6:
                break

    seen: set[str] = set()
    out: list[str] = []
    for c in all_candidates:
        key = c.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(c)
    out = out[:15]

    result = {
        "row": row["row"],
        "company": company,
        "key_brands": ", ".join(out) if out else "",
        "pages_checked": pages_checked,
    }
    ckpt_path.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    return result


async def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}

    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        company = ws.cell(row=r, column=idx["Company"]).value
        if not company:
            continue
        current_brands = str(ws.cell(row=r, column=idx["Key Brands Represented"]).value or "").strip()
        if current_brands and current_brands.lower() != NOT_DISCLOSED.lower():
            continue  # already has real brands — don't touch
        rows.append(
            {
                "row": r,
                "company": str(company).strip(),
                "website": str(ws.cell(row=r, column=idx["Website"]).value or "").strip(),
            }
        )

    limit = int(os.getenv("BRANDS_LIMIT", "0") or "0")
    if limit > 0:
        rows = rows[:limit]

    print(f"Companies needing brand lookup: {len(rows)}", flush=True)

    sem = asyncio.Semaphore(CONCURRENCY)
    completed = 0

    async def _one(row: dict) -> None:
        nonlocal completed
        async with sem:
            result = await research_company(row)
            completed += 1
            print(
                f"[{completed}/{len(rows)}] row {row['row']} {row['company']!r} "
                f"found={'yes' if result.get('key_brands') else 'no'} "
                f"pages={result.get('pages_checked', 0)}",
                flush=True,
            )

    await asyncio.gather(*[_one(row) for row in rows])

    print("Done. Writing results back into workbook...", flush=True)
    written = 0
    for row in rows:
        ckpt_path = CKPT_DIR / f"{row['row']:04d}.json"
        if not ckpt_path.exists():
            continue
        result = json.loads(ckpt_path.read_text(encoding="utf-8"))
        if result.get("key_brands"):
            ws.cell(row=row["row"], column=idx["Key Brands Represented"], value=result["key_brands"])
            written += 1

    wb.save(XLSX)
    print(f"Saved {XLSX} — {written} rows given a real brand list.", flush=True)


if __name__ == "__main__":
    asyncio.run(main())
