from collections import Counter
import re
from openpyxl import load_workbook

FILES = [
    ("smartwatch", r"output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL.xlsx", "Smartwatch Companies"),
    ("green", r"output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL.xlsx", "Green Chemical Companies"),
    ("firewall", r"output/run4/global_firewall_market_global/global_firewall_market_FINAL.xlsx", "Firewall Companies"),
    ("smartphone", r"output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx", "Companies"),
    ("avocado", r"output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL.xlsx", "Avocado Oil Distributors"),
]
DIST = re.compile(
    r"distributor|wholesale|retail|channel|dealer|reseller|vad|var|partner|authorized|marketplace|foodservice|importer|integrator|mssp|buying|e-commerce|cash",
    re.I,
)

for key, path, sheet in FILES:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    hi = None
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            hi = i
            hdr = vals
            break
    data = []
    for r in rows[hi + 1 :]:
        if r and r[hdr.index("Company")]:
            data.append(
                {
                    hdr[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
                    for j in range(len(hdr))
                    if hdr[j]
                }
            )
    print(f"\n{key} total={len(data)}")
    for needle in [
        "ingram micro",
        "exclusive networks",
        "arrow ecs",
        "td synnex",
        "redington",
        "also holding",
        "infinigate",
        "westcon",
    ]:
        hits = [d["Company"] for d in data if needle in d["Company"].lower()]
        if hits:
            print(f"  {needle}: {len(hits)} -> {hits}")
    non = [
        f"{d['Company']}|{d.get('Business Type')}"
        for d in data
        if (d.get("Business Type") or "").strip() and not DIST.search(d.get("Business Type") or "")
    ]
    print(f"  non-channel-looking BT: {len(non)}")
    for x in non[:8]:
        print("   ", x)
    print("  BT:", Counter(d.get("Business Type") or "" for d in data).most_common(6))
    if key == "smartphone":
        print("  micromax:", [d["Company"] for d in data if "micromax" in d["Company"].lower()])
