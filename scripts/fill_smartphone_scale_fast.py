#!/usr/bin/env python3
"""FAST fill Employees + Turnover/Scale as ranges — Wikipedia + web regex (no per-row LLM)."""
from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import sys
import urllib.parse
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

FINAL = (
    ROOT
    / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx"
)
SHEET = "Companies"
NPD = "Not publicly disclosed"
WORKERS = int(os.getenv("SCALE_FAST_WORKERS") or "10")

EMP_RANGES = (
    "1–10",
    "11–50",
    "51–200",
    "201–500",
    "501–1,000",
    "1,001–5,000",
    "5,001–10,000",
    "10,001–50,000",
    "50,001+",
)
TO_RANGES = (
    "<$1M",
    "$1–10M",
    "$10–50M",
    "$50–100M",
    "$100–500M",
    "$500M–1B",
    "$1–5B",
    "$5–10B",
    "$10–50B",
    "$50B+",
)

KNOWN: dict[str, dict[str, str]] = {
    "amazon": {"employees": "50,001+", "turnover": "$50B+"},
    "best buy": {"employees": "50,001+", "turnover": "$10–50B"},
    "currys": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "cdw": {"employees": "10,001–50,000", "turnover": "$10–50B"},
    "also holding": {"employees": "1,001–5,000", "turnover": "$10–50B"},
    "ab sa": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "ingram micro": {"employees": "10,001–50,000", "turnover": "$50B+"},
    "td synnex": {"employees": "10,001–50,000", "turnover": "$50B+"},
    "d&h distributing": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "esprinet": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "computacenter": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "bechtle": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "dicker data": {"employees": "501–1,000", "turnover": "$1–5B"},
    "elko group": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "brightstar": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "exertis": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "exertis ireland": {"employees": "501–1,000", "turnover": "$500M–1B"},
    "lazada": {"employees": "5,001–10,000", "turnover": "$1–5B"},
    "erajaya": {"employees": "5,001–10,000", "turnover": "$1–5B"},
    "insight enterprises": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "softcat": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "shi international": {"employees": "5,001–10,000", "turnover": "$10–50B"},
    "softchoice": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "zones": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "ban leong technologies": {"employees": "201–500", "turnover": "$100–500M"},
    "asbis": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "redington": {"employees": "1,001–5,000", "turnover": "$5–10B"},
    "rashi peripherals": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "westcoast": {"employees": "501–1,000", "turnover": "$1–5B"},
    "3g mobile": {"employees": "51–200", "turnover": "$100–500M"},
    "reliance jio": {"employees": "50,001+", "turnover": "$10–50B"},
    "croma": {"employees": "5,001–10,000", "turnover": "$1–5B"},
    "fido solutions inc.": {"employees": "1,001–5,000", "turnover": "$1–5B"},
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _emptyish(v: str) -> bool:
    return _norm(v) in {
        "", "n/a", "na", "none", "-", "—", "not publicly disclosed",
        "unknown", "tbd", "todo", "seed_file",
    }


def _emp_num_to_range(n: int) -> str:
    if n <= 0:
        return ""
    if n <= 10:
        return "1–10"
    if n <= 50:
        return "11–50"
    if n <= 200:
        return "51–200"
    if n <= 500:
        return "201–500"
    if n <= 1000:
        return "501–1,000"
    if n <= 5000:
        return "1,001–5,000"
    if n <= 10000:
        return "5,001–10,000"
    if n <= 50000:
        return "10,001–50,000"
    return "50,001+"


def _usd_to_range(usd: float) -> str:
    if usd < 1_000_000:
        return "<$1M"
    if usd < 10_000_000:
        return "$1–10M"
    if usd < 50_000_000:
        return "$10–50M"
    if usd < 100_000_000:
        return "$50–100M"
    if usd < 500_000_000:
        return "$100–500M"
    if usd < 1_000_000_000:
        return "$500M–1B"
    if usd < 5_000_000_000:
        return "$1–5B"
    if usd < 10_000_000_000:
        return "$5–10B"
    if usd < 50_000_000_000:
        return "$10–50B"
    return "$50B+"


def _parse_employees(text: str) -> int | None:
    if not text:
        return None
    patterns = [
        r"([\d,]+(?:\.\d+)?)\s*(?:thousand|k)\+?\s*(?:employees|staff|people|workers)",
        r"([\d,]+)\s*\+?\s*(?:employees|staff|people|workers)",
        r"(?:employees|staff|headcount|number of employees)[:\s]+([\d,]+)",
        r"employs\s+(?:about\s+|approximately\s+|over\s+|more than\s+)?([\d,]+)",
        r"workforce of\s+([\d,]+)",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if not m:
            continue
        raw = m.group(1).replace(",", "")
        try:
            n = float(raw)
        except ValueError:
            continue
        if "thousand" in (m.group(0).lower()) or re.search(r"\bk\b", m.group(0), re.I):
            n *= 1000
        n = int(n)
        if 1 <= n <= 2_000_000 and n not in range(1900, 2031):
            return n
    return None


def _parse_revenue_usd(text: str) -> float | None:
    if not text:
        return None
    patterns = [
        r"(?:revenue|turnover|sales|net sales)[^.]{0,40}?(\$|€|£|usd|eur|gbp)?\s*([\d.,]+)\s*(billion|million|bn|mn|b|m)\b",
        r"(\$|€|£)\s*([\d.,]+)\s*(billion|million|bn|mn|b|m)\b[^.]{0,30}(?:revenue|turnover|sales)",
        r"([\d.,]+)\s*(billion|million)\s+(?:dollars|usd|euros)?[^.]{0,20}(?:revenue|turnover|sales)",
    ]
    best = None
    for pat in patterns:
        for m in re.finditer(pat, text, re.I):
            groups = m.groups()
            # find number + unit
            num = None
            unit = None
            cur = None
            for g in groups:
                if g is None:
                    continue
                gl = g.lower()
                if gl in ("$", "€", "£", "usd", "eur", "gbp"):
                    cur = gl
                elif gl in ("billion", "million", "bn", "mn", "b", "m"):
                    unit = gl
                elif re.match(r"^[\d.,]+$", g):
                    num = g
            if not num or not unit:
                continue
            try:
                n = float(num.replace(",", ""))
            except ValueError:
                continue
            if unit in ("billion", "bn", "b"):
                usd = n * 1_000_000_000
            else:
                usd = n * 1_000_000
            # crude FX
            if cur in ("€", "eur"):
                usd *= 1.08
            elif cur in ("£", "gbp"):
                usd *= 1.27
            if best is None or usd > best:
                best = usd
    return best


def _emp_from_cell(v: str) -> str:
    if _emptyish(v):
        return ""
    s = v.strip()
    for r in EMP_RANGES:
        if _norm(s).replace(",", "").replace("-", "–") == _norm(r).replace(",", ""):
            return r
    if re.fullmatch(r"20\d{2}", s):
        return ""
    s2 = s.replace(",", "")
    m = re.search(r"(\d+)\s*[-–to]+\s*(\d+)", s2, re.I)
    if m:
        mid = (int(m.group(1)) + int(m.group(2))) // 2
        return _emp_num_to_range(mid)
    m = re.search(r"(\d+)\+?", s2)
    if m:
        n = int(m.group(1))
        if n in range(1900, 2031):
            return ""
        return _emp_num_to_range(n)
    return ""


def _to_from_cell(v: str) -> str:
    if _emptyish(v):
        return ""
    s = v.strip()
    for r in TO_RANGES:
        if _norm(s).replace(" ", "") == _norm(r).replace(" ", ""):
            return r
    if re.fullmatch(r"20\d{2}", s) or re.fullmatch(r"\d{1,4}", s):
        # bare small ints are junk for turnover
        return ""
    usd = _parse_revenue_usd(s)
    if usd:
        return _usd_to_range(usd)
    # try $X–Y patterns already
    m = re.search(r"\$?\s*([\d.]+)\s*[-–]\s*([\d.]+)\s*([mb])", s, re.I)
    if m:
        a, b, u = float(m.group(1)), float(m.group(2)), m.group(3).lower()
        mid = (a + b) / 2
        usd = mid * (1_000_000_000 if u == "b" else 1_000_000)
        return _usd_to_range(usd)
    return ""


def wiki_fetch(name: str) -> str:
    """Fetch Wikipedia extract text (English)."""
    try:
        q = urllib.parse.quote(name)
        url = (
            "https://en.wikipedia.org/w/api.php?action=query&prop=extracts"
            f"&exintro=0&explaintext=1&redirects=1&format=json&titles={q}"
        )
        req = urllib.request.Request(
            url,
            headers={"User-Agent": "VendorIntelScaleBot/1.0 (research)"},
        )
        with urllib.request.urlopen(req, timeout=12) as resp:
            data = json.loads(resp.read().decode("utf-8", errors="ignore"))
        pages = (data.get("query") or {}).get("pages") or {}
        texts = []
        for p in pages.values():
            if int(p.get("pageid") or 0) < 0:
                continue
            t = p.get("extract") or ""
            if t and len(t) > 40:
                texts.append(t[:5000])
        if texts:
            return "\n".join(texts)
        # search API fallback
        surl = (
            "https://en.wikipedia.org/w/api.php?action=opensearch&limit=1&namespace=0"
            f"&format=json&search={q}"
        )
        req2 = urllib.request.Request(
            surl, headers={"User-Agent": "VendorIntelScaleBot/1.0 (research)"}
        )
        with urllib.request.urlopen(req2, timeout=10) as resp:
            arr = json.loads(resp.read().decode("utf-8", errors="ignore"))
        if isinstance(arr, list) and len(arr) > 1 and arr[1]:
            return wiki_fetch(arr[1][0])
    except Exception:
        return ""
    return ""


def web_snippets(name: str, website: str) -> str:
    host = ""
    if website:
        host = re.sub(r"^https?://", "", website, flags=re.I).split("/")[0].replace("www.", "")
    queries = [
        f'"{name}" employees OR headcount OR "number of employees"',
        f'"{name}" revenue OR turnover OR "annual sales" OR billion OR million',
    ]
    if host:
        queries.append(f'"{name}" {host} employees OR revenue')
    parts: list[str] = []
    try:
        from ddgs import DDGS

        for q in queries[:2]:
            try:
                hits = list(DDGS().text(q, max_results=4)) or []
            except Exception:
                hits = []
            for h in hits:
                parts.append(f"{h.get('title', '')} {h.get('body', '')}")
    except Exception:
        pass
    return "\n".join(parts)[:6000]


def resolve_row(row: dict[str, str]) -> dict[str, Any]:
    name = row.get("Company") or ""
    known = KNOWN.get(_norm(name), {})
    emp = known.get("employees") or _emp_from_cell(row.get("Employees") or "")
    turn = known.get("turnover") or _to_from_cell(row.get("Turnover / Scale") or "")

    wiki = ""
    web = ""
    if not emp or not turn:
        wiki = wiki_fetch(name)
        blob = wiki
        if not emp:
            n = _parse_employees(blob)
            if n:
                emp = _emp_num_to_range(n)
        if not turn:
            usd = _parse_revenue_usd(blob)
            if usd:
                turn = _usd_to_range(usd)

    if not emp or not turn:
        web = web_snippets(name, row.get("Website") or "")
        blob = web
        if not emp:
            n = _parse_employees(blob)
            if n:
                emp = _emp_num_to_range(n)
        if not turn:
            usd = _parse_revenue_usd(blob)
            if usd:
                turn = _usd_to_range(usd)

    # final fallback: convert leftover exact employee cells even if tiny default
    if not emp:
        emp = _emp_from_cell(row.get("Employees") or "")
    if not turn:
        turn = _to_from_cell(row.get("Turnover / Scale") or "")

    return {
        "company": name,
        "employees": emp or NPD,
        "turnover": turn or NPD,
        "sources": ";".join(
            x for x in [
                "known" if known else "",
                "wikipedia" if wiki else "",
                "web" if web else "",
                "cell_map",
            ]
            if x
        ),
    }


def load_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers: list[str] = []
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            header_idx = i
            headers = vals
            break
    data = []
    for r in rows[header_idx + 1 :]:
        d = {
            headers[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
            for j in range(len(headers))
            if headers[j]
        }
        if d.get("Company"):
            data.append(d)
    return [h for h in headers if h], data


def write_xlsx(path: Path, headers: list[str], rows: list[dict[str, str]], title: str, meta: dict) -> None:
    wb = load_workbook(path)
    ws = wb[SHEET]
    hr = 2
    for r in range(1, 6):
        if ws.cell(r, 1).value == "Company":
            hr = r
            break
    if ws.max_row > hr:
        ws.delete_rows(hr + 1, ws.max_row - hr)
    for j, h in enumerate(headers, 1):
        ws.cell(hr, j).value = h
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    for i, row in enumerate(rows):
        rr = hr + 1 + i
        for j, h in enumerate(headers, 1):
            cell = ws.cell(rr, j, row.get(h, ""))
            cell.alignment = wrap
            cell.border = thin
    ws.cell(1, 1).value = title
    name = "Scale Fast Audit"
    if name in wb.sheetnames:
        del wb[name]
    aus = wb.create_sheet(name)
    aus.append(["field", "value"])
    for k, v in meta.items():
        aus.append([k, str(v)])
    wb.save(path)


def main() -> int:
    headers, rows = load_rows(FINAL)
    print(f"=== FAST scale fill (Wikipedia + web) ===", flush=True)
    print(f"  companies: {len(rows)}  workers={WORKERS}", flush=True)

    results: dict[str, dict[str, Any]] = {}
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = {ex.submit(resolve_row, row): row for row in rows}
        done = 0
        for fut in as_completed(futs):
            row = futs[fut]
            done += 1
            try:
                res = fut.result()
            except Exception as err:
                res = {
                    "company": row.get("Company"),
                    "employees": _emp_from_cell(row.get("Employees") or "") or NPD,
                    "turnover": _to_from_cell(row.get("Turnover / Scale") or "") or NPD,
                    "sources": f"error:{err}",
                }
            results[_norm(res["company"])] = res
            if done % 20 == 0 or done == len(rows):
                print(f"  progress {done}/{len(rows)}", flush=True)

    out = []
    audits = []
    for row in rows:
        res = results[_norm(row.get("Company") or "")]
        fixed = dict(row)
        fixed["Employees"] = res["employees"]
        fixed["Turnover / Scale"] = res["turnover"]
        src = fixed.get("Data Sources") or ""
        if "wiki_web_scale" not in src:
            fixed["Data Sources"] = (src + "; wiki_web_scale").strip("; ")
        out.append(fixed)
        audits.append(res)

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    bak = FINAL.with_suffix(FINAL.suffix + f".bak_scale_fast_{ts}")
    shutil.copy2(FINAL, bak)
    emp_c = Counter(r["Employees"] for r in out)
    to_c = Counter(r["Turnover / Scale"] for r in out)
    audit_path = FINAL.parent / f"smartphone_scale_fast_{ts}.json"
    audit_path.write_text(
        json.dumps(
            {
                "ts": ts,
                "employees_counts": dict(emp_c),
                "turnover_counts": dict(to_c),
                "audits": audits,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    write_xlsx(
        FINAL,
        headers,
        out,
        f"Global Smartphone Market — channel partners (scale ranges fast {ts})",
        {
            "ts": ts,
            "method": "wikipedia+web+known+cell_map",
            "rows": len(out),
            "employees_counts": dict(emp_c),
            "turnover_counts": dict(to_c),
            "backup": str(bak),
            "audit": str(audit_path),
        },
    )
    print("=== DONE ===", flush=True)
    print(f"  backup: {bak.name}", flush=True)
    print(f"  Emp:", dict(emp_c.most_common(12)), flush=True)
    print(f"  TO :", dict(to_c.most_common(12)), flush=True)
    # samples
    for want in ("Amazon", "ALSO Holding", "AB SA", "Best Buy", "CDW", "3G Mobile", "Brightstar"):
        r = results.get(_norm(want))
        if r:
            print(f"  {want}: {r['employees']} | {r['turnover']} ({r['sources']})", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
