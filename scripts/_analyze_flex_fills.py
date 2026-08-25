"""Analyze whether firmographic fields look templated/hallucinated."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
PROBE = FP.parent / "_website_probe.tsv"

wb = load_workbook(FP, read_only=True, data_only=True)
ws = wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
rows = list(ws.iter_rows(min_row=2, values_only=True))
wb.close()

for col_name in ["Founded", "Ownership", "Employees", "Core Categories", "Specialty Focus", "Distribution Type", "Retail / E-commerce / Both"]:
    i = headers.index(col_name)
    c = Counter(str(r[i] or "") for r in rows)
    print(f"\n=== {col_name} top 15 ===")
    for val, n in c.most_common(15):
        print(f"  {n:4d} | {val[:120]!r}")

# join probe status
status = {}
for line in PROBE.read_text(encoding="utf-8").splitlines()[1:]:
    parts = line.split("\t")
    if len(parts) >= 4:
        status[int(parts[0])] = parts[3]

ok_rows = []
bad_rows = []
for idx, r in enumerate(rows, start=2):
    st = status.get(idx, "")
    if st.startswith("http_2") or st.startswith("http_3"):
        ok_rows.append(r)
    else:
        bad_rows.append((idx, r[0], r[1], st))

print(f"\nOK_ROWS {len(ok_rows)} BAD_OR_FAIL {len(bad_rows)}")
print("\nOK company+website sample (first 40):")
for r in ok_rows[:40]:
    print(f"  {r[0]} | {r[1]}")

# fabricated name patterns
import re
fab_re = re.compile(
    r"^(Embalagens|Empaques|Embalajes|Plásticos|Canadian|Vietnamese|Thai|Singapore|Philippine|"
    r"Turkish|Egyptian|Algerian|Accra|Dar es Salaam|Cape Packaging|Casablanca|"
    r"Flexible Packaging|Packaging (Direct|Distributors|Plus|Products|Solutions|Wholesale|World)|"
    r"Bunzl|Poly (Pak|Packaging)|Polytex)",
    re.I,
)
fab = [r for r in rows if fab_re.search(str(r[0] or ""))]
print(f"\nPATTERN_SUSPECT_NAMES {len(fab)}")
