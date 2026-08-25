# -*- coding: utf-8 -*-
"""Pass 3: apply verified cores for remaining weak rows; blank unverified weak templates."""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"

HAND = {
    "Anda Pharmaceuticals": (
        "Pharmaceuticals; Generics; Specialty Pharmaceuticals; Medical Supplies; OTC",
        "Specialty pharmaceutical distributor supplying branded, generic, biologic, specialty drugs, diagnostics, and medical supplies to pharmacies, clinics, hospitals, and physician customers.",
    ),
    "Medhold": (
        "Medical Devices; Surgical Systems; Theatre & ICU Equipment; Point-of-Care Diagnostics",
        "Independent medical device distributor supplying surgical technology, theatre and ICU capital equipment, orthopaedics, infusion therapy, cardiology, maternity, infection control, and related clinical solutions to healthcare providers.",
    ),
    "Shenzhen Biocare Bio-Medical Equipment": (
        "ECG Systems; Patient Monitors; Fetal Monitors; Diagnostic Ultrasound",
        "Manufactures and supplies diagnostic instruments including digital ECG systems, multi-parameter patient monitors, fetal monitors, and diagnostic ultrasound systems for clinical use.",
    ),
    "Shenzhen XFT Medical": (
        "Neurorehabilitation Devices; EMG Biofeedback; Functional Electrical Stimulation; Wearable Rehab Systems",
        "Develops neurological rehabilitation solutions including wearable EMG-triggered functional electrical stimulation systems, biofeedback devices, and robotic glove training systems for post-stroke and neuromuscular recovery.",
    ),
    "Zephyr Technology": (
        "Wearable Physiological Monitoring; ECG & Respiration Sensors; Performance Monitoring Software",
        "Develops wearable physiological status monitoring systems such as BioHarness/BioModule sensors and OmniSense software that track ECG, heart rate, respiration, activity, and related metrics for training, occupational, and remote monitoring use.",
    ),
    "AED Superstore": (
        "AEDs; AED Accessories; CPR Training Products; First Aid & EMS Supplies",
        "Specialty reseller distributing automated external defibrillators, electrode pads and batteries, CPR manikins, first aid products, and AED program management solutions.",
    ),
    "AED USA": (
        "AEDs; AED Accessories; AED Compliance Management; CPR/AED Training",
        "Provides automated external defibrillators and accessories with compliance management, medical direction support, and training programs for sudden cardiac arrest preparedness.",
    ),
    "Response Ready": (
        "AEDs; AED Accessories; First Aid & Bleed Kits; Emergency Preparedness Training",
        "Supplies automated external defibrillators, accessories, first aid and bleed kits, plus training and maintenance programs for cardiac emergency preparedness.",
    ),
    "HSI Emergency Care": (
        "CPR/AED/First Aid Training; Emergency Care Education; Training Supplies",
        "Provides CPR, AED, and first aid training programs and related instructional supplies for workplace and professional emergency care certification.",
    ),
    "Vitality Medical": (
        "Homecare Medical Supplies; Incontinence & Ostomy; Wound Care; Respiratory Supplies",
        "Online distributor of medical supplies and equipment for consumers and professionals, covering incontinence, ostomy, wound care, respiratory, and related homecare product categories.",
    ),
    "SpinLife": (
        "Power Wheelchairs; Mobility Scooters; Manual Wheelchairs; Accessibility Equipment",
        "Retail division of Numotion specializing in durable medical equipment including power and manual wheelchairs, scooters, ramps, stair lifts, and related mobility products.",
    ),
    "Claflin Company": (
        "Medical-Surgical Distribution; Healthcare Logistics; Physician & Acute Care Supplies",
        "Independent medical distributor and logistics provider supplying medical-surgical products and supply-chain services to healthcare customers.",
    ),
    "Praxisdienst": (
        "Practice Medical Supplies; Diagnostic Devices; Clinical Consumables; Veterinary & Dental Supplies",
        "Online mail-order supplier of medical supplies and practice equipment for human, dental, and veterinary professionals, including diagnostic devices and clinical consumables.",
    ),
    "Buffalo Supply": (
        "Medical Supplies; Hospital Consumables; Healthcare Distribution",
        "Distributor of medical supplies and related healthcare products to institutional and professional customers.",
    ),
    "EMS Supply Co.": (
        "EMS Supplies; Emergency Medical Equipment; First Responder Products",
        "Supplier of emergency medical service equipment and first-responder consumables for prehospital care teams.",
    ),
}

WEAK = {
    "Pharmaceuticals; Medical Supplies; Healthcare Distribution",
    "Medical Devices; Specialty Healthcare Products; Clinical Supplies",
    "Medical Supplies; Healthcare Products; Specialty Retail",
}

GENERIC_FOCUS = (
    "broadline",
    "specialty distributor supplying medical devices and related clinical",
    "specialty reseller of medical and healthcare products",
    "broadline healthcare distributor supplying medicines",
)


def clean_spec(t: str) -> str:
    t = re.sub(r"\s+", " ", (t or "").strip())
    if not t:
        return ""
    t = re.sub(r"[,;]?\s*others\.?$", "", t, flags=re.I).strip()
    t = re.sub(
        r"\s*(?:across|throughout|in|within|for|serving|covering)\s+(?:the\s+)?[A-Za-z ,&\-]{2,50}\.?$",
        "",
        t,
        flags=re.I,
    ).strip()
    t = re.sub(
        r"\s+(?:United States|United Kingdom|Saudi Arabia|South Africa|Mexico|Canada|"
        r"India|China|Japan|Brazil|Germany|France|Spain|Italy|Australia|Korea|"
        r"Europe|Asia|Africa|Middle East)\.?$",
        ".",
        t,
        flags=re.I,
    )
    if t and t[-1] not in ".!?":
        t += "."
    return t.strip()


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = ROOT / f"global_wearable_medical_devices_market_global_FINAL_BEFORE_CORE_SPEC_PASS3_{stamp}.xlsx"
    shutil.copy2(XLSX, bak)

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers) if h}

    filled = blanked = 0
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        core = str(ws.cell(r, ic["Core Categories"]).value or "").strip()
        if co in HAND:
            c, f = HAND[co]
            ws.cell(r, ic["Core Categories"]).value = c
            ws.cell(r, ic["Specialty Focus"]).value = clean_spec(f)
            filled += 1
        elif core in WEAK:
            ws.cell(r, ic["Core Categories"]).value = ""
            sp = str(ws.cell(r, ic["Specialty Focus"]).value or "").strip()
            low = sp.lower()
            if len(sp) < 50 or any(g in low for g in GENERIC_FOCUS):
                ws.cell(r, ic["Specialty Focus"]).value = ""
            else:
                ws.cell(r, ic["Specialty Focus"]).value = clean_spec(sp)
            blanked += 1

    for r in range(2, ws.max_row + 1):
        ws.cell(r, ic["Specialty Focus"]).value = clean_spec(
            str(ws.cell(r, ic["Specialty Focus"]).value or "")
        )

    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    hf = PatternFill("solid", fgColor="1F4E79")
    for i in range(1, len(headers) + 1):
        cell = ws.cell(1, i)
        cell.fill = hf
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
    for r in range(2, ws.max_row + 1):
        for i in range(1, len(headers) + 1):
            cell = ws.cell(r, i)
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin

    wb.save(XLSX)
    empty_c = empty_s = weak = 0
    for r in range(2, ws.max_row + 1):
        c = str(ws.cell(r, ic["Core Categories"]).value or "").strip()
        s = str(ws.cell(r, ic["Specialty Focus"]).value or "").strip()
        if not c:
            empty_c += 1
        if not s:
            empty_s += 1
        if c in WEAK:
            weak += 1
    print("backup", bak.name)
    print("filled", filled, "blanked_weak", blanked)
    print("empty_core", empty_c, "empty_spec", empty_s, "weak_left", weak)


if __name__ == "__main__":
    main()
