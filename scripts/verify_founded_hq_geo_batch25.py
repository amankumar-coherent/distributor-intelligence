#!/usr/bin/env python3
"""Fast batch verify Founded + Headquarters + Continent/Geography via Google AI.

Batches 25 companies per Google AI query using:
  \"{company} founded year, headquarter and all region served mention all country\"

Adds/updates column: Continent / Geography (all countries served).
Updates Final SLIM Excels only. Checkpointed for resume.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")
os.environ["GOOGLE_AI_SCRAPER_ENABLED"] = "true"
os.environ["GOOGLE_AI_SCRAPER_FORCE_OVERVIEW"] = "true"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRIES"] = "2"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRY_SLEEP_SEC"] = "1.2"

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

SLIM_FILES = [
    ROOT / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL_SLIM.xlsx",
    ROOT / "output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL_SLIM.xlsx",
    ROOT / "output/run4/global_firewall_market_global/global_firewall_market_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_antibody_drug_conjugates_market_global/global_antibody_drug_conjugates_market_FINAL_SLIM.xlsx",
]

BATCH_SIZE = 25
GEO_COL = "Continent / Geography"
PROGRESS = ROOT / "scripts/_founded_hq_geo_progress.jsonl"
RESULTS = ROOT / "scripts/_founded_hq_geo_results.json"
CURRENT_YEAR = 2026


def norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def detect_header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [norm(c.value) for c in ws[r] if c.value not in (None, "")]
        if "company" in vals:
            return r
    return 1


def collect_rows() -> list[dict]:
    out: list[dict] = []
    for f in SLIM_FILES:
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        hr = detect_header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        fi = headers.index("Founded") if "Founded" in headers else None
        hi = headers.index("Headquarters") if "Headquarters" in headers else None
        wi = headers.index("Website") if "Website" in headers else None
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if not row or not row[ci]:
                continue
            company = str(row[ci]).strip()
            out.append(
                {
                    "file": str(f),
                    "company": company,
                    "founded": str(row[fi] or "").strip() if fi is not None else "",
                    "hq": str(row[hi] or "").strip() if hi is not None else "",
                    "website": str(row[wi] or "").strip() if wi is not None else "",
                    "key": f"{norm(f.name)}::{norm(company)}",
                }
            )
        wb.close()
    return out


def load_done() -> set[str]:
    done: set[str] = set()
    if PROGRESS.exists():
        for line in PROGRESS.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
                if obj.get("key") and obj.get("ok"):
                    done.add(obj["key"])
            except Exception:
                pass
    return done


def build_batch_query(batch: list[dict]) -> str:
    lines = [
        "For EACH company below, answer in clear structure:",
        "1) Founded year (YYYY only)",
        "2) Headquarters city and country",
        "3) ALL countries / regions served — list EVERY country name, even if more than 50. Do not summarize as 'worldwide' only; expand to country list when known. If truly global with unknown full list, say Global and list major countries known.",
        "",
        "Use this exact style of question for each:",
        '"<company> founded year, headquarter and all region served mention all country"',
        "",
        "Companies:",
    ]
    for i, item in enumerate(batch, 1):
        extra = []
        if item.get("website"):
            extra.append(f"website={item['website']}")
        if item.get("hq"):
            extra.append(f"current_hq={item['hq']}")
        if item.get("founded"):
            extra.append(f"current_founded={item['founded']}")
        suffix = f" ({'; '.join(extra)})" if extra else ""
        lines.append(f"{i}. {item['company']}{suffix}")
    lines.append("")
    lines.append("Return facts only. Do not invent countries.")
    return "\n".join(lines)


def deepseek_parse(batch: list[dict], markdown: str) -> list[dict]:
    api_key = os.getenv("OPENAI_API_KEY") or ""
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    client = OpenAI(api_key=api_key, base_url=base)
    names = [b["company"] for b in batch]
    prompt = f"""Extract structured firmographic data from Google AI Overview text.

Companies in this batch (keep these exact names):
{json.dumps(names, ensure_ascii=False)}

Google AI text:
```
{(markdown or '')[:12000]}
```

Return JSON array ONLY, one object per company that appears in the text (skip if absent):
[
  {{
    "company": "exact name from list",
    "founded_year": 2019 or null,
    "headquarters": "City, Country" or null,
    "countries_served": ["Country1", "Country2", ...]  // ALL country names mentioned; expand regions to countries only if explicitly listed; keep as given if already countries; if text says Global/Worldwide with no list use ["Global"]
  }}
]
Rules:
- founded_year must be 1800-{CURRENT_YEAR} or null. Never use copyright year alone.
- Do not invent. Prefer null over guess.
- countries_served: list every country named; no duplicates; preserve order.
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    text = (resp.choices[0].message.content or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I | re.M).strip()
    try:
        data = json.loads(text)
        if isinstance(data, dict) and "companies" in data:
            data = data["companies"]
        if not isinstance(data, list):
            return []
        return [x for x in data if isinstance(x, dict)]
    except Exception:
        return []


def normalize_countries(vals: object) -> str:
    if vals is None:
        return ""
    if isinstance(vals, str):
        parts = re.split(r"[,;/|]|\band\b", vals)
        items = [p.strip() for p in parts if p.strip()]
    elif isinstance(vals, list):
        items = [str(x).strip() for x in vals if str(x).strip()]
    else:
        return ""
    seen = set()
    out = []
    for it in items:
        key = norm(it)
        if key in seen:
            continue
        seen.add(key)
        out.append(it)
    return "; ".join(out)


def year_ok(y: object) -> int | None:
    try:
        yi = int(y)
    except Exception:
        m = re.search(r"(18|19|20)\d{2}", str(y or ""))
        yi = int(m.group(0)) if m else None
    if yi is None:
        return None
    if 1800 <= yi <= CURRENT_YEAR:
        return yi
    return None


async def process_batch(batch: list[dict]) -> list[dict]:
    query = build_batch_query(batch)
    # Also append one example single-style line to nudge AI Overview
    sample = batch[0]["company"]
    query_alt = (
        f"{sample} founded year, headquarter and all region served mention all country. "
        f"Then do the same for these {len(batch)} companies:\n"
        + "\n".join(f"- {b['company']}" for b in batch)
    )
    data = await google_ai_ask(query, llm_clean=True)
    md = str(data.get("markdown") or data.get("markdown_raw") or "").strip()
    if len(md) < 120:
        data2 = await google_ai_ask(query_alt, llm_clean=True)
        md2 = str(data2.get("markdown") or data2.get("markdown_raw") or "").strip()
        if len(md2) > len(md):
            md = md2
            data = data2

    parsed = deepseek_parse(batch, md) if md else []
    by_name = {norm(p.get("company")): p for p in parsed if p.get("company")}

    missing = [item for item in batch if norm(item["company"]) not in by_name]
    if missing:
        # One re-ask for misses only (still batched, not 1-by-1)
        q_miss = build_batch_query(missing)
        d_miss = await google_ai_ask(q_miss, llm_clean=True)
        md_miss = str(d_miss.get("markdown") or d_miss.get("markdown_raw") or "").strip()
        if md_miss:
            for p in deepseek_parse(missing, md_miss):
                if p.get("company"):
                    by_name[norm(p["company"])] = p
            # keep richest markdown for snippet
            if len(md_miss) > len(md):
                md = md_miss

    # Last resort: single asks only for still-missing (cap 5 per batch to stay fast)
    still = [item for item in batch if norm(item["company"]) not in by_name]
    for item in still[:5]:
        q1 = (
            f"{item['company']} founded year, headquarter and all region served "
            f"mention all country"
        )
        d1 = await google_ai_ask(q1, llm_clean=True)
        md1 = str(d1.get("markdown") or d1.get("markdown_raw") or "").strip()
        if md1:
            parsed1 = deepseek_parse([item], md1)
            if parsed1 and parsed1[0].get("company"):
                by_name[norm(item["company"])] = parsed1[0]

    results: list[dict] = []
    for item in batch:
        p = by_name.get(norm(item["company"]), {})
        founded = year_ok(p.get("founded_year"))
        hq = str(p.get("headquarters") or "").strip() or None
        geo = normalize_countries(p.get("countries_served"))
        ok = bool(founded or hq or geo)
        results.append(
            {
                "key": item["key"],
                "file": item["file"],
                "company": item["company"],
                "old_founded": item["founded"],
                "old_hq": item["hq"],
                "founded": str(founded) if founded else None,
                "headquarters": hq,
                "continent_geography": geo or None,
                "ok": ok,
                "snippet": (md[:300] if md else ""),
            }
        )
    return results


def ensure_geo_column(ws, header_row: int, headers: list[str]) -> tuple[list[str], int]:
    """Ensure Continent / Geography exists; return headers and 0-based geo index."""
    if GEO_COL in headers:
        return headers, headers.index(GEO_COL)

    # insert after Headquarters if present, else after Founded, else at end
    if "Headquarters" in headers:
        insert_at = headers.index("Headquarters") + 1
    elif "Founded" in headers:
        insert_at = headers.index("Founded") + 1
    else:
        insert_at = len(headers)

    # shift cells right from insert_at
    max_col = ws.max_column
    for r in range(header_row, ws.max_row + 1):
        for c in range(max_col, insert_at, -1):
            ws.cell(r, c + 1).value = ws.cell(r, c).value
        ws.cell(r, insert_at + 1).value = None

    headers = headers[:insert_at] + [GEO_COL] + headers[insert_at:]
    cell = ws.cell(header_row, insert_at + 1, GEO_COL)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(insert_at + 1)].width = 40
    return headers, insert_at


def apply_all(results: list[dict]) -> None:
    by_file: dict[str, dict[str, dict]] = {}
    for fx in results:
        if not fx.get("ok"):
            continue
        by_file.setdefault(fx["file"], {})[norm(fx["company"])] = fx

    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for fpath, mapping in by_file.items():
        wb = load_workbook(fpath)
        ws = wb.active
        hr = detect_header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        headers, geo_i = ensure_geo_column(ws, hr, headers)
        # refresh headers after possible insert
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        fi = headers.index("Founded") if "Founded" in headers else None
        hi = headers.index("Headquarters") if "Headquarters" in headers else None
        gi = headers.index(GEO_COL)

        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            company = ws.cell(r, ci + 1).value
            if not company:
                continue
            fx = mapping.get(norm(str(company)))
            if not fx:
                continue
            if fi is not None and fx.get("founded"):
                old = str(ws.cell(r, fi + 1).value or "").strip()
                # only overwrite if new year and (old empty/NPD/different)
                if not old or old.upper() == "NPD" or old != fx["founded"]:
                    # protect huge deltas unless old NPD
                    om = re.search(r"(18|19|20)\d{2}", old)
                    if om and old.upper() != "NPD":
                        oy = int(om.group(0))
                        ny = int(fx["founded"])
                        if abs(oy - ny) > 100:
                            pass
                        else:
                            ws.cell(r, fi + 1).value = fx["founded"]
                    else:
                        ws.cell(r, fi + 1).value = fx["founded"]
            if hi is not None and fx.get("headquarters"):
                old_hq = str(ws.cell(r, hi + 1).value or "").strip()
                if not old_hq or old_hq.upper() == "NPD" or len(fx["headquarters"]) >= 5:
                    # update HQ when we have a solid answer
                    if fx["headquarters"] and fx["headquarters"].lower() not in {"n/a", "npd", "unknown"}:
                        ws.cell(r, hi + 1).value = fx["headquarters"]
            if fx.get("continent_geography"):
                cell = ws.cell(r, gi + 1, fx["continent_geography"])
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = thin
            n += 1
        wb.save(fpath)
        print(f"Updated {n} rows -> {Path(fpath).name}")


async def main() -> None:
    if not await google_ai_ping():
        raise SystemExit("Google AI scraper not reachable on :15551")

    rows = collect_rows()
    done = load_done()
    todo = [r for r in rows if r["key"] not in done]
    print(f"Companies: {len(rows)} | done: {len(done)} | remaining: {len(todo)}")
    print(f"Batch size: {BATCH_SIZE} | batches left: {(len(todo) + BATCH_SIZE - 1) // BATCH_SIZE}")

    PROGRESS.parent.mkdir(parents=True, exist_ok=True)
    all_new: list[dict] = []

    with PROGRESS.open("a", encoding="utf-8") as prog:
        for bi in range(0, len(todo), BATCH_SIZE):
            batch = todo[bi : bi + BATCH_SIZE]
            bnum = bi // BATCH_SIZE + 1
            btot = (len(todo) + BATCH_SIZE - 1) // BATCH_SIZE
            print(f"\n=== Batch {bnum}/{btot} ({len(batch)} companies) ===", flush=True)
            for b in batch:
                print(f"  - {b['company']}", flush=True)
            try:
                results = await process_batch(batch)
            except Exception as e:
                print(f"  BATCH ERROR: {e}", flush=True)
                results = [
                    {
                        "key": x["key"],
                        "file": x["file"],
                        "company": x["company"],
                        "ok": False,
                        "error": str(e),
                    }
                    for x in batch
                ]
            for fx in results:
                prog.write(json.dumps(fx, ensure_ascii=False) + "\n")
                prog.flush()
                all_new.append(fx)
                print(
                    f"  -> {fx['company']}: founded={fx.get('founded')} | hq={fx.get('headquarters')} | "
                    f"geo={str(fx.get('continent_geography') or '')[:80]} | ok={fx.get('ok')}",
                    flush=True,
                )
            await asyncio.sleep(0.5)

    # merge all progress for apply
    merged: list[dict] = []
    if PROGRESS.exists():
        for line in PROGRESS.read_text(encoding="utf-8").splitlines():
            if line.strip():
                merged.append(json.loads(line))
    RESULTS.write_text(json.dumps(merged, indent=2, ensure_ascii=False), encoding="utf-8")
    print("\nApplying updates to SLIM Excels...")
    apply_all(merged)
    ok_n = sum(1 for x in merged if x.get("ok"))
    print(f"Done. ok={ok_n}/{len(merged)} | results={RESULTS}")


if __name__ == "__main__":
    asyncio.run(main())
