#!/usr/bin/env python3
"""Fill Key Brands via official-site fetch + DeepSeek only (no Google AI scraper)."""
from __future__ import annotations

import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import openpyxl
import requests
from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")

XLSX = ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_OFFICE_SLIM.xlsx"
CACHE = ROOT / "output/web_expand/global_avocado_oil_market_global/_avocado_brands_ds_web_cache.json"
AUDIT = ROOT / "output/web_expand/global_avocado_oil_market_global/_avocado_brands_ds_web_audit.json"
NPD = "Not publicly disclosed"

# Agent web-verified overrides (this session + prior)
VERIFIED = {
    "Labatt Food Service": NPD,  # no private labels by policy
    "Dole Europe": "Dole, others",
    "Nisshin OilliO Distribution": "Nisshin, Healthy Resetta, others",
    "Sime Darby Oils Distribution": "Alif, Bagus, Cap Bendera, Chief, Hand Brand, Sunbeam, Crispa Gold, Leila, Delico, NATRIEO, others",
    "The Olive Oil Company": "Bertolli, Carapelli, Carbonell, Hojiblanca, Koipe, Figaro, others",
    "Tree of Life": "CADIA, MADE WITH, others",
    "Spinneys Distribution": "SpinneysFOOD, SpinneysHOME, SpinneysWELLNESS, Fine Food, Waitrose, Discovery by Spinneys, others",
    "YHS Distribution": "Yeo's, others",
    "Mina": "Mina, others",
    "Naissance": "Naissance, others",
    "Horizon Harvest": NPD,
    "The Bulk Cart": "The Bulk Cart, others",
    "Tradin Organic": "Tradin Organic, others",
    "Pro*Act": NPD,
    "Sid Wainer & Son": "Sid Wainer & Son, others",
    "P. Brändle": "Braendle, others",
    "P. Brändle GmbH": "Braendle, others",
    "Foodservice Direct": NPD,
    "Dot Foods": NPD,
    "Core-Mark": NPD,
    "Vistar": NPD,
    "Martin-Brower": NPD,
    "Reyes Holdings": NPD,
    "Matthew Clark": NPD,
    "Fresh Produce Consortium Members Distribution": NPD,
    "Gourmet Gulf Distribution": NPD,  # bad website in sheet
    "Food Supplies Company": NPD,  # bad tripadvisor website
    "Horeca Trade": NPD,  # bad website
    "Jan Oskam": NPD,  # bad website askjan.org
    "Scavuzzo's": NPD,  # bad website
    "Patties Food Group Distribution": NPD,  # bad website
    "GreenGold USA": NPD,  # greengold.tv not food
    "Woods Foodservice": "Spence, others",
    "Qualifirst Foods": "Qualifirst, Castelanotti, Oliveio, Royal Command, others",
}


def ensure_others(text: str) -> str:
    t = (text or "").strip().rstrip(",")
    if not t:
        return t
    if re.search(r",\s*others\s*$", t, re.I):
        return re.sub(r",\s*others\s*$", ", others", t, flags=re.I)
    return t + ", others"


def clean_brand_list(raw: str) -> str:
    if not raw:
        return ""
    parts = re.split(r"[,;/|]+", raw)
    out, seen = [], set()
    junk = {
        "others",
        "private label",
        "house brands",
        "various",
        "multiple",
        "n/a",
        "none",
        "unknown",
        "brands",
        "brand",
        "principal",
        "suppliers",
        "portfolio",
    }
    for p in parts:
        b = re.sub(r"\s+", " ", p).strip(" .")
        if not b or len(b) < 2 or b.lower() in junk:
            continue
        if re.search(r"principal|private label|house brand|not publicly|inventory|assortment", b, re.I):
            continue
        if len(b) > 60:
            continue
        key = b.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(b)
    if not out:
        return ""
    return ensure_others(", ".join(out[:12]))


def ds_client() -> tuple[OpenAI, str]:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    if not key:
        raise SystemExit("OPENAI_API_KEY missing")
    return OpenAI(api_key=key, base_url=base, timeout=90.0), model


def fetch_text(url: str, timeout: float = 20.0) -> str:
    if not url or not url.lower().startswith("http"):
        return ""
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; VendorIntelBot/1.0; +https://example.local)"
    }
    try:
        r = requests.get(url, headers=headers, timeout=timeout, allow_redirects=True)
        if r.status_code >= 400:
            return ""
        html = r.text or ""
    except Exception:
        return ""
    # strip scripts/styles
    html = re.sub(r"(?is)<(script|style|noscript).*?>.*?</\1>", " ", html)
    text = re.sub(r"(?is)<[^>]+>", " ", html)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:12000]


def candidate_urls(website: str) -> list[str]:
    if not website or not website.lower().startswith("http"):
        return []
    base = website.rstrip("/") + "/"
    paths = [
        "",
        "about",
        "about-us",
        "our-brands",
        "brands",
        "products",
        "private-label",
        "own-brand",
        "en/about",
        "en/brands",
    ]
    out = []
    for p in paths:
        out.append(urljoin(base, p))
    # dedupe preserve order
    seen, uniq = set(), []
    for u in out:
        if u not in seen:
            seen.add(u)
            uniq.append(u)
    return uniq[:6]


def gather_evidence(company: str, website: str) -> str:
    parts = [f"Company: {company}", f"Website: {website}"]
    for u in candidate_urls(website):
        txt = fetch_text(u)
        if len(txt) < 200:
            continue
        parts.append(f"SOURCE {u}:\n{txt[:3500]}")
        if sum(len(p) for p in parts) > 9000:
            break
        time.sleep(0.25)
    return "\n\n".join(parts)


def deepseek_extract(client: OpenAI, model: str, company: str, evidence: str) -> dict:
    system = (
        "Extract ONLY real consumer/trade brand names this company owns or officially "
        "distributes as exclusive/private-label/house brands. Prefer named brands shown "
        "in evidence. Do NOT invent. Do NOT return generic phrases. If evidence does not "
        "clearly name brands, found=false. Return ONLY JSON."
    )
    user = (
        f"Company: {company}\n\nEVIDENCE:\n{(evidence or '')[:9000]}\n\n"
        'JSON: {"found":true|false,"brands":["Brand A","Brand B"],'
        '"confidence":"high|medium|low","reason":"brief"}'
    )
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=350,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        raw = resp.choices[0].message.content or "{}"
        if raw.startswith("```"):
            raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
        data = json.loads(raw)
    except Exception as ex:
        return {"found": False, "error": str(ex)[:160]}

    conf = str(data.get("confidence") or "").lower()
    if not data.get("found") or conf not in {"high", "medium"}:
        return {"found": False, "confidence": conf, "raw": data}

    brands = data.get("brands") or []
    if isinstance(brands, str):
        brands = [brands]
    brands = [str(b).strip() for b in brands if str(b).strip()]
    evid_l = (evidence or "").lower()
    kept = []
    for b in brands:
        token = re.sub(r"[^a-z0-9]", "", b.lower())
        evid_tok = re.sub(r"[^a-z0-9]", "", evid_l)
        if len(token) >= 4 and token in evid_tok:
            kept.append(b)
        elif conf == "high" and any(
            w.lower() in evid_l for w in re.findall(r"[A-Za-z0-9][A-Za-z0-9&''\-]{2,}", b)
        ):
            kept.append(b)
    if not kept:
        return {"found": False, "confidence": conf, "reason": "brands not evidenced"}
    return {"found": True, "brands": clean_brand_list(", ".join(kept)), "confidence": conf}


def main() -> None:
    limit = None
    if len(sys.argv) > 1 and sys.argv[1].isdigit():
        limit = int(sys.argv[1])

    client, model = ds_client()
    cache: dict = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

    # Phase A verified
    applied = 0
    for row in range(3, ws.max_row + 1):
        co = str(ws.cell(row, ic["Company"]).value or "").strip()
        if co not in VERIFIED:
            continue
        ws.cell(row, ic["Key Brands Represented"]).value = VERIFIED[co]
        applied += 1
        print(f"verified {co}", flush=True)
    wb.save(XLSX)
    print(f"phaseA verified={applied}", flush=True)

    gaps = []
    for row in range(3, ws.max_row + 1):
        co = str(ws.cell(row, ic["Company"]).value or "").strip()
        web = str(ws.cell(row, ic["Website"]).value or "").strip()
        brands = str(ws.cell(row, ic["Key Brands Represented"]).value or "").strip()
        if brands in ("", NPD):
            gaps.append({"row": row, "company": co, "website": web})
    if limit:
        gaps = gaps[:limit]
    print(f"phaseB gaps={len(gaps)} model={model}", flush=True)

    filled = 0
    for i, g in enumerate(gaps, 1):
        key = f"dsweb::{g['company']}"
        if key in cache and cache[key].get("done"):
            hit = cache[key]
            if hit.get("brands"):
                ws.cell(g["row"], ic["Key Brands Represented"]).value = hit["brands"]
                filled += 1
            print(f"[{i}/{len(gaps)}] cache {g['company'][:42]}", flush=True)
            continue

        print(f"[{i}/{len(gaps)}] DS+WEB {g['company'][:42]}", flush=True)
        evidence = gather_evidence(g["company"], g["website"])
        if len(evidence) < 120:
            hit = {"found": False, "reason": "no_evidence"}
        else:
            hit = deepseek_extract(client, model, g["company"], evidence)

        cache[key] = {"done": True, "ts": datetime.now(timezone.utc).isoformat(), **hit}
        if hit.get("found") and hit.get("brands"):
            ws.cell(g["row"], ic["Key Brands Represented"]).value = hit["brands"]
            filled += 1
            print(f"  -> {hit['brands'][:110]}", flush=True)
        else:
            print(f"  -> NPD ({hit.get('reason') or hit.get('error') or hit.get('confidence')})", flush=True)

        if i % 8 == 0:
            CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
            wb.save(XLSX)
        time.sleep(0.35)

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
    wb.save(XLSX)
    wb.close()

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr2 = [str(c or "") for c in rows[1]]
    ic2 = {h: i for i, h in enumerate(hdr2) if h}
    real = sum(
        1
        for r in rows[2:]
        if str(r[ic2["Company"]] or "").strip()
        and str(r[ic2["Key Brands Represented"]] or "").strip() not in ("", NPD)
    )
    total = sum(1 for r in rows[2:] if str(r[ic2["Company"]] or "").strip())
    wb.close()
    audit = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "phaseA_verified": applied,
        "phaseB_attempted": len(gaps),
        "phaseB_filled": filled,
        "total_real_brands": real,
        "total_rows": total,
        "model": model,
        "source": "agent_verified+deepseek_website_only",
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2), flush=True)


if __name__ == "__main__":
    main()
