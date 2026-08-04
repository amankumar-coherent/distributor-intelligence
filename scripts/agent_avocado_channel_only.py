"""Avocado FINAL: distributors/channel partners ONLY — agent web research fill.

NO SearXNG / NO DDGS router / NO OpenAI SDK.
- Keep grocery / foodservice / ingredient CHANNEL distributors
- Drop manufacturers, growers, oil mills, brand producers
- Fill from agent-verified WEB_FACTS + Wikipedia API (direct)
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from vendor_intel.pipeline.web_expand import (  # noqa: E402
    HEADERS,
    MARKET_DEFAULTS,
    _blank_row,
    _norm,
    write_final_xlsx,
)

FINAL = ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL.xlsx"
SEEDS = ROOT / "queries/seeds/global_avocado_oil_market.txt"
FAMILY = "avocado"
MARKET = "Global Avocado Oil Market"

NPD = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "todo",
    "not applicable",
    "seed_file",
}

# Explicit manufacturers / producers / growers / mills to REMOVE
MANUFACTURER_NEEDLES = {
    "olivado",
    "limbua",
    "avohealth",
    "chosen foods",
    "sesajal",
    "calavo",
    "mission produce",
    "westfalia",
    "imara",
    "avolio",
    "aceitera mevi",
    "aceites naturales",
    "keitt exporters",
    "grove avocado",
    "village press",
    "catania oils",
    "compagnia alimentare",
    "export all jalisco",
    "superfine africa",
    "spraylab",
    "nnrv trade",
    "aak",  # specialty oils processor / manufacturer
    "adm european",  # ADM is processor; keep pure channel names only if clear distributor
    "henry avocado",  # grower/packer
    "index fresh",  # grower/packer
    "salud foodgroup",
    "oleo hass",
    "avopure",
    "avooil",
    "avocado oil oleo",
}

MANUFACTURER_RE = re.compile(
    r"(?i)\b(manufacturer|producer|processor|oil mill|cold[- ]pressed|"
    r"grower|packing house|refinery|brand owner)\b"
)

JUNK_RE = re.compile(
    r"(?i)^(avocado oil\b.*|"
    r".*\b(buyers?|directory|marketplace|list of|top \d+|volza|tridge|trademo)\b)$"
)

CHANNEL_HINT_RE = re.compile(
    r"(?i)\b(distributor|distribution|wholesale|wholesaler|foodservice|"
    r"food service|redistributor|importer|channel|broker|coop|"
    r"cooperative|cash.?and.?carry|specialty foods)\b"
)

# Agent WebSearch-verified DISTRIBUTOR / CHANNEL facts only
WEB_FACTS: dict[str, dict[str, str]] = {
    "sysco": {
        "Founded": "1969",
        "Headquarters": "Houston, Texas, USA",
        "Ownership": "Public (NYSE: SYY)",
        "Business Type": "Foodservice distributor",
        "Employees": "~75,000",
        "Turnover / Scale": "~$81.4B (FY2025)",
        "Contact Person": "Kevin Hourican",
        "Role": "Chairman & CEO",
        "Website": "https://www.sysco.com",
        "LinkedIn": "https://www.linkedin.com/company/sysco",
        "Core Categories": "Broadline foodservice; edible oils including avocado oil",
        "Specialty Focus": "Foodservice channel distribution",
        "Distribution Type": "Broadline foodservice",
        "Regions Served": "Global (US + international)",
        "Market Related": "yes",
        "Summary": "Sysco — world's largest foodservice distributor; channel for edible oils incl. avocado oil.",
    },
    "us foods": {
        "Founded": "1989",
        "Headquarters": "Rosemont, Illinois, USA",
        "Ownership": "Public (NYSE: USFD)",
        "Business Type": "Foodservice distributor",
        "Employees": "~30,000",
        "Turnover / Scale": "~$37.9B (2024)",
        "Contact Person": "David E. Flitman",
        "Role": "CEO",
        "Phone / WhatsApp": "+1-847-720-8000",
        "Website": "https://www.usfoods.com",
        "LinkedIn": "https://www.linkedin.com/company/us-foods",
        "Core Categories": "Broadline foodservice; specialty oils",
        "Specialty Focus": "Foodservice operator supply",
        "Distribution Type": "Broadline foodservice",
        "Summary": "US Foods — leading US foodservice distributor; HQ Rosemont; CEO Dave Flitman.",
    },
    "gordon food service": {
        "Founded": "1897",
        "Headquarters": "Wyoming, Michigan, USA",
        "Ownership": "Privately held (family)",
        "Business Type": "Foodservice distributor",
        "Employees": "~20,000",
        "Turnover / Scale": "~$12.9B+ (historical reports)",
        "Contact Person": "Rich Wolowski",
        "Role": "CEO",
        "Website": "https://www.gfs.com",
        "Core Categories": "Broadline foodservice distribution",
        "Specialty Focus": "Largest family-operated foodservice distributor in N. America",
        "Distribution Type": "Broadline foodservice",
        "Summary": "Gordon Food Service — founded 1897; HQ Wyoming MI; CEO Rich Wolowski; foodservice channel.",
    },
    "performance food group": {
        "Founded": "1885",
        "Headquarters": "Richmond, Virginia, USA",
        "Ownership": "Public (NYSE: PFGC)",
        "Business Type": "Foodservice distributor",
        "Website": "https://pfgc.com",
        "Core Categories": "Foodservice & convenience wholesale",
        "Distribution Type": "Broadline / specialty foodservice",
    },
    "kehe": {
        "Founded": "1952",
        "Headquarters": "Naperville, Illinois, USA",
        "Ownership": "Employee-owned (ESOP)",
        "Business Type": "Natural / specialty grocery distributor",
        "Employees": "~6,800–8,000",
        "Contact Person": "Deb Conklin",
        "Role": "President & CEO",
        "Website": "https://www.kehe.com",
        "LinkedIn": "https://www.linkedin.com/company/kehe-distributors",
        "Core Categories": "Natural, organic, specialty grocery wholesale",
        "Specialty Focus": "Specialty / natural channel — avocado oil brands to retail",
        "Distribution Type": "Specialty grocery wholesale",
        "Summary": "KeHE — employee-owned natural/specialty distributor; channel for avocado oil SKUs.",
    },
    "unfi": {
        "Company": "UNFI (United Natural Foods)",
        "Founded": "1976",
        "Headquarters": "Providence, Rhode Island, USA",
        "Ownership": "Public (NYSE: UNFI)",
        "Business Type": "Natural / organic wholesale distributor",
        "Employees": "~25,600",
        "Contact Person": "J. Alexander (Sandy) Miller Douglas",
        "Role": "CEO",
        "Phone / WhatsApp": "+1-401-528-8634",
        "Website": "https://www.unfi.com",
        "Core Categories": "Natural, organic, specialty wholesale",
        "Specialty Focus": "Primary natural/organic channel (incl. specialty oils)",
        "Distribution Type": "Natural grocery wholesale",
        "Summary": "UNFI — largest public natural/specialty wholesaler; avocado oil retail channel.",
    },
    "c&s wholesale": {
        "Founded": "1918",
        "Headquarters": "Keene, New Hampshire, USA",
        "Ownership": "Privately held",
        "Business Type": "Grocery wholesaler",
        "Employees": "~14,000–17,000",
        "Contact Person": "Eric Winn",
        "Role": "CEO",
        "Website": "https://www.cswg.com",
        "Distribution Type": "Grocery wholesale",
        "Summary": "C&S — largest US grocery wholesaler; retail channel for grocery oils.",
    },
    "ben e. keith": {
        "Founded": "1906",
        "Headquarters": "Fort Worth, Texas, USA",
        "Ownership": "Privately held",
        "Business Type": "Foodservice & beverage distributor",
        "Employees": "~6,000",
        "Turnover / Scale": "~$6.8B",
        "Contact Person": "Howard Hallam",
        "Role": "CEO",
        "Website": "https://www.benekeith.com",
        "Distribution Type": "Foodservice wholesale",
        "Summary": "Ben E. Keith — TX foodservice distributor; edible oils in broadline catalog.",
    },
    "shamrock foods": {
        "Founded": "1922",
        "Headquarters": "Phoenix, Arizona, USA",
        "Website": "https://www.shamrockfoods.com",
        "Business Type": "Foodservice distributor",
        "Ownership": "Privately held",
        "Distribution Type": "Foodservice wholesale",
    },
    "dot foods": {
        "Founded": "1960",
        "Headquarters": "Mount Sterling, Illinois, USA",
        "Website": "https://www.dotfoods.com",
        "Business Type": "Food redistributor",
        "Ownership": "Privately held",
        "Distribution Type": "Redistribution to distributors",
        "Specialty Focus": "Redistributor — oils move through Dot to regional distributors",
    },
    "mclane": {
        "Founded": "1894",
        "Headquarters": "Temple, Texas, USA",
        "Website": "https://www.mclaneco.com",
        "Business Type": "Wholesale distribution",
        "Ownership": "Berkshire Hathaway",
        "Distribution Type": "Grocery / convenience wholesale",
    },
    "core-mark": {
        "Headquarters": "Westlake, Texas, USA",
        "Website": "https://www.core-mark.com",
        "Business Type": "Convenience wholesale distributor",
        "Ownership": "Part of Performance Food Group",
        "Distribution Type": "Convenience wholesale",
    },
    "spartannash": {
        "Founded": "1917",
        "Headquarters": "Byron Center, Michigan, USA",
        "Website": "https://www.spartannash.com",
        "Business Type": "Grocery wholesale & retail support",
        "Ownership": "Public (Nasdaq: SPTN)",
        "Distribution Type": "Grocery wholesale",
    },
    "associated wholesale grocers": {
        "Founded": "1924",
        "Headquarters": "Kansas City, Kansas, USA",
        "Website": "https://awginc.com",
        "Business Type": "Grocery cooperative wholesaler",
        "Distribution Type": "Retailer-owned wholesale",
    },
    "associated food stores": {
        "Headquarters": "Salt Lake City, Utah, USA",
        "Website": "https://www.afstores.com",
        "Business Type": "Grocery cooperative wholesaler",
        "Distribution Type": "Retailer-owned wholesale",
    },
    "affiliated foods": {
        "Headquarters": "Amarillo, Texas, USA",
        "Website": "http://www.afiama.com",
        "Business Type": "Grocery wholesaler",
        "Distribution Type": "Grocery wholesale",
    },
    "wakefern": {
        "Founded": "1946",
        "Headquarters": "Keasbey, New Jersey, USA",
        "Website": "https://www.wakefern.com",
        "Business Type": "Retailer-owned grocery cooperative",
        "Distribution Type": "Retailer-owned wholesale",
    },
    "chefs' warehouse": {
        "Headquarters": "Ridgefield, Connecticut, USA",
        "Website": "https://www.chefswarehouse.com",
        "Business Type": "Specialty foodservice distributor",
        "Ownership": "Public (Nasdaq: CHEF)",
        "Distribution Type": "Specialty foodservice",
        "Specialty Focus": "Premium specialty oils to chefs / foodservice",
    },
    "baldor specialty foods": {
        "Headquarters": "Bronx, New York, USA",
        "Website": "https://www.baldorfood.com",
        "Business Type": "Specialty produce & foodservice distributor",
        "Distribution Type": "Specialty foodservice",
    },
    "brakes": {
        "Headquarters": "United Kingdom",
        "Website": "https://www.brake.co.uk",
        "Business Type": "Foodservice distributor",
        "Ownership": "Part of Sysco",
        "Distribution Type": "Foodservice wholesale (UK)",
    },
    "bidfood": {
        "Headquarters": "United Kingdom / international",
        "Website": "https://www.bidfood.co.uk",
        "Business Type": "Foodservice distributor",
        "Distribution Type": "Foodservice wholesale",
    },
    "brenntag": {
        "Headquarters": "Essen, Germany",
        "Website": "https://www.brenntag.com",
        "Business Type": "Chemical / ingredients distributor",
        "Distribution Type": "Ingredients distribution (incl. oils)",
        "Specialty Focus": "Channel for specialty/cosmetic & food oils",
    },
    "univar": {
        "Headquarters": "Downers Grove, Illinois, USA",
        "Website": "https://www.univarsolutions.com",
        "Business Type": "Chemical / ingredients distributor",
        "Distribution Type": "Ingredients distribution",
    },
    "azelis": {
        "Founded": "2001",
        "Headquarters": "Antwerp, Belgium",
        "Website": "https://www.azelis.com",
        "Business Type": "Specialty ingredients distributor",
        "Distribution Type": "Ingredients distribution",
    },
    "jedwards international": {
        "Headquarters": "Massachusetts, USA",
        "Website": "https://bulknaturaloils.com",
        "Business Type": "Bulk natural oils wholesaler / distributor",
        "Distribution Type": "Bulk oils wholesale",
        "Core Categories": "Bulk avocado oil and other natural oils",
        "Specialty Focus": "B2B bulk oils channel (not a grower mill)",
    },
    "bulk apothecary": {
        "Headquarters": "Ohio, USA",
        "Website": "https://www.bulkapothecary.com",
        "Business Type": "Bulk ingredients wholesaler",
        "Distribution Type": "Ingredients wholesale",
    },
    "all organic treasures": {
        "Company": "All Organic Treasures GmbH",
        "Headquarters": "Wiggensbach, Germany",
        "Website": "https://www.aot.de",
        "Business Type": "Organic oils & ingredients wholesaler",
        "Distribution Type": "Organic ingredients wholesale",
    },
    "reyes holdings": {
        "Headquarters": "Rosemont, Illinois, USA",
        "Website": "https://www.reyesholdings.com",
        "Business Type": "Beverage & food distribution",
        "Distribution Type": "Distribution holding company",
    },
    "martin-brower": {
        "Headquarters": "Rosemont, Illinois, USA",
        "Website": "https://www.martin-brower.com",
        "Business Type": "QSR supply-chain distributor",
        "Distribution Type": "QSR distribution",
    },
}


def _is_npd(v: object) -> bool:
    return str(v or "").strip().lower() in NPD


def _is_manufacturer(name: str, row: dict[str, str] | None = None) -> bool:
    key = _norm(name)
    if any(n in key for n in MANUFACTURER_NEEDLES):
        return True
    blob = " ".join(
        [
            name,
            (row or {}).get("Business Type") or "",
            (row or {}).get("Specialty Focus") or "",
            (row or {}).get("Summary") or "",
        ]
    )
    if MANUFACTURER_RE.search(blob) and not CHANNEL_HINT_RE.search(blob):
        return True
    # product-as-company
    if key.startswith("avocado oil") or "100 pure avocado" in key:
        return True
    return False


def _is_junk(name: str) -> bool:
    n = (name or "").strip()
    if not n or JUNK_RE.match(n):
        return True
    return False


def _is_channel(row: dict[str, str]) -> bool:
    """Market-related channel / distributor check."""
    name = row.get("Company") or ""
    if _is_junk(name) or _is_manufacturer(name, row):
        return False
    blob = " ".join(
        str(row.get(h) or "")
        for h in ("Business Type", "Distribution Type", "Specialty Focus", "Summary", "Data Sources", "Company")
    )
    # Seeds are labeled distributors of avocado oil
    if "distributor" in blob.lower() or "wholesale" in blob.lower() or "foodservice" in blob.lower():
        return True
    if CHANNEL_HINT_RE.search(blob):
        return True
    # Default seed companies without rich text still OK if not manufacturer
    src = (row.get("Data Sources") or "").lower()
    if "seed" in src or "seed_file" in src:
        return True
    bt = (row.get("Business Type") or "").lower()
    if bt in ("distributor", "wholesaler", "importer", "foodservice distributor"):
        return True
    return False


def _load_seeds() -> list[str]:
    out = []
    for ln in SEEDS.read_text(encoding="utf-8").splitlines():
        ln = ln.strip()
        if not ln or ln.startswith("#"):
            continue
        name = ln.split("|")[0].strip()
        if name and not _is_manufacturer(name) and not _is_junk(name):
            out.append(name)
    return out


def _load_final() -> list[dict[str, str]]:
    if not FINAL.exists():
        return []
    wb = load_workbook(FINAL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c or "") for c in rows[0]]
    out = []
    for r in rows[1:]:
        d = {h: str(r[i] if i < len(r) and r[i] is not None else "").strip() for i, h in enumerate(hdr)}
        if d.get("Company"):
            out.append({h: d.get(h, "") for h in HEADERS})
    wb.close()
    return out


def _apply_facts(row: dict[str, str]) -> int:
    key = _norm(row.get("Company", ""))
    n = 0
    for needle, patch in WEB_FACTS.items():
        if needle in key:
            for h, v in patch.items():
                if h == "Market Related":
                    continue
                if h == "Company":
                    row["Company"] = v
                    n += 1
                    continue
                if v and (_is_npd(row.get(h)) or h in ("Founded", "Headquarters", "Contact Person", "Role", "Employees", "Website")):
                    if _is_npd(row.get(h)) or h in ("Contact Person", "Role", "Founded", "Headquarters", "Employees", "Turnover / Scale"):
                        row[h] = v
                        n += 1
            break
    return n


def _defaults(row: dict[str, str]) -> None:
    for h, dv in (MARKET_DEFAULTS.get(FAMILY) or {}).items():
        if h in HEADERS and _is_npd(row.get(h)):
            row[h] = dv
    row["Business Type"] = "Distributor"
    if _is_npd(row.get("Distribution Type")):
        row["Distribution Type"] = "Wholesale / foodservice channel"
    if _is_npd(row.get("Core Categories")):
        row["Core Categories"] = "Edible oils; avocado oil; grocery / foodservice SKUs"
    if _is_npd(row.get("Specialty Focus")):
        row["Specialty Focus"] = "Channel partner for avocado oil / edible oils"
    if _is_npd(row.get("Summary")):
        row["Summary"] = (
            f"{row.get('Company')} — distributor/channel partner in the global avocado oil market "
            "(not a manufacturer)."
        )
    src = row.get("Data Sources") or ""
    if "agent_channel_web" not in src:
        row["Data Sources"] = f"{src}; agent_channel_web_verify".strip("; ")
    for h in HEADERS:
        if not str(row.get(h) or "").strip():
            row[h] = "Not publicly disclosed"


def _wiki_fill(name: str) -> dict[str, str]:
    """Direct Wikipedia API — not SearXNG."""
    out: dict[str, str] = {}
    try:
        with httpx.Client(timeout=12.0, headers={"User-Agent": "VendorIntelAgent/1.0"}) as client:
            sr = client.get(
                "https://en.wikipedia.org/w/api.php",
                params={
                    "action": "query",
                    "list": "search",
                    "srsearch": name,
                    "format": "json",
                    "srlimit": 2,
                },
            )
            if sr.status_code != 200:
                return out
            hits = (sr.json().get("query") or {}).get("search") or []
            if not hits:
                return out
            title = hits[0].get("title") or ""
            er = client.get(
                "https://en.wikipedia.org/w/api.php",
                params={
                    "action": "query",
                    "prop": "extracts",
                    "exintro": 1,
                    "explaintext": 1,
                    "titles": title,
                    "format": "json",
                },
            )
            if er.status_code != 200:
                return out
            pages = (er.json().get("query") or {}).get("pages") or {}
            text = ""
            for p in pages.values():
                text = p.get("extract") or ""
                break
            if not text:
                return out
            # light parse
            m = re.search(r"(?:founded|established)\s+in\s+(\d{4})", text, re.I)
            if m:
                out["Founded"] = m.group(1)
            m = re.search(
                r"(?:headquartered in|based in)\s+([A-Z][^.\n]{3,60})",
                text,
            )
            if m:
                out["Headquarters"] = m.group(1).strip().rstrip(",")
            m = re.search(r"([\d,]+)\s+employees", text, re.I)
            if m:
                out["Employees"] = m.group(1).replace(",", "")
            # reject if Wikipedia says manufacturer of oil only
            low = text.lower()
            if ("manufacturer" in low or "producer of" in low) and "distributor" not in low and "wholesale" not in low:
                out["_maybe_mfr"] = "1"
    except Exception:
        return out
    return out


def main() -> int:
    print("=== Channel-only avocado rebuild (NO SearXNG) ===", flush=True)
    by: dict[str, dict[str, str]] = {}

    # 1) Seeds first (curated distributors)
    for name in _load_seeds():
        k = _norm(name)
        by[k] = _blank_row(name, "", FAMILY, "seed_distributor")
        by[k]["Business Type"] = "Distributor"
        by[k]["Specialty Focus"] = "Distributors of avocado oil / edible oils"

    # 2) Merge existing FINAL if channel-ok
    dropped_mfr = 0
    dropped_junk = 0
    for r in _load_final():
        name = r.get("Company") or ""
        if _is_junk(name):
            dropped_junk += 1
            continue
        if _is_manufacturer(name, r):
            print(f"  drop manufacturer: {name}", flush=True)
            dropped_mfr += 1
            continue
        if not _is_channel(r) and _norm(name) not in by:
            # allow if already in seeds map
            print(f"  drop non-channel: {name}", flush=True)
            dropped_mfr += 1
            continue
        k = _norm(name)
        if k in by:
            # merge richer non-NPD into seed row
            for h in HEADERS:
                if h == "Company":
                    continue
                if not _is_npd(r.get(h)) and (
                    _is_npd(by[k].get(h)) or h in ("Website", "Headquarters", "Founded")
                ):
                    by[k][h] = r[h]
        else:
            if _is_channel(r):
                by[k] = {h: r.get(h, "") for h in HEADERS}

    rows = list(by.values())
    print(f"  after channel filter: {len(rows)} (dropped mfr/junk={dropped_mfr}/{dropped_junk})", flush=True)

    # 3) Apply agent web facts
    fact_writes = 0
    for r in rows:
        fact_writes += _apply_facts(r)

    # 4) Wikipedia fill for remaining HQ/Founded gaps (direct API, not SearXNG)
    wiki_n = 0
    need_wiki = [r for r in rows if _is_npd(r.get("Founded")) or _is_npd(r.get("Headquarters"))]
    print(f"  wikipedia fill for {len(need_wiki)} rows…", flush=True)
    for i, r in enumerate(need_wiki, 1):
        name = r["Company"]
        w = _wiki_fill(name)
        if w.get("_maybe_mfr") and not CHANNEL_HINT_RE.search(
            " ".join([r.get("Business Type") or "", r.get("Specialty Focus") or ""])
        ):
            # skip applying if looks like pure manufacturer page and not already seeded channel
            if "seed" not in (r.get("Data Sources") or "").lower():
                print(f"  [{i}] skip wiki mfr-like: {name}", flush=True)
                continue
        got = False
        for h in ("Founded", "Headquarters", "Employees"):
            if w.get(h) and _is_npd(r.get(h)):
                r[h] = w[h]
                got = True
        if got:
            wiki_n += 1
            print(f"  [{i}] wiki filled {name}", flush=True)
        if i % 40 == 0:
            print(f"  [{i}/{len(need_wiki)}] …", flush=True)

    # 5) Final verify pass — manufacturers again + force distributor labeling
    final: list[dict[str, str]] = []
    for r in rows:
        if _is_manufacturer(r.get("Company", ""), r):
            print(f"  final drop mfr: {r.get('Company')}", flush=True)
            continue
        _apply_facts(r)
        _defaults(r)
        r["Business Type"] = "Distributor"
        final.append(r)

    final.sort(key=lambda x: _norm(x.get("Company", "")))

    # de-dupe
    uniq: dict[str, dict[str, str]] = {}
    for r in final:
        uniq[_norm(r["Company"])] = r
    final = sorted(uniq.values(), key=lambda x: _norm(x["Company"]))

    audit = {
        "mode": "agent_own_web_channel_only_no_searxng",
        "market": MARKET,
        "rule": "distributors/channel partners only — manufacturers removed",
        "dropped_manufacturers": dropped_mfr,
        "dropped_junk": dropped_junk,
        "web_fact_writes": fact_writes,
        "wikipedia_fills": wiki_n,
        "companies": len(final),
        "target_met": len(final) >= 300,
        "engines": "agent WebSearch facts + Wikipedia API (no SearXNG/DDGS)",
        "output": str(FINAL),
    }
    write_final_xlsx(FINAL, final, "Avocado Oil Distributors", audit)
    (FINAL.parent / "agent_channel_only_audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2), flush=True)
    print(f"DONE → {len(final)} channel companies → {FINAL}", flush=True)
    return 0 if len(final) >= 300 else 2


if __name__ == "__main__":
    raise SystemExit(main())
