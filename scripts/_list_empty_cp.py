# -*- coding: utf-8 -*-
"""List empty Contact Person rows for one or more Excels."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}


def list_path(path: Path) -> None:
    if not path.exists():
        print("MISSING", path)
        return
    wb = load_workbook(path, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and ("Contact Person" in vals or "Contact" in vals):
            headers = vals
            hr = r
            break
    print("===", path.name, "sheet", ws.title, "header_row", hr)
    tot = cp = li_url = empty_cp = named_npd = 0
    empties = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {
            headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
        }
        co = d.get("Company") or ""
        if not co:
            continue
        tot += 1
        person = (d.get("Contact Person") or d.get("Contact") or "").strip()
        role = (d.get("Role") or "").strip()
        li = (d.get("LinkedIn") or "").strip()
        web = (d.get("Website") or "").strip()
        hq = (d.get("Headquarters") or d.get("HQ") or "").strip()
        has_cp = bool(person) and person.lower() not in NPD
        if has_cp:
            cp += 1
            if "linkedin.com/in/" in li.lower():
                li_url += 1
            else:
                named_npd += 1
        else:
            empty_cp += 1
            empties.append((co, web, hq, person, role, li))
    print(f"rows={tot} cp={cp} li_url={li_url} named_npd={named_npd} empty_cp={empty_cp}")
    for i, e in enumerate(empties, 1):
        print(f"{i}\t{e[0]}\t{e[1]}\t{e[2]}")
    print("empty_total", len(empties))
    wb.close()


if __name__ == "__main__":
    paths = [Path(p) for p in sys.argv[1:]]
    if not paths:
        paths = [
            Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market.xlsx"),
            Path(r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market.xlsx"),
        ]
    for p in paths:
        list_path(p)
        print()
