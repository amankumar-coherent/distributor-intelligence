"""Avocado FINAL: keep distributors / channel partners only; drop manufacturers.

Removes:
- Pure brands / D2C / private-label brand rows (not a distributor entity)
- Oil / food producers & mills primarily manufacturing
- Fresh produce growers / estate producers (not oil channel partners)
- Commodity marketing boards
- Product-as-company placeholder rows (Bulk Avocado Oil, etc.)
- Chemical / ingredient manufacturers (non-channel)

Keeps:
- Foodservice / grocery / cash & carry / wholesale distributors
- Ingredient & specialty chemical distributors (channel)
- Retail / warehouse-club channel partners
- Agri distribution arms clearly labeled as distribution (e.g. Cargill Food Distribution)
"""
from __future__ import annotations

import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl

PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_global_FINAL.xlsx"
)

# Exact company names to remove (case-insensitive after norm)
REMOVE_EXACT = {
    # Brands / D2C / private label (not distributor entities)
    "avohass",
    "crisco",
    "kirkland signature",
    "healthy harvest",
    "nature s way",
    "nutrablast",
    "nutrabio",
    "nutricost wholesale",
    "la tourangelle distribution",
    "oils of aloha",
    "spectrum naturals hain celestial channel",
    "borges international group",
    "deoleo distribution partners",
    # Marketing board
    "california avocado commission",
    # Producers / processors
    "musim mas distribution",
    "wilmar international distribution",
    "crest foods co inc",
    "crest foods ingredients",
    "hydrite chemical",
    "essential wholesale",
    "spectrum chemical mfg corp",
    # Growers / estate producers
    "bonnysa",
    "compagnie fruitiere",
    "frutas montosa",
    "g s fresh",
    "reyes gutierrez",
    "sanlucar",
    # Product-as-company / bulk product listings
    "bulk avocado oil",
    "bulk avocado culinary oil",
    "organic avocado oil",
    "refined avocado oil",
    "wholesale avocado oil",
    "olve true avocado oil",
    "kevala s avocado oil bulk",
    "tropical avocado",
    "hbno bulk",
}

# Substring needles (matched against normalized company key)
REMOVE_NEEDLES = {
    "chosen foods",
    "olivado",
    "limbua",
    "avohealth",
    "sesajal",
    "calavo",
    "mission produce",
    "westfalia",
    "avolio",
    "avopure",
    "avooil",
    "village press",
    "catania oils",
    "aceitera",
    "oil mill",
    "cold pressed",
}

# Business types that are not distributor/channel (unless company is allowlisted)
REMOVE_BT = re.compile(
    r"(?i)^("
    r"d2c\b.*|"
    r".*\bbrand channel\b.*|"
    r"private label brand channel|"
    r"edible oils brand channel|"
    r"natural oils brand channel|"
    r"olive oil brand distribution|"
    r"commodity marketing board|"
    r"specialty oils producer / distributor|"
    r"specialty oils brand / distributor|"
    r"edible oils / foods producer & distributor|"
    r"palm / oils producer / distributor|"
    r"sports nutrition brand / wholesale|"
    r"supplements brand / wholesale|"
    r"fresh produce grower / distributor|"
    r"fresh produce producer / distributor|"
    r"food ingredients manufacturer / distributor|"
    r"cosmetic ingredients manufacturer / wholesale|"
    r"chemical distributor / manufacturer|"
    r"bulk avocado oil supplier"  # product listings / mill-style bulk brand rows
    r")$"
)

# Keep these even if BT matches a gray area (true channel partners)
KEEP_EXACT = {
    "trader joe s",  # specialty grocery retail channel
    "jetro restaurant depot",
    "restaurant depot",
    "costco mexico wholesale",
    "sam s club mexico",
    "bunge loders croklaan distribution",
    "cargill food distribution europe",
    "spectrum chemical distribution",  # lab chemical distributor arm
    "henry lamotte oils",
    "new directions aromatics",
    "the olive oil company",
    "nature s best",
    "nature s pride",
    "nature s produce",
    "from nature with love",
    "bulk apothecary",
    "bulk by cho",
}


def norm_key(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower().replace("&", " and ").replace("™", " ").replace("®", " ")
    s = s.replace("(", " ").replace(")", " ").replace("/", " ").replace(".", " ")
    s = s.replace(",", " ").replace("'", " ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def should_remove(company: str, btype: str) -> str | None:
    key = norm_key(company)
    if not key:
        return "empty"
    if key in KEEP_EXACT:
        return None
    if key in REMOVE_EXACT:
        return "exact_mfr_brand"
    if any(n in key for n in REMOVE_NEEDLES):
        return "needle_mfr"
    # product-as-company patterns
    if re.match(
        r"^(bulk |organic |refined |wholesale |cold pressed )?(avocado oil)\b",
        key,
    ) or key.endswith(" avocado oil") and len(key) < 40:
        if "distribution" not in key and "distributor" not in key:
            return "product_as_company"
    if REMOVE_BT.match((btype or "").strip()):
        return "business_type"
    # BT contains manufacturer/producer without clear distributor-only label
    bt = (btype or "").lower()
    if "manufacturer" in bt and "distributor" in bt:
        return "hybrid_manufacturer"
    if re.search(r"\b(producer|grower)\b", bt) and "distributor" in bt:
        return "hybrid_producer_grower"
    if re.search(r"\bd2c\b", bt):
        return "d2c_brand"
    return None


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = PATH.with_name(f"{PATH.stem}_pre_channel_only_{stamp}{PATH.suffix}")
    shutil.copy2(PATH, bak)

    wb = openpyxl.load_workbook(PATH)
    ws = wb["Avocado Oil Distributors"]
    hdr = [c.value for c in ws[1]]
    cols = {h: i + 1 for i, h in enumerate(hdr) if h}

    removed: list[tuple[str, str, str]] = []
    keep_rows: list[list] = []

    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, len(hdr) + 1)]
        company = str(row[cols["Company"] - 1] or "").strip()
        btype = str(row[cols["Business Type"] - 1] or "").strip()
        reason = should_remove(company, btype)
        if reason:
            removed.append((company, btype, reason))
        else:
            keep_rows.append(row)

    # rewrite sheet
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for i, row in enumerate(keep_rows, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(i, c).value = val

    wb.save(PATH)

    # reason tally
    from collections import Counter

    reasons = Counter(r[2] for r in removed)
    print("backup", bak.name)
    print("before", len(keep_rows) + len(removed), "after", len(keep_rows), "removed", len(removed))
    print("reasons", dict(reasons))
    print("--- removed ---")
    for co, bt, reason in sorted(removed, key=lambda x: x[0].lower()):
        print(f"  [{reason}] {co} | {bt}")


if __name__ == "__main__":
    main()
