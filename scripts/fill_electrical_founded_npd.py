"""Fill NPD Founded years in electrical FINAL via known brands + web-verified years.

Only writes years backed by public company history / Wikipedia / official About pages.
Templated fake rows (Casa Eléctrica de X, Elektro Großhandel X, PT Mega…, etc.) stay NPD.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from vendor_intel.enrichment.known_founded import (
    resolve_known_founded,
    sanitize_founded_value,
)

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)
NPD_SET = {
    "",
    "not publicly disclosed",
    "n/a",
    "npd",
    "unknown",
    "-",
}

# Exact-name fills from agent web verification (sources in audit JSON)
WEB_VERIFIED: dict[str, tuple[str, str]] = {
    # (year, short_source)
    "Easa Saleh Al Gurg Group": ("1960", "algurg.com timeline / Family Business Histories"),
    "Juma Al Majid Group": ("1950", "al-majid.com / Wikipedia"),
    "Khimji Ramdas": ("1870", "khimji.com about / Oman Observer 150 years"),
    "Bin Dasmal Group": ("1976", "bindasmal.com group profile"),
    "Hussain Ali Alireza & Sons": ("1906", "Wikipedia Haji Husein Alireza & Co."),
    "L&H Group": ("1886", "lh.com.au since 1886 / Sonepar AU"),
    "Middy's Electrical (NZ)": ("1928", "middys.com.au about / family since 1928"),
    "PlaceMakers": ("1981", "Wikipedia PlaceMakers / brand timeline"),
    "Bunzl Australia": ("1854", "bunzl.com Our history origins 1854"),
    "Element14": ("1939", "Wikipedia / Farnell heritage A.C. Farnell 1939"),
    "Element14 (Asia Pacific)": ("1939", "Wikipedia / Farnell heritage 1939"),
    "Plumb Center (Wolseley UK)": ("1985", "Wolseley Centers / Plumb Center 1985"),
    "Wolseley Canada": ("1958", "PrivCo / Wolseley Canada founding"),
    "Wesfarmers Industrial & Safety": ("1914", "wesfarmers.com history 1914"),
    "Würth Elektro Großhandel": ("1945", "wuerth.com History / Wikipedia"),
    "Wuerth Elektro Grosshandel": ("1945", "wuerth.com History / Wikipedia"),
    "Capitol Light": ("1926", "capitollight.com About since 1926"),
    "Corys Electrical": ("1920", "corys.co.nz history Cory-Wright & Salmon"),
    "Elektroland24": ("2003", "elektroland24.de Über uns / ecommerceDB"),
    "AWM Electrical": ("1916", "Metal Manufactures Ltd / MMEM 1916"),
    "AWM Electrical NZ": ("1916", "Metal Manufactures Ltd / MMEM 1916"),
    "Haymans Electrical": ("1916", "Metal Manufactures Ltd / MMEM 1916"),
    "TLE Electrical": ("1916", "Metal Manufactures Ltd / MMEM 1916"),
    "Litecor": ("1985", "Yellosa Litecor / Voltex division listing"),
    "Vietnam Electrical Equipment Joint Stock Company (VEE)": (
        "1960",
        "veecp.com about Cam Pha 1960",
    ),
    "Laser Electrical": ("1983", "laserelectrical.com.au Our Story"),
    "Weidmuller South Africa": ("1850", "Weidmüller group founding"),
}


def _is_npd(val: object) -> bool:
    return str(val or "").strip().lower() in NPD_SET


def main() -> None:
    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    fi = headers.index("Founded") + 1

    changes: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        cur = str(ws.cell(row_idx, fi).value or "").strip()
        new = None
        reason = ""
        source = ""

        # 1) Exact web-verified map (only fill empty/NPD, or fix via known sanitize)
        if name in WEB_VERIFIED:
            year, source = WEB_VERIFIED[name]
            if _is_npd(cur) or cur != year:
                if _is_npd(cur) or resolve_known_founded(name) is not None:
                    new, reason = year, "web_verified_fill"

        # 2) known_founded registry (Rexel/Sonepar/Al-Futtaim/…)
        if new is None:
            known = resolve_known_founded(name)
            if known is not None:
                target = str(known)
                if _is_npd(cur) or cur != target:
                    new, reason, source = target, "known_founded_registry", "known_founded.py"

        # 3) Sanitize any existing wrong known-brand year even if not NPD
        if new is None and not _is_npd(cur):
            fixed = sanitize_founded_value(name, cur)
            if fixed and fixed != cur and fixed.isdigit():
                new, reason, source = fixed, "known_founded_sanitize", "known_founded.py"

        if new is not None:
            ws.cell(row_idx, fi).value = new
            changes.append(
                {
                    "company": name,
                    "old": cur or "(empty)",
                    "new": new,
                    "reason": reason,
                    "source": source,
                }
            )

    for cell in ws[1]:
        cell.font = Font(bold=True)

    wb.save(XLSX)
    dump = XLSX.parent / "_founded_npd_fill_audit.json"
    dump.write_text(json.dumps(changes, indent=2, ensure_ascii=False), encoding="utf-8")

    # Recount
    npd_left = 0
    with_year = 0
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        founded = str(ws.cell(row_idx, fi).value or "").strip()
        if _is_npd(founded):
            npd_left += 1
        else:
            with_year += 1

    print(f"filled/corrected {len(changes)} Founded values")
    for c in changes:
        print(f"  {c['company']}: {c['old']!r} -> {c['new']!r} ({c['reason']})")
    print(f"now with_year={with_year} npd={npd_left} total={with_year + npd_left}")
    print(f"wrote {XLSX}")
    print(f"audit {dump}")


if __name__ == "__main__":
    main()
