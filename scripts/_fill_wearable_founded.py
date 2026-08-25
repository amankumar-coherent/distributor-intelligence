# -*- coding: utf-8 -*-
"""Fill Founded year for Wearable Medical FINAL from company websites + optional DeepSeek extract.

No Google AI scraper. Never invent: only write a year when evidence is found.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
CACHE = ROOT / "_wearable_founded_harvest_cache.json"
AUDIT = ROOT / "_wearable_founded_fill_audit.json"

load_dotenv(Path(r"D:\vendor-intel\.env"))

NPD = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "null",
}

YEAR_PATTERNS = [
    re.compile(
        r"(?i)\b(?:founded|established|est\.?|incorporated|incorporation|started|since|dating\s+back\s+to|formed|created)\b"
        r"[^\d]{0,40}\b((?:18|19|20)\d{2})\b"
    ),
    re.compile(
        r"(?i)\b((?:18|19|20)\d{2})\b[^\d]{0,20}\b(?:founded|established|est\.?|incorporated)\b"
    ),
    re.compile(r"(?i)\bsince\s+((?:18|19|20)\d{2})\b"),
    re.compile(r"(?i)\best(?:ablished|\.)?\s*(?:in\s*)?((?:18|19|20)\d{2})\b"),
    re.compile(r"(?i)\bfounding\s+year[:\s]+((?:18|19|20)\d{2})\b"),
]

ABOUT_PATHS = [
    "",
    "/about",
    "/about-us",
    "/aboutus",
    "/about/",
    "/about-us/",
    "/company",
    "/company/",
    "/company/about",
    "/company/profile",
    "/our-company",
    "/our-story",
    "/our-history",
    "/history",
    "/who-we-are",
    "/en/about",
    "/en/about-us",
    "/en/company",
    "/en/company/about",
    "/corporate/about",
    "/corporate/about-us",
    "/about/company",
    "/pages/about-us",
]


def is_empty_founded(v: str) -> bool:
    s = (v or "").strip().lower()
    return s in NPD or s.startswith("not public")


def valid_year(y: int) -> bool:
    return 1800 <= y <= datetime.now().year


def extract_years(text: str) -> list[tuple[int, str]]:
    if not text:
        return []
    # strip scripts/styles roughly
    t = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", text)
    t = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    t = re.sub(r"\s+", " ", t)
    hits: list[tuple[int, str]] = []
    for pat in YEAR_PATTERNS:
        for m in pat.finditer(t):
            try:
                y = int(m.group(1))
            except Exception:
                continue
            if valid_year(y):
                ctx = t[max(0, m.start() - 40) : m.end() + 40]
                hits.append((y, ctx.strip()))
    return hits


def pick_year(hits: list[tuple[int, str]]) -> tuple[str, str]:
    if not hits:
        return "", ""
    # prefer earliest plausible founding year among hits (company founding usually oldest)
    # but filter out copyright-only years by requiring founding keywords in context
    keyed = []
    for y, ctx in hits:
        low = ctx.lower()
        if any(
            k in low
            for k in (
                "founded",
                "established",
                "est.",
                "est ",
                "incorporated",
                "since",
                "started",
                "formed",
                "created",
                "dating back",
                "founding",
            )
        ):
            keyed.append((y, ctx))
    pool = keyed or hits
    # modal / earliest: use earliest among keyed
    y = min(x[0] for x in pool)
    ctx = next(c for yy, c in pool if yy == y)
    return str(y), ctx


def base_url(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    p = urlparse(u)
    return f"{p.scheme}://{p.netloc}"


def fetch(url: str, client: httpx.Client) -> str:
    try:
        r = client.get(url, follow_redirects=True)
        if r.status_code >= 400:
            return ""
        ctype = (r.headers.get("content-type") or "").lower()
        if "html" not in ctype and "text" not in ctype and ctype:
            return ""
        return r.text[:250000]
    except Exception:
        return ""


def harvest_one(company: str, website: str) -> dict:
    root = base_url(website)
    if not root:
        return {"company": company, "website": website, "founded": "", "source": "", "evidence": "", "ok": False}
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml",
        "Accept-Language": "en-US,en;q=0.9",
    }
    best_year = ""
    best_ev = ""
    best_src = ""
    with httpx.Client(timeout=12.0, headers=headers, verify=False) as client:
        for path in ABOUT_PATHS:
            url = root.rstrip("/") + path if path.startswith("/") or path == "" else urljoin(root + "/", path)
            if path == "":
                url = root
            html = fetch(url, client)
            if not html or len(html) < 200:
                continue
            hits = extract_years(html)
            y, ctx = pick_year(hits)
            if y:
                # prefer about pages over homepage if both found
                if not best_year or (path and not best_src):
                    best_year, best_ev, best_src = y, ctx, url
                elif path and ("about" in path or "compan" in path or "histor" in path or "story" in path):
                    best_year, best_ev, best_src = y, ctx, url
                    break
    return {
        "company": company,
        "website": website,
        "founded": best_year,
        "source": best_src,
        "evidence": best_ev[:240],
        "ok": bool(best_year),
    }


def deepseek_extract(company: str, website: str, snippets: str) -> str:
    """Extract founded year ONLY from provided snippets; return '' if not present."""
    api_key = os.getenv("OPENAI_API_KEY") or ""
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    if not api_key or not snippets.strip():
        return ""
    prompt = (
        "From the website text snippets below, extract the company founding/establishment year "
        f"for '{company}' ({website}). "
        "Return ONLY a 4-digit year if clearly stated as founded/established/since/incorporated. "
        "If not clearly present, return NONE. Do not guess.\n\n"
        f"SNIPPETS:\n{snippets[:6000]}"
    )
    try:
        with httpx.Client(timeout=40.0) as client:
            r = client.post(
                f"{base}/chat/completions",
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json={
                    "model": model,
                    "temperature": 0,
                    "messages": [
                        {"role": "system", "content": "Extract only explicitly stated founding years. No guessing."},
                        {"role": "user", "content": prompt},
                    ],
                },
            )
            r.raise_for_status()
            text = r.json()["choices"][0]["message"]["content"].strip()
        m = re.search(r"\b((?:18|19|20)\d{2})\b", text)
        if not m or "NONE" in text.upper():
            return ""
        y = int(m.group(1))
        return str(y) if valid_year(y) else ""
    except Exception:
        return ""


def main() -> None:
    import urllib3

    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = XLSX.with_name(XLSX.stem + f"_BEFORE_FOUNDED_{ts}.xlsx")
    shutil.copy2(XLSX, backup)

    cache: dict = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    wb = load_workbook(XLSX)
    ws = wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(hdr) if h}

    targets = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        founded = str(ws.cell(r, ic["Founded"]).value or "").strip()
        if is_empty_founded(founded):
            targets.append({"row": r, "company": co, "website": web})

    print("need", len(targets), "already cached", sum(1 for t in targets if t["company"] in cache and cache[t["company"]].get("ok")))

    # harvest missing
    todo = [t for t in targets if not (cache.get(t["company"], {}).get("ok"))]
    print("harvesting", len(todo))
    done = 0
    with ThreadPoolExecutor(max_workers=16) as ex:
        futs = {ex.submit(harvest_one, t["company"], t["website"]): t for t in todo}
        for fut in as_completed(futs):
            t = futs[fut]
            try:
                res = fut.result()
            except Exception as e:
                res = {"company": t["company"], "website": t["website"], "founded": "", "source": "", "evidence": str(e), "ok": False}
            cache[t["company"]] = res
            done += 1
            if done % 25 == 0:
                CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
                print("harvested", done, "/", len(todo), "hits", sum(1 for v in cache.values() if v.get("ok")))

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")

    # DeepSeek pass for misses: re-fetch about text briefly and ask extract-only
    misses = [t for t in targets if not cache.get(t["company"], {}).get("founded")]
    print("deepseek candidates", len(misses))
    ds_filled = 0
    for i, t in enumerate(misses[:180], 1):  # cap API volume
        # gather raw snippets again quickly from homepage+about
        root = base_url(t["website"])
        snippets = []
        headers = {"User-Agent": "Mozilla/5.0", "Accept": "text/html"}
        try:
            with httpx.Client(timeout=10.0, headers=headers, verify=False) as client:
                for path in ["", "/about", "/about-us", "/company", "/our-story", "/history"]:
                    url = root if path == "" else root.rstrip("/") + path
                    html = fetch(url, client)
                    if html:
                        hits = extract_years(html)
                        for y, ctx in hits[:3]:
                            snippets.append(ctx)
                        # also keep a compressed text window
                        plain = re.sub(r"(?is)<[^>]+>", " ", html)
                        plain = re.sub(r"\s+", " ", plain)
                        if "founded" in plain.lower() or "established" in plain.lower() or "since 1" in plain.lower() or "since 2" in plain.lower():
                            snippets.append(plain[:2500])
        except Exception:
            pass
        if not snippets:
            continue
        y = deepseek_extract(t["company"], t["website"], "\n---\n".join(snippets))
        if y:
            cache[t["company"]] = {
                "company": t["company"],
                "website": t["website"],
                "founded": y,
                "source": "deepseek_extract_from_site_text",
                "evidence": snippets[0][:240] if snippets else "",
                "ok": True,
            }
            ds_filled += 1
        if i % 20 == 0:
            CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
            print("deepseek progress", i, "filled", ds_filled)
        time.sleep(0.15)

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")

    # write to excel
    filled = 0
    still = []
    for t in targets:
        res = cache.get(t["company"], {})
        y = str(res.get("founded") or "").strip()
        if y and valid_year(int(y)):
            ws.cell(t["row"], ic["Founded"]).value = y
            filled += 1
        else:
            # leave as Not publicly disclosed if empty
            cur = str(ws.cell(t["row"], ic["Founded"]).value or "").strip()
            if is_empty_founded(cur):
                ws.cell(t["row"], ic["Founded"]).value = "Not publicly disclosed"
            still.append({"company": t["company"], "website": t["website"]})

    wb.save(XLSX)
    wb.close()

    # recount
    wb2 = load_workbook(XLSX, data_only=True)
    ws2 = wb2.active
    hdr2 = [str(c.value or "").strip() for c in next(ws2.iter_rows(min_row=1, max_row=1, values_only=True))]
    ic2 = {h: i for i, h in enumerate(hdr2)}
    have = need = 0
    for r in ws2.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        f = str(r[ic2["Founded"]] or "").strip()
        if is_empty_founded(f):
            need += 1
        else:
            have += 1
    wb2.close()

    audit = {
        "timestamp": ts,
        "backup": str(backup),
        "targets": len(targets),
        "filled_this_pass": filled,
        "deepseek_filled": ds_filled,
        "have_after": have,
        "still_npd_or_empty": need,
        "still_sample": still[:40],
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print("backup", backup)
    print("filled_this_pass", filled)
    print("have_after", have, "still", need)
    print("audit", AUDIT)


if __name__ == "__main__":
    main()
