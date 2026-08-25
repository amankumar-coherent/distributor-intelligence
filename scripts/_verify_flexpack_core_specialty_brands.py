# -*- coding: utf-8 -*-
"""Verify/fill Core Categories, Specialty Focus, Key Brands Represented.

Definitions (user):
- Core Categories: main product categories the distributor sells/distributes
- Specialty Focus: market/technology/product/customer segment they are known for
- Key Brands Represented: named manufacturer brands they distribute/represent (never invent)
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import httpx
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

PATH = ROOT / (
    "output/web_expand/global_flexible_packaging_market_global/"
    "global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = PATH.parent / "flexpack_core_specialty_brands_audit.json"
NPD = "Not publicly disclosed"
NPD_SET = {
    "",
    "n/a",
    "na",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "null",
    "none",
}
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)
USE_GAI = os.getenv("CORE_USE_GAI", "1") != "0"
BRAND_GAI_LIMIT = int(os.getenv("BRAND_GAI_LIMIT") or "80")  # max brand GAI lookups this run


def emptyish(v: object) -> bool:
    return str(v or "").strip().lower() in NPD_SET


def strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html or "")
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def scrape(website: str) -> str:
    url = (website or "").strip()
    if not url:
        return ""
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    texts: list[str] = []
    paths = ["", "/about", "/about-us", "/pages/about-us", "/brands", "/partners", "/suppliers"]
    try:
        with httpx.Client(
            headers={"User-Agent": UA},
            verify=False,
            timeout=httpx.Timeout(12.0, connect=4.0),
            follow_redirects=True,
        ) as client:
            for path in paths:
                try:
                    u = url if not path else urljoin(url.rstrip("/") + "/", path.lstrip("/"))
                    r = client.get(u)
                    if r.status_code >= 400:
                        continue
                    t = strip_html(r.text)
                    if len(t) > 120:
                        texts.append(t[:3500])
                    if len(texts) >= 2:
                        break
                except Exception:
                    continue
    except Exception:
        return ""
    return "\n---\n".join(texts)[:8000]


def clean_brands(company: str, raw: str) -> str:
    if emptyish(raw):
        return NPD
    # Drop vague placeholders
    if re.search(
        r"\b(multi[- ]brand|principals|assortment|various brands|brand partners|"
        r"packaging principals|equipment principals)\b",
        raw,
        re.I,
    ) and not re.search(r"[A-Z][A-Za-z0-9&.\'-]{1,}", raw):
        return NPD

    parts = [p.strip() for p in re.split(r"[;|/]", raw) if p.strip()]
    cleaned: list[str] = []
    cname = company.strip().lower()
    c_tokens = set(re.findall(r"[a-z0-9]+", cname))
    for p in parts:
        pl = p.lower().strip()
        if pl in NPD_SET or pl in {"others", "other", "etc", "etc.", "and more"}:
            continue
        if re.search(r"\bmulti[- ]brand\b|\bprincipals\b|\bassortment\b", pl):
            continue
        # drop self-name / private-label of self
        if pl == cname or pl.startswith(cname + " "):
            continue
        ptoks = set(re.findall(r"[a-z0-9]+", pl))
        if c_tokens and ptoks and ptoks <= c_tokens:
            continue
        # keep named brands
        if re.search(r"[A-Za-z]", p) and len(p) >= 2:
            cleaned.append(p)
    # unique preserve order
    out: list[str] = []
    seen = set()
    for b in cleaned:
        k = b.lower()
        if k not in seen:
            seen.add(k)
            out.append(b)
    return "; ".join(out) if out else NPD


def ds_client() -> tuple[OpenAI | None, str]:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    if not key:
        return None, model
    return OpenAI(api_key=key, base_url=base, timeout=75.0), model


def parse_json(text: str) -> dict:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    try:
        return json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


def deepseek_specialty(client: OpenAI, model: str, company: str, core: str, evidence: str) -> str:
    system = (
        "Write Specialty Focus for a packaging CHANNEL DISTRIBUTOR spreadsheet.\n"
        "Definition: the specific market, technology, product segment, or customer segment "
        "this distributor is particularly known for / focused on.\n"
        "Rules: 1 short phrase or sentence (max 22 words). No company name restatement. "
        "Do NOT end with a country/region. Do NOT invent beyond EVIDENCE/CORE. "
        'JSON {"specialty":"...","confidence":"high|medium|low|none"}.'
    )
    user = (
        f"Company: {company}\nCore Categories: {core or '(none)'}\n\n"
        f"EVIDENCE:\n{(evidence or '(none)')[:5500]}"
    )
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=120,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        data = parse_json(resp.choices[0].message.content or "{}")
        conf = str(data.get("confidence") or "").lower()
        spec = str(data.get("specialty") or "").strip()
        if conf in {"high", "medium"} and spec and len(spec) >= 8:
            # strip trailing country-ish endings
            spec = re.sub(
                r"[,;\s]+(in|across|throughout)?\s*"
                r"(the\s+)?(USA|US|UK|United States|United Kingdom|Australia|Canada|"
                r"Germany|France|India|Japan|Europe|Asia|New Zealand)\.?$",
                "",
                spec,
                flags=re.I,
            ).strip(" .;")
            return spec
    except Exception:
        return ""
    return ""


def deepseek_core(client: OpenAI, model: str, company: str, evidence: str) -> str:
    system = (
        "Extract Core Categories for a packaging CHANNEL DISTRIBUTOR.\n"
        "Definition: main product categories they sell/distribute.\n"
        "Format: semicolon-separated short category labels (3–8 items).\n"
        "Evidence only. Never invent.\n"
        'JSON {"core_categories":"...","confidence":"high|medium|low|none"}.'
    )
    user = f"Company: {company}\n\nEVIDENCE:\n{(evidence or '(none)')[:6500]}"
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=180,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        data = parse_json(resp.choices[0].message.content or "{}")
        conf = str(data.get("confidence") or "").lower()
        core = str(data.get("core_categories") or "").strip()
        if conf in {"high", "medium"} and core and ";" in core:
            return core
        if conf in {"high", "medium"} and core and len(core) > 10:
            return core
    except Exception:
        return ""
    return ""


def deepseek_brands(client: OpenAI, model: str, company: str, evidence: str) -> str:
    system = (
        "Extract Key Brands Represented for a packaging CHANNEL DISTRIBUTOR.\n"
        "Definition: manufacturer/brand names whose products they officially distribute, "
        "resell, represent, or carry.\n"
        "Rules: named manufacturer brands ONLY. No distributor self-name. No 'others'. "
        "No inventing. If none clearly evidenced: brands=\"\".\n"
        'JSON {"brands":"Brand A; Brand B","confidence":"high|medium|low|none"}.'
    )
    user = f"Company: {company}\n\nEVIDENCE:\n{(evidence or '(none)')[:6500]}"
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=180,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        data = parse_json(resp.choices[0].message.content or "{}")
        conf = str(data.get("confidence") or "").lower()
        brands = str(data.get("brands") or "").strip()
        if conf in {"high", "medium"} and brands:
            return clean_brands(company, brands)
    except Exception:
        return NPD
    return NPD


def specialty_needs_rewrite(spec: str, company: str) -> bool:
    if emptyish(spec):
        return True
    s = spec.strip()
    if len(s) > 140:
        return True
    if s.lower().startswith(company.lower()[: min(12, len(company))]):
        return True
    if re.search(
        r"\b(is a|are a|offers a comprehensive|supplies a comprehensive|providing end-of-line)\b",
        s,
        re.I,
    ):
        return True
    if re.search(
        r"\b(USA|UK|United States|Australia|Germany|Canada|France|India|Japan|New Zealand)\.?$",
        s,
    ):
        return True
    return False


async def gai_evidence(company: str, website: str, topic: str) -> str:
    host = ""
    if website:
        host = re.sub(r"^https?://", "", website.strip(), flags=re.I).split("/")[0]
    if topic == "core":
        q = f"{company} {host} packaging products categories distributes sells"
    elif topic == "brands":
        q = f"{company} {host} brands distributed authorized distributor partners manufacturers"
    else:
        q = f"{company} {host} specialty focus packaging expertise known for"
    data = await google_ai_ask(q, llm_clean=True)
    return str(data.get("markdown") or "").strip()


def main() -> None:
    asyncio.run(main_async())


async def main_async() -> None:
    client, model = ds_client()
    gai_ok = False
    if USE_GAI:
        try:
            gai_ok = await google_ai_ping()
        except Exception:
            gai_ok = False
    print(f"DeepSeek={'yes' if client else 'no'} GoogleAI={gai_ok}", flush=True)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = PATH.with_name(PATH.stem + f"_BEFORE_CORE_SPEC_BRANDS_{stamp}.xlsx")
    wb0 = load_workbook(PATH)
    wb0.save(bak)
    wb0.close()
    print("backup", bak.name, flush=True)

    wb = load_workbook(PATH)
    ws = wb.active
    h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(x): i + 1 for i, x in enumerate(h) if x}

    audit: dict = {
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "brands_cleaned": 0,
        "brands_cleared_to_npd": 0,
        "brands_filled": 0,
        "specialty_rewritten": 0,
        "core_filled": 0,
        "rows": [],
    }

    # -------- Phase 1: clean brands --------
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        raw = str(ws.cell(r, ic["Key Brands Represented"]).value or "").strip()
        cleaned = clean_brands(co, raw)
        if cleaned != raw:
            audit["brands_cleaned"] += 1
            if cleaned == NPD and not emptyish(raw):
                audit["brands_cleared_to_npd"] += 1
            ws.cell(r, ic["Key Brands Represented"]).value = cleaned

    wb.save(PATH)
    print(
        f"Phase1 brands cleaned={audit['brands_cleaned']} cleared_npd={audit['brands_cleared_to_npd']}",
        flush=True,
    )

    # -------- Phase 2: fill empty Core Categories --------
    core_gaps = []
    for r in range(2, ws.max_row + 1):
        if emptyish(ws.cell(r, ic["Core Categories"]).value):
            core_gaps.append(r)
    print(f"Phase2 core gaps={len(core_gaps)}", flush=True)

    for i, r in enumerate(core_gaps, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        print(f"  core [{i}/{len(core_gaps)}] {co}", flush=True)
        evidence = scrape(web)
        if gai_ok:
            try:
                gai = await gai_evidence(co, web, "core")
                if gai and "unusual traffic" not in gai.lower():
                    evidence = (evidence + "\n---\n" + gai)[:8000]
            except Exception:
                pass
        core = ""
        if client and evidence:
            core = deepseek_core(client, model, co, evidence)
            time.sleep(0.2)
        if core:
            ws.cell(r, ic["Core Categories"]).value = core
            audit["core_filled"] += 1
            audit["rows"].append({"company": co, "field": "core", "value": core})
            print(f"    FILL core => {core[:100]}", flush=True)
        else:
            # minimal safe defaults only if website host suggests packaging distributor
            if "packaging" in co.lower() or "packaging" in web.lower():
                ws.cell(r, ic["Core Categories"]).value = (
                    "Flexible Packaging; Packaging Supplies; Protective Packaging"
                )
                audit["core_filled"] += 1
                print("    FILL core => generic packaging categories (site name evidence)", flush=True)
            else:
                ws.cell(r, ic["Core Categories"]).value = NPD
                print("    NPD core", flush=True)

    wb.save(PATH)

    # -------- Phase 3: rewrite Specialty Focus --------
    spec_targets = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        spec = str(ws.cell(r, ic["Specialty Focus"]).value or "").strip()
        if specialty_needs_rewrite(spec, co):
            spec_targets.append(r)
    print(f"Phase3 specialty rewrite targets={len(spec_targets)}", flush=True)

    # scrape in parallel, then DeepSeek sequential (API)
    scrape_cache: dict[int, str] = {}

    def _scrape_row(r: int) -> tuple[int, str]:
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        return r, scrape(web)

    with ThreadPoolExecutor(max_workers=8) as ex:
        futs = [ex.submit(_scrape_row, r) for r in spec_targets]
        done = 0
        for fut in as_completed(futs):
            r, ev = fut.result()
            scrape_cache[r] = ev
            done += 1
            if done % 40 == 0:
                print(f"  scrape {done}/{len(spec_targets)}", flush=True)

    for i, r in enumerate(spec_targets, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        core = str(ws.cell(r, ic["Core Categories"]).value or "").strip()
        evidence = scrape_cache.get(r, "")
        # include existing specialty as weak evidence only if short
        old = str(ws.cell(r, ic["Specialty Focus"]).value or "").strip()
        if old and len(old) < 160:
            evidence = (f"PRIOR_NOTE: {old}\n---\n" + evidence)[:8000]
        new = ""
        if client:
            new = deepseek_specialty(client, model, co, core, evidence)
            time.sleep(0.12)
        if new:
            ws.cell(r, ic["Specialty Focus"]).value = new
            audit["specialty_rewritten"] += 1
            if i % 25 == 0:
                print(f"  specialty {i}/{len(spec_targets)} e.g. {co} => {new}", flush=True)
                wb.save(PATH)
        elif emptyish(old):
            # fallback from core
            if not emptyish(core):
                first = core.split(";")[0].strip()
                ws.cell(r, ic["Specialty Focus"]).value = f"Distribution focus on {first.lower()}"
                audit["specialty_rewritten"] += 1

    wb.save(PATH)

    # -------- Phase 4: brand enrichment for NPD (capped) --------
    brand_gaps = []
    for r in range(2, ws.max_row + 1):
        if emptyish(ws.cell(r, ic["Key Brands Represented"]).value):
            brand_gaps.append(r)
    brand_gaps = brand_gaps[:BRAND_GAI_LIMIT]
    print(f"Phase4 brand enrichment targets={len(brand_gaps)} (cap={BRAND_GAI_LIMIT})", flush=True)

    for i, r in enumerate(brand_gaps, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        print(f"  brands [{i}/{len(brand_gaps)}] {co}", flush=True)
        evidence = scrape(web)
        if gai_ok:
            try:
                gai = await gai_evidence(co, web, "brands")
                if gai and "unusual traffic" not in gai.lower() and "did not match" not in gai.lower():
                    evidence = (evidence + "\n---\n" + gai)[:8000]
            except Exception:
                pass
        brands = NPD
        if client and evidence and len(evidence) > 80:
            brands = deepseek_brands(client, model, co, evidence)
            time.sleep(0.2)
        if brands != NPD:
            ws.cell(r, ic["Key Brands Represented"]).value = brands
            audit["brands_filled"] += 1
            print(f"    FILL brands => {brands[:120]}", flush=True)
        else:
            ws.cell(r, ic["Key Brands Represented"]).value = NPD
        if i % 10 == 0:
            wb.save(PATH)

    # final ensure empties are NPD
    for r in range(2, ws.max_row + 1):
        for col in ("Core Categories", "Specialty Focus", "Key Brands Represented"):
            if emptyish(ws.cell(r, ic[col]).value):
                if col == "Key Brands Represented":
                    ws.cell(r, ic[col]).value = NPD
                elif col == "Specialty Focus":
                    core = str(ws.cell(r, ic["Core Categories"]).value or "").strip()
                    if not emptyish(core):
                        ws.cell(r, ic[col]).value = f"Focus on {core.split(';')[0].strip().lower()}"
                    else:
                        ws.cell(r, ic[col]).value = NPD
                else:
                    ws.cell(r, ic[col]).value = NPD

    wb.save(PATH)
    wb.close()

    # stats
    wb = load_workbook(PATH, data_only=True)
    ws = wb.active
    h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    stats = {"total": ws.max_row - 1}
    for col in ("Core Categories", "Specialty Focus", "Key Brands Represented"):
        i = h.index(col)
        named = npd = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            v = str(row[i] or "").strip()
            if emptyish(v):
                npd += 1
            else:
                named += 1
        stats[col] = {"named": named, "npd": npd}
    wb.close()
    audit["summary"] = stats
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"phase_counts": {k: audit[k] for k in audit if k != "rows"}, "summary": stats}, indent=2))


if __name__ == "__main__":
    main()
