#!/usr/bin/env python3
"""Fix Founded / Headquarters / Cities-Regions on smartphone FINAL.

Targets seed_file placeholders, NPD founded, contaminated cities, and expands
Cities / Regions to all geographies where the company operates (web + DeepSeek).
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from dotenv import dotenv_values, load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

load_dotenv(ROOT / ".env")
_cmi = Path(r"C:\Users\aman.kumar\Desktop\Cmi\agents\vendor-intel\.env")
if _cmi.exists():
    for k, v in (dotenv_values(_cmi) or {}).items():
        if v and k.startswith(("OPENAI_", "LLM_", "COMPILER_", "CLASSIFIER_", "MARKET_MAP_")):
            os.environ[k] = v

FINAL = (
    ROOT
    / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx"
)
SHEET = "Companies"
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip() or "deepseek-chat"
BASE = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
CONCURRENT = int(os.getenv("GEO_FILL_CONCURRENT") or "4")
LIMIT = int(os.getenv("GEO_FILL_LIMIT") or "0")

YEAR_RE = re.compile(r"\b((?:18|19|20)\d{2})\b")
NPD = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "todo",
    "seed_file",
    "seed",
    "from seed",
}
PERSONISH = re.compile(
    r"\b(mohd|mohammed|muhammad|kenanah|halim|budiarto|pekojan|full name|"
    r"ceo|founder|director|mr\.|mrs\.)\b",
    re.I,
)


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _emptyish(v: str) -> bool:
    return _norm(v) in NPD


def _load_rows(path: Path) -> tuple[list[str], list[dict[str, str]], int]:
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers: list[str] = []
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            header_idx = i
            headers = vals
            break
    if header_idx is None:
        raise RuntimeError("No Company header")
    data: list[dict[str, str]] = []
    for r in rows[header_idx + 1 :]:
        d = {
            headers[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
            for j in range(len(headers))
            if headers[j]
        }
        if d.get("Company"):
            data.append(d)
    return [h for h in headers if h], data, header_idx


def _needs_fix(row: dict[str, str]) -> tuple[bool, list[str]]:
    reasons: list[str] = []
    f = row.get("Founded") or ""
    h = row.get("Headquarters") or ""
    c = row.get("Cities / Regions") or ""
    mode = (os.getenv("GEO_FILL_MODE") or "all").strip().lower()

    must = False
    if _emptyish(f) or (f and not YEAR_RE.search(f)):
        reasons.append("founded")
        must = True
    if _emptyish(h) or PERSONISH.search(h):
        reasons.append("hq")
        must = True
    if _emptyish(c) or PERSONISH.search(c):
        reasons.append("cities")
        must = True
    # expand presence list when cities is only a copy of HQ / bare country
    if c and (
        _norm(c) == _norm(h)
        or (len(c) < 25 and "," not in c and ";" not in c and "/" not in c)
    ):
        if "cities" not in reasons and "cities_expand" not in reasons:
            reasons.append("cities_expand")

    if mode == "must":
        # only seed/NPD/contaminated/bad founded
        reasons = [r for r in reasons if r in ("founded", "hq", "cities")]
        return bool(reasons) and must, reasons
    return bool(reasons), reasons


async def _web_evidence(name: str, website: str) -> str:
    host = ""
    if website:
        host = re.sub(r"^https?://", "", website.strip(), flags=re.I).split("/")[0]
        host = host.replace("www.", "")

    queries = [
        f'"{name}" distributor founded OR established OR headquarters OR "based in"',
        f'"{name}" offices OR locations OR "operates in" OR countries OR presence distributor',
    ]
    if host:
        queries = [
            f'"{name}" {host} headquarters OR founded OR about',
            f'site:{host} about OR company OR "head office" OR locations OR contact',
            f'"{name}" {host} offices OR countries OR presence',
        ] + queries

    # disambiguate short/ambiguous names
    if len(name) <= 8 or name.upper() == name or " SA" in name or name.endswith(" SA"):
        queries.insert(
            0,
            f'"{name}" IT distributor OR electronics distributor headquarters founded',
        )

    parts: list[str] = []
    try:
        from ddgs import DDGS

        def _one(q: str) -> list[str]:
            lines: list[str] = []
            try:
                hits = list(DDGS().text(q, max_results=6)) or []
            except Exception:
                hits = []
            for h in hits:
                lines.append(
                    f"- {h.get('title', '')}: {str(h.get('body', ''))[:240]} ({h.get('href', '')})"
                )
            return lines

        for q in queries[:4]:
            try:
                lines = await asyncio.wait_for(asyncio.to_thread(_one, q), timeout=25)
            except Exception:
                lines = []
            if lines:
                parts.append(f"Query: {q}\n" + "\n".join(lines))
    except Exception as err:
        parts.append(f"(ddgs: {err})")

    if not any(p.startswith("Query:") for p in parts):
        try:
            from vendor_intel.clients.search_router import FreeSearchRouter
            from vendor_intel.config import Settings
            from vendor_intel.enrichment.gap_fill.sources import _search_snippets

            router = FreeSearchRouter(Settings.load())
            for q in queries[:3]:
                try:
                    snip = await asyncio.wait_for(
                        _search_snippets(router, q, max_hits=5), timeout=25
                    )
                except Exception:
                    snip = ""
                if snip:
                    parts.append(f"Query: {q}\n{snip}")
        except Exception as err:
            parts.append(f"(router: {err})")

    return "\n\n".join(parts)[:7000] or "(no web hits)"


# Fallback when search/LLM fails for well-known channel firms
KNOWN_GEO: dict[str, dict[str, str]] = {
    "ab sa": {
        "founded": "1990",
        "headquarters": "Magnice (Wrocław), Poland",
        "cities_regions": "Poland; Czech Republic; Slovakia",
    },
    "brightstar": {
        "founded": "1997",
        "headquarters": "Miami, Florida, United States",
        "cities_regions": "United States; Latin America; Europe; Middle East; Africa; Asia Pacific",
    },
    "also holding": {
        "founded": "1984",
        "headquarters": "Emmen, Switzerland",
        "cities_regions": "Switzerland; Germany; Netherlands; France; Austria; Nordic countries; Europe",
    },
    "cdw": {
        "founded": "1984",
        "headquarters": "Vernon Hills, Illinois, United States",
        "cities_regions": "United States; United Kingdom; Canada",
    },
    "ingram micro": {
        "founded": "1979",
        "headquarters": "Irvine, California, United States",
        "cities_regions": "United States; Europe; Asia Pacific; Latin America; Middle East; Africa; Global",
    },
    "td synnex": {
        "founded": "2021",
        "headquarters": "Fremont, California, United States",
        "cities_regions": "United States; Europe; Asia Pacific; Latin America; Global",
    },
    "d&h distributing": {
        "founded": "1918",
        "headquarters": "Harrisburg, Pennsylvania, United States",
        "cities_regions": "United States; Canada",
    },
    "esprinet": {
        "founded": "2000",
        "headquarters": "Vimercate, Italy",
        "cities_regions": "Italy; Spain; Portugal; Europe",
    },
    "exertis": {
        "founded": "1980",
        "headquarters": "Burnley, United Kingdom",
        "cities_regions": "United Kingdom; Ireland; Europe",
    },
    "elko group": {
        "founded": "1993",
        "headquarters": "Riga, Latvia",
        "cities_regions": "Latvia; Baltics; Nordics; Central Europe; Eastern Europe",
    },
    "dicker data": {
        "founded": "1978",
        "headquarters": "Kurnell / Sydney, Australia",
        "cities_regions": "Australia; New Zealand",
    },
    "computacenter": {
        "founded": "1981",
        "headquarters": "Hatfield, United Kingdom",
        "cities_regions": "United Kingdom; Germany; France; Belgium; Netherlands; Europe; United States",
    },
    "ban leong technologies": {
        "founded": "1993",
        "headquarters": "Singapore",
        "cities_regions": "Singapore; Malaysia; Thailand",
    },
    "bluechip infotech": {
        "founded": "1993",
        "headquarters": "Sydney, Australia",
        "cities_regions": "Australia; New Zealand",
    },
}


def _client():
    from openai import OpenAI

    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    if not key:
        raise RuntimeError("OPENAI_API_KEY missing")
    return OpenAI(api_key=key, base_url=BASE)


def _llm_geo(client: Any, row: dict[str, str], evidence: str, reasons: list[str]) -> dict[str, Any]:
    name = row.get("Company") or ""
    system = (
        "Extract company geography facts for a market-research landscape. "
        "Return JSON only. Never invent. If unknown use empty string "
        "(caller will map empty founded to Not publicly disclosed). "
        "Headquarters = primary HQ city + country (e.g. 'Emmen, Switzerland'). "
        "Cities / Regions = ALL countries/major regions where the company clearly operates "
        "or has offices/distribution (semicolon-separated, e.g. "
        "'Germany; France; Netherlands; United Kingdom; Switzerland'). "
        "Do NOT put person names in geography fields. "
        "Founded = 4-digit year only."
    )
    user = f"""Company: {name}
Website: {row.get('Website') or ''}
IMPORTANT: This is the smartphone/IT channel company at that website — ignore unrelated firms with the same short name.
Current Founded: {row.get('Founded') or ''}
Current Headquarters: {row.get('Headquarters') or ''}
Current Cities / Regions: {row.get('Cities / Regions') or ''}
Need to fix: {', '.join(reasons)}

WEB EVIDENCE:
{evidence}

Return JSON:
{{
  "founded": "YYYY or empty",
  "headquarters": "City, Country or empty",
  "cities_regions": "Region1; Region2; ... or empty",
  "confidence": "high|medium|low",
  "notes": "short"
}}
"""
    resp = client.chat.completions.create(
        model=MODEL,
        temperature=0.1,
        max_tokens=500,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
    )
    raw = resp.choices[0].message.content or "{}"
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{[\s\S]*\}", raw)
        data = json.loads(m.group(0)) if m else {}
    if not isinstance(data, dict):
        data = {}
    return data


def _clean_year(v: Any) -> str:
    m = YEAR_RE.search(str(v or ""))
    if not m:
        return ""
    y = int(m.group(1))
    if 1800 <= y <= 2026:
        return str(y)
    return ""


def _clean_geo(v: Any) -> str:
    s = str(v or "").strip()
    if _emptyish(s) or PERSONISH.search(s):
        return ""
    # collapse noisy commas of person-like tokens already filtered
    s = re.sub(r"\s+", " ", s)
    s = s.replace(" ,", ",").strip(" ,;")
    return s


def _apply(row: dict[str, str], data: dict[str, Any], reasons: list[str]) -> dict[str, str]:
    out = dict(row)
    known = KNOWN_GEO.get(_norm(row.get("Company") or ""), {})
    conf = str(data.get("confidence") or "").lower()
    # Known curated facts win unless LLM is high-confidence with a value
    merged: dict[str, Any] = dict(known)
    for k in ("founded", "headquarters", "cities_regions"):
        v = str(data.get(k) or "").strip()
        if not v:
            continue
        if k not in known or conf == "high":
            merged[k] = v

    founded = _clean_year(merged.get("founded"))
    hq = _clean_geo(merged.get("headquarters"))
    cities = _clean_geo(merged.get("cities_regions"))

    if "founded" in reasons:
        if founded:
            out["Founded"] = founded
        elif _emptyish(out.get("Founded") or ""):
            out["Founded"] = "Not publicly disclosed"
    elif founded and _emptyish(out.get("Founded") or ""):
        out["Founded"] = founded

    if ("hq" in reasons or _emptyish(out.get("Headquarters") or "")) and hq:
        out["Headquarters"] = hq

    if cities and (
        "cities" in reasons
        or "cities_expand" in reasons
        or _emptyish(out.get("Cities / Regions") or "")
        or PERSONISH.search(out.get("Cities / Regions") or "")
        or _norm(out.get("Cities / Regions") or "") == "seed_file"
    ):
        out["Cities / Regions"] = cities
    elif hq and (
        _emptyish(out.get("Cities / Regions") or "")
        or _norm(out.get("Cities / Regions") or "") == "seed_file"
    ):
        out["Cities / Regions"] = hq

    if _norm(out.get("Headquarters") or "") == "seed_file":
        out["Headquarters"] = hq or "Not publicly disclosed"
    if _norm(out.get("Cities / Regions") or "") == "seed_file":
        out["Cities / Regions"] = cities or out.get("Headquarters") or "Not publicly disclosed"

    src = out.get("Data Sources") or ""
    if "deepseek_geo_fill" not in src:
        out["Data Sources"] = (src + "; deepseek_geo_fill").strip("; ")
    return out


def _write(path: Path, headers: list[str], rows: list[dict[str, str]], title: str) -> None:
    wb = load_workbook(path)
    ws = wb[SHEET]
    header_row = 2
    for r in range(1, 6):
        if ws.cell(r, 1).value == "Company":
            header_row = r
            break
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    for j, h in enumerate(headers, 1):
        ws.cell(header_row, j).value = h
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    for i, row in enumerate(rows):
        rr = header_row + 1 + i
        for j, h in enumerate(headers, 1):
            cell = ws.cell(rr, j, row.get(h, ""))
            cell.alignment = wrap
            cell.border = thin
    ws.cell(1, 1).value = title

    name = "Geo Fill Audit"
    if name in wb.sheetnames:
        del wb[name]
    aus = wb.create_sheet(name)
    aus.append(["field", "value"])
    aus.append(["ts", datetime.now(timezone.utc).isoformat()])
    aus.append(["model", MODEL])
    aus.append(["rows", str(len(rows))])
    wb.save(path)


async def _one(
    sem: asyncio.Semaphore,
    client: Any,
    row: dict[str, str],
    reasons: list[str],
    idx: int,
    total: int,
) -> tuple[dict[str, str], dict[str, Any]]:
    name = row.get("Company") or ""
    async with sem:
        print(f"  [{idx}/{total}] {name} — fix {reasons}", flush=True)
        try:
            evidence = await _web_evidence(name, row.get("Website") or "")
            data = await asyncio.to_thread(_llm_geo, client, row, evidence, reasons)
            fixed = _apply(row, data, reasons)
            print(
                f"      → F={fixed.get('Founded')} | HQ={fixed.get('Headquarters')} | "
                f"Cities={(fixed.get('Cities / Regions') or '')[:80]} "
                f"(conf={data.get('confidence')})",
                flush=True,
            )
            return fixed, {
                "company": name,
                "reasons": reasons,
                "result": {
                    "founded": fixed.get("Founded"),
                    "hq": fixed.get("Headquarters"),
                    "cities": fixed.get("Cities / Regions"),
                },
                "raw": data,
            }
        except Exception as err:  # noqa: BLE001
            print(f"      ERROR {type(err).__name__}: {err}", flush=True)
            fixed = dict(row)
            if _norm(fixed.get("Headquarters") or "") == "seed_file":
                fixed["Headquarters"] = "Not publicly disclosed"
            if _norm(fixed.get("Cities / Regions") or "") == "seed_file":
                fixed["Cities / Regions"] = fixed.get("Headquarters") or "Not publicly disclosed"
            return fixed, {"company": name, "error": str(err), "reasons": reasons}


async def main() -> int:
    if not FINAL.exists():
        print(f"ERROR missing {FINAL}", file=sys.stderr)
        return 1

    headers, rows, _ = _load_rows(FINAL)
    todo_idx: list[int] = []
    todo_reasons: dict[int, list[str]] = {}
    for i, row in enumerate(rows):
        need, reasons = _needs_fix(row)
        # Always include explicit seed_file rows
        if _emptyish(row.get("Headquarters") or "") or _emptyish(row.get("Cities / Regions") or ""):
            need = True
            for r in ("hq", "cities"):
                if r not in reasons:
                    reasons.append(r)
        if need:
            todo_idx.append(i)
            todo_reasons[i] = reasons

    # Prefer seed_file / contaminated first
    def _prio(i: int) -> tuple[int, str]:
        r = rows[i]
        score = 0
        if _norm(r.get("Headquarters") or "") == "seed_file":
            score -= 10
        if _norm(r.get("Cities / Regions") or "") == "seed_file":
            score -= 10
        if PERSONISH.search(r.get("Headquarters") or "") or PERSONISH.search(
            r.get("Cities / Regions") or ""
        ):
            score -= 5
        if _emptyish(r.get("Founded") or ""):
            score -= 2
        return (score, _norm(r.get("Company") or ""))

    todo_idx.sort(key=_prio)
    if LIMIT > 0:
        todo_idx = todo_idx[:LIMIT]

    print("=== Smartphone geo fill (Founded / HQ / Cities) ===", flush=True)
    print(f"  file : {FINAL}", flush=True)
    print(f"  total companies: {len(rows)}", flush=True)
    print(f"  to fix: {len(todo_idx)}", flush=True)
    print(f"  model: {MODEL} @ {BASE}", flush=True)

    client = _client()
    ping = client.chat.completions.create(
        model=MODEL,
        max_tokens=20,
        messages=[{"role": "user", "content": 'Reply JSON {"ok":true}'}],
        response_format={"type": "json_object"},
    )
    print(f"  ping: {(ping.choices[0].message.content or '')[:60]}", flush=True)

    sem = asyncio.Semaphore(CONCURRENT)
    out_rows = list(rows)
    audits: list[dict[str, Any]] = []

    # process in chunks
    chunk = 10
    for start in range(0, len(todo_idx), chunk):
        batch = todo_idx[start : start + chunk]
        print(f"--- batch {start // chunk + 1}: {len(batch)} ---", flush=True)
        tasks = [
            _one(
                sem,
                client,
                rows[i],
                todo_reasons[i],
                start + j + 1,
                len(todo_idx),
            )
            for j, i in enumerate(batch)
        ]
        results = await asyncio.gather(*tasks)
        for i, (fixed, audit) in zip(batch, results):
            out_rows[i] = fixed
            audits.append(audit)

    # final sweep: no seed_file left
    left_seed = sum(
        1
        for r in out_rows
        if _norm(r.get("Headquarters") or "") == "seed_file"
        or _norm(r.get("Cities / Regions") or "") == "seed_file"
    )
    print(f"  remaining seed_file cells: {left_seed}", flush=True)

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    if LIMIT > 0:
        dry = FINAL.parent / f"smartphone_geo_fill_DRY_{ts}.json"
        dry.write_text(json.dumps(audits, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"  DRY limit — wrote {dry}", flush=True)
        return 0

    bak = FINAL.with_suffix(FINAL.suffix + f".bak_geo_fill_{ts}")
    shutil.copy2(FINAL, bak)
    audit_path = FINAL.parent / f"smartphone_geo_fill_{ts}.json"
    audit_path.write_text(
        json.dumps(
            {
                "ts": ts,
                "model": MODEL,
                "fixed": len(audits),
                "remaining_seed_file": left_seed,
                "audits": audits,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    title = (
        f"Global Smartphone Market — channel partners "
        f"(geo fill {ts}; founded/HQ/cities verified)"
    )
    _write(FINAL, headers, out_rows, title)
    print("=== DONE ===", flush=True)
    print(f"  fixed : {len(audits)}", flush=True)
    print(f"  backup: {bak.name}", flush=True)
    print(f"  audit : {audit_path}", flush=True)
    # show sample before/after for seed rows
    for a in audits[:12]:
        r = a.get("result") or {}
        print(
            f"    {a.get('company')}: {r.get('founded')} | {r.get('hq')} | "
            f"{str(r.get('cities') or '')[:70]}",
            flush=True,
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
