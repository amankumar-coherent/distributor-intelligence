"""Append agent-verified packaging distributors mid-batch (safe if domains not yet present)."""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "agent_manual_verified_add_audit2.json"
ND = "Not publicly disclosed"

ROWS = [
    {
        "Company": "Fastpak Systems",
        "Website": "https://fastpaksystems.com",
        "Founded": "1986",
        "Headquarters": "1450 Williams Dr NE, Marietta, Georgia 30066, USA",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "Southeast US (Marietta GA HQ; Jacksonville FL; fulfillment Buford GA)",
        "Ownership": "Private",
        "Employees": "11–50",
        "Core Categories": "Industrial packaging supplies; stretch wrap; strapping; shrink; packaging equipment; fulfillment",
        "Specialty Focus": "Value-added packaging distributor for manufacturers and distributors (PDA member)",
        "Key Brands Represented": "Multi-brand packaging equipment and consumables",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": "https://www.linkedin.com/company/fastpak-systems-inc-",
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "Southeast US packaging distributor founded 1986; PDA member; HQ Marietta GA.",
        "_evidence": "pdachain.com welcome Fastpak founded 1986; fastpaksystems.com distributor",
    },
    {
        "Company": "Sigma Supply of North America",
        "Website": "https://www.sigmasupply.com",
        "Founded": "1970",
        "Headquarters": "Hot Springs, Arkansas, USA",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "United States (coast-to-coast packaging engineers/technicians noted)",
        "Ownership": "Private (WBENC women-owned; part of BradyPLUS network noted publicly)",
        "Employees": ND,
        "Core Categories": "Packaging solutions; packaging equipment; technical support; shipping supplies",
        "Specialty Focus": "Packaging distributor with equipment sales, technical support, and large SKU online assortment",
        "Key Brands Represented": "Multi-brand packaging assortment",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "US packaging distributor founded 1970 in Hot Springs AR; WBENC; packaging equipment and supplies.",
        "_evidence": "sigmasupply.com/about/timeline founded 1970 Martha and Dan Hamby",
    },
    {
        "Company": "SupplyOne Canada (Crownhill Packaging)",
        "Website": "https://www.crownhillpackaging.com",
        "Founded": "1986",
        "Headquarters": "8905 Goreway Drive, Brampton, Ontario, Canada",
        "Continent / Geography": "NA(Canada)",
        "Operational Presence": "North America (Toronto-area facilities; PDA network access)",
        "Ownership": "Private (SupplyOne Canada; formerly Crownhill Packaging)",
        "Employees": ND,
        "Core Categories": "Corrugated boxes; bubble wrap; foam; stretch/shrink film; tape; mailers; void fill; packaging equipment",
        "Specialty Focus": "Full-service packaging distributor and custom packaging solutions supplier",
        "Key Brands Represented": "Multi-supplier packaging materials",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "CA",
        "Region Code": "NA",
        "Summary": "Canadian packaging distributor (Crownhill/SupplyOne Canada) founded ~1986; Brampton ON; PDA member history.",
        "_evidence": "crownhillpackaging.com about — over 35 years; PDA partnership page founded Canada 1986",
    },
]

HEADERS = list(ROWS[0].keys())
HEADERS = [h for h in HEADERS if not h.startswith("_")]


def domain_of(url: str) -> str:
    try:
        return urlparse(url if "://" in url else "https://" + url).netloc.lower().replace("www.", "")
    except Exception:
        return ""


def main():
    backup = FP.with_name(FP.stem + f"_pre_agent_manual2_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(FP, backup)
    wb = load_workbook(FP)
    ws = wb.active
    names, domains = set(), set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        names.add((r[0] or "").strip().lower())
        domains.add(domain_of(str(r[1] or "")))
    added, skipped = [], []
    for row in ROWS:
        d = domain_of(row["Website"])
        if row["Company"].lower() in names or d in domains:
            skipped.append(row["Company"])
            continue
        ws.append([row[h] for h in HEADERS])
        names.add(row["Company"].lower())
        domains.add(d)
        rec = {k: row[k] for k in HEADERS}
        rec["_evidence"] = row.get("_evidence")
        added.append(rec)
    for col in ws.columns:
        for cell in col:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[col[0].column_letter].width = 22
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(FP)
    audit = {
        "ts": datetime.now().isoformat(),
        "backup": str(backup),
        "added": added,
        "skipped": skipped,
        "final_rows": ws.max_row - 1,
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print("added", len(added), "skipped", skipped, "final", audit["final_rows"])


if __name__ == "__main__":
    main()
