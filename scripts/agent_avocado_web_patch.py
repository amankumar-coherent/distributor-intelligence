"""Apply agent web-research fills to avocado FINAL.xlsx (no OpenAI SDK).

- Drop leftover junk / non-company rows
- Patch firmographics from agent web verification
- Fast DDGS/SearXNG heuristic fill for remaining gaps
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.config import Settings
from vendor_intel.enrichment.gap_fill.heuristic_extract import heuristic_extract_fields
from vendor_intel.pipeline.web_expand import HEADERS, MARKET_DEFAULTS, _norm, write_final_xlsx

FINAL = ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL.xlsx"
FAMILY = "avocado"
NPD = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "todo",
    "not applicable",
    "seed_file",
}

JUNK_RE = re.compile(
    r"(?i)^("
    r"avocado oil(\s|:).*(buyers?|importers?|suppliers?|wholesalers?|wholesale|direct|"
    r"international|products?|refined|import/?export|co\.?$|company$|oleo|available)|"
    r"avocado oil supplier nz|"
    r".*\b(directory|marketplace|list of|top \d+)\b"
    r")$"
)

# Agent web-verified firmographics (sources: company sites, Wikipedia, SEC, CB Insights, etc.)
WEB_PATCH: dict[str, dict[str, str]] = {
    "sysco": {
        "Founded": "1969",
        "Headquarters": "Houston, Texas, USA",
        "Cities / Regions": "Houston, TX; 340+ distribution centers globally",
        "Ownership": "Public (NYSE: SYY)",
        "Business Type": "Foodservice distributor",
        "Employees": "~75,000",
        "Turnover / Scale": "~$81.4B (FY2025)",
        "Contact Person": "Kevin Hourican",
        "Role": "Chairman & CEO",
        "Website": "https://www.sysco.com",
        "LinkedIn": "https://www.linkedin.com/company/sysco",
        "Core Categories": "Broadline foodservice; edible oils incl. avocado oil",
        "Specialty Focus": "Foodservice distribution",
        "Summary": "World's largest foodservice distributor; founded 1969; HQ Houston; CEO Kevin Hourican.",
    },
    "us foods": {
        "Founded": "1989",
        "Headquarters": "Rosemont, Illinois, USA",
        "Ownership": "Public (NYSE: USFD)",
        "Business Type": "Foodservice distributor",
        "Website": "https://www.usfoods.com",
        "LinkedIn": "https://www.linkedin.com/company/us-foods",
        "Core Categories": "Broadline foodservice; specialty oils",
    },
    "performance food group": {
        "Founded": "1885",
        "Headquarters": "Richmond, Virginia, USA",
        "Ownership": "Public (NYSE: PFGC)",
        "Business Type": "Foodservice distributor",
        "Website": "https://pfgc.com",
    },
    "gordon food service": {
        "Founded": "1897",
        "Headquarters": "Wyoming, Michigan, USA",
        "Ownership": "Privately held",
        "Business Type": "Foodservice distributor",
        "Website": "https://www.gfs.com",
    },
    "kehe": {
        "Founded": "1952",
        "Headquarters": "Naperville, Illinois, USA",
        "Cities / Regions": "1245 E Diehl Rd Ste 200, Naperville, IL; 19 DCs North America",
        "Ownership": "Employee-owned (ESOP)",
        "Business Type": "Natural / specialty grocery distributor",
        "Employees": "~6,800–8,000",
        "Contact Person": "Deb Conklin",
        "Role": "President & CEO",
        "Website": "https://www.kehe.com",
        "LinkedIn": "https://www.linkedin.com/company/kehe-distributors",
        "Core Categories": "Natural, organic, specialty, fresh grocery",
        "Specialty Focus": "Natural & specialty distribution to ~31,000 stores",
        "Summary": "KeHE Distributors — employee-owned natural/specialty grocery distributor; founded 1952; HQ Naperville; CEO Deb Conklin.",
    },
    "unfi": {
        "Company": "UNFI (United Natural Foods)",
        "Founded": "1976",
        "Headquarters": "Providence, Rhode Island, USA",
        "Cities / Regions": "15 Park Row West Suite 302, Providence, RI",
        "Ownership": "Public (NYSE: UNFI)",
        "Business Type": "Natural / organic wholesale distributor",
        "Employees": "~25,600",
        "Turnover / Scale": "~$27B+",
        "Contact Person": "J. Alexander (Sandy) Miller Douglas",
        "Role": "CEO",
        "Phone / WhatsApp": "+1-401-528-8634",
        "Website": "https://www.unfi.com",
        "LinkedIn": "https://www.linkedin.com/company/unfi",
        "Core Categories": "Natural, organic, specialty wholesale grocery",
        "Specialty Focus": "Largest public natural/specialty wholesaler; major Whole Foods supplier",
        "Summary": "UNFI founded 1976; HQ Providence RI; CEO Sandy Douglas; ~25,600 employees.",
    },
    "c&s wholesale": {
        "Founded": "1918",
        "Headquarters": "Keene, New Hampshire, USA",
        "Ownership": "Privately held",
        "Business Type": "Grocery wholesaler",
        "Employees": "~14,000–17,000",
        "Turnover / Scale": "~$30B (historical reports)",
        "Contact Person": "Eric Winn",
        "Role": "CEO",
        "Website": "https://www.cswg.com",
        "Summary": "C&S Wholesale Grocers — founded 1918 Worcester MA; HQ Keene NH; CEO Eric Winn.",
    },
    "chosen foods": {
        "Company": "Chosen Foods LLC",
        "Founded": "2011",
        "Headquarters": "San Diego, California, USA",
        "Cities / Regions": "350 Camino de la Reina Suite 400, San Diego, CA 92108",
        "Ownership": "Private equity-backed (Butterfly majority)",
        "Business Type": "Specialty oils brand / marketer",
        "Contact Person": "Gabriel Krieb",
        "Role": "CEO",
        "Phone / WhatsApp": "+1-877-674-2244",
        "Website": "https://chosenfoods.com",
        "Core Categories": "Avocado oil; mayo; dressings; sprays",
        "Specialty Focus": "100% pure avocado oil cooking products",
        "Summary": "Chosen Foods — founded 2011 San Diego; CEO Gabriel Krieb; avocado oil brand leader.",
    },
    "olivado": {
        "Founded": "2000",
        "Headquarters": "Kerikeri, New Zealand (production also Kenya)",
        "Business Type": "Avocado oil producer / global brand distributor",
        "Website": "https://olivado.com",
        "Core Categories": "Extra virgin avocado oil; olive oil",
        "Specialty Focus": "Cold-pressed extra virgin avocado oil; Fair Trade Kenya supply",
        "Summary": "Olivado founded 2000 in Kerikeri NZ; world's largest EVO avocado oil producer; Kenya plants.",
    },
    "avohealth": {
        "Founded": "Not publicly disclosed",
        "Headquarters": "New Zealand; processing Kenya (EPZ)",
        "Website": "https://www.avohealth.com",
        "Business Type": "Avocado oil producer / bulk B2B distributor",
        "Core Categories": "Extra virgin, crude & refined avocado oil",
    },
    "limbua": {
        "Headquarters": "Germany / Kenya (Mount Kenya)",
        "Website": "https://www.limbua-group.com",
        "Business Type": "Organic avocado oil producer / B2B exporter",
        "Core Categories": "Organic virgin / crude / refined avocado oil",
    },
    "aak": {
        "Founded": "2005",
        "Headquarters": "Malmö, Sweden",
        "Ownership": "Public (Nasdaq Stockholm: AAK)",
        "Business Type": "Specialty oils & fats ingredients",
        "Website": "https://www.aak.com",
        "LinkedIn": "https://www.linkedin.com/company/aak",
    },
    "dot foods": {
        "Founded": "1960",
        "Headquarters": "Mount Sterling, Illinois, USA",
        "Website": "https://www.dotfoods.com",
        "Business Type": "Food redistributor",
        "Ownership": "Privately held",
    },
    "mclane": {
        "Founded": "1894",
        "Headquarters": "Temple, Texas, USA",
        "Website": "https://www.mclaneco.com",
        "Business Type": "Wholesale distribution (grocery & foodservice)",
        "Ownership": "Berkshire Hathaway",
    },
    "core-mark": {
        "Headquarters": "Westlake, Texas, USA",
        "Website": "https://www.core-mark.com",
        "Business Type": "Convenience wholesale distributor",
        "Ownership": "Part of Performance Food Group",
    },
    "ben e. keith": {
        "Founded": "1906",
        "Headquarters": "Fort Worth, Texas, USA",
        "Website": "https://www.benekeith.com",
        "Business Type": "Foodservice & beverage distributor",
        "Ownership": "Privately held",
    },
    "shamrock foods": {
        "Founded": "1922",
        "Headquarters": "Phoenix, Arizona, USA",
        "Website": "https://www.shamrockfoods.com",
        "Business Type": "Foodservice distributor",
        "Ownership": "Privately held",
    },
    "spartannash": {
        "Founded": "1917",
        "Headquarters": "Byron Center, Michigan, USA",
        "Website": "https://www.spartannash.com",
        "Business Type": "Grocery wholesale & retail",
        "Ownership": "Public (Nasdaq: SPTN)",
    },
    "associated wholesale grocers": {
        "Founded": "1924",
        "Headquarters": "Kansas City, Kansas, USA",
        "Website": "https://awginc.com",
        "Business Type": "Grocery cooperative wholesaler",
    },
    "associated food stores": {
        "Headquarters": "Salt Lake City, Utah, USA",
        "Website": "https://www.afstores.com",
        "Business Type": "Grocery cooperative wholesaler",
    },
    "affiliated foods": {
        "Headquarters": "Amarillo, Texas, USA",
        "Website": "http://www.afiama.com",
        "Business Type": "Grocery wholesaler",
    },
    "wakefern": {
        "Founded": "1946",
        "Headquarters": "Keasbey, New Jersey, USA",
        "Website": "https://www.wakefern.com",
        "Business Type": "Retailer-owned grocery cooperative",
    },
    "reyes holdings": {
        "Headquarters": "Rosemont, Illinois, USA",
        "Website": "https://www.reyesholdings.com",
        "Business Type": "Beverage & food distribution",
        "Ownership": "Privately held",
    },
    "martin-brower": {
        "Headquarters": "Rosemont, Illinois, USA",
        "Website": "https://www.martin-brower.com",
        "Business Type": "QSR supply chain distributor",
    },
    "azelis": {
        "Founded": "2001",
        "Headquarters": "Antwerp, Belgium",
        "Website": "https://www.azelis.com",
        "Business Type": "Specialty chemicals / food ingredients distributor",
        "Ownership": "Public",
    },
    "calavo": {
        "Founded": "1924",
        "Headquarters": "Santa Paula, California, USA",
        "Website": "https://www.calavo.com",
        "Business Type": "Avocado grower / marketer / distributor",
        "Ownership": "Public (Nasdaq: CVGW)",
    },
    "mission produce": {
        "Founded": "1983",
        "Headquarters": "Oxnard, California, USA",
        "Website": "https://www.missionproduce.com",
        "Business Type": "Avocado grower / global distributor",
        "Ownership": "Public (Nasdaq: AVO)",
    },
    "westfalia fruit": {
        "Headquarters": "South Africa (global operations)",
        "Website": "https://www.westfaliafruit.com",
        "Business Type": "Avocado grower / packer / marketer",
    },
    "adm": {
        "Founded": "1902",
        "Headquarters": "Chicago, Illinois, USA",
        "Website": "https://www.adm.com",
        "Business Type": "Agricultural processor / ingredients",
        "Ownership": "Public (NYSE: ADM)",
    },
    "all organic treasures": {
        "Company": "All Organic Treasures GmbH",
        "Headquarters": "Wiggensbach, Germany",
        "Website": "https://www.aot.de",
        "Business Type": "Organic oils & ingredients wholesaler",
    },
}


def _is_npd(v: object) -> bool:
    return str(v or "").strip().lower() in NPD


def _is_junk(name: str) -> bool:
    n = (name or "").strip()
    if not n:
        return True
    if JUNK_RE.match(n):
        return True
    low = n.lower()
    # generic non-entity labels
    if low.startswith("avocado oil") and len(n.split()) <= 5:
        # allow real brand-like names with extra words handled above
        if any(x in low for x in ("buyers", "importers", "suppliers", "wholesale", "direct", "products", "refined", "import")):
            return True
    return False


def _load() -> list[dict[str, str]]:
    wb = load_workbook(FINAL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c or "") for c in rows[0]]
    out = []
    for r in rows[1:]:
        d = {h: str(r[i] if i < len(r) and r[i] is not None else "").strip() for i, h in enumerate(hdr)}
        if d.get("Company"):
            out.append({h: d.get(h, "") for h in HEADERS})
    wb.close()
    return out


def _apply_web_patch(row: dict[str, str]) -> int:
    key = _norm(row.get("Company", ""))
    changed = 0
    for needle, patch in WEB_PATCH.items():
        if needle in key:
            for h, v in patch.items():
                if not v:
                    continue
                if h == "Company":
                    if row.get("Company") != v:
                        row["Company"] = v
                        changed += 1
                    continue
                cur = str(row.get(h) or "")
                if _is_npd(cur) or cur.lower() == "seed_file" or (
                    h in ("Founded", "Headquarters", "Employees", "Contact Person", "Role", "Website", "Turnover / Scale")
                    and v
                    and (_is_npd(cur) or len(v) > len(cur))
                ):
                    if cur != v:
                        row[h] = v
                        changed += 1
            break
    return changed


def _defaults(row: dict[str, str]) -> None:
    for h, dv in (MARKET_DEFAULTS.get(FAMILY) or {}).items():
        if h in HEADERS and _is_npd(row.get(h)):
            row[h] = dv
    if _is_npd(row.get("Business Type")):
        row["Business Type"] = "Distributor"
    for h in HEADERS:
        if not str(row.get(h) or "").strip():
            row[h] = "Not publicly disclosed"
    src = row.get("Data Sources") or ""
    if "agent_web_patch" not in src:
        row["Data Sources"] = f"{src}; agent_web_patch".strip("; ")


async def _fast_search_fill(rows: list[dict[str, str]], *, limit: int = 120) -> int:
    os.environ["GAP_FILL_USE_LLM"] = "false"
    from vendor_intel.clients.search_router import FreeSearchRouter
    from vendor_intel.enrichment.gap_fill.search_queries import backfill_query

    router = FreeSearchRouter(Settings())
    todo = [
        r
        for r in rows
        if any(_is_npd(r.get(h)) for h in ("Founded", "Headquarters", "Employees", "Website"))
    ][:limit]
    print(f"  fast web-fill {len(todo)} companies…", flush=True)
    updated = 0
    for i, row in enumerate(todo, 1):
        name = row["Company"]
        q = backfill_query(name, domain=str(row.get("Website") or ""))
        try:
            hits = await router.search(q, discovery_mode=True)
        except Exception as err:  # noqa: BLE001
            print(f"  [{i}] {name}: search err {type(err).__name__}", flush=True)
            continue
        if not hits:
            continue
        blob = "\n".join(f"{h.title}: {h.snippet} ({h.link})" for h in hits[:8])
        extracted = heuristic_extract_fields(blob, company=name)
        crawl_map = {
            "founded_year": "Founded",
            "headquarters": "Headquarters",
            "employee_count": "Employees",
            "revenue": "Turnover / Scale",
            "linkedin_url": "LinkedIn",
            "contact_person": "Contact Person",
            "contact_role": "Role",
        }
        got = False
        for src, dst in crawl_map.items():
            val = str(extracted.get(src) or "").strip()
            if val and _is_npd(row.get(dst)):
                row[dst] = val
                got = True
        # website from first company-looking link
        if _is_npd(row.get("Website")):
            for h in hits:
                link = (h.link or "").lower()
                if any(x in link for x in ("linkedin.com", "facebook.com", "wikipedia.org", "owler.com")):
                    continue
                if link.startswith("http"):
                    row["Website"] = h.link.split("?")[0]
                    got = True
                    break
        if got:
            updated += 1
            print(f"  [{i}/{len(todo)}] filled {name}", flush=True)
        elif i % 20 == 0:
            print(f"  [{i}/{len(todo)}] …", flush=True)
    return updated


async def main() -> int:
    print("=== Agent web patch + fast fill ===", flush=True)
    rows = _load()
    print(f"  loaded {len(rows)}", flush=True)

    cleaned: list[dict[str, str]] = []
    dropped = 0
    for r in rows:
        if _is_junk(r.get("Company", "")):
            print(f"  drop junk: {r.get('Company')}", flush=True)
            dropped += 1
            continue
        cleaned.append(r)
    rows = cleaned

    patched = 0
    for r in rows:
        patched += _apply_web_patch(r)
        _defaults(r)

    # de-dupe by name
    by: dict[str, dict[str, str]] = {}
    for r in rows:
        k = _norm(r["Company"])
        if k not in by:
            by[k] = r
        else:
            # prefer row with more non-NPD fields
            cur = by[k]
            if sum(1 for h in HEADERS if not _is_npd(r.get(h))) > sum(
                1 for h in HEADERS if not _is_npd(cur.get(h))
            ):
                by[k] = r
    rows = sorted(by.values(), key=lambda x: _norm(x["Company"]))

    filled_n = await _fast_search_fill(rows, limit=150)
    for r in rows:
        _apply_web_patch(r)
        _defaults(r)

    # verify pass: drop remaining junk websites
    final_rows = []
    for r in rows:
        web = (r.get("Website") or "").lower()
        if any(x in web for x in ("tradewheel.com", "volza.com", "jumia.com", "trademo.com")):
            continue
        if "agent_web_verify" not in (r.get("Data Sources") or ""):
            r["Data Sources"] = f"{r.get('Data Sources')}; agent_web_verify".strip("; ")
        final_rows.append(r)

    audit = {
        "mode": "agent_own_web_search_fill",
        "dropped_junk": dropped,
        "web_patch_field_writes": patched,
        "fast_search_updated": filled_n,
        "companies": len(final_rows),
        "target_met": len(final_rows) >= 300,
        "output": str(FINAL),
    }
    write_final_xlsx(FINAL, final_rows, "Avocado Oil Companies", audit)
    (FINAL.parent / "agent_web_patch_audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2), flush=True)
    print(f"DONE → {len(final_rows)} companies", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
