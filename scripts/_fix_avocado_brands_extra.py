#!/usr/bin/env python3
from pathlib import Path
import openpyxl

p = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
    r"\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
wb = openpyxl.load_workbook(p)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

EXTRA = {
    "Anecoop": "Not publicly disclosed",
    "Booker Wholesale": "Chef's Larder, others",
    "Eosta": "Nature & More, others",
    "Emirates Snack Foods Distribution": "Barilla, Indomie, Hershey's, Bonne Maman, Alpro, others",
    "Emirates Snack Foods LLC": "Barilla, Indomie, Hershey's, Bonne Maman, Alpro, others",
}

for row in range(3, ws.max_row + 1):
    co = str(ws.cell(row, ic["Company"]).value or "").strip()
    if co in EXTRA:
        ws.cell(row, ic["Key Brands Represented"]).value = EXTRA[co]
        print("set", co)
    b = str(ws.cell(row, ic["Key Brands Represented"]).value or "")
    if "grower brands" in b.lower():
        ws.cell(row, ic["Key Brands Represented"]).value = "Not publicly disclosed"
        print("cleared soft", co)

wb.save(p)
wb.close()

wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c or "") for c in rows[1]]
ic = {h: i for i, h in enumerate(hdr) if h}
print("--- real brands ---")
n = 0
for r in rows[2:]:
    b = str(r[ic["Key Brands Represented"]] or "").strip()
    if b and b != "Not publicly disclosed":
        print(r[ic["Company"]], "|", b)
        n += 1
print("real brand rows", n, "of", len(rows) - 2)
wb.close()
