#!/usr/bin/env python3
"""Retry ONLY previously skipped Website/Founded/HQ companies.

Loads skips from _verify_wfh_simple_progress.jsonl, re-queries Google AI with
simple search-bar queries, then DeepSeek residual. Applies Excel fixes.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import load_workbook
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")
os.environ["GOOGLE_AI_SCRAPER_ENABLED"] = "true"
os.environ["GOOGLE_AI_SCRAPER_FORCE_OVERVIEW"] = "false"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRIES"] = "1"

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

PREV = ROOT / "scripts/_verify_wfh_simple_progress.jsonl"
OUT = ROOT / "scripts/_verify_wfh_skip_retry_progress.jsonl"
SUMMARY = ROOT / "scripts/_verify_wfh_skip_retry_summary.json"
CURRENT_YEAR = 2026
CONCURRENCY = 4


def norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [norm(c.value) for c in ws[r] if c.value not in (None, "")]
        if "company" in vals:
            return r
    return 1


def clean_url(u: object) -> str | None:
    s = str(u or "").strip()
    if not s or s.upper() in {"NPD", "N/A", "NA", "NONE", "UNKNOWN"}:
        return None
    if not re.match(r"^https?://", s, re.I):
        s = "https://" + s.lstrip("/")
    try:
        p = urlparse(s)
        if not p.netloc or "." not in p.netloc:
            return None
        host = p.netloc.lower()
        junk = ("reddit.com", "facebook.com", "archive.org", "google.com", "wikipedia.org", "tiktok.com")
        if any(j in host for j in junk):
            return None
        return f"{p.scheme}://{p.netloc}".rstrip("/")
    except Exception:
        return None


def year_ok(y: object) -> int | None:
    if y is None:
        return None
    if isinstance(y, int):
        yi = y
    else:
        m = re.search(r"(18|19|20)\d{2}", str(y))
        yi = int(m.group(0)) if m else None
    if yi is None:
        return None
    return yi if 1800 <= yi <= CURRENT_YEAR else None


def queries_for(company: str) -> list[str]:
    # Multiple simple search-bar styles (like the user)
    return [
        f"{company} official website founded year headquarters",
        f"{company} founded year?",
        f"{company} headquarters official website",
        f"{company} company website founded established",
    ]


def deepseek() -> OpenAI:
    return OpenAI(
        api_key=os.getenv("OPENAI_API_KEY") or "",
        base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/"),
    )


def extract(company: str, text: str, source: str) -> dict:
    if not (text or "").strip():
        return {}
    low = text.lower()
    if "did not match any documents" in low or "for each company, give" in low:
        return {}
    client = deepseek()
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    prompt = f"""Extract official website, founded year, headquarters for ONE company from search text.

Company: {company}

Text:
```
{text[:9000]}
```

JSON only:
{{"website":"https://..." or null,"founded_year":1994 or null,"headquarters":"City, Country" or null,"confidence":"high"|"medium"|"low"}}
Do not invent. Ignore copyright years.
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    raw = (resp.choices[0].message.content or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.M).strip()
    try:
        data = json.loads(raw)
    except Exception:
        return {}
    if not isinstance(data, dict):
        return {}
    return {
        "website": clean_url(data.get("website")),
        "founded": year_ok(data.get("founded_year")),
        "hq": (str(data.get("headquarters") or "").strip() or None),
        "confidence": data.get("confidence") or "medium",
        "source": source,
    }


def ddg(company: str) -> str:
    try:
        from ddgs import DDGS
    except ImportError:
        try:
            from duckduckgo_search import DDGS  # type: ignore
        except ImportError:
            return ""
    bits = []
    try:
        with DDGS() as ddgs:
            for q in queries_for(company)[:2]:
                for r in ddgs.text(q, max_results=4):
                    bits.append(
                        f"- {r.get('title','')}: {r.get('body') or r.get('description') or ''} | {r.get('href') or ''}"
                    )
    except Exception as e:
        bits.append(f"(search error: {e})")
    return "\n".join(bits)


def deepseek_residual(company: str, old: dict) -> dict:
    snippets = ddg(company)
    client = deepseek()
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    prompt = f"""Find official website, founded year, headquarters.

Company: {company}
Current website: {old.get('website') or 'empty'}
Current founded: {old.get('founded') or 'empty'}
Current HQ: {old.get('hq') or 'empty'}

Web snippets:
{snippets or '(none)'}

JSON only:
{{"website":"https://..." or null,"founded_year":1994 or null,"headquarters":"City, Country" or null,"confidence":"high"|"medium"|"low"}}
Do not invent. If unknown return nulls.
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    raw = (resp.choices[0].message.content or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.M).strip()
    try:
        data = json.loads(raw)
    except Exception:
        return {}
    if not isinstance(data, dict):
        return {}
    return {
        "website": clean_url(data.get("website")),
        "founded": year_ok(data.get("founded_year")),
        "hq": (str(data.get("headquarters") or "").strip() or None),
        "confidence": data.get("confidence") or "medium",
        "source": "deepseek",
    }


def load_current_row(file: str, company: str) -> dict:
    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    hr = header_row(ws)
    headers = [str(c.value or "").strip() for c in ws[hr]]
    ci = headers.index("Company")
    wi = headers.index("Website") if "Website" in headers else None
    fi = headers.index("Founded") if "Founded" in headers else None
    hi = headers.index("Headquarters") if "Headquarters" in headers else None
    out = {"website": "", "founded": "", "hq": ""}
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        if not row or not row[ci]:
            continue
        if norm(row[ci]) == norm(company):
            out = {
                "website": str(row[wi] or "").strip() if wi is not None else "",
                "founded": str(row[fi] or "").strip() if fi is not None else "",
                "hq": str(row[hi] or "").strip() if hi is not None else "",
            }
            break
    wb.close()
    return out


def decide(old: dict, new: dict) -> dict:
    conf = str(new.get("confidence") or "medium").lower()
    changes: dict[str, str] = {}
    if new.get("website") and conf in {"high", "medium"}:
        if clean_url(old.get("website")) != new["website"]:
            changes["Website"] = new["website"]
    if new.get("founded") is not None and conf in {"high", "medium"}:
        oy = year_ok(old.get("founded"))
        if oy is None or str(old.get("founded") or "").upper() == "NPD" or new["founded"] != oy:
            if oy is None or abs(new["founded"] - oy) <= 100 or conf == "high":
                changes["Founded"] = str(new["founded"])
    if new.get("hq") and conf in {"high", "medium"}:
        oh = str(old.get("hq") or "").strip()
        if not oh or oh.upper() in {"NPD", "N/A", "GLOBAL", "UNKNOWN"} or norm(oh) != norm(new["hq"]):
            changes["Headquarters"] = new["hq"]
    if changes:
        return {"action": "fix", "changes": changes}
    if new.get("website") or new.get("founded") is not None or new.get("hq"):
        return {"action": "confirmed", "changes": {}}
    return {"action": "skip", "changes": {}}


async def google_try(company: str) -> dict:
    best: dict = {}
    for q in queries_for(company):
        data = await google_ai_ask(q, llm_clean=True)
        md = str(data.get("markdown") or data.get("markdown_raw") or "").strip()
        parsed = await asyncio.to_thread(extract, company, md, "google_ai")
        if parsed and any([parsed.get("website"), parsed.get("founded"), parsed.get("hq")]):
            # prefer higher confidence / more fields
            score = sum(1 for k in ("website", "founded", "hq") if parsed.get(k))
            best_score = sum(1 for k in ("website", "founded", "hq") if best.get(k))
            if score > best_score or (score == best_score and parsed.get("confidence") == "high"):
                best = parsed
                best["query"] = q
            if score >= 2:
                break
        await asyncio.sleep(0.2)
    return best


async def process_one(item: dict, sem: asyncio.Semaphore, google_ok: bool) -> dict:
    async with sem:
        company = item["company"]
        file = item["file"]
        old = await asyncio.to_thread(load_current_row, file, company)
        parsed: dict = {}
        phase = "google_ai"
        if google_ok:
            try:
                parsed = await google_try(company)
            except Exception as e:
                print(f"  google err {company}: {e}", flush=True)
        if not parsed or not any([parsed.get("website"), parsed.get("founded"), parsed.get("hq")]):
            phase = "deepseek"
            try:
                parsed = await asyncio.to_thread(deepseek_residual, company, old)
            except Exception as e:
                print(f"  deepseek err {company}: {e}", flush=True)
                parsed = {}
        if not parsed:
            rec = {
                "key": item["key"],
                "file": file,
                "company": company,
                "processed": True,
                "phase": phase,
                "action": "skip",
                "changes": {},
            }
        else:
            d = decide(old, parsed)
            rec = {
                "key": item["key"],
                "file": file,
                "company": company,
                "processed": True,
                "phase": phase,
                "query": parsed.get("query"),
                **d,
                "parsed": {
                    "website": parsed.get("website"),
                    "founded": parsed.get("founded"),
                    "hq": parsed.get("hq"),
                    "confidence": parsed.get("confidence"),
                },
            }
        return rec


def apply(results: list[dict]) -> None:
    by_file: dict[str, dict[str, dict]] = {}
    for fx in results:
        if fx.get("action") != "fix" or not fx.get("changes"):
            continue
        by_file.setdefault(fx["file"], {})[norm(fx["company"])] = fx["changes"]
    for fpath, mapping in by_file.items():
        wb = load_workbook(fpath)
        ws = wb.active
        hr = header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            company = ws.cell(r, ci + 1).value
            if not company:
                continue
            ch = mapping.get(norm(str(company)))
            if not ch:
                continue
            for col, val in ch.items():
                if col in headers:
                    ws.cell(r, headers.index(col) + 1).value = val
            n += 1
        wb.save(fpath)
        print(f"Applied {n} fixes -> {Path(fpath).name}")


async def main() -> None:
    skips = []
    for line in PREV.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        obj = json.loads(line)
        if obj.get("action") == "skip":
            skips.append(obj)
    # de-dupe by key
    by_key = {s["key"]: s for s in skips if s.get("key")}
    skips = list(by_key.values())
    print(f"Retrying skipped companies: {len(skips)}")

    google_ok = await google_ai_ping()
    print(f"Google AI scraper: {'connected' if google_ok else 'DOWN'}")

    sem = asyncio.Semaphore(CONCURRENCY)
    results: list[dict] = []
    if OUT.exists():
        OUT.unlink()

    with OUT.open("a", encoding="utf-8") as prog:
        for i in range(0, len(skips), 25):
            batch = skips[i : i + 25]
            print(f"\n=== Retry batch {i//25+1}/{(len(skips)+24)//25} ({len(batch)}) ===", flush=True)
            for s in batch:
                print(f"  {s['company']}", flush=True)
            recs = await asyncio.gather(*[process_one(s, sem, google_ok) for s in batch])
            for rec in recs:
                prog.write(json.dumps(rec, ensure_ascii=False) + "\n")
                prog.flush()
                results.append(rec)
                print(
                    f"  -> {rec['action'].upper()} {rec['company']} [{rec.get('phase')}] "
                    f"{rec.get('changes') or 'no_change/skip'}",
                    flush=True,
                )

    print("\nApplying Excel updates...")
    apply(results)

    fixed = sum(1 for r in results if r.get("action") == "fix")
    confirmed = sum(1 for r in results if r.get("action") == "confirmed")
    still = sum(1 for r in results if r.get("action") == "skip")
    gfix = sum(1 for r in results if r.get("action") == "fix" and r.get("phase") == "google_ai")
    dfix = sum(1 for r in results if r.get("action") == "fix" and r.get("phase") == "deepseek")
    summary = {
        "retried": len(results),
        "fixed": fixed,
        "confirmed": confirmed,
        "still_skipped": still,
        "fixed_via_google_ai": gfix,
        "fixed_via_deepseek": dfix,
    }
    SUMMARY.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print("\n=== SKIP RETRY SUMMARY ===")
    for k, v in summary.items():
        print(f"{k}: {v}")


if __name__ == "__main__":
    asyncio.run(main())
