"""Pass2: tighten Specialty Focus to short focus phrases; refine weak Core Categories."""
from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = PATH.parent / "flexpack_specialty_tighten_audit.json"
NPD = "Not publicly disclosed"

# Evidence-backed Core overrides (from company summary / known product focus)
CORE_OVERRIDES = {
    "Bulk Packaging": "Bulk Bags & FIBCs; Bulk Liners; Flexible Bulk Packaging",
    "Global Packaging": "Flexible Packaging Films; Poly Bags; Packaging Materials",
    "Kiwi Packaging": "Flexible Packaging; Retail Packaging; Packaging Supplies",
    "Packaging Warehouse": "E-commerce Packaging; Flexible Packaging; Packaging Supplies",
}


def emptyish(v) -> bool:
    s = str(v or "").strip()
    return (not s) or s.lower() in {NPD.lower(), "n/a", "na", "-", "none", "unknown"}


def parse_json(text: str) -> dict:
    text = (text or "").strip()
    try:
        return json.loads(text)
    except Exception:
        m = re.search(r"\{[\s\S]*\}", text)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
    return {}


def needs_tighten(spec: str, company: str) -> bool:
    if emptyish(spec):
        return True
    s = spec.strip()
    words = re.findall(r"\b[\w'-]+\b", s)
    if len(words) > 18:
        return True
    if len(s) > 110:
        return True
    if s.lower().startswith(company.lower()[: min(10, len(company))]):
        return True
    if re.search(
        r"\b(is a|are a|offers a|supplies a comprehensive|distributor of|wholesale distributor)\b",
        s,
        re.I,
    ):
        return True
    if re.search(
        r"[,;\s]+(in|across|throughout)\s+(the\s+)?"
        r"(USA|US|UK|United States|Australia|Canada|Germany|France|India|New Zealand|Europe)\.?$",
        s,
        re.I,
    ):
        return True
    return False


def deepseek_specialty(client: OpenAI, model: str, company: str, core: str, summary: str, old: str) -> str:
    system = (
        "Rewrite Specialty Focus for a packaging CHANNEL DISTRIBUTOR spreadsheet cell.\n"
        "Definition: the specific market, technology, product segment, OR customer segment "
        "this distributor is particularly known for.\n"
        "Output: ONE short focus phrase, MAX 14 words. Examples of good form:\n"
        '- \"Foodservice disposables and janitorial packaging for commercial kitchens\"\n'
        '- \"Industrial stretch film and protective packaging for manufacturers\"\n'
        '- \"E-commerce fulfilment packaging and mailers\"\n'
        "Forbidden: company name, country/region endings, long blurbs, 'is a distributor of'.\n"
        "Use only the provided evidence. If unclear, return confidence none.\n"
        'JSON {"specialty":"...","confidence":"high|medium|low|none"}.'
    )
    user = (
        f"Company: {company}\n"
        f"Core Categories: {core or '(none)'}\n"
        f"Current Specialty (may be too long): {old or '(none)'}\n"
        f"Summary: {(summary or '')[:900]}"
    )
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=80,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        data = parse_json(resp.choices[0].message.content or "{}")
        conf = str(data.get("confidence") or "").lower()
        spec = str(data.get("specialty") or "").strip().strip(" .;")
        if conf not in {"high", "medium"} or len(spec) < 8:
            return ""
        spec = re.sub(
            r"[,;\s]+(in|across|throughout)?\s*"
            r"(the\s+)?(USA|US|UK|United States|United Kingdom|Australia|Canada|"
            r"Germany|France|India|Japan|Europe|Asia|New Zealand)\.?$",
            "",
            spec,
            flags=re.I,
        ).strip(" .;")
        # hard trim if model ignored word limit
        words = re.findall(r"\S+", spec)
        if len(words) > 16:
            spec = " ".join(words[:14])
        return spec
    except Exception:
        return ""


def main() -> None:
    api_key = os.getenv("OPENAI_API_KEY") or os.getenv("DEEPSEEK_API_KEY")
    base = os.getenv("OPENAI_BASE_URL", "https://api.deepseek.com/v1")
    model = os.getenv("OPENAI_MODEL", "deepseek-chat")
    if not api_key:
        raise SystemExit("Missing DeepSeek/OpenAI API key")
    client = OpenAI(api_key=api_key, base_url=base)

    bak = PATH.with_name(
        PATH.stem + f"_BEFORE_SPECIALTY_TIGHTEN_{datetime.now().strftime('%Y%m%d_%H%M%S')}" + PATH.suffix
    )
    wb = load_workbook(PATH)
    wb.save(bak)
    print(f"backup {bak.name}", flush=True)

    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {h: i + 1 for i, h in enumerate(headers)}

    audit = {"core_overrides": 0, "specialty_tightened": 0, "skipped": 0, "failed": 0}

    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co in CORE_OVERRIDES:
            ws.cell(r, ic["Core Categories"]).value = CORE_OVERRIDES[co]
            audit["core_overrides"] += 1
            print(f"  core override {co} => {CORE_OVERRIDES[co]}", flush=True)

    targets = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        spec = str(ws.cell(r, ic["Specialty Focus"]).value or "").strip()
        if needs_tighten(spec, co):
            targets.append(r)
    print(f"specialty tighten targets={len(targets)}", flush=True)

    for i, r in enumerate(targets, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        core = str(ws.cell(r, ic["Core Categories"]).value or "").strip()
        summary = str(ws.cell(r, ic.get("Summary", 0)).value or "").strip() if "Summary" in ic else ""
        old = str(ws.cell(r, ic["Specialty Focus"]).value or "").strip()
        new = deepseek_specialty(client, model, co, core, summary, old)
        time.sleep(0.15)
        if new and new.lower() != old.lower():
            ws.cell(r, ic["Specialty Focus"]).value = new
            audit["specialty_tightened"] += 1
            if i % 25 == 0 or i <= 5:
                print(f"  [{i}/{len(targets)}] {co} => {new}", flush=True)
        elif not new:
            # fallback: first core category as focus
            if not emptyish(core):
                first = core.split(";")[0].strip()
                ws.cell(r, ic["Specialty Focus"]).value = f"{first} distribution focus"
                audit["specialty_tightened"] += 1
            else:
                audit["failed"] += 1
        else:
            audit["skipped"] += 1
        if i % 40 == 0:
            wb.save(PATH)

    wb.save(PATH)
    wb.close()
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2))
    print("done", flush=True)


if __name__ == "__main__":
    main()
