#!/usr/bin/env python3
"""Fill remaining avocado Office No. / Office Email from website contact pages.

Strategy:
1) Fetch homepage + /contact, /contact-us, /contactus, /about, imprint, etc.
2) Also follow any href on homepage that looks like a contact page.
3) Regex extract tel:/mailto: + office-style emails.
4) DeepSeek extracts ONLY from scraped page text (must appear in text).
5) Never invent; leave Not publicly disclosed if not found.
"""
from __future__ import annotations

import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
import openpyxl
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl.styles import Alignment

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

OUT = ROOT / "output/web_expand/global_avocado_oil_market_global"
XLSX = OUT / "global_avocado_oil_market_OFFICE_SLIM.xlsx"
CACHE = OUT / "_avocado_contactpage_fill_cache.json"
AUDIT = OUT / "_avocado_contactpage_fill_audit.json"

NPD = "Not publicly disclosed"
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

CONTACT_PATHS = [
    "",
    "/contact",
    "/contact-us",
    "/contactus",
    "/contacts",
    "/contact-us/",
    "/en/contact",
    "/en/contact-us",
    "/us/contact",
    "/about/contact",
    "/company/contact",
    "/pages/contact",
    "/pages/contact-us",
    "/about",
    "/about-us",
    "/aboutus",
    "/impressum",
    "/imprint",
    "/legal/imprint",
    "/get-in-touch",
    "/reach-us",
    "/customer-service",
    "/customer-care",
    "/support/contact",
    "/locations",
    "/our-locations",
    "/find-us",
]

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
TEL_HREF_RE = re.compile(r"tel:([+\d][\d\-().\s]{6,})", re.I)
MAILTO_RE = re.compile(r"mailto:([a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,})", re.I)
HREF_RE = re.compile(r"""href=["']([^"']+)["']""", re.I)
PHONE_HINT_RE = re.compile(
    r"(?:tel:|phone|tel\.|call|fax|t\s*[\(:]|office|contact|hotline|customer\s*service|"
    r"telefon|telephone|teléfono|telefone|contacto)"
    r"[^0-9+]{0,40}(\+?[\d][\d\s().\-/]{7,22}\d)",
    re.I,
)

BAD_EMAIL_SUB = (
    "example.com",
    "exemple.com",
    "w3schools",
    "mysite.com",
    "opencart",
    "sentry",
    "wixpress",
    "schema.org",
    "cloudflare",
    "financebuzz",
    "noreply",
    "no-reply",
    "donotreply",
    "yourcompany",
    "domain.com",
    "email.com",
    "sentry.io",
    "github.com",
    "googleapis",
)

OFFICE_LOCAL = re.compile(
    r"^(info|contact|sales|office|hello|enquiry|enquiries|support|customer|customerservice|"
    r"customerservices|customersupport|orders|export|mail|admin|team|trade|help|press|media|"
    r"technical|inquiry|kontakt|infouk|communications|customercare|service|promo|"
    r"reception|general|wholesale|distributor|distributors|procurement|"
    r"info\.|sales\.|contact\.)([._-]|$)",
    re.I,
)

JUNK_PHONE_DIGITS = {
    "6666666667",
    "6666666666",
    "1234567890",
    "0000000000",
    "1111111111",
}


def emptyish(v: object) -> bool:
    s = str(v or "").strip().lower()
    return s in {
        "",
        "n/a",
        "na",
        "none",
        "-",
        "—",
        "not publicly disclosed",
        "unknown",
        "tbd",
        "null",
    }


def normalize_website(url: str) -> str:
    u = (url or "").strip()
    if emptyish(u):
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    return u.rstrip("/")


def clean_phone(raw: str) -> str:
    s = re.sub(r"\s+", " ", str(raw or "").strip()).replace("tel:", "").strip()
    digits = re.sub(r"\D", "", s)
    if len(digits) < 8 or len(digits) > 15:
        return ""
    if digits in JUNK_PHONE_DIGITS:
        return ""
    # reject mostly-same-digit spam
    if len(set(digits)) <= 2 and len(digits) >= 8:
        return ""
    return s[:40]


def clean_email(raw: str) -> str:
    e = str(raw or "").strip().lower().lstrip("%20").rstrip(".,;)>")
    m = re.search(r"([a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,})", e)
    if m:
        e = m.group(1)
    if not e or "@" not in e:
        return ""
    if any(b in e for b in BAD_EMAIL_SUB):
        return ""
    if e.endswith((".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".css", ".js")):
        return ""
    local, _, dom = e.partition("@")
    if not local or not dom or "." not in dom:
        return ""
    return e


def is_office_email(email: str) -> bool:
    e = clean_email(email)
    if not e:
        return False
    local = e.split("@")[0]
    if OFFICE_LOCAL.search(local):
        return True
    return bool(
        re.search(
            r"(info|contact|sales|office|support|enquir|customer|hello|export|orders|service|wholesale)",
            local,
            re.I,
        )
    )


def is_junk_phone(phone: str) -> bool:
    if emptyish(phone):
        return True
    digits = re.sub(r"\D", "", phone)
    if digits in JUNK_PHONE_DIGITS:
        return True
    if len(digits) < 8:
        return True
    if len(set(digits)) <= 2:
        return True
    # weird spaced zeros like "0 0 512 512"
    if re.fullmatch(r"[\d\s]+", phone) and phone.count("0") >= 3 and len(digits) <= 10:
        if digits.startswith("00") or "00512" in digits:
            return True
    return False


def html_to_text(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<noscript.*?>.*?</noscript>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    t = re.sub(r"&\w+;", " ", t)
    t = re.sub(r"\s+", " ", t)
    return t.strip()


def discover_contact_urls(base: str, html: str) -> list[str]:
    out: list[str] = []
    seen = set()
    host = urlparse(base).netloc.lower()
    keys = (
        "contact",
        "contacto",
        "kontakt",
        "impressum",
        "imprint",
        "get-in-touch",
        "reach-us",
        "customer-service",
        "customer-care",
        "locations",
        "find-us",
        "about-us",
        "aboutus",
    )
    for href in HREF_RE.findall(html or ""):
        h = href.strip()
        if not h or h.startswith(("#", "javascript:", "mailto:", "tel:")):
            continue
        low = h.lower()
        if not any(k in low for k in keys):
            continue
        full = urljoin(base + "/", h)
        p = urlparse(full)
        if p.scheme not in {"http", "https"}:
            continue
        if host and p.netloc and host not in p.netloc.lower() and p.netloc.lower() not in host:
            # allow www vs non-www
            if urlparse(base).netloc.replace("www.", "") not in p.netloc.replace("www.", ""):
                continue
        key = full.split("#")[0].rstrip("/")
        if key in seen:
            continue
        seen.add(key)
        out.append(key)
        if len(out) >= 12:
            break
    return out


def uniq_list(xs: list[str], phone: bool = False) -> list[str]:
    seen = set()
    out = []
    for x in xs:
        k = re.sub(r"\D", "", x) if phone else x.lower()
        if not k or k in seen:
            continue
        seen.add(k)
        out.append(x)
    return out


def extract_from_html(html: str, base_url: str) -> tuple[list[str], list[str], str]:
    phones: list[str] = []
    emails: list[str] = []
    for m in TEL_HREF_RE.finditer(html):
        p = clean_phone(m.group(1))
        if p:
            phones.append(p)
    for m in MAILTO_RE.finditer(html):
        e = clean_email(m.group(1))
        if e and is_office_email(e):
            emails.append(e)
    for m in EMAIL_RE.finditer(html):
        e = clean_email(m.group(0))
        if e and is_office_email(e):
            emails.append(e)
    txt = html_to_text(html)
    for m in PHONE_HINT_RE.finditer(html):
        p = clean_phone(m.group(1))
        if p:
            phones.append(p)
    for m in PHONE_HINT_RE.finditer(txt):
        p = clean_phone(m.group(1))
        if p:
            phones.append(p)
    return uniq_list(phones, phone=True), uniq_list(emails), txt


def pick_best(phones: list[str], emails: list[str], base_url: str) -> tuple[str, str]:
    phone = phones[0] if phones else ""
    host = urlparse(base_url).netloc.lower().removeprefix("www.").split(":")[0]
    domain = [e for e in emails if host and (host in e.split("@")[-1] or e.split("@")[-1] in host)]
    prefer_locals = {
        "info",
        "contact",
        "sales",
        "office",
        "hello",
        "enquiries",
        "enquiry",
        "support",
        "customerservice",
        "customerservices",
        "orders",
        "export",
        "wholesale",
        "customercare",
    }
    pool = domain or emails
    prefer = [e for e in pool if e.split("@")[0].split("+")[0] in prefer_locals]
    email = (prefer or pool)[0] if pool else ""
    return phone, email


def scrape_site(website: str) -> dict:
    url = normalize_website(website)
    if not url:
        return {"phone": "", "email": "", "text": "", "url": "", "ok": False, "pages": 0}
    phones: list[str] = []
    emails: list[str] = []
    texts: list[str] = []
    pages = 0
    try:
        with httpx.Client(
            headers={"User-Agent": UA, "Accept-Language": "en-US,en;q=0.9"},
            verify=False,
            timeout=httpx.Timeout(12.0, connect=5.0),
            follow_redirects=True,
        ) as client:
            # homepage first for discovery
            try:
                home = client.get(url)
                if home.status_code < 400:
                    pages += 1
                    p, e, t = extract_from_html(home.text or "", str(home.url))
                    phones.extend(p)
                    emails.extend(e)
                    if t:
                        texts.append(t[:5000])
                    extra = discover_contact_urls(str(home.url).rstrip("/"), home.text or "")
                else:
                    extra = []
            except Exception:
                extra = []

            paths = list(CONTACT_PATHS[1:])  # skip "" already fetched
            urls = []
            for path in paths:
                urls.append(urljoin(url + "/", path.lstrip("/")))
            urls.extend(extra)

            seen_u = {url.rstrip("/")}
            for u in urls:
                key = u.split("#")[0].rstrip("/")
                if key in seen_u:
                    continue
                seen_u.add(key)
                try:
                    r = client.get(u)
                    if r.status_code >= 400:
                        continue
                    pages += 1
                    p, e, t = extract_from_html(r.text or "", str(r.url))
                    phones.extend(p)
                    emails.extend(e)
                    if t:
                        texts.append(t[:5000])
                    if phones and emails and pages >= 3:
                        break
                except Exception:
                    continue
    except Exception as ex:  # noqa: BLE001
        return {
            "phone": "",
            "email": "",
            "text": "",
            "url": url,
            "ok": False,
            "pages": 0,
            "error": str(ex)[:120],
        }

    phones = uniq_list(phones, phone=True)
    emails = uniq_list(emails)
    phone, email = pick_best(phones, emails, url)
    return {
        "phone": phone,
        "email": email,
        "phones": phones[:8],
        "emails": emails[:8],
        "text": "\n---\n".join(texts)[:14000],
        "url": url,
        "ok": bool(phone or email or texts),
        "pages": pages,
    }


def ds_client() -> tuple[OpenAI, str]:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    if not key:
        raise SystemExit("OPENAI_API_KEY missing")
    return OpenAI(api_key=key, base_url=base, timeout=60.0), model


def _parse_json(text: str) -> dict:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    try:
        return json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


def deepseek_extract(client: OpenAI, model: str, company: str, website: str, page_text: str) -> dict:
    if not (page_text or "").strip():
        return {"phone": "", "email": ""}
    system = (
        "Extract PUBLIC office phone and office email from the website text only. "
        "Prefer info@/sales@/contact@/support@/customerservice@. "
        "Do NOT invent. If not clearly in the text, found_*=false. Return ONLY JSON."
    )
    user = (
        f"Company: {company}\nWebsite: {website}\n\n"
        f"TEXT:\n{page_text[:9000]}\n\n"
        '{"found_phone":true|false,"office_phone":"","found_email":true|false,'
        '"office_email":"","confidence":"high|medium|low","phone_in_text":true|false,'
        '"email_in_text":true|false}'
    )
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=300,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        data = _parse_json(resp.choices[0].message.content or "{}")
    except Exception as ex:  # noqa: BLE001
        return {"phone": "", "email": "", "error": str(ex)[:120]}

    conf = str(data.get("confidence") or "").lower()
    out = {"phone": "", "email": "", "confidence": conf, "raw": data}
    if conf not in {"high", "medium"}:
        return out

    digits_page = re.sub(r"\D", "", page_text)
    text_l = page_text.lower()

    if data.get("found_phone"):
        phone = clean_phone(str(data.get("office_phone") or ""))
        if phone:
            d = re.sub(r"\D", "", phone)
            if d[-7:] in digits_page:
                out["phone"] = phone

    if data.get("found_email"):
        email = clean_email(str(data.get("office_email") or ""))
        if email and email in text_l and is_office_email(email):
            out["email"] = email
    return out


def main() -> None:
    client, model = ds_client()
    print(f"DeepSeek model={model}", flush=True)

    cache: dict = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[2]]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    gaps = []
    for row in range(3, ws.max_row + 1):
        co = str(ws.cell(row, ic["Company"]).value or "").strip()
        if not co:
            continue
        web = str(ws.cell(row, ic["Website"]).value or "").strip()
        ph = str(ws.cell(row, ic["Office No."]).value or "").strip()
        em = str(ws.cell(row, ic["Office Email"]).value or "").strip()
        need_p = emptyish(ph) or is_junk_phone(ph)
        need_e = emptyish(em)
        if need_p or need_e:
            gaps.append(
                {
                    "row": row,
                    "company": co,
                    "website": web,
                    "need_phone": need_p,
                    "need_email": need_e,
                    "cur_phone": ph,
                    "cur_email": em,
                }
            )

    print(f"gap rows={len(gaps)} (phone and/or email NPD/junk)", flush=True)

    scrape_map: dict[str, dict] = {}
    todo = []
    for g in gaps:
        web = normalize_website(g["website"])
        if web:
            todo.append((g["company"], web))

    print(f"scraping contact pages for {len(todo)} sites ...", flush=True)
    with ThreadPoolExecutor(max_workers=10) as ex:
        futs = {ex.submit(scrape_site, web): (co, web) for co, web in todo}
        done = 0
        for fut in as_completed(futs):
            co, web = futs[fut]
            try:
                scrape_map[web] = fut.result()
            except Exception as exn:  # noqa: BLE001
                scrape_map[web] = {
                    "phone": "",
                    "email": "",
                    "text": "",
                    "url": web,
                    "ok": False,
                    "error": str(exn)[:80],
                }
            done += 1
            if done % 25 == 0 or done == len(todo):
                print(f"  scraped {done}/{len(todo)} last={co[:40]}", flush=True)

    filled_phone = filled_email = 0
    ds_calls = 0
    details = []

    for i, g in enumerate(gaps, 1):
        row = g["row"]
        co = g["company"]
        web = normalize_website(g["website"])
        scrap = scrape_map.get(web, {})
        phone = scrap.get("phone") or ""
        email = scrap.get("email") or ""

        still_need = (g["need_phone"] and not phone) or (g["need_email"] and not email)
        cache_key = f"ds::{co}::{web}"
        ds_hit = cache.get(cache_key) or {}

        if still_need and scrap.get("text"):
            # refresh DS when prior cache empty or still missing
            if not (ds_hit.get("phone") or ds_hit.get("email")) or still_need:
                print(f"[{i}/{len(gaps)}] DeepSeek {co[:45]}", flush=True)
                ds_hit = deepseek_extract(client, model, co, web, scrap.get("text") or "")
                cache[cache_key] = {
                    **ds_hit,
                    "ts": datetime.now(timezone.utc).isoformat(),
                    "pages": scrap.get("pages"),
                }
                ds_calls += 1
                if ds_calls % 15 == 0:
                    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
                time.sleep(0.1)

        if g["need_phone"] and not phone and ds_hit.get("phone"):
            phone = ds_hit["phone"]
        if g["need_email"] and not email and ds_hit.get("email") and is_office_email(ds_hit["email"]):
            email = ds_hit["email"]

        changed = {}
        if g["need_phone"]:
            if phone and not is_junk_phone(phone):
                ws.cell(row, ic["Office No."]).value = phone
                filled_phone += 1
                changed["phone"] = phone
            else:
                ws.cell(row, ic["Office No."]).value = NPD
        if g["need_email"]:
            if email and is_office_email(email):
                ws.cell(row, ic["Office Email"]).value = email
                filled_email += 1
                changed["email"] = email
            else:
                ws.cell(row, ic["Office Email"]).value = NPD

        if changed:
            cell_p = ws.cell(row, ic["Office No."])
            cell_e = ws.cell(row, ic["Office Email"])
            cell_p.alignment = Alignment(wrap_text=True, vertical="top")
            cell_e.alignment = Alignment(wrap_text=True, vertical="top")
            details.append({"company": co, "website": web, **changed, "source_pages": scrap.get("pages")})

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
    wb.save(XLSX)
    wb.close()

    # recount
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    ic2 = {str(h): i for i, h in enumerate(hdr) if h}
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if str(r[ic2["Company"]] or "").strip()]
    phone_n = sum(1 for r in rows if not emptyish(r[ic2["Office No."]]) and not is_junk_phone(str(r[ic2["Office No."]])))
    email_n = sum(1 for r in rows if not emptyish(r[ic2["Office Email"]]))
    phone_npd = len(rows) - phone_n
    email_npd = len(rows) - email_n
    wb.close()

    audit = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "file": str(XLSX),
        "gap_rows": len(gaps),
        "deepseek_calls": ds_calls,
        "newly_filled_phone": filled_phone,
        "newly_filled_email": filled_email,
        "total_rows": len(rows),
        "total_phone_filled": phone_n,
        "total_email_filled": email_n,
        "remaining_phone_npd": phone_npd,
        "remaining_email_npd": email_npd,
        "filled_samples": details[:40],
        "model": model,
        "rule": "Contact-page scrape + DeepSeek extract only values present in page text; never invent",
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: audit[k] for k in audit if k != "filled_samples"}, indent=2))
    print(f"samples_filled={len(details)}")
    for d in details[:25]:
        print(" +", d.get("company"), d.get("phone", ""), d.get("email", ""))


if __name__ == "__main__":
    main()
