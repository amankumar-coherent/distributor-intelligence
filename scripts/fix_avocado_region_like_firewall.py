"""Restore avocado FINAL columns and segment like global firewall market.

Layout (same as firewall):
  Company | Region Segment | Founded | Headquarters | ...all other detail columns...

Sheets:
  Avocado Oil Distributors (main, all companies, full columns)
  Region Summary
  Region - North America / Europe / Asia Pacific / ... / Global
"""
from __future__ import annotations

import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FINAL = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_global_FINAL.xlsx"
)
# Clean backup: Company-first, all detail cols, exclusivity already removed
RESTORE_FROM = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_global_FINAL_pre_region_segment_20260803_185518.xlsx"
)

REGION_ORDER = [
    "North America",
    "Europe",
    "Asia Pacific",
    "Middle East",
    "Latin America",
    "Africa",
    "Global",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
BODY_FONT = Font(name="Calibri", size=10)
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
ALT = PatternFill("solid", fgColor="F2F2F2")

COL_WIDTHS = {
    "Company": 28,
    "Region Segment": 16,
    "Founded": 10,
    "Headquarters": 28,
    "Cities / Regions": 32,
    "Ownership": 20,
    "Business Type": 26,
    "Employees": 12,
    "Turnover / Scale": 14,
    "Contact Person": 20,
    "Role": 22,
    "Email": 26,
    "Phone / WhatsApp": 16,
    "LinkedIn": 32,
    "Website": 28,
    "Core Categories": 36,
    "Specialty Focus": 38,
    "Price Segment": 20,
    "Key Brands Represented": 36,
    "Retail Chains": 36,
    "Specialty Grocers": 36,
    "Foodservice / HoReCa": 36,
    "E-commerce": 28,
    "Channel Strength": 28,
    "Distribution Type": 28,
    "Regions Served": 28,
    "Regional Extensions": 28,
    "Warehouse / Logistics": 32,
    "Delivery / Storage": 28,
    "Competitive Benchmark": 32,
    "Data Confidence": 12,
    "Quality Score": 12,
    "Data Sources": 40,
    "Summary": 48,
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def infer_region(hq: str, cities: str, regions_served: str, regional_ext: str = "") -> str:
    """Firewall-style: Global if multi-region; else HQ/geo based."""
    blob = _norm(f"{regions_served} {regional_ext} {hq} {cities}")

    # Explicit Global (firewall pattern)
    if any(
        x in blob
        for x in (
            "global",
            "worldwide",
            "70+ countries",
            "30+ countries",
            "multi-region",
            "emea hq; americas",
        )
    ):
        # Country-specific Bidfood stays local
        return "Global"

    # Latin America (before generic america)
    if any(
        x in blob
        for x in (
            "brazil",
            "chile",
            "colombia",
            "argentina",
            "peru",
            "uruguay",
            "latam",
            "latin america",
            "costa rica",
        )
    ):
        return "Latin America"
    if "mexico" in blob and not any(x in blob for x in ("united states", "usa", "canada", "u.s.")):
        # Firewall groups Mexico under North America
        return "North America"

    # Africa
    if any(x in blob for x in ("south africa", "nigeria", "kenya", "egypt", "africa")) and "middle east" not in blob:
        if any(x in blob for x in ("south africa", "nigeria", "kenya", "egypt", "africa")):
            # MEA mixed — prefer Africa keywords alone
            if not any(x in blob for x in ("uae", "dubai", "saudi", "gcc", "qatar")):
                return "Africa"

    # Middle East
    if any(
        x in blob
        for x in (
            "uae",
            "dubai",
            "abu dhabi",
            "saudi",
            "gcc",
            "middle east",
            "qatar",
            "bahrain",
            "kuwait",
            "oman",
            "israel",
            "turkey",
            "mena",
        )
    ):
        return "Middle East"

    # Asia Pacific
    if any(
        x in blob
        for x in (
            "asia",
            "australia",
            "new zealand",
            "japan",
            "china",
            "singapore",
            "malaysia",
            "indonesia",
            "thailand",
            "korea",
            "india",
            "pakistan",
            "hong kong",
            "taiwan",
            "philippines",
            "vietnam",
            "asia-pacific",
            "asia pacific",
            "anz",
        )
    ):
        return "Asia Pacific"

    # Europe
    if any(
        x in blob
        for x in (
            "united kingdom",
            " uk",
            "uk ",
            "england",
            "scotland",
            "netherlands",
            "germany",
            "france",
            "spain",
            "belgium",
            "poland",
            "portugal",
            "italy",
            "sweden",
            "ireland",
            "switzerland",
            "czech",
            "austria",
            "denmark",
            "norway",
            "finland",
            "greece",
            "hungary",
            "romania",
            "bulgaria",
            "baltics",
            "latvia",
            "lithuania",
            "estonia",
            "europe",
            "benelux",
        )
    ):
        return "Europe"

    # North America (incl. Mexico like firewall)
    if any(
        x in blob
        for x in (
            "united states",
            "usa",
            "u.s.",
            "canada",
            "mexico",
            "québec",
            "quebec",
            "north america",
            "texas",
            "california",
        )
    ):
        return "North America"

    return "Global"


def _style_header_row(ws, ncols: int, row: int = 1) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = "C2"


def _style_body(ws, headers: list[str], n_data_rows: int) -> None:
    for r in range(2, n_data_rows + 2):
        ws.row_dimensions[r].height = 45
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r % 2 == 0:
                cell.fill = ALT
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS.get(h, 18)
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last}{n_data_rows + 1}"


def _write_sheet(wb, title: str, headers: list[str], rows: list[dict], index: int | None = None) -> None:
    name = title[:31]
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, index=index)
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c).value = h
    for r, d in enumerate(rows, start=2):
        for c, h in enumerate(headers, start=1):
            ws.cell(r, c).value = d.get(h)
    _style_header_row(ws, len(headers))
    if rows:
        _style_body(ws, headers, len(rows))


def main() -> None:
    if not RESTORE_FROM.exists():
        raise SystemExit(f"Missing restore source: {RESTORE_FROM}")

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = FINAL.with_name(f"{FINAL.stem}_pre_firewall_style_region_{stamp}{FINAL.suffix}")
    if FINAL.exists():
        shutil.copy2(FINAL, bak)

    # Start from clean Company-first workbook
    shutil.copy2(RESTORE_FROM, FINAL)
    wb = openpyxl.load_workbook(FINAL)
    ws = wb["Avocado Oil Distributors"]
    hdr = [str(c.value or "") for c in ws[1]]
    if hdr[0] != "Company":
        raise SystemExit(f"Expected Company first, got {hdr[:3]}")

    # Drop exclusivity leftovers if present
    drop = {"Exclusivity", "Partnership Duration", "Region Segment", "Primary Region"}
    keep_hdr = [h for h in hdr if h and h not in drop]

    companies: list[dict] = []
    for r in range(2, ws.max_row + 1):
        d = {}
        for i, h in enumerate(hdr, start=1):
            if h in keep_hdr:
                d[h] = ws.cell(r, i).value
        company = str(d.get("Company") or "").strip()
        if not company:
            continue
        d["Region Segment"] = infer_region(
            str(d.get("Headquarters") or ""),
            str(d.get("Cities / Regions") or ""),
            str(d.get("Regions Served") or ""),
            str(d.get("Regional Extensions") or ""),
        )
        # Bidfood country entities: force local region from company name
        m = re.match(r"(?i)^bidfood\s+(.+)$", company)
        if m:
            region_name = m.group(1).strip().lower()
            mapping = {
                "uk": "Europe",
                "australia": "Asia Pacific",
                "new zealand": "Asia Pacific",
                "singapore": "Asia Pacific",
                "asia": "Asia Pacific",
                "chile": "Latin America",
                "south africa": "Africa",
                "belgium": "Europe",
                "netherlands": "Europe",
                "spain": "Europe",
                "portugal": "Europe",
                "poland farutex": "Europe",
                "czech": "Europe",
                "baltics": "Europe",
                "turkey": "Middle East",
            }
            d["Region Segment"] = mapping.get(region_name, d["Region Segment"])
            if region_name not in mapping:
                if any(x in region_name for x in ("poland", "czech", "baltics", "belgium", "spain", "portugal", "netherlands", "uk")):
                    d["Region Segment"] = "Europe"
                elif any(x in region_name for x in ("australia", "zealand", "asia", "singapore")):
                    d["Region Segment"] = "Asia Pacific"
                elif "chile" in region_name:
                    d["Region Segment"] = "Latin America"
                elif "africa" in region_name:
                    d["Region Segment"] = "Africa"
        companies.append(d)

    new_headers = ["Company", "Region Segment"] + [h for h in keep_hdr if h != "Company"]

    # Remove old region helper sheets
    for name in list(wb.sheetnames):
        if name.startswith("Region -") or name in {
            "By Region Summary",
            "Region Summary",
            "Global Split Note",
            "All Distributors",
        }:
            del wb[name]

    # Rewrite main sheet completely (Company | Region Segment | ...)
    # Clear sheet
    if ws.max_column:
        ws.delete_cols(1, ws.max_column)
    if ws.max_row > 1:
        ws.delete_rows(1, ws.max_row)
    for c, h in enumerate(new_headers, start=1):
        ws.cell(1, c).value = h
    companies_sorted = sorted(companies, key=lambda x: _norm(str(x.get("Company") or "")))
    for r, d in enumerate(companies_sorted, start=2):
        for c, h in enumerate(new_headers, start=1):
            ws.cell(r, c).value = d.get(h)
    _style_header_row(ws, len(new_headers))
    _style_body(ws, new_headers, len(companies_sorted))

    # Region Summary (firewall style)
    counts = Counter(d["Region Segment"] for d in companies_sorted)
    sum_ws = wb.create_sheet("Region Summary", 1)
    sum_ws["A1"] = "Global Avocado Oil Market — Region-wise segmentation"
    sum_ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    sum_ws.merge_cells("A1:D1")
    sum_ws["A3"] = "Region Segment"
    sum_ws["B3"] = "Company Count"
    sum_ws["C3"] = "% of Total"
    sum_ws["D3"] = "Notes"
    for col in range(1, 5):
        cell = sum_ws.cell(3, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    total = len(companies_sorted)
    notes = {
        "Global": "Regions Served = Global / multi-region (HQ may still be local)",
        "North America": "USA, Canada, Mexico",
        "Europe": "UK + EU / wider Europe",
        "Asia Pacific": "India, China, Japan, SEA, ANZ, Pakistan, Korea",
        "Middle East": "GCC, UAE, Saudi, Turkey",
        "Latin America": "Brazil, Chile, Colombia, Argentina, Peru, Costa Rica",
        "Africa": "South Africa and other African markets",
    }
    r = 4
    for region in REGION_ORDER:
        n = counts.get(region, 0)
        if n == 0:
            continue
        sum_ws.cell(r, 1, region).border = THIN
        sum_ws.cell(r, 2, n).border = THIN
        sum_ws.cell(r, 3, f"{100.0 * n / total:.1f}%").border = THIN
        sum_ws.cell(r, 4, notes.get(region, "")).border = THIN
        r += 1
    sum_ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    sum_ws.cell(r, 2, total).font = Font(bold=True)
    sum_ws.column_dimensions["A"].width = 18
    sum_ws.column_dimensions["B"].width = 14
    sum_ws.column_dimensions["C"].width = 12
    sum_ws.column_dimensions["D"].width = 55

    # Per-region sheets with FULL same columns
    by_region: dict[str, list[dict]] = defaultdict(list)
    for d in companies_sorted:
        by_region[d["Region Segment"]].append(d)

    for region in REGION_ORDER:
        rows = by_region.get(region) or []
        if not rows:
            continue
        _write_sheet(wb, f"Region - {region}", new_headers, rows)

    # Keep Web Expand Audit if present
    wb.save(FINAL)
    print("restored_from", RESTORE_FROM.name)
    print("backup_current", bak.name)
    print("file", FINAL.name)
    print("total", total)
    print("headers", new_headers[:6], "...", f"({len(new_headers)} cols)")
    for region in REGION_ORDER:
        if region in counts:
            print(f"  {region}: {counts[region]}")
    print("sheets", wb.sheetnames)

    # sanity: Company is col1, Region Segment col2, Founded col3
    wb2 = openpyxl.load_workbook(FINAL, data_only=True)
    main = wb2["Avocado Oil Distributors"]
    h = [c.value for c in main[1]]
    print("VERIFY header", h[:5])
    print("VERIFY row2", [main.cell(2, i).value for i in range(1, 6)])
    # empties check on a few detail cols
    empty = Counter()
    for r in range(2, main.max_row + 1):
        for f in ("Company", "Founded", "Headquarters", "Business Type", "Core Categories", "Summary"):
            if f in h and not str(main.cell(r, h.index(f) + 1).value or "").strip():
                empty[f] += 1
    print("empty_spotcheck", dict(empty))


if __name__ == "__main__":
    main()
