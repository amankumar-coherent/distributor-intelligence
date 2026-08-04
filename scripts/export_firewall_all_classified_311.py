"""Export all_classified (311) companies for global firewall market to Excel."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_global.json")
OUT = Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_all_classified_311.xlsx")

PREFERRED = [
    "company",
    "brand",
    "name",
    "founded",
    "year_founded",
    "headquarters",
    "hq",
    "hq_city",
    "hq_country",
    "cities",
    "cities_regions",
    "ownership",
    "business_type",
    "role",
    "employees",
    "turnover",
    "scale",
    "contact_person",
    "role_title",
    "email",
    "phone",
    "linkedin",
    "website",
    "domain",
    "core_categories",
    "specialty_focus",
    "price_segment",
    "key_brands",
    "regions_served",
    "distribution_type",
    "is_relevant",
    "confidence",
    "quality_score",
    "classification",
    "channel_role",
    "summary",
    "sources",
    "data_sources",
]

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
body_font = Font(name="Calibri", size=10)
wrap = Alignment(wrap_text=True, vertical="top")
header_align = Alignment(wrap_text=True, vertical="center", horizontal="center")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def cell_val(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, (dict, list)):
        try:
            s = json.dumps(v, ensure_ascii=False)
        except Exception:
            s = str(v)
        return s[:500] + ("..." if len(s) > 500 else "")
    return str(v)


def pick(row: dict, *keys):
    for k in keys:
        if row.get(k) not in (None, ""):
            return row.get(k)
    return ""


def style_header(ws, ncols: int) -> None:
    for col_i in range(1, ncols + 1):
        cell = ws.cell(1, col_i)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin
    ws.row_dimensions[1].height = 32


def style_body(ws, ncols: int) -> None:
    for r in range(2, ws.max_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin
        ws.row_dimensions[r].height = 28


def main() -> None:
    data = json.loads(SRC.read_text(encoding="utf-8"))
    rows = [r for r in (data.get("all_classified") or []) if isinstance(r, dict)]
    print(f"all_classified count: {len(rows)}")
    if not rows:
        raise SystemExit("no all_classified rows")

    seen: set[str] = set()
    all_keys: list[str] = []
    for r in rows:
        for k in r.keys():
            if k.startswith("_"):
                continue
            if k not in seen:
                seen.add(k)
                all_keys.append(k)

    cols = [k for k in PREFERRED if k in seen]
    cols += [k for k in all_keys if k not in cols]

    wb = openpyxl.Workbook()

    # Summary first
    sum_ws = wb.active
    sum_ws.title = "Run Summary"
    sum_ws["A1"] = "Global Firewall — all classified companies"
    sum_ws["A2"] = f"Total rows: {len(rows)}"
    sum_ws["A3"] = "Source: all_classified from global_firewall_market_global.json"
    sum_ws["A4"] = f"Created: {datetime.now():%Y-%m-%d %H:%M}"
    sum_ws["A5"] = "Note: This is the full classified pool (311), NOT the FINAL Excel export subset."
    sum_ws["A6"] = f"relevant_companies in JSON: {len(data.get('relevant_companies') or [])}"
    sum_ws["A7"] = f"export_rejected_count: {data.get('export_rejected_count')}"
    sum_ws.column_dimensions["A"].width = 95

    # Simple sheet
    simple = wb.create_sheet("Companies Simple")
    simple_headers = ["#", "Company", "Role / Business Type", "Website", "Is Relevant", "HQ / Location", "Founded"]
    simple.append(simple_headers)
    style_header(simple, len(simple_headers))
    for i, r in enumerate(rows, 1):
        company = pick(r, "company", "Company", "brand", "name")
        role = pick(r, "role", "business_type", "Business Type", "channel_role", "classification")
        website = pick(r, "website", "Website", "domain")
        rel = pick(r, "is_relevant", "Is_Relevant")
        if isinstance(rel, bool):
            rel = "yes" if rel else "no"
        hq = pick(r, "headquarters", "Headquarters", "hq", "hq_city", "cities", "Cities / Regions")
        founded = pick(r, "founded", "Founded", "year_founded")
        simple.append(
            [
                i,
                company,
                cell_val(role),
                cell_val(website),
                cell_val(rel),
                cell_val(hq),
                cell_val(founded),
            ]
        )
    style_body(simple, len(simple_headers))
    for c, w in enumerate([6, 32, 26, 34, 12, 28, 12], 1):
        simple.column_dimensions[get_column_letter(c)].width = w
    simple.freeze_panes = "B2"
    simple.auto_filter.ref = f"A1:G{simple.max_row}"

    # Full sheet
    full = wb.create_sheet("All Classified 311")
    full.append(cols)
    style_header(full, len(cols))
    for r in rows:
        full.append([cell_val(r.get(c)) for c in cols])
    style_body(full, len(cols))
    for c, h in enumerate(cols, 1):
        hl = h.lower()
        if "company" in hl or hl in {"brand", "name"}:
            w = 30
        elif any(x in hl for x in ("website", "linkedin", "domain")):
            w = 32
        elif any(x in hl for x in ("summary", "source", "categories")):
            w = 36
        elif any(x in hl for x in ("headquarter", "cities")):
            w = 28
        else:
            w = 18
        full.column_dimensions[get_column_letter(c)].width = w
    full.freeze_panes = "B2"
    full.auto_filter.ref = f"A1:{get_column_letter(full.max_column)}{full.max_row}"

    wb.save(OUT)
    print(f"Wrote: {OUT}")
    print(f"Companies Simple: {simple.max_row - 1}")
    print(f"All Classified 311: {full.max_row - 1}")


if __name__ == "__main__":
    main()
