# -*- coding: utf-8 -*-
"""
Rewrite Continent / Geography + Operational Presence for Flexible Packaging FINAL.

USER DEFINITIONS:
- Continent / Geography = PHYSICAL ops — list ALL countries AND cities
  (HQ, offices, warehouses, plants, stores) that are evidenced.
  Format: NA(United States: Pleasant Prairie WI [HQ], Braselton GA; Canada: Milton ON)
- Operational Presence = COMMERCIAL sell/serve countries (all evidenced).
  Format: United States; Canada; Mexico

Never invent cities or countries.
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "geo_op_cities_audit.json"

COUNTRY_REGION = {
    "United States": "NA", "Canada": "NA", "Mexico": "NA", "Puerto Rico": "NA", "Panama": "NA",
    "Brazil": "LATAM", "Argentina": "LATAM", "Chile": "LATAM", "Colombia": "LATAM",
    "Peru": "LATAM", "Bolivia": "LATAM", "Uruguay": "LATAM", "Paraguay": "LATAM",
    "United Kingdom": "EU", "Ireland": "EU", "France": "EU", "Germany": "EU", "Spain": "EU",
    "Portugal": "EU", "Italy": "EU", "Netherlands": "EU", "Belgium": "EU", "Luxembourg": "EU",
    "Switzerland": "EU", "Austria": "EU", "Sweden": "EU", "Norway": "EU", "Denmark": "EU",
    "Finland": "EU", "Poland": "EU", "Czechia": "EU", "Romania": "EU", "Hungary": "EU",
    "Greece": "EU", "Lithuania": "EU", "Latvia": "EU", "Estonia": "EU", "Slovakia": "EU",
    "Slovenia": "EU", "Croatia": "EU", "Bulgaria": "EU", "Turkey": "EU",
    "South Africa": "AF", "Tunisia": "AF", "Angola": "AF",
    "Israel": "ME", "Saudi Arabia": "ME", "United Arab Emirates": "ME",
    "Japan": "APAC", "China": "APAC", "Hong Kong": "APAC", "Taiwan": "APAC",
    "South Korea": "APAC", "India": "APAC", "Singapore": "APAC", "Malaysia": "APAC",
    "Thailand": "APAC", "Indonesia": "APAC", "Philippines": "APAC", "Vietnam": "APAC",
    "Australia": "OC", "New Zealand": "OC",
}
REGION_ORDER = ["NA", "LATAM", "EU", "AF", "ME", "APAC", "OC"]

# country -> ordered unique cities (physical). Empty city list = country-only (ops exist, city NPD).
# Sources noted in comments.
PHYSICAL: dict[str, dict[str, list[str]]] = {
    # uline.com/Corporate/About_Locations
    "Uline": {
        "United States": [
            "Pleasant Prairie WI [HQ]", "Braselton GA", "Plainfield CT", "Kenosha WI",
            "Etna OH", "DFW Airport TX", "Ontario CA", "Naples FL", "Hudson WI",
            "Allentown PA", "Reno NV", "Lacey WA",
        ],
        "Canada": ["Milton ON", "St. Albert AB"],
        "Mexico": ["Mexicali BC", "Apodaca NL"],
    },
    # associatedbag.com/contact-us
    "Associated Bag": {
        "United States": [
            "Milwaukee WI [HQ]", "New Kingstown PA", "Dallas TX", "Sparks NV",
        ],
    },
    # hoffmann-verpackung.de/service/standorte
    "Hoffmann Verpackung": {
        "Germany": [
            "Kaarst [HQ]", "Wetter/Ruhr", "Schkeuditz", "Mittenwalde",
            "Bad Camberg", "Erding", "Ahrensburg",
        ],
        "Switzerland": ["Nesselnbach"],
    },
    # rajapack.co.uk European presence (addresses listed)
    "RAJA": {
        "France": [
            "Roissy / Paris Nord II [HQ area]", "Clichy", "Roissy-en-France",
            "Survilliers", "Neuville-en-Ferrain", "Villeneuve-Loubet",
        ],
        "United Kingdom": ["Brogborough Bedford", "Harlow Essex", "Leicester"],
        "Germany": ["Ettlingen", "Duisburg", "Grossostheim"],
        "Austria": ["Vienna", "Grodig/Glanegg", "Pregarten"],
        "Belgium": ["Tongeren", "Mouscron", "Antwerp"],
        "Netherlands": ["Breda", "Venlo"],
        "Spain": ["Ripollet / Barcelona", "Galdakao"],
        "Italy": ["Castel San Giovanni", "Castelletto Cervo"],
        "Switzerland": ["Pratteln", "Baar", "Aarau"],
        "Poland": ["Warsaw"],
        "Czechia": ["Ricany u Prahy / Nupaky"],
        "Denmark": ["Soborg"],
        "Sweden": ["Landvetter"],
        "Norway": ["Vestby"],
        "Portugal": ["Cacem"],
        "Slovakia": ["Bratislava"],
        "Hungary": ["Budapest"],
        "Ireland": ["Dublin"],
        "Luxembourg": ["Differdange/Niederkorn"],
    },
    "RAJAPACK UK": {
        "United Kingdom": ["Brogborough Bedford [HQ/DC]", "Harlow Essex"],
    },
    # nnz.com/locations — countries with owned subsidiaries (cities often not on page)
    "NNZ": {
        "Netherlands": ["Groningen [HQ]"],
        "Germany": ["Zentrale", "Lueneburg (Nord)", "Wesel (West)", "Mutterstadt (Sued)"],
        "Austria": [],
        "Switzerland": [],
        "Denmark": [],
        "Norway": [],
        "Finland": [],
        "Poland": [],
        "Estonia": [],
        "Latvia": [],
        "Lithuania": [],
        "United Kingdom": [],
        "United States": [],
        "Canada": [],  # Alberta / West / East offices
        "South Africa": [],
        "China": [],  # purchase office
    },
    "NNZ GmbH": {
        "Germany": ["Lüneburg", "Wesel", "Mutterstadt"],
    },
    # Ahlsell countries of operation (cities from that page / reports)
    "Ahlsell": {
        "Sweden": ["Stockholm [HQ]", "Hallsberg (logistics)"],
        "Norway": ["Oslo / Stavanger"],
        "Finland": ["Hyvinkaa"],
        "Denmark": ["Brondby"],
        "Estonia": ["Tallinn"],
        "Poland": [],
        "China": [],  # sourcing office
    },
    # Antalis ESG operating entities — countries; HQ city FR
    "Antalis": {
        "France": ["Boulogne-Billancourt [HQ]"],
        "Austria": [], "Belgium": [], "Bolivia": [], "Bulgaria": [], "Canada": [],
        "Chile": [], "Czechia": [], "Denmark": [], "Estonia": [], "Finland": [],
        "Germany": [], "Hungary": [], "Ireland": [], "Italy": [], "Latvia": [],
        "Lithuania": [], "Mexico": [], "Netherlands": [], "Norway": [], "Peru": [],
        "Poland": [], "Portugal": [], "Romania": [], "Slovakia": [], "Spain": [],
        "Sweden": [], "Switzerland": [], "Turkey": [], "United Kingdom": [],
    },
    # Lyreco direct ops countries (cities: Marly HQ)
    "Lyreco": {
        "France": ["Marly [HQ]"],
        "Austria": [], "Belgium": [], "Czechia": [], "Denmark": [], "Finland": [],
        "Germany": [], "Hungary": [], "Ireland": [], "Italy": [], "Luxembourg": [],
        "Netherlands": [], "Norway": [], "Poland": [], "Portugal": [], "Slovakia": [],
        "Spain": [], "Sweden": [], "Switzerland": [], "United Kingdom": [],
        "Hong Kong": [], "South Korea": [], "Malaysia": [], "Singapore": [], "Thailand": [],
    },
    # Bunzl Wikipedia countries + known regional HQs (not inventing 150+ warehouses)
    "Bunzl": {
        "United Kingdom": ["London [HQ]"],
        "Ireland": [],
        "United States": ["Creve Coeur / St. Louis MO (Bunzl Distribution USA HQ)"],
        "Canada": ["Burlington ON (Bunzl Canada HQ)"],
        "Mexico": [],
        "Puerto Rico": [],
        "Australia": ["Mount Waverley VIC (Asia Pacific HQ)"],
        "New Zealand": [],
        "Austria": [], "Belgium": [], "Brazil": [], "Chile": [], "China": [],
        "Colombia": [], "Czechia": [], "Denmark": [], "France": [], "Germany": [],
        "Hong Kong": [], "Hungary": [], "Israel": [], "Italy": [], "Netherlands": [],
        "Peru": [], "Poland": [], "Romania": [], "Singapore": [], "Slovakia": [],
        "Spain": [], "Switzerland": [], "Turkey": [],
    },
    "Veritiv": {
        "United States": ["Atlanta GA [HQ]"],
        "Mexico": [],
    },
    "Imperial Dade": {
        "United States": ["Jersey City NJ [HQ]"],
        "Canada": [],
    },
    "Global Industrial": {
        "United States": ["Port Washington NY [HQ]"],
    },
    "Global Industrial Canada": {
        "Canada": ["Markham ON [HQ]"],
    },
    "Fastenal": {
        "United States": ["Winona MN [HQ]"],
        "Canada": [],
        "Mexico": [],
        "Puerto Rico": [],
        "Panama": [],
        "China": [],
    },
    "Manutan": {
        "France": ["Gonesse [HQ]", "Niort", "Veauche", "Strasbourg"],
        "United Kingdom": [],
        "Belgium": [],
        "Netherlands": [],
        "Germany": [],
        "Italy": [],
        "Spain": [],
        "Portugal": [],
        "Denmark": [],
        "Sweden": [],
        "Norway": [],
        "Finland": [],
        "Czechia": [],
        "Slovakia": [],
        "Hungary": [],
        "Poland": [],
        "Switzerland": [],
    },
    "Brangs + Heinrich": {
        "Germany": ["Solingen [HQ]"],
        "Netherlands": ["Utrecht"],
        "Belgium": ["Dison"],
    },
    "Verpacken Online": {"Germany": ["Solingen [HQ]"]},
    "Inapa": {
        "Portugal": ["Sintra [HQ]"],
        "Germany": ["Hamburg"],
        "France": ["Dourdan"],
        "Spain": [], "Belgium": [], "Luxembourg": [], "Turkey": [], "Angola": [],
    },
    "Papier Union": {"Germany": ["Hamburg"]},
    "Askul": {"Japan": ["Tokyo [HQ]"]},
    "MonotaRO": {"Japan": ["Osaka [HQ]"]},
    "Staples": {"United States": ["Framingham MA [HQ]"]},
    "Quill": {"United States": ["Lincolnshire IL [HQ]"]},
    "The Home Depot": {"United States": ["Atlanta GA [HQ]"]},
    "McMaster-Carr": {"United States": ["Elmhurst IL [HQ]"]},
    "Officeworks": {"Australia": ["Chadstone VIC [HQ]"]},
    "Blackwoods": {"Australia": ["Macquarie Park NSW [HQ]"]},
    "Winc": {"Australia": ["Mascot NSW [HQ]"]},
    "Ball & Doggett": {"Australia": ["Dandenong South VIC [HQ]"]},
    "Spicers": {"Australia": ["Dandenong South VIC [HQ]"]},
    "Spicers Paper": {"Canada": ["Vaughan ON [HQ]"]},
    "Cospak": {"Australia": ["Minto NSW [HQ]"]},
    "Signet": {"Australia": ["Wakerley QLD [HQ]"]},
    "Pluspak": {
        "Australia": ["Scoresby VIC [HQ]", "Milperra NSW", "Crestmead QLD", "Canning Vale WA"],
    },
    "Ozpack": {"Australia": ["Silverwater NSW [HQ/warehouse]"]},
    "Macfarlane Packaging": {"United Kingdom": ["Glasgow [HQ]"]},
    "National Flexible": {"United Kingdom": ["Birkenshaw / Bradford [HQ]"]},
    "Abco Kovex": {"Ireland": ["Swords / Dublin [HQ]"]},
    "The Packaging Centre": {
        "Ireland": ["Dublin [HQ]"],
        "United Kingdom": ["Saintfield / County Down"],
    },
    "McLernon Packaging": {"Ireland": ["Dublin [HQ]"]},
    "Viking Direct": {"Ireland": ["Dublin [HQ]"]},
    "Mondoffice": {"France": ["Survilliers [HQ]"]},
    "Kalamazoo": {"Spain": ["Galdakao [HQ]"]},
    "JPG": {"France": ["Survilliers [HQ]"]},
    "Syspex": {"Singapore": ["Joo Koon [HQ]"], "Malaysia": []},
    "Seow Khim Polythene": {"Singapore": ["Pasir Ris / Loyang [HQ]"]},
    "Tantraco": {"Singapore": ["Ang Mo Kio [HQ]"]},
    "FlexPRO": {"Bulgaria": ["Sofia [HQ]"]},
    "Flexi-Films": {"United Kingdom": ["Annesley / Nottingham [HQ]"]},
    "Empaques del Norte": {"Mexico": ["San Nicolas de los Garza NL [HQ]"]},
    "Transpak": {"Germany": ["Solms [HQ]"]},
    "ESWE Versandpack": {"Germany": ["Sternenfels [HQ]"]},
    "Packaging2Buy": {"United Kingdom": ["Stevenage Hertfordshire [ops]"]},
    "Packaging Express": {"United Kingdom": ["Hayes Middlesex [main depot/HQ]"]},
    "Swiftpak": {"United Kingdom": ["Theale / Reading [HQ]"]},
    "Polythene UK": {"United Kingdom": ["Witney [HQ]"]},
    "UK Packaging Supplies": {"United Kingdom": ["Tottenham / London [HQ]"]},
    "Unicorr Packaging": {"United States": ["North Haven CT [HQ]"]},
    "Uboxes": {"United States": ["Miramar FL [HQ]"]},
    "The Bottle Crew": {"United States": ["Farmington Hills MI [HQ]"]},
    "The Boxery": {"United States": ["Brooklyn NY [HQ]"]},
    "Stamar Packaging": {"United States": ["Elgin IL [HQ]"]},
    "Schermerhorn Bros.": {"United States": ["Lombard IL [HQ]"]},
    "Preferred Plastics & Packaging": {"United States": ["Lyndhurst NJ [HQ]"]},
    "Plastic Bag Source": {"United States": ["Monsey NY [HQ]"]},
    "Packaging Partners": {"United States": ["Franklin WI [HQ]"]},
    "Packaging Price": {"United States": ["Warrenville IL [HQ]"]},
    "Great Pacific Packaging": {"United States": ["Fife WA [HQ]"]},
    "Innovative Packaging Group": {"United States": ["Dallas TX [HQ]"]},
    "Imperial Supplies": {"United States": ["Green Bay WI [HQ]"]},
    "R.J. Schinner Co.": {"United States": ["Menomonee Falls WI [HQ]"]},
    "Gertex Solutions": {"Canada": ["Etobicoke ON [HQ]"]},
    "Poly Pak": {"Canada": ["Vancouver BC [HQ]"]},
    "Super Poly": {"Canada": ["Edmonton AB [HQ]"]},
    "Unipac Packaging Products": {"Canada": ["Edmonton AB [HQ]"]},
    "PakFactory": {"Canada": ["Markham ON [HQ]"]},
    "Atlantic Packaging Products": {"Canada": ["Scarborough ON [HQ]"]},
    "Richards Packaging": {"Canada": ["Mississauga ON [HQ]"]},
    "Crawford Packaging": {"Canada": ["London ON [HQ]"]},
    "Polyair": {
        "Canada": ["Etobicoke ON"],
        "United States": ["Sarasota FL (IPG parent HQ via itape.com)"],
    },
    "Furodet Emballages": {
        "France": ["Sennecey-le-Grand [HQ]", "Ormes (warehouse)"],
    },
    "Gerdmans": {"Sweden": ["Markaryd [HQ]"]},
    "Ratioform": {"Germany": ["Pliening [HQ]"]},
    "Papyrus": {"Sweden": ["Molndal [HQ]"]},
    "Igepa Group": {"Germany": ["Hamburg [HQ]"]},
    "Karton.eu": {"Germany": ["Spremberg [HQ]"]},
    "Network Packaging": {"United Kingdom": ["Wolverhampton [HQ]"]},
    "GWP Group": {"United Kingdom": ["Cricklade Wiltshire [HQ]"]},
    "GM Packaging": {"United Kingdom": ["Murton / Seaham [HQ]"]},
    "Industrial Packaging": {"United Kingdom": ["Southport [HQ]"]},
    "Lesters Packaging": {"United Kingdom": ["Burntwood Staffordshire [HQ]"]},
    "Multiplastics": {"United Kingdom": ["Northwich Cheshire [HQ]"]},
    "Complete Packaging Solutions": {"United Kingdom": ["Long Crendon Buckinghamshire [HQ]"]},
    "Diamond Packaging": {"United Kingdom": ["Canvey Island Essex [HQ]"]},
    "Able Packaging": {"United Kingdom": ["Leeds [HQ]"]},
    "Pack Supplies": {"United Kingdom": ["Thatcham Berkshire [ops]"]},
    "Packaging World": {"United Kingdom": ["Swansea [HQ]"]},
    "Seal Packaging": {"United Kingdom": ["Luton Bedfordshire [ops]"]},
    "Transpack": {"United Kingdom": ["Totton / Southampton [ops]"]},
    "Packaging House": {"New Zealand": ["Wiri / Auckland [HQ]"]},
    "Packaging Products": {"New Zealand": ["Lower Hutt / Wellington [HQ]"]},
    "Premier Packaging": {"New Zealand": ["Mount Wellington / Auckland [HQ]"]},
    "Moglix": {"India": ["Noida [HQ]"]},
    "IndustryBuying": {"India": ["New Delhi [HQ]"]},
    "India Pack": {"India": ["Pune [HQ]"]},
    "Allied Distributors": {"South Africa": []},
    "Office Depot": {"Mexico": []},
    "SupplyOne": {"United States": ["Newtown Square PA [HQ]"]},
    "Packaging Corporation of America": {"United States": ["Lake Forest IL [HQ]"]},
    "Shorr Packaging": {"United States": ["Aurora IL [HQ]"]},
    "WebstaurantStore": {"United States": ["Lititz PA [HQ]"]},
    "Aramsco": {"United States": ["Radnor PA [HQ]"]},
    "W.B. Mason": {"United States": ["Brockton MA [HQ]"]},
    "Hisco": {"United States": ["Houston TX [HQ]"]},
    "Kaufman Container": {"United States": ["Cleveland OH [HQ]"]},
    "O.Berk": {"United States": ["Union NJ [HQ]"]},
    "SKS Bottle & Packaging": {"United States": ["Saratoga Springs NY [HQ]"]},
    "HD Supply": {"United States": ["Atlanta GA [HQ]"]},
    "BradyPLUS": {"United States": ["Las Vegas NV [HQ]"]},
    "Zoro Tools": {"United States": ["Buffalo Grove IL [HQ]"]},
    "Acme Paper & Supply": {"United States": ["Savage MD [HQ]"]},
    "All American Containers": {"United States": ["Miami FL [HQ]"]},
}

# Commercial countries (full lists). May be wider than physical (partners/exports).
OP: dict[str, list[str]] = {
    "Uline": ["United States", "Canada", "Mexico"],
    "Associated Bag": ["United States"],
    "Hoffmann Verpackung": ["Germany", "Switzerland"],
    "RAJA": [
        "Austria", "Belgium", "Czechia", "Denmark", "France", "Germany", "Hungary",
        "Ireland", "Italy", "Luxembourg", "Netherlands", "Norway", "Poland", "Portugal",
        "Slovakia", "Spain", "Sweden", "Switzerland", "United Kingdom",
    ],
    "RAJAPACK UK": ["United Kingdom"],
    "NNZ": [
        "Netherlands", "Germany", "Austria", "Switzerland", "Denmark", "Norway",
        "Finland", "Poland", "Estonia", "Latvia", "Lithuania", "United Kingdom",
        "United States", "Canada", "South Africa", "China",
    ],
    "NNZ GmbH": ["Germany"],
    "Ahlsell": ["Sweden", "Norway", "Finland", "Denmark", "Estonia", "Poland"],
    "Antalis": [
        "Austria", "Belgium", "Bolivia", "Bulgaria", "Canada", "Chile", "Czechia",
        "Denmark", "Estonia", "Finland", "France", "Germany", "Hungary", "Ireland",
        "Italy", "Latvia", "Lithuania", "Mexico", "Netherlands", "Norway", "Peru",
        "Poland", "Portugal", "Romania", "Slovakia", "Spain", "Sweden", "Switzerland",
        "Turkey", "United Kingdom",
    ],
    "Lyreco": [
        "Austria", "Belgium", "Czechia", "Denmark", "France", "Finland", "Germany",
        "Hungary", "Ireland", "Italy", "Luxembourg", "Netherlands", "Norway", "Poland",
        "Portugal", "Slovakia", "Spain", "Sweden", "Switzerland", "United Kingdom",
        "Hong Kong", "South Korea", "Malaysia", "Singapore", "Thailand",
        "United States", "Canada", "Mexico", "Brazil", "Argentina",
        "Latvia", "Lithuania", "Estonia", "Romania", "Turkey",
        "Taiwan", "China", "India", "Australia", "New Zealand",
    ],
    "Bunzl": [
        "United Kingdom", "Ireland", "United States", "Canada", "Mexico", "Puerto Rico",
        "Australia", "New Zealand", "Austria", "Belgium", "Brazil", "Chile", "China",
        "Colombia", "Czechia", "Denmark", "France", "Germany", "Hong Kong", "Hungary",
        "Israel", "Italy", "Netherlands", "Peru", "Poland", "Romania", "Singapore",
        "Slovakia", "Spain", "Switzerland", "Turkey",
    ],
    "Veritiv": ["United States", "Mexico"],
    "Imperial Dade": ["United States", "Canada"],
    "Global Industrial": ["United States"],
    "Global Industrial Canada": ["Canada"],
    "Fastenal": ["United States", "Canada", "Mexico", "Puerto Rico", "Panama"],
    "Manutan": [
        "France", "United Kingdom", "Belgium", "Netherlands", "Germany", "Italy",
        "Spain", "Portugal", "Denmark", "Sweden", "Norway", "Finland", "Czechia",
        "Slovakia", "Hungary", "Poland", "Switzerland",
    ],
    "Brangs + Heinrich": ["Germany", "Netherlands", "Belgium"],
    "Inapa": [
        "Portugal", "Germany", "France", "Spain", "Belgium", "Luxembourg", "Turkey", "Angola",
    ],
    "The Packaging Centre": ["Ireland", "United Kingdom"],
    "Pluspak": ["Australia"],
    "Syspex": ["Singapore", "Malaysia"],
    "Polyair": ["Canada", "United States"],
    "Mondoffice": ["France"],
    "Kalamazoo": ["Spain"],
}


def canon_country(token: str) -> str | None:
    t = (token or "").strip()
    if not t:
        return None
    low = t.lower().rstrip(".")
    alias = {
        "usa": "United States", "us": "United States", "u.s.": "United States",
        "uk": "United Kingdom", "u.k.": "United Kingdom", "england": "United Kingdom",
        "czech republic": "Czechia", "czechia": "Czechia",
        "holland": "Netherlands", "the netherlands": "Netherlands",
        "espagne": "Spain", "deutschland": "Germany",
        "hong-kong": "Hong Kong", "korea": "South Korea",
        "uae": "United Arab Emirates",
    }
    if low in alias:
        return alias[low]
    specials = {
        "united states": "United States", "united kingdom": "United Kingdom",
        "united arab emirates": "United Arab Emirates", "south africa": "South Africa",
        "south korea": "South Korea", "new zealand": "New Zealand",
        "hong kong": "Hong Kong", "saudi arabia": "Saudi Arabia",
        "puerto rico": "Puerto Rico",
    }
    if low in specials:
        return specials[low]
    if t in COUNTRY_REGION:
        return t
    titled = t.title() if t.islower() or t.isupper() else t
    return titled if titled in COUNTRY_REGION else None


US_STATE_RE = re.compile(
    r"\b(AL|AK|AZ|AR|CA|CO|CT|DE|FL|GA|HI|ID|IL|IN|IA|KS|KY|LA|ME|MD|MA|MI|MN|MS|"
    r"MO|MT|NE|NV|NH|NJ|NM|NY|NC|ND|OH|OK|OR|PA|RI|SC|SD|TN|TX|UT|VT|VA|WA|WV|WI|WY|DC)\b",
    re.I,
)
CA_PROV_RE = re.compile(
    r"\b(Ontario|Quebec|Alberta|British Columbia|Manitoba|Saskatchewan|"
    r"Nova Scotia|New Brunswick|ON|QC|AB|BC|MB|SK)\b",
    re.I,
)


def parse_hq(hq: str) -> tuple[str | None, str | None]:
    """Return (country, city_label) from Headquarters cell."""
    s = (hq or "").strip()
    if not s or "not publicly" in s.lower():
        return None, None
    s_clean = re.sub(r"\([^)]*\)", " ", s)
    low = s_clean.lower()
    country = None
    if "canada" in low:
        country = "Canada"
    elif any(x in low for x in ("united states", "usa", "u.s.")):
        country = "United States"
    elif any(x in low for x in ("united kingdom", "england", "scotland", "wales")):
        country = "United Kingdom"
    if "," in s_clean:
        parts = [p.strip() for p in s_clean.split(",") if p.strip()]
        if not country:
            for part in reversed(parts):
                c = canon_country(re.split(r"[;/]", part)[0].strip())
                if c:
                    country = c
                    break
        # city = first segment(s) before country
        city_parts = []
        for part in parts:
            c = canon_country(re.split(r"[;/]", part)[0].strip())
            if c and c == country:
                break
            if canon_country(part):
                continue
            city_parts.append(part)
        city = ", ".join(city_parts).strip(" ,") if city_parts else None
        if city and country:
            return country, f"{city} [HQ]"
        if country:
            return country, None
    if not country:
        country = canon_country(s_clean)
    if not country and US_STATE_RE.search(s_clean):
        country = "United States"
        # "Mansfield, OH" / "New York, NY"
        m = re.match(r"^([^,]+),\s*[A-Z]{2}\b", s_clean.strip())
        if m:
            return country, f"{m.group(1).strip()} [HQ]"
        return country, f"{s_clean.strip()} [HQ]"
    if not country and CA_PROV_RE.search(s_clean):
        country = "Canada"
    if country and "," not in s_clean:
        # bare city sometimes
        if not canon_country(s_clean):
            return country, f"{s_clean} [HQ]"
    return country, None


def format_geo(physical: dict[str, list[str]]) -> str:
    if not physical:
        return "Not publicly disclosed"
    buckets: dict[str, list[str]] = {r: [] for r in REGION_ORDER}
    for country, cities in physical.items():
        cc = canon_country(country) or country
        if cc not in COUNTRY_REGION:
            continue
        cities = [c for c in cities if c]
        if cities:
            buckets[COUNTRY_REGION[cc]].append(f"{cc}: {', '.join(cities)}")
        else:
            buckets[COUNTRY_REGION[cc]].append(cc)
    parts = []
    for r in REGION_ORDER:
        if buckets[r]:
            parts.append(f"{r}({'; '.join(buckets[r])})")
    return "; ".join(parts) if parts else "Not publicly disclosed"


def format_op(countries: list[str]) -> str:
    seen = set()
    out = []
    for c in countries:
        cc = canon_country(c)
        if cc and cc not in seen and cc in COUNTRY_REGION:
            seen.add(cc)
            out.append(cc)
    return "; ".join(out) if out else "Not publicly disclosed"


def main() -> None:
    # Ensure Bulgaria in map for FlexPRO
    COUNTRY_REGION.setdefault("Bulgaria", "EU")

    wb = load_workbook(FP)
    ws = wb.active
    assert ws.cell(1, 5).value == "Continent / Geography"
    assert ws.cell(1, 6).value == "Operational Presence"

    audit = {"updated_at": datetime.now().isoformat(timespec="seconds"), "rows": []}
    n_city_rich = n_hq = n_weak = 0

    for row in ws.iter_rows(min_row=2):
        name = str(row[0].value or "").strip()
        hq = str(row[3].value or "").strip()
        if not name:
            continue
        old_geo = str(row[4].value or "").strip()
        old_op = str(row[5].value or "").strip()

        hq_country, hq_city = parse_hq(hq)
        source = "hq_city"

        if name in PHYSICAL:
            phys = {k: list(v) for k, v in PHYSICAL[name].items()}
            # ensure HQ city present
            if hq_country and hq_country in phys and hq_city:
                if not any("[HQ]" in c for c in phys[hq_country]):
                    phys[hq_country] = [hq_city] + phys[hq_country]
            elif hq_country and hq_country not in phys and hq_country in COUNTRY_REGION:
                phys[hq_country] = [hq_city] if hq_city else []
            source = "verified_locations"
            n_city_rich += 1
        elif hq_country:
            phys = {hq_country: [hq_city] if hq_city else []}
            source = "hq_city"
            n_hq += 1
        else:
            phys = {}
            source = "none"
            n_weak += 1

        if name in OP:
            op_list = list(OP[name])
        else:
            op_list = list(phys.keys())

        new_geo = format_geo(phys)
        new_op = format_op(op_list)
        row[4].value = new_geo
        row[5].value = new_op

        audit["rows"].append({
            "company": name,
            "source": source,
            "hq": hq,
            "old_geo": old_geo,
            "new_geo": new_geo,
            "old_op": old_op,
            "new_op": new_op,
        })

    wb.save(FP)
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"saved {FP}")
    print(f"verified_location_overlays={n_city_rich} hq_derived={n_hq} weak={n_weak}")
    print(f"audit {AUDIT}")


if __name__ == "__main__":
    main()
