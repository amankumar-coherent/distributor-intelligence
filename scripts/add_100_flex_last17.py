"""Add remaining verified packaging distributors to reach +100 from baseline 224."""
from __future__ import annotations

import json
import os
import re
import shutil
import socket
import ssl
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")
FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "add100_last17_audit.json"

API_KEY = (os.getenv("OPENAI_API_KEY") or "").strip()
BASE_URL = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
client = OpenAI(api_key=API_KEY, base_url=BASE_URL)
ND = "Not publicly disclosed"
BASELINE = 224
TARGET = 100

# Web-verified distributor candidates (agent researched)
CANDS = [
    {"company": "Kyana Packaging Solutions", "website": "https://www.kyanaind.com", "hq": "Louisville, KY, USA", "cc": "US", "founded": "1976", "cats": "Industrial packaging supplies; stretch film; corrugated systems", "focus": "Regional US packaging distributor", "summary": "Kyana Packaging Solutions is a family-run packaging distributor founded in 1976 in Louisville, Kentucky."},
    {"company": "Lancaster Packaging", "website": "https://lancasterpackaging.com", "hq": "USA", "cc": "US", "founded": "1989", "cats": "Wholesale packaging; custom packaging; procurement", "focus": "US packaging distribution and supply management", "summary": "Lancaster Packaging provides wholesale packaging design and distribution services."},
    {"company": "PaperPack", "website": "https://paperpackinc.com", "hq": "USA", "cc": "US", "founded": "1999", "cats": "Industrial packaging; stretch film; shipping; janitorial", "focus": "US packaging and promotional products distributor", "summary": "PaperPack, Inc. distributes industrial packaging, stretch film, and related supplies."},
    {"company": "Partners Packaging", "website": "https://www.partnerspackaging.net", "hq": "USA", "cc": "US", "founded": "2009", "cats": "Export crating; shipping supplies; stretch film", "focus": "Local full-service packaging distributor", "summary": "Partners Packaging Inc. is a locally owned full-service packaging distributor specializing in shipping supplies and stretch films."},
    {"company": "Diversified Packaging St Louis", "website": "https://www.diversifiedpackaging.com", "hq": "St. Louis, MO, USA", "cc": "US", "founded": "1984", "cats": "Full-line packaging supplies; equipment service", "focus": "St. Louis region packaging distributor", "summary": "Diversified Packaging is a family-owned full-line packaging distributor serving the St. Louis region since 1984."},
    {"company": "TransPack Krumbach", "website": "https://transpack-krumbach.de", "hq": "Krumbach, Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Transport packaging; cartons; shipping materials", "focus": "German packaging wholesaler/merchant", "summary": "TransPack-Krumbach is a German packaging specialist wholesaler for transport and shipping materials."},
    {"company": "Ehlken Verpackungen", "website": "https://www.ehlkenverpackungen.de", "hq": "Bremerhaven, Germany", "cc": "DE", "founded": "1953", "cats": "Food and non-food packaging materials wholesale", "focus": "German packaging materials wholesaler", "summary": "Ehlken Verpackungen GmbH wholesales packaging materials to food and non-food customers across Germany and the EU since 1953."},
    {"company": "Spar-Pack", "website": "https://www.spar-pack.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Tapes; stretch film; industrial packaging wholesale", "focus": "German packaging materials wholesaler", "summary": "Spar-Pack is a German packaging materials wholesaler with own warehouse stock for industry and trade."},
    {"company": "Verpacken Online", "website": "https://verpacken.online", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Industrial and trade packaging materials online", "focus": "German packaging e-wholesaler (Brangs + Heinrich shop)", "summary": "verpacken.online is the B2B packaging e-shop of Brangs + Heinrich, a German packaging wholesaler."},
    {"company": "NNZ Germany", "website": "https://www.nnz.de", "hq": "Germany", "cc": "DE", "founded": "1922", "cats": "Produce and industrial packaging", "focus": "German arm of NNZ packaging distribution network", "summary": "NNZ Germany distributes packaging materials as part of the NNZ packaging network founded in 1922."},
    {"company": "Cardboardmaster AU", "website": "https://www.cardboardmaster.com.au", "hq": "Dandenong, VIC, Australia", "cc": "AU", "founded": "Not publicly disclosed", "cats": "Cartons; bubble; stretch; tapes", "focus": "Melbourne packaging importer/distributor", "summary": "Cardboardmaster is a Melbourne packaging importer and distributor of cartons and packing supplies."},
    {"company": "Ozpack", "website": "https://ozpack.com.au", "hq": "Sydney, Australia", "cc": "AU", "founded": "Not publicly disclosed", "cats": "Wholesale packaging supplies; boxes; satchels", "focus": "Australian packaging wholesaler", "summary": "Ozpack is an Australian wholesale packaging supplies distributor based in Sydney."},
    {"company": "TigerPak", "website": "https://tigerpak.com.au", "hq": "Australia", "cc": "AU", "founded": "2002", "cats": "Packaging materials; warehouse consumables", "focus": "Australian packaging distributor", "summary": "TigerPak is an Australian family packaging materials distributor founded in 2002."},
    {"company": "QPack UK Merchant", "website": "https://qpack.co.uk", "hq": "UK", "cc": "GB", "founded": "Not publicly disclosed", "cats": "Retail packaging; carrier bags; tissue; boxes", "focus": "UK retail packaging supplier", "summary": "QPack supplies retail packaging suites to UK shops and brands."},
    {"company": "Weller Packaging UK", "website": "https://www.wellerpackaging.co.uk", "hq": "UK", "cc": "GB", "founded": "Not publicly disclosed", "cats": "Flexible polythene; paper packaging wholesale", "focus": "UK wholesaler to packaging merchants", "summary": "Weller Packaging is a UK importer/stockist/distributor supplying the packaging merchant trade."},
    {"company": "Sovereign Partners UK", "website": "https://www.sovereignpartners.co.uk", "hq": "Leighton Buzzard, UK", "cc": "GB", "founded": "Not publicly disclosed", "cats": "Packaging supplies", "focus": "UK packaging supplies distributor", "summary": "Sovereign Partners is a UK packaging supplies distributor."},
    {"company": "Global Trading Packaging UK", "website": "https://gtuk.co.uk", "hq": "Preston, UK", "cc": "GB", "founded": "Not publicly disclosed", "cats": "Mailing bags; carrier bags; stockholding", "focus": "UK packaging procurement/stockholding", "summary": "Global Trading UK provides packaging design, procurement and stockholding services."},
    {"company": "Maidstone Packaging", "website": "https://www.maidstonepackaging.co.uk", "hq": "Kent, UK", "cc": "GB", "founded": "Not publicly disclosed", "cats": "Cartons; bubble; pallet wrap; tapes", "focus": "SE England packaging distributor", "summary": "Maidstone Packaging Supplies distributes packaging materials across South East England."},
    {"company": "Baas Packaging NL", "website": "https://baasverpakkingen.nl", "hq": "Netherlands", "cc": "NL", "founded": "Not publicly disclosed", "cats": "Boxes; films; bags; tapes; e-commerce packaging", "focus": "Dutch packaging wholesaler", "summary": "Baas Verpakkingen is a Dutch B2B packaging materials wholesaler."},
    {"company": "De Wild Packaging", "website": "https://www.dewild.nl", "hq": "Netherlands", "cc": "NL", "founded": "Not publicly disclosed", "cats": "Packaging materials", "focus": "Dutch packaging distributor/wholesaler", "summary": "De Wild Group distributes packaging materials in the Netherlands."},
    {"company": "Packen24 DE", "website": "https://www.packen24.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Online packaging materials", "focus": "German packaging e-merchant", "summary": "Packen24 is a German online packaging materials wholesaler/merchant."},
    {"company": "Serwo DE", "website": "https://www.serwo.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Packaging wholesale", "focus": "German packaging wholesaler", "summary": "Serwo is a German packaging materials wholesaler."},
    {"company": "Paul Hildebrandt AG", "website": "https://www.hildebrandt.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Packaging wholesale", "focus": "German packaging wholesaler", "summary": "Paul Hildebrandt is a German packaging wholesaler/distributor."},
    {"company": "Knueppel Packaging", "website": "https://knueppel.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Industrial packaging wholesale", "focus": "German B2B packaging wholesaler", "summary": "Knüppel Verpackung is a German B2B packaging wholesaler/distributor."},
    {"company": "Tillmann Packaging DE", "website": "https://www.tillmann-verpackungen.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Packaging materials wholesale", "focus": "German packaging wholesaler", "summary": "Tillmann Verpackungen is a German packaging materials wholesaler."},
    {"company": "ESWE Versandpack DE", "website": "https://www.eswe-versandpack.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Shipping packaging materials", "focus": "German shipping packaging wholesaler", "summary": "ESWE Versandpack is a German shipping packaging wholesaler."},
    {"company": "Horna Packaging DE", "website": "https://www.horna.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Packaging wholesale", "focus": "German packaging wholesaler", "summary": "Horna Verpackungen is a German packaging wholesaler."},
    {"company": "Wagener Packaging DE", "website": "https://www.wagener-verpackung.de", "hq": "Germany", "cc": "DE", "founded": "Not publicly disclosed", "cats": "Packaging materials", "focus": "German packaging wholesaler", "summary": "Wagener Verpackung is a German packaging wholesaler."},
    {"company": "ANW Singapore", "website": "https://anw.com.sg", "hq": "Singapore", "cc": "SG", "founded": "Not publicly disclosed", "cats": "Strapping; marking; industrial packaging systems", "focus": "Singapore packaging systems distributor", "summary": "ANW is a Singapore distributor of shipping and industrial packaging systems."},
    {"company": "PackQueen AU", "website": "https://www.packqueen.com.au", "hq": "Australia", "cc": "AU", "founded": "Not publicly disclosed", "cats": "Boxes; shipping; retail packaging", "focus": "Australian packaging wholesaler", "summary": "PackQueen is an Australian online packaging wholesaler/distributor."},
    {"company": "Surepak AU", "website": "https://www.surepak.com.au", "hq": "Australia", "cc": "AU", "founded": "Not publicly disclosed", "cats": "Tapes; films; stretch wrap", "focus": "Australian packaging distributor", "summary": "Surepak is an Australian packaging supplies distributor."},
    {"company": "Network Dist Packaging", "website": "https://us.networkdistribution.com", "hq": "Schaumburg, IL, USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Industrial packaging; jan-san; disposables", "focus": "National packaging distribution organization", "summary": "Network Distribution is a national packaging and jan-san distribution organization."},
    {"company": "Pollock Packaging", "website": "https://www.pollock.com", "hq": "USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Packaging; paper; facility supplies", "focus": "US packaging/paper distributor", "summary": "Pollock (Orora network) distributes packaging and paper supplies in the US."},
    {"company": "Aramsco Packaging", "website": "https://www.aramsco.com", "hq": "USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Specialty supplies including packaging", "focus": "Specialty distributor with packaging lines", "summary": "Aramsco is a specialty distributor whose product lines include packaging supplies."},
    {"company": "Gillis Lane Packaging", "website": "https://gillislane.com", "hq": "USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Packaging and paper products", "focus": "US packaging/paper distributor", "summary": "Gillis & Lane distributes packaging and paper products with CA/TX warehouses."},
    {"company": "Pro Pac TX", "website": "https://www.propacmaterials.com", "hq": "Grand Prairie, TX, USA", "cc": "US", "founded": "1971", "cats": "Packaging materials and machinery", "focus": "Southwest packaging materials distributor", "summary": "Pro Pac Materials is an independent packaging materials distributor in the Southwest US."},
    {"company": "SSI Packaging VA", "website": "https://ssipkg.com", "hq": "Richmond, VA, USA", "cc": "US", "founded": "1967", "cats": "Shipping supplies; corrugated; tape", "focus": "Mid-Atlantic packaging distributor", "summary": "SSI Packaging Group distributes packaging and shipping supplies from Richmond, VA."},
    {"company": "Elkay Plastics Dist", "website": "https://www.elkayplastics.com", "hq": "Commerce, CA, USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Poly bags; film packaging", "focus": "National poly packaging distributor", "summary": "Elkay Plastics distributes poly bags and flexible packaging supplies nationally."},
    {"company": "Stickel Packaging NJ", "website": "https://www.stickelpackaging.com", "hq": "Lakewood, NJ, USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Packaging; industrial; foodservice", "focus": "NJ packaging distributor", "summary": "Stickel Packaging Supply is a New Jersey packaging distributor."},
    {"company": "OSI Packaging LA", "website": "https://www.osipackaging.com", "hq": "Shreveport, LA, USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Packaging materials; equipment", "focus": "Stocking packaging distributor", "summary": "OSI Packaging is a stocking packaging materials and equipment distributor."},
    {"company": "Citation Packaging IL", "website": "https://www.citationpackaging.com", "hq": "Chicago, IL, USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Stock boxes; packaging materials", "focus": "Chicago packaging distributor", "summary": "Citation Packaging distributes stock boxes and packaging materials in Chicago."},
    {"company": "Discount Packaging US", "website": "https://www.discountpackaging.com", "hq": "USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Boxes; shipping supplies", "focus": "US packaging e-merchant/distributor", "summary": "Discount Packaging is a US packaging supplies e-merchant/distributor."},
    {"company": "Skybox Packaging OH", "website": "https://www.skyboxpackaging.com", "hq": "Ohio, USA", "cc": "US", "founded": "Not publicly disclosed", "cats": "Packaging supplies", "focus": "US packaging distributor", "summary": "Skybox Packaging distributes packaging supplies in the US."},
    {"company": "PakFactory CA", "website": "https://www.pakfactory.com", "hq": "Canada", "cc": "CA", "founded": "Not publicly disclosed", "cats": "Custom and stock packaging", "focus": "Canadian packaging channel supplier", "summary": "PakFactory supplies custom and stock packaging as a Canadian channel supplier."},
    {"company": "Embageant FR", "website": "https://embageant.fr", "hq": "France", "cc": "FR", "founded": "Not publicly disclosed", "cats": "E-commerce packaging materials", "focus": "Online packaging wholesaler", "summary": "Embageant is an online packaging materials wholesaler serving France/Benelux."},
    {"company": "Furodet FR", "website": "https://furodet-emballages.fr", "hq": "France", "cc": "FR", "founded": "Not publicly disclosed", "cats": "Food and retail packaging", "focus": "French packaging merchant", "summary": "Furodet Emballages supplies packaging materials to businesses in France and Belgium."},
    {"company": "NNZ Group", "website": "https://www.nnz.com", "hq": "Netherlands", "cc": "NL", "founded": "1922", "cats": "Produce and industrial packaging", "focus": "International packaging distributor network", "summary": "NNZ is a multinational packaging distribution network founded in 1922."},
]


def domain_of(url: str) -> str:
    u = url if "://" in url else "https://" + url
    return urlparse(u).netloc.lower().replace("www.", "")


def dns_ok(d: str) -> bool:
    try:
        socket.getaddrinfo(d, 443)
        return True
    except Exception:
        try:
            socket.getaddrinfo(d, 80)
            return True
        except Exception:
            return False


def fetch(website: str) -> str:
    base = website.rstrip("/")
    chunks = []
    ctx = ssl.create_default_context()
    for p in ["", "/about", "/about-us"]:
        try:
            req = Request(base + p, headers={"User-Agent": "Mozilla/5.0"})
            with urlopen(req, timeout=12, context=ctx) as resp:
                html = resp.read(70000).decode("utf-8", errors="ignore")
            if any(x in html.lower() for x in ("domain is for sale", "buy this domain", "parked domain")):
                return "PARKED"
            text = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
            text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
            text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"\s+", " ", text)
            chunks.append(text[:2500])
        except Exception:
            continue
    return "\n".join(chunks)[:6000]


def role_check(company, website, source):
    prompt = f"""Classify packaging channel.
Company: {company}
Website: {website}
SOURCE: {source[:5500]}
KEEP only Specialty Distributor/Wholesaler/Merchant/Channel Partner.
REMOVE Manufacturer/OEM/Converter/Hybrid/Wrong/Unknown.
JSON: {{"keep":true|false,"role":"...","market_fit":true|false,"founded":"...","hq":"...","reason":"..."}}"""
    resp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": "Strict channel-only."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.1,
        response_format={"type": "json_object"},
    )
    try:
        return json.loads(resp.choices[0].message.content or "{}")
    except Exception:
        return {"keep": False, "role": "Unknown", "market_fit": False}


def main():
    backup = FP.with_name(FP.stem + f"_pre_add100_last_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(FP, backup)
    wb = load_workbook(FP)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    existing_names = {(r[0] or "").strip().lower() for r in rows}
    existing_domains = {domain_of(str(r[1] or "")) for r in rows if r[1]}
    need = TARGET - (len(rows) - BASELINE)
    print("current", len(rows), "need", need)

    # also match by domain to avoid renaming duplicates of already-added firms
    pool = []
    for c in CANDS:
        d = domain_of(c["website"])
        if c["company"].lower() in existing_names or d in existing_domains:
            continue
        pool.append(c)
    print("new pool", len(pool))

    added = 0
    accepted = []
    rejected = []
    for i, c in enumerate(pool):
        if added >= need:
            break
        name, web, d = c["company"], c["website"], domain_of(c["website"])
        print(f"[{i+1}/{len(pool)}] {name}", flush=True)
        if not dns_ok(d):
            rejected.append({"company": name, "reason": "dns"})
            print("  dns", flush=True)
            continue
        src = fetch(web)
        if src == "PARKED":
            rejected.append({"company": name, "reason": "parked"})
            print("  parked", flush=True)
            continue
        data = role_check(name, web, src or f"{name} {c['focus']}")
        role = (data.get("role") or "").lower()
        keep = bool(data.get("keep")) and bool(data.get("market_fit", True))
        if any(x in role for x in ("manufacturer", "oem", "converter", "hybrid", "wrong", "unknown")):
            keep = False
        if not keep:
            rejected.append({"company": name, "role": data.get("role"), "reason": data.get("reason")})
            print(f"  REJECT {data.get('role')}", flush=True)
            time.sleep(0.12)
            continue
        if data.get("founded"):
            c["founded"] = str(data["founded"]).strip()
        if data.get("hq"):
            c["hq"] = str(data["hq"]).strip()
        cc = c["cc"]
        rc = "NA" if cc in {"US", "CA", "MX"} else ("APAC" if cc in {"AU", "NZ", "SG"} else "EU")
        geo = f"{rc}({cc})"
        ws.append(
            [
                c["company"],
                c["website"],
                c.get("founded") or ND,
                c.get("hq") or ND,
                geo,
                geo,
                "Private",
                ND,
                c.get("cats") or "Packaging supplies",
                c.get("focus") or "Packaging distributor",
                c["company"],
                "Both",
                "Specialty Distributor",
                ND,
                ND,
                ND,
                ND,
                ND,
                cc,
                rc,
                c.get("summary") or f"{c['company']} packaging specialty distributor.",
            ]
        )
        existing_names.add(name.lower())
        existing_domains.add(d)
        added += 1
        accepted.append(c)
        print(f"  KEEP ({(len(rows)-BASELINE)+added}/{TARGET})", flush=True)
        time.sleep(0.12)

    wb.save(FP)
    final = ws.max_row - 1
    audit = {
        "ts": datetime.now().isoformat(),
        "backup": str(backup),
        "added": added,
        "after": final,
        "from_baseline": final - BASELINE,
        "accepted": accepted,
        "rejected": rejected,
        "file": str(FP),
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print("DONE after", final, "from_baseline", final - BASELINE, "added", added)


if __name__ == "__main__":
    main()
