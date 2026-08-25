# -*- coding: utf-8 -*-
"""Finish last-13 HQ fills: agent-verified + DeepSeek on hard cases only."""
from __future__ import annotations

import json
import os
import shutil
from datetime import datetime
from pathlib import Path

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
AUDIT = ROOT / "_wearable_hq_last13_audit.json"

VERIFIED = {
    "Primary Care Supplies": "Aylesford, United Kingdom",
    "BioD Medica": "Subang Jaya, Malaysia",
    "MyBanzu Healthcare": "Puchong, Malaysia",
    "AED USA": "Fort Worth, TX, United States",
    "Farmacias Carol": "Santo Domingo, Dominican Republic",
    "Farmacias La Paz": "Tegucigalpa, Honduras",
    "Farmacias Siman": "San Pedro Sula, Honduras",
}

SNIPPETS_FOR_DS = {
    "Al-Watania Medical": (
        "Sheet website https://www.alwatania.com. A similarly named Al Wattan Medical Group "
        "(alwattanmed.com) lists Riyadh, Saudi Arabia. Domains may differ. Only return an HQ "
        "if you are confident it is the same entity; otherwise empty hq."
    ),
    "ProCare Pharmacy": (
        "Website https://www.procarepharmacy.com often errors. Multiple unrelated ProCare Pharmacy "
        "entities exist (Orlando FL; Garden Grove CA; San Diego CA; Gainesville GA PBM). "
        "Do not guess. Return empty hq unless text clearly matches this exact website entity."
    ),
    "Tecnomed Brasil": (
        "Website https://www.tecnomed.com.br. Multiple Brazil entities use Tecnomed/Tecnomedi/Tecmedic "
        "(Volta Redonda RJ retail, Curitiba Tecmedic, Sao Paulo occupational health). "
        "Do not guess. Empty hq unless domain evidence is clear."
    ),
    "Provita Medical": (
        "Sheet website https://www.provita.com, country United States. German manufacturer "
        "provita.de imprint is Wermelskirchen, Germany. If this row is the German company, "
        "use Wermelskirchen, Germany. If it is a distinct US entity and city unknown, empty hq."
    ),
    "Veloce Solutions": (
        "https://www.veloce-solutions.com appears as a logistics/freight template site with no "
        "clear headquarters city on the homepage. Return empty hq."
    ),
    "EMS Supply Co.": (
        "https://www.emssupply.com shows USA/AU/NZ phone numbers only; no HQ city on index. "
        "Dixie EMS Brooklyn uses dixieems.com — different company. Return empty hq."
    ),
}


def ok_hq(hq: str) -> bool:
    s = (hq or "").strip()
    if "," not in s:
        return False
    parts = [p.strip() for p in s.split(",") if p.strip()]
    if len(parts) < 2:
        return False
    bad = {
        "united states",
        "usa",
        "united kingdom",
        "uk",
        "saudi arabia",
        "malaysia",
        "brazil",
        "honduras",
        "el salvador",
        "puerto rico",
        "europe",
        "belgium",
        "germany",
        "singapore",
    }
    if s.lower() == "singapore, singapore":
        return True
    return parts[0].lower() not in bad


def has_city(hq: str) -> bool:
    s = (hq or "").strip()
    if not s or "," not in s:
        return False
    low = s.lower()
    if low in {
        "singapore, singapore",
        "hong kong sar, china",
        "hong kong, china",
    }:
        return True
    parts = [p.strip() for p in s.split(",") if p.strip()]
    if len(parts) < 2:
        return False
    countries_only = {
        "united states",
        "usa",
        "united kingdom",
        "uk",
        "germany",
        "france",
        "japan",
        "china",
        "india",
        "brazil",
        "canada",
        "australia",
        "singapore",
        "malaysia",
        "south africa",
        "sweden",
        "spain",
        "italy",
        "netherlands",
        "belgium",
        "switzerland",
        "denmark",
        "finland",
        "norway",
        "mexico",
        "korea",
        "south korea",
        "uae",
        "united arab emirates",
        "hong kong",
        "saudi arabia",
        "europe",
        "iraq",
        "honduras",
        "el salvador",
        "bolivia",
        "puerto rico",
    }
    return parts[0].lower() not in countries_only


def deepseek(company: str, website: str, country: str, snippet: str) -> dict:
    key = (os.getenv("OPENAI_API_KEY") or os.getenv("DEEPSEEK_API_KEY") or "").strip()
    if not key:
        return {"hq": "", "confidence": "none", "reason": "no_api_key"}
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    system = (
        "Extract HQ city+country from evidence only. Never invent. "
        'JSON keys: hq, confidence (high|medium|low|none), reason. '
        "US format: City, ST, United States. Empty hq if unsure."
    )
    user = (
        f"Company: {company}\nWebsite: {website}\nSheet HQ: {country}\nEvidence:\n{snippet}"
    )
    try:
        with httpx.Client(timeout=60.0) as c:
            r = c.post(
                f"{base}/chat/completions",
                headers={
                    "Authorization": f"Bearer {key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": model,
                    "temperature": 0,
                    "messages": [
                        {"role": "system", "content": system},
                        {"role": "user", "content": user},
                    ],
                    "response_format": {"type": "json_object"},
                },
            )
            r.raise_for_status()
            return json.loads(r.json()["choices"][0]["message"]["content"])
    except Exception as e:
        return {"hq": "", "confidence": "none", "reason": str(e)}


def main() -> None:
    load_dotenv(r"D:\vendor-intel\.env")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = ROOT / f"global_wearable_medical_devices_market_global_FINAL_BEFORE_HQ_LAST13B_{stamp}.xlsx"
    shutil.copy2(XLSX, bak)

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers) if h}
    co_i, hq_i, web_i = ic["Company"], ic["Headquarters"], ic["Website"]

    updated: list[str] = []
    ds_log: list[dict] = []
    targets = list(dict.fromkeys([*VERIFIED.keys(), *SNIPPETS_FOR_DS.keys()]))

    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, co_i).value or "").strip()
        if co not in targets:
            continue
        old = str(ws.cell(r, hq_i).value or "").strip()
        web = str(ws.cell(r, web_i).value or "").strip()
        chosen = ""
        source = ""
        if co in VERIFIED:
            chosen = VERIFIED[co]
            source = "agent_web_verified"
        if co in SNIPPETS_FOR_DS:
            ds = deepseek(co, web, old, SNIPPETS_FOR_DS[co])
            ds_log.append({"company": co, "deepseek": ds})
            print(co, "deepseek", ds, flush=True)
            conf = str(ds.get("confidence") or "").lower()
            hq = str(ds.get("hq") or "").strip()
            # Accept high only for ambiguous hard cases (avoid inventing)
            if conf == "high" and ok_hq(hq):
                chosen = hq
                source = f"deepseek_high:{ds.get('reason')}"
            # Provita: allow medium if clearly Wermelskirchen Germany from imprint evidence
            if (
                co == "Provita Medical"
                and conf in {"high", "medium"}
                and "wermelskirchen" in hq.lower()
            ):
                chosen = "Wermelskirchen, Germany"
                source = f"deepseek:{conf}:provita.de imprint"
        if chosen and ok_hq(chosen) and chosen != old:
            ws.cell(r, hq_i).value = chosen
            updated.append(f"{co}: {old} -> {chosen} ({source})")
            print("UPDATED", updated[-1], flush=True)
        elif chosen and ok_hq(chosen):
            print("already", co, chosen, flush=True)

    wb.save(XLSX)

    city = 0
    still = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, co_i).value or "").strip()
        hq = str(ws.cell(r, hq_i).value or "").strip()
        if has_city(hq):
            city += 1
        else:
            still.append(
                {
                    "company": co,
                    "hq": hq,
                    "website": str(ws.cell(r, web_i).value or "").strip(),
                }
            )

    audit = {
        "backup": bak.name,
        "updated": updated,
        "deepseek_log": ds_log,
        "city_country_total": city,
        "country_only_remaining": len(still),
        "still_all": still,
        "rows": ws.max_row - 1,
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print("DONE city", city, "still", len(still))
    for s in still:
        print(" STILL", s["hq"], "|", s["company"], "|", s["website"])


if __name__ == "__main__":
    main()
