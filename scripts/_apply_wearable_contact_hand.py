# -*- coding: utf-8 -*-
"""Apply web-verified sales/procurement contacts for major wearable channel firms."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
    r"\global_wearable_medical_devices_market_global_FINAL.xlsx"
)
NPD = "Not publicly disclosed"

# name, role, linkedin — only web-verified this session
HAND = {
    "Cardinal Health": (
        "Shannon Kelso",
        "Senior Vice President, Acute Sales (US Medical Products and Distribution)",
        "https://www.linkedin.com/in/shannon-kelso-5616876",
    ),
    "Zuellig Pharma": (
        "Jan Rask Christensen",
        "VP Sales, Customer and Client Excellence",
        "https://www.linkedin.com/in/janraskchristensen",
    ),
    "Boots UK": (
        "Steve Ager",
        "VP, Commercial Director, Boots UK and ROI",
        NPD,  # appointment news; no verified personal LI in this pass
    ),
    "DKSH": (
        "Stefan P. Räbsamen",
        "Head of Procurement",
        NPD,  # may already have LI from Apollo
    ),
}


def empty(v: object) -> bool:
    return str(v or "").strip().lower() in {"", "n/a", "na", "-", "not publicly disclosed"}


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(f"{XLSX.stem}_BEFORE_CONTACT_HAND_{stamp}.xlsx")
    shutil.copy2(XLSX, bak)
    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers) if h}
    n = 0
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co not in HAND:
            continue
        name, role, li = HAND[co]
        cur = str(ws.cell(r, ic["Contact Person"]).value or "").strip()
        # fill if empty or upgrade with verified full name + LI
        if empty(cur) or (co in HAND and "linkedin.com/in/" in li.lower()):
            ws.cell(r, ic["Contact Person"]).value = name
            ws.cell(r, ic["Role"]).value = role
            if li != NPD and "linkedin.com/in/" in li.lower():
                ws.cell(r, ic["LinkedIn"]).value = li
            elif empty(ws.cell(r, ic["LinkedIn"]).value) or "/company/" in str(
                ws.cell(r, ic["LinkedIn"]).value or ""
            ).lower():
                if empty(ws.cell(r, ic["LinkedIn"]).value) or "/company/" in str(
                    ws.cell(r, ic["LinkedIn"]).value or ""
                ):
                    # keep existing person LI if present
                    existing_li = str(ws.cell(r, ic["LinkedIn"]).value or "")
                    if "linkedin.com/in/" not in existing_li.lower():
                        ws.cell(r, ic["LinkedIn"]).value = NPD
            n += 1
            print("filled", co, "=>", name, "|", role)
        elif empty(ws.cell(r, ic["LinkedIn"]).value) and "linkedin.com/in/" in li.lower():
            ws.cell(r, ic["LinkedIn"]).value = li
            n += 1
            print("LI only", co)
    wb.save(XLSX)
    print("backup", bak.name, "updated", n)


if __name__ == "__main__":
    main()
