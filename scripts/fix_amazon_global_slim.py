#!/usr/bin/env python3
"""Normalize Amazon row in smartwatch (and smartphone) SLIM to global Amazon."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")

GEO = "; ".join(
    [
        "United States",
        "Canada",
        "Mexico",
        "Brazil",
        "United Kingdom",
        "Germany",
        "France",
        "Italy",
        "Spain",
        "Netherlands",
        "Belgium",
        "Poland",
        "Sweden",
        "Turkey",
        "United Arab Emirates",
        "Saudi Arabia",
        "Egypt",
        "India",
        "Singapore",
        "Japan",
        "Australia",
        "China",
    ]
)

SMARTWATCH_UPDATES = {
    "Company": "Amazon",
    "Website": "https://www.amazon.com",
    "Founded": "1994",
    "Headquarters": "Seattle, Washington, United States",
    "Continent / Geography": GEO,
    "Ownership": "Public",
    "Employees": "1000001+",
    "Core Categories": "Smartwatches; Wearables; Fitness Trackers; Consumer Electronics; Accessories, others",
    "Specialty Focus": "Global e-commerce marketplace for smartwatches, wearables and consumer electronics, others",
    "Key Brands Represented": "Apple; Samsung; Garmin; Fitbit; Amazfit; Noise; boAt; Google; Huawei, others",
    "Retail / E-commerce / Both": "Ecommerce",
    "Distribution Type": "Online retail / marketplace",
    "Regions Served": "Global",
    "Regional Extensions": GEO,
    "Summary": (
        "Amazon (founded 1994, HQ Seattle, Washington, United States) is a global e-commerce "
        "marketplace and a major retail channel for smartwatches and wearables (Apple Watch, "
        "Samsung, Garmin, Fitbit and others) across dedicated country storefronts worldwide."
    ),
}


def header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [str(c.value or "").strip().lower() for c in ws[r] if c.value not in (None, "")]
        if "company" in vals:
            return r
    return 1


def main() -> None:
    f = ROOT / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL_SLIM.xlsx"
    wb = load_workbook(f)
    ws = wb.active
    hr = header_row(ws)
    headers = [str(c.value or "").strip() for c in ws[hr]]
    ci = headers.index("Company")
    target = None
    for r in range(hr + 1, ws.max_row + 1):
        name = str(ws.cell(r, ci + 1).value or "")
        if "amazon" in name.lower():
            target = r
            print(f"Smartwatch row {r}: {name}")
            break
    if not target:
        raise SystemExit("Amazon row not found in smartwatch SLIM")

    for col_name, val in SMARTWATCH_UPDATES.items():
        if col_name not in headers:
            print("skip missing", col_name)
            continue
        c = headers.index(col_name) + 1
        old = ws.cell(target, c).value
        ws.cell(target, c).value = val
        shown = val if len(str(val)) < 90 else str(val)[:87] + "..."
        print(f"  {col_name}: {old!r} -> {shown!r}")
    wb.save(f)
    print("Saved", f.name)

    # Align smartphone Amazon HQ/geo
    f2 = ROOT / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL_SLIM.xlsx"
    wb2 = load_workbook(f2)
    ws2 = wb2.active
    hr2 = header_row(ws2)
    headers2 = [str(c.value or "").strip() for c in ws2[hr2]]
    ci2 = headers2.index("Company")
    for r in range(hr2 + 1, ws2.max_row + 1):
        if str(ws2.cell(r, ci2 + 1).value or "").strip().lower() != "amazon":
            continue
        hi = headers2.index("Headquarters")
        gi = headers2.index("Continent / Geography")
        old_hq = ws2.cell(r, hi + 1).value
        old_geo = ws2.cell(r, gi + 1).value
        ws2.cell(r, hi + 1).value = "Seattle, Washington, United States"
        ws2.cell(r, gi + 1).value = GEO
        if "Founded" in headers2:
            ws2.cell(r, headers2.index("Founded") + 1).value = "1994"
        if "Website" in headers2:
            ws2.cell(r, headers2.index("Website") + 1).value = "https://www.amazon.com"
        print(f"Smartphone Amazon HQ: {old_hq!r} -> Seattle...")
        print(f"Smartphone Amazon Geo: {old_geo!r} -> {GEO[:80]}...")
        break
    wb2.save(f2)
    print("Saved", f2.name)
    print("Market relevance: YES — Amazon is a major smartwatch/wearables e-commerce retail channel.")


if __name__ == "__main__":
    main()
