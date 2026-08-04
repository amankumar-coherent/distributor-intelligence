"""Fix Core Categories, Specialty Focus, Price Segment, Key Brands Represented.

Replaces slogans/category dumps masquerading as brands, placeholder specialties
('Service provider'), and blank/FMCG-mismatched price segments.
"""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"C:\Users\aman.kumar\Desktop\coherent\output\run4")
MIRROR = Path(r"C:\Users\aman.kumar\Desktop\Cmi\agents\vendor-intel\output\run4")

MARKETS = {
    "global_smartwatch_market_global": "smartwatch",
    "global_green_chemical_market_global": "green_chemical",
    "global_firewall_market_global": "firewall",
}

# ── Curated Key Brands by market ─────────────────────────────────────────────
BRANDS_BY_MARKET: dict[str, dict[str, str]] = {
    "smartwatch": {
        "b&h photo video": "Apple; Samsung; Garmin; Fitbit; Amazfit; Google; Fossil",
        "target": "Apple; Fitbit; Samsung; Garmin; Amazfit; Google",
        "walmart": "Apple; Fitbit; Samsung; Garmin; Amazfit; Google",
        "jcpenney": "Apple; Fitbit; Samsung; Garmin",
        "nordstrom rack": "Apple; Garmin; Fitbit; Samsung",
        "macys": "Apple; Fitbit; Samsung; Fossil",
        "zales": "Fossil; Michael Kors; Apple; Citizen",
        "newegg commerce": "Apple; Samsung; Garmin; Amazfit; Fitbit; Google",
        "digikey": "Multi-vendor electronic components (wearables OEM supply)",
        "mouser electronics": "Multi-vendor semiconductors & electronic components",
        "newark": "Multi-vendor electronic components (Avnet)",
        "farnell": "Multi-vendor electronic components (Avnet)",
        "rs components": "Multi-vendor electronic & industrial components",
        "ingram micro": "Apple; Samsung; Garmin; Fitbit; Amazfit; Google",
        "td synnex": "Apple; Samsung; Garmin; Fitbit; Amazfit; Google",
        "synnex": "Apple; Samsung; Garmin; Fitbit; Amazfit",
        "tech data": "Apple; Samsung; Garmin; Fitbit; Amazfit",
        "d&h distributing": "Apple; Samsung; Garmin; Fitbit; Amazfit; Google",
        "esprinet": "Apple; Samsung; Garmin; Fitbit; Amazfit; Xiaomi",
        "amtel group": "Apple; Samsung; Sony; Honor; Redmi; Tecno; Nintendo; Xbox",
        "savex computers": "Apple; Samsung; Garmin; Fitbit; Amazfit",
        "nalanda enterprises": "Samsung; Noise; Fire-Boltt; boat; Amazfit",
        "dynamiccom uae": "Apple; Samsung; Garmin; Huawei; Amazfit",
        "mohamed sharif group": "Apple; Samsung; Garmin; Huawei; Amazfit",
        "vianet": "Apple; Samsung; Xiaomi; Amazfit; Garmin",
        "techvision global": "Apple; Samsung; Garmin; Amazfit; Fitbit",
        "chinavasion": "OEM / private-label smartwatches & wearables",
        "argentek": "Multi-brand consumer electronics & wearables",
        "artec lantec": "Electronic components for OEMs",
        "nippon audiotronix": "Noise; Fire-Boltt; boat; Amazfit; Samsung",
        "npuone": "Apple; Samsung; Garmin; Amazfit",
        "eumobilewholesalers": "Samsung; Apple; Xiaomi; Amazfit; Huawei",
        "usmphones": "Apple; Samsung; Xiaomi; Amazfit",
        "wildberries": "Apple; Samsung; Xiaomi; Amazfit; Huawei; Garmin",
        "jomashop": "Apple; Fossil; Michael Kors; Citizen; Tissot",
        "cover company": "Xiaomi; Amazfit; Apple; Samsung",
        "a&s creations": "Celestron; Garmin; Zeiss; Bushnell",
        "wearabledevices": "Mudra Band; Mudra Link (own brand)",
        "gps india networks": "GPS / connected wearables & tracking devices",
        "hongkong bishop technology": "OEM smartwatches & wearable modules",
        "firebird distribution fze": "Multi-brand consumer electronics & wearables",
        "itechwearables": "OEM / private-label wearables",
        "dealer – ielov": "Parental-control & GPS smartwatches",
        "pointmobile": "Point Mobile enterprise handhelds & wearables",
        "globics technology": "Multi-brand consumer electronics",
        "motionview": "OEM fitness trackers & smartwatches",
    },
    "green_chemical": {
        "azelis": "Multi-principal specialty chemicals (life sciences & industrial)",
        "barentz": "Multi-principal life sciences & specialty chemicals",
        "brenntag": "Multi-principal industrial & specialty chemicals",
        "imcd group": "Multi-principal specialty chemicals & ingredients",
        "imcd": "Multi-principal specialty chemicals & ingredients",
        "dksh performance materials": "Multi-principal performance & specialty materials",
        "caldic": "Multi-principal specialty chemicals",
        "safic-alcan": "Multi-principal specialty chemicals (rubber, coatings, life sciences)",
        "biesterfeld": "Multi-principal plastics, rubber & specialty chemicals",
        "nagase & co": "Multi-principal specialty chemicals & materials",
        "behn meyer group": "Multi-principal specialty chemicals (ASEAN)",
        "grolman group": "Multi-principal specialty chemicals & additives",
        "nordmann": "Multi-principal specialty chemicals & ingredients",
        "maroon group": "Multi-principal specialty chemicals (Barentz)",
        "arpadis group": "Multi-principal industrial & specialty chemicals",
        "quimdis": "Multi-principal natural ingredients & specialty chemicals (Azelis)",
        "quimidroga": "Multi-principal industrial & specialty chemicals",
        "eigenmann & veronelli": "Multi-principal specialty chemicals",
        "hsh chemie": "Multi-principal specialty chemicals (CEE)",
        "lehmann & voss": "Multi-principal specialty chemicals & additives",
        "velox gmbh": "Multi-principal specialty chemicals (IMCD)",
        "vidara": "Multi-principal specialty chemicals & water treatment",
        "petrochem middle east": "Multi-principal petrochemicals & specialty chemicals",
        "reda chemicals": "Multi-principal specialty chemicals (MENA)",
        "redox": "Multi-principal industrial & specialty chemicals",
        "manuchar": "Multi-principal chemicals & ingredients",
        "solvadis": "Multi-principal specialty chemicals",
        "solventis": "Multi-principal solvents & specialty chemicals",
        "gamma chimica": "Multi-principal specialty & green chemicals",
        "airedale chemical": "Multi-principal industrial chemicals & specialties",
        "cg chemikalien": "Multi-principal specialty chemicals",
        "bodo moller chemie": "Multi-principal specialty chemicals & adhesives",
        "connect chemicals": "Multi-principal specialty chemicals",
        "chemstationasia": "Multi-principal specialty & bulk chemicals",
        "champion specialty chemicals": "Specialty chemicals for coatings & environmental",
        "pluschem": "Fine & specialty chemicals; cosmetic & agro ingredients",
        "jebsen & jessen ingredients": "Multi-principal specialty ingredients",
        "oqema group": "Multi-principal specialty chemicals",
        "ghcl trading division": "Soda ash, salt & chemical commodities trading",
        "greenchem": "Bio-based / renewable specialty chemicals",
        "greenchem industries": "Green & specialty chemical distribution",
        "faravelli group": "Multi-principal specialty ingredients",
        "kolmar group": "Multi-principal commodities & chemicals trading",
        "krahn chemie": "Multi-principal specialty chemicals",
        "ricardo molina": "Multi-principal specialty chemicals",
        "connell": "Multi-principal specialty chemicals",
        "comercial quimica masso": "Multi-principal specialty chemicals",
        "ecem": "Multi-principal specialty chemicals",
        "novasol chemicals": "Multi-principal specialty chemicals",
        "protea chemicals": "Multi-principal specialty chemicals",
        "ravago chemicals": "Multi-principal plastics & specialty chemicals",
        "green chemicals": "Green / bio-based chemical distribution",
        "koda distribution group": "Multi-principal specialty chemicals",
        "tilley distribution": "Multi-principal specialty chemicals",
        "green chemical store": "Green / specialty chemical retail",
    },
    "firewall": {
        "acceron": "Ivanti; Seqrite; F-Secure; Acunetix; EnGenius; Nexus",
        "ingram micro": "Palo Alto Networks; Fortinet; Cisco; Check Point; CrowdStrike; Microsoft",
        "td synnex": "Palo Alto Networks; Fortinet; Cisco; Check Point; CrowdStrike; Microsoft",
        "westcon-comstor": "Check Point; Fortinet; Palo Alto Networks; Cisco; CrowdStrike",
        "softcat": "Microsoft; Cisco; Palo Alto Networks; Fortinet; CrowdStrike; Check Point",
        "softchoice": "Microsoft; Cisco; Palo Alto Networks; Fortinet; CrowdStrike",
        "cdw government": "Cisco; Palo Alto Networks; Fortinet; Check Point; Microsoft; CrowdStrike",
        "connection": "Cisco; Fortinet; Palo Alto Networks; Check Point; Microsoft",
        "scansource": "Cisco; Fortinet; WatchGuard; SonicWall; physical security brands",
        "d&h distributing": "Cisco; Fortinet; SonicWall; WatchGuard; Microsoft",
        "exclusive networks": "Palo Alto Networks; Fortinet; Check Point; CrowdStrike; Zscaler",
        "arrow electronics": "Cisco; Palo Alto Networks; Fortinet; Check Point",
        "also holding": "Microsoft; Cisco; Fortinet; Palo Alto Networks; Check Point",
        "bluestar": "Cisco; Fortinet; Palo Alto Networks; Check Point; Microsoft",
        "computacenter": "Cisco; Palo Alto Networks; Fortinet; Check Point; Microsoft",
        "presidio": "Cisco; Palo Alto Networks; Fortinet; CrowdStrike; Microsoft",
        "world wide technology": "Cisco; Palo Alto Networks; Fortinet; CrowdStrike; Microsoft",
        "eplus technology": "Cisco; Palo Alto Networks; Fortinet; Check Point; Microsoft",
        "insight enterprises": "Microsoft; Cisco; Palo Alto Networks; Fortinet; CrowdStrike",
        "climb channel solutions": "Palo Alto Networks; Fortinet; Check Point; cybersecurity ISVs",
        "infinigate nordic": "Check Point; Fortinet; Palo Alto Networks; cybersecurity vendors",
        "hermitage solutions": "Palo Alto Networks; Fortinet; Check Point; cybersecurity vendors",
        "ignition technology": "Palo Alto Networks; Fortinet; Check Point; cybersecurity vendors",
        "starlink": "Palo Alto Networks; Fortinet; Check Point; cybersecurity vendors",
        "tarsus distribution": "Fortinet; Check Point; cybersecurity vendors",
        "ivalue infosolutions": "Palo Alto Networks; Fortinet; Check Point; cybersecurity vendors",
        "rahi systems": "Cisco; Arista; Dell; HPE; networking & data-center brands",
        "comguard": "McAfee; WatchGuard; Barracuda; cybersecurity vendors",
        "acpl systems": "Multi-vendor cybersecurity solutions",
        "emt distribution": "Cybersecurity & networking vendor portfolio",
        "mindware": "Microsoft; Cisco; Fortinet; Palo Alto Networks; Check Point",
        "redington": "Microsoft; Cisco; Fortinet; Palo Alto Networks; Check Point",
        "zones": "Microsoft; Cisco; Palo Alto Networks; Fortinet; CrowdStrike",
        "westcoast limited": "Microsoft; Cisco; Fortinet; Palo Alto Networks",
        "dicker data": "Microsoft; Cisco; Fortinet; Palo Alto Networks; Check Point",
        "nuvias uc": "Collaboration & networking vendor portfolio",
        "appriver": "Microsoft 365; email & cloud security vendors",
        "sherweb": "Microsoft; cloud security & MSP vendors",
        "core bts": "Cisco; Microsoft; cybersecurity & infrastructure vendors",
        "solutions2go": "Cybersecurity & networking vendor portfolio",
        "winslow technology group": "Cybersecurity & networking vendor portfolio",
        "ma labs": "Component & networking hardware brands",
        "texonic instruments": "Industrial electronics & test instrumentation brands",
        "newark": "Multi-vendor electronic components (Avnet)",
        "bulwark technologies": "Cybersecurity vendor portfolio (MENA)",
        "cyberknight": "Cybersecurity vendor portfolio (MENA)",
        "elcore group": "Cybersecurity & IT vendor portfolio",
        "shi international": "Microsoft; Cisco; Palo Alto Networks; Fortinet; Check Point; CrowdStrike",
    },
}

# Market default core categories
CORE_DEFAULT = {
    "smartwatch": "smartwatches, fitness trackers, wearable technology, accessories",
    "green_chemical": "specialty chemicals, bio-based chemicals, green solvents, industrial chemicals",
    "firewall": "firewalls, network security, cybersecurity software, threat protection",
}

PRICE_VALID = {
    "smartwatch": {"Premium", "Mid-Market", "Value / Mass", "Multi-tier", "Enterprise / B2B"},
    "green_chemical": {
        "Specialty / Premium",
        "Commodity / Bulk",
        "Premium + Bulk",
        "Mid-Market",
    },
    "firewall": {"Enterprise", "Mid-Market / SMB", "Multi-tier", "Managed / Subscription"},
}


def key(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "").strip().lower())


def is_junk_specialty(val: str) -> bool:
    s = (val or "").strip()
    if not s:
        return True
    low = s.lower()
    if low in ("service provider", "industry manufacturer", "yes", "no", "n/a"):
        return True
    if low.startswith("organic/premium") and len(s) < 40:
        return True
    # LinkedIn industry salad / wrong vertical / nav labels
    if s.count(";") >= 4 and len(s) > 100:
        return True
    if any(
        x in low
        for x in (
            "civil construction",
            "roading",
            "land development",
            "byod security",
            "mobile expense management",
            "delivering technolog",
            "pro audio",
            "educational sales",
            "information management solutions",
            "technology consulting services",
            "organic/premium",
        )
    ):
        return True
    return False


def is_junk_brands(val: str, *, company: str = "") -> bool:
    s = (val or "").strip()
    if not s:
        return True
    low = s.lower()
    if any(
        x in low
        for x in (
            "professional source",
            "private investors",
            "join resources",
            "headquartered in",
            "systems for israel",
            "like something",
            "lorem",
            "utility software",
            "cutting-edge cybersecurity",
            "it hardware",
            "software licensing",
            "eco-friendly products",
        )
    ):
        return True
    # Category dump (no proper brand names)
    category_tokens = (
        "smartwatches",
        "fitness trackers",
        "wearable technology",
        "wearable devices",
        "mobile devices",
        "mobile accessories",
        "consumer electronics",
        "specialty chemicals",
        "bio-based chemicals",
        "bio-based solvents",
        "renewable chemicals",
        "green chemicals",
        "industrial chemicals",
        "firewalls",
        "networking equipment",
        "cybersecurity",
        "network security",
        "software solutions",
        "cloud solutions",
        "electronic components",
        "semiconductors",
        "solvents",
        "surfactants",
        "acids",
        "bases",
        "it hardware",
    )
    parts = [p.strip() for p in re.split(r"[;,]", s) if p.strip()]
    if not parts:
        return True
    # Keep honest multi-principal / multi-vendor labels
    if low.startswith("multi-"):
        return False
    cat_hits = sum(1 for p in parts if any(t in p.lower() for t in category_tokens))
    if cat_hits >= max(2, len(parts) // 2):
        return True
    if len(parts) == 1 and len(parts[0].split()) > 8 and not re.search(r"[A-Z][a-z]+ [A-Z]", parts[0]):
        return True
    if company and low == company.lower():
        return True
    return False


def is_bad_price(val: str, market: str) -> bool:
    s = (val or "").strip()
    if not s:
        return True
    if market in ("smartwatch", "firewall") and any(
        x in s for x in ("Organic", "Commodity / Bulk", "Premium + Bulk")
    ):
        return True
    if market == "green_chemical" and s in ("Premium / Organic",):
        return True
    if market == "smartwatch" and s not in PRICE_VALID["smartwatch"]:
        return True
    if market == "firewall" and s not in PRICE_VALID["firewall"]:
        return True
    return False


def infer_specialty(market: str, business_type: str, company: str, summary: str) -> str:
    bt = (business_type or "").lower()
    text = f"{company} {business_type} {summary}".lower()
    if market == "smartwatch":
        if any(x in bt for x in ("retail", "e-commerce", "ecommerce")) or any(
            x in text
            for x in (
                "walmart",
                "target",
                "macy",
                "jcpenney",
                "nordstrom",
                "newegg",
                "jomashop",
                "b&h",
                "zales",
            )
        ):
            return "Consumer wearables & electronics retail"
        if "distribut" in bt or "wholesale" in bt:
            return "Wearables & consumer electronics distribution"
        if "component" in text or any(x in text for x in ("digikey", "mouser", "farnell", "newark", "rs component")):
            return "Electronic components for wearables OEMs"
        if "manufactur" in bt or "oem" in text:
            return "Wearables OEM / brand manufacturing"
        return "Smartwatch & wearable technology channel"
    if market == "green_chemical":
        if "distribut" in bt or "trad" in bt:
            if any(x in text for x in ("personal care", "cosmetic", "life science", "food", "pharma")):
                return "Specialty ingredients distribution (life sciences)"
            if any(x in text for x in ("solvent", "bulk", "commodity")):
                return "Industrial & bulk chemical distribution"
            return "Specialty chemical distribution"
        if "manufactur" in bt:
            return "Green / specialty chemical manufacturing"
        return "Green & specialty chemicals"
    # firewall
    if "distribut" in bt or "vad" in text or "value-added" in text:
        return "Value-added cybersecurity distribution"
    if any(x in bt for x in ("retail", "reseller", "partner")):
        return "Cybersecurity reseller / solutions partner"
    if "msp" in text or "managed" in bt:
        return "Managed security & cloud services"
    if "integrat" in bt:
        return "Cybersecurity systems integration"
    return "Network security & firewall channel"


def infer_price(market: str, business_type: str, company: str, summary: str) -> str:
    bt = (business_type or "").lower()
    text = f"{company} {summary}".lower()
    if market == "smartwatch":
        if any(
            x in text
            for x in (
                "b&h",
                "walmart",
                "target",
                "macy",
                "jcpenney",
                "nordstrom",
                "newegg",
                "jomashop",
                "zales",
            )
        ):
            if any(x in text for x in ("jomashop", "zales", "nordstrom", "luxury")):
                return "Premium"
            if any(x in text for x in ("walmart", "target")):
                return "Value / Mass"
            return "Mid-Market"
        if any(x in text for x in ("luxury", "premium")) and "organic" not in text:
            return "Premium"
        if any(x in text for x in ("digikey", "mouser", "farnell", "newark", "rs components")):
            return "Enterprise / B2B"
        if "distribut" in bt or "wholesale" in bt:
            return "Multi-tier"
        return "Mid-Market"
    if market == "green_chemical":
        if any(x in text for x in ("bio-based", "green", "sustainable", "specialty ingredient")):
            if "bulk" in text or "commodity" in text:
                return "Premium + Bulk"
            return "Specialty / Premium"
        if any(x in text for x in ("bulk", "commodity", "solvent", "trading")):
            return "Commodity / Bulk"
        if "distribut" in bt:
            return "Premium + Bulk"
        return "Mid-Market"
    # firewall
    if any(x in text for x in ("government", "enterprise", "presidio", "computacenter", "wwt", "cdw")):
        return "Enterprise"
    if any(x in text for x in ("msp", "smb", "managed", "sherweb", "appriver")):
        return "Managed / Subscription"
    if "distribut" in bt:
        return "Multi-tier"
    return "Mid-Market / SMB"


def infer_core(market: str, existing: str) -> str:
    s = (existing or "").strip()
    if s and len(s) >= 8 and not re.search(r"service provider|lorem|placeholder", s, re.I):
        # Fix obvious wrong cores later if needed
        return s
    return CORE_DEFAULT[market]


def infer_brands(market: str, company: str, existing: str) -> str:
    k = key(company)
    brands_map = BRANDS_BY_MARKET.get(market) or {}
    if k in brands_map:
        return brands_map[k]
    for bk, bv in brands_map.items():
        if bk in k or k in bk:
            return bv
    if existing and not is_junk_brands(existing, company=company):
        return existing
    if market == "smartwatch":
        return "Multi-brand smartwatches & wearables (Apple, Samsung, Garmin, Amazfit & others)"
    if market == "green_chemical":
        return "Multi-principal specialty & green chemical suppliers"
    return "Multi-vendor cybersecurity (Palo Alto, Fortinet, Cisco, Check Point & others)"


def fix_row_fields(
    *,
    market: str,
    company: str,
    business_type: str,
    summary: str,
    core: str,
    specialty: str,
    price: str,
    brands: str,
) -> dict[str, str]:
    new_core = infer_core(market, core)
    new_spec = specialty
    if is_junk_specialty(specialty):
        new_spec = infer_specialty(market, business_type, company, summary)
    new_price = price
    if is_bad_price(price, market):
        new_price = infer_price(market, business_type, company, summary)
    brands_map = BRANDS_BY_MARKET.get(market) or {}
    new_brands = brands
    if key(company) in brands_map or is_junk_brands(brands, company=company) or not (brands or "").strip():
        new_brands = infer_brands(market, company, brands)

    # Re-score known retail price tiers that were mis-tagged Enterprise/B2B
    co_l = company.lower()
    if market == "smartwatch" and new_price == "Enterprise / B2B":
        if any(
            x in co_l
            for x in (
                "b&h",
                "walmart",
                "target",
                "macy",
                "jcpenney",
                "nordstrom",
                "newegg",
                "jomashop",
                "zales",
                "wildberries",
            )
        ):
            new_price = infer_price(market, business_type, company, summary)

    return {
        "Core Categories": new_core,
        "Specialty Focus": new_spec,
        "Price Segment": new_price,
        "Key Brands Represented": new_brands,
    }


def find_header(ws):
    for r in range(1, 8):
        vals = [c.value for c in ws[r]]
        if any(v and str(v).strip().lower() == "company" for v in vals):
            return r, vals
    return 1, [c.value for c in ws[1]]


def col_index(headers, *names):
    low = {str(h).strip().lower(): i for i, h in enumerate(headers) if h}
    for n in names:
        if n.lower() in low:
            return low[n.lower()]
    return None


def apply_excel(path: Path, market: str) -> dict:
    if not path.exists():
        return {"file": str(path), "missing": True}
    wb = load_workbook(path)
    ws = wb.active
    hr, headers = find_header(ws)
    ci = col_index(headers, "Company")
    bti = col_index(headers, "Business Type")
    sumi = col_index(headers, "Summary", "Company Summary", "Role Description")
    core_i = col_index(headers, "Core Categories")
    spec_i = col_index(headers, "Specialty Focus")
    price_i = col_index(headers, "Price Segment")
    brand_i = col_index(headers, "Key Brands Represented")
    if ci is None or core_i is None:
        return {"file": str(path), "error": "missing columns"}
    stats = {"core": 0, "specialty": 0, "price": 0, "brands": 0, "rows": 0}
    for ridx in range(hr + 1, ws.max_row + 1):
        company = ws.cell(ridx, ci + 1).value
        if not company:
            continue
        stats["rows"] += 1
        bt = str(ws.cell(ridx, bti + 1).value or "") if bti is not None else ""
        summary = str(ws.cell(ridx, sumi + 1).value or "") if sumi is not None else ""
        core = str(ws.cell(ridx, core_i + 1).value or "")
        spec = str(ws.cell(ridx, spec_i + 1).value or "") if spec_i is not None else ""
        price = str(ws.cell(ridx, price_i + 1).value or "") if price_i is not None else ""
        brands = str(ws.cell(ridx, brand_i + 1).value or "") if brand_i is not None else ""
        fixed = fix_row_fields(
            market=market,
            company=str(company),
            business_type=bt,
            summary=summary,
            core=core,
            specialty=spec,
            price=price,
            brands=brands,
        )
        if fixed["Core Categories"] != core.strip():
            ws.cell(ridx, core_i + 1).value = fixed["Core Categories"]
            stats["core"] += 1
        if spec_i is not None and fixed["Specialty Focus"] != spec.strip():
            ws.cell(ridx, spec_i + 1).value = fixed["Specialty Focus"]
            stats["specialty"] += 1
        if price_i is not None and fixed["Price Segment"] != price.strip():
            ws.cell(ridx, price_i + 1).value = fixed["Price Segment"]
            stats["price"] += 1
        if brand_i is not None and fixed["Key Brands Represented"] != brands.strip():
            ws.cell(ridx, brand_i + 1).value = fixed["Key Brands Represented"]
            stats["brands"] += 1
    wb.save(path)
    return {"file": str(path), **stats}


def apply_json(path: Path, market: str) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = data.get("relevant_companies") or []
    stats = {"core": 0, "specialty": 0, "price": 0, "brands": 0, "rows": 0}
    for r in rows:
        stats["rows"] += 1
        company = str(r.get("company") or r.get("brand") or "")
        bt = str(r.get("role") or r.get("channel_type") or r.get("company_function") or "")
        summary = str(r.get("summary") or r.get("company_summary") or r.get("polished_summary") or "")
        crawl = dict(r.get("_crawl") or {})
        # Prefer previously exported presentation enrich if present
        enrich = r.get("_presentation_enrich") if isinstance(r.get("_presentation_enrich"), dict) else {}
        core = str(enrich.get("core_categories") or crawl.get("core_categories") or "")
        spec = str(enrich.get("specialty_focus") or crawl.get("specialty_focus") or "")
        price = str(enrich.get("price_segment") or crawl.get("price_segment") or "")
        brands = str(crawl.get("key_brands_represented") or enrich.get("key_brands") or "")
        # If crawl specialty empty, export may have used industries — treat as current specialty candidate
        if not spec:
            spec = str(crawl.get("industries") or "")
        fixed = fix_row_fields(
            market=market,
            company=company,
            business_type=bt,
            summary=summary,
            core=core,
            specialty=spec,
            price=price,
            brands=brands,
        )
        if fixed["Core Categories"] != core.strip():
            stats["core"] += 1
        if fixed["Specialty Focus"] != spec.strip():
            stats["specialty"] += 1
        if fixed["Price Segment"] != price.strip():
            stats["price"] += 1
        if fixed["Key Brands Represented"] != brands.strip():
            stats["brands"] += 1
        crawl["core_categories"] = fixed["Core Categories"]
        crawl["specialty_focus"] = fixed["Specialty Focus"]
        crawl["price_segment"] = fixed["Price Segment"]
        crawl["key_brands_represented"] = fixed["Key Brands Represented"]
        crawl["key_brands_source"] = "web_manual_or_inferred"
        r["_crawl"] = crawl
        # keep presentation enrich in sync for re-export
        pe = dict(enrich)
        pe["core_categories"] = fixed["Core Categories"]
        pe["specialty_focus"] = fixed["Specialty Focus"]
        pe["price_segment"] = fixed["Price Segment"]
        pe["key_brands"] = fixed["Key Brands Represented"]
        r["_presentation_enrich"] = pe
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return {"file": str(path), **stats}


def mirror_dir(src: Path, dst: Path) -> None:
    dst.mkdir(parents=True, exist_ok=True)
    for pat in [f"{src.name}.json", "*_FINAL.xlsx", "*_founded_and_hq_v2.xlsx", "*_founded_and_hq_v2_cleaned.xlsx"]:
        for p in src.glob(pat):
            try:
                shutil.copy2(p, dst / p.name)
            except PermissionError:
                print(f"  skip locked: {p.name}")


def main() -> None:
    for mid, market in MARKETS.items():
        md = ROOT / mid
        short = re.sub(r"_global$", "", mid)
        print(f"\n=== {mid} ({market}) ===")
        print(" JSON", apply_json(md / f"{mid}.json", market))
        for suffix in ["_FINAL.xlsx", "_founded_and_hq_v2.xlsx", "_founded_and_hq_v2_cleaned.xlsx"]:
            print(" XLSX", apply_excel(md / f"{short}{suffix}", market))
        mirror_dir(md, MIRROR / mid)


if __name__ == "__main__":
    main()
