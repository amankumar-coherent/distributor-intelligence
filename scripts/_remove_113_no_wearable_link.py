#!/usr/bin/env python3
"""One-off removal of 113 companies confirmed to have no wearable-device
market link (Category A: legit distributor but no wearable evidence;
Category B: wrong industry entirely; Category C: AED cluster).

Verifies each company name against its row before deleting, then deletes
bottom-to-top so row numbers stay valid during the loop.
"""
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/global_wearable_medical_devices_market_global_FINAL.xlsx"
SHEET = "Wearable Medical Companies"

TO_REMOVE = [
    "AAH Pharmaceuticals", "Africa Health Supplies", "Al Noor Medical", "Al-Ahli Medical",
    "Al-Bahar Medical", "Al-Etihad Medical Supplies", "Al-Futtaim Health", "Al-Hayat Medical Company",
    "Al-Mansour Medical", "Al-Moammar Information Systems (MIS)", "Al-Salam Medical", "Alliance Healthcare",
    "Almedio Inc.", "Anda Pharmaceuticals", "Arjo", "As One Corporation", "Asahi Kasei Medical Asia Pacific",
    "Australian Medical Supplies", "B. Braun", "Bidvest Medical", "BML Inc.", "Boryung",
    "Brasil Medical Distribuidora", "Bunzl", "Canon Medical Systems", "Carex Health Brands", "CERP Rouen",
    "Cipla Medpro South Africa", "Comen Medical Asia Pacific", "CuraScript SD", "Daehwa Pharmaceutical",
    "Daiichi Sankyo Healthcare", "Darby Group Companies", "DocMorris", "Durbin PLC", "Farmacias del Dr. Simi",
    "Gehe Pharma Handel", "Grifols", "Gulf Drug LLC", "Gulf Medical Company", "HD Supply",
    "Health Products For You", "Healthcare Direct", "Hefame", "Itochu Corporation", "Kohlpharma",
    "Kojima Co., Ltd.", "Laboratoires Gilbert", "Lincare Holdings", "Medhold", "Medi-Dyne Healthcare Products",
    "MediLink", "Mediq", "Medisource Supply (Asia)", "MediSupplies Ltd", "Medline Industries", "Omnicell",
    "Owens & Minor", "Patterson Companies", "Phoenix Group", "Practo", "Rakuten Medical",
    "Toho Holdings Co., Ltd.", "Zuellig Pharma", "Buffalo Hospital Supply", "Cirúrgica Fernandes",
    "Claflin Company", "Clinisupplies", "Dealmed", "Defries Industries", "Device Technologies", "Pocket Nurse",
    "Sarnova", "School Health Corporation", "Symbiomed", "Tri-anim Health Services", "WorldPoint",
    "Zahrawi Group", "Brocacef", "Concordance Healthcare Solutions", "Drogueria del Sur", "Nadro", "Oriola",
    "Shanghai Pharmaceuticals Distribution", "Tamro", "Welcia Holdings", "Cosmos Pharmaceutical Retail",
    "Buffalo Supply", "Edan Instruments Asia Pacific",
    # Category B -- wrong industry entirely
    "Anixter", "Arrow Electronics", "Berner Group", "Fastenal", "Fisher Scientific", "Grainger",
    "Hakuto Co., Ltd.", "Mouser Electronics", "Movianto", "MSC Industrial Supply", "Murakami Corporation",
    "SHI International Corp.", "VWR (Avantor)", "Wesco International", "Zones, LLC", "Zoro",
    "Shoprite Holdings", "Alloga", "Imperial Health Sciences",
    # Category C -- AED cluster
    "AED USA", "AED Brands", "Response Ready", "Enerspect Medical Solutions", "HSI Emergency Care",
]


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}
    comp_col = idx["Company"]

    remaining = set(TO_REMOVE)
    rows_to_delete = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=comp_col).value
        if name in remaining:
            rows_to_delete.append((r, name))
            remaining.discard(name)

    print(f"Matched {len(rows_to_delete)} of {len(TO_REMOVE)} names in the sheet.")
    if remaining:
        print("NOT FOUND (skipped):")
        for n in sorted(remaining):
            print(" -", n)

    before = ws.max_row
    for r, name in sorted(rows_to_delete, key=lambda t: t[0], reverse=True):
        ws.delete_rows(r, 1)

    wb.save(XLSX)
    print(f"Row count before: {before}, after: {ws.max_row}, deleted: {len(rows_to_delete)}")


if __name__ == "__main__":
    main()
