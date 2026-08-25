"""Final push: add remaining distributors to reach exactly +100 from baseline 224."""
from __future__ import annotations

import json
import os
import re
import shutil
import socket
import ssl
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook
from openpyxl.styles import Alignment

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")
FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "add100_final_push_audit.json"

API_KEY = (os.getenv("OPENAI_API_KEY") or "").strip()
BASE_URL = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
client = OpenAI(api_key=API_KEY, base_url=BASE_URL)
ND = "Not publicly disclosed"
BASELINE = 224
TARGET = 100

# Hand-picked, previously web-verified channel companies likely not in sheet
SEEDS = [
    ("Citation Box & Paper", "https://www.citationpackaging.com", "Chicago, IL, USA", "US", "Stock boxes; packaging materials", "Chicago packaging distributor"),
    ("OSI Packaging", "https://www.osipackaging.com", "Shreveport, LA, USA", "US", "Packaging materials; equipment", "Louisiana stocking packaging distributor"),
    ("Stickel Packaging Supply", "https://www.stickelpackaging.com", "Lakewood, NJ, USA", "US", "Packaging; industrial; foodservice", "NJ packaging distributor"),
    ("Elkay Plastics", "https://www.elkayplastics.com", "Commerce, CA, USA", "US", "Poly bags; film packaging", "National poly packaging distributor"),
    ("SSI Packaging Group", "https://ssipkg.com", "Richmond, VA, USA", "US", "Shipping supplies; corrugated; tape", "Mid-Atlantic packaging distributor"),
    ("Weller Packaging", "https://www.wellerpackaging.co.uk", "UK", "GB", "Flexible polythene; paper packaging wholesale", "UK merchant-trade packaging wholesaler"),
    ("Maidstone Packaging Supplies", "https://www.maidstonepackaging.co.uk", "Kent, UK", "GB", "Cartons; bubble; pallet wrap; tapes", "SE England packaging distributor"),
    ("QPack", "https://qpack.co.uk", "UK", "GB", "Retail packaging; carrier bags; tissue", "UK retail packaging supplier"),
    ("Global Trading UK", "https://gtuk.co.uk", "Preston, UK", "GB", "Mailing bags; carrier bags; stockholding", "UK packaging procurement/stockholding distributor"),
    ("Baas Verpakkingen", "https://baasverpakkingen.nl", "Netherlands", "NL", "Boxes; films; bags; tapes", "Dutch packaging wholesaler"),
    ("Knueppel Verpackung", "https://knueppel.de", "Germany", "DE", "Industrial packaging wholesale", "German B2B packaging wholesaler"),
    ("Packen24", "https://www.packen24.de", "Germany", "DE", "Online packaging materials", "German packaging e-merchant"),
    ("Serwo", "https://www.serwo.de", "Germany", "DE", "Packaging wholesale", "German packaging wholesaler"),
    ("Paul Hildebrandt", "https://www.hildebrandt.de", "Germany", "DE", "Packaging wholesale", "German packaging wholesaler"),
    ("PackQueen", "https://www.packqueen.com.au", "Australia", "AU", "Boxes; shipping; retail packaging", "Australian packaging wholesaler"),
    ("Surepak", "https://www.surepak.com.au", "Australia", "AU", "Tapes; films; stretch wrap", "Australian packaging distributor"),
    ("ANW Packaging", "https://anw.com.sg", "Singapore", "SG", "Strapping; marking; industrial packaging", "Singapore packaging systems distributor"),
    ("Network Distribution", "https://us.networkdistribution.com", "Schaumburg, IL, USA", "US", "Industrial packaging; jan-san; disposables", "National packaging distribution organization"),
    ("Pollock Orora", "https://www.pollock.com", "USA", "US", "Packaging; paper; facility supplies", "US packaging/paper distributor"),
    ("Aramsco", "https://www.aramsco.com", "USA", "US", "Specialty supplies including packaging", "Specialty distributor with packaging lines"),
    ("Gillis & Lane", "https://gillislane.com", "USA", "US", "Packaging and paper products", "US packaging/paper distributor"),
    ("Pro Pac Materials", "https://www.propacmaterials.com", "Grand Prairie, TX, USA", "US", "Packaging materials and machinery", "Southwest packaging materials distributor"),
    ("NNZ", "https://www.nnz.com", "Netherlands", "NL", "Produce and industrial packaging", "International packaging distributor network"),
    ("Embageant", "https://embageant.fr", "France", "FR", "E-commerce packaging materials", "Online packaging wholesaler"),
    ("Furodet Emballages", "https://furodet-emballages.fr", "France", "FR", "Food and retail packaging", "French packaging merchant"),
    ("Brangs Heinrich", "https://www.brangs-heinrich.de", "Germany", "DE", "Industrial packaging; tapes; films", "German packaging wholesaler"),
    ("PakFactory", "https://www.pakfactory.com", "Canada", "CA", "Custom and stock packaging", "Canadian packaging channel supplier"),
    ("Skybox Packaging", "https://www.skyboxpackaging.com", "Ohio, USA", "US", "Packaging supplies", "US packaging distributor"),
    ("Discount Packaging", "https://www.discountpackaging.com", "USA", "US", "Boxes; shipping supplies", "US packaging e-merchant/distributor"),
]


def domain_of(url: str) -> str:
    u = url if "://" in url else "https://" + url
    return urlparse(u).netloc.lower().replace("www.", "")


def dns_ok(d: str) -> bool:
    try:
        socket.getaddrinfo(d, 443)
        return True
    except Exception:
        try:
            socket.getaddrinfo(d, 80)
            return True
        except Exception:
            return False


def fetch(website: str) -> str:
    base = website.rstrip("/")
    chunks = []
    ctx = ssl.create_default_context()
    for p in ["", "/about", "/about-us"]:
        try:
            req = Request(base + p, headers={"User-Agent": "Mozilla/5.0"})
            with urlopen(req, timeout=12, context=ctx) as resp:
                html = resp.read(70000).decode("utf-8", errors="ignore")
            if any(x in html.lower() for x in ("domain is for sale", "buy this domain", "parked domain")):
                return "PARKED"
            text = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
            text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
            text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"\s+", " ", text)
            chunks.append(text[:2500])
        except Exception:
            continue
    return "\n".join(chunks)[:6000]


def role_check(company, website, source):
    prompt = f"""Classify packaging channel.
Company: {company}
Website: {website}
SOURCE: {source[:5500]}
KEEP only Specialty Distributor/Wholesaler/Merchant/Channel Partner.
REMOVE Manufacturer/OEM/Converter/Hybrid/Wrong/Unknown.
JSON: {{"keep":true|false,"role":"...","market_fit":true|false,"founded":"...","hq":"...","reason":"..."}}"""
    resp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": "Strict channel-only."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.1,
        response_format={"type": "json_object"},
    )
    try:
        return json.loads(resp.choices[0].message.content or "{}")
    except Exception:
        return {"keep": False, "role": "Unknown", "market_fit": False}


def suggest(existing):
    sample = ", ".join(sorted(list(existing))[:90])
    out = []
    for region in [
        "United States independent packaging distributors with official websites",
        "United Kingdom packaging merchants and wholesalers with official websites",
        "Germany and Benelux packaging wholesalers with official websites",
        "Canada and Australia packaging distributors with official websites",
    ]:
        prompt = f"""List 20 REAL packaging distributors/wholesalers/merchants for {region}.
Real websites only. No manufacturers/hybrids. Avoid: {sample}
JSON: {{"companies":[{{"company":"","website":"https://","hq":"","country_code":"","founded":"Not publicly disclosed","core_categories":"","specialty_focus":"","summary":""}}]}}"""
        try:
            resp = client.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": "Compact JSON; real channel companies only."},
                    {"role": "user", "content": prompt},
                ],
                temperature=0.15,
                response_format={"type": "json_object"},
                max_tokens=3000,
            )
            comps = json.loads(resp.choices[0].message.content or "{}").get("companies") or []
            print(region[:30], len(comps), flush=True)
            out.extend(comps)
        except Exception as e:
            print("fail", e, flush=True)
    return out


def main():
    backup = FP.with_name(FP.stem + f"_pre_add100_push_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]

    # remove Staples Canada if Staples exists
    has_staples = any((r[0] or "").strip().lower() == "staples" for r in rows)
    cleaned = []
    removed = []
    for r in rows:
        name = (r[0] or "").strip()
        if has_staples and name.lower() == "staples canada":
            removed.append(name)
            continue
        cleaned.append(r)
    if removed:
        header = [c.value for c in ws[1]]
        ws.delete_rows(1, ws.max_row)
        ws.append(header)
        for r in cleaned:
            ws.append(r)
        rows = cleaned
        print("removed", removed)

    existing_names = {(r[0] or "").strip().lower() for r in rows}
    existing_domains = {domain_of(str(r[1] or "")) for r in rows if r[1]}
    added_now_count = len(rows) - BASELINE
    print("current", len(rows), "added", added_now_count, "need", TARGET - added_now_count)

    pool = []
    for item in SEEDS:
        name, web, hq, cc, cats, focus = item
        if name.lower() in existing_names or domain_of(web) in existing_domains:
            continue
        pool.append(
            {
                "company": name,
                "website": web,
                "hq": hq,
                "cc": cc,
                "founded": ND,
                "cats": cats,
                "focus": focus,
                "summary": f"{name} — {focus}.",
            }
        )

    for raw in suggest(existing_names):
        name = (raw.get("company") or "").strip()
        web = (raw.get("website") or "").strip()
        if not name or not web:
            continue
        if not web.startswith("http"):
            web = "https://" + web
        if name.lower() in existing_names or domain_of(web) in existing_domains:
            continue
        # skip staples/office depot clones
        if any(x in name.lower() for x in ("staples", "office depot", "uline", "bunzl", "raja", "manutan")):
            continue
        pool.append(
            {
                "company": name,
                "website": web,
                "hq": raw.get("hq") or ND,
                "cc": (raw.get("country_code") or "XX").upper()[:2],
                "founded": raw.get("founded") or ND,
                "cats": raw.get("core_categories") or "Packaging supplies",
                "focus": raw.get("specialty_focus") or "Packaging distributor",
                "summary": raw.get("summary") or "",
            }
        )

    # unique
    seen = set()
    uniq = []
    for c in pool:
        k = c["company"].lower()
        d = domain_of(c["website"])
        if k in seen or k in existing_names or d in existing_domains:
            continue
        seen.add(k)
        uniq.append(c)
    print("pool", len(uniq), flush=True)

    added = 0
    for i, c in enumerate(uniq):
        if (len(rows) - BASELINE) + added >= TARGET:
            break
        name, web, d = c["company"], c["website"], domain_of(c["website"])
        print(f"[{i+1}/{len(uniq)}] {name}", flush=True)
        if not dns_ok(d):
            print("  dns", flush=True)
            continue
        src = fetch(web)
        if src == "PARKED":
            print("  parked", flush=True)
            continue
        data = role_check(name, web, src or f"{name} {c['focus']}")
        role = (data.get("role") or "").lower()
        keep = bool(data.get("keep")) and bool(data.get("market_fit", True))
        if any(x in role for x in ("manufacturer", "oem", "converter", "hybrid", "wrong", "unknown")):
            keep = False
        if not keep:
            print(f"  REJECT {data.get('role')}", flush=True)
            time.sleep(0.12)
            continue
        if data.get("founded"):
            c["founded"] = str(data["founded"]).strip()
        if data.get("hq"):
            c["hq"] = str(data["hq"]).strip()
        cc = c.get("cc") or "XX"
        rc = "NA" if cc in {"US", "CA", "MX"} else ("APAC" if cc in {"AU", "NZ", "SG"} else "EU")
        geo = f"{rc}({cc})"
        ws.append(
            [
                c["company"],
                c["website"],
                c.get("founded") or ND,
                c.get("hq") or ND,
                geo,
                geo,
                "Private",
                ND,
                c.get("cats") or "Packaging supplies",
                c.get("focus") or "Packaging distributor",
                c["company"],
                "Both",
                "Specialty Distributor",
                ND,
                ND,
                ND,
                ND,
                ND,
                cc,
                rc,
                c.get("summary") or f"{c['company']} packaging specialty distributor.",
            ]
        )
        existing_names.add(name.lower())
        existing_domains.add(d)
        added += 1
        print(f"  KEEP ({(len(rows)-BASELINE)+added}/{TARGET})", flush=True)
        time.sleep(0.12)

    wb.save(FP)
    final = ws.max_row - 1
    audit = {
        "ts": datetime.now().isoformat(),
        "backup": str(backup),
        "removed": removed,
        "added": added,
        "after": final,
        "from_baseline": final - BASELINE,
        "file": str(FP),
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print("DONE after", final, "from_baseline", final - BASELINE)


if __name__ == "__main__":
    main()
