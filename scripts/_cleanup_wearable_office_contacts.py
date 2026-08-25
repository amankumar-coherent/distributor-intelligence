# -*- coding: utf-8 -*-
"""Cleanup junk Office No. / Office Email and apply verified hand fills."""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\vendor-intel")
OUT = ROOT / "output/web_expand/global_wearable_medical_devices_market_global"
XLSX = OUT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
NPD = "Not publicly disclosed"

# Exact junk values (scrap IPs, Excel floats, truncated garbage)
JUNK_EXACT = {
    "+0102-0103",
    "0 0 29.84 29.813",
    "10.132.175.19",
    "4.347826087",
    "00036 174",
    "0 0 512 512",
    "1.42857143",
    "10.202.52.71",
    "31.66666667",
    "192.168.3.173",
    "331.000000",
}

# Company → (office_phone or None to leave, office_email or None to leave)
# Only values verified from public contact/IR pages. None = do not overwrite.
HAND = {
    "Cardinal Health": ("+1 614-757-5000", "media.relations@cardinalhealth.com"),
    "Cardinal Health Canada": ("1-888-291-5033", None),
    "Henry Schein": ("1-631-843-5500", "custserv@henryschein.com"),
    "AmerisourceBergen": ("+1-610-727-7000", "InvestorRelations@cencora.com"),
    "DKSH": ("+41 44 386 7272", "marketing@dksh.com"),
    "Owens & Minor": ("+1 (804) 723-7000", None),
    "Owens & Minor Distribution": ("+1 (804) 723-7000", None),
    "Medline Industries": ("+1 847-949-5500", "service@medline.com"),
    "Alliance Healthcare": ("0330 100 0448", "customerservice@alliance-healthcare.co.uk"),
    "Zuellig Pharma": ("+65 6708 8188", None),
}

# Normalize dotted NANP phones like 877.599.7748 → 877-599-7748
DOT_NANP = re.compile(
    r"^(?:1[.\-\s]?)?(\d{3})[.\-\s](\d{3})[.\-\s](\d{4})$"
)


def emptyish(v: object) -> bool:
    s = str(v or "").strip().lower()
    return s in {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown"}


def is_ip(s: str) -> bool:
    return bool(re.fullmatch(r"\d{1,3}(?:\.\d{1,3}){3}", s.strip()))


def is_floatish_junk(s: str) -> bool:
    t = s.strip()
    if re.fullmatch(r"\d+\.\d+", t):
        try:
            f = float(t)
            # Excel serials / ratios, not phone-like
            if f < 1000 and len(t.replace(".", "")) <= 10:
                digits = t.replace(".", "")
                if not (len(digits) == 10 or (len(digits) == 11 and digits.startswith("1"))):
                    return True
        except ValueError:
            return False
    return False


def normalize_phone(ph: str) -> str:
    s = re.sub(r"\s+", " ", str(ph or "").strip())
    if emptyish(s) or s in JUNK_EXACT or is_ip(s) or is_floatish_junk(s):
        return NPD
    # Belgian style keep spaces
    m = DOT_NANP.match(s.replace(" ", ""))
    if m and "." in s:
        a, b, c = m.groups()
        return f"{a}-{b}-{c}"
    # 1.800.xxx.xxxx
    m2 = re.fullmatch(r"1[.\-\s](\d{3})[.\-\s](\d{3})[.\-\s](\d{4})", s.replace(" ", ""))
    if m2 and "." in s:
        a, b, c = m2.groups()
        return f"1-{a}-{b}-{c}"
    digits = re.sub(r"\D", "", s)
    if len(digits) < 8:
        return NPD
    if "666666" in digits:
        return NPD
    if re.fullmatch(r"\+?0{2,}\d{0,6}", s.replace(" ", "").replace("-", "")):
        return NPD
    return s


def normalize_email(em: str) -> str:
    e = str(em or "").strip().lower()
    if emptyish(e):
        return NPD
    bad = ("api-services-support@", "example.com", "noreply", "no-reply", "sentry")
    if any(b in e for b in bad):
        return NPD
    if "@" not in e:
        return NPD
    return e


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = XLSX.with_name(XLSX.stem + f"_BEFORE_OFFICE_CLEAN_{stamp}.xlsx")
    wb0 = openpyxl.load_workbook(XLSX)
    wb0.save(backup)
    wb0.close()

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    cleared_phone = 0
    reformatted = 0
    cleared_email = 0
    hand_n = 0

    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        ph = str(ws.cell(r, ic["Office No."]).value or "").strip()
        em = str(ws.cell(r, ic["Office Email"]).value or "").strip()

        if not emptyish(ph) and ph != NPD:
            new_ph = normalize_phone(ph)
            if new_ph != ph:
                if new_ph == NPD:
                    cleared_phone += 1
                else:
                    reformatted += 1
                ws.cell(r, ic["Office No."]).value = new_ph

        if not emptyish(em) and em != NPD:
            new_em = normalize_email(em)
            if new_em != em:
                if new_em == NPD:
                    cleared_email += 1
                ws.cell(r, ic["Office Email"]).value = new_em

        if co in HAND:
            hp, he = HAND[co]
            if hp:
                ws.cell(r, ic["Office No."]).value = hp
                hand_n += 1
            if he:
                cur = str(ws.cell(r, ic["Office Email"]).value or "").strip()
                if emptyish(cur) or cur == NPD:
                    ws.cell(r, ic["Office Email"]).value = he
                    hand_n += 1

        # blanks → NPD
        if emptyish(ws.cell(r, ic["Office No."]).value):
            ws.cell(r, ic["Office No."]).value = NPD
        if emptyish(ws.cell(r, ic["Office Email"]).value):
            ws.cell(r, ic["Office Email"]).value = NPD

        for col in (ic["Office No."], ic["Office Email"]):
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")

    for name in ("Office No.", "Office Email"):
        ws.column_dimensions[get_column_letter(ic[name])].width = 28

    wb.save(XLSX)
    wb.close()

    # recount
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    oi = headers.index("Office No.")
    ei = headers.index("Office Email")
    phone_n = email_n = 0
    total = ws.max_row - 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        ph = str(row[oi] or "").strip()
        em = str(row[ei] or "").strip()
        if ph and ph != NPD:
            phone_n += 1
        if em and em != NPD and "@" in em:
            email_n += 1
    wb.close()

    summary = {
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "backup": str(backup),
        "cleared_phone": cleared_phone,
        "reformatted_phone": reformatted,
        "cleared_email": cleared_email,
        "hand_updates": hand_n,
        "office_no_filled": phone_n,
        "office_email_filled": email_n,
        "total_rows": total,
    }
    (OUT / "wearable_office_cleanup_summary.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
