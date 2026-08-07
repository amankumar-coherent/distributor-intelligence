# -*- coding: utf-8 -*-
"""Remap smartwatch Region Segment to the six allowed regions only.

Invalid values fixed:
  - Eurasia / CIS  -> Europe  (Russia / CIS under Europe)
  - Global         -> primary HQ region (iStore -> Asia Pacific)
"""
from __future__ import annotations

import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SIX = [
    "Asia Pacific",
    "Europe",
    "North America",
    "Middle East",
    "Latin America",
    "Africa",
]
NOTES = {
    "Asia Pacific": "India, China, Japan, SEA, ANZ, Korea",
    "Europe": "UK + EU / wider Europe (incl. Russia / CIS)",
    "North America": "USA, Canada, Mexico",
    "Middle East": "GCC, UAE, Saudi, Turkey, Israel",
    "Latin America": "Brazil, Chile, Colombia, Argentina, Peru, Uruguay",
    "Africa": "South Africa, Nigeria, Egypt, Kenya, etc.",
}

PATHS = [
    (
        Path(
            r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL.xlsx"
        ),
        2,
        False,
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL_320plus.xlsx"
        ),
        3,
        True,
    ),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
CELL_FONT = Font(name="Calibri", size=9)
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP = Alignment(wrap_text=True, vertical="top")


def resolve_from_text(hq: str, cities: str, served: str) -> str:
    blob = f"{hq} {cities} {served}".lower()
    checks = [
        (
            "Asia Pacific",
            [
                "australia",
                "india",
                "china",
                "japan",
                "singapore",
                "malaysia",
                "korea",
                "taiwan",
                "hong kong",
                "indonesia",
                "thailand",
                "vietnam",
                "philippines",
                "new zealand",
            ],
        ),
        (
            "Europe",
            [
                "united kingdom",
                "uk",
                "germany",
                "france",
                "italy",
                "spain",
                "netherlands",
                "poland",
                "russia",
                "sweden",
                "switzerland",
                "ireland",
                "cis",
            ],
        ),
        ("North America", ["united states", "usa", "u.s.", "canada", "mexico"]),
        (
            "Middle East",
            ["uae", "dubai", "saudi", "qatar", "kuwait", "bahrain", "oman", "israel", "turkey"],
        ),
        ("Latin America", ["brazil", "chile", "colombia", "argentina", "peru", "uruguay"]),
        ("Africa", ["south africa", "nigeria", "egypt", "kenya", "morocco"]),
    ]
    for region, keys in checks:
        if any(k in blob for k in keys):
            return region
    return "Europe"


def write_rows(ws, headers, rows) -> None:
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = 18 if h != "Summary" else 36
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
        rr = ws.max_row
        for i in range(1, len(headers) + 1):
            cell = ws.cell(rr, i)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = WRAP
    last = get_column_letter(len(headers))
    if rows:
        ws.auto_filter.ref = f"A1:{last}{ws.max_row}"


def fix_file(path: Path, hdr_row: int, rebuild: bool) -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_pre_fix_eurasia_region_{stamp}{path.suffix}")
    shutil.copy2(path, bak)

    wb = openpyxl.load_workbook(path)
    ws = wb["Smartwatch Companies"]
    headers = [str(c.value or "").strip() for c in ws[hdr_row]]
    cols = {h: i + 1 for i, h in enumerate(headers) if h}
    seg_col = cols["Region Segment"]
    co_col = cols["Company"]
    hq_col = cols.get("Headquarters")
    cities_col = cols.get("Cities / Regions")
    served_col = cols.get("Regions Served")

    fixes: list[tuple[str, str, str]] = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        co = str(ws.cell(r, co_col).value or "").strip()
        if not co:
            continue
        old = str(ws.cell(r, seg_col).value or "").strip()
        if old in SIX:
            continue

        hq = str(ws.cell(r, hq_col).value or "") if hq_col else ""
        cities = str(ws.cell(r, cities_col).value or "") if cities_col else ""
        served = str(ws.cell(r, served_col).value or "") if served_col else ""

        if old in {"Eurasia / CIS", "Eurasia", "CIS"}:
            new = "Europe"
        elif co.lower() == "istore":
            new = "Asia Pacific"  # HQ Australia
        else:
            new = resolve_from_text(hq, cities, served)

        ws.cell(r, seg_col).value = new
        fixes.append((co, old, new))

    left = Counter()
    for r in range(hdr_row + 1, ws.max_row + 1):
        co = str(ws.cell(r, co_col).value or "").strip()
        if not co:
            continue
        left[str(ws.cell(r, seg_col).value or "").strip()] += 1

    if rebuild:
        companies: list[dict[str, str]] = []
        for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            if not row or not row[co_col - 1]:
                continue
            d = {
                headers[i]: str(row[i] if i < len(row) and row[i] is not None else "")
                for i in range(len(headers))
            }
            companies.append(d)

        for name in list(wb.sheetnames):
            if name.startswith("Region - ") or name == "Region Summary":
                del wb[name]

        counts = Counter(d.get("Region Segment", "") for d in companies)
        total = len(companies)
        sum_ws = wb.create_sheet("Region Summary", 1)
        sum_ws["A1"] = "Global Smartwatch Market — Region-wise segmentation (6 regions)"
        sum_ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
        sum_ws.merge_cells("A1:D1")
        for i, h in enumerate(["Region Segment", "Company Count", "% of Total", "Notes"], 1):
            cell = sum_ws.cell(3, i, h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = THIN
        rr = 4
        for region in SIX:
            n = counts.get(region, 0)
            sum_ws.cell(rr, 1, region).border = THIN
            sum_ws.cell(rr, 2, n).border = THIN
            pct = f"{100.0 * n / total:.1f}%" if total else "0%"
            sum_ws.cell(rr, 3, pct).border = THIN
            sum_ws.cell(rr, 4, NOTES[region]).border = THIN
            rr += 1
        sum_ws.cell(rr, 1, "TOTAL").font = Font(bold=True)
        sum_ws.cell(rr, 2, total).font = Font(bold=True)
        sum_ws.column_dimensions["A"].width = 22
        sum_ws.column_dimensions["B"].width = 16
        sum_ws.column_dimensions["C"].width = 12
        sum_ws.column_dimensions["D"].width = 55

        by_region: dict[str, list[dict[str, str]]] = defaultdict(list)
        for d in companies:
            by_region[d.get("Region Segment", "")].append(d)
        for region in SIX:
            rows = by_region.get(region) or []
            if not rows:
                continue
            rws = wb.create_sheet(f"Region - {region}"[:31])
            write_rows(rws, headers, rows)

        desired = ["Region Summary"] + [f"Region - {r}"[:31] for r in SIX if by_region.get(r)]
        for offset, name in enumerate(desired, start=1):
            if name in wb.sheetnames:
                current = wb.sheetnames.index(name)
                if current != offset:
                    wb.move_sheet(name, offset=offset - current)

    an = "Region Six Fix Audit"
    if an in wb.sheetnames:
        del wb[an]
    aws = wb.create_sheet(an)
    aws.append(["Company", "Old Region", "New Region"])
    for co, old, new in fixes:
        aws.append([co, old, new])

    wb.save(path)
    print(f"=== {path.name} ===")
    print(f"  backup: {bak.name}")
    print(f"  fixed {len(fixes)}:")
    for co, old, new in fixes:
        print(f"    {co}: {old} -> {new}")
    print(f"  region counts: {dict((k, left[k]) for k in SIX)}")
    bad = [k for k in left if k not in SIX]
    print(f"  non-six left: {bad or 'none'}")


def main() -> None:
    for path, hdr, rebuild in PATHS:
        if not path.exists():
            print("MISSING", path)
            continue
        fix_file(path, hdr, rebuild)


if __name__ == "__main__":
    main()
