# -*- coding: utf-8 -*-
"""Pass 2: fill remaining Online/Offline/Both NPD rows from evidence."""
from __future__ import annotations

import json
import time
from pathlib import Path

from openpyxl import load_workbook

import verify_flexpack_online_offline as m

FP = m.FP
CKPT = m.CKPT
NPD = m.NPD
LIMIT = 110


def main() -> None:
    ckpt = {}
    if CKPT.exists():
        try:
            ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
        except Exception:
            ckpt = {}

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers)}
    ci, wi, mi = ic["Company"], ic["Website"], ic[m.NEW_COL]

    rows = []
    for row in range(2, ws.max_row + 1):
        name = str(ws.cell(row, ci).value or "").strip()
        val = str(ws.cell(row, mi).value or "").strip()
        if name and name not in m.HAND and val == NPD:
            rows.append(row)

    print("npd_todo", len(rows), flush=True)
    updated = skipped = 0
    for idx, row in enumerate(rows[:LIMIT]):
        name = str(ws.cell(row, ci).value or "").strip()
        web = str(ws.cell(row, wi).value or "")
        print(f"[{idx+1}/{min(LIMIT,len(rows))}] {name} …", flush=True)
        evidence = m.fetch_evidence(web)
        mode = m.deepseek_classify(name, web, evidence) or m.heuristic(evidence)
        if mode in ("Online", "Offline", "Both"):
            ws.cell(row, mi).value = mode
            ckpt[name] = {"mode": mode}
            CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")
            updated += 1
            print("  ->", mode, flush=True)
        else:
            skipped += 1
        time.sleep(0.3)

    wb.save(FP)
    from collections import Counter

    counts = Counter(str(ws.cell(r, mi).value or "") for r in range(2, ws.max_row + 1))
    print("updated", updated, "skipped", skipped, "counts", dict(counts))
    print("saved", FP)


if __name__ == "__main__":
    main()
