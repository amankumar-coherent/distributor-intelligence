"""Fill Continent/Geography + Operational Presence for electrical FINAL Excel.

Reference style (global smartphone / firewall Excels):
  Operational Presence = semicolon-separated place names (HQ city + markets)
  Continent / Geography = REGION(Country[, Country…]); …
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from vendor_intel.enrichment.known_presence import (
    presence_to_operational_text,
    resolve_presence,
)
from vendor_intel.export.landscape_slim import (
    country_and_region_codes,
    format_continent_geography,
    parse_country_list,
)

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)


def main() -> None:
    backup = XLSX.with_name(XLSX.stem + "_BEFORE_GEO_OP_FILL.xlsx")
    if not backup.exists():
        backup.write_bytes(XLSX.read_bytes())

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    need = [
        "Company",
        "Headquarters",
        "Continent / Geography",
        "Operational Presence",
    ]
    for h in need:
        if h not in headers:
            raise SystemExit(f"missing column {h}")
    ci = headers.index("Company") + 1
    hi = headers.index("Headquarters") + 1
    gi = headers.index("Continent / Geography") + 1
    oi = headers.index("Operational Presence") + 1
    cci = headers.index("Country Code") + 1 if "Country Code" in headers else None
    rci = headers.index("Region Code") + 1 if "Region Code" in headers else None

    changes: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        hq = str(ws.cell(row_idx, hi).value or "").strip()
        old_geo = str(ws.cell(row_idx, gi).value or "").strip()
        old_op = str(ws.cell(row_idx, oi).value or "").strip()

        places = resolve_presence(name, hq)
        # Merge any already-good multi-place op values (keep extras)
        if old_op and ";" in old_op:
            for part in old_op.split(";"):
                p = part.strip()
                if p and p not in places:
                    places.append(p)
            # re-dedupe via helper
            places = resolve_presence(name, hq) + [
                p.strip() for p in old_op.split(";") if p.strip()
            ]
            # unique preserve order
            seen = set()
            uniq = []
            for p in places:
                k = p.lower()
                if k in seen:
                    continue
                seen.add(k)
                uniq.append(p)
            places = uniq

        op_text = presence_to_operational_text(places)
        countries = parse_country_list(op_text)
        if not countries:
            countries = parse_country_list(hq)
        geo_text = format_continent_geography(countries) if countries else old_geo

        ws.cell(row_idx, oi).value = op_text
        ws.cell(row_idx, gi).value = geo_text

        if cci and rci:
            cc, rc = country_and_region_codes(hq)
            if cc:
                ws.cell(row_idx, cci).value = cc
            if rc:
                ws.cell(row_idx, rci).value = rc

        if op_text != old_op or geo_text != old_geo:
            changes.append(
                {
                    "company": name,
                    "old_geo": old_geo,
                    "new_geo": geo_text,
                    "old_op": old_op,
                    "new_op": op_text,
                }
            )

    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(XLSX)

    dump = XLSX.parent / "_geo_op_fill_audit.json"
    dump.write_text(
        json.dumps(
            {
                "when": datetime.now().isoformat(timespec="seconds"),
                "changed": len(changes),
                "changes": changes[:50],
                "changes_total": len(changes),
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(f"updated {len(changes)} rows")
    for c in changes[:12]:
        print(f"  {c['company']}")
        print(f"    GEO: {c['old_geo']!r} -> {c['new_geo']!r}")
        print(f"    OP : {c['old_op']!r} -> {c['new_op']!r}")
    print(f"wrote {XLSX}")
    print(f"audit {dump}")
    print(f"backup {backup}")


if __name__ == "__main__":
    main()
