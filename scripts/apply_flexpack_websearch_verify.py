# -*- coding: utf-8 -*-
"""Apply WebSearch-subagent verification results (Core Categories, Specialty Focus,
Key Brands Represented) to the flexpack FINAL workbook. Backs up first.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
WORKDIR = ROOT / "output" / "web_expand" / "global_flexible_packaging_market_global"
XLSX = WORKDIR / "global_flexible_packaging_market_global_FINAL.xlsx"
NPD = "Not publicly disclosed"


def clean_brands(raw: str) -> list[str]:
    parts = re.split(r"[;,]+", raw or "")
    out = []
    for p in parts:
        b = p.strip(" .")
        if b and b.lower() not in {"none", "n/a", "not publicly disclosed"}:
            out.append(b)
    return out


def main(results_path: str) -> None:
    results = json.loads(Path(results_path).read_text(encoding="utf-8"))
    by_name = {r["company"].strip().lower(): r for r in results}

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_WEBSEARCH_VERIFY_{stamp}" + XLSX.suffix)
    shutil.copy2(XLSX, bak)
    print(f"backup -> {bak.name}")

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers) if h}

    core_changed = spec_changed = brand_changed = matched = 0
    unmatched = []

    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        rec = by_name.get(co.lower())
        if not rec:
            unmatched.append(co)
            continue
        matched += 1

        if not rec.get("core_categories_accurate", True) and rec.get("core_categories_corrected"):
            ws.cell(r, ic["Core Categories"]).value = rec["core_categories_corrected"].strip()
            core_changed += 1

        if not rec.get("specialty_focus_accurate", True) and rec.get("specialty_focus_corrected"):
            ws.cell(r, ic["Specialty Focus"]).value = rec["specialty_focus_corrected"].strip()
            spec_changed += 1

        cur_brand = str(ws.cell(r, ic["Key Brands Represented"]).value or "").strip()
        is_junk = bool(rec.get("current_brand_is_placeholder_junk", False))
        found = rec.get("brands_found", "") or ""
        found_list = clean_brands(found)

        if found_list:
            if cur_brand and cur_brand.lower() != NPD.lower() and not is_junk:
                existing = clean_brands(cur_brand)
                existing_lower = {e.lower() for e in existing}
                additions = [b for b in found_list if b.lower() not in existing_lower]
                if additions:
                    ws.cell(r, ic["Key Brands Represented"]).value = "; ".join(existing + additions)
                    brand_changed += 1
            else:
                ws.cell(r, ic["Key Brands Represented"]).value = "; ".join(found_list)
                brand_changed += 1
        elif is_junk:
            ws.cell(r, ic["Key Brands Represented"]).value = NPD
            brand_changed += 1
        elif not cur_brand:
            ws.cell(r, ic["Key Brands Represented"]).value = NPD

    wb.save(XLSX)
    print(json.dumps({
        "matched": matched,
        "unmatched_in_results": unmatched,
        "core_changed": core_changed,
        "spec_changed": spec_changed,
        "brand_changed": brand_changed,
    }, indent=2))


if __name__ == "__main__":
    main(sys.argv[1])
