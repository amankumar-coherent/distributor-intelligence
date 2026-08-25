from collections import Counter
import json
from pathlib import Path
from openpyxl import load_workbook

p = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)
wb = load_workbook(p, read_only=True, data_only=True)
ws = wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
idx = {h: i for i, h in enumerate(headers)}
cp = idx.get("Contact Person")
filled = empty = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    person = str((row[cp] or "") if cp is not None else "").strip()
    if not person or "*" in person:
        empty += 1
    else:
        filled += 1
print(f"electrical excel: filled={filled} empty={empty} total={filled + empty}")
wb.close()

prog = p.with_name(p.stem + "_gfirst_contact_progress.jsonl")
c = Counter()
n = 0
for line in prog.read_text(encoding="utf-8").splitlines():
    if not line.strip():
        continue
    rec = json.loads(line)
    c[rec.get("status")] += 1
    n += 1
print(f"progress lines={n} {dict(c)}")
