# -*- coding: utf-8 -*-
"""ADC Contact Person / Role / LinkedIn — sales or procurement only. No CEO."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

EMPTY = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}

ADC_PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global"
    r"\global_antibody_drug_conjugates_market.xlsx"
)

# Exact Company names. Sales / procurement / commercial / BD only.
ADC = {
    "Biologics by McKesson": {
        "person": "Patrick McGlinchey",
        "role": "Vice President, Provider Sales",
        "li": "https://www.linkedin.com/in/patrick-mcglinchey-3359711a",
    },
    "McKesson Medical-Surgical": {
        "person": "Deborah Haywood",
        "role": "Senior Vice President, Enterprise Sales",
        "li": "https://www.linkedin.com/in/deborah-s-haywood-79468111",
    },
    "McKesson Canada": {
        "person": "Yves Blanchette",
        "role": "Vice President, National Sales and Commercial Strategy",
        "li": "https://www.linkedin.com/in/yves-blanchette-06660846",
    },
    "McKesson Specialty Care Distribution": {
        "person": "Michael Martin",
        "role": "Vice President, National Accounts, Specialty Pharmacy",
        "li": "https://www.linkedin.com/in/michaelmartin718",
    },
    "McKesson Specialty Health": {
        "person": "Mark Jolly",
        "role": "Vice President, Payor Sales",
        "li": "https://www.linkedin.com/in/mark-jolly-1161097",
    },
    "Cardinal Health Specialty Pharmaceutical Distribution": {
        "person": "Kraig Corwin",
        "role": "Vice President, National Markets Sales (Mail and Specialty)",
        "li": "https://www.linkedin.com/in/kcorwin",
    },
    "Cardinal Health Specialty Solutions": {
        "person": "Jon McWhinney",
        "role": "Director of Sales, Specialty Solutions",
        "li": "https://www.linkedin.com/in/jon-mcwhinney-a9848711",
    },
    "Azelis Group NV": {
        "person": "Marnik Tack",
        "role": "Chief Commercial Officer, EMEA",
        "li": "https://www.linkedin.com/in/marnik-tack-620aa13",
    },
    "Avantor Inc.": {
        "person": "Bonnie Smith",
        "role": "Vice President of Sales, Bioprocessing Americas",
        "li": "https://www.linkedin.com/in/bonnie-smith-056b2640",
    },
    "Fisher Scientific": {
        "person": "Suzanne Sheridan",
        "role": "Vice President, Commercial Operations, Fisher Scientific Channel",
        "li": "https://www.linkedin.com/in/suzanne-sheridan",
    },
    "Henry Schein Medical": {
        "person": "Ty Ford",
        "role": "Vice President and General Manager, U.S. Medical Sales",
        "li": "https://www.linkedin.com/in/ty-ford-5b53162",
    },
    "CVS Specialty": {
        "person": "Christine Walsh",
        "role": "Vice President of Specialty and Infusion Sales, Business Development",
        "li": "https://www.linkedin.com/in/christine-walsh-aa047b9",
    },
    "Coram CVS Specialty Infusion": {
        "person": "Christine Walsh",
        "role": "Vice President of Specialty and Infusion Sales, Business Development",
        "li": "https://www.linkedin.com/in/christine-walsh-aa047b9",
    },
    "AllianceRx Walgreens Pharmacy": {
        "person": "David Duerst",
        "role": "Vice President, Specialty Sales",
        "li": "https://www.linkedin.com/in/davidduerst",
    },
    "Walgreens Specialty Pharmacy": {
        "person": "David Duerst",
        "role": "Vice President, Specialty Sales",
        "li": "https://www.linkedin.com/in/davidduerst",
    },
    "Accredo Health Group": {
        "person": "Vic Perini",
        "role": "Vice President, Sales and Account Management",
        "li": "https://www.linkedin.com/in/vic-perini-24aa93b",
    },
    "Express Scripts Accredo network": {
        "person": "Vic Perini",
        "role": "Vice President, Sales and Account Management",
        "li": "https://www.linkedin.com/in/vic-perini-24aa93b",
    },
    "Cigna Specialty Pharmacy / Accredo": {
        "person": "Vic Perini",
        "role": "Vice President, Sales and Account Management",
        "li": "https://www.linkedin.com/in/vic-perini-24aa93b",
    },
    "Anda Inc.": {
        "person": "Daniel Shannon",
        "role": "Senior Director, Pharmacy Sales and GPO Relations",
        "li": "https://www.linkedin.com/in/daniel-shannon-7946a648",
    },
    "Anda Pharmaceuticals": {
        "person": "Daniel Shannon",
        "role": "Senior Director, Pharmacy Sales and GPO Relations",
        "li": "https://www.linkedin.com/in/daniel-shannon-7946a648",
    },
    "Galenica / Alloga": {
        "person": "Alistair Binns",
        "role": "Senior Director Europe, Client and Commercial Development",
        "li": "https://www.linkedin.com/in/alistair-binns-7187121",
    },
    "Zuellig Pharma": {
        "person": "Jan Christensen",
        "role": "Vice President Sales, Customer and Client Excellence",
        "li": "https://www.linkedin.com/in/janraskchristensen",
    },
    "Zuellig Pharma Philippines": {
        "person": "Larry Cabrera",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/larry-cabrera-2177aa20",
    },
    "Diethelm Thailand / DKSH Thailand": {
        "person": "Nannaphat Limkome",
        "role": "General Manager, Healthcare, Pharma Commercialization and Business Development",
        "li": "https://www.linkedin.com/in/nannaphat-limkome-1606b146",
    },
    "Brocacef": {
        "person": "Helene Lantsheer",
        "role": "Group Commercial Director",
        "li": "https://www.linkedin.com/in/helene-lantsheer-3626203",
    },
    "Diplomat Specialty Pharmacy (Optum)": {
        "person": "Michele Ringler",
        "role": "Regional Vice President of Sales, Optum Specialty Pharmacy",
        "li": "https://www.linkedin.com/in/micheleringler",
    },
    "Optum Specialty Pharmacy": {
        "person": "Michele Ringler",
        "role": "Regional Vice President of Sales",
        "li": "https://www.linkedin.com/in/micheleringler",
    },
    "CenterWell Pharmacy": {
        "person": "Lindsey Irvin",
        "role": "Market Director, Specialty Sales",
        "li": "https://www.linkedin.com/in/lindsey-irvin-3b6a4a304",
    },
    "Harvard Drug Group": {
        "person": "Randy Saal",
        "role": "Director of Sales, Retail Division",
        "li": "https://www.linkedin.com/in/randy-saal-93992732",
    },
    "Gulf Drug LLC": {
        "person": "Mohamed Aly Rohyime",
        "role": "Sales Manager",
        "li": "https://www.linkedin.com/in/mohamed-aly-rohyime-19354085",
    },
    "Medline Industries": {
        "person": "Doug Golwas",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/doug-golwas-b7132754",
    },
    "Morris & Dickson Co., LLC": {
        "person": "Todd Laser",
        "role": "Senior Vice President, Retail Sales",
        "li": "https://www.linkedin.com/in/todd-laser-a09b3618",
    },
    "Morris & Dickson Specialty Distribution": {
        "person": "Todd Laser",
        "role": "Senior Vice President, Retail Sales",
        "li": "https://www.linkedin.com/in/todd-laser-a09b3618",
    },
    "Option Care Health specialty infusion": {
        "person": "Prateek Bhatia",
        "role": "Executive Vice President, Specialty Sales and Solutions",
        "li": "https://www.linkedin.com/in/prateek-bhatia",
    },
    "Kohl & Frisch": {
        "person": "David Lang",
        "role": "Senior Director, National Independent Sales",
        "li": "https://www.linkedin.com/in/david-lang-a04845141",
    },
    "Mediq": {
        "person": "Maurits Bisschop",
        "role": "Sales Director, Netherlands",
        "li": "https://www.linkedin.com/in/maurits-bisschop-03636b4b",
    },
}


def _headers(ws):
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            return r, {h: i + 1 for i, h in enumerate(vals) if h}
    raise RuntimeError("headers not found")


def _empty(v):
    return (str(v or "").strip().lower() in EMPTY) or (v is None)


def save_atomic(wb, path: Path):
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def apply(path: Path, contacts: dict) -> int:
    wb = load_workbook(path)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    hr, col = _headers(ws)
    c_co = col["Company"]
    c_cp = col["Contact Person"]
    c_role = col["Role"]
    c_li = col["LinkedIn"]
    n = 0
    unmatched = set(contacts)
    for row in range(hr + 1, ws.max_row + 1):
        co = str(ws.cell(row, c_co).value or "").strip()
        if co not in contacts:
            continue
        unmatched.discard(co)
        if not _empty(ws.cell(row, c_cp).value):
            print("SKIP already filled", co)
            continue
        rec = contacts[co]
        ws.cell(row, c_cp).value = rec["person"]
        ws.cell(row, c_role).value = rec["role"]
        ws.cell(row, c_li).value = rec["li"]
        n += 1
        print(f"FILL {co} | {rec['person']} | {rec['role']}")
    save_atomic(wb, path)
    wb.close()
    print("filled", n, path.name)
    if unmatched:
        print("UNMATCHED", unmatched)
    return n


if __name__ == "__main__":
    apply(ADC_PATH, ADC)
