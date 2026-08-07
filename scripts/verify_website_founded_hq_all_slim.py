#!/usr/bin/env python3
"""Verify Website, Founded, Headquarters for all FINAL_SLIM rows.

Search bar style (same as user):
  \"{company} official website founded year headquarters\"

Pipeline:
  1) Google AI scraper — batches of 25
  2) DeepSeek (OPENAI_* / DeepSeek) for residuals / misses

Reports fixed / confirmed / skipped counts. Checkpointed for resume.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import load_workbook
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")
os.environ["GOOGLE_AI_SCRAPER_ENABLED"] = "true"
os.environ["GOOGLE_AI_SCRAPER_FORCE_OVERVIEW"] = "true"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRIES"] = "2"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRY_SLEEP_SEC"] = "1.0"

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

SLIM_FILES = [
    ROOT / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL_SLIM.xlsx",
    ROOT / "output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL_SLIM.xlsx",
    ROOT / "output/run4/global_firewall_market_global/global_firewall_market_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_antibody_drug_conjugates_market_global/global_antibody_drug_conjugates_market_FINAL_SLIM.xlsx",
]

BATCH = 25
CURRENT_YEAR = 2026
PROGRESS = ROOT / "scripts/_verify_web_founded_hq_progress.jsonl"
SUMMARY = ROOT / "scripts/_verify_web_founded_hq_summary.json"


def norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [norm(c.value) for c in ws[r] if c.value not in (None, "")]
        if "company" in vals:
            return r
    return 1


def collect_rows() -> list[dict]:
    rows: list[dict] = []
    for f in SLIM_FILES:
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        hr = header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        wi = headers.index("Website") if "Website" in headers else None
        fi = headers.index("Founded") if "Founded" in headers else None
        hi = headers.index("Headquarters") if "Headquarters" in headers else None
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if not row or not row[ci]:
                continue
            company = str(row[ci]).strip()
            rows.append(
                {
                    "file": str(f),
                    "company": company,
                    "website": str(row[wi] or "").strip() if wi is not None else "",
                    "founded": str(row[fi] or "").strip() if fi is not None else "",
                    "hq": str(row[hi] or "").strip() if hi is not None else "",
                    "key": f"{norm(f.name)}::{norm(company)}",
                }
            )
        wb.close()
    return rows


def load_done() -> set[str]:
    done: set[str] = set()
    if PROGRESS.exists():
        for line in PROGRESS.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
                if obj.get("key") and obj.get("processed"):
                    done.add(obj["key"])
            except Exception:
                pass
    return done


def clean_url(u: object) -> str | None:
    s = str(u or "").strip()
    if not s or s.upper() in {"NPD", "N/A", "NA", "NONE", "UNKNOWN"}:
        return None
    if not re.match(r"^https?://", s, re.I):
        s = "https://" + s.lstrip("/")
    try:
        p = urlparse(s)
        if not p.netloc or "." not in p.netloc:
            return None
        # drop obvious junk hosts
        host = p.netloc.lower()
        junk = ("reddit.com", "facebook.com", "archive.org", "genius.com", "zhihu.com", "tiktok.com")
        if any(j in host for j in junk):
            return None
        return f"{p.scheme}://{p.netloc}{p.path}".rstrip("/")
    except Exception:
        return None


def year_ok(y: object) -> int | None:
    if y is None:
        return None
    if isinstance(y, int):
        yi = y
    else:
        m = re.search(r"(18|19|20)\d{2}", str(y))
        yi = int(m.group(0)) if m else None
    if yi is None:
        return None
    if 1800 <= yi <= CURRENT_YEAR:
        return yi
    return None


def deepseek() -> OpenAI:
    return OpenAI(
        api_key=os.getenv("OPENAI_API_KEY") or "",
        base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/"),
    )


def parse_batch(companies: list[dict], text: str, source: str) -> dict[str, dict]:
    names = [c["company"] for c in companies]
    client = deepseek()
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    prompt = f"""Extract official website, founded year, and headquarters for each company.

Companies:
{json.dumps(names, ensure_ascii=False)}

Source text ({source}):
```
{(text or '')[:14000]}
```

Return JSON array ONLY:
[
  {{
    "company": "exact name from list",
    "website": "https://..." or null,
    "founded_year": 1994 or null,
    "headquarters": "City, Country" or null,
    "confidence": "high"|"medium"|"low"
  }}
]
Rules:
- Prefer official company website (not LinkedIn/Wikipedia/news unless that is all available).
- founded_year 1800-{CURRENT_YEAR}; never use copyright footer year alone.
- headquarters: city + country when possible.
- Do not invent. Skip company if nothing useful (omit from array).
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    raw = (resp.choices[0].message.content or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.M).strip()
    try:
        data = json.loads(raw)
    except Exception:
        return {}
    if isinstance(data, dict):
        data = data.get("companies") or data.get("results") or []
    out: dict[str, dict] = {}
    if not isinstance(data, list):
        return out
    for row in data:
        if not isinstance(row, dict) or not row.get("company"):
            continue
        out[norm(row["company"])] = {
            "website": clean_url(row.get("website")),
            "founded": year_ok(row.get("founded_year")),
            "hq": (str(row.get("headquarters") or "").strip() or None),
            "confidence": (row.get("confidence") or "medium"),
            "source": source,
        }
    return out


def search_query(company: str) -> str:
    # Same style as user's Google search bar examples
    return f"{company} official website founded year headquarters"


def build_batch_query(batch: list[dict]) -> str:
    lines = [
        "For EACH company, give:",
        "1) Official website URL",
        "2) Founded year (YYYY)",
        "3) Headquarters (city, country)",
        "",
        "Search style for each: \"<company> official website founded year headquarters\"",
        "",
        "Companies:",
    ]
    for i, c in enumerate(batch, 1):
        hint = []
        if c.get("website"):
            hint.append(f"current_website={c['website']}")
        if c.get("founded"):
            hint.append(f"current_founded={c['founded']}")
        if c.get("hq"):
            hint.append(f"current_hq={c['hq']}")
        extra = f" ({'; '.join(hint)})" if hint else ""
        lines.append(f"{i}. {c['company']}{extra}")
    lines.append("Do not invent. Prefer official sources.")
    return "\n".join(lines)


async def google_batch(batch: list[dict]) -> dict[str, dict]:
    q = build_batch_query(batch)
    data = await google_ai_ask(q, llm_clean=True)
    md = str(data.get("markdown") or data.get("markdown_raw") or "").strip()
    found = parse_batch(batch, md, "google_ai") if md else {}

    missing = [c for c in batch if norm(c["company"]) not in found]
    if missing:
        q2 = build_batch_query(missing)
        data2 = await google_ai_ask(q2, llm_clean=True)
        md2 = str(data2.get("markdown") or data2.get("markdown_raw") or "").strip()
        if md2:
            found.update(parse_batch(missing, md2, "google_ai_retry"))

    still = [c for c in batch if norm(c["company"]) not in found]
    for c in still[:6]:
        d1 = await google_ai_ask(search_query(c["company"]), llm_clean=True)
        md1 = str(d1.get("markdown") or d1.get("markdown_raw") or "").strip()
        if md1:
            found.update(parse_batch([c], md1, "google_ai_single"))
    return found


def ddg_snippets(company: str) -> str:
    try:
        from ddgs import DDGS
    except ImportError:
        try:
            from duckduckgo_search import DDGS  # type: ignore
        except ImportError:
            return ""
    bits: list[str] = []
    try:
        with DDGS() as ddgs:
            for r in ddgs.text(search_query(company), max_results=5):
                bits.append(
                    f"- {r.get('title','')}: {r.get('body') or r.get('description') or ''} | {r.get('href') or r.get('link') or ''}"
                )
    except Exception as e:
        bits.append(f"(search error: {e})")
    return "\n".join(bits)


def deepseek_one(c: dict) -> dict:
    snippets = ddg_snippets(c["company"])
    client = deepseek()
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    prompt = f"""Verify firmographics for this company.

Company: {c['company']}
Current website: {c.get('website') or 'empty'}
Current founded: {c.get('founded') or 'empty'}
Current HQ: {c.get('hq') or 'empty'}

Web snippets:
{snippets or '(none)'}

Return JSON only:
{{
  "website": "https://..." or null,
  "founded_year": 1994 or null,
  "headquarters": "City, Country" or null,
  "confidence": "high"|"medium"|"low"
}}
Rules: official website preferred; founded 1800-{CURRENT_YEAR}; do not invent.
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    raw = (resp.choices[0].message.content or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.M).strip()
    try:
        data = json.loads(raw)
    except Exception:
        return {}
    if not isinstance(data, dict):
        return {}
    return {
        "website": clean_url(data.get("website")),
        "founded": year_ok(data.get("founded_year")),
        "hq": (str(data.get("headquarters") or "").strip() or None),
        "confidence": data.get("confidence") or "medium",
        "source": "deepseek",
    }


def decide_updates(old: dict, new: dict) -> dict:
    """Return fields to write + action flags. Skip low confidence invents."""
    conf = str(new.get("confidence") or "medium").lower()
    if conf == "low" and not any([new.get("website"), new.get("founded"), new.get("hq")]):
        return {"action": "skip", "changes": {}, "reason": "low_empty"}

    changes: dict[str, str] = {}
    reasons: list[str] = []

    # Website
    new_web = new.get("website")
    old_web = clean_url(old.get("website"))
    if new_web:
        if not old_web or old_web != new_web:
            # replace junk or empty; also replace if different domain and conf high/medium
            if conf in {"high", "medium"}:
                changes["Website"] = new_web
                reasons.append("website")

    # Founded
    new_y = new.get("founded")
    old_y = year_ok(old.get("founded"))
    if new_y is not None:
        if old_y is None or str(old.get("founded") or "").upper() == "NPD":
            changes["Founded"] = str(new_y)
            reasons.append("founded_fill")
        elif new_y != old_y and conf in {"high", "medium"}:
            # protect huge jumps unless high
            if abs(new_y - old_y) <= 80 or conf == "high":
                changes["Founded"] = str(new_y)
                reasons.append("founded_fix")

    # HQ
    new_hq = new.get("hq")
    old_hq = str(old.get("hq") or "").strip()
    if new_hq and conf in {"high", "medium"}:
        if (
            not old_hq
            or old_hq.upper() in {"NPD", "N/A", "GLOBAL", "UNKNOWN"}
            or (norm(new_hq) != norm(old_hq) and len(new_hq) >= 5)
        ):
            # avoid tiny noisy flips when both look detailed and similar-ish
            if not old_hq or old_hq.upper() in {"NPD", "N/A", "GLOBAL", "UNKNOWN"} or conf == "high" or abs(len(new_hq) - len(old_hq)) > 3:
                changes["Headquarters"] = new_hq
                reasons.append("hq")

    if not changes:
        if new_web or new_y is not None or new_hq:
            return {"action": "confirmed", "changes": {}, "reason": "already_ok", "source": new.get("source")}
        return {"action": "skip", "changes": {}, "reason": "no_data"}

    return {
        "action": "fix",
        "changes": changes,
        "reason": ",".join(reasons),
        "source": new.get("source"),
        "confidence": conf,
    }


def apply_results(results: list[dict]) -> None:
    by_file: dict[str, dict[str, dict]] = {}
    for fx in results:
        if fx.get("action") != "fix" or not fx.get("changes"):
            continue
        by_file.setdefault(fx["file"], {})[norm(fx["company"])] = fx["changes"]

    for fpath, mapping in by_file.items():
        wb = load_workbook(fpath)
        ws = wb.active
        hr = header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            company = ws.cell(r, ci + 1).value
            if not company:
                continue
            ch = mapping.get(norm(str(company)))
            if not ch:
                continue
            for col, val in ch.items():
                if col in headers:
                    ws.cell(r, headers.index(col) + 1).value = val
            n += 1
        wb.save(fpath)
        print(f"Applied fixes to {n} rows -> {Path(fpath).name}")


async def main() -> None:
    rows = collect_rows()
    done = load_done()
    todo = [r for r in rows if r["key"] not in done]
    print(f"Companies: {len(rows)} | already processed: {len(done)} | remaining: {len(todo)}")
    print('Search bar style: "{company} official website founded year headquarters"')

    google_ok = await google_ai_ping()
    print(f"Google AI scraper: {'connected' if google_ok else 'DOWN — DeepSeek only'}")

    results: list[dict] = []
    PROGRESS.parent.mkdir(parents=True, exist_ok=True)

    # reload previous results for apply at end
    prev: list[dict] = []
    if PROGRESS.exists():
        for line in PROGRESS.read_text(encoding="utf-8").splitlines():
            if line.strip():
                try:
                    prev.append(json.loads(line))
                except Exception:
                    pass

    with PROGRESS.open("a", encoding="utf-8") as prog:
        # Phase 1 Google AI
        if google_ok and todo:
            for i in range(0, len(todo), BATCH):
                batch = todo[i : i + BATCH]
                bnum = i // BATCH + 1
                btot = (len(todo) + BATCH - 1) // BATCH
                print(f"\n[Google AI] Batch {bnum}/{btot} ({len(batch)})", flush=True)
                try:
                    found = await google_batch(batch)
                except Exception as e:
                    print(f"  batch error: {e}", flush=True)
                    found = {}

                need_ds: list[dict] = []
                for c in batch:
                    parsed = found.get(norm(c["company"]))
                    if not parsed:
                        need_ds.append(c)
                        continue
                    decision = decide_updates(c, parsed)
                    rec = {
                        "key": c["key"],
                        "file": c["file"],
                        "company": c["company"],
                        "processed": True,
                        "phase": "google_ai",
                        **decision,
                        "parsed": {
                            "website": parsed.get("website"),
                            "founded": parsed.get("founded"),
                            "hq": parsed.get("hq"),
                            "confidence": parsed.get("confidence"),
                        },
                    }
                    prog.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    prog.flush()
                    results.append(rec)
                    print(
                        f"  {decision['action'].upper()} {c['company']}: {decision.get('changes') or decision.get('reason')}",
                        flush=True,
                    )

                # Phase 1b DeepSeek for batch misses immediately
                for c in need_ds:
                    print(f"  [DeepSeek] {c['company']}", flush=True)
                    try:
                        parsed = deepseek_one(c)
                    except Exception as e:
                        parsed = {}
                        print(f"    error: {e}", flush=True)
                    if not parsed:
                        rec = {
                            "key": c["key"],
                            "file": c["file"],
                            "company": c["company"],
                            "processed": True,
                            "phase": "deepseek",
                            "action": "skip",
                            "changes": {},
                            "reason": "no_data",
                        }
                    else:
                        decision = decide_updates(c, parsed)
                        rec = {
                            "key": c["key"],
                            "file": c["file"],
                            "company": c["company"],
                            "processed": True,
                            "phase": "deepseek",
                            **decision,
                            "parsed": {
                                "website": parsed.get("website"),
                                "founded": parsed.get("founded"),
                                "hq": parsed.get("hq"),
                                "confidence": parsed.get("confidence"),
                            },
                        }
                    prog.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    prog.flush()
                    results.append(rec)
                    print(
                        f"    {rec['action'].upper()}: {rec.get('changes') or rec.get('reason')}",
                        flush=True,
                    )
                await asyncio.sleep(0.3)

        elif todo:
            # Google down — DeepSeek all
            for i, c in enumerate(todo, 1):
                print(f"[DeepSeek {i}/{len(todo)}] {c['company']}", flush=True)
                try:
                    parsed = deepseek_one(c)
                except Exception as e:
                    parsed = {}
                    print(f"  error: {e}")
                if not parsed:
                    rec = {
                        "key": c["key"],
                        "file": c["file"],
                        "company": c["company"],
                        "processed": True,
                        "phase": "deepseek",
                        "action": "skip",
                        "changes": {},
                        "reason": "no_data",
                    }
                else:
                    decision = decide_updates(c, parsed)
                    rec = {
                        "key": c["key"],
                        "file": c["file"],
                        "company": c["company"],
                        "processed": True,
                        "phase": "deepseek",
                        **decision,
                        "parsed": parsed,
                    }
                prog.write(json.dumps(rec, ensure_ascii=False) + "\n")
                prog.flush()
                results.append(rec)

    all_recs = prev + results
    # de-dupe by key keeping last
    by_key: dict[str, dict] = {}
    for r in all_recs:
        if r.get("key"):
            by_key[r["key"]] = r
    final = list(by_key.values())

    print("\nApplying Excel updates...")
    apply_results(final)

    fixed = sum(1 for r in final if r.get("action") == "fix")
    confirmed = sum(1 for r in final if r.get("action") == "confirmed")
    skipped = sum(1 for r in final if r.get("action") == "skip")
    by_phase = {"google_ai": 0, "deepseek": 0}
    for r in final:
        if r.get("action") == "fix":
            ph = r.get("phase") or "unknown"
            by_phase[ph] = by_phase.get(ph, 0) + 1

    summary = {
        "total_companies": len(rows),
        "processed": len(final),
        "fixed": fixed,
        "confirmed_already_ok": confirmed,
        "skipped": skipped,
        "fixed_via_google_ai": by_phase.get("google_ai", 0),
        "fixed_via_deepseek": by_phase.get("deepseek", 0),
        "search_query_style": "{company} official website founded year headquarters",
    }
    SUMMARY.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print("\n=== SUMMARY ===")
    for k, v in summary.items():
        print(f"{k}: {v}")
    print(f"Wrote {SUMMARY}")


if __name__ == "__main__":
    asyncio.run(main())
