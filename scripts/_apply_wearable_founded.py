# -*- coding: utf-8 -*-
"""Apply Founded to Wearable FINAL from verified overrides + cleaned harvest cache.

No Google AI scraper. Never invent years.
"""
from __future__ import annotations

import json
import re
import shutil
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import httpx
from openpyxl import load_workbook

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
CACHE = ROOT / "_wearable_founded_harvest_cache.json"
AUDIT = ROOT / "_wearable_founded_fill_audit.json"
REMAINING = ROOT / "_wearable_founded_still_needed.json"

# Agent-verified overrides (About / company history / reputable sources)
VERIFIED: dict[str, str] = {
    "AAH Pharmaceuticals": "1892",
    "AddLife AB": "2014",
    "AddLife": "2014",
    "Adcock Ingram": "1890",
    "Adcock Ingram Critical Care Distribution": "1890",
    "180 Medical": "2002",
    "Advanced Diabetes Supply": "2002",
    "ActivStyle": "1997",
    "AdaptHealth": "2012",
    "Medline": "1966",
    "Medline Industries": "1966",
    "Henry Schein": "1932",
    "Cardinal Health": "1971",
    "McKesson": "1833",
    "McKesson Corporation": "1833",
    "McKesson Canada": "1833",
    "Suzuken": "1932",
    "Alfresa Holdings Corporation": "2003",
    "Alfresa Holdings": "2003",
    "Toho Holdings": "1948",
    "Bound Tree Medical": "1978",
    "Life-Assist": "1977",
    "Vitality Medical": "2000",
    "Concordance Healthcare Solutions": "2016",
    "Mediq": "1899",
    "Grifols": "1909",
    "Medipal Holdings": "1898",
    "MEDIPAL HOLDINGS": "1898",
    "Medipal": "1898",
    "Alliance Healthcare": "1997",
    "Alliance Healthcare UK": "1997",
    "Alliance Healthcare France": "1997",
    "Alliance Healthcare Deutschland": "1997",
    "Alliance Healthcare Netherlands": "1997",
    "Alliance Healthcare Spain": "1997",
    "Byram Healthcare": "1968",
    "Edgepark": "1928",
    "Edgepark Medical Supplies": "1928",
    "Owens & Minor": "1882",
    "Owens and Minor": "1882",
    "DKSH": "1865",
    "Zuellig Pharma": "1922",
    "Phoenix Pharmahandel": "1994",
    "PHOENIX Group": "1994",
    "Phoenix Group": "1994",
    "Numotion": "2001",
    "Canadian Hospital Specialties": "1967",
    "CHS Ltd": "1967",
    "Zahrawi Group": "1989",
    "Williams Medical Supplies": "1986",
    "Shields Health Solutions": "2012",
    "Walgreens": "1901",
    "Walgreens Boots Alliance": "2014",
    "Invacare": "1885",
    "Invacare Europe": "1885",
    "Invacare Corporation": "1885",
}

# Drop bad harvest hits (wrong entity / product year / page noise)
REJECT_CACHE = {
    "Health Care Australia",  # staffing agency mismatch risk
    "MediLink",  # Medical Fair Asia noise
    "Carex Health Brands",  # product "since 2008"
    "Numotion",  # override
    "McKesson Canada",  # override
    "Canadian Hospital Specialties",  # override
    "ActivStyle",  # override 1997 not 1986
    "Alliance Healthcare France",  # override
    "Invacare Europe",  # override
    "Newegg Business",  # not a med channel founding signal we trust here
    "Zoro",
    "Rakuten Medical",
    "Shoprite Holdings",
}

REJECT_EVIDENCE = [
    "medical fair asia",
    "tens 7000",
    "copyright",
    "all rights reserved",
    "customer since",
]

NPD = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "null",
}

YEAR_PATTERNS = [
    re.compile(
        r"(?i)\b(?:founded|established|est\.?|incorporated|incorporation|started|since|dating\s+back\s+to|formed|created|establishment)\b"
        r"[^\d]{0,40}\b((?:18|19|20)\d{2})\b"
    ),
    re.compile(
        r"(?i)\b((?:18|19|20)\d{2})\b[^\d]{0,20}\b(?:founded|established|est\.?|incorporated)\b"
    ),
    re.compile(r"(?i)\bsince\s+((?:18|19|20)\d{2})\b"),
]

ABOUT_PATHS = [
    "",
    "/about",
    "/about-us",
    "/about/",
    "/about-us/",
    "/company",
    "/our-story",
    "/history",
    "/who-we-are",
    "/en/about",
    "/pages/about-us",
]


def is_empty_founded(v: str) -> bool:
    s = (v or "").strip().lower()
    return s in NPD or s.startswith("not public")


def valid_year(y: int) -> bool:
    return 1800 <= y <= datetime.now().year


def cache_hit_ok(name: str, res: dict) -> bool:
    if name in REJECT_CACHE:
        return False
    if not res or not res.get("ok"):
        return False
    y = str(res.get("founded") or "").strip()
    if not y.isdigit() or not valid_year(int(y)):
        return False
    ev = (res.get("evidence") or "").lower()
    if any(bad in ev for bad in REJECT_EVIDENCE):
        return False
    if not any(
        k in ev
        for k in (
            "founded",
            "established",
            "est.",
            "since",
            "incorporated",
            "started",
            "formed",
            "created",
            "establishment",
        )
    ):
        return False
    return True


def extract_years(text: str) -> list[tuple[int, str]]:
    if not text:
        return []
    t = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", text)
    t = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    t = re.sub(r"\s+", " ", t)
    hits = []
    for pat in YEAR_PATTERNS:
        for m in pat.finditer(t):
            try:
                y = int(m.group(1))
            except Exception:
                continue
            if valid_year(y):
                ctx = t[max(0, m.start() - 40) : m.end() + 40]
                hits.append((y, ctx.strip()))
    return hits


def pick_year(hits: list[tuple[int, str]]) -> tuple[str, str]:
    keyed = []
    for y, ctx in hits:
        low = ctx.lower()
        if any(bad in low for bad in REJECT_EVIDENCE):
            continue
        if any(
            k in low
            for k in (
                "founded",
                "established",
                "est.",
                "est ",
                "incorporated",
                "since",
                "started",
                "formed",
                "created",
                "dating back",
                "establishment",
            )
        ):
            keyed.append((y, ctx))
    if not keyed:
        return "", ""
    y = min(x[0] for x in keyed)
    ctx = next(c for yy, c in keyed if yy == y)
    return str(y), ctx


def base_url(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    p = urlparse(u)
    return f"{p.scheme}://{p.netloc}"


def fetch(url: str, client: httpx.Client) -> str:
    try:
        r = client.get(url, follow_redirects=True)
        if r.status_code >= 400:
            return ""
        ctype = (r.headers.get("content-type") or "").lower()
        if "html" not in ctype and "text" not in ctype and ctype:
            return ""
        return r.text[:200000]
    except Exception:
        return ""


def harvest_one(company: str, website: str) -> dict:
    root = base_url(website)
    if not root:
        return {"company": company, "website": website, "founded": "", "source": "", "evidence": "", "ok": False}
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml",
        "Accept-Language": "en-US,en;q=0.9",
    }
    best_year = best_ev = best_src = ""
    with httpx.Client(timeout=8.0, headers=headers, verify=False) as client:
        for path in ABOUT_PATHS:
            url = root if path == "" else root.rstrip("/") + path
            html = fetch(url, client)
            if not html or len(html) < 200:
                continue
            y, ctx = pick_year(extract_years(html))
            if y:
                best_year, best_ev, best_src = y, ctx, url
                if path and any(k in path for k in ("about", "histor", "story", "compan")):
                    break
    return {
        "company": company,
        "website": website,
        "founded": best_year,
        "source": best_src,
        "evidence": best_ev[:240],
        "ok": bool(best_year),
    }


def main() -> None:
    import urllib3

    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = XLSX.with_name(XLSX.stem + f"_BEFORE_FOUNDED_{ts}.xlsx")
    shutil.copy2(XLSX, backup)

    cache: dict = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    for name, year in VERIFIED.items():
        prev = cache.get(name, {})
        cache[name] = {
            "company": name,
            "website": prev.get("website", ""),
            "founded": year,
            "source": "agent_web_verified",
            "evidence": f"verified founded {year}",
            "ok": True,
            "verified": True,
        }

    wb = load_workbook(XLSX)
    ws = wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(hdr) if h}

    targets = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        founded = str(ws.cell(r, ic["Founded"]).value or "").strip()
        if is_empty_founded(founded):
            targets.append({"row": r, "company": co, "website": web})

    todo = [
        t
        for t in targets
        if t["company"] not in VERIFIED
        and not cache_hit_ok(t["company"], cache.get(t["company"], {}))
    ]
    print(f"need={len(targets)} harvest={len(todo)}", flush=True)

    done = 0
    with ThreadPoolExecutor(max_workers=24) as ex:
        futs = {ex.submit(harvest_one, t["company"], t["website"]): t for t in todo}
        for fut in as_completed(futs):
            t = futs[fut]
            try:
                res = fut.result()
            except Exception as e:
                res = {
                    "company": t["company"],
                    "website": t["website"],
                    "founded": "",
                    "source": "",
                    "evidence": str(e),
                    "ok": False,
                }
            if not cache.get(t["company"], {}).get("verified"):
                cache[t["company"]] = res
            done += 1
            if done % 30 == 0 or done == len(todo):
                CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
                okc = sum(1 for n, v in cache.items() if cache_hit_ok(n, v) or n in VERIFIED)
                print(f"harvested {done}/{len(todo)} cache_ok~{okc}", flush=True)

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")

    filled_verified = filled_harvest = 0
    still = []
    for t in targets:
        name = t["company"]
        y = ""
        if name in VERIFIED:
            y = VERIFIED[name]
            filled_verified += 1
        elif cache_hit_ok(name, cache.get(name, {})):
            y = str(cache[name].get("founded") or "").strip()
            filled_harvest += 1
        if y and y.isdigit() and valid_year(int(y)):
            ws.cell(t["row"], ic["Founded"]).value = y
        else:
            cur = str(ws.cell(t["row"], ic["Founded"]).value or "").strip()
            if is_empty_founded(cur):
                ws.cell(t["row"], ic["Founded"]).value = "Not publicly disclosed"
            still.append({"company": name, "website": t["website"]})

    wb.save(XLSX)
    wb.close()

    # recount
    wb2 = load_workbook(XLSX, data_only=True)
    ws2 = wb2.active
    hdr2 = [str(c.value or "").strip() for c in ws2[1]]
    fi = hdr2.index("Founded") + 1
    numeric = npd = empty = 0
    for r in range(2, ws2.max_row + 1):
        f = str(ws2.cell(r, fi).value or "").strip()
        if f.isdigit():
            numeric += 1
        elif is_empty_founded(f):
            empty += 1
        else:
            npd += 1
    wb2.close()

    REMAINING.write_text(json.dumps(still, indent=2, ensure_ascii=False), encoding="utf-8")
    audit = {
        "timestamp": ts,
        "backup": str(backup),
        "targets": len(targets),
        "filled_verified": filled_verified,
        "filled_harvest": filled_harvest,
        "numeric_founded_total": numeric,
        "npd_total": npd,
        "empty_total": empty,
        "still_needed": len(still),
        "still_sample": still[:60],
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print("backup", backup, flush=True)
    print(
        f"filled verified={filled_verified} harvest={filled_harvest} | total numeric={numeric} npd={npd} empty={empty}",
        flush=True,
    )
    print("still", len(still), "file", REMAINING, flush=True)


if __name__ == "__main__":
    main()
