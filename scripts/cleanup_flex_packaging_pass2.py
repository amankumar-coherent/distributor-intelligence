"""Second pass: drop residual fakes, fix broken names, merge dupes, blank unverified geo."""
from __future__ import annotations

import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
OUT_TSV = FP.parent / "unique_companies_cleaned.tsv"
REPORT = FP.parent / "cleanup_unique_report_pass2.txt"

DROP_NAME = re.compile(
    r"(?ix)^("
    r"embalagens\b|"
    r"flex\s+embalagens\b|"
    r"flexopack\s+brasil\b|"
    r"pack\s+distribuidora\b|"
    r"envases\s+flexibles\b|"
    r"envases\s+y\s+empaques\s+de\b|"
    r"envases\s+del\s+plata\b|"
    r"pl[aá]sticos\s+del\s+(pac[ií]fico|sur)\b|"
    r"grupo\s+empresarial\s+envases\b|"
    r"indian\s+packaging\b|"
    r"new\s+zealand\s+packaging\b|"
    r"pakistani\s+packaging\b|"
    r"palestinian\s+packaging\b|"
    r"syrian\s+packaging\b|"
    r"middle\s+east\s+packaging\b|"
    r"kuwait\s+packaging\s+trading\b|"
    r"jordan\s+packaging\s+trading\b|"
    r"iranian\s+packaging\b|"
    r"gulf\s+flexible\s+packaging\b|"
    r"flexible\s+packaging\s+supply\b|"
    r"u\.?s\.?\s+packaging\s+supply\b"
    r")"
)

# Manual verified restores / renames (name -> website). Only confirmed.
FORCE_UPSERT = {
    "ProAmpac": "https://www.proampac.com",
}

MERGE_KEYS = {
    "tci": "TCI",
    "transport corporation of india": "TCI",
    "papier mettler": "Papier-Mettler",
    "papier mettler kg": "Papier-Mettler",
    "orora": "Orora",
    "orora limited": "Orora",
    "orora packaging solutions": "Orora",
    "sankoco": "Sanko Co. Ltd.",
    "sanko": "Sanko Co. Ltd.",
}


def keyify(name: str) -> str:
    n = re.sub(r"[^\w\s]", " ", name.lower())
    return re.sub(r"\s+", " ", n).strip()


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = FP.with_name(FP.stem + f"_backup_before_pass2_{stamp}.xlsx")
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    raw = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]

    kept = []
    dropped = []
    for row in raw:
        name = str(row[0] or "").strip()
        # broken strip leftovers
        if name.endswith(" de") or name.endswith(" de.") or name in {"Jindal", "Temkin"}:
            dropped.append((name, "broken_or_ambiguous_name"))
            continue
        if DROP_NAME.search(name):
            dropped.append((name, "residual_fabricated"))
            continue
        kept.append(row)

    # merge groups
    groups: dict[str, list[list]] = defaultdict(list)
    for row in kept:
        k = keyify(str(row[0]))
        # normalize known merge keys
        for mk, canon in MERGE_KEYS.items():
            if k == mk or k.startswith(mk + " "):
                k = keyify(canon)
                row[0] = canon
                break
        if k.startswith("orora"):
            k = "orora"
            row[0] = "Orora"
            row[1] = "https://www.orora.com"
        if k.startswith("papier mettler"):
            k = "papier mettler"
            row[0] = "Papier-Mettler"
            row[1] = "https://www.papier-mettler.com"
        if k in {"tci", "transport corporation of india"} or "transport corporation of india" in k:
            k = "tci"
            row[0] = "TCI"
            row[1] = "https://www.tcil.com"
        groups[k].append(row)

    # upsert verified
    for name, web in FORCE_UPSERT.items():
        k = keyify(name)
        if k not in groups:
            blank = [None] * (ws.max_column)
            blank[0] = name
            blank[1] = web
            groups[k] = [blank]

    merged = []
    merge_log = []
    for k, items in groups.items():
        best = items[0]
        for r in items[1:]:
            for i, v in enumerate(r):
                if i >= len(best):
                    break
                if (best[i] in (None, "")) and v not in (None, ""):
                    best[i] = v
        # blank unverified geo / codes (were mass-filled)
        for idx in (3, 4, 5, 18, 19):  # HQ, Continent, Ops presence, CC, RC
            if idx < len(best):
                best[idx] = None
        # blank still-generic specialty
        if len(best) > 9 and best[9] and re.match(
            r"(?i)^(distributes|imports and distributes)\s+flexible packaging",
            str(best[9]),
        ):
            best[9] = None
        merged.append(best)
        if len(items) > 1:
            merge_log.append(f"{best[0]} <= {[x[0] for x in items]}")

    merged.sort(key=lambda r: str(r[0] or "").lower())

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in merged:
        # pad
        while len(r) < ws.max_column:
            r.append(None)
        ws.append(r[: ws.max_column])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(FP)

    with OUT_TSV.open("w", encoding="utf-8") as f:
        f.write("company\twebsite\n")
        for r in merged:
            f.write(f"{r[0]}\t{r[1] or ''}\n")

    lines = [
        f"backup: {backup}",
        f"before: {len(raw)}",
        f"dropped: {len(dropped)}",
        f"after_unique: {len(merged)}",
        "",
        "MERGES:",
        *merge_log,
        "",
        "DROPS:",
        *[f"{n}\t{why}" for n, why in dropped],
        "",
        "UNIQUE:",
        *[f"{r[0]}\t{r[1]}" for r in merged],
    ]
    REPORT.write_text("\n".join(lines), encoding="utf-8")
    print(f"UNIQUE={len(merged)} dropped={len(dropped)} merges={len(merge_log)}")
    for r in merged:
        print(f"{r[0]}\t{r[1]}")


if __name__ == "__main__":
    main()
