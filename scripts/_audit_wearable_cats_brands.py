# -*- coding: utf-8 -*-
from pathlib import Path
import openpyxl

p = Path(r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global\global_wearable_medical_devices_market_global_FINAL.xlsx")
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active
h = [c.value for c in ws[1]]
ic = {str(x): i + 1 for i, x in enumerate(h) if x}

def emptyish(v):
    s = str(v or "").strip().lower()
    return s in {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd"} or s.startswith("verify:")

def weak_brands(v):
    s = str(v or "").strip()
    if emptyish(s):
        return True
    low = s.lower()
    generics = (
        "multi-vendor pharmaceutical principals, others",
        "multi-brand pharmacy assortment, others",
        "multi-principal medtech brands via addlife subsidiaries, others",
        "various manufacturers, others",
        "not publicly disclosed",
    )
    if low in generics:
        return True
    if low.startswith("multi-vendor") or low.startswith("multi-brand pharmacy"):
        # no concrete brand tokens
        parts = [p.strip() for p in s.replace(";", ",").split(",") if p.strip()]
        real = [p for p in parts if p.lower() not in {"others", "etc", "etc."} and "multi-" not in p.lower() and "assortment" not in p.lower() and "principals" not in p.lower()]
        return len(real) == 0
    return False

print("EMPTY Core or Specialty:")
for r in range(2, ws.max_row + 1):
    co = ws.cell(r, ic["Company"]).value
    core = ws.cell(r, ic["Core Categories"]).value
    spec = ws.cell(r, ic["Specialty Focus"]).value
    brands = ws.cell(r, ic["Key Brands Represented"]).value
    web = ws.cell(r, ic["Website"]).value
    if emptyish(core) or emptyish(spec):
        print("---", co)
        print("  core:", core)
        print("  spec:", (str(spec or ""))[:100])
        print("  brands:", (str(brands or ""))[:100])
        print("  web:", web)

empty = weak = ok = 0
empty_list = []
weak_list = []
for r in range(2, ws.max_row + 1):
    co = ws.cell(r, ic["Company"]).value
    brands = ws.cell(r, ic["Key Brands Represented"]).value
    web = ws.cell(r, ic["Website"]).value
    if emptyish(brands):
        empty += 1
        empty_list.append((co, web))
    elif weak_brands(brands):
        weak += 1
        weak_list.append((co, brands, web))
    else:
        ok += 1
print("\nbrands empty", empty, "weak", weak, "ok", ok)
print("\nEMPTY brands sample:")
for co, web in empty_list[:30]:
    print(" ", co, "|", web)
print("\nWEAK brands sample:")
for co, b, web in weak_list[:25]:
    print(" ", co, "=>", b)
