# -*- coding: utf-8 -*-
"""Retry Key Brands Represented via Google AI scraper + DeepSeek (evidence-gated)."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

OUT = ROOT / "output/web_expand/global_wearable_medical_devices_market_global"
XLSX = OUT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
AUDIT = OUT / "wearable_brands_gai_audit.json"
SUMMARY = OUT / "wearable_brands_gai_summary.json"
NPD = "Not publicly disclosed"
NPD_SET = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd"}


def emptyish(v: object) -> bool:
    return str(v or "").strip().lower() in NPD_SET


def ensure_others(brands: str) -> str:
    t = (brands or "").strip().rstrip(",")
    if not t or emptyish(t):
        return NPD
    if re.search(r",\s*others\s*$", t, re.I):
        return re.sub(r",\s*others\s*$", ", others", t, flags=re.I)
    return t + ", others"


def clean_brands(raw: str) -> str:
    if emptyish(raw) or str(raw).strip().lower() in {"none", "null", "unknown", "[]"}:
        return NPD
    parts = re.split(r"[,;/|]+", str(raw))
    out: list[str] = []
    seen: set[str] = set()
    junk = {
        "others",
        "etc",
        "various",
        "multiple",
        "n/a",
        "none",
        "unknown",
        "brands",
        "brand",
        "principals",
        "suppliers",
        "portfolio",
        "multi-vendor",
        "multi-brand",
        "assortment",
        "many",
        "numerous",
    }
    for p in parts:
        b = re.sub(r"\s+", " ", p).strip(" .\"'")
        if not b or len(b) < 2 or b.lower() in junk:
            continue
        if re.search(r"multi-vendor|multi-brand|assortment|not publicly|principals|thousands", b, re.I):
            continue
        if len(b) > 55:
            continue
        key = b.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(b)
    if not out:
        return NPD
    return ensure_others(", ".join(out[:14]))


def brands_in_evidence(brands: str, evidence: str) -> str:
    if emptyish(brands) or brands == NPD:
        return NPD
    ev = (evidence or "").lower()
    kept: list[str] = []
    for b in [x.strip() for x in brands.replace(";", ",").split(",")]:
        if not b or b.lower() == "others":
            continue
        token = re.sub(r"[^A-Za-z0-9&.\\- ]", "", b).strip().lower()
        if len(token) < 2:
            continue
        if token in ev or token.replace(" ", "") in ev.replace(" ", ""):
            kept.append(b)
            continue
        first = token.split()[0] if token.split() else ""
        if len(first) >= 3 and first in ev:
            kept.append(b)
    if not kept:
        return NPD
    return ensure_others(", ".join(kept[:14]))


def ds_client() -> tuple[OpenAI, str]:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    return OpenAI(api_key=key, base_url=base, timeout=90.0), model


def parse_json(text: str) -> dict:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    try:
        return json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


async def ask_brands(company: str, website: str, hq: str, core: str) -> str:
    geo = f" ({hq})" if hq else ""
    q1 = f"key brands represented by {company}{geo} distributor partners principals manufacturers"
    q2 = f"{company} authorized distributor brands products portfolio {website}"
    chunks: list[str] = []
    for q in (q1, q2):
        data = await google_ai_ask(q, llm_clean=True)
        md = str(data.get("markdown") or "").strip()
        if md and len(md) >= 40:
            chunks.append(f"Q: {q}\n{md}")
        await asyncio.sleep(0.5)
        if len("\n".join(chunks)) > 3500:
            break
    return "\n\n----\n\n".join(chunks)[:10000]


def deepseek_brands(client: OpenAI, model: str, company: str, evidence: str) -> dict:
    system = (
        "Extract manufacturer/brand names this distributor officially distributes or represents. "
        "EVIDENCE ONLY. Never invent.\n"
        'JSON: {"found":bool,"brands":"comma-separated brand names or empty","confidence":"high|medium|low|none"}\n'
        "If evidence only says they carry many brands without naming them, found=false."
    )
    user = f"Company: {company}\n\nEVIDENCE:\n{(evidence or '(none)')[:8500]}"
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=250,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        return parse_json(resp.choices[0].message.content or "{}")
    except Exception as e:  # noqa: BLE001
        return {"found": False, "confidence": "none", "error": str(e)[:120]}


# Prefer device/specialty distributors over pure full-line pharmacy wholesalers for GAI pass
PRIORITY_HINTS = (
    "device",
    "aed",
    "diagnostic",
    "monitoring",
    "surgical",
    "medical equipment",
    "medical supplies",
    "hospital",
    "ems",
    "defibrillator",
    "imaging",
    "laboratory",
    "home care",
    "dme",
    "consumable",
)


async def main_async() -> None:
    client, model = ds_client()
    ok = await google_ai_ping()
    print(f"GoogleAI ping={ok} DeepSeek=yes model={model}", flush=True)
    if not ok:
        print("ERROR: Google AI scraper not up", flush=True)
        return

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_BRANDS_GAI_{stamp}.xlsx")
    wb0 = openpyxl.load_workbook(XLSX)
    wb0.save(bak)
    wb0.close()
    print("backup", bak.name, flush=True)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    gaps: list[int] = []
    for r in range(2, ws.max_row + 1):
        brands = ws.cell(r, ic["Key Brands Represented"]).value
        if not emptyish(brands):
            continue
        core = str(ws.cell(r, ic["Core Categories"]).value or "").lower()
        spec = str(ws.cell(r, ic["Specialty Focus"]).value or "").lower()
        blob = core + " " + spec
        # prioritize specialty/device distributors; still include others but later
        score = sum(1 for h in PRIORITY_HINTS if h in blob)
        gaps.append((score, r))
    gaps.sort(reverse=True)
    rows = [r for _, r in gaps]
    # Cap this pass to keep runtime reasonable; remaining stay NPD
    MAX = int(os.getenv("WEARABLE_BRANDS_GAI_MAX", "80"))
    rows = rows[:MAX]
    print(f"NPD brand gaps prioritized={len(gaps)} processing={len(rows)}", flush=True)

    filled = left = 0
    audit: list[dict] = []
    for i, r in enumerate(rows, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        hq = str(ws.cell(r, ic["Headquarters"]).value or "").strip() if "Headquarters" in ic else ""
        core = str(ws.cell(r, ic["Core Categories"]).value or "").strip()
        print(f"[{i}/{len(rows)}] GoogleAI brands {co} ...", flush=True)
        evidence = await ask_brands(co, web, hq, core)
        data = deepseek_brands(client, model, co, evidence)
        conf = str(data.get("confidence") or "").lower()
        brands = clean_brands(str(data.get("brands") or ""))
        if brands != NPD:
            brands = brands_in_evidence(brands, evidence)
        if data.get("found") and conf in {"high", "medium"} and brands != NPD:
            ws.cell(r, ic["Key Brands Represented"]).value = brands
            filled += 1
            audit.append({"company": co, "brands": brands, "source": "google_ai_deepseek", "conf": conf})
            print(f"  FILL {brands}", flush=True)
        else:
            left += 1
            audit.append({"company": co, "brands": NPD, "source": "still_npd", "conf": conf, "preview": evidence[:300]})
            print("  NPD", flush=True)
        time.sleep(0.15)
        if i % 15 == 0:
            wb.save(XLSX)
            print("  checkpoint", flush=True)

    wb.save(XLSX)
    wb.close()

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    bi = headers.index("Key Brands Represented")
    ci = headers.index("Core Categories")
    si = headers.index("Specialty Focus")
    bf = be = 0
    cf = ce = sf = se = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        b = str(row[bi] or "").strip()
        c = str(row[ci] or "").strip()
        s = str(row[si] or "").strip()
        if emptyish(b):
            be += 1
        else:
            bf += 1
        if emptyish(c):
            ce += 1
        else:
            cf += 1
        if emptyish(s):
            se += 1
        else:
            sf += 1
    wb.close()

    stats = {
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "backup": str(bak),
        "processed": len(rows),
        "filled_brands": filled,
        "still_npd_brands_this_pass": left,
        "Key Brands Represented": {"filled": bf, "empty_or_npd": be},
        "Core Categories": {"filled": cf, "empty_or_npd": ce},
        "Specialty Focus": {"filled": sf, "empty_or_npd": se},
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    SUMMARY.write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(stats, indent=2), flush=True)


def main() -> None:
    asyncio.run(main_async())


if __name__ == "__main__":
    main()
