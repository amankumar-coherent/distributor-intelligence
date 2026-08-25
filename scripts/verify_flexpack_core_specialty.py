# -*- coding: utf-8 -*-
"""
Verify / fill Core Categories + Specialty Focus for flexpack FINAL.

Rules:
- Core Categories: real product/service categories, semicolon-separated
- Specialty Focus: detailed functional description of what the channel does
- NO country/region at end of Specialty Focus
- Never invent; only evidence from web / hand overrides
- DeepSeek used only to extract/polish from provided evidence text
"""
from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "core_specialty_verify_audit.json"
CKPT = FP.parent / "core_specialty_ckpt.json"

load_dotenv(r"D:\vendor-intel\.env")

# Hand-verified from official sites / product pages (this pass)
HAND: dict[str, dict[str, str]] = {
    "Uline": {
        "core": "Shipping Boxes; Shipping Supplies; Packaging Materials; Material Handling; Warehouse Equipment; Janitorial & Safety",
        "focus": "B2B catalog distributor of shipping boxes, packaging materials, material-handling equipment, warehouse fixtures, and facility supplies, emphasizing in-stock breadth and rapid fulfillment from regional distribution centers.",
    },
    "Associated Bag": {
        "core": "Plastic Bags & Pouches; Corrugated Boxes; Bubble & Cushioning; Stretch Wrap & Strapping; Tape & Adhesives; Liners & Covers; Static Control",
        "focus": "Wholesale manufacturer-distributor of bags, boxes, cushioning, stretch wrap, tape, and related shipping and workplace packaging, with stock and custom printing options for commercial customers.",
    },
    "Bunzl": {
        "core": "Foodservice Packaging; Retail Packaging; Cleaning & Hygiene; Safety Consumables; Industrial Packaging",
        "focus": "Specialist international distributor of foodservice and retail packaging, cleaning and hygiene products, and safety consumables to grocery, food processing, foodservice, and other business customers.",
    },
    "RAJA": {
        "core": "Cartons & Boxes; Protective Packaging; Films & Stretch Wrap; Tapes; Mailers; Packaging Equipment; Workplace Consumables",
        "focus": "Multi-channel B2B distributor of packaging and workplace supplies offering cartons, protective materials, films, tapes, mailers, and packaging equipment with catalog and e-commerce fulfillment across Europe.",
    },
    "Antalis": {
        "core": "Papers & Boards; Packaging Supplies; Visual Communications; Postal Packaging; Cushioning & Void Fill; Tapes & Films",
        "focus": "Merchant distributor of papers, packaging supplies, and visual communication materials to print, retail, office, industrial packaging, and publishing customers, including boxes, tapes, cushioning, postal products, and packaging papers.",
    },
    "Lyreco": {
        "core": "Office Supplies; Workplace Equipment; Packaging Consumables; Facility & Hygiene Supplies; PPE",
        "focus": "B2B workplace solutions distributor supplying office products, facility and hygiene supplies, PPE, and packaging consumables to businesses through multi-country catalog and e-commerce channels.",
    },
    "Manutan": {
        "core": "Industrial Equipment; Warehouse & Handling; Packaging Supplies; Office & Workshop Furnishings; Safety Equipment",
        "focus": "European B2B e-commerce distributor of industrial equipment, warehouse and handling products, packaging supplies, and workplace furnishings for enterprises and local authorities.",
    },
    "Veritiv": {
        "core": "Specialty Packaging; Facility Solutions; Print & Publishing Products; Packaging Equipment",
        "focus": "B2B specialty packaging and facility solutions distributor providing packaging materials, equipment, and supply-chain programs to manufacturing, healthcare, food and beverage, logistics, and public-sector customers.",
    },
    "Imperial Dade": {
        "core": "Foodservice Packaging; Janitorial & Sanitation; Industrial Packaging; Facility Supplies",
        "focus": "North American distributor of foodservice disposables, janitorial and sanitation products, and industrial packaging to foodservice, healthcare, education, and commercial facility customers.",
    },
    "BradyPLUS": {
        "core": "Janitorial & Sanitation; Foodservice Packaging; Industrial Packaging; Facility Maintenance Supplies",
        "focus": "Distributor of JanSan, foodservice packaging, and industrial packaging supplies with regional fulfillment supporting commercial facility and foodservice operations (now part of Imperial Brady).",
    },
    "Macfarlane Packaging": {
        "core": "Protective Packaging; Corrugated & Cartons; Foam & Cushioning; Industrial Packaging; Stretch Wrap & Tape; Custom Packaging Design",
        "focus": "UK/Ireland protective packaging distributor and manufacturer supplying stock packaging materials plus designed industrial packaging (corrugated, timber, foam, returnable) for transit protection across commercial and industrial customers.",
    },
    "GWP Group": {
        "core": "Protective Packaging; Corrugated Packaging; Foam Packaging; Conductive Packaging; Coatings",
        "focus": "Protective packaging manufacturer and distributor producing corrugated, foam, Correx, anti-static, and related specialist packaging solutions, operating as part of Macfarlane Group.",
    },
    "Ratioform": {
        "core": "Cartons; Void Fill & Cushioning; Postal Packaging; Films & Tapes; Packaging Machines; Shipping Labels",
        "focus": "B2B mail-order and e-commerce packaging specialist supplying cartons, cushioning, postal packaging, films, tapes, and packaging equipment for warehouse and shipping operations under the TAKKT/kaiserkraft network.",
    },
    "Davpack": {
        "core": "Cardboard Boxes; Postal Packaging; Protective Packaging; Polythene; Retail Packaging; Packaging Consumables",
        "focus": "UK mail-order packaging supplier offering in-stock cardboard boxes, postal and protective packaging, polythene, and related consumables with same-day dispatch capability as part of TAKKT packaging brands.",
    },
    "NNZ": {
        "core": "Produce Retail Packaging; Produce Transport Packaging; Industrial Big Bags; Nets & Flexible Packaging",
        "focus": "Independent packaging advisor-distributor specializing in fresh-produce retail and transport packaging plus industrial packaging such as big bags and nets for food, feed, and non-food markets.",
    },
    "Shorr Packaging": {
        "core": "Packaging Materials; Packaging Equipment; Stretch Film; Corrugated; Protective Packaging; Automation Support",
        "focus": "Employee-owned full-service packaging distributor supplying materials and equipment with consultative packaging programs to optimize customer packing lines and supply chains.",
    },
    "Richards Packaging": {
        "core": "Rigid Containers; Plastic Bottles & Closures; Glass Packaging; Pharma & Healthcare Packaging; Cosmetic Packaging",
        "focus": "North American rigid packaging distributor supplying plastic, glass, and related containers and closures to healthcare, cosmetic, food and beverage, and other product marketers, with pharmacy/medical device adjacent offerings.",
    },
    "All American Containers": {
        "core": "Glass Containers; Plastic Containers; Metal Containers; Closures & Dispensers; Rigid Packaging",
        "focus": "Rigid packaging distributor of glass, plastic, and metal containers with closures and dispensers for food, beverage, personal care, and industrial packers (now under TricorBraun).",
    },
    "O.Berk": {
        "core": "Plastic Bottles; Glass Bottles & Jars; Metal Containers; Closures; Value-Added Packaging Services",
        "focus": "Rigid packaging distributor serving beauty, personal care, pharmaceutical, nutraceutical, household, and food markets with containers, closures, and related packaging services (Berlin Packaging company).",
    },
    "Novvia Group": {
        "core": "Rigid Metal & Plastic Containers; Drums & Pails; Bottles & Cans; Life Sciences Packaging",
        "focus": "Platform of regional rigid packaging distributors supplying cans, pails, drums, bottles, and specialty life-sciences packaging to industrial, chemical, food, and healthcare customers.",
    },
    "MonotaRO": {
        "core": "MRO Supplies; Industrial Consumables; Packaging & Packing Materials; Logistics & Storage Products; Safety Products",
        "focus": "B2B e-commerce platform for manufacturing and construction MRO products, including packing tapes, bags, wrapping materials, logistics/storage items, and related industrial consumables.",
    },
    "Askul": {
        "core": "Office Supplies; Business Consumables; Shipping & Packaging Supplies; Facility Products; Medical & Living Goods",
        "focus": "Japanese B2B e-commerce distributor of office and business supplies including shipping packaging, retail bags, facility products, and related consumables with nationwide next-day logistics.",
    },
    "Fastenal": {
        "core": "Fasteners; Industrial MRO; Safety Products; Packaging & Shipping Supplies; Tools",
        "focus": "Industrial distributor operating branch and vending networks for fasteners, MRO, safety, and packaging/shipping supplies sold to manufacturing and construction customers.",
    },
    "The Home Depot": {
        "core": "Building Materials; Tools; Packaging & Moving Supplies; Home Improvement Retail",
        "focus": "Home-improvement retailer supplying packaging and moving materials (boxes, packing paper, tape, stretch wrap) alongside tools and building products through stores and e-commerce.",
    },
    "Staples": {
        "core": "Office Supplies; Shipping Supplies; Packaging Materials; Business Services",
        "focus": "Office products retailer and B2B supplier of shipping and packaging materials, mailroom supplies, and workplace consumables through retail and delivery channels.",
    },
    "Quill": {
        "core": "Office Supplies; Shipping & Packaging Supplies; Breakroom Products; Business Essentials",
        "focus": "B2B catalog and e-commerce supplier of office, shipping, and packaging supplies serving small and mid-size business customers as a Staples company.",
    },
    "WebstaurantStore": {
        "core": "Foodservice Equipment; Foodservice Packaging; Shipping Boxes; Stretch Wrap; Poly Bags; Packaging Equipment",
        "focus": "Online foodservice equipment and supplies distributor offering restaurant packaging, shipping boxes, stretch wrap, poly bags, and related packaging equipment for foodservice operators.",
    },
    "Zoro Tools": {
        "core": "Industrial Supplies; MRO; Packaging Supplies; Tools; Safety Products",
        "focus": "Online industrial and MRO marketplace (Grainger) supplying tools, safety, and packaging/shipping products to business customers.",
    },
    "McMaster-Carr": {
        "core": "Industrial Hardware; Raw Materials; Packaging Components; Maintenance Supplies; Tools",
        "focus": "Private industrial supplier catalog covering hardware, materials, packaging components, and maintenance products with fast fulfillment for engineering and plant customers.",
    },
    "Ahlsell": {
        "core": "Installation Products; Tools; Personal Protective Equipment; Industrial Supplies; Technical Products",
        "focus": "Nordic B2B distributor of installation, construction, and industrial technical products including tools, PPE, and related supplies for professional trades and industry.",
    },
    "Blackwoods": {
        "core": "Industrial MRO; Safety Products; Tools; Packaging & Consumables; Engineering Supplies",
        "focus": "Wesfarmers industrial and safety distributor supplying MRO, PPE, tools, and related consumables including packaging items to enterprise and trade customers.",
    },
    "Officeworks": {
        "core": "Office Supplies; Technology; Print & Copy; Packaging & Shipping Supplies; Furniture",
        "focus": "Australian office products retailer offering packaging and shipping supplies alongside stationery, technology, print services, and furniture for home and business customers.",
    },
    "Viking Direct": {
        "core": "Office Supplies; Workplace Furniture; Packaging Consumables; Print Supplies; Facility Products",
        "focus": "European B2B e-commerce office and workplace supplies distributor (RAJA Office) offering stationery, furniture, packaging consumables, and facility products to SME customers.",
    },
    "JPG": {
        "core": "Office Supplies; Workplace Equipment; Packaging Consumables; Facility Supplies",
        "focus": "French B2B online reseller of office supplies and workplace equipment (Mondoffice France / RAJA Office) including packaging and facility consumables for business customers.",
    },
    "Kalamazoo": {
        "core": "Office Supplies; Workplace Equipment; Packaging Consumables; Facility Supplies",
        "focus": "Spanish B2B reseller of office and workplace supplies (Mondoffice Spain / RAJA Office) serving businesses via e-commerce, catalog, and field sales.",
    },
    "Mondoffice": {
        "core": "Office Supplies; Workplace Equipment; Packaging Consumables; Facility Supplies",
        "focus": "RAJA Office brand distributing office supplies, workplace equipment, and related packaging/facility consumables to business customers via catalog and e-commerce channels.",
    },
    "Hoffmann Verpackung": {
        "core": "Protective Packaging; Cartons; Filling Materials; Films; Tapes; Custom Packaging Design",
        "focus": "Family-owned packaging merchant and producer supplying protective and shipping packaging materials with design/sample support for industrial and trade customers.",
    },
    "Papier-Mettler": {
        "core": "Paper Carrier Bags; Plastic Carrier Bags; Flexible Films; Shrink & Stretch Films; Foodservice Packaging; E-commerce Mailers",
        "focus": "Manufacturer and supplier of paper and plastic service packaging and flexible films for retail, foodservice, and e-commerce applications, with international sales offices.",
    },
    "Brangs + Heinrich": {
        "core": "Packaging Materials; Paper & Film Converting; Industrial Packaging; Protective Packaging",
        "focus": "Family-owned packaging materials trader with in-house paper and film converting, supplying industrial and trade customers with films, papers, and protective packaging.",
    },
    "Igepa Group": {
        "core": "Graphic Papers; Packaging Materials; Visual Communication Media; Print Accessories",
        "focus": "European specialist wholesale group distributing paper, packaging, and advertising-technology products to printers, packagers, and visual communication businesses.",
    },
    "Papyrus": {
        "core": "Graphic Papers; Office Papers; Industrial Packaging; Facility Supplies; Visual Communication Materials",
        "focus": "European B2B merchant (OptiGroup) distributing paper, packaging materials, and facility supplies to printers, publishers, industrial, and service customers.",
    },
    "Inapa": {
        "core": "Paper; Packaging; Visual Communication; Logistics Services",
        "focus": "Paper and packaging merchant group distributing graphic papers, packaging products, and visual communication materials with related logistics services.",
    },
    "Spicers": {
        "core": "Paper; Packaging; Sign & Display Media; Graphic Supplies",
        "focus": "KPP Oceania merchant distributing paper, packaging, and visual communication products to print, packaging, and signage customers.",
    },
    "Spicers Paper": {
        "core": "Paper; Packaging; Graphic Arts Supplies; Wide-Format Media; Facility Supplies",
        "focus": "Canadian paper and packaging merchant (CNG) distributing commercial print papers, packaging materials, and graphic supplies through a national warehouse network.",
    },
    "Signet": {
        "core": "Packaging Supplies; Protective Papers; Marking Products; Business Supplies",
        "focus": "Spicers Group packaging and marking supplies distributor serving industrial and commercial customers with packaging, protective papers, and related business supplies.",
    },
    "Cospak": {
        "core": "Rigid Packaging; Glass & Plastic Containers; Closures; Food & Beverage Packaging; Custom Packaging Design",
        "focus": "Australasian packaging distributor and manufacturer (SMYA) supplying containers and packaging solutions for wine, food, beverage, and related markets with design and warehousing support.",
    },
    "Abco Kovex": {
        "core": "Stretch Film; Strapping; Protective Packaging; Edge Boards; Packaging Machinery",
        "focus": "Bunzl packaging company manufacturing and distributing stretch films, strapping, protective packaging, and packaging machinery to industrial and commercial customers.",
    },
    "Hisco": {
        "core": "Specialty Chemicals; Adhesives & Tapes; Electronic Assembly Materials; ESD & Thermal Materials; Packaging Consumables",
        "focus": "Specialty industrial distributor (DSG) supplying adhesives, chemicals, tapes, and electronics-assembly materials with value-added converting, VMI, and specialty warehousing including cold storage.",
    },
    "Crawford Packaging": {
        "core": "Stretch Film; Shrink Film; Protective Packaging; Tape; Packaging Equipment; Automation",
        "focus": "North American packaging distributor specializing in stretch and shrink films, protective packaging, tape, and packaging equipment/automation with technical service support.",
    },
    "Polyair": {
        "core": "Bubble Cushioning; Foam Packaging; Mailers; Air Pillow Systems; Protective Packaging",
        "focus": "Protective packaging manufacturer (IPG/Clearlake) producing bubble, foam, mailers, and air-pillow systems for e-commerce and industrial shipping protection.",
    },
    "Atlantic Packaging Products": {
        "core": "Corrugated Packaging; Paper Bags; Retail Packaging & Displays; Packaging Supplies; Recycling",
        "focus": "Integrated packaging manufacturer-distributor of corrugated products, paper bags, retail packaging, and related supplies, with converting and distribution operations supporting retail and industrial customers.",
    },
    "Transpack": {
        "core": "Polythene Bags; Cellophane Bags; Postal & Mailing Packaging; Cardboard Boxes; Warehouse Packaging; Pallet Wrap; Tapes; Eco-friendly & Food Packaging",
        "focus": "Family-run wholesale packaging supplier offering polythene and cellophane bags, postal packaging, boxes, warehouse protective materials, pallet wrap, tapes, and eco-friendly/food packaging with bespoke and bulk options for trade and business customers.",
    },
    "Serwo": {
        "core": "Industrial Films & Bags; Stretch Film; Adhesive Tapes; Strapping; Carrier Bags; Agricultural Packaging; Automotive Workshop Protection",
        "focus": "German producer-wholesaler of industrial packaging (films, bags, tapes, strapping, and specialty sector packaging) plus automotive workshop protection products, stocking standard ranges for short-lead delivery and offering custom printing on many items.",
    },
    "Box Brothers": {
        "core": "Moving Boxes; Shipping Boxes; Packaging Supplies; Custom Crating; Professional Packing; Shipping Services",
        "focus": "Packing and shipping retailer network supplying moving and shipping boxes and packaging materials, with professional packing, custom crating, and carrier shipping services for consumers and small businesses.",
    },
    "Plastic Bag Source": {
        "core": "Custom Printed Plastic Bags; Laminated Bags; Reusable Bags; T-Shirt Bags; Pouch Bags; Trade Show Bags; Flexible Films",
        "focus": "Manufacturer-importer-supplier focused on custom printed plastic, laminated, and reusable bags plus pouch and trade-show bags, supporting retailers and industrial customers with designed flexible packaging.",
    },
    "The Packaging Centre": {
        "core": "Food & Service Packaging; Glass Jars & Bottles; Paper & Cardboard Packaging; Gift & Retail Packaging; Industrial Packaging; Everyday Essentials",
        "focus": "Irish packaging merchant supplying foodservice packaging, glass jars and bottles, paper and cardboard packaging, gift and retail packaging, and industrial packaging solutions with stocked essentials for trade customers.",
    },
    "Packaging Resources": {
        "core": "Custom Polybags; Films; Rigid Food Containers; Automated Closures",
        "focus": "Packaging supplier providing custom polybags, films, rigid food containers, and automated closure packaging products for commercial packaging and processing customers.",
    },
    "Nashville Wraps": {
        "core": "Gift Wrap; Shopping Bags; Tissue Paper; Ribbon & Bows; Gift Boxes; Cellophane; Retail Shipping Supplies; Shrink Wrap",
        "focus": "Wholesale gift, gourmet, and retail packaging supplier offering shopping bags, gift wrap, tissue, ribbon, boxes, cellophane, and related shipping supplies, with custom printing and fast fulfillment of in-stock lines for retailers and small brands.",
    },
    "Shippers Supply": {
        "core": "Corrugated Boxes; Poly Bags; Stretch & Shrink Film; Tapes; Mailers; Protective Foam; Packaging Equipment; Custom Packaging Design",
        "focus": "Strategic packaging distributor supplying stock industrial packaging materials plus custom packaging design, foam fabrication, and packaging equipment/automation support to manufacturing, food and beverage, e-commerce, and medical customers.",
    },
    "Acme Paper & Supply": {
        "core": "Foodservice Packaging; Retail Packaging; Industrial Packaging; Janitorial Equipment & Cleaning Supplies; Restaurant Equipment & Smallwares",
        "focus": "Wholesale packaging distributor supplying foodservice, retail, and industrial packaging plus janitorial and restaurant equipment/smallwares, with creative packaging solutions and regional delivery for commercial customers.",
    },
    "BOX Partners": {
        "core": "Corrugated Boxes; Shipping Supplies; Packaging Materials; Moving Boxes",
        "focus": "Wholesale distributor of corrugated boxes and shipping/packaging supplies serving business customers with large in-stock assortments and volume fulfillment.",
    },
    "Global Industrial": {
        "core": "Material Handling; Storage & Shelving; Packaging Supplies; Facility Maintenance; Safety Products",
        "focus": "Industrial supplies e-commerce distributor offering material handling, storage, packaging, facility maintenance, and safety products to business and public-sector customers.",
    },
    "HD Supply": {
        "core": "Facilities Maintenance; Construction Supplies; Packaging & MRO Consumables",
        "focus": "Home Depot company distributing facilities maintenance and construction supplies including packaging and MRO consumables to professional customers.",
    },
    "Kaufman Container": {
        "core": "Plastic Containers; Glass Containers; Closures; Decorating & Decorating Services",
        "focus": "Privately held rigid packaging supplier and decorator providing containers, closures, and value-added decorating services to brand owners.",
    },
    "SKS Bottle & Packaging": {
        "core": "Glass Bottles & Jars; Plastic Bottles & Jars; Metal Containers; Closures; Custom Decorating",
        "focus": "Online and wholesale container distributor (Pipeline Packaging) specializing in glass, plastic, and metal containers with closures and decorating services for personal care, nutraceutical, and craft markets.",
    },
    "The Cary Company": {
        "core": "Containers; Packaging; Spill Control; Filtration; Industrial Products",
        "focus": "Long-established packaging and industrial products supplier offering containers, packaging, spill control, filtration, and related sourcing/warehousing services to professional customers.",
    },
    "Unicorr Packaging": {
        "core": "Corrugated Containers; Retail Displays; Protective Packaging; Foam; Wood Crates & Pallets",
        "focus": "Family-owned corrugated converter and protective packaging manufacturer-distributor serving Northeast customers with custom boxes, displays, foam, and related packaging.",
    },
    "Allied Packaging": {
        "core": "Flexible Packaging; Industrial Packaging Supplies; Protective Packaging",
        "focus": "Packaging distributor supplying flexible and industrial packaging materials to commercial and industrial customers, with a focus on stock packaging supplies for trade buyers.",
    },
    "Packaging House": {
        "core": "Packaging Materials; Shipping Supplies; Industrial Packaging",
        "focus": "New Zealand packaging materials supplier providing shipping and industrial packaging products to trade and business customers.",
    },
    "Waltons": {
        "core": "Stationery; Office Supplies; Packaging & Envelopes; Office Furniture; Technology Accessories",
        "focus": "Bidvest retail and commercial office products chain supplying stationery, packaging/envelopes, furniture, and related workplace products through stores and e-commerce.",
    },
    "Office Depot": {
        "core": "Office Supplies; Technology; Furniture; Packaging & Shipping Supplies",
        "focus": "Office products retail and B2B brand (Grupo Gigante in Mexico) offering stationery, technology, furniture, and packaging/shipping supplies.",
    },
    "Moglix": {
        "core": "Industrial MRO; Safety Products; Packaging & Consumables; Electrical & Tools; Facility Supplies",
        "focus": "B2B industrial e-commerce platform supplying MRO, safety, packaging consumables, and related industrial products to manufacturing and enterprise buyers.",
    },
    "IndustryBuying": {
        "core": "Industrial Supplies; MRO; Packaging Materials; Safety Products; Tools",
        "focus": "Indian B2B industrial marketplace (MonotaRO group) offering MRO, packaging materials, safety products, and tools to SME and industrial buyers.",
    },
    "Gerdmans": {
        "core": "Office Furniture; Warehouse Storage; Industrial Furnishings; Packaging-Related Workplace Equipment",
        "focus": "Nordic B2B catalog supplier (TAKKT) of office furniture, warehouse shelving, and industrial workplace equipment supporting packing and storage environments.",
    },
    "Seow Khim Polythene": {
        "core": "Disposable Food Packaging; Party Supplies; Plastic Packaging; Retail Packaging",
        "focus": "Singapore manufacturer-distributor of disposable food packaging and party supplies with retail outlets and export of packaging products.",
    },
    "Syspex": {
        "core": "Industrial Packaging; Stretch Film; Protective Packaging; Packaging Equipment",
        "focus": "Southeast Asian packaging distributor supplying stretch film, protective packaging, and related industrial packaging equipment and materials.",
    },
    "Tantraco": {
        "core": "Shrink Packaging Machinery; Packaging Equipment; Food Processing Equipment; Labeling & Coding",
        "focus": "Singapore packaging machinery distributor designing and supplying shrink-bundling and related packaging/food-processing equipment with regional manufacturing support.",
    },
    "FlexPRO": {
        "core": "Flexible Packaging; Pre-made Pouches; Plastic Containers; Printed Soft Packaging",
        "focus": "Balkan sales and logistics partner for rotogravure flexible packaging, pre-made pouches, and plastic containers sourced from manufacturing partners.",
    },
    "Karton.eu": {
        "core": "Cardboard Boxes; Custom Cartons; Packaging Materials; Warehouse Equipment",
        "focus": "German e-commerce supplier of cardboard boxes, custom cartons, and packaging materials for B2B and B2C shipping and warehouse use.",
    },
    "Packaging Direct B.V.": {
        "core": "Foodservice Disposables; Takeaway Packaging; Cups & Containers; Custom-Printed Packaging",
        "focus": "Dutch e-commerce company supplying disposables and packaging for hospitality and takeaway, including custom-printed options for Netherlands and Belgium customers.",
    },
    "Furodet Emballages": {
        "core": "Food Packaging Trays; Heat-Seal Trays; Meal Containers; Fruit & Vegetable Packaging; Takeaway Packaging",
        "focus": "French food packaging distributor supplying trays, meal containers, and related food packaging formats to foodservice and produce customers.",
    },
    "PRODINGER Verpackung": {
        "core": "Industrial Packaging Materials; Films; Cartons; Adhesive Technology; Hygiene Products",
        "focus": "German industrial packaging group distributing films, cartons, adhesives, and hygiene products with consultative packaging solutions for SME and industrial customers.",
    },
    "U.S. Plastic Corp.": {
        "core": "Plastic Containers; Labware; Industrial Plastic Products; Packaging Containers",
        "focus": "Distributor of plastic containers, labware, and industrial plastic products for laboratory, industrial, and packaging applications.",
    },
    "Container and Packaging Supply": {
        "core": "Plastic Bottles; Jars; Closures; Packaging Containers",
        "focus": "Online wholesale supplier of plastic bottles, jars, closures, and related packaging containers for product fillers and brand owners.",
    },
    "Freund Container & Supply": {
        "core": "Plastic Containers; Closures; Packaging Supplies",
        "focus": "Chicago-area distributor of plastic containers, closures, and packaging supplies for commercial and industrial packagers.",
    },
    "The Boxery": {
        "core": "Corrugated Boxes; Shipping Supplies; Moving Boxes; Packaging Materials",
        "focus": "Brooklyn-based supplier of corrugated boxes and shipping/moving packaging materials for local and e-commerce customers.",
    },
}

WEAK_CORE = {
    "Flexible packaging supplies; films; bags; pouches; stretch/shrink; protective packaging",
    "Flexible packaging; packaging films; shipping supplies",
    "Flexible packaging supplies; films; bags; stretch/shrink; packaging materials",
    "Packaging supplies",
    "Flexible packaging supplies; films; bags; packaging materials (channel/distributor)",
    "Packaging supplies; shipping materials",
    "Industrial packaging supplies",
    "Packaging materials wholesale",
    "Packaging materials",
    "Flexible packaging supplies; films; bags; pouches; stretch/shrink; packaging materials",
}

WEAK_SPEC_SUBSTR = (
    "channel partner for flexible packaging",
    "business distributor/wholesaler of packaging supplies",
    "us packaging distributor",
    "german packaging wholesaler",
    "uk packaging supplies distributor",
    "australian packaging supplies distributor",
    "distributes packaging materials and supplies",
    "european paper and packaging merchant",
    "regional us packaging distributor",
    "german packaging materials wholesaler",
    "offers a range of flexible packaging solutions",
    "supplies flexible packaging solutions",
    "distributes packaging materials including flexible",
)


# Full trailing geography clause only (avoid stripping last country word and leaving "and.")
COUNTRY_TAIL = re.compile(
    r"(?:,?\s+)?(?:in|across|throughout|serving|based in)\s+"
    r"(?:the\s+)?"
    r"(?:United States|USA|U\.S\.A\.|U\.S\.|Canada|United Kingdom|UK|England|Scotland|"
    r"Australia|New Zealand|Germany|France|Ireland|Japan|India|Singapore|Mexico|Netherlands|"
    r"Sweden|Norway|Denmark|Finland|Spain|Italy|Belgium|Switzerland|Austria|Poland|China|"
    r"South Africa|Brazil|Portugal|Malaysia|Thailand|Indonesia|South Korea|Hong Kong|"
    r"North America|Southern Europe|Europe|Asia|Asia Pacific|APAC|EMEA|LATAM|"
    r"South Australia|New South Wales|Victoria|Queensland|Western Australia|"
    r"Ontario|British Columbia|California|Texas|Florida|Illinois)"
    r"(?:,?\s+(?:the\s+)?(?:United States|USA|Canada|United Kingdom|UK|Ireland|Germany|France|"
    r"Australia|New Zealand|Netherlands|Belgium|Switzerland|Austria|China|Europe|"
    r"North America|Asia)(?:\s+and\s+(?:the\s+)?[A-Za-z. ]+)?)?"
    r"\.?\s*$",
    re.I,
)


def clean_core(t: str) -> str:
    t = re.sub(r"\s+", " ", (t or "").strip())
    t = re.sub(r";\s*;", "; ", t)
    t = t.strip(" ;")
    return t


def clean_specialty(t: str) -> str:
    t = re.sub(r"\s+", " ", (t or "").strip())
    if not t:
        return ""
    t = re.sub(r"[,;]?\s*others\.?$", "", t, flags=re.I).strip()
    for _ in range(3):
        nt = COUNTRY_TAIL.sub("", t).strip(" ,;")
        if nt == t:
            break
        t = nt
    # Remove dangling prepositions/conjunctions left by geography strip
    t = re.sub(
        r"\s+(?:in|across|throughout|serving|based in|for|and|or|the|of|to|with|including)\s*[.,;]*$",
        "",
        t,
        flags=re.I,
    ).strip(" ,;")
    t = re.sub(r"\s+in\s+(?:Southern|Northern|Eastern|Western|Central)\.?$", "", t, flags=re.I).strip()
    if t and t[0].islower():
        t = t[0].upper() + t[1:]
    if t and t[-1] not in ".!?":
        t += "."
    return t


def is_weak(core: str, spec: str) -> bool:
    c = (core or "").strip()
    s = (spec or "").strip()
    if not c or c in WEAK_CORE:
        return True
    if "Flexible packaging supplies; films; bags" in c:
        return True
    if not s or len(s) < 55:
        return True
    low = s.lower()
    if any(x in low for x in WEAK_SPEC_SUBSTR):
        return True
    if re.fullmatch(
        r"(US|UK|NJ|German|Australian|Canadian|French|Dutch|Indian|Italian|Swiss|Ireland)\s+packaging\s+(distributor|wholesaler|merchant|supplies).{0,40}",
        s,
        re.I,
    ):
        return True
    if re.search(r"\b(distributor|wholesaler|merchant)\.?$", s, re.I) and len(s) < 90:
        return True
    return False


def fetch_evidence(url: str, timeout: float = 12.0) -> str:
    if not url or not str(url).startswith("http"):
        return ""
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; VendorIntelBot/1.0; +research)",
        "Accept": "text/html,application/xhtml+xml",
    }
    texts = []
    try:
        with httpx.Client(follow_redirects=True, timeout=timeout, headers=headers) as client:
            r = client.get(url)
            if r.status_code < 400:
                html = r.text
                # crude text extract
                html = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", html)
                html = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", html)
                html = re.sub(r"(?is)<[^>]+>", " ", html)
                html = re.sub(r"&\w+;", " ", html)
                html = re.sub(r"\s+", " ", html)
                texts.append(html[:6000])
            for path in ("/about", "/about-us", "/about-us/", "/company", "/products", "/our-products"):
                try:
                    u = urljoin(url if url.endswith("/") else url + "/", path.lstrip("/"))
                    rr = client.get(u)
                    if rr.status_code < 400 and len(rr.text) > 400:
                        h = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", rr.text)
                        h = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", h)
                        h = re.sub(r"(?is)<[^>]+>", " ", h)
                        h = re.sub(r"\s+", " ", h)
                        texts.append(h[:4000])
                        break
                except Exception:
                    pass
    except Exception:
        return ""
    return "\n---\n".join(texts)[:10000]


def deepseek_extract(company: str, website: str, evidence: str) -> dict[str, str] | None:
    key = os.getenv("OPENAI_API_KEY")
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    if not key or not evidence or len(evidence) < 200:
        return None
    prompt = f"""Extract Core Categories and Specialty Focus for packaging CHANNEL company "{company}" ({website}).

STRICT RULES:
- Use ONLY facts supported by the Evidence below. If evidence is insufficient, return conf="low" and empty strings.
- Core Categories: 4-8 real product/service categories separated by "; " (no country names).
- Specialty Focus: 1-2 sentences in ENGLISH describing WHAT they do functionally (products + how they serve customers). Detailed. NO country/region at the end. No trailing "others".
- This is a distributor/wholesaler/retailer channel row where applicable — do not invent OEM manufacturing unless evidence says so.
- Translate product names to English if evidence is non-English.

Return ONLY JSON: {{"core":"...","focus":"...","conf":"high|medium|low"}}

Evidence:
{evidence[:8500]}
"""
    try:
        with httpx.Client(timeout=60.0) as client:
            r = client.post(
                f"{base}/chat/completions",
                headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
                json={
                    "model": model,
                    "temperature": 0,
                    "messages": [{"role": "user", "content": prompt}],
                },
            )
            r.raise_for_status()
            content = r.json()["choices"][0]["message"]["content"]
            m = re.search(r"\{[\s\S]*\}", content)
            if not m:
                return None
            data = json.loads(m.group(0))
            core = clean_core(str(data.get("core") or ""))
            focus = clean_specialty(str(data.get("focus") or ""))
            conf = str(data.get("conf") or "low").lower()
            if conf == "low" or not core or not focus or len(focus) < 40:
                return None
            # reject if still generic
            if is_weak(core, focus):
                return None
            return {"core": core, "focus": focus, "conf": conf}
    except Exception:
        return None


def main() -> None:
    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i for i, h in enumerate(headers)}
    ci, wi = ic["Company"] + 1, ic["Website"] + 1
    cci, sfi = ic["Core Categories"] + 1, ic["Specialty Focus"] + 1

    ckpt = {}
    if CKPT.exists():
        try:
            ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
        except Exception:
            ckpt = {}

    updated = []
    skipped = []
    # 1) Apply HAND
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, ci).value
        if name not in HAND:
            continue
        info = HAND[name]
        core = clean_core(info["core"])
        focus = clean_specialty(info["focus"])
        old_c = ws.cell(row, cci).value
        old_s = ws.cell(row, sfi).value
        ws.cell(row, cci).value = core
        ws.cell(row, sfi).value = focus
        updated.append({"company": name, "source": "hand", "core": core, "focus": focus, "old_core": old_c, "old_focus": old_s})

    # 2) Clean country tails on ALL specialty cells
    for row in range(2, ws.max_row + 1):
        sp = clean_specialty(str(ws.cell(row, sfi).value or ""))
        if sp:
            ws.cell(row, sfi).value = sp

    # 3) DeepSeek for remaining weak (cap batch for runtime)
    weak_rows = []
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, ci).value
        if not name or name in HAND:
            continue
        core = str(ws.cell(row, cci).value or "")
        spec = str(ws.cell(row, sfi).value or "")
        if is_weak(core, spec):
            weak_rows.append(row)

    # process up to 100 weak with evidence per run
    limit = 100
    for idx, row in enumerate(weak_rows[:limit]):
        name = ws.cell(row, ci).value
        web = str(ws.cell(row, wi).value or "")
        key = str(name)
        if key in ckpt and ckpt[key].get("core") and ckpt[key].get("focus"):
            ws.cell(row, cci).value = clean_core(ckpt[key]["core"])
            ws.cell(row, sfi).value = clean_specialty(ckpt[key]["focus"])
            updated.append({"company": name, "source": "ckpt", **ckpt[key]})
            continue
        print(f"[{idx+1}/{min(limit,len(weak_rows))}] {name} …", flush=True)
        evidence = fetch_evidence(web)
        result = deepseek_extract(str(name), web, evidence)
        if result:
            ws.cell(row, cci).value = result["core"]
            ws.cell(row, sfi).value = result["focus"]
            ckpt[key] = result
            updated.append({"company": name, "source": "deepseek", **result})
            CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")
        else:
            skipped.append({"company": name, "website": web, "reason": "insufficient_evidence"})
        time.sleep(0.4)

    # Final pass: clean all specialty again
    for row in range(2, ws.max_row + 1):
        sp = clean_specialty(str(ws.cell(row, sfi).value or ""))
        if sp:
            ws.cell(row, sfi).value = sp

    wb.save(FP)
    # stats
    weak_left = 0
    empty_c = empty_s = 0
    for row in range(2, ws.max_row + 1):
        core = str(ws.cell(row, cci).value or "")
        spec = str(ws.cell(row, sfi).value or "")
        if not core.strip():
            empty_c += 1
        if not spec.strip():
            empty_s += 1
        if is_weak(core, spec):
            weak_left += 1

    audit = {
        "ts": datetime.now().isoformat(timespec="seconds"),
        "hand": len(HAND),
        "updated": len(updated),
        "skipped": skipped,
        "weak_left": weak_left,
        "empty_core": empty_c,
        "empty_spec": empty_s,
        "rows": updated[:200],
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"updated={len(updated)} weak_left={weak_left} empty_core={empty_c} empty_spec={empty_s}")
    print(f"saved {FP}")
    print(f"audit {AUDIT}")


if __name__ == "__main__":
    main()
