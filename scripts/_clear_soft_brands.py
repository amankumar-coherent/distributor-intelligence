#!/usr/bin/env python3
from pathlib import Path
import openpyxl

p = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
    r"\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
wb = openpyxl.load_workbook(p)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

CLEAR = {
    "Mystic Moments",
    "Nutra Food Ingredients",
    "Patties Food Group Distribution",
    "Performance Food Group",
    "Pro*Act",
    "Sojitz",
    "Tradin Organic",
    "Olam Food Ingredients Distribution",
}
FIX = {
    "Savola Distribution": "Afia, Al-Arabiya, others",
    "Nature's Pride": "Nature's Pride, others",
    "Pampa Foods": "Pampa Foods, others",
    "Prinova": "Prinova, others",
}

for row in range(3, ws.max_row + 1):
    co = str(ws.cell(row, ic["Company"]).value or "").strip()
    if co in CLEAR:
        ws.cell(row, ic["Key Brands Represented"]).value = "Not publicly disclosed"
        print("cleared", co)
    if co in FIX:
        ws.cell(row, ic["Key Brands Represented"]).value = FIX[co]
        print("fixed", co)

wb.save(p)
wb.close()

wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c or "") for c in rows[1]]
ic = {h: i for i, h in enumerate(hdr) if h}
real = [
    (r[ic["Company"]], r[ic["Key Brands Represented"]])
    for r in rows[2:]
    if str(r[ic["Key Brands Represented"]] or "").strip() != "Not publicly disclosed"
]
print(f"real brands {len(real)} / {len(rows)-2}")
for co, b in real:
    print(co, "|", b)
wb.close()
