# -*- coding: utf-8 -*-
"""Audit LinkedIn cells on the four done markets."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}
FILES = {
    "electrical": Path(
        r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
        r"\global_electrical_distribution_market_global_FINAL.xlsx"
    ),
    "avocado": Path(
        r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
        r"\global_avocado_oil_market.xlsx"
    ),
    "smartphone": Path(
        r"D:\vendor-intel\output\web_expand\global_smartphone_market_global"
        r"\global_smartphone_market.xlsx"
    ),
    "firewall": Path(
        r"D:\vendor-intel\output\run4\global_firewall_market_global"
        r"\global_firewall_market.xlsx"
    ),
}

PERSON_RE = re.compile(r"https?://([a-z]{2}\.)?linkedin\.com/in/[^?\s]+", re.I)


def classify_li(li: str) -> str:
    s = (li or "").strip()
    low = s.lower()
    if not s or low in NPD:
        return "npd"
    if "/company/" in low or "/school/" in low:
        return "company_page"
    if "/pub/" in low or "/search/" in low or "linkedin.com/in/dir/" in low:
        return "search"
    if PERSON_RE.search(s):
        return "ok_person"
    if "linkedin.com" in low:
        return "bad_format"
    return "not_linkedin"


def main():
    need = []
    bad_existing = []
    counts = {}
    for mkt, path in FILES.items():
        wb = load_workbook(path, data_only=True)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        headers = None
        hr = 1
        for r in range(1, 6):
            vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
            if "Company" in vals and "Contact Person" in vals:
                headers = vals
                hr = r
                break
        c = {"ok_person": 0, "npd": 0, "company_page": 0, "search": 0, "bad_format": 0, "not_linkedin": 0, "cp_no_li": 0}
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            d = {
                headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
                for i in range(len(headers))
            }
            co = d.get("Company") or ""
            if not co:
                continue
            person = (d.get("Contact Person") or "").strip()
            role = (d.get("Role") or "").strip()
            li = (d.get("LinkedIn") or "").strip()
            has_cp = bool(person) and person.lower() not in NPD
            kind = classify_li(li)
            if has_cp:
                c[kind] = c.get(kind, 0) + 1
                rec = {
                    "m": mkt,
                    "company": co,
                    "person": person,
                    "role": role,
                    "li": li,
                    "hq": d.get("Headquarters") or "",
                    "kind": kind,
                }
                if kind != "ok_person":
                    c["cp_no_li"] += 1
                    need.append(rec)
                elif kind in ("company_page", "search", "bad_format", "not_linkedin"):
                    bad_existing.append(rec)
            elif kind not in ("npd",) and kind != "ok_person":
                bad_existing.append({
                    "m": mkt, "company": co, "person": person, "role": role,
                    "li": li, "hq": d.get("Headquarters") or "", "kind": kind,
                })
        counts[mkt] = c
        wb.close()
    out = Path(r"D:\vendor-intel\output\_linkedin_need.json")
    out.write_text(json.dumps({"counts": counts, "need": need, "bad": bad_existing}, ensure_ascii=False, indent=2), encoding="utf-8")
    print("counts", json.dumps(counts, indent=2))
    print("need_person_li", len(need))
    print("bad_existing", len(bad_existing))
    for m in FILES:
        rows = [x for x in need if x["m"] == m]
        print("====", m, len(rows), "====")
        for x in rows[:40]:
            print(f"{x['person'][:40]:40} | {x['role'][:36]:36} | {x['company'][:32]} | {x['kind']}")


if __name__ == "__main__":
    main()
