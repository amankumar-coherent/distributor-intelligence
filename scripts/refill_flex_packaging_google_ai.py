"""Re-fill Global Flexible Packaging FINAL via Google AI scraper + DeepSeek residual.

Only fills empty / NPD cells — does not invent contacts. Company + Website kept.
"""
from __future__ import annotations

import asyncio
import os
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))

FP = ROOT / (
    "output/web_expand/global_flexible_packaging_market_global"
    "/global_flexible_packaging_market_global_FINAL.xlsx"
)
CKPT = FP.parent / "refill_google_ai_ckpt.json"
MARKET = "Global Flexible Packaging Market"
FAMILY = "flexible_packaging"


def _load_rows() -> tuple[list[str], list[dict[str, str]]]:
    wb = load_workbook(FP, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, str]] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = {
            headers[i]: ("" if i >= len(r) or r[i] is None else str(r[i]).strip())
            for i in range(len(headers))
            if headers[i]
        }
        if d.get("Company"):
            rows.append(d)
    wb.close()
    return headers, rows


def _save_rows(headers: list[str], rows: list[dict[str, str]]) -> None:
    wb = load_workbook(FP)
    ws = wb.active
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for row in rows:
        ws.append([row.get(h, "") or None for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(FP)


async def main() -> None:
    os.environ.setdefault("PYTHONIOENCODING", "utf-8")
    os.environ["GOOGLE_AI_SCRAPER_ENABLED"] = "true"
    os.environ["SEARCH_STACK_SOURCES"] = "google_ai"
    os.environ["SEARCH_STACK_OPENAI_RESIDUAL"] = "true"
    os.environ["GOOGLE_AI_GAP_FILL_MAX_QUERIES"] = os.environ.get(
        "GOOGLE_AI_GAP_FILL_MAX_QUERIES", "22"
    )
    os.environ["GOOGLE_AI_SCRAPER_LLM_CLEAN"] = "true"
    os.environ["GOOGLE_AI_SCRAPER_LLM_CLEAN_ONLY_WHEN_MISSING"] = "false"
    os.environ["GOOGLE_AI_SCRAPER_LLM_MODEL"] = os.environ.get(
        "GOOGLE_AI_SCRAPER_LLM_MODEL", "deepseek-chat"
    )
    os.environ["EXPAND_FILL_CONCURRENT"] = "1"
    os.environ["GAP_FILL_GOOGLE_AI_TIMEOUT_SEC"] = os.environ.get(
        "GAP_FILL_GOOGLE_AI_TIMEOUT_SEC", "900"
    )

    from vendor_intel.pipeline.openai_expand_checkpoint import ExpandCheckpoint
    from vendor_intel.pipeline.search_stack_fill import search_stack_fill_rows

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = FP.with_name(FP.stem + f"_backup_before_refill_{stamp}.xlsx")
    shutil.copy2(FP, backup)
    print(f"backup: {backup}", flush=True)

    headers, rows = _load_rows()
    print(f"companies: {len(rows)}", flush=True)

    ckpt, resumed = ExpandCheckpoint.load_or_create(
        FP.parent,
        query=MARKET,
        country="global",
        batch="refill_flex",
        target=len(rows),
        resume=True,
        meta={"refill": True, "source": "google_ai"},
    )
    print(f"checkpoint resumed={resumed} path={ckpt.path}", flush=True)

    filled = await search_stack_fill_rows(
        query=MARKET,
        family=FAMILY,
        rows=rows,
        ckpt=ckpt,
        sources=("google_ai",),
        use_apollo=False,
        openai_residual=True,
    )

    # preserve original order by company name key
    by_name = {str(r.get("Company") or "").strip().lower(): r for r in filled}
    out = []
    for orig in rows:
        key = str(orig.get("Company") or "").strip().lower()
        out.append(by_name.get(key) or orig)

    _save_rows(headers, out)

    # fill rates
    print("--- fill rates after ---", flush=True)
    for h in headers:
        n = sum(1 for r in out if str(r.get(h) or "").strip())
        print(f"  {n}/{len(out)}  {h}", flush=True)
    print(f"saved: {FP}", flush=True)


if __name__ == "__main__":
    asyncio.run(main())
