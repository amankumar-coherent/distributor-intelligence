#!/usr/bin/env python3
import re
from pathlib import Path

import openpyxl

X = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
NPD = "Not publicly disclosed"
EXTRA = {
    "Grocers Supply": "Best Yet, Food Club, Full Circle, TopCare, others",
    "Nature's Best": "CADIA, MADE WITH, others",
    "Tony's Fine Foods": "Alcatraz Bay Seafoods, North Coast Fisheries, Gold Rush Creamery, Diana Supreme, Wild Harvest, Field Day, Essential Everyday, others",
    "DPI Specialty Foods": "CADIA, MADE WITH, others",
    "Londis Wholesale": "Jack's, Euro Shopper, Chef's Larder, others",
    "Parfetts": "Go Local, Xenergy, others",
    "World Finer Foods": "Reese, DaVinci, Ty Ling, El Rio, London Pub, Wellington, Candoni, others",
    "Liberty Richter": "Reese, DaVinci, Ty Ling, others",
    "JJ Foodservice": "Signature, Tabaki, Royal Roast, Super Crisp, Rice Majestic, others",
    "PriceSmart": "Member's Selection, others",
    "Cheney Brothers": "Grand Western, Crazy Fish International, Duck Deli, Fronte, Bari Brand, Acento Latino, Tres Hermanos, Cheney Brothers Estate, Grumpy Potato, others",
    "Doerle Food Services": "Sysco Classic, Sysco Premium, Sysco Essentials, others",
    "Foodstuffs North Island Wholesale": "Pams, Pams Value, Pams Finest, others",
    "Gilmours": "Pams, Pams Value, Pams Finest, others",
    "McLane Company": "CVP, Hometown Market, YumBees, Beau Dacious, Excursion Beef Jerky, Road-Tech, Work Fare, Divinely Sweet, Fly Guys Pizza, others",
    "Piggly Wiggly Alabama Distributing": "Piggly Wiggly, Sol Select, others",
    "H.T. Hackney": "Uncle Ray's, Great American Deli, others",
    "Reynolds Catering Supplies": "Carnivore, Shoal, Davin Foods, Solstice Select, Braehead Foods, others",
    "The SYGMA Network": "Sysco Classic, Sysco Premium, Sysco Essentials, others",
    "FreshPoint": "Sysco Classic, Sysco Premium, others",
    "FreshPoint Central California": "Sysco Classic, Sysco Premium, others",
    "Selgros Cash & Carry": "Transgourmet Economy, Transgourmet Quality, Transgourmet Ursprung, Transgourmet Natura, others",
    "Fonterra Foodservice Distribution": "Anchor Food Professionals, NZMP, others",
    "Labatt Food Service": NPD,
    "So African Natural Oils": "Seed Essential, SANO, others",
}

wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}
n = 0
for r in range(3, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co in EXTRA:
        ws.cell(r, ic["Key Brands Represented"]).value = EXTRA[co]
        n += 1
wb.save(X)
wb.close()
print("extra_applied", n)

wb = openpyxl.load_workbook(X, data_only=True)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i for i, h in enumerate(hdr) if h}
real = npd = 0
npd_list = []
off_re = re.compile(
    r"cosmetic|soapmak|aromatherapy|oleochemical|specialty chemical|personal care raw|candle",
    re.I,
)
food_re = re.compile(
    r"edible|cooking oil|foodservice|horeca|grocery|fmcg|avocado|vegetable oil|specialty oil|olive|gourmet|cash.?carry|wholesale food|cold-pressed|seed oil",
    re.I,
)
suspect = []
for row in ws.iter_rows(min_row=3, values_only=True):
    c = str(row[ic["Company"]] or "").strip()
    if not c:
        continue
    b = str(row[ic["Key Brands Represented"]] or "").strip()
    core = str(row[ic["Core Categories"]] or "")
    focus = str(row[ic["Specialty Focus"]] or "")
    if b in ("", NPD):
        npd += 1
        npd_list.append(c)
    else:
        real += 1
    blob = f"{core} {focus}"
    if off_re.search(blob) and not food_re.search(blob):
        suspect.append((c, core[:70]))
print("TOTAL", real + npd, "REAL", real, "NPD", npd)
print("suspect_count", len(suspect))
for s in suspect:
    print("suspect", s)
print("NPD_HEAD")
for x in npd_list[:40]:
    print(x)
wb.close()
