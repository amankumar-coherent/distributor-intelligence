"""Fill Founded years for remaining Flexible Packaging rows.

Pipeline (no hallucination):
1) Google AI scraper ask: "<company> packaging founded OR established year"
2) DeepSeek extracts candidate year + confidence from that text (and optional web snippet)
3) Accept ONLY if:
   - year is 1800-2026 AND
   - year appears in Google AI markdown (or company About page fetch) AND
   - DeepSeek confidence is high / year matches scraped text
4) Never write DeepSeek-only guesses.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote_plus
from urllib.request import Request, urlopen

from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "founded_fill_google_deepseek_audit.json"
CKPT = FP.parent / "founded_fill_google_deepseek_ckpt.json"

SCRAPER = (os.getenv("GOOGLE_AI_SCRAPER_BASE_URL") or "http://127.0.0.1:15551").rstrip("/")
API_KEY = (os.getenv("OPENAI_API_KEY") or "").strip()
BASE_URL = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()

YEAR_RE = re.compile(r"\b((?:18|19|20)\d{2})\b")
BATCH_LIMIT = int(os.getenv("FOUNDED_FILL_LIMIT") or "80")  # half-batch


def http_get_json(url: str, timeout: int = 100) -> dict:
    req = Request(url, headers={"User-Agent": "Mozilla/5.0 VendorIntelFounded/1.0"})
    with urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8", errors="replace"))


def google_ai_ask(query: str) -> str:
    url = f"{SCRAPER}/ask?q={quote_plus(query)}"
    try:
        data = http_get_json(url, timeout=120)
    except Exception as e:
        return f"[scraper_error] {e}"
    if data.get("error"):
        return f"[scraper_error] {data.get('error')}"
    md = (data.get("markdown") or data.get("text") or "").strip()
    return md


def fetch_about_text(website: str) -> str:
    if not website:
        return ""
    base = website if website.startswith("http") else "https://" + website
    base = base.rstrip("/")
    paths = [
        "",
        "/about",
        "/about-us",
        "/about-us/",
        "/company",
        "/en/about",
        "/our-story",
        "/history",
    ]
    chunks = []
    for p in paths[:4]:
        try:
            req = Request(
                base + p,
                headers={"User-Agent": "Mozilla/5.0"},
            )
            with urlopen(req, timeout=12) as resp:
                html = resp.read(120_000).decode("utf-8", errors="ignore")
            # crude strip tags
            text = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
            text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
            text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"\s+", " ", text)
            chunks.append(text[:4000])
        except Exception:
            continue
    return "\n".join(chunks)[:8000]


def deepseek_extract(company: str, website: str, google_md: str, about_text: str) -> dict:
    """Return {year, confidence, reason} — year only if evidenced in provided text."""
    from openai import OpenAI

    client = OpenAI(api_key=API_KEY, base_url=BASE_URL)
    prompt = f"""You extract ONLY a founding/established year that is explicitly supported by the SOURCE TEXT below.

Company: {company}
Website: {website}

Rules:
- Return JSON only: {{"year":"YYYY"|null,"confidence":"high"|"medium"|"low","evidence":"short quote","reason":"..."}}
- year MUST appear in the source text (Google or About).
- Prefer phrases like founded/established/incorporated/since/opened in YEAR.
- If multiple years, pick company founding (not plant expansion / acquisition year unless that's the company formation).
- Reject evidence about a DIFFERENT legal entity with a similar name (wrong country, wrong website, wrong brand).
- Prefer evidence that mentions this company's website domain or exact company name.
- If unclear or not stated, year=null. NEVER invent.
- confidence=high only if explicit founded/established/incorporated year for this company.

SOURCE TEXT:
=== GOOGLE AI ===
{google_md[:6000]}

=== ABOUT PAGE SNIPPETS ===
{about_text[:5000]}
"""
    resp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": "You are a careful research extractor. Never invent years."},
            {"role": "user", "content": prompt},
        ],
        temperature=0,
        response_format={"type": "json_object"},
    )
    raw = resp.choices[0].message.content or "{}"
    try:
        return json.loads(raw)
    except Exception:
        return {"year": None, "confidence": "low", "reason": "bad_json", "raw": raw[:300]}


def year_in_text(year: str, *texts: str) -> bool:
    blob = "\n".join(texts or [])
    return bool(re.search(rf"\b{re.escape(year)}\b", blob))


def company_tokens(company: str) -> list[str]:
    stop = {
        "inc", "llc", "ltd", "ltda", "co", "company", "corp", "corporation",
        "sa", "sas", "gmbh", "bv", "plc", "pty", "the", "and", "of", "group",
        "packaging", "package", "packages", "emballage", "emballages", "empaque",
        "empaques", "flexible", "plastics", "plastic", "supply", "supplies",
        "industries", "industry", "solutions", "international", "global",
    }
    toks = re.findall(r"[A-Za-z0-9]{3,}", company.lower())
    return [t for t in toks if t not in stop]


def entity_match(company: str, *texts: str) -> bool:
    """Require distinctive company tokens in evidence to avoid name collisions."""
    blob = "\n".join(texts or []).lower()
    toks = company_tokens(company)
    if not toks:
        # fallback: first significant word from original name
        toks = re.findall(r"[A-Za-z0-9]{4,}", company.lower())[:2]
    if not toks:
        return False
    hits = sum(1 for t in toks if t in blob)
    need = 1 if len(toks) == 1 else min(2, len(toks))
    return hits >= need


def host_of(website: str) -> str:
    try:
        from urllib.parse import urlparse

        u = website if "://" in website else f"https://{website}"
        return (urlparse(u).hostname or "").lower().replace("www.", "")
    except Exception:
        return ""


def validate_candidate(
    company: str, website: str, year: str | None, conf: str, google_md: str, about: str
) -> tuple[bool, str]:
    if not year or not re.fullmatch(r"(18|19|20)\d{2}", str(year)):
        return False, "bad_year"
    y = int(year)
    if y < 1800 or y > 2026:
        return False, "out_of_range"
    if conf not in ("high", "medium"):
        return False, f"conf_{conf}"
    # Must appear in sources
    if not year_in_text(str(year), google_md, about):
        return False, "year_not_in_sources"
    if not entity_match(company, google_md, about):
        return False, "entity_mismatch"
    host = host_of(website)
    host_ok = bool(host) and host in (google_md + "\n" + about).lower()
    about_year_ok = bool(about) and year_in_text(str(year), about)
    # Require website domain OR About-page year to reduce cross-company collisions
    if not host_ok and not about_year_ok:
        return False, "need_domain_or_about"
    # Prefer high; medium only if About page also contains year near founded words
    if conf == "medium":
        if not about or not year_in_text(str(year), about):
            return False, "medium_needs_about"
        if not re.search(
            rf"(founded|established|incorporated|since|est\.?)\D{{0,40}}{year}|{year}\D{{0,40}}(founded|established|incorporated)",
            about,
            re.I,
        ):
            # still allow if google has clear founded phrase
            if not re.search(
                rf"(founded|established|incorporated|since)\D{{0,40}}{year}|{year}\D{{0,20}}(founded|established)",
                google_md,
                re.I,
            ):
                return False, "medium_weak_context"
    # high still needs founded-ish context somewhere
    if conf == "high":
        if not re.search(
            rf"(founded|established|incorporated|since|est\.?|opened|began|started)\D{{0,60}}{year}|{year}\D{{0,40}}(founded|established|incorporated|anniversary)",
            google_md + "\n" + about,
            re.I,
        ):
            return False, "high_missing_founded_context"
    return True, "ok"


def main() -> None:
    if not API_KEY:
        raise SystemExit("OPENAI_API_KEY missing")

    # health
    try:
        health = http_get_json(f"{SCRAPER}/health", timeout=10)
        print("scraper", health.get("status"), "ext", health.get("extension_connected"))
        if not health.get("extension_connected"):
            print("WARNING: extension not connected — Google AI may fail")
    except Exception as e:
        raise SystemExit(f"Google AI scraper not reachable: {e}")

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = FP.with_name(FP.stem + f"_backup_before_founded_gds_{stamp}.xlsx")
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = {str(h): i + 1 for i, h in enumerate(headers) if h}

    rows = []
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, col["Company"]).value or "").strip()
        site = str(ws.cell(r, col["Website"]).value or "").strip()
        founded = ws.cell(r, col["Founded"]).value
        fs = "" if founded is None else str(founded).strip()
        if not fs or fs in ("None", "N/A", "Not publicly disclosed", "-"):
            rows.append({"row": r, "company": name, "website": site})

    ckpt = {"done": {}, "rejected": [], "errors": []}
    if CKPT.exists():
        try:
            ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
        except Exception:
            pass

    # Manual verified from this session / prior web research (always allowed)
    MANUAL = {
        "Emballages Carrousel": ("1971", "carrousel.ca bio PDF Denis Jalbert founded 1971"),
        "Allied Packaging": ("2000", "ASIC Allied Packaging Co Pty Limited registered 2000-06-28"),
    }

    filled = []
    rejected = list(ckpt.get("rejected") or [])
    errors = list(ckpt.get("errors") or [])
    processed = 0

    for item in rows:
        if processed >= BATCH_LIMIT:
            break
        name = item["company"]
        site = item["website"]
        r = item["row"]
        if name in (ckpt.get("done") or {}):
            continue
        if name in (ckpt.get("leave_blank") or {}):
            print(f"  SKIP leave_blank {name}", flush=True)
            continue

        # Manual first
        if name in MANUAL:
            year, src = MANUAL[name]
            ws.cell(r, col["Founded"]).value = year
            filled.append({"company": name, "founded": year, "source": src, "method": "manual_verified"})
            ckpt.setdefault("done", {})[name] = {"year": year, "source": src}
            processed += 1
            continue

        host = ""
        try:
            from urllib.parse import urlparse

            host = (urlparse(site if "://" in site else f"https://{site}").hostname or "").replace(
                "www.", ""
            )
        except Exception:
            host = ""
        q = f'"{name}" {host} packaging founded OR established OR "since" year'.strip()
        print(f"[{processed+1}/{BATCH_LIMIT}] {name}", flush=True)
        google_md = google_ai_ask(q)
        print(f"  google_chars={len(google_md)}", flush=True)
        time.sleep(1.0)
        about = fetch_about_text(site)

        try:
            extracted = deepseek_extract(name, site, google_md, about)
        except Exception as e:
            errors.append({"company": name, "error": str(e)[:200]})
            print(f"  ERR {e}", flush=True)
            processed += 1
            continue

        year = extracted.get("year")
        conf = str(extracted.get("confidence") or "low").lower()
        ok, why = validate_candidate(name, site, str(year) if year else None, conf, google_md, about)
        if ok:
            ws.cell(r, col["Founded"]).value = str(year)
            rec = {
                "company": name,
                "founded": str(year),
                "confidence": conf,
                "evidence": extracted.get("evidence"),
                "reason": extracted.get("reason"),
                "method": "google_ai+deepseek_verified",
            }
            filled.append(rec)
            ckpt.setdefault("done", {})[name] = rec
            print(f"  OK {year} ({conf})", flush=True)
        else:
            rejected.append(
                {
                    "company": name,
                    "candidate": year,
                    "confidence": conf,
                    "why": why,
                    "evidence": extracted.get("evidence"),
                    "reason": extracted.get("reason"),
                }
            )
            print(f"  REJECT {year} ({conf}) {why}", flush=True)
        processed += 1

        if processed % 3 == 0:
            # Reload sheet so concurrent manual edits are not wiped
            wb2 = load_workbook(FP)
            ws2 = wb2.active
            headers2 = [c.value for c in ws2[1]]
            col2 = {str(h): i + 1 for i, h in enumerate(headers2) if h}
            for rec in filled:
                name = rec["company"]
                year = rec["founded"]
                for rr in range(2, ws2.max_row + 1):
                    if str(ws2.cell(rr, col2["Company"]).value or "").strip() == name:
                        ws2.cell(rr, col2["Founded"]).value = str(year)
                        break
            wb2.save(FP)
            wb, ws, col = wb2, ws2, col2
            CKPT.write_text(json.dumps(ckpt, indent=2), encoding="utf-8")

    # Final save: merge only this run's fills into latest file
    wb2 = load_workbook(FP)
    ws2 = wb2.active
    headers2 = [c.value for c in ws2[1]]
    col2 = {str(h): i + 1 for i, h in enumerate(headers2) if h}
    for rec in filled:
        name = rec["company"]
        year = rec["founded"]
        for rr in range(2, ws2.max_row + 1):
            if str(ws2.cell(rr, col2["Company"]).value or "").strip() == name:
                ws2.cell(rr, col2["Founded"]).value = str(year)
                break
    wb2.save(FP)
    wb, ws, col = wb2, ws2, col2
    empty = has = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col["Founded"]).value
        if v is None or str(v).strip() in ("", "None", "N/A", "Not publicly disclosed", "-"):
            empty += 1
        else:
            has += 1

    audit = {
        "backup": str(backup),
        "filled": filled,
        "filled_count": len(filled),
        "rejected_sample": rejected[-40:],
        "rejected_count": len(rejected),
        "errors": errors,
        "founded_filled_total": has,
        "founded_empty_total": empty,
        "total_rows": ws.max_row - 1,
        "batch_limit": BATCH_LIMIT,
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    ckpt["rejected"] = rejected
    ckpt["errors"] = errors
    CKPT.write_text(json.dumps(ckpt, indent=2), encoding="utf-8")
    print(
        f"FILLED_NOW={len(filled)} TOTAL_FILLED={has} STILL_EMPTY={empty} "
        f"REJECTED={len(rejected)} ERRORS={len(errors)}"
    )


if __name__ == "__main__":
    main()
