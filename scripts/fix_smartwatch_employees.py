"""Fix smartwatch FINAL Employees column → verified ranges (no garbage / raw mega-numbers)."""
from __future__ import annotations

import re
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

PATHS = [
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL_320plus.xlsx"),
]

# Verified overrides (agent web research)
OVERRIDES = {
    "d&h distributing": "1001-5000",  # ~1,600 (Forbes/Wikipedia Nov 2024)
    "macys": "50001-100000",  # ~90k–94k (Macy's SEC)
    "macy's": "50001-100000",
    "target": "100001-500000",  # ~415k–440k
    "walmart": "1000001+",  # ~2.1M
}

NPD = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd"}

BUCKETS = [
    (10, "1-10"),
    (50, "11-50"),
    (200, "51-200"),
    (500, "201-500"),
    (1000, "501-1000"),
    (5000, "1001-5000"),
    (10000, "5001-10000"),
    (50000, "10001-50000"),
    (100000, "50001-100000"),
    (500000, "100001-500000"),
    (1000000, "500001-1000000"),
]


def norm_name(s: str) -> str:
    return re.sub(r"[^a-z0-9&]+", " ", (s or "").lower()).strip()


def parse_int(tok: str) -> int | None:
    tok = tok.replace(",", "").replace(" ", "")
    if not tok.isdigit():
        return None
    return int(tok)


def number_to_range(n: int) -> str:
    for hi, label in BUCKETS:
        if n <= hi:
            return label
    return "1000001+"


def looks_like_garbage(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if re.search(
        r"(?i)\b(trust|foundation|bank|llc|inc\.?|ltd|gmbh|corp|partners|holdings)\b",
        t,
    ):
        if not re.search(r"\d", t):
            return True
    if re.fullmatch(r"[A-Za-z][A-Za-z &.'-]{2,}", t) and t.lower() not in NPD:
        return True
    return False


def normalize_employees(raw: str, company: str) -> tuple[str, str]:
    key = norm_name(company)
    for needle, rng in OVERRIDES.items():
        if needle in key or key in needle:
            return rng, f"verified_override:{needle}"

    v = (raw or "").strip()
    if v.lower() in NPD:
        return "Not publicly disclosed", "npd"

    if looks_like_garbage(v):
        return "Not publicly disclosed", f"cleared_garbage:{v}"

    v2 = re.sub(r"(?i)\s*(employees?|staff|team members?)\s*$", "", v).strip()
    v2 = v2.replace("–", "-").replace("—", "-").replace("−", "-")
    v2 = re.sub(r"\s+", "", v2)

    m = re.fullmatch(r"([\d,]+)\s*-\s*([\d,]+)", v2)
    if m:
        a, b = parse_int(m.group(1)), parse_int(m.group(2))
        if a is not None and b is not None and a <= b:
            mid = (a + b) // 2
            return number_to_range(mid), f"normalized_range:{a}-{b}"

    m = re.fullmatch(r"([\d,]+)\+", v2)
    if m:
        n = parse_int(m.group(1))
        if n is not None:
            # "10000+" means at least 10k → use next bucket, not the <=10k bucket
            return number_to_range(n + 1), f"open_ended:{n}+"

    m = re.fullmatch(r"[\d,]+", v2)
    if m:
        n = parse_int(v2)
        if n is not None:
            return number_to_range(n), f"point_to_range:{n}"

    nums = re.findall(r"[\d,]+", v2)
    if nums:
        n = parse_int(nums[0])
        if n is not None:
            return number_to_range(n), f"extracted:{n}"

    return "Not publicly disclosed", f"unparsed:{v}"


def fix_file(path: Path) -> dict:
    bak = path.with_name(
        path.stem + f"_pre_emp_fix_{datetime.now().strftime('%Y%m%d_%H%M%S')}" + path.suffix
    )
    shutil.copy2(path, bak)

    wb = openpyxl.load_workbook(path)
    ws = wb["Smartwatch Companies"] if "Smartwatch Companies" in wb.sheetnames else wb.active

    hdr_row = None
    hdr = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), 1):
        vals = [str(c or "") for c in row]
        if "Company" in vals and "Employees" in vals:
            hdr_row = i
            hdr = vals
            break
    if not hdr_row:
        raise RuntimeError(f"No header in {path}")

    emp_col = hdr.index("Employees") + 1
    co_col = hdr.index("Company") + 1

    changes = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        company = str(ws.cell(r, co_col).value or "").strip()
        if not company:
            continue
        old = str(ws.cell(r, emp_col).value or "").strip()
        new, reason = normalize_employees(old, company)
        if new != old:
            ws.cell(r, emp_col).value = new
            changes.append({"company": company, "old": old, "new": new, "reason": reason})

    wb.save(path)
    return {"file": str(path), "backup": str(bak), "changed": len(changes), "samples": changes}


def main() -> None:
    for p in PATHS:
        rep = fix_file(p)
        print("FILE", rep["file"])
        print(" backup", rep["backup"])
        print(" changed", rep["changed"])
        for c in rep["samples"]:
            if (
                "garbage" in c["reason"]
                or "override" in c["reason"]
                or "point_to_range" in c["reason"]
                or int(re.search(r"\d+", c["old"] or "0").group()) >= 100000
                if re.search(r"\d{5,}", c["old"] or "")
                else False
            ):
                print(f"  {c['company']}: {c['old']!r} -> {c['new']!r} ({c['reason']})")
        # always print key fixes
        for c in rep["samples"]:
            if c["company"] in {"D&H Distributing", "Macys", "Target", "Walmart"} or c["old"] in {
                "BENEFIT TRUST",
                "175000",
                "400000",
                "2300000",
            }:
                print(f"  KEY {c['company']}: {c['old']!r} -> {c['new']!r}")
        print()

    wb = openpyxl.load_workbook(PATHS[0], data_only=True)
    ws = wb["Smartwatch Companies"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    hi = None
    for i, r in enumerate(rows[:5]):
        vals = [str(c or "") for c in r]
        if "Employees" in vals:
            hi = i
            hdr = vals
            break
    ei = hdr.index("Employees")
    c = Counter(str(r[ei] or "") for r in rows[hi + 1 :] if r and r[0])
    print("After unique Employees:")
    for v, n in c.most_common():
        print(f"  {n:4d} | {v}")
    wb.close()


if __name__ == "__main__":
    main()
