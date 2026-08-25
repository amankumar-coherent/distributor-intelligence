#!/usr/bin/env python3
"""Final formatting pass for the Global Wearable Medical Devices Market sheet.

Core Categories:
  - Rows whose CURRENT Core Categories already contain a wearable-specific
    label (written by the earlier verify_wearable_categories_brands.py pass)
    get remapped onto the fixed 5-term controlled vocabulary:
        Medical Wearable Devices; Health Monitoring Devices;
        Remote Patient Monitoring; Diagnostic Wearables; Therapeutic Wearables
  - Rows with NO wearable-specific signal keep their real, honest categories
    (e.g. "Pharmaceuticals; OTC Healthcare; Retail Pharmacy") -- never forced
    into the wearable vocabulary, per explicit instruction.

Specialty Focus:
  - Rebuilt as: "[Application/clinical focus area(s)] for [customer segment]"
  - Application areas come from clinical-focus keyword rules (Cardiac
    Monitoring, Diabetes Monitoring, etc.) -- a different dimension from
    Core Categories, not a restatement of it.
  - Source text is the ORIGINAL research prose in the checkpoint JSON files
    (untouched by earlier reformat passes) + the sheet's Summary column.

Rows flagged "Needs manual verification" are left untouched. Pure keyword
rules, no LLM, no new web calls.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import openpyxl

XLSX = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/global_wearable_medical_devices_market_global_FINAL.xlsx"
SHEET = "Wearable Medical Companies"
CKPT_DIR = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/verify_checkpoints"
FLAG_PREFIX = "Needs manual verification"

# --- Core Categories: wearable-specific label -> 5-term controlled vocabulary ---
_WEARABLE_LABEL_MAP: dict[str, str] = {
    "continuous glucose monitors (cgm)": "Diagnostic Wearables",
    "wearable ecg monitors": "Diagnostic Wearables",
    "pulse oximeters": "Diagnostic Wearables",
    "blood pressure monitors": "Health Monitoring Devices",
    "smartwatches": "Health Monitoring Devices",
    "fitness trackers": "Health Monitoring Devices",
    "sleep trackers": "Health Monitoring Devices",
    "smart rings": "Health Monitoring Devices",
    "hearables / hearing aids": "Therapeutic Wearables",
    "remote patient monitoring (rpm) devices": "Remote Patient Monitoring",
    "patient monitors": "Remote Patient Monitoring",
    "wearable devices (general)": "Medical Wearable Devices",
}
_VOCAB_ORDER = [
    "Medical Wearable Devices",
    "Health Monitoring Devices",
    "Remote Patient Monitoring",
    "Diagnostic Wearables",
    "Therapeutic Wearables",
]


def remap_core_categories(current: str) -> str | None:
    """Return the 5-term vocabulary subset if any wearable-specific label is
    present, else None (meaning: leave the row's real categories untouched)."""
    items = [s.strip().lower() for s in current.split(";") if s.strip()]
    mapped = {
        _WEARABLE_LABEL_MAP[item]
        for item in items
        if item in _WEARABLE_LABEL_MAP
    }
    if not mapped:
        return None
    return "; ".join(v for v in _VOCAB_ORDER if v in mapped)


# --- Specialty Focus: clinical/application focus area rules ---
_FOCUS_RULES: list[tuple[str, str]] = [
    # NOTE: bare "cgm" deliberately excluded -- "CMA CGM" (a shipping group)
    # false-triggers it; require the unambiguous full phrase or a plain
    # diabetes-related word instead.
    (r"continuous glucose monitor|diabetes|glucose meter|insulin", "Diabetes Monitoring"),
    (r"cardiac|\becg\b|\bekg\b|heart|cardiovascular|holter", "Cardiac Monitoring"),
    (r"sleep (?:apnea|tracking|monitor)", "Sleep Monitoring"),
    (r"respiratory|pulmonary|oxygen therapy|ventilat|\bcpap\b", "Respiratory Care"),
    (r"neonatal|maternal|paediatric|pediatric", "Neonatal & Maternal Care"),
    (r"wound care|wound management", "Wound Care"),
    (r"surgical|surgery|operating room", "Surgical Support"),
    (r"rehabilitation|physiotherapy|physical therapy|mobility aid|wheelchair", "Rehabilitation & Mobility"),
    (r"infection (?:control|prevention)|sterili[sz]ation", "Infection Control"),
    (r"emergency|first aid|resuscitation|\baed\b|defibrillat", "Emergency Response"),
    (r"\bdental\b", "Dental Care"),
    (r"veterinar", "Veterinary Care"),
    (r"home\s?care|home\s?health", "Home Healthcare"),
    (r"hospice|palliative", "Palliative Care"),
    (r"diagnostic|laboratory|pathology", "Diagnostics"),
    (r"orthopedic|orthopaedic", "Orthopedic Care"),
    (r"mental health|behavioral health|psychiatric", "Behavioral Health"),
    (r"oncology|\bcancer\b", "Oncology Care"),
    (r"\brenal\b|dialysis|kidney", "Renal Care"),
    (r"vaccin|immuni[sz]ation", "Immunization"),
    (r"ostomy|urology|continence", "Urology & Continence Care"),
    (r"hearing|audiolog", "Hearing Care"),
    (r"\bvision\b|ophthalmic|\boptical\b", "Vision Care"),
    (r"fitness tracker|activity tracking|fitness monitoring", "Fitness & Wellness Monitoring"),
    (r"chronic (?:disease|condition)", "Chronic Disease Management"),
]
_FALLBACK_RULES: list[tuple[str, str]] = [
    # "broadline"/"wholesale" alone is NOT pharma-specific -- Fastenal calls
    # itself a "broadline industrial distributor". Require an explicit
    # pharma/medical/healthcare qualifier before assuming this fallback.
    (
        r"(?:broadline|full.line|wholesale distribut|wholesaler)"
        r"(?:(?!industrial|electronic|hardware|tool|fastener|construction).){0,40}"
        r"(?:pharma|drug|medicine|medical|healthcare)",
        "General Pharmaceutical & Medical Distribution",
    ),
    (r"retail pharmac|pharmacy chain|drugstore", "OTC & Retail Health Products"),
    (
        r"information technology|\bit systems\b|\bit services\b|electronic component|"
        r"industrial distribut|industrial suppl|fastener|hardware suppl|tool distribut",
        "General Industrial/IT Distribution",
    ),
]
_DEFAULT_FALLBACK = "General Distribution Services"

# --- Specialty Focus: customer segment rules ---
_SEGMENT_RULES: list[tuple[str, str]] = [
    (r"neonatal|maternal|paediatric|pediatric", "neonatal, maternal, and pediatric care providers"),
    (r"veterinar", "veterinary clinics"),
    (r"\bdental\b", "dental practices"),
    (r"home\s?care|home\s?health|homecare", "home healthcare patients and caregivers"),
    (r"retail pharmac|pharmacy chain|drugstore|drug store", "retail pharmacy consumers"),
    (r"community pharmac", "community pharmacies"),
    (r"laborator(?:y|ies)|diagnostic (?:center|centre)|clinical lab", "clinical and diagnostic laboratories"),
    (r"government|public sector|ministry of health|\bngo\b", "government health institutions and NGOs"),
    (r"direct.to.consumer|\bd2c\b|consumers directly|online pharmacy|telehealth", "consumers directly"),
    (r"hospitals? and clinics|hospitals?,? clinics?|healthcare providers", "hospitals and healthcare providers"),
    (r"animal (?:health|production)|livestock", "veterinary and animal-health providers"),
]
_DEFAULT_SEGMENT = "hospitals and healthcare providers"


def derive_focus_areas(text: str, max_items: int = 3) -> str:
    low = text.lower()
    hits: list[str] = []
    for pattern, label in _FOCUS_RULES:
        if label in hits:
            continue
        if re.search(pattern, low):
            hits.append(label)
        if len(hits) >= max_items:
            break
    if hits:
        return "; ".join(hits)
    for pattern, label in _FALLBACK_RULES:
        if re.search(pattern, low):
            return label
    # No clinical-focus signal at all (e.g. a freight/logistics or
    # industrial-supply company with no healthcare specialty) -- use a
    # generic, honest, non-healthcare-assuming label. Deliberately NOT a
    # copy of Core Categories, to avoid restating the same column twice.
    return _DEFAULT_FALLBACK


def derive_segment(text: str) -> str:
    low = text.lower()
    for pattern, phrase in _SEGMENT_RULES:
        if re.search(pattern, low):
            return phrase
    return _DEFAULT_SEGMENT


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}

    core_remapped = 0
    focus_updated = 0
    skipped_flagged = 0

    for r in range(2, ws.max_row + 1):
        company = ws.cell(row=r, column=idx["Company"]).value
        if not company:
            continue

        current_core = str(ws.cell(row=r, column=idx["Core Categories"]).value or "").strip()
        if current_core.startswith(FLAG_PREFIX):
            skipped_flagged += 1
            continue

        ckpt_path = CKPT_DIR / f"{r:04d}.json"
        original_focus = ""
        if ckpt_path.exists():
            try:
                original_focus = json.loads(ckpt_path.read_text(encoding="utf-8")).get("specialty_focus", "")
            except Exception:
                original_focus = ""
        if original_focus.startswith(FLAG_PREFIX):
            skipped_flagged += 1
            continue

        summary = str(ws.cell(row=r, column=idx["Summary"]).value or "").strip()
        source_text = f"{original_focus} {summary}"

        new_core = remap_core_categories(current_core)
        if new_core:
            ws.cell(row=r, column=idx["Core Categories"], value=new_core)
            core_remapped += 1

        application = derive_focus_areas(source_text)
        segment = derive_segment(source_text)
        new_focus = f"{application} for {segment}."
        ws.cell(row=r, column=idx["Specialty Focus"], value=new_focus)
        focus_updated += 1

    wb.save(XLSX)
    print(
        f"Core Categories remapped to wearable vocabulary: {core_remapped}\n"
        f"Specialty Focus rebuilt: {focus_updated}\n"
        f"Skipped (flagged rows): {skipped_flagged}"
    )


if __name__ == "__main__":
    main()
