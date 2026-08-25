# -*- coding: utf-8 -*-
"""Web-search fallback for Contact Person / Role / LinkedIn on the Global Flexible
Packaging Market FINAL, for companies Apollo's free people search couldn't find
(scripts/_apollo_fill_flexpack_contacts.py).

Only sales / procurement leadership (same role filter as the Apollo script:
Head/VP Sales, Sales Director, Commercial Director, Head/VP Procurement,
Procurement Director, CRO/CPO, etc.).

Evidence sources (no DeepSeek — currently out of balance):
  1. Google AI Overview scraper (dedicated instance) — asks directly for a
     named sales/procurement leader + their LinkedIn profile.
  2. DDGS site:linkedin.com/in search — real search-result titles/snippets
     often already contain "Name - Title - Company | LinkedIn".

All-or-nothing per row: Name + Role come from real evidence text, but are only
WRITTEN once a matching linkedin.com/in/ URL for that exact person is also
confirmed (name tokens cross-checked against the LinkedIn URL slug itself,
the most reliable ground truth available). Freeform AI-overview text alone
produced too many false positives (title phrases, company-name fragments,
truncated words) to trust without that cross-check — so no LinkedIn match
means no contact written for that row, rather than a name we can't verify.

No subagents — plain script, run directly:
  .venv\\Scripts\\python.exe scripts\\web_search_flexpack_contacts.py
  .venv\\Scripts\\python.exe scripts\\web_search_flexpack_contacts.py --limit 5   # smoke test
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from dotenv import load_dotenv  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

load_dotenv(ROOT / ".env")

from vendor_intel.clients.ddgs_engines import normalize_text_backends  # noqa: E402
from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

# Reuse the exact role/name filters already validated on this dataset.
import importlib

_apollo_mod = importlib.import_module("_apollo_fill_flexpack_contacts")
_role_is_target = _apollo_mod._role_is_target
_name_ok = _apollo_mod._name_ok
_empty = _apollo_mod._empty
BAD_NAME = _apollo_mod.BAD_NAME

WORKDIR = ROOT / "output" / "web_expand" / "global_flexible_packaging_market_global"
XLSX = WORKDIR / "global_flexible_packaging_market_global_FINAL.xlsx"
CKPT = WORKDIR / "_flexpack_web_contacts_ckpt.json"
AUDIT = WORKDIR / "_flexpack_web_contacts_audit.json"
NPD = "Not publicly disclosed"

BACKEND_PARAM, BACKEND_NAMES, _DROPPED = normalize_text_backends()
GAI_BASE_URL = (os.getenv("GAI_BASE_URL") or "http://127.0.0.1:15551").rstrip("/")

TITLE_WORDS = (
    "VP Sales", "Vice President of Sales", "Head of Sales", "Sales Director",
    "Director of Sales", "Commercial Director", "Chief Revenue Officer", "CRO",
    "Head of Business Development", "VP Business Development", "Channel Sales",
    "Head of Procurement", "VP Procurement", "Vice President of Procurement",
    "Procurement Director", "Director of Procurement", "Chief Procurement Officer",
    "Procurement Manager", "Purchasing Manager", "Head of Purchasing",
)

LI_PERSON_RE = re.compile(r"https?://(?:[a-z]{2,3}\.)?linkedin\.com/in/[A-Za-z0-9\-_%]+/?", re.I)

# "Name - Title - Company | LinkedIn" or "Name | LinkedIn" style result titles
NAME_TITLE_RE = re.compile(
    r"^([A-Z][a-zA-Z.'\-]+(?:\s+[A-Z][a-zA-Z.'\-]+){0,3})\s*[-|\u2013\u2014]\s*([^|]{4,80})",
)


def _needs_contact(cp: str, role: str, li: str) -> bool:
    if _empty(cp) or BAD_NAME.match(cp):
        return True
    if _empty(role) or not _role_is_target(role):
        return True
    if "linkedin.com/in/" not in li.lower():
        return True
    return False


def ddg_search(query: str, max_results: int = 6) -> list[dict]:
    try:
        from ddgs import DDGS

        with DDGS() as d:
            return list(d.text(query, backend=BACKEND_PARAM, max_results=max_results)) or []
    except Exception as e:  # noqa: BLE001
        print(f"    ddgs error ({query[:60]}...): {e}", flush=True)
        return []


def extract_from_ddgs(hits: list[dict], company: str) -> dict | None:
    for h in hits:
        href = str(h.get("href") or "")
        title = str(h.get("title") or "")
        body = str(h.get("body") or "")
        m = LI_PERSON_RE.search(href) or LI_PERSON_RE.search(body)
        if not m:
            continue
        li_url = m.group(0).rstrip("/")
        blob = f"{title} {body}"
        nt = NAME_TITLE_RE.match(title.strip())
        name = ""
        role = ""
        if nt:
            name = nt.group(1).strip()
            role = nt.group(2).strip()
        else:
            for tw in TITLE_WORDS:
                if tw.lower() in blob.lower():
                    role = tw
                    first_words = title.split(" - ")[0].split(" | ")[0].strip()
                    if _name_ok(first_words):
                        name = first_words
                    break
        slug_name = name_from_linkedin_slug(li_url)
        if (
            name
            and len(name.split()) >= 2
            and role
            and _name_ok(name)
            and _role_is_target(role)
            and not _is_company_fragment(name, company)
            and slug_name
            and _names_closely_match(name, slug_name)
        ):
            final_name = slug_name if len(slug_name.split()) > len(name.split()) else name
            return {"contact_person": final_name, "contact_role": role, "linkedin": li_url, "source": "ddgs"}
    return None


_GAI_BAD_MARKERS = ("unusual traffic", "did not match", "ai overview is not available", "no ai overview")

_COMPANY_STOPWORDS = {
    "the", "and", "of", "co", "inc", "ltd", "llc", "gmbh", "group", "corp", "corporation",
    "company", "packaging", "packaging.", "supply", "supplies", "solutions", "distribution",
    "distributors", "international", "global", "products", "industries",
}

# Common non-name words that can end up capitalized at a sentence/line break
# (e.g. "...profile.\nAt Acme Paper...") — a real personal name never contains
# these as a standalone token.
_NON_NAME_WORDS = {
    "at", "in", "on", "for", "by", "with", "from", "to", "as", "is", "are",
    "was", "were", "a", "an", "this", "that", "their", "its", "his", "her",
    # job-title vocabulary — a real personal name never consists of these
    "vice", "president", "director", "head", "sales", "procurement", "commercial",
    "revenue", "chief", "business", "development", "channel", "purchasing",
    "manager", "officer", "national", "regional", "svp", "vp", "cro", "cpo",
    "partner", "team", "department", "division",
}


def _tokens(s: str) -> set[str]:
    return {t.lower() for t in re.findall(r"[A-Za-z']+", s or "") if len(t) > 1}


def _has_non_name_word(name: str) -> bool:
    return bool(_tokens(name) & _NON_NAME_WORDS)


def _is_company_fragment(name: str, company: str) -> bool:
    """True if `name` is really a piece of the company name / a title phrase,
    not a real person."""
    if _has_non_name_word(name):
        return True
    name_tokens = _tokens(name) - _COMPANY_STOPWORDS
    if not name_tokens:
        return True
    company_tokens = _tokens(company)
    if name_tokens.issubset(company_tokens):
        return True
    # A distinctive (long, non-generic) word from the company's own proper
    # noun leaking into the "name" — e.g. "Antalis Manage" for company
    # "Antalis" — is a sign of a bad extraction, not a real surname match.
    distinctive = {t for t in company_tokens if len(t) >= 5} - _COMPANY_STOPWORDS
    if name_tokens & distinctive:
        return True
    return False


def name_from_linkedin_slug(url: str) -> str:
    """.../in/ronald-attman-8a232656 -> 'Ronald Attman'"""
    m = re.search(r"/in/([A-Za-z0-9\-]+)", url)
    if not m:
        return ""
    slug = m.group(1)
    parts = [p for p in slug.split("-") if p and not re.fullmatch(r"[0-9a-f]{6,}", p, re.I)]
    parts = [p for p in parts if not p.isdigit()]
    if len(parts) < 2:
        return ""
    words = [p.capitalize() for p in parts[:3]]
    return " ".join(words)


def _names_share_token(a: str, b: str) -> bool:
    ta, tb = _tokens(a), _tokens(b)
    return bool(ta & tb)


async def gai_find_name_role(company: str, website: str) -> tuple[str, str] | None:
    """Phase 1: find a plausible sales/procurement leader's name + role.
    Google's AI Overview often names the LinkedIn profile as a link label
    ("Andy Attman LinkedIn Profile") without the raw URL appearing in the
    plain text — so this phase does NOT require a URL match."""
    host = re.sub(r"^https?://", "", (website or "").strip(), flags=re.I).split("/")[0]
    q = (
        f"{company} {host} current VP of Sales OR Head of Sales OR Sales Director OR "
        f"Commercial Director OR VP of Procurement OR Procurement Director name"
    )
    try:
        data = await google_ai_ask(q, base_url=GAI_BASE_URL, llm_clean=False)
    except Exception as e:  # noqa: BLE001
        print(f"    gai error: {e}", flush=True)
        return None
    md = str(data.get("markdown") or "").strip()
    if not md or any(m in md.lower() for m in _GAI_BAD_MARKERS):
        return None

    role_found = ""
    role_idx = -1
    for tw in TITLE_WORDS:
        idx = md.lower().find(tw.lower())
        if idx != -1:
            role_found = tw
            role_idx = idx
            break
    if not role_found or not _role_is_target(role_found):
        return None

    window = md[max(0, role_idx - 150) : role_idx + 50]
    name_candidates = re.findall(r"\b([A-Z][a-z]+(?:\s+[A-Z][a-z.\-]+){1,2})\b", window)
    for raw_cand in name_candidates:
        # Stop at the first sentence-ending period so a name never swallows
        # the start of the next sentence (e.g. "Brad Kroehn. The ..." -> "Brad Kroehn").
        cand = re.sub(r"\.\s.*$", "", raw_cand)
        cand = re.sub(r"\s+", " ", cand).strip().rstrip(".")
        if len(cand.split()) < 2:
            continue  # require a first + last name, not a single stray word
        if not _name_ok(cand) or BAD_NAME.match(cand):
            continue
        if _is_company_fragment(cand, company):
            continue
        return cand, role_found
    return None


def _names_closely_match(a: str, b: str) -> bool:
    """Stricter than sharing any single token — requires matching on the
    first name at minimum, and both names' tokens must substantially overlap."""
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return False
    a_first = a.strip().split()[0].lower()
    if a_first not in tb:
        return False
    return len(ta & tb) >= min(len(ta), len(tb))


async def find_linkedin_url(name: str, company: str) -> tuple[str, str]:
    """Phase 2: given a known real name, look up their specific LinkedIn URL.
    Returns (url, slug_name) or ("", "") if no confirmed match."""
    q = f'"{name}" "{company}" site:linkedin.com/in'
    hits = await asyncio.to_thread(ddg_search, q, 5)
    for h in hits:
        href = str(h.get("href") or "")
        m = LI_PERSON_RE.search(href)
        if not m:
            continue
        slug_name = name_from_linkedin_slug(m.group(0))
        if slug_name and _names_closely_match(name, slug_name):
            return m.group(0).rstrip("/"), slug_name

    # DDGS may be down — try GAI for the same targeted lookup.
    try:
        data = await google_ai_ask(
            f"{name} {company} LinkedIn profile URL", base_url=GAI_BASE_URL, llm_clean=False
        )
    except Exception:
        return "", ""
    md = str(data.get("markdown") or "")
    candidates = []
    for cite in data.get("citations") or []:
        u = str(cite.get("url") or cite) if isinstance(cite, dict) else str(cite)
        candidates.append(u)
    m = LI_PERSON_RE.search(md)
    if m:
        candidates.append(m.group(0))
    for u in candidates:
        um = LI_PERSON_RE.search(u)
        if not um:
            continue
        slug_name = name_from_linkedin_slug(um.group(0))
        if slug_name and _names_closely_match(name, slug_name):
            return um.group(0).rstrip("/"), slug_name
    return "", ""


async def gai_find_contact(company: str, website: str) -> dict | None:
    found = await gai_find_name_role(company, website)
    if not found:
        return None
    name, role = found
    li_url, slug_name = await find_linkedin_url(name, company)
    if not li_url:
        # No confirmed LinkedIn match for this name — too risky to write an
        # un-cross-validated name pulled from freeform AI-overview text.
        return None
    # Prefer the LinkedIn-derived name when it's more complete (e.g. handles
    # hyphenated surnames the free-text regex cut short).
    final_name = slug_name if len(slug_name.split()) > len(name.split()) else name
    return {"contact_person": final_name, "contact_role": role, "linkedin": li_url, "source": "gai"}


async def find_contact_for(company: str, website: str, gai_ok: bool) -> dict | None:
    if gai_ok:
        hit = await gai_find_contact(company, website)
        if hit:
            return hit
    for q in (
        f'site:linkedin.com/in "{company}" ("VP Sales" OR "Head of Sales" OR "Sales Director" OR "Commercial Director")',
        f'site:linkedin.com/in "{company}" ("VP Procurement" OR "Head of Procurement" OR "Procurement Director")',
    ):
        hits = await asyncio.to_thread(ddg_search, q, 6)
        hit = extract_from_ddgs(hits, company)
        if hit:
            return hit
    return None


def main() -> int:
    return asyncio.run(main_async())


async def main_async() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0)
    args = ap.parse_args()

    gai_ok = False
    try:
        gai_ok = await google_ai_ping(base_url=GAI_BASE_URL)
    except Exception:
        gai_ok = False
    print(f"Google AI scraper @ {GAI_BASE_URL}: {'UP' if gai_ok else 'DOWN'}", flush=True)
    print(f"DDGS backends: {BACKEND_NAMES}", flush=True)

    if not XLSX.exists():
        print(f"ERROR: not found: {XLSX}", file=sys.stderr)
        return 1

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_WEB_CONTACTS_{stamp}.xlsx")
    if not bak.exists():
        bak.write_bytes(XLSX.read_bytes())
        print(f"backup -> {bak.name}", flush=True)

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    targets: list[tuple[int, str]] = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        cp = str(ws.cell(r, ic["Contact Person"]).value or "").strip()
        role = str(ws.cell(r, ic["Role"]).value or "").strip()
        li = str(ws.cell(r, ic["LinkedIn"]).value or "").strip()
        if _needs_contact(cp, role, li):
            targets.append((r, co))
    print(f"targets still needing a sales/procurement contact: {len(targets)}", flush=True)

    ckpt: dict = {"done": {}}
    if CKPT.exists():
        try:
            ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
        except Exception:
            pass
    done = ckpt.setdefault("done", {})

    todo = [(r, co) for r, co in targets if co not in done]
    if args.limit:
        todo = todo[: args.limit]
    print(f"to process this run: {len(todo)} (already done: {len(targets) - len(todo)})", flush=True)

    CONCURRENCY = 6
    filled = 0
    processed = 0
    for start in range(0, len(todo), CONCURRENCY):
        batch = todo[start : start + CONCURRENCY]
        results = await asyncio.gather(
            *(
                find_contact_for(co, str(ws.cell(r, ic["Website"]).value or "").strip(), gai_ok)
                for r, co in batch
            )
        )
        for (r, co), hit in zip(batch, results):
            processed += 1
            if hit:
                ws.cell(r, ic["Contact Person"]).value = hit["contact_person"]
                ws.cell(r, ic["Role"]).value = hit["contact_role"]
                ws.cell(r, ic["LinkedIn"]).value = hit["linkedin"]
                filled += 1
                done[co] = {"applied": True, **hit}
                print(
                    f"[{processed}/{len(todo)}] {co} -> {hit['contact_person']} | "
                    f"{hit['contact_role']} | {hit['linkedin']} ({hit['source']})",
                    flush=True,
                )
            else:
                done[co] = {"applied": False}
                print(f"[{processed}/{len(todo)}] {co} -> no match", flush=True)

        wb.save(XLSX)
        CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"  checkpoint saved at {processed}/{len(todo)}", flush=True)

    wb.save(XLSX)
    CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")

    audit = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "processed_this_run": len(todo),
        "filled_this_run": filled,
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(audit, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
