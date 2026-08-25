#!/usr/bin/env python3
"""Re-apply web-verified overrides onto avocado OFFICE SLIM Excel (no re-fetch)."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global")
XLSX = ROOT / "global_avocado_oil_market_OFFICE_SLIM.xlsx"
OVERRIDES = ROOT / "_avocado_web_verified_overrides.json"
NPD = "Not publicly disclosed"

COUNTRY_RE = re.compile(
    r"\b(United States|United Kingdom|South Africa|New Zealand|Saudi Arabia|Hong Kong|South Korea|"
    r"Czech Republic|United Arab Emirates|USA|U\.S\.A\.|U\.S\.|UK|UAE|EU|EMEA|APAC|LATAM|GCC|"
    r"Australia|Canada|Mexico|Chile|Peru|Spain|Italy|France|Germany|Netherlands|India|China|"
    r"Japan|Brazil|Israel|Argentina|Colombia|Kenya|Morocco|Portugal|Poland|Sweden|Denmark|"
    r"Norway|Belgium|Switzerland|Ireland|Singapore|Malaysia|Thailand|Indonesia|Philippines|"
    r"Vietnam|Turkey|Greece|Egypt|Korea|Taiwan|Austria|Finland|Hungary|Romania|regional|"
    r"nationwide|global|American|Australian|Canadian|Mexican|Chilean|Peruvian|British|"
    r"German|French|Italian|Dutch|Indian|Chinese|Japanese|Brazilian|Israeli|Spanish|"
    r"African|European|Asian|Middle East(?:ern)?)\b",
    re.I,
)


def ensure_others(text: str) -> str:
    t = re.sub(r"\s+", " ", (text or "").strip(" ,;"))
    if not t:
        return "Food / ingredients distribution, others"
    if not re.search(r",?\s*others\s*$", t, re.I):
        t = t.rstrip(".,;") + ", others"
    return re.sub(r"(,\s*others)+\s*$", ", others", t, flags=re.I)


def strip_countries(text: str) -> str:
    t = COUNTRY_RE.sub(" ", text or "")
    t = re.sub(r"\s{2,}", " ", t)
    t = re.sub(r",\s*,+", ",", t)
    return t.strip(" ,;-")


def main() -> None:
    ov = json.loads(OVERRIDES.read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[2]]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}
    n = 0
    for row in range(3, ws.max_row + 1):
        co = ws.cell(row, ic["Company"]).value
        if not co or co not in ov:
            continue
        hit = ov[co]
        if hit.get("specialty"):
            ws.cell(row, ic["Specialty Focus"]).value = ensure_others(strip_countries(hit["specialty"]))
        if hit.get("office_no"):
            ws.cell(row, ic["Office No."]).value = hit["office_no"]
        if hit.get("office_email"):
            ws.cell(row, ic["Office Email"]).value = hit["office_email"]
        n += 1
    # ensure all specialty end with others and no country
    fixed_spec = 0
    for row in range(3, ws.max_row + 1):
        sp = str(ws.cell(row, ic["Specialty Focus"]).value or "")
        cleaned = ensure_others(strip_countries(sp))
        if cleaned != sp:
            ws.cell(row, ic["Specialty Focus"]).value = cleaned
            fixed_spec += 1
        # ensure office cols not blank
        if not str(ws.cell(row, ic["Office No."]).value or "").strip():
            ws.cell(row, ic["Office No."]).value = NPD
        if not str(ws.cell(row, ic["Office Email"]).value or "").strip():
            ws.cell(row, ic["Office Email"]).value = NPD
    wb.save(XLSX)
    wb.close()

    # recount
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c or "") for c in rows[1]]
    ic2 = {h: i for i, h in enumerate(hdr) if h}
    phone = email = country = 0
    for r in rows[2:]:
        if r[ic2["Office No."]] and str(r[ic2["Office No."]]).strip() != NPD:
            phone += 1
        if r[ic2["Office Email"]] and str(r[ic2["Office Email"]]).strip() != NPD:
            email += 1
        if COUNTRY_RE.search(str(r[ic2["Specialty Focus"]] or "")):
            country += 1
    print(
        json.dumps(
            {
                "overrides_applied": n,
                "specialty_rewrites": fixed_spec,
                "office_no_filled": phone,
                "office_email_filled": email,
                "specialty_country_left": country,
                "total": len(rows) - 2,
                "file": str(XLSX),
                "headers": hdr,
            },
            indent=2,
        )
    )
    # sample first 8
    for r in rows[2:10]:
        print(
            "-",
            r[ic2["Company"]],
            "|",
            r[ic2["Specialty Focus"]][:70],
            "|",
            r[ic2["Office No."]],
            "|",
            r[ic2["Office Email"]],
        )
    wb.close()


if __name__ == "__main__":
    main()
