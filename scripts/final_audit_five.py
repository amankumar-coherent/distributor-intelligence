#!/usr/bin/env python3
"""Final audit report for 5 market FINALs."""
from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
NPD = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd"}

FILES = [
    ("smartwatch", ROOT / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL.xlsx", "Smartwatch Companies",
     ("smartwatch", "wearable", "watch", "fitness tracker", "garmin", "amazfit", "fitbit", "band")),
    ("green", ROOT / "output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL.xlsx", "Green Chemical Companies",
     ("chemical", "green", "bio", "solvent", "surfactant", "specialty", "ingredient", "sustainable")),
    ("firewall", ROOT / "output/run4/global_firewall_market_global/global_firewall_market_FINAL.xlsx", "Firewall Companies",
     ("firewall", "ngfw", "security", "cyber", "network", "utm", "sase", "fortinet", "checkpoint")),
    ("smartphone", ROOT / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx", "Companies",
     ("smartphone", "mobile", "handset", "phone", "cellular", "android", "iphone", "telecom")),
    ("avocado", ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL.xlsx", "Avocado Oil Distributors",
     ("avocado", "oil", "foodservice", "grocery", "ingredient", "edible", "produce", "fmcg", "cosmetic")),
]

OEM = {
    "apple", "samsung", "garmin", "fitbit", "amazfit", "huawei", "xiaomi", "google", "cisco",
    "fortinet", "palo alto", "check point", "sophos", "watchguard", "sonicwall", "barracuda",
    "oppo", "vivo", "oneplus", "motorola", "nokia", "sony", "realme", "honor", "basf", "dow",
}
DIST = re.compile(
    r"distributor|wholesale|retail|channel|dealer|reseller|vad|var|partner|authorized|"
    r"marketplace|foodservice|importer|integrator|mssp|buying|e-commerce|cash|cooperative|"
    r"supplier|trader|broker|hypermarket",
    re.I,
)
MFR = re.compile(r"\bmanufacturer\b|\boem\b|\bdevice maker\b", re.I)


def empty(v: str) -> bool:
    return str(v or "").strip().lower() in NPD


def load(path: Path, sheet: str):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hi = None
    hdr = None
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            hi = i
            hdr = vals
            break
    data = []
    for r in rows[hi + 1 :]:
        if not r or not r[hdr.index("Company")]:
            continue
        data.append(
            {
                hdr[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
                for j in range(len(hdr))
                if hdr[j]
            }
        )
    return hdr, data


def main():
    report = []
    for key, path, sheet, kws in FILES:
        hdr, data = load(path, sheet)
        cols = [
            "Contact Person",
            "Role",
            "Email",
            "Phone / WhatsApp",
            "Founded",
            "Headquarters",
            "Website",
            "Business Type",
            "Specialty Focus",
            "Summary",
        ]
        empties = {c: sum(1 for d in data if empty(d.get(c))) for c in cols if c in (hdr or []) or True}
        mfrs = []
        oems = []
        weak = []
        for d in data:
            name = d.get("Company") or ""
            bt = d.get("Business Type") or ""
            blob = f"{name} {bt} {d.get('Specialty Focus') or ''} {d.get('Summary') or ''} {d.get('Core Categories') or ''}".lower()
            nl = re.sub(r"\s+", " ", name.lower())
            if MFR.search(bt) and not DIST.search(bt):
                mfrs.append(name)
            for oem in OEM:
                if nl == oem or nl.startswith(oem + " "):
                    if not DIST.search(name + " " + bt):
                        oems.append(name)
                        break
            if not any(k in blob for k in kws) and not DIST.search(bt):
                weak.append(name)
            elif not DIST.search(f"{bt} {name} {d.get('Specialty Focus') or ''}"):
                if key != "avocado":  # avocado has supplier/cooperative wording
                    weak.append(name)

        item = {
            "market": key,
            "path": str(path),
            "count": len(data),
            "empty": empties,
            "manufacturers_bt": mfrs,
            "oem_names": oems,
            "weak_or_non_channel_sample": weak[:20],
            "weak_count": len(set(weak)),
            "bt_top": Counter(d.get("Business Type") or "" for d in data).most_common(8),
            "contact_fill_pct": round(100 * (len(data) - empties.get("Contact Person", 0)) / max(1, len(data)), 1),
        }
        report.append(item)
        print(f"\n=== {key} count={item['count']} contact_fill={item['contact_fill_pct']}% ===")
        print(f"  empty contact={empties.get('Contact Person')} role={empties.get('Role')} email={empties.get('Email')} founded={empties.get('Founded')} hq={empties.get('Headquarters')} website={empties.get('Website')}")
        print(f"  mfr={mfrs} oem={oems} weak={item['weak_count']}")
        print(f"  BT={item['bt_top'][:5]}")

    out = ROOT / "output/five_markets_final_audit.json"
    out.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nWrote {out}")
    print("TOTAL", sum(r["count"] for r in report))


if __name__ == "__main__":
    main()
