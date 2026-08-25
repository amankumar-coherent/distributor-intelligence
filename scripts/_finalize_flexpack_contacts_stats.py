# -*- coding: utf-8 -*-
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

P = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
NPD = "Not publicly disclosed"


def main() -> None:
    wb = load_workbook(P)
    ws = wb.active
    h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {n: h.index(n) + 1 for n in h}

    for r in range(2, ws.max_row + 1):
        name = (ws.cell(r, idx["Company"]).value or "").strip()
        if name == "Office Depot":
            ws.cell(r, idx["Contact Person"]).value = "Ian Foster"
            ws.cell(r, idx["Role"]).value = "Regional Sales Director"
            ws.cell(r, idx["LinkedIn"]).value = (
                "https://www.linkedin.com/in/ian-foster-06427717"
            )
            print("set Office Depot")
        li = (ws.cell(r, idx["LinkedIn"]).value or "").strip()
        if "linkedin.com/company/" in li.lower():
            ws.cell(r, idx["LinkedIn"]).value = NPD

    wb.save(P)

    named = []
    person_li = 0
    npd = 0
    for r in range(2, ws.max_row + 1):
        name = (ws.cell(r, idx["Company"]).value or "").strip()
        c = (ws.cell(r, idx["Contact Person"]).value or "").strip()
        role = (ws.cell(r, idx["Role"]).value or "").strip()
        li = (ws.cell(r, idx["LinkedIn"]).value or "").strip()
        if c and c.lower() != NPD.lower() and not c.startswith("http") and "@" not in c:
            named.append(
                {"company": name, "contact": c, "role": role, "linkedin": li}
            )
            if "linkedin.com/in/" in li.lower():
                person_li += 1
        else:
            npd += 1

    summary = {
        "market": "global_flexible_packaging",
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "total_rows": ws.max_row - 1,
        "named_contacts": len(named),
        "with_person_linkedin": person_li,
        "npd_or_empty": npd,
        "policy": (
            "Head/VP Sales, Sales Director, Head/VP Procurement only; "
            "no invented emails; person /in/ only"
        ),
        "apollo_status": (
            "rate_limited_429_retry_after_~9h; skipped further Apollo this pass"
        ),
        "contacts": named,
    }
    out = P.parent / "flexpack_contacts_hand_verify_summary.json"
    out.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    wb.close()
    print("named", len(named), "person_li", person_li, "npd", npd)
    print("wrote", out)
    for x in named:
        print(x["company"], "|", x["contact"], "|", x["role"])


if __name__ == "__main__":
    main()
