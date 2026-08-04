from openpyxl import load_workbook
from pathlib import Path
import re

ROOT = Path(r"C:\Users\aman.kumar\Desktop\coherent\output\run4")


def count(path: Path):
    if not path.exists():
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    hr = 1
    headers = []
    for r in range(1, 8):
        vals = [c.value for c in ws[r]]
        if any(v and str(v).strip().lower() == "company" for v in vals):
            hr = r
            headers = vals
            break
    ci = next(i for i, h in enumerate(headers) if h and str(h).strip().lower() == "company")
    n = sum(1 for row in ws.iter_rows(min_row=hr + 1, values_only=True) if row and row[ci])
    wb.close()
    return n


for mid in [
    "global_smartwatch_market_global",
    "global_green_chemical_market_global",
    "global_firewall_market_global",
]:
    short = re.sub(r"_global$", "", mid)
    print(mid)
    for s in ["_FINAL.xlsx", "_founded_and_hq_v2.xlsx", "_founded_and_hq_v2_cleaned.xlsx"]:
        p = ROOT / mid / f"{short}{s}"
        print(" ", s, count(p))
