"""Remove Exclusivity and Partnership Duration from firewall FINAL Excel."""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

DROP = {"Exclusivity", "Partnership Duration"}
PATHS = [
    Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL.xlsx"),
    Path(
        r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL_320plus.xlsx"
    ),
]


def _header_row(ws) -> int:
    for r in range(1, 6):
        vals = [str(c.value or "") for c in ws[r]]
        if "Exclusivity" in vals or "Partnership Duration" in vals or "Company" in vals:
            return r
    return 3


def drop_cols(path: Path) -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_pre_drop_exclusivity_{stamp}{path.suffix}")
    shutil.copy2(path, bak)
    wb = openpyxl.load_workbook(path)
    touched: list[str] = []
    for name in wb.sheetnames:
        if name not in {"Firewall Companies"} and not name.startswith("Region -"):
            continue
        ws = wb[name]
        hr = _header_row(ws)
        hdr = [c.value for c in ws[hr]]
        to_del = [i for i, h in enumerate(hdr, 1) if str(h or "") in DROP]
        if not to_del:
            continue
        for i in reversed(to_del):
            ws.delete_cols(i, 1)
        new_hdr = [c.value for c in ws[hr]]
        touched.append(
            f"{name}: removed {[hdr[i - 1] for i in to_del]}; "
            f"cols now {len([h for h in new_hdr if h])}"
        )
    wb.save(path)
    print(path.name)
    print("  backup", bak.name)
    for line in touched:
        print(" ", line)
    print()


if __name__ == "__main__":
    for p in PATHS:
        if p.exists():
            drop_cols(p)
        else:
            print("missing", p)
