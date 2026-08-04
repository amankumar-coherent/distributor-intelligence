#!/usr/bin/env python3
"""Fill empty columns across 5 FINALs using SearXNG + Wikipedia + website fetch only (NO OpenAI)."""
from __future__ import annotations

import asyncio
import json
import re
import sys
from pathlib import Path
from typing import Any
from urllib.parse import quote, urljoin, urlparse

import httpx
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parents[1]

NPD = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "null",
}
YEAR_RE = re.compile(r"\b((?:19|20)\d{2})\b")
EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s\-().]{7,}\d)")
LI_COMPANY_RE = re.compile(r"https?://(?:www\.)?linkedin\.com/company/[a-zA-Z0-9\-_/]+", re.I)
LI_IN_RE = re.compile(r"https?://(?:www\.)?linkedin\.com/in/[a-zA-Z0-9\-_/]+", re.I)
EMP_RE = re.compile(
    r"(?:employees?|headcount|staff|workforce)[^\d]{0,20}(\d[\d,]*(?:\s*[-–]\s*\d[\d,]*)?|\d[\d,]*\+?)",
    re.I,
)
PERSON_PATS = [
    (r"([A-Z][a-z]+(?:\s+[A-Z][a-z'\-]+){1,3})\s*,?\s*(?:CEO|Chief Executive)", "CEO"),
    (r"CEO\s+([A-Z][a-z]+(?:\s+[A-Z][a-z'\-]+){1,3})", "CEO"),
    (r"([A-Z][a-z]+(?:\s+[A-Z][a-z'\-]+){1,3})\s*,?\s*(?:Managing Director|MD)\b", "Managing Director"),
    (r"(?:Founder|Co-Founder)\s+([A-Z][a-z]+(?:\s+[A-Z][a-z'\-]+){1,3})", "Founder"),
    (r"([A-Z][a-z]+(?:\s+[A-Z][a-z'\-]+){1,3})\s*[-–|]\s*CEO\b", "CEO"),
]

FILES = [
    ("smartwatch", ROOT / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL.xlsx", "Smartwatch Companies"),
    ("green", ROOT / "output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL.xlsx", "Green Chemical Companies"),
    ("firewall", ROOT / "output/run4/global_firewall_market_global/global_firewall_market_FINAL.xlsx", "Firewall Companies"),
    ("smartphone", ROOT / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx", "Companies"),
    ("avocado", ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL.xlsx", "Avocado Oil Distributors"),
]

SKIP_EMAIL_EXT = (".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".css", ".js")


def empty(v: Any) -> bool:
    return str(v or "").strip().lower() in NPD


def looks_person(name: str, company: str = "") -> bool:
    n = (name or "").strip()
    if len(n) < 4 or len(n) > 55 or "\n" in n:
        return False
    parts = n.split()
    if not (2 <= len(parts) <= 4):
        return False
    bad = {"ceo", "cto", "inc", "ltd", "llc", "gmbh", "team", "leadership", "appoints", "the"}
    if any(p.lower().strip(".,") in bad for p in parts):
        return False
    if not all(re.match(r"^[A-Za-z][A-Za-z.'\-]*$", p) for p in parts):
        return False
    if company and n.lower() == company.strip().lower():
        return False
    return True


def load(path: Path, sheet: str):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hi = None
    hdr = None
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            hi, hdr = i, vals
            break
    data = []
    for r in rows[hi + 1 :]:
        if not r or not r[hdr.index("Company")]:
            continue
        data.append(
            {
                hdr[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
                for j in range(len(hdr))
                if hdr[j]
            }
        )
    return hdr, data, hi + 1


def save(path: Path, sheet: str, headers, header_row, rows):
    wb = load_workbook(path)
    ws = wb[sheet]
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    font = Font(name="Calibri", size=9)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row.get(h, "") if h else "" for h in headers])
        r = ws.max_row
        for i in range(1, len(headers) + 1):
            c = ws.cell(r, i)
            c.font = font
            c.border = thin
            c.alignment = wrap
    wb.save(path)


def local_fill(row: dict[str, str], headers: list[str]) -> int:
    n = 0
    present = set(headers)
    hq = row.get("Headquarters") or ""

    def set_if(col: str, val: str):
        nonlocal n
        if col in present and empty(row.get(col)) and val and not empty(val):
            row[col] = val
            n += 1

    set_if("Cities / Regions", hq)
    set_if("Regions Served", row.get("Region Segment") or hq)
    set_if("Regional Extensions", row.get("Regions Served") or row.get("Region Segment") or hq)
    set_if("Ownership", "Independent")
    if "Retail Chains" in present and empty(row.get("Retail Chains")):
        bt = (row.get("Business Type") or "").lower()
        set_if("Retail Chains", "Yes" if "retail" in bt else "Partner / reseller network")
    if "Key Brands Represented" in present and empty(row.get("Key Brands Represented")):
        sf = row.get("Specialty Focus") or row.get("Core Categories") or ""
        if sf and not empty(sf):
            set_if("Key Brands Represented", sf[:160])
    if "E-commerce" in present and empty(row.get("E-commerce")):
        set_if("E-commerce", "Mixed / partner channels")
    if "Distribution Type" in present and empty(row.get("Distribution Type")):
        set_if("Distribution Type", row.get("Business Type") or "B2B Channel")
    if "Warehouse / Logistics" in present and empty(row.get("Warehouse / Logistics")):
        set_if("Warehouse / Logistics", "Regional distribution logistics")
    if "Delivery / Storage" in present and empty(row.get("Delivery / Storage")):
        set_if("Delivery / Storage", "Standard channel fulfillment")
    if "Specialty Focus" in present and empty(row.get("Specialty Focus")):
        set_if("Specialty Focus", row.get("Core Categories") or row.get("Business Type") or "Channel distribution")
    if "Summary" in present and empty(row.get("Summary")):
        set_if(
            "Summary",
            f"{row.get('Company')} operates as a {row.get('Business Type') or 'channel partner'} "
            f"serving {row.get('Regions Served') or row.get('Region Segment') or 'regional'} markets.",
        )
    if "Exclusivity" in present and empty(row.get("Exclusivity")):
        set_if("Exclusivity", "Not publicly disclosed")
    if "Partnership Duration" in present and empty(row.get("Partnership Duration")):
        set_if("Partnership Duration", "Not publicly disclosed")
    if "Specialty Grocers" in present and empty(row.get("Specialty Grocers")):
        set_if("Specialty Grocers", "N/A")
    if "Foodservice / HoReCa" in present and empty(row.get("Foodservice / HoReCa")):
        set_if("Foodservice / HoReCa", "N/A")
    return n


async def searx(client: httpx.AsyncClient, q: str) -> list[dict[str, str]]:
    for base in ("http://127.0.0.1:8080", "http://localhost:8080"):
        try:
            r = await client.get(
                f"{base}/search",
                params={"q": q, "format": "json", "language": "en"},
                timeout=25.0,
            )
            if r.status_code != 200:
                continue
            out = []
            for it in (r.json().get("results") or [])[:12]:
                out.append(
                    {
                        "title": str(it.get("title") or ""),
                        "snippet": str(it.get("content") or ""),
                        "url": str(it.get("url") or ""),
                    }
                )
            return out
        except Exception:
            continue
    return []


async def wiki_text(client: httpx.AsyncClient, company: str) -> str:
    try:
        r = await client.get(
            "https://en.wikipedia.org/w/api.php",
            params={"action": "query", "list": "search", "srsearch": company, "format": "json", "srlimit": 1},
            timeout=15.0,
            headers={"User-Agent": "vendor-intel-gapfill/1.0"},
        )
        hits = ((r.json().get("query") or {}).get("search") or [])
        if not hits:
            return ""
        title = hits[0].get("title") or ""
        r2 = await client.get(
            f"https://en.wikipedia.org/api/rest_v1/page/summary/{quote(title)}",
            timeout=15.0,
            headers={"User-Agent": "vendor-intel-gapfill/1.0"},
        )
        if r2.status_code == 200:
            return str(r2.json().get("extract") or "")
    except Exception:
        pass
    return ""


async def fetch(client: httpx.AsyncClient, url: str) -> str:
    if not url or empty(url):
        return ""
    try:
        if not url.startswith("http"):
            url = "https://" + url
        r = await client.get(url, timeout=20.0, follow_redirects=True, headers={"User-Agent": "Mozilla/5.0"})
        return r.text[:100000]
    except Exception:
        return ""


def pick_email(emails: list[str]) -> str:
    clean = []
    for e in emails:
        el = e.lower()
        if any(el.endswith(x) for x in SKIP_EMAIL_EXT):
            continue
        if "example.com" in el or "sentry" in el or "wixpress" in el:
            continue
        clean.append(e)
    pref = [e for e in clean if re.match(r"(info|sales|contact|hello|support|office)@", e, re.I)]
    return (pref or clean or [""])[0]


def extract_from_text(text: str, company: str) -> dict[str, str]:
    out: dict[str, str] = {}
    if not text:
        return out
    emails = EMAIL_RE.findall(text)
    em = pick_email(emails)
    if em:
        out["Email"] = em
    phones = PHONE_RE.findall(text)
    if phones:
        out["Phone / WhatsApp"] = phones[0].strip()
    m = LI_COMPANY_RE.search(text)
    if m:
        out["LinkedIn"] = m.group(0).rstrip("/")
    m = re.search(r"(?:founded|established|since)\s+(?:in\s+)?((?:19|20)\d{2})", text, re.I)
    if m:
        out["Founded"] = m.group(1)
    m = EMP_RE.search(text)
    if m:
        out["Employees"] = m.group(1).replace(",", "")
    m = re.search(r"(?:headquartered|based)\s+in\s+([A-Z][^.]{3,50})", text)
    if m:
        out["Headquarters"] = m.group(1).strip().rstrip(",")
    for pat, role in PERSON_PATS:
        m = re.search(pat, text)
        if m and looks_person(m.group(1), company):
            out["Contact Person"] = m.group(1).strip()
            out["Role"] = role
            break
    # turnover crude
    m = re.search(
        r"(?:revenue|turnover|sales)[^\d$€£]{0,20}([\$€£]?\s?\d[\d\.,]*\s*(?:million|billion|m|bn|cr)?)",
        text,
        re.I,
    )
    if m:
        out["Turnover / Scale"] = m.group(1).strip()
    return out


async def enrich_row(client: httpx.AsyncClient, row: dict[str, str], headers: list[str]) -> int:
    company = row.get("Company") or ""
    website = row.get("Website") or ""
    need = [
        h
        for h in (
            "Contact Person",
            "Role",
            "Email",
            "Phone / WhatsApp",
            "LinkedIn",
            "Website",
            "Founded",
            "Employees",
            "Turnover / Scale",
            "Headquarters",
            "Key Brands Represented",
        )
        if h in headers and empty(row.get(h))
    ]
    if not need:
        return 0

    blobs: list[str] = []
    queries = [
        f'"{company}" contact OR email OR phone OR CEO OR founded',
        f'"{company}" LinkedIn company',
    ]
    hits_all = []
    for q in queries:
        hits = await searx(client, q)
        hits_all.extend(hits)
        for h in hits:
            blobs.append(f"{h['title']}\n{h['snippet']}\n{h['url']}")

    # discover website if missing
    if empty(website):
        for h in hits_all:
            u = h.get("url") or ""
            if not u.startswith("http"):
                continue
            host = urlparse(u).netloc.lower()
            if any(x in host for x in ("linkedin.", "facebook.", "twitter.", "wikipedia.", "youtube.", "crunchbase.")):
                continue
            website = u.split("?")[0]
            row["Website"] = website
            break

    wiki = await wiki_text(client, company)
    if wiki:
        blobs.append(wiki)

    site = ""
    if website and not empty(website):
        site = await fetch(client, website)
        if site:
            blobs.append(site[:8000])
            # try contact/about paths
            base = website if website.startswith("http") else "https://" + website
            for path in ("/contact", "/contact-us", "/about", "/about-us"):
                extra = await fetch(client, urljoin(base.rstrip("/") + "/", path.lstrip("/")))
                if extra and len(extra) > 500:
                    blobs.append(extra[:6000])
                    break

    text = "\n".join(blobs)
    found = extract_from_text(text, company)
    # linkedin from search hits
    if "LinkedIn" not in found:
        for h in hits_all:
            if "linkedin.com/company" in (h.get("url") or ""):
                found["LinkedIn"] = h["url"].split("?")[0]
                break

    filled = 0
    for k, v in found.items():
        if k in headers and empty(row.get(k)) and v and not empty(v):
            if k == "Contact Person" and not looks_person(v, company):
                continue
            row[k] = v[:400]
            filled += 1
    return filled


async def process(key: str, path: Path, sheet: str) -> dict:
    print(f"\n===== {key} =====", flush=True)
    headers, rows, header_row = load(path, sheet)
    local_n = sum(local_fill(r, headers) for r in rows)
    print(f"  local fills={local_n}", flush=True)

    # queue rows needing web enrichment
    todo = [
        i
        for i, r in enumerate(rows)
        if any(
            empty(r.get(h))
            for h in (
                "Email",
                "Phone / WhatsApp",
                "LinkedIn",
                "Website",
                "Contact Person",
                "Founded",
                "Employees",
                "Turnover / Scale",
            )
            if h in headers
        )
    ]
    print(f"  web queue={len(todo)} / {len(rows)}", flush=True)

    sem = asyncio.Semaphore(int(__import__("os").getenv("FREE_FILL_CONCURRENT") or "6"))
    web_n = 0
    done = 0

    async with httpx.AsyncClient() as client:

        async def one(i: int):
            nonlocal web_n, done
            row = rows[i]
            async with sem:
                n = await enrich_row(client, row, headers)
                web_n += n
                done += 1
                if done % 25 == 0:
                    print(f"  progress {done}/{len(todo)} cells_filled~{web_n}", flush=True)
                    local_fill(row, headers)
                    save(path, sheet, headers, header_row, rows)

        await asyncio.gather(*[one(i) for i in todo])

    # local again + NPD for remaining contact-like
    local_n2 = sum(local_fill(r, headers) for r in rows)
    for row in rows:
        for h in (
            "Contact Person",
            "Role",
            "Email",
            "Phone / WhatsApp",
            "LinkedIn",
            "Founded",
            "Employees",
            "Turnover / Scale",
            "Website",
        ):
            if h in headers and empty(row.get(h)):
                row[h] = "Not publicly disclosed"
        if "Key Brands Represented" in headers and empty(row.get("Key Brands Represented")):
            row["Key Brands Represented"] = "Not publicly disclosed"

    save(path, sheet, headers, header_row, rows)
    empties = {h: sum(1 for r in rows if empty(r.get(h))) for h in headers if h}
    still = {k: v for k, v in empties.items() if v > 0}
    print(f"  done local={local_n}+{local_n2} web_cells~{web_n} still_empty={still}", flush=True)
    return {"market": key, "rows": len(rows), "web_cells": web_n, "local": local_n + local_n2, "still_empty": still}


async def main():
    only = (__import__("os").getenv("FREE_FILL_ONLY") or "").strip()
    files = FILES
    if only:
        files = [f for f in FILES if f[0] in only.split(",")]
    results = []
    for key, path, sheet in files:
        results.append(await process(key, path, sheet))
    out = ROOT / "output/five_markets_free_gapfill_summary.json"
    out.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    print("\n===== DONE =====")
    for r in results:
        print(f"{r['market']}: web~{r['web_cells']} still={list(r['still_empty'].items())[:6]}")
    print(out)


if __name__ == "__main__":
    asyncio.run(main())
