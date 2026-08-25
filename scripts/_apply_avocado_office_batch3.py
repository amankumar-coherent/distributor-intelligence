#!/usr/bin/env python3
"""Apply web-verified office contacts batch 3."""
from __future__ import annotations

import openpyxl
from pathlib import Path

p = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
    r"\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
wb = openpyxl.load_workbook(p)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

FILLS = {
    "US Foods": {
        "office_no": "(847) 720-8000",
        "office_email": "ir@usfoods.com",
    },
    "US Foods CHEF'STORE": {
        "office_no": "1-877-583-9659",
    },
    "KeHE Distributors": {
        "office_no": "(630) 343-0000",
        "office_email": "contactus@kehe.com",
    },
    "KeHE Specialty": {
        "office_no": "(630) 343-0000",
        "office_email": "contactus@kehe.com",
    },
    "IMCD": {
        "office_no": "+31 10 290 8684",
        "office_email": "info@imcdgroup.com",
    },
    "IMCD US": {
        "office_no": "+31 10 290 8684",
        "office_email": "info@imcdgroup.com",
    },
    "Metro AG": {
        "office_no": "+49 211 6886-0",
        "office_email": "imprint@metro.de",
    },
    "Sligro Food Group": {
        "office_no": "+31 413 34 35 00",
        "office_email": "info@sligrofoodgroup.nl",
    },
    "Univar Solutions": {
        "office_no": "+1 331-777-6000",
    },
    "Gordon Food Service": {
        "office_no": "(616) 530-7000",
    },
}


n = 0
for row in range(3, ws.max_row + 1):
    co = str(ws.cell(row, ic["Company"]).value or "").strip()
    if co not in FILLS:
        continue
    hit = FILLS[co]
    if hit.get("office_no"):
        ws.cell(row, ic["Office No."]).value = hit["office_no"]
    if hit.get("office_email"):
        ws.cell(row, ic["Office Email"]).value = hit["office_email"]
    n += 1
    print("filled", co)

wb.save(p)
wb.close()

wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c or "") for c in rows[1]]
ic2 = {h: i for i, h in enumerate(hdr) if h}
phone = sum(
    1
    for r in rows[2:]
    if r[ic2["Office No."]] and str(r[ic2["Office No."]]).strip() != "Not publicly disclosed"
)
email = sum(
    1
    for r in rows[2:]
    if r[ic2["Office Email"]] and str(r[ic2["Office Email"]]).strip() != "Not publicly disclosed"
)
print("TOTAL phone", phone, "email", email, "of", len(rows) - 2, "applied", n)
wb.close()
