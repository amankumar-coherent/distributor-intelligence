#!/usr/bin/env python3
"""Fast web-verified Key Brands fills + avocado-oil channel relevance cleanup."""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
AUDIT = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_relevance_audit.json"
)
NPD = "Not publicly disclosed"

# Agent web-verified Key Brands (edible oils / foodservice / grocery channel)
VERIFIED_BRANDS: dict[str, str] = {
    "Cheney Brothers": "Grand Western, Crazy Fish International, Duck Deli, Fronte, Bari Brand, Acento Latino, Tres Hermanos, Cheney Brothers Estate, Grumpy Potato, others",
    "Doerle Food Services": "Sysco Classic, Sysco Premium, Sysco Essentials, others",
    "Foodstuffs North Island Wholesale": "Pams, Pams Value, Pams Finest, others",
    "Gilmours": "Pams, Pams Value, Pams Finest, others",
    "PriceSmart": "Member's Selection, others",
    "JJ Foodservice": "Signature, Tabaki, Royal Roast, Super Crisp, Rice Majestic, others",
    "Reynolds Catering Supplies": "Carnivore, Shoal, Davin Foods, Solstice Select, Braehead Foods, others",
    "McLane Company": "CVP, Hometown Market, YumBees, Beau Dacious, Excursion Beef Jerky, Road-Tech, Work Fare, Divinely Sweet, Fly Guys Pizza, others",
    "Piggly Wiggly Alabama Distributing": "Piggly Wiggly, Sol Select, others",
    "H.T. Hackney": "Uncle Ray's, Great American Deli, others",
    "Associated Food Stores": "Food Club, Full Circle, TopCare, Buckley Farms, Culinary Tours, BirchBerry, Crav'n Flavor, others",
    "Affiliated Foods": "IGA, others",
    "Bozzuto's": "IGA, Hy-Top, Seven Farms, others",
    "C&S Wholesale Grocers": "Best Yet, Our Family, Fresh & Finest by Our Family, IGA, others",
    "Alex Lee": "Brown Bag, IGA, others",
    "Chefs' Warehouse": "BelAria Essentials, Grand Reserve, Allen Brothers, Proper Stock, others",
    "Associated Wholesale Grocers": "Always Save, Best Choice, Clearly by Best Choice, Superior Selections by Best Choice, Pure Wonder by Best Choice, Hello World by Best Choice, others",
    "Bestway Wholesale": "Best-in, others",
    "KeHE Distributors": "CADIA, MADE WITH, others",
    "KeHE Specialty": "CADIA, MADE WITH, others",
    "Wakefern Food Corp": "Bowl & Basket, Wholesome Pantry, Paperbird, others",
    "SpartanNash": "Our Family, Open Acres, Fresh & Finest by Our Family, others",
    "Metro AG": "METRO Chef, Aro, METRO Professional, Fine Life, Rioba, METRO Premium, others",
    "Makro": "METRO Chef, Aro, METRO Professional, Fine Life, Rioba, others",
    "Selgros Cash & Carry": "Transgourmet Economy, Transgourmet Quality, Transgourmet Ursprung, Transgourmet Natura, others",
    "Transgourmet": "Transgourmet Economy, Transgourmet Quality, Transgourmet Ursprung, Transgourmet Natura, Transgourmet Premium, others",
    "Restaurant Depot": "Chef's Quality, Sabor Nuestro, Big C, Supremo Italiano, Skyline, James Farm, others",
    "Jetro Restaurant Depot": "Chef's Quality, Sabor Nuestro, Big C, Supremo Italiano, Skyline, James Farm, others",
    "Sligro Food Group": "Dinnerland, Meesterschap van het huis, Kern, Sanissimo, others",
    "Brakes": "Sysco Essentials, Sysco Classic, Sysco Premium, La Boulangerie, Portico, Wholesome Farms, House Recipe, Citavo, Earth Plus, Country Choice, others",
    "Brake Bros": "Sysco Essentials, Sysco Classic, Sysco Premium, La Boulangerie, Portico, Wholesome Farms, House Recipe, Citavo, Earth Plus, Country Choice, others",
    "US Foods": "Chef's Line, Metro Deli, Rykoff Sexton, Stock Yards, Cross Valley Farms, Monarch, Monogram, Glenview Farms, Harbor Banks, Roseli, Molly's Kitchen, Harvest Value, Hilltop Hearth, others",
    "US Foods CHEF'STORE": "Chef's Line, Metro Deli, Rykoff Sexton, Stock Yards, Cross Valley Farms, Monarch, Monogram, others",
    "Cash & Carry Smart Foodservice": "Chef's Line, Metro Deli, Rykoff Sexton, Stock Yards, Cross Valley Farms, Monarch, Monogram, others",
    "Walmart": "Great Value, Marketside, Sam's Choice, others",
    "Gordon Food Service": "Markon First Crop, Ready-Set-Serve, Halperns', GFS, others",
    "Ben E. Keith Foods": "Markon First Crop, Ready-Set-Serve, Markon Essentials, others",
    "Shamrock Foods": "Markon First Crop, Ready-Set-Serve, Markon Essentials, others",
    "Nicholas and Company": "Markon First Crop, Ready-Set-Serve, Markon Essentials, others",
    "Performance Food Group": "Peak Fresh Produce, others",
    "PFG Customized Distribution": "Peak Fresh Produce, others",
    "Trader Joe's": "Trader Joe's, others",
    "The SYGMA Network": "Sysco Classic, Sysco Premium, Sysco Essentials, others",
    "FreshPoint": "Sysco Classic, Sysco Premium, others",
    "FreshPoint Central California": "Sysco Classic, Sysco Premium, others",
    "Fonterra Foodservice Distribution": "Anchor Food Professionals, NZMP, others",
    "Labatt Food Service": NPD,  # publicly states no private labels
    "Londis Wholesale": "Jack's, Euro Shopper, Chef's Larder, others",
    "Parfetts": "Go Local, Xenergy, others",
    "World Finer Foods": "Reese, DaVinci, Ty Ling, El Rio, London Pub, Wellington, Candoni, others",
    "Liberty Richter": "Reese, DaVinci, Ty Ling, others",
}

# Core-category tokens that strongly indicate NOT avocado-oil food channel
OFF_MARKET_CORE = re.compile(
    r"\b("
    r"cosmetic|cosmetics|personal care|soapmak|aromatherapy|fragrance|"
    r"oleochemical|specialty chemical|industrial chemical|laboratory|"
    r"pharma(?:ceutical)? ingredient|nutraceutical ingredient|"
    r"petrochemical|surfactant|emulsifier only|candle making|"
    r"diy beauty|skincare formulation|makeup ingredient"
    r")\s*",
    re.I,
)

# Companies known off avocado-oil edible channel (cosmetics chem / non-food)
OFF_MARKET_COMPANIES = {
    "Acme-Hardesty",
    "AIDP",
    "Aromantic",
    "Base Formula",
    "Bulk Apothecary",
    "Camden-Grey Essential Oils",
    "Charkit Chemical",
    "Charkit Chemical Company LLC",
    "Circle Ingredients",
    "From Nature With Love",
    "Liberty Natural Products",
    "Lotioncrafter",
    "MakingCosmetics",
    "Mystic Moments",
    "Naturallythinking",
    "Pacific Coast Chemicals",
    "KIC Chemicals",
    "Norkem",
    "Quimidroga",
    "Atlantic Chemicals Trading",
    "Azelis",  # specialty chem distributor — not edible avocado oil channel
    "IMCD",
    "IMCD US",
    "Univar Solutions",
    "Imperial Dade",  # packaging / janitorial, not oils food channel
}


def related_to_avocado_oil(core: str, focus: str, brands: str, company: str) -> tuple[bool, str]:
    if company in OFF_MARKET_COMPANIES:
        return False, "known_off_market_company"
    blob = f"{core} {focus} {brands}".lower()
    if OFF_MARKET_CORE.search(blob) and not re.search(
        r"edible oil|cooking oil|foodservice|grocery|fmcg|avocado|vegetable oil|specialty oil|olive oil|gourmet oil",
        blob,
        re.I,
    ):
        return False, "core/focus cosmetics-or-chem without edible-oil signal"
    # Positive signals keep
    if re.search(
        r"edible|cooking oil|foodservice|horeca|grocery|fmcg|avocado|vegetable oil|specialty oil|olive|gourmet|cash.?carry|wholesale food|distributor",
        blob,
        re.I,
    ):
        return True, "edible/food channel signal"
    if brands and brands != NPD:
        return True, "has brands; keep pending soft review"
    return True, "default keep"


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

    brand_fills = 0
    for r in range(3, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co not in VERIFIED_BRANDS:
            continue
        ws.cell(r, ic["Key Brands Represented"]).value = VERIFIED_BRANDS[co]
        brand_fills += 1

    removed = []
    kept_stats = {"related": 0, "removed": 0}
    # iterate bottom-up when deleting
    for r in range(ws.max_row, 2, -1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        core = str(ws.cell(r, ic["Core Categories"]).value or "")
        focus = str(ws.cell(r, ic["Specialty Focus"]).value or "")
        brands = str(ws.cell(r, ic["Key Brands Represented"]).value or "")
        ok, reason = related_to_avocado_oil(core, focus, brands, co)
        if not ok:
            removed.append({"company": co, "reason": reason, "core": core[:120], "brands": brands[:120]})
            ws.delete_rows(r, 1)
            kept_stats["removed"] += 1
        else:
            kept_stats["related"] += 1

    wb.save(XLSX)
    wb.close()

    # recount
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    ic2 = {str(h): i for i, h in enumerate(hdr) if h}
    real = npd = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        c = str(row[ic2["Company"]] or "").strip()
        if not c:
            continue
        b = str(row[ic2["Key Brands Represented"]] or "").strip()
        if b in ("", NPD):
            npd += 1
        else:
            real += 1
    total = real + npd
    wb.close()

    audit = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "brand_overrides_applied": brand_fills,
        "removed_off_market": removed,
        "removed_count": kept_stats["removed"],
        "kept_count": kept_stats["related"],
        "total_rows": total,
        "key_brands_real": real,
        "key_brands_npd": npd,
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: audit[k] for k in audit if k != "removed_off_market"}, indent=2))
    print("removed_names:", [x["company"] for x in removed])


if __name__ == "__main__":
    main()
