#!/usr/bin/env python3
"""Integrity audit of delivered markets (output/final/<batch>/<batch> distributors
(status checked)/). Runs automatically at the end of run_fire_10_browsers
finalize(); run by hand for all batches with no argument.

Per final workbook it checks:
  * row 2 carries the market's own "MARKET DEFINITION: ..." (not blank/generic)
  * no blank company names or websites
  * no company left twice (same name, or same website + HQ country)
  * no flagged company also in the main list
  * every HQ inside the market's required countries (COUNTRY_GATE below)
  * the JSON has the same number of companies as the Excel
  * no .zip anywhere in the batch folder (user: Excel + JSON only)

Exit code 0 when every file passes, 1 otherwise. Added 2026-10-21 after a
manual audit found duplicates the build had missed and stale Wound files.

Usage:
    python scripts/audit_final.py               # every batch
    python scripts/audit_final.py --batch soft  # one batch
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / "output" / "final"
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

# batch -> substrings one of which every Headquarters value must contain.
# A bare city (e.g. "Bangalore") is accepted when ACCEPTED_BARE_CITY allows it.
COUNTRY_GATE = {
    "geo": ("India",), "h2": ("India",), "spice": ("India",), "food": ("India",),
    "fire": ("India",), "brazil": ("Brazil", "Brasil"), "ent": ("United States", "USA"),
    "mg": ("Saudi", "United Arab Emirates", "UAE", "Kuwait", "Qatar", "Oman", "Bahrain",
           "Egypt", "Morocco", "Algeria", "Libya", "Sudan", "Tunisia", "Jordan", "Lebanon",
           "Iraq", "Yemen"),
}
ACCEPTED_BARE_CITY = {"india": ("bangalore", "bengaluru", "mumbai", "delhi", "new delhi", "chennai",
                                "hyderabad", "pune", "kolkata", "ahmedabad", "noida", "gurugram",
                                "gurgaon", "kochi", "jaipur", "lucknow", "indore", "surat")}


def _clean(s: object) -> str:
    return re.sub(r"\(acquired by.*\)$", "", str(s or ""), flags=re.I).strip()


def _dom(u: object) -> str:
    u = re.sub(r"^(https?://)?(www\.)?", "", str(u or "").strip().lower())
    return u.split("/")[0]


def audit_batch(batch: str) -> list[str]:
    """Problems found for one batch, one line each (empty = clean)."""
    import status_check_folder as S

    problems: list[str] = []
    batch_dir = FINAL / batch
    for z in batch_dir.rglob("*.zip"):
        problems.append(f"{batch}: zip file present: {z.name}")
    for out in sorted(batch_dir.glob("* distributors (status checked)")):
        flagged = set()
        ff = out / "flagged_for_review.xlsx"
        if ff.is_file():
            for r in list(openpyxl.load_workbook(ff, read_only=True).active.iter_rows(values_only=True))[1:]:
                if r and len(r) > 1 and r[1]:
                    flagged.add(S.name_key(_clean(r[1])))
        for x in sorted(out.glob("*_verified_distributors.xlsx")):
            tag = f"{batch}/{x.stem.replace('_verified_distributors', '')}"
            rows = list(openpyxl.load_workbook(x, read_only=True).active.iter_rows(values_only=True))
            a2 = str(rows[1][0] or "") if len(rows) > 1 else ""
            if not a2.startswith("MARKET DEFINITION:") or len(a2) < 150:
                problems.append(f"{tag}: row 2 market definition missing or too short")
            hdr = next((i for i, r in enumerate(rows[:12]) if r and "Company" in [str(c) for c in r]), None)
            if hdr is None:
                problems.append(f"{tag}: no header row")
                continue
            h = [str(c) for c in rows[hdr]]
            ci = h.index("Company")
            wi = h.index("Website") if "Website" in h else None
            hi = h.index("Headquarters") if "Headquarters" in h else None
            data = [r for r in rows[hdr + 1:] if r and any(r)]
            if any(not r[ci] for r in data):
                problems.append(f"{tag}: {sum(1 for r in data if not r[ci])} blank company names")
            if wi is not None and any(not r[wi] for r in data):
                problems.append(f"{tag}: {sum(1 for r in data if not r[wi])} blank websites")
            keys = [S.name_key(_clean(r[ci])) for r in data if r[ci]]

            def _country(r) -> str:
                hq = str(r[hi] or "") if hi is not None else ""
                return hq.split(",")[-1].strip().lower() if "," in hq else ""

            # Same name counts as a duplicate only inside one country (or when
            # a row has no country) -- the build's merge rule; same short name
            # in two countries is usually two companies ("OSP", "Sun Medical").
            seen_nc: dict[str, set] = {}
            dups = set()
            for r in data:
                if not r[ci]:
                    continue
                k, c = S.name_key(_clean(r[ci])), _country(r)
                prev = seen_nc.setdefault(k, set())
                # Same 4-character minimum as the build's merge: shorter names
                # ("OSP") are shared by unrelated companies too often.
                if len(k) >= 4 and (c in prev or (prev and (not c or "" in prev))):
                    dups.add(k)
                prev.add(c)
            if dups:
                problems.append(f"{tag}: same name twice: {sorted(dups)[:5]}")
            if wi is not None and hi is not None:
                pairs = [(_dom(r[wi]), _country(r)) for r in data if r[wi] and _country(r)]
                same = sorted({p for p in pairs if pairs.count(p) > 1})
                if same:
                    problems.append(f"{tag}: same website + country twice: {same[:3]}")
            leak = set(keys) & flagged
            if leak:
                problems.append(f"{tag}: {len(leak)} flagged companies also in the main list")
            gate = COUNTRY_GATE.get(batch)
            if gate and hi is not None:
                bare = {c for g in gate for c in ACCEPTED_BARE_CITY.get(g.lower(), ())}
                outside = [r[ci] for r in data
                           if not any(g.lower() in str(r[hi]).lower() for g in gate)
                           and str(r[hi]).split(",")[0].strip().lower() not in bare]
                if outside:
                    problems.append(f"{tag}: {len(outside)} HQ outside {'/'.join(gate[:3])}: {outside[:3]}")
            j = x.with_suffix(".json")
            if not j.is_file():
                problems.append(f"{tag}: JSON missing")
            else:
                d = json.loads(j.read_text(encoding="utf-8"))
                n = len(d.get("companies", [])) if isinstance(d, dict) else len(d)
                if n != len(data):
                    problems.append(f"{tag}: JSON has {n} companies, Excel has {len(data)}")
    return problems


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--batch", default="")
    args = ap.parse_args()
    batches = [args.batch] if args.batch else sorted(p.name for p in FINAL.iterdir() if p.is_dir())
    all_problems = []
    for b in batches:
        probs = audit_batch(b)
        all_problems += probs
        print(f"AUDIT {b}: " + ("OK" if not probs else f"{len(probs)} problem(s)"))
        for p in probs:
            print(f"   - {p}")
    sys.exit(1 if all_problems else 0)


if __name__ == "__main__":
    main()
