"""DDGS founded fill for remaining rows — NO Google AI scraper.

Accept year only if strong founded/since pattern appears in DDGS snippets
or About page text.
"""
from __future__ import annotations

import json
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
MISS = FP.parent / "founded_missing_152.json"
CKPT = FP.parent / "founded_fill_ddgs_only_ckpt.json"
AUDIT = FP.parent / "founded_fill_ddgs_only_audit.json"

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
STRONG = re.compile(
    r"(?:founded|established|est\.?|since|incorporated|gegr[uü]ndet|seit|"
    r"opgericht|fond[eée]|in business since|family[- ]owned since|"
    r"serving.{0,30}since|commenced trading|registered on|incorporation date)"
    r"\D{0,30}\b((?:18|19|20)\d{2})\b|"
    r"\b((?:18|19|20)\d{2})\b\D{0,20}(?:founded|established|gegr[uü]ndet)",
    re.I,
)


def fetch_about(website: str) -> str:
    if not website:
        return ""
    base = website if website.startswith("http") else "https://" + website
    base = base.rstrip("/")
    chunks = []
    for p in ["", "/about", "/about-us", "/history", "/company", "/unternehmen", "/ueber-uns"]:
        try:
            req = Request(base + p, headers={"User-Agent": UA})
            with urlopen(req, timeout=10) as resp:
                html = resp.read(100_000).decode("utf-8", errors="ignore")
            text = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>|<[^>]+>", " ", html, flags=re.I)
            text = re.sub(r"\s+", " ", text)
            chunks.append(text[:4000])
            if len(chunks) >= 2:
                break
        except Exception:
            continue
    return "\n".join(chunks)


def ddgs_text(query: str) -> str:
    try:
        from ddgs import DDGS
    except Exception as e:
        return f"[err]{e}"
    lines = []
    try:
        with DDGS() as ddg:
            for r in ddg.text(query, max_results=8):
                lines.append(f"{r.get('title','')} {r.get('body','')} {r.get('href','')}")
    except Exception as e:
        return f"[err]{e}"
    return "\n".join(lines)


def years(text: str) -> list[tuple[str, str]]:
    out, seen = [], set()
    for m in STRONG.finditer(text or ""):
        y = m.group(1) or m.group(2)
        if not y or y in seen:
            continue
        yi = int(y)
        if yi < 1800 or yi > 2026:
            continue
        seen.add(y)
        sn = text[max(0, m.start() - 40) : min(len(text), m.end() + 40)]
        out.append((y, sn.strip()))
    return out


def main() -> None:
    missing = json.loads(MISS.read_text(encoding="utf-8"))
    ckpt = json.loads(CKPT.read_text(encoding="utf-8")) if CKPT.exists() else {"done": {}}
    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    fi = headers.index("Founded") + 1
    row_by = {str(ws.cell(r, ci).value or "").strip(): r for r in range(2, ws.max_row + 1)}

    filled = 0
    audit = []
    for i, m in enumerate(missing, 1):
        company = str(m["company"])
        website = str(m.get("website") or "")
        if company in ckpt.get("done", {}):
            continue
        row = row_by.get(company)
        if not row:
            continue
        cur = str(ws.cell(row, fi).value or "").strip()
        if cur.isdigit():
            ckpt["done"][company] = {"year": cur, "src": "already"}
            continue

        print(f"[{i}/{len(missing)}] {company}", flush=True)
        about = fetch_about(website)
        domain = urlparse(website if "://" in website else "https://" + website).netloc.replace("www.", "")
        search = ddgs_text(f'"{company}" (founded OR established OR since OR gegründet) {domain}')
        time.sleep(1.0)

        about_hits = years(about)
        search_hits = years(search)
        year = None
        src = None
        if len(about_hits) == 1:
            year, sn = about_hits[0]
            src = "about:" + sn[:120]
        elif about_hits and search_hits:
            agree = {y for y, _ in about_hits} & {y for y, _ in search_hits}
            if len(agree) == 1:
                year = next(iter(agree))
                src = "agree"
        elif len(search_hits) == 1:
            year, sn = search_hits[0]
            src = "ddgs:" + sn[:120]
        elif len({y for y, _ in search_hits}) == 1 and search_hits:
            year = search_hits[0][0]
            src = "ddgs_consistent"

        rec = {"company": company, "year": year, "src": src, "about_hits": about_hits[:3], "search_hits": search_hits[:3]}
        audit.append(rec)
        ckpt["done"][company] = rec
        if year:
            ws.cell(row, fi).value = year
            filled += 1
            print(f"  -> {year} ({src})", flush=True)
        else:
            print("  -> NONE", flush=True)

        if i % 10 == 0:
            wb.save(FP)
            CKPT.write_text(json.dumps(ckpt, indent=2), encoding="utf-8")

    wb.save(FP)
    CKPT.write_text(json.dumps(ckpt, indent=2), encoding="utf-8")
    AUDIT.write_text(
        json.dumps(
            {"filled": filled, "updated": datetime.now().isoformat(timespec="seconds"), "audit": audit},
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"DONE filled={filled}")


if __name__ == "__main__":
    main()
