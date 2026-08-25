# -*- coding: utf-8 -*-
from pathlib import Path
import re
from openpyxl import load_workbook

PATHS = [
    Path(r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global\global_electrical_distribution_market_global_FINAL.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market.xlsx"),
]

CEO_RE = re.compile(r"\bceo\b|chief executive", re.I)
SALES_RE = re.compile(r"sales|commercial|procurement|purchas|revenue|\bcco\b|\bcro\b", re.I)


def load_rows(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals:
            headers = vals
            hr = r
            break
    rows = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {
            headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
        }
        if d.get("Company"):
            rows.append(d)
    wb.close()
    return rows


def main() -> None:
    print(f"{'market':<58} {'rows':>6} {'role':>6} {'ceo':>6}")
    tot_rows = tot_role = tot_ceo = 0
    for p in PATHS:
        rows = load_rows(p)
        role_n = ceo_n = 0
        for d in rows:
            role = d.get("Role") or ""
            if role:
                role_n += 1
            if CEO_RE.search(role):
                ceo_n += 1
        tot_rows += len(rows)
        tot_role += role_n
        tot_ceo += ceo_n
        print(f"{p.name:<58} {len(rows):6d} {role_n:6d} {ceo_n:6d}")
    print("-" * 80)
    print(f"{'TOTAL':<58} {tot_rows:6d} {tot_role:6d} {tot_ceo:6d}")


if __name__ == "__main__":
    main()
