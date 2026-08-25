"""Fast About-page founded extractor (no Google AI scraper, no DDGS).

Fetches homepage/about/history and accepts year only on strong founded/since patterns.
"""
from __future__ import annotations

import json
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
MISS = FP.parent / "founded_missing_152.json"
AUDIT = FP.parent / "founded_fill_about_fast_audit.json"

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 VendorIntelFounded/3.0"
STRONG = re.compile(
    r"(?:founded|established|est\.?|since|incorporated|gegr[uü]ndet|seit|"
    r"opgericht|fond[eée]|family[- ]owned since|in business since|"
    r"serving.{0,40}since|gegr\.?\s*(?:im)?)\D{0,25}\b((?:18|19|20)\d{2})\b|"
    r"\b((?:18|19|20)\d{2})\b\D{0,20}(?:founded|established|gegr[uü]ndet)",
    re.I,
)

# Agent web-verified (About / Companies House / Wikipedia / official history)
AGENT: dict[str, str] = {
    "Acme Paper & Supply": "1946",
    "Atlantic Packaging Products": "1945",
    "Ball & Doggett": "2017",
    "Ernest Packaging": "1946",
    "Abco Kovex": "1987",
    "Fillmore Container": "1995",
    "BigDug": "2004",
    "Box Brothers": "1985",
    "Spicers": "1896",
    "Winc": "1986",
    "Burch Bottle & Packaging": "1978",
    "All Packaging Company": "1946",
    "Durable Packaging": "1944",
    "Aropack": "2016",
    "Davpack": "1968",
    "Excel Pac": "1997",
    "GWP Group": "1990",
    "Hoffmann Verpackung": "1920",
}


def fetch_text(website: str) -> str:
    if not website:
        return ""
    base = website if website.startswith("http") else "https://" + website
    base = base.rstrip("/")
    paths = ["", "/about", "/about-us", "/about-us/", "/company", "/our-story", "/history",
             "/unternehmen", "/ueber-uns", "/a-propos", "/en/about"]
    chunks = []
    for p in paths:
        try:
            req = Request(base + p, headers={"User-Agent": UA})
            with urlopen(req, timeout=12) as resp:
                html = resp.read(140_000).decode("utf-8", errors="ignore")
            text = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
            text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
            text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"\s+", " ", text)
            chunks.append(text[:4500])
            if len(chunks) >= 3:
                break
        except Exception:
            continue
        time.sleep(0.15)
    return "\n".join(chunks)


def strong_years(text: str) -> list[tuple[str, str]]:
    out = []
    seen = set()
    for m in STRONG.finditer(text or ""):
        y = m.group(1) or m.group(2)
        if not y:
            continue
        yi = int(y)
        if yi < 1800 or yi > 2026 or y in seen:
            continue
        seen.add(y)
        sn = text[max(0, m.start() - 50): min(len(text), m.end() + 50)].strip()
        out.append((y, sn))
    return out


def main() -> None:
    missing = json.loads(MISS.read_text(encoding="utf-8"))
    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    fi = headers.index("Founded") + 1

    # map company -> excel row
    row_by = {}
    for r in range(2, ws.max_row + 1):
        row_by[str(ws.cell(r, ci).value or "").strip()] = r

    audit = {"filled": [], "ambiguous": [], "none": []}
    n = 0
    for i, m in enumerate(missing, 1):
        company = m["company"]
        website = m["website"]
        row = row_by.get(company)
        if not row:
            continue
        cur = str(ws.cell(row, fi).value or "").strip()
        if cur.isdigit():
            continue

        if company in AGENT:
            ws.cell(row, fi).value = AGENT[company]
            audit["filled"].append({"company": company, "year": AGENT[company], "src": "agent"})
            n += 1
            print(f"[{i}] AGENT {company} -> {AGENT[company]}", flush=True)
            continue

        print(f"[{i}/{len(missing)}] fetch {company}", flush=True)
        text = fetch_text(website)
        hits = strong_years(text)
        if len(hits) == 1:
            y, sn = hits[0]
            ws.cell(row, fi).value = y
            audit["filled"].append({"company": company, "year": y, "src": "about", "snip": sn[:160]})
            n += 1
            print(f"  ABOUT {y}", flush=True)
        elif len(hits) > 1:
            # if all same year
            years = {h[0] for h in hits}
            if len(years) == 1:
                y = next(iter(years))
                ws.cell(row, fi).value = y
                audit["filled"].append({"company": company, "year": y, "src": "about_multi_same"})
                n += 1
                print(f"  ABOUT {y}", flush=True)
            else:
                audit["ambiguous"].append({"company": company, "years": list(years), "hits": hits[:5]})
                print(f"  AMBIG {years}", flush=True)
        else:
            audit["none"].append(company)
            print("  NONE", flush=True)

        if i % 20 == 0:
            wb.save(FP)

    wb.save(FP)
    audit["filled_count"] = n
    audit["updated"] = datetime.now().isoformat(timespec="seconds")
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(f"DONE filled={n} ambiguous={len(audit['ambiguous'])} none={len(audit['none'])}")


if __name__ == "__main__":
    main()
