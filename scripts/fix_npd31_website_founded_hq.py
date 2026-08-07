#!/usr/bin/env python3
"""Fix remaining 31 NPD companies: curated web-verified + deep Google/DDG search."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import load_workbook
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")
os.environ["GOOGLE_AI_SCRAPER_ENABLED"] = "true"
os.environ["GOOGLE_AI_SCRAPER_FORCE_OVERVIEW"] = "false"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRIES"] = "1"

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

TODO = ROOT / "scripts/_npd31_todo.json"
OUT = ROOT / "scripts/_npd31_fix_progress.jsonl"
SUMMARY = ROOT / "scripts/_npd31_fix_summary.json"
CURRENT_YEAR = 2026
CONCURRENCY = 3

# Web-verified curated overrides (high confidence)
CURATED: dict[str, dict[str, str]] = {
    "PC Factory": {
        "Website": "https://www.pcfactory.cl",
        "Founded": "1996",
        "Headquarters": "Santiago, Chile",
    },
    "BTech": {
        "Website": "https://www.btech.com",
        "Founded": "1997",
        "Headquarters": "Cairo, Egypt",
    },
    "Brakes": {
        "Website": "https://www.brake.co.uk",
        "Founded": "1958",
        "Headquarters": "Ashford, Kent, United Kingdom",
    },
    "Bramble Berry": {
        "Website": "https://www.brambleberry.com",
        "Founded": "1999",
        "Headquarters": "Bellingham, Washington, United States",
    },
    "Henry Schein Medical": {
        "Website": "https://www.henryschein.com",
        "Founded": "1932",
        "Headquarters": "Melville, New York, United States",
    },
    "Aramark Middle East Distribution": {
        "Website": "https://www.aramark.com",
        "Founded": "1959",
        "Headquarters": "Philadelphia, Pennsylvania, United States",
    },
    "Mi Brasil": {
        "Website": "https://www.mi.com/br",
        "Founded": "2010",
        "Headquarters": "Beijing, China",
    },
    "Walmart Specialty Pharmacy": {
        "Website": "https://www.walmart.com",
        "Founded": "1962",
        "Headquarters": "Bentonville, Arkansas, United States",
    },
    "Alliance Healthcare Ireland": {
        "Website": "https://www.alliance-healthcare.com",
        "Founded": "1938",
        "Headquarters": "Dublin, Ireland",
    },
    "Alliance Healthcare Polska": {
        "Website": "https://www.alliance-healthcare.com",
        "Founded": "1938",
        "Headquarters": "Warsaw, Poland",
    },
    "Alliance Healthcare Portugal": {
        "Website": "https://www.alliance-healthcare.com",
        "Founded": "1938",
        "Headquarters": "Lisbon, Portugal",
    },
    "Pharmalog Austria": {
        "Website": "https://www.pharma-logistik.at",
        "Founded": "2006",
        "Headquarters": "Wels, Austria",
    },
}

MARKET_HINT = {
    "smartwatch": "electronics retail smartwatch distributor",
    "green_chemical": "green chemical distributor supplier",
    "avocado": "avocado oil foodservice distributor",
    "smartphone": "smartphone mobile phone distributor retailer",
    "antibody": "pharmaceutical specialty pharmacy drug distributor",
}

NAME_VARIANTS = {
    "NpuOne": ["NpuOne Malaysia", "NPU One electronics"],
    "Jaber & Sons / JABER": ["Jaber & Sons chemical Lebanon", "JABER chemicals trading"],
    "Horizon Harvest": ["Horizon Harvest avocado oil", "Horizon Harvest foods company"],
    "Pueblo Foods Distribution Europe": ["Pueblo Foods Europe distributor", "Pueblo Foods distribution"],
    "ASAR General Trading FZC": ["ASAR General Trading FZC UAE smartphones"],
    "B2BMobiles EU": ["B2BMobiles EU", "B2B Mobiles wholesale phones Europe"],
    "BuskaPhone": ["BuskaPhone mobile distributor"],
    "Kikaku Exports": ["Kikaku Exports phones Japan"],
    "MO2MO Wireless LLC": ["MO2MO Wireless LLC"],
    "Ariel Specialty Pharmacy": ["Ariel Specialty Pharmacy"],
    "BluePharmaMed": ["BluePharmaMed pharmacy"],
    "Drogueria Almus Uruguay": ["Drogueria Almus Uruguay"],
    "Drogueria Americana Paraguay": ["Drogueria Americana Paraguay"],
    "Highmark Specialty Pharmacy": ["Highmark Specialty Pharmacy"],
    "InnoveMed Healthcare": ["InnoveMed Healthcare"],
    "LeaderNET specialty pharmacy network": ["LeaderNET specialty pharmacy network"],
    "Pharmalog Austria": ["Pharmalog Austria GmbH"],
    "Victória Medicamentos": ["Victoria Medicamentos Brasil", "Victória Medicamentos"],
    "Les Aliments Dubeau": ["Les Aliments Dubeau Québec", "Aliments Dubeau"],
    "Modena Medicamentos": ["Modena Medicamentos Brasil"],
}


def market_key(filepath: str) -> str:
    f = filepath.lower().replace("\\", "/")
    for k in MARKET_HINT:
        if k in f or (k == "antibody" and "antibody" in f):
            return k
    if "green_chemical" in f:
        return "green_chemical"
    if "smartwatch" in f:
        return "smartwatch"
    if "avocado" in f:
        return "avocado"
    if "smartphone" in f:
        return "smartphone"
    if "antibody" in f:
        return "antibody"
    return ""


def norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [norm(c.value) for c in ws[r] if c.value not in (None, "")]
        if "company" in vals:
            return r
    return 1


def clean_url(u: object) -> str | None:
    s = str(u or "").strip()
    if not s or s.upper() in {"NPD", "N/A", "NA", "NONE", "UNKNOWN"}:
        return None
    if not re.match(r"^https?://", s, re.I):
        s = "https://" + s.lstrip("/")
    try:
        p = urlparse(s)
        if not p.netloc or "." not in p.netloc:
            return None
        host = p.netloc.lower()
        junk = (
            "reddit.com", "facebook.com", "archive.org", "google.com",
            "wikipedia.org", "tiktok.com", "linkedin.com", "crunchbase.com",
            "bloomberg.com", "zoominfo.com", "dnb.com", "yelp.com",
        )
        if any(j in host for j in junk):
            return None
        return f"{p.scheme}://{p.netloc}".rstrip("/")
    except Exception:
        return None


def year_ok(y: object) -> int | None:
    if y is None:
        return None
    if isinstance(y, int):
        yi = y
    else:
        m = re.search(r"(18|19|20)\d{2}", str(y))
        yi = int(m.group(0)) if m else None
    if yi is None:
        return None
    return yi if 1800 <= yi <= CURRENT_YEAR else None


def company_tokens(name: str) -> list[str]:
    stop = {
        "ltd", "llc", "inc", "gmbh", "co", "company", "corp", "the", "and",
        "group", "sa", "ag", "bv", "plc", "distribution", "distributor",
        "pharmacy", "specialty", "international", "global", "europe",
        "middle", "east", "network", "healthcare", "medical", "sons",
    }
    toks = re.findall(r"[a-z0-9]+", norm(name))
    return [t for t in toks if len(t) >= 3 and t not in stop]


def mentioned(text: str, company: str) -> bool:
    low = (text or "").lower()
    toks = company_tokens(company)
    if not toks:
        return True
    hits = sum(1 for t in toks[:4] if t in low)
    return hits >= (1 if len(toks) == 1 else min(2, len(toks)))


def not_found_signal(text: str) -> bool:
    low = (text or "").lower()
    patterns = [
        r"did not match any documents",
        r"no (?:major |official |widely documented )?(?:company|entity|website)",
        r"not (?:publicly )?(?:indexed|documented|available|found)",
        r"could not find",
        r"no results",
        r"unable to (?:find|locate)",
    ]
    return any(re.search(p, low) for p in patterns)


def deepseek() -> OpenAI:
    return OpenAI(
        api_key=os.getenv("OPENAI_API_KEY") or "",
        base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/"),
    )


def extract(company: str, text: str, source: str, hint: str) -> dict:
    if not (text or "").strip():
        return {"status": "empty", "source": source}
    if not_found_signal(text) and not mentioned(text, company):
        return {"status": "not_found", "source": source}

    client = deepseek()
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    prompt = f"""Extract firmographics ONLY if clearly about this exact company.

Company: {company}
Market context: {hint or "unknown"}

Search/AI text:
```
{text[:12000]}
```

Return JSON only:
{{
  "status": "found"|"not_found"|"unclear",
  "website": "https://..." or null,
  "founded_year": 1994 or null,
  "headquarters": "City, Country" or null,
  "confidence": "high"|"medium"|"low",
  "reason": "short"
}}
Rules:
- Prefer official corporate site of THIS company.
- For clear brand units (Xiaomi/Mi, Aramark, Henry Schein, Walmart, Alliance Healthcare, Brakes):
  official brand site + real founding year + HQ is OK.
- Wrong/similar-name company → status=unclear.
- No solid evidence → status=not_found.
- No LinkedIn/Wikipedia/directories as website.
- founded_year 1800-{CURRENT_YEAR}; never copyright-only years.
- Do not invent.
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    raw = (resp.choices[0].message.content or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.M).strip()
    try:
        data = json.loads(raw)
    except Exception:
        return {"status": "unclear", "source": source}
    if not isinstance(data, dict):
        return {"status": "unclear", "source": source}
    status = str(data.get("status") or "unclear").lower()
    conf = str(data.get("confidence") or "low").lower()
    out = {
        "status": status,
        "website": clean_url(data.get("website")),
        "founded": year_ok(data.get("founded_year")),
        "hq": (str(data.get("headquarters") or "").strip() or None),
        "confidence": conf,
        "reason": data.get("reason") or "",
        "source": source,
    }
    if conf == "low" and status == "found":
        out["status"] = "unclear"
    if status == "found" and not any([out["website"], out["founded"], out["hq"]]):
        out["status"] = "unclear"
    return out


def build_queries(company: str, filepath: str) -> list[str]:
    mk = market_key(filepath)
    hint = MARKET_HINT.get(mk, "")
    variants = NAME_VARIANTS.get(company, [company])
    qs: list[str] = []
    for v in variants[:3]:
        qs.append(f"{v} official website founded year headquarters")
        qs.append(f"{v} {hint} founded headquarters website".strip())
        qs.append(f"{v} company founded year?")
    seen, out = set(), []
    for q in qs:
        qn = q.strip().lower()
        if qn and qn not in seen:
            seen.add(qn)
            out.append(q.strip())
    return out[:6]


def ddg_text(company: str, filepath: str) -> str:
    try:
        from ddgs import DDGS
    except ImportError:
        try:
            from duckduckgo_search import DDGS  # type: ignore
        except ImportError:
            return ""
    bits = []
    try:
        with DDGS() as ddgs:
            for q in build_queries(company, filepath):
                for r in ddgs.text(q, max_results=5):
                    bits.append(
                        f"- {r.get('title','')}: {r.get('body') or r.get('description') or ''} | {r.get('href') or ''}"
                    )
    except Exception as e:
        bits.append(f"(ddg error: {e})")
    return "\n".join(bits)


async def google_bundle(company: str, filepath: str) -> str:
    texts = []
    for q in build_queries(company, filepath)[:3]:
        try:
            data = await google_ai_ask(q, llm_clean=True)
            md = str(data.get("markdown") or data.get("markdown_raw") or "").strip()
            if md:
                texts.append(f"[query] {q}\n{md}")
        except Exception as e:
            texts.append(f"[query] {q}\n(error {e})")
        await asyncio.sleep(0.3)
    return "\n\n".join(texts)


async def process_one(item: dict, sem: asyncio.Semaphore, google_ok: bool) -> dict:
    async with sem:
        company = item["company"]
        filepath = item["filepath"]
        if not Path(filepath).is_absolute():
            filepath = str(ROOT / filepath)

        if company in CURATED:
            print(f"→ {company} [CURATED]", flush=True)
            return {
                "company": company,
                "file": filepath,
                "parsed": {"status": "found", "confidence": "high", "source": "curated_web"},
                "action": "fix",
                "changes": dict(CURATED[company]),
            }

        hint = MARKET_HINT.get(market_key(filepath), "")
        print(f"→ {company}", flush=True)

        gtext = ""
        parsed: dict = {"status": "empty"}
        if google_ok:
            gtext = await google_bundle(company, filepath)
            parsed = await asyncio.to_thread(extract, company, gtext, "google_ai", hint)
            print(f"  google status={parsed.get('status')} conf={parsed.get('confidence')}", flush=True)

        if parsed.get("status") != "found":
            dtext = await asyncio.to_thread(ddg_text, company, filepath)
            bundle = f"GOOGLE:\n{gtext or '(none)'}\n\nDDG:\n{dtext}"
            parsed2 = await asyncio.to_thread(extract, company, bundle, "web+deepseek", hint)
            print(f"  residual status={parsed2.get('status')} conf={parsed2.get('confidence')}", flush=True)
            if parsed2.get("status") == "found":
                parsed = parsed2
            elif parsed.get("status") in {"empty", "unclear"}:
                parsed = parsed2

        changes: dict[str, str] = {}
        action = "still_npd"
        if parsed.get("status") == "found" and parsed.get("confidence") in {"high", "medium"}:
            if parsed.get("website"):
                changes["Website"] = parsed["website"]
            if parsed.get("founded") is not None:
                changes["Founded"] = str(parsed["founded"])
            if parsed.get("hq"):
                changes["Headquarters"] = parsed["hq"]
            if changes:
                for col in ("Website", "Founded", "Headquarters"):
                    if col not in changes:
                        changes[col] = "NPD"
                action = "fix"
        else:
            changes = {"Website": "NPD", "Founded": "NPD", "Headquarters": "NPD"}
            action = "still_npd"

        return {
            "company": company,
            "file": filepath,
            "parsed": parsed,
            "action": action,
            "changes": changes,
        }


def apply(results: list[dict]) -> None:
    by_file: dict[str, dict[str, dict]] = {}
    for r in results:
        if r.get("action") != "fix" or not r.get("changes"):
            continue
        by_file.setdefault(r["file"], {})[norm(r["company"])] = r["changes"]

    for fpath, mapping in by_file.items():
        wb = load_workbook(fpath)
        ws = wb.active
        hr = header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            company = ws.cell(r, ci + 1).value
            if not company:
                continue
            ch = mapping.get(norm(str(company)))
            if not ch:
                continue
            for col, val in ch.items():
                if col in headers:
                    ws.cell(r, headers.index(col) + 1).value = val
            n += 1
        wb.save(fpath)
        print(f"Updated {n} rows -> {Path(fpath).name}")


async def main() -> None:
    items = json.loads(TODO.read_text(encoding="utf-8"))
    for it in items:
        fp = it["filepath"]
        if not Path(fp).is_absolute():
            it["filepath"] = str(ROOT / fp)

    print(f"Deep re-search {len(items)} NPD companies ({len(CURATED)} curated)")
    google_ok = await google_ai_ping()
    print(f"Google AI scraper: {'connected' if google_ok else 'DOWN'}")

    if OUT.exists():
        OUT.unlink()

    sem = asyncio.Semaphore(CONCURRENCY)
    results: list[dict] = []
    with OUT.open("a", encoding="utf-8") as prog:
        for i in range(0, len(items), 8):
            batch = items[i : i + 8]
            print(f"\n=== Batch {i//8+1}/{(len(items)+7)//8} ===", flush=True)
            recs = await asyncio.gather(*[process_one(x, sem, google_ok) for x in batch])
            for rec in recs:
                prog.write(json.dumps(rec, ensure_ascii=False) + "\n")
                prog.flush()
                results.append(rec)
                print(f"  {rec['action'].upper()} {rec['company']}: {rec.get('changes')}", flush=True)

    print("\nApplying fixes to Excel...")
    apply(results)

    summary = {
        "total": len(results),
        "fixed_with_data": sum(1 for r in results if r["action"] == "fix"),
        "still_npd": sum(1 for r in results if r["action"] == "still_npd"),
        "curated": sum(1 for r in results if (r.get("parsed") or {}).get("source") == "curated_web"),
    }
    SUMMARY.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print("\n=== SUMMARY ===")
    for k, v in summary.items():
        print(f"{k}: {v}")


if __name__ == "__main__":
    asyncio.run(main())
