# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}

def cnt(p: Path) -> None:
    wb = load_workbook(p, data_only=True)
    ws = wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals:
            headers = vals
            hr = r
            break
    tot = cp = li = 0
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {
            headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
        }
        if not d.get("Company"):
            continue
        tot += 1
        person = (d.get("Contact Person") or "").strip()
        liu = (d.get("LinkedIn") or "").strip()
        if person and person.lower() not in NPD:
            cp += 1
            if "linkedin.com/in/" in liu.lower():
                li += 1
    wb.close()
    print(p.name, "rows", tot, "cp", cp, "li", li, "empty", tot - cp)


cnt(Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market.xlsx"))
cnt(Path(r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market.xlsx"))
