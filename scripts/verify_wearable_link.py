#!/usr/bin/env python3
"""Full-sheet re-audit: does each remaining company have a CONFIRMED link to
wearable medical devices (CGMs, smartwatches/trackers with medical claims,
wearable ECG/patches, wearable insulin pumps, hearing aids, sleep/respiratory
wearables, remote-patient-monitoring wearables, etc.) -- or is it just a
broadline pharma/medical distributor with no evidence of a wearable line?

Same stack as the rest of this task -- Google AI Overview scraper + ddgs,
no LLM anywhere, no subagents. This script only COLLECTS raw research text
into per-company checkpoints; it does NOT decide or write anything to the
workbook. A manual review pass reads the checkpoints and decides keep/remove.

Resumable: one JSON checkpoint per company under
output/web_expand/global_wearable_medical_devices_market_global/verify_checkpoints_wearlink/
"""
from __future__ import annotations

import asyncio
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

import openpyxl
from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping
from vendor_intel.enrichment.gap_fill.gaps import identity_text_matches

import verify_wearable_categories_brands as base  # reuse ddgs_search + filters

XLSX = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/global_wearable_medical_devices_market_global_FINAL.xlsx"
SHEET = "Wearable Medical Companies"
CKPT_DIR = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/verify_checkpoints_wearlink"
CKPT_DIR.mkdir(parents=True, exist_ok=True)

CONCURRENCY = int(os.getenv("WEARLINK_CONCURRENCY", "4"))

WEARABLE_TAGS = {
    "Medical Wearable Devices", "Health Monitoring Devices", "Diagnostic Wearables",
    "Remote Patient Monitoring", "Therapeutic Wearables",
}


async def research_company(row: dict) -> dict:
    company = row["company"]
    website = row["website"] or ""
    ckpt_path = CKPT_DIR / f"{row['row']:04d}.json"
    if ckpt_path.exists():
        try:
            return json.loads(ckpt_path.read_text(encoding="utf-8"))
        except Exception:
            pass

    chunks: list[str] = []
    if USE_GOOGLE_AI:
        queries = [
            f"does {company} sell or distribute wearable medical devices",
            f"{company} continuous glucose monitor OR smartwatch OR wearable ECG OR sleep tracker OR remote patient monitoring",
        ]
        for q in queries:
            try:
                data = await google_ai_ask(q, llm_clean=False)
                md = str(data.get("markdown") or "").strip()
                if (
                    data.get("ai_overview_missing")
                    or base._BLOCKED_PAGE_RE.search(md)
                    or base._AMBIGUOUS_ENTITY_RE.search(md)
                ):
                    continue
                if md and identity_text_matches(company, website, md):
                    chunks.append(md)
            except Exception as exc:
                print(f"  [google_ai] error for {company} ({q!r}): {exc}", flush=True)

    ddgs_hits = base.ddgs_search(
        f'"{company}" wearable OR "continuous glucose monitor" OR smartwatch OR "remote patient monitoring"',
        max_results=6,
    )
    lines = []
    for h in ddgs_hits:
        title = base._SITE_SUFFIX_RE.sub("", h.get("title", "")).strip()
        line = f"{title}: {h.get('body', '')}"
        if base._AMBIGUOUS_ENTITY_RE.search(line):
            continue
        if identity_text_matches(company, website, line):
            lines.append(line)
    if lines:
        chunks.append("\n".join(lines))

    context = "\n\n".join(chunks)
    lc = context.lower()
    wearable_kw = [
        "wearable", "smartwatch", "smart watch", "continuous glucose", "cgm",
        "glucose sensor", "wearable ecg", "wearable patch", "biosensor",
        "sleep tracker", "activity tracker", "remote patient monitoring",
        "wearable insulin", "hearing aid", "fitness tracker", "wrist-worn",
        "wrist worn", "patch sensor", "biopatch",
    ]
    hit_kw = sorted({kw for kw in wearable_kw if kw in lc})

    result = {
        "row": row["row"],
        "company": company,
        "had_research_text": bool(context.strip()),
        "wearable_keyword_hits": hit_kw,
        "raw_context": context[:4000],
    }
    ckpt_path.write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    return result


USE_GOOGLE_AI = True


async def main() -> None:
    global USE_GOOGLE_AI
    if not await google_ai_ping():
        print("WARNING: google-ai-scraper backend not reachable at :15551 -- "
              "falling back to ddgs-only for this run.", flush=True)
        USE_GOOGLE_AI = False

    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}

    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        company = ws.cell(row=r, column=idx["Company"]).value
        if not company:
            continue
        cc = str(ws.cell(row=r, column=idx["Core Categories"]).value or "")
        first_cat = cc.split(";")[0].strip()
        if first_cat in WEARABLE_TAGS:
            continue  # already confirmed wearable-tagged, skip
        rows.append(
            {
                "row": r,
                "company": str(company).strip(),
                "website": str(ws.cell(row=r, column=idx["Website"]).value or "").strip(),
            }
        )

    limit = int(os.getenv("WEARLINK_LIMIT", "0") or "0")
    if limit > 0:
        rows = rows[:limit]

    print(f"Companies to research: {len(rows)}", flush=True)

    sem = asyncio.Semaphore(CONCURRENCY)
    completed = 0

    async def _one(row: dict) -> None:
        nonlocal completed
        async with sem:
            result = await research_company(row)
            completed += 1
            print(
                f"[{completed}/{len(rows)}] row {row['row']} {row['company']!r} "
                f"kw_hits={result.get('wearable_keyword_hits')}",
                flush=True,
            )

    await asyncio.gather(*[_one(row) for row in rows])
    print("Research complete. Checkpoints written -- run manual review pass next.", flush=True)


if __name__ == "__main__":
    asyncio.run(main())
