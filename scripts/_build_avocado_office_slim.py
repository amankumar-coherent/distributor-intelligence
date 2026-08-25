#!/usr/bin/env python3
"""Build avocado oil OFFICE SLIM Excel (fast concurrent harvest + web overrides).

Source: global_avocado_oil_market.xlsx
Removes: Contact Person, Email, LinkedIn, Retail / E-commerce / Both
Adds: Office Email after Office No.
Specialty Focus: strip country tokens; end with ', others'; no invented phones/emails.
"""
from __future__ import annotations

import json
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global")
SRC = ROOT / "global_avocado_oil_market.xlsx"
OUT = ROOT / "global_avocado_oil_market_OFFICE_SLIM.xlsx"
CACHE = ROOT / "_avocado_office_harvest_cache.json"
OVERRIDES = ROOT / "_avocado_web_verified_overrides.json"
AUDIT = ROOT / "_avocado_office_slim_audit.json"

NPD = "Not publicly disclosed"

NEW_HEADERS = [
    "Company",
    "Website",
    "Founded",
    "Headquarters",
    "Continent / Geography",
    "Operational Presence",
    "Ownership",
    "Employees",
    "Core Categories",
    "Specialty Focus",
    "Key Brands Represented",
    "Distribution Type",
    "Role",
    "Office No.",
    "Office Email",
    "Country Code",
    "Region Code",
    "Summary",
]

WIDTHS = {
    "Company": 28,
    "Website": 32,
    "Specialty Focus": 44,
    "Office No.": 22,
    "Office Email": 30,
    "Summary": 40,
    "Headquarters": 24,
    "Operational Presence": 28,
    "Core Categories": 36,
    "Key Brands Represented": 32,
    "Distribution Type": 24,
    "Role": 22,
}

COUNTRY_RE = re.compile(
    r"\b("
    r"United States|United Kingdom|South Africa|New Zealand|Saudi Arabia|"
    r"Hong Kong|South Korea|Czech Republic|United Arab Emirates|"
    r"USA|U\.S\.A\.|U\.S\.|UK|UAE|EU|EMEA|APAC|LATAM|GCC|"
    r"Australia|Canada|Mexico|Chile|Peru|Spain|Italy|France|Germany|"
    r"Netherlands|India|China|Japan|Brazil|Israel|Argentina|Colombia|"
    r"Kenya|Morocco|Portugal|Poland|Sweden|Denmark|Norway|Belgium|"
    r"Switzerland|Ireland|Singapore|Malaysia|Thailand|Indonesia|"
    r"Philippines|Vietnam|Turkey|Greece|Egypt|Korea|Taiwan|Austria|"
    r"Finland|Hungary|Romania|Bulgaria|Croatia|Slovenia|Slovakia|"
    r"Ukraine|Russia|Nigeria|Ghana|Tanzania|Uganda|Ethiopia|"
    r"Ecuador|Guatemala|Costa Rica|Panama|Dominican|Puerto Rico|"
    r"Spanish|American|Australian|Canadian|Mexican|Chilean|Peruvian|"
    r"British|German|French|Italian|Dutch|Indian|Chinese|Japanese|"
    r"Brazilian|Israeli|Korean|Turkish|Moroccan|Kenyan|African|"
    r"European|Asian|Middle East(?:ern)?|Latin American|North American|"
    r"Southern African|Southeast Asian|regional|nationwide|global"
    r")\b",
    re.I,
)

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
PHONE_HINT_RE = re.compile(
    r"(?:tel:|phone|tel\.|call|fax|office|contact|hotline|customer\s*service)"
    r"[^0-9+]{0,24}(\+?[\d][\d\s().\-/]{7,22}\d)",
    re.I,
)
TEL_HREF_RE = re.compile(r"tel:([+\d][\d\-().\s]{6,})", re.I)
MAILTO_RE = re.compile(r"mailto:([a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,})", re.I)

BAD_EMAIL_PARTS = (
    "example.com",
    "sentry.",
    "wixpress",
    "schema.org",
    ".png",
    ".jpg",
    "google",
    "facebook",
    "cloudflare",
    "wordpress",
    "squarespace",
    "domain.com",
    "email.com",
    "yourcompany",
    "noreply",
    "no-reply",
    "donotreply",
    "sentry-next",
)

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)


def emptyish(v: object) -> bool:
    s = str(v or "").strip().lower()
    return s in {
        "",
        "n/a",
        "na",
        "none",
        "-",
        "—",
        "not publicly disclosed",
        "unknown",
        "tbd",
        "null",
    }


def normalize_website(url: str) -> str:
    u = (url or "").strip()
    if emptyish(u):
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    return u.rstrip("/")


def clean_phone(raw: str) -> str:
    s = re.sub(r"\s+", " ", str(raw or "").strip()).replace("tel:", "").strip()
    if len(re.sub(r"\D", "", s)) < 8:
        return ""
    return s[:40]


def clean_email(raw: str) -> str:
    e = str(raw or "").strip().lower().rstrip(".")
    if not e or "@" not in e:
        return ""
    if any(b in e for b in BAD_EMAIL_PARTS):
        return ""
    if e.endswith((".png", ".jpg", ".jpeg", ".gif", ".svg", ".css", ".js")):
        return ""
    return e


def strip_countries(text: str) -> str:
    t = COUNTRY_RE.sub(" ", text or "")
    t = re.sub(r"\s{2,}", " ", t)
    t = re.sub(r"\s+,", ",", t)
    t = re.sub(r",\s*,+", ",", t)
    t = re.sub(r"^\s*,\s*|\s*,\s*$", "", t)
    return t.strip(" ,;-")


def ensure_others(text: str) -> str:
    t = re.sub(r"\s+", " ", (text or "").strip(" ,;"))
    if not t:
        return "Food / ingredients distribution, others"
    if not re.search(r",?\s*others\s*$", t, re.I):
        t = t.rstrip(".,;") + ", others"
    t = re.sub(r"(,\s*others)+\s*$", ", others", t, flags=re.I)
    return t


def rewrite_specialty(existing: str, about_snip: str = "") -> str:
    base = strip_countries(existing or "")
    base = re.sub(r"\b(including|incl\.?|for|in|across|of)\s+(the\s+)?\b", " ", base, flags=re.I)
    base = re.sub(r"\s{2,}", " ", base).strip(" ,;-")
    blob = f"{base} {about_snip}".lower()
    bits: list[str] = []

    def add(x: str) -> None:
        xl = x.lower()
        if xl and all(xl not in b.lower() for b in bits):
            bits.append(x)

    if any(k in blob for k in ("oleochemical", "castor", "fatty acid", "glycerine", "glycerin")):
        add("Bio-based oleochemical & vegetable oil derivatives distribution")
    if any(k in blob for k in ("cosmetic", "personal care", "skincare", "skin care", "aromatherapy", "carrier oil")):
        add("Cosmetic / personal-care botanical & carrier oils")
    if any(k in blob for k in ("foodservice", "horeca", "chef", "culinary", "catering", "hospitality")):
        add("Foodservice / culinary edible oils supply")
    if any(k in blob for k in ("organic", "non-gmo", "non gmo")):
        add("Organic / natural edible oil assortment")
    if any(k in blob for k in ("ingredient", "formulat", "life science", "nutraceut", "supplement")):
        add("Food & specialty ingredient oils for manufacturers")
    if any(k in blob for k in ("wholesale grocer", "retailer-owned", "cash & carry", "cash and carry", "self-service wholesale")):
        add("Wholesale grocery oils assortment")
    if any(k in blob for k in ("import", "gourmet", "specialty food", "specialty oil")):
        add("Specialty / gourmet edible oils import & wholesale")
    if any(k in blob for k in ("avocado oil", "avocado")) and not bits:
        add("Edible avocado oil wholesale / distribution")
    if any(k in blob for k in ("produce", "fresh avocado", "fresh fruit")) and not any("oil" in b.lower() for b in bits):
        add("Fresh avocado / produce distribution (culinary oil adjacency)")
    if any(k in blob for k in ("fmcg", "packaged oil", "consumer packaged")):
        add("FMCG packaged edible oils distribution")

    if bits:
        return ensure_others("; ".join(bits[:3]))
    if base and len(base) >= 12:
        return ensure_others(base)
    return ensure_others("Food / ingredients & edible oils distribution")


def extract_from_html(html: str, base_url: str) -> tuple[str, str, str]:
    phones: list[str] = []
    emails: list[str] = []
    for m in TEL_HREF_RE.finditer(html):
        p = clean_phone(m.group(1))
        if p:
            phones.append(p)
    for m in MAILTO_RE.finditer(html):
        e = clean_email(m.group(1))
        if e:
            emails.append(e)
    for m in PHONE_HINT_RE.finditer(html):
        p = clean_phone(m.group(1))
        if p:
            phones.append(p)
    for m in EMAIL_RE.finditer(html):
        e = clean_email(m.group(0))
        if e:
            emails.append(e)

    text = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    text = re.sub(r"(?is)<style.*?>.*?</style>", " ", text)
    text = re.sub(r"(?is)<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text)
    snip = text[:2500]

    def uniq(xs: list[str]) -> list[str]:
        seen = set()
        out = []
        for x in xs:
            k = re.sub(r"\D", "", x) if "@" not in x else x
            if k in seen:
                continue
            seen.add(k)
            out.append(x)
        return out

    phone = uniq(phones)[0] if phones else ""
    ems = uniq(emails)
    host = urlparse(base_url).netloc.lower().removeprefix("www.")
    domain_match = [e for e in ems if host and host.split(":")[0] in e.split("@")[-1]]
    prefer_names = {
        "info",
        "contact",
        "sales",
        "office",
        "hello",
        "enquiries",
        "enquiry",
        "support",
        "customerservice",
        "customer.service",
        "ahcsc",
    }
    pool = domain_match or ems
    prefer = [e for e in pool if e.split("@")[0] in prefer_names]
    email = (prefer or pool)[0] if pool else ""
    return phone, email, snip


def fetch_one(website: str) -> dict:
    url = normalize_website(website)
    if not url:
        return {"phone": "", "email": "", "snip": "", "ok": False, "error": "no website", "url": ""}
    paths = ["", "/contact", "/contact-us", "/about"]
    phone = email = snip = ""
    last_err = ""
    try:
        with httpx.Client(
            headers={"User-Agent": UA},
            verify=False,
            timeout=httpx.Timeout(8.0, connect=4.0),
            follow_redirects=True,
        ) as client:
            for path in paths:
                try:
                    u = url if not path else urljoin(url + "/", path.lstrip("/"))
                    r = client.get(u)
                    if r.status_code >= 400:
                        last_err = f"HTTP {r.status_code}"
                        continue
                    p, e, s = extract_from_html(r.text or "", str(r.url))
                    phone = phone or p
                    email = email or e
                    snip = snip or s
                    if phone and email:
                        break
                except Exception as ex:  # noqa: BLE001
                    last_err = str(ex)[:120]
                    continue
    except Exception as ex:  # noqa: BLE001
        last_err = str(ex)[:120]
    return {
        "phone": phone,
        "email": email,
        "snip": snip[:2000],
        "ok": bool(phone or email or snip),
        "error": last_err,
        "url": url,
    }


def load_rows() -> list[dict]:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c or "").strip() for c in rows[1]]
    ic = {h: i for i, h in enumerate(hdr) if h}
    out = []
    for r in rows[2:]:
        if not r or not r[ic["Company"]]:
            continue
        out.append({h: ("" if r[ic[h]] is None else str(r[ic[h]]).strip()) for h in ic})
    wb.close()
    return out


def write_xlsx(rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Avocado Oil Distributors"
    title = "Global Avocado Oil Market — Distributors — OFFICE SLIM (no personal contact cols)"
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(NEW_HEADERS))
    ws.append(NEW_HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    body_font = Font(name="Calibri", size=9)

    ws["A1"].font = title_font
    for col, h in enumerate(NEW_HEADERS, 1):
        cell = ws.cell(2, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in rows:
        ws.append([row.get(h, "") for h in NEW_HEADERS])
        r = ws.max_row
        for c in range(1, len(NEW_HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.font = body_font
            cell.border = thin
            cell.alignment = wrap

    for i, h in enumerate(NEW_HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 18)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "A3"
    wb.save(OUT)


def main() -> None:
    source = load_rows()
    cache: dict = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            cache = {}

    overrides = {}
    if OVERRIDES.exists():
        overrides = json.loads(OVERRIDES.read_text(encoding="utf-8"))

    # harvest missing websites concurrently
    todo = []
    for row in source:
        web = row.get("Website", "")
        key = normalize_website(web) or row["Company"]
        if key not in cache:
            todo.append((row["Company"], web, key))

    print(f"rows={len(source)} cache={len(cache)} todo_fetch={len(todo)}", flush=True)

    done = 0
    with ThreadPoolExecutor(max_workers=12) as ex:
        futs = {ex.submit(fetch_one, web): (co, key) for co, web, key in todo}
        for fut in as_completed(futs):
            co, key = futs[fut]
            try:
                hit = fut.result()
            except Exception as exn:  # noqa: BLE001
                hit = {"phone": "", "email": "", "snip": "", "ok": False, "error": str(exn)[:120], "url": ""}
            cache[key] = hit
            done += 1
            if done % 25 == 0 or done == len(todo):
                CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
                print(f"fetched {done}/{len(todo)} latest={co[:40]}", flush=True)

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")

    out_rows = []
    phone_filled = email_filled = override_applied = 0
    for row in source:
        co = row["Company"]
        web = row.get("Website", "")
        key = normalize_website(web) or co
        hit = cache.get(key, {})

        existing_phone = "" if emptyish(row.get("Office No.")) else clean_phone(row.get("Office No.", ""))
        phone = existing_phone or hit.get("phone") or ""
        email = hit.get("email") or ""
        specialty = rewrite_specialty(row.get("Specialty Focus", ""), hit.get("snip", ""))

        ov = overrides.get(co)
        if ov:
            override_applied += 1
            if ov.get("specialty"):
                specialty = ensure_others(strip_countries(ov["specialty"]))
            if ov.get("office_no"):
                phone = ov["office_no"]
            if ov.get("office_email"):
                email = ov["office_email"]

        if phone:
            phone_filled += 1
        if email:
            email_filled += 1

        out_rows.append(
            {
                "Company": co,
                "Website": web,
                "Founded": row.get("Founded", ""),
                "Headquarters": row.get("Headquarters", ""),
                "Continent / Geography": row.get("Continent / Geography", ""),
                "Operational Presence": row.get("Operational Presence", ""),
                "Ownership": row.get("Ownership", ""),
                "Employees": row.get("Employees", ""),
                "Core Categories": row.get("Core Categories", ""),
                "Specialty Focus": specialty,
                "Key Brands Represented": row.get("Key Brands Represented", ""),
                "Distribution Type": row.get("Distribution Type", ""),
                "Role": row.get("Role", ""),
                "Office No.": phone if phone else NPD,
                "Office Email": email if email else NPD,
                "Country Code": row.get("Country Code", ""),
                "Region Code": row.get("Region Code", ""),
                "Summary": row.get("Summary", ""),
            }
        )

    write_xlsx(out_rows)

    # country leftover check
    with_country = sum(1 for r in out_rows if COUNTRY_RE.search(r["Specialty Focus"] or ""))

    audit = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "source": str(SRC),
        "output": str(OUT),
        "rows": len(out_rows),
        "removed_columns": ["Contact Person", "Email", "LinkedIn", "Retail / E-commerce / Both"],
        "added_columns": ["Office Email"],
        "office_no_filled": phone_filled,
        "office_email_filled": email_filled,
        "office_no_npd": len(out_rows) - phone_filled,
        "office_email_npd": len(out_rows) - email_filled,
        "web_overrides_applied": override_applied,
        "specialty_still_has_country_token": with_country,
        "rule": "No invented phones/emails — website extraction + verified overrides only; else Not publicly disclosed",
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(audit, indent=2))


if __name__ == "__main__":
    main()
