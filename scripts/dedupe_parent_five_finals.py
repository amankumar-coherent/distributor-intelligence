#!/usr/bin/env python3
"""Deduplicate parent/subsidiary variants across 5 FINALs; keep distributors/channel only."""
from __future__ import annotations

import json
import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parents[1]

FILES = [
    {
        "key": "smartwatch",
        "path": ROOT
        / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL.xlsx",
        "sheet": "Smartwatch Companies",
    },
    {
        "key": "green",
        "path": ROOT
        / "output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL.xlsx",
        "sheet": "Green Chemical Companies",
    },
    {
        "key": "firewall",
        "path": ROOT
        / "output/run4/global_firewall_market_global/global_firewall_market_FINAL.xlsx",
        "sheet": "Firewall Companies",
    },
    {
        "key": "smartphone",
        "path": ROOT
        / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx",
        "sheet": "Companies",
    },
    {
        "key": "avocado",
        "path": ROOT
        / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL.xlsx",
        "sheet": "Avocado Oil Distributors",
    },
]

# Longest-first parent keys for matching regional subsidiaries.
PARENT_KEYS: list[tuple[str, str]] = [
    # (match_key lowercase, canonical parent display name)
    ("exclusive networks", "Exclusive Networks"),
    ("network distributor (secureway)", "Network Distributor (Secureway)"),
    ("network distributors (secureway)", "Network Distributor (Secureway)"),
    ("secureway", "Network Distributor (Secureway)"),
    ("westcon-comstor", "Westcon-Comstor"),
    ("westcon comstor", "Westcon-Comstor"),
    ("westcon", "Westcon-Comstor"),
    ("arrow ecs", "Arrow ECS"),
    ("arrow electronics", "Arrow Electronics"),
    ("arrow enterprise computing", "Arrow ECS"),
    ("arrowecs", "Arrow ECS"),
    ("td synnex", "TD SYNNEX"),
    ("tdsynnex", "TD SYNNEX"),
    ("synnex westcon", "TD SYNNEX"),
    ("synnex", "TD SYNNEX"),
    ("tech data", "TD SYNNEX"),
    ("ingram micro", "Ingram Micro"),
    ("aptec (ingram", "Ingram Micro"),
    ("aptec distribution", "Ingram Micro"),
    ("aptec egypt", "Ingram Micro"),
    ("aptec ", "Ingram Micro"),
    ("also holding", "ALSO Holding"),
    ("also latvia", "ALSO Holding"),
    ("also schweiz", "ALSO Holding"),
    ("also cloud", "ALSO Holding"),
    ("infinigate", "Infinigate"),
    ("esprinet", "Esprinet"),
    ("redington", "Redington"),
    ("dicker data", "Dicker Data"),
    ("computacenter", "Computacenter"),
    ("bechtle", "Bechtle"),
    ("cancom", "Cancom"),
    ("softcat", "Softcat"),
    ("cdw", "CDW"),
    ("shi international", "SHI International"),
    ("presidio", "Presidio"),
    ("optiv", "Optiv Security"),
    ("world wide technology", "World Wide Technology"),
    ("wwt", "World Wide Technology"),
    ("softchoice", "Softchoice"),
    ("insight enterprises", "Insight Enterprises"),
    ("muk distribution", "MUK Distribution"),
    ("muk global", "MUK Distribution"),
    ("ndc pro", "NDC PRO"),
    ("softprom", "Softprom"),
    ("erc distribution", "ERC Distribution"),
    ("erc -", "ERC Distribution"),
    ("softline", "Softline"),
    ("elko group", "ELKO Group"),
    ("elko ", "ELKO Group"),
    ("asbis", "ASBIS"),
    ("scan source", "ScanSource"),
    ("scansource", "ScanSource"),
    ("carahsoft", "Carahsoft"),
    ("intcomex", "Intcomex"),
    ("logicom", "Logicom"),
    ("d&h distributing", "D&H Distributing"),
    ("d and h", "D&H Distributing"),
    ("avnet", "Avnet"),
    ("mouser", "Mouser Electronics"),
    ("rs components", "RS Components"),
    ("conrad electronic", "Conrad Electronic"),
    ("mediamarktsaturn", "MediaMarktSaturn"),
    ("media markt", "MediaMarktSaturn"),
    ("mediamarkt", "MediaMarktSaturn"),
    ("mediaworld", "MediaMarktSaturn"),
    ("saturn austria", "MediaMarktSaturn"),
    ("harvey norman", "Harvey Norman"),
    ("jb hi-fi", "JB Hi-Fi"),
    ("jb hifi", "JB Hi-Fi"),
    ("best buy", "Best Buy"),
    ("walmart", "Walmart"),
    ("amazon business", "Amazon"),
    ("amazon easy", "Amazon"),
    ("amazon ", "Amazon"),
    ("reliance digital", "Reliance Digital"),
    ("croma", "Croma"),
    ("yodobashi", "Yodobashi Camera"),
    ("bic camera", "Bic Camera"),
    ("the gioi di dong", "The Gioi Di Dong"),
    ("fpt shop", "FPT Shop"),
    ("f.biz", "FPT Shop"),
    ("coolblue", "Coolblue"),
    ("bol.com", "Bol.com"),
    ("fnac", "Fnac"),
    ("darty", "Darty"),
    ("currys", "Currys"),
    ("dixons", "Currys"),
    ("expert international", "Expert"),
    ("euronics", "Euronics"),
    ("elkjop", "Elkjop Nordic"),
    ("power nordic", "Power Nordic"),
    ("komplett", "Komplett Group"),
    ("netonnet", "NetOnNet"),
    ("alternate", "Alternate"),
    ("cyberport", "Cyberport"),
    ("notebooksbilliger", "Notebooksbilliger"),
    ("sharaf dg", "Sharaf DG"),
    ("virgin megastore", "Virgin Megastore"),
    ("lulu hypermarket", "Lulu Hypermarket"),
    ("jarir", "Jarir Bookstore"),
    ("magazine luiza", "Magazine Luiza"),
    ("americanas", "Americanas"),
    ("falabella", "Falabella"),
    ("liverpool", "Liverpool"),
    ("coppel", "Coppel"),
    ("yamada denki", "Yamada Denki"),
    ("edion", "Edion"),
    ("joshin", "Joshin Denki"),
    ("gome", "GOME"),
    ("suning", "Suning"),
    ("jd.com", "JD.com"),
    ("tmall", "Tmall"),
    ("challenger technologies", "Challenger"),
    ("courts asia", "Courts Asia"),
    ("power buy", "Power Buy"),
    ("jaymart", "Jaymart"),
    ("erafone", "Erafone"),
    ("digimap", "Digimap"),
    ("ibox indonesia", "iBox Indonesia"),
    ("beyond the box", "Beyond the Box"),
    ("power mac center", "Power Mac Center"),
    ("cellphones", "CellphoneS"),
    ("hoang ha mobile", "Hoang Ha Mobile"),
    ("nguyen kim", "Nguyen Kim"),
    ("worten", "Worten"),
    ("pccomponentes", "PcComponentes"),
    ("directd", "DIRECTD"),
    ("mmg distribution", "MMG Distribution"),
    ("vivid telecommunication", "Vivid Telecommunication"),
    ("ss solution", "SS Solution"),
    ("solutions 2 go", "Solutions 2 GO"),
    ("mountpole", "MountPole"),
    ("distrihub", "DistriHUB"),
    ("kapra falconeri", "Kapra Falconeri"),
    ("oss infocom", "OSS Infocom"),
    ("amit gps", "AMIT GPS & Navigation"),
    ("novitel", "Novitel"),
    ("digisanse", "Digisanse"),
    ("duna electronics", "Duna Electronics"),
    ("rec-line", "Rec-Line"),
    ("costco", "Costco"),
    ("sam's club", "Sam's Club"),
    ("sams club", "Sam's Club"),
    ("b&h photo", "B&H Photo"),
    ("adorama", "Adorama"),
    ("micro center", "Micro Center"),
    ("newegg", "Newegg"),
    ("target electronics", "Target"),
    ("target ", "Target"),
    ("fry's", "Fry's"),
    ("frys ", "Fry's"),
    ("axians", "Axians"),
    ("dimension data", "NTT / Dimension Data"),
    ("ntt data", "NTT / Dimension Data"),
    ("ntt ", "NTT / Dimension Data"),
    ("vodafone", "Vodafone Business"),
    ("vodacom", "Vodacom Business"),
]

OEM_EXACT = {
    "apple",
    "samsung",
    "garmin",
    "fitbit",
    "amazfit",
    "huawei",
    "xiaomi",
    "google",
    "cisco",
    "fortinet",
    "palo alto networks",
    "palo alto",
    "check point",
    "sophos",
    "juniper",
    "watchguard",
    "sonicwall",
    "barracuda",
    "zscaler",
    "cloudflare",
    "oppo",
    "vivo",
    "oneplus",
    "motorola",
    "nokia",
    "sony",
    "realme",
    "honor",
}

DIST_RE = re.compile(
    r"\b(distributor|distribution|wholesaler|wholesale|reseller|vad|var|channel|"
    r"partner|importer|authorized|value[\s-]?added|retailer|retail|dealer|"
    r"integrator|marketplace|buying group|mssp|cash\s*&\s*carry|foodservice|"
    r"broker|e-commerce|hypermarket|cooperative)\b",
    re.I,
)
MFR_RE = re.compile(
    r"\b(manufacturer|oem brand|device maker|handset maker|produces? only)\b",
    re.I,
)


def _norm(s: str) -> str:
    s = (s or "").lower().strip()
    s = s.replace("&", " and ")
    s = re.sub(r"[^\w\s\-./()]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


# Companies that contain a parent key but are distinct legal parents — do not collapse.
DISTINCT_KEEP = {
    "synnex technology international",
    "synnex technology international corp",
    "synnex technology international corporation",
    "micromax informatics",
    "micromax",
}


def parent_of(name: str) -> str | None:
    n = _norm(name)
    if n in DISTINCT_KEEP:
        return None
    # Require key as whole tokens / bounded phrase — avoid "croma" inside "micromax"
    for key, canon in sorted(PARENT_KEYS, key=lambda x: -len(x[0])):
        k = key.strip()
        if not k:
            continue
        # exact match or starts with key + space/punct
        if n == k or n.startswith(k + " ") or n.startswith(k + "-") or n.startswith(k + "("):
            return canon
        # bounded phrase: non-alnum or start before, non-alnum or end after
        pat = r"(?:^|[^a-z0-9])" + re.escape(k) + r"(?:[^a-z0-9]|$)"
        if re.search(pat, n):
            return canon
    return None


def is_oem(name: str) -> bool:
    n = _norm(name)
    for oem in OEM_EXACT:
        if n == oem or n.startswith(oem + " ") or f" {oem} " in f" {n} ":
            if not DIST_RE.search(name):
                # allow "Apple authorized retailer" style
                if any(x in n for x in ("authorized", "reseller", "distributor", "retail", "partner", "dealer")):
                    return False
                return True
    return False


def is_channel(row: dict[str, str]) -> bool:
    bt = row.get("Business Type") or ""
    focus = row.get("Specialty Focus") or ""
    summary = row.get("Summary") or ""
    blob = f"{bt} {focus} {summary} {row.get('Company') or ''}"
    if MFR_RE.search(bt) and not DIST_RE.search(bt):
        return False
    if is_oem(row.get("Company") or ""):
        return False
    if DIST_RE.search(blob):
        return True
    # if BT empty but known parent distributor family, keep
    if parent_of(row.get("Company") or ""):
        return True
    return False


def score_row(row: dict[str, str], canon: str) -> tuple:
    """Higher = better keep candidate for parent group."""
    name = row.get("Company") or ""
    n = _norm(name)
    c = _norm(canon)
    exact = 1 if n == c or n.startswith(c) and len(n) <= len(c) + 3 else 0
    # prefer shorter / cleaner parent-like names
    shortness = -len(name)
    filled = sum(
        1
        for k in (
            "Website",
            "Headquarters",
            "Founded",
            "Contact Person",
            "Email",
            "LinkedIn",
            "Summary",
        )
        if (row.get(k) or "").strip()
        and (row.get(k) or "").strip().lower() not in {"n/a", "not publicly disclosed", "-", ""}
    )
    # penalize region suffixes
    regional_penalty = 0
    if re.search(
        r"\b(gmbh|ltd|llc|s\.?a\.?|s\.?r\.?l\.?|bv|oy|ab|pty|sdn|fzco|fze|"
        r"africa|austria|germany|france|spain|uk|ireland|norway|sweden|denmark|"
        r"poland|slovakia|slovenia|croatia|finland|iceland|israel|mauritius|"
        r"nigeria|kenya|egypt|bahrain|oman|pakistan|jordan|lebanon|kuwait|"
        r"india|brazil|china|japan|korea|singapore|anz|latam|gulf|nordic|"
        r"north africa|south africa|belgium|netherlands|switzerland|portugal|"
        r"czech|romania|turkey|uae|middle east|asia|commercial|channel|security|"
        r"wearables|electronics|consumer)\b",
        n,
        re.I,
    ):
        regional_penalty = -5
    if exact:
        regional_penalty = 10
    return (exact, filled + regional_penalty, shortness)


def load_sheet(path: Path, sheet: str) -> tuple[list[str], list[dict[str, str]], int]:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hi = None
    headers: list[str] = []
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            hi = i
            headers = vals
            break
    assert hi is not None
    data: list[dict[str, str]] = []
    for r in rows[hi + 1 :]:
        if not r or not r[headers.index("Company")]:
            continue
        d = {
            headers[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
            for j in range(len(headers))
            if headers[j]
        }
        data.append(d)
    return headers, data, hi + 1  # 1-based header row


def rewrite_company_sheet(
    path: Path,
    sheet: str,
    headers: list[str],
    header_row: int,
    rows: list[dict[str, str]],
    audit: dict,
) -> None:
    wb = load_workbook(path)
    ws = wb[sheet]
    # delete data rows
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    font = Font(name="Calibri", size=9)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        live = [row.get(h, "") if h else "" for h in headers]
        ws.append(live)
        r = ws.max_row
        for i in range(1, len(headers) + 1):
            cell = ws.cell(r, i)
            cell.font = font
            cell.border = thin
            cell.alignment = wrap
        ws.row_dimensions[r].height = 30

    aname = "Dedupe Audit"
    if aname in wb.sheetnames:
        del wb[aname]
    aws = wb.create_sheet(aname)
    aws.append(["field", "value"])
    for k, v in audit.items():
        aws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v])
    wb.save(path)


def process(m: dict) -> dict:
    path: Path = m["path"]
    sheet = m["sheet"]
    print(f"\n===== {m['key']} | {path.name} =====", flush=True)
    headers, data, header_row = load_sheet(path, sheet)
    before = len(data)

    removed_mfr: list[str] = []
    channel_rows: list[dict[str, str]] = []
    for row in data:
        if not is_channel(row):
            removed_mfr.append(row.get("Company") or "")
            continue
        channel_rows.append(row)

    # Group by parent
    groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    singles: list[dict[str, str]] = []
    for row in channel_rows:
        p = parent_of(row.get("Company") or "")
        if p:
            groups[p].append(row)
        else:
            singles.append(row)

    kept: list[dict[str, str]] = list(singles)
    removed_dups: list[dict[str, str]] = []
    for canon, members in groups.items():
        if len(members) == 1:
            only = dict(members[0])
            # still normalize regional/affiliate naming to parent brand
            if _norm(only.get("Company") or "") != _norm(canon):
                removed_dups.append(
                    {
                        "kept": canon,
                        "removed": only.get("Company"),
                        "reason": "renamed_to_parent",
                    }
                )
                only["Company"] = canon
                src = only.get("Data Sources") or ""
                if "parent_dedupe" not in src:
                    only["Data Sources"] = (src + "; parent_dedupe").strip("; ")
            kept.append(only)
            continue
        ranked = sorted(members, key=lambda r: score_row(r, canon), reverse=True)
        winner = ranked[0]
        # normalize winner company name to parent
        winner = dict(winner)
        winner["Company"] = canon
        # merge regions served
        regions = []
        for r in members:
            for k in ("Region Segment", "Regions Served", "Headquarters"):
                v = (r.get(k) or "").strip()
                if v and v not in regions:
                    regions.append(v)
        if regions:
            winner["Regions Served"] = "; ".join(regions[:12])
            if not (winner.get("Regional Extensions") or "").strip():
                winner["Regional Extensions"] = winner["Regions Served"]
        # note in sources
        src = winner.get("Data Sources") or ""
        if "parent_dedupe" not in src:
            winner["Data Sources"] = (src + "; parent_dedupe").strip("; ")
        winner["Summary"] = (
            f"{canon} (parent channel company; consolidated {len(members)} regional/affiliate listings). "
            + (winner.get("Summary") or "")
        )[:450]
        kept.append(winner)
        for loser in ranked[1:]:
            removed_dups.append(
                {"kept": canon, "removed": loser.get("Company"), "reason": "subsidiary_or_regional_duplicate"}
            )

    # Exact-name dedupe safety
    seen = set()
    final: list[dict[str, str]] = []
    for row in kept:
        n = _norm(row.get("Company") or "")
        if not n or n in seen:
            removed_dups.append(
                {"kept": row.get("Company"), "removed": row.get("Company"), "reason": "exact_duplicate"}
            )
            continue
        seen.add(n)
        final.append(row)

    # stable sort by company
    final.sort(key=lambda r: (r.get("Company") or "").lower())

    bak = path.with_suffix(path.suffix + f".bak_dedupe_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
    shutil.copy2(path, bak)
    audit = {
        "before": before,
        "after": len(final),
        "removed_manufacturer_or_oem": removed_mfr,
        "removed_duplicates": removed_dups,
        "duplicate_groups_collapsed": len([g for g in groups.values() if len(g) > 1]),
        "backup": str(bak),
    }
    rewrite_company_sheet(path, sheet, headers, header_row, final, audit)
    # also write CLEANED copy
    cleaned = path.with_name(path.stem + "_DEDUPED.xlsx")
    shutil.copy2(path, cleaned)

    print(
        f"  {before} -> {len(final)}  "
        f"(removed mfr/oem={len(removed_mfr)}, collapsed_dups={len(removed_dups)})",
        flush=True,
    )
    if removed_dups[:8]:
        print("  sample dups removed:", flush=True)
        for d in removed_dups[:8]:
            print(f"    keep={d.get('kept')}  drop={d.get('removed')}", flush=True)
    return {"market": m["key"], **audit}


def main() -> None:
    results = [process(m) for m in FILES]
    out = ROOT / "output/five_markets_parent_dedupe_summary.json"
    out.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    print("\n===== TOTALS =====")
    for r in results:
        print(f"{r['market']}: {r['before']} -> {r['after']}")
    print(f"Summary: {out}")


if __name__ == "__main__":
    main()
