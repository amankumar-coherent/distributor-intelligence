"""Sort FINAL company sheets: Asia Pacific first, then other regions, companies A–Z within region.

Also rebuild Region Summary and Region sheets in the same order.
"""
from __future__ import annotations

import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Asia Pacific first per user request
REGION_ORDER = [
    "Asia Pacific",
    "Europe",
    "North America",
    "Middle East",
    "Latin America",
    "Africa",
    "Eurasia / CIS",
    "Global",
    "Other - Unclear",
]

WORKBOOKS = [
    (
        Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL.xlsx"),
        "Smartwatch Companies",
        "Global Smartwatch Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL_320plus.xlsx"
        ),
        "Smartwatch Companies",
        "Global Smartwatch Market — Region-wise segmentation",
    ),
    (
        Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL.xlsx"),
        "Firewall Companies",
        "Global Firewall Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL_320plus.xlsx"
        ),
        "Firewall Companies",
        "Global Firewall Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL.xlsx"
        ),
        "Green Chemical Companies",
        "Global Green Chemical Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL_320plus.xlsx"
        ),
        "Green Chemical Companies",
        "Global Green Chemical Market — Region-wise segmentation",
    ),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
CELL_FONT = Font(name="Calibri", size=9)
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP = Alignment(wrap_text=True, vertical="top")

NOTES = {
    "Asia Pacific": "India, China, Japan, SEA, ANZ, Korea",
    "Europe": "UK + EU / wider Europe (incl. Baltics, CEE)",
    "North America": "USA, Canada, Mexico",
    "Middle East": "GCC, UAE, Saudi, Turkey, Israel",
    "Latin America": "Brazil, Chile, Colombia, Argentina, Peru, Uruguay",
    "Africa": "South Africa, Nigeria, Egypt, Kenya, etc.",
    "Eurasia / CIS": "Russia / CIS channel partners",
    "Global": "Regions Served = Global / multi-region",
    "Other - Unclear": "Insufficient HQ/region text to classify",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _region_rank(region: str) -> int:
    try:
        return REGION_ORDER.index(region)
    except ValueError:
        return len(REGION_ORDER)


def _find_header_row(ws) -> int:
    for r in range(1, 6):
        vals = [str(c.value or "") for c in ws[r]]
        if "Company" in vals and "Region Segment" in vals:
            return r
    raise RuntimeError("Company/Region Segment header not found")


def _style_header(ws, headers: list[str]) -> None:
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = 18 if h != "Summary" else 36
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _write_rows(ws, headers: list[str], rows: list[dict[str, str]]) -> None:
    _style_header(ws, headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
        rr = ws.max_row
        for i in range(1, len(headers) + 1):
            cell = ws.cell(rr, i)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = WRAP
        ws.row_dimensions[rr].height = 32
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last}{ws.max_row}"


def sort_workbook(path: Path, sheet_name: str, summary_title: str) -> dict:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_pre_region_sort_{stamp}{path.suffix}")
    shutil.copy2(path, bak)

    wb = openpyxl.load_workbook(path)
    ws = wb[sheet_name]
    hr = _find_header_row(ws)
    headers = [str(c.value or "") for c in ws[hr]]
    ncols = len(headers)
    seg_i = headers.index("Region Segment")
    co_i = headers.index("Company")

    # Capture title rows above header (preserve formatting values)
    title_rows: list[list] = []
    for r in range(1, hr):
        title_rows.append([ws.cell(r, c).value for c in range(1, ncols + 1)])

    companies: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        if not row or not row[co_i]:
            continue
        d = {
            headers[i]: str(row[i] if i < len(row) and row[i] is not None else "")
            for i in range(len(headers))
        }
        companies.append(d)

    companies.sort(
        key=lambda d: (_region_rank(d.get("Region Segment", "")), _norm(d.get("Company", "")))
    )

    # Clear data area and rewrite sorted
    if ws.max_row > hr:
        ws.delete_rows(hr + 1, ws.max_row - hr)

    for d in companies:
        vals = [d.get(h, "") for h in headers]
        ws.append(vals)
        rr = ws.max_row
        for c in range(1, ncols + 1):
            cell = ws.cell(rr, c)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = WRAP

    # Rebuild Region Summary + Region sheets
    for name in list(wb.sheetnames):
        if name.startswith("Region - ") or name in {"Region Summary", "By Region"}:
            del wb[name]

    counts = Counter(d.get("Region Segment", "") for d in companies)
    total = len(companies)
    sum_ws = wb.create_sheet("Region Summary", 1)
    sum_ws["A1"] = summary_title
    sum_ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    sum_ws.merge_cells("A1:D1")
    sum_ws["A3"] = "Region Segment"
    sum_ws["B3"] = "Company Count"
    sum_ws["C3"] = "% of Total"
    sum_ws["D3"] = "Notes"
    for col in range(1, 5):
        c = sum_ws.cell(3, col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = THIN

    r = 4
    for region in REGION_ORDER:
        n = counts.get(region, 0)
        if n == 0:
            continue
        sum_ws.cell(r, 1, region).border = THIN
        sum_ws.cell(r, 2, n).border = THIN
        sum_ws.cell(r, 3, f"{100.0 * n / total:.1f}%").border = THIN
        sum_ws.cell(r, 4, NOTES.get(region, "")).border = THIN
        r += 1
    sum_ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    sum_ws.cell(r, 2, total).font = Font(bold=True)
    sum_ws.column_dimensions["A"].width = 22
    sum_ws.column_dimensions["B"].width = 16
    sum_ws.column_dimensions["C"].width = 12
    sum_ws.column_dimensions["D"].width = 55

    by_region: dict[str, list[dict[str, str]]] = defaultdict(list)
    for d in companies:
        by_region[d.get("Region Segment", "Other - Unclear")].append(d)

    # Create region sheets in Asia-Pacific-first order (append after existing sheets)
    for region in REGION_ORDER:
        rows = by_region.get(region) or []
        if not rows:
            continue
        safe = region.replace("/", "-")
        title = f"Region - {safe}"[:31]
        rws = wb.create_sheet(title)
        _write_rows(rws, headers, rows)

    # Move Region sheets to follow Region Summary, APAC first already by creation order
    # Ensure Region Summary is index 1 and region sheets follow in order
    desired = ["Region Summary"] + [
        f"Region - {r.replace('/', '-')}"[:31]
        for r in REGION_ORDER
        if by_region.get(r)
    ]
    # Move sheets to positions after main company sheet (index 0)
    for offset, name in enumerate(desired, start=1):
        if name in wb.sheetnames:
            current = wb.sheetnames.index(name)
            if current != offset:
                wb.move_sheet(name, offset=offset - current)

    wb.save(path)

    # First few companies for report
    head = [(d.get("Region Segment"), d.get("Company")) for d in companies[:8]]
    return {
        "file": path.name,
        "backup": bak.name,
        "total": total,
        "counts": {r: counts[r] for r in REGION_ORDER if counts.get(r)},
        "head": head,
    }


def main() -> None:
    for path, sheet, title in WORKBOOKS:
        if not path.exists():
            print("missing", path)
            continue
        rep = sort_workbook(path, sheet, title)
        print("FILE", rep["file"])
        print("  backup", rep["backup"])
        print("  order:", " → ".join(f"{k}({v})" for k, v in rep["counts"].items()))
        print("  first rows:")
        for reg, co in rep["head"]:
            print(f"    {reg} | {co}")
        print()


if __name__ == "__main__":
    main()
