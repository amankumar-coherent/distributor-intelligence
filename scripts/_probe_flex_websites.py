"""Probe website reachability + spot fake domains for flex packaging sheet."""
from __future__ import annotations

import re
import socket
import ssl
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)

wb = load_workbook(FP, read_only=True, data_only=True)
ws = wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
rows = list(ws.iter_rows(min_row=2, values_only=True))
wb.close()

# fill rates
print("FILL_RATES")
for i, h in enumerate(headers):
    filled = sum(1 for r in rows if r[i] not in (None, ""))
    print(f"  {h}: {filled}/{len(rows)}")

LOC_WORDS = {
    "india", "usa", "us", "uk", "uae", "china", "japan", "germany", "france",
    "italy", "spain", "brazil", "mexico", "canada", "australia", "korea",
    "thailand", "vietnam", "indonesia", "malaysia", "singapore", "taiwan",
    "philippines", "turkey", "poland", "netherlands", "belgium", "switzerland",
    "sweden", "norway", "denmark", "finland", "austria", "ireland", "portugal",
    "russia", "egypt", "argentina", "chile", "colombia", "peru", "pakistan",
    "bangladesh", "new zealand", "hong kong", "europe", "asia", "emea",
    "latam", "middle east", "global", "worldwide", "international",
}


def looks_location_variant(name: str) -> bool:
    n = name.strip().lower()
    for w in LOC_WORDS:
        if n.endswith(" " + w) or n.endswith(" " + w + ".") or f" {w} " in n:
            return True
        if n.endswith(" of " + w) or n.endswith(" de " + w):
            return True
    if re.search(r"\b(packaging|flexibles?|empaques|embalagens)\b.+\b(" + "|".join(map(re.escape, LOC_WORDS)) + r")\b", n):
        return True
    return False


def domain(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    host = urlparse(u).netloc.lower().split("@")[-1]
    return host[4:] if host.startswith("www.") else host


def probe(url: str) -> tuple[str, str]:
    """Return (status, note)."""
    u = (url or "").strip()
    if not u:
        return "empty", ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    host = domain(url)
    if not host:
        return "bad_url", ""
    try:
        socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
    except Exception as e:
        return "dns_fail", str(e)[:80]
    try:
        req = Request(u, headers={"User-Agent": "Mozilla/5.0"}, method="GET")
        with urlopen(req, timeout=8, context=ssl.create_default_context()) as resp:
            code = getattr(resp, "status", 200)
            return f"http_{code}", host
    except Exception as e:
        # try http
        try:
            u2 = "http://" + host
            req = Request(u2, headers={"User-Agent": "Mozilla/5.0"}, method="GET")
            with urlopen(req, timeout=8) as resp:
                code = getattr(resp, "status", 200)
                return f"http_{code}_via_http", host
        except Exception as e2:
            return "http_fail", str(e2)[:100]


suspects = [(i + 2, r[0], r[1]) for i, r in enumerate(rows) if looks_location_variant(str(r[0] or ""))]
print("LOCATION_VARIANT_ROWS", len(suspects))

# probe all websites (bounded)
items = [(i + 2, str(r[0] or ""), str(r[1] or "")) for i, r in enumerate(rows)]
results = {}
with ThreadPoolExecutor(max_workers=24) as ex:
    futs = {ex.submit(probe, w): (row, n, w) for row, n, w in items}
    for fut in as_completed(futs):
        row, n, w = futs[fut]
        results[row] = (n, w, fut.result())

dns_fail = [(row, *results[row]) for row in sorted(results) if results[row][2][0] == "dns_fail"]
http_fail = [(row, *results[row]) for row in sorted(results) if results[row][2][0] == "http_fail"]
ok = sum(1 for row in results if results[row][2][0].startswith("http_2") or results[row][2][0].startswith("http_3"))
print("OK_2xx_3xx", ok)
print("DNS_FAIL", len(dns_fail))
print("HTTP_FAIL", len(http_fail))
print("\nDNS_FAIL sample:")
for row, n, w, st in dns_fail[:80]:
    print(f"  r{row}: {n} | {w} | {st}")

print("\nLOCATION variants with dns/http fail:")
for row, n, w in suspects:
    st = results[row][2]
    if st[0] in ("dns_fail", "http_fail", "empty"):
        print(f"  r{row}: {n} | {w} | {st}")

# dump probe tsv
out = FP.parent / "_website_probe.tsv"
with out.open("w", encoding="utf-8") as f:
    f.write("row\tcompany\twebsite\tstatus\tnote\n")
    for row in sorted(results):
        n, w, (st, note) = results[row]
        f.write(f"{row}\t{n}\t{w}\t{st}\t{note}\n")
print("WROTE", out)
