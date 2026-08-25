# -*- coding: utf-8 -*-
from pathlib import Path
import re
from openpyxl import load_workbook

PATHS = [
    Path(r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global\global_electrical_distribution_market_global_FINAL.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market.xlsx"),
]

CEO_RE = re.compile(r"\bceo\b|chief executive", re.I)
KEEP_RE = re.compile(
    r"sales|commercial|procurement|purchas|revenue|\bcco\b|\bcro\b|\bcso\b|business development|\bbd\b",
    re.I,
)


def main() -> None:
    for p in PATHS:
        wb = load_workbook(p, data_only=True)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        headers = None
        hr = 1
        for r in range(1, 6):
            vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
            if "Company" in vals:
                headers = vals
                hr = r
                break
        print("===", p.name)
        n = 0
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            d = {
                headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
                for i in range(len(headers))
            }
            if not d.get("Company"):
                continue
            role = d.get("Role") or ""
            if not CEO_RE.search(role):
                continue
            keep = bool(KEEP_RE.search(role))
            n += 1
            print(f"  {'KEEP' if keep else 'DROP'}\t{d.get('Company')}\t{d.get('Contact Person')}\t{role}")
        print("count", n)
        wb.close()


if __name__ == "__main__":
    main()
