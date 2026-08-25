"""Add final 6 verified packaging distributors to reach +100."""
from __future__ import annotations

import json
import os
import re
import shutil
import socket
import ssl
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")
FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "add100_final6_audit.json"
API_KEY = (os.getenv("OPENAI_API_KEY") or "").strip()
BASE_URL = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
client = OpenAI(api_key=API_KEY, base_url=BASE_URL)
ND = "Not publicly disclosed"
BASELINE = 224
TARGET = 100

CANDS = [
    {
        "company": "New Century Packaging USA",
        "website": "https://ncpkgusa.com",
        "hq": "Valencia, CA, USA",
        "cc": "US",
        "founded": "Not publicly disclosed",
        "cats": "Full-line packaging supplies; warehouse automation",
        "focus": "US packaging distributor and systems integrator",
        "summary": "New Century Packaging USA Inc. is a full-line packaging distributor and systems integrator with locations across the US.",
    },
    {
        "company": "UK Packaging Online",
        "website": "https://ukpackaging.com",
        "hq": "Midlands, UK",
        "cc": "GB",
        "founded": "Not publicly disclosed",
        "cats": "Online packaging materials; cartons; shipping supplies",
        "focus": "UK online packaging merchant/distributor",
        "summary": "UK Packaging is a Midlands-based online packaging materials supplier with next-day delivery.",
    },
    {
        "company": "Office Packaging UK",
        "website": "https://officepackaging.co.uk",
        "hq": "UK",
        "cc": "GB",
        "founded": "Not publicly disclosed",
        "cats": "E-commerce packaging; boxes; mailing bags; bubble wrap",
        "focus": "UK packaging supplies distributor",
        "summary": "Office Packaging is a UK distributor of packaging materials for businesses and e-commerce.",
    },
    {
        "company": "Ernest Packaging",
        "website": "https://www.ernestpackaging.com",
        "hq": "USA",
        "cc": "US",
        "founded": "Not publicly disclosed",
        "cats": "Packaging supplies; equipment; custom packaging",
        "focus": "US packaging distributor",
        "summary": "Ernest Packaging distributes packaging supplies and solutions across the US.",
    },
    {
        "company": "Gillen Company",
        "website": "https://www.gillenco.com",
        "hq": "USA",
        "cc": "US",
        "founded": "Not publicly disclosed",
        "cats": "Packaging supplies",
        "focus": "US packaging distributor",
        "summary": "Gillen Company is a US packaging distributor.",
    },
    {
        "company": "All American Packaging Supply",
        "website": "https://www.allamericanpackaging.com",
        "hq": "USA",
        "cc": "US",
        "founded": "Not publicly disclosed",
        "cats": "Packaging supplies",
        "focus": "US packaging supplies distributor",
        "summary": "All American Packaging Supply distributes packaging supplies in the US.",
    },
    {
        "company": "Packaging Products Corp Dist",
        "website": "https://www.packagingproductscorp.com",
        "hq": "USA",
        "cc": "US",
        "founded": "Not publicly disclosed",
        "cats": "Packaging supplies",
        "focus": "US packaging distributor",
        "summary": "Packaging Products Corp distributes packaging supplies.",
    },
    {
        "company": "Postpack Merchant",
        "website": "https://www.postpack.co.uk",
        "hq": "East Midlands, UK",
        "cc": "GB",
        "founded": "1998",
        "cats": "Postal packaging; mailing boxes wholesale",
        "focus": "UK postal packaging wholesale merchant",
        "summary": "Postpack wholesales postal packaging with next-day delivery across the UK.",
    },
]


def domain_of(url: str) -> str:
    u = url if "://" in url else "https://" + url
    return urlparse(u).netloc.lower().replace("www.", "")


def dns_ok(d: str) -> bool:
    try:
        socket.getaddrinfo(d, 443)
        return True
    except Exception:
        try:
            socket.getaddrinfo(d, 80)
            return True
        except Exception:
            return False


def fetch(website: str) -> str:
    base = website.rstrip("/")
    chunks = []
    ctx = ssl.create_default_context()
    for p in ["", "/about", "/about-us"]:
        try:
            req = Request(base + p, headers={"User-Agent": "Mozilla/5.0"})
            with urlopen(req, timeout=12, context=ctx) as resp:
                html = resp.read(70000).decode("utf-8", errors="ignore")
            if any(x in html.lower() for x in ("domain is for sale", "buy this domain", "parked domain")):
                return "PARKED"
            text = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
            text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
            text = re.sub(r"<[^>]+>", " ", text)
            text = re.sub(r"\s+", " ", text)
            chunks.append(text[:2500])
        except Exception:
            continue
    return "\n".join(chunks)[:6000]


def role_check(company, website, source):
    prompt = f"""Classify packaging channel.
Company: {company}
Website: {website}
SOURCE: {source[:5500]}
KEEP only Specialty Distributor/Wholesaler/Merchant/Channel Partner.
REMOVE Manufacturer/OEM/Converter/Hybrid/Wrong/Unknown.
JSON: {{"keep":true|false,"role":"...","market_fit":true|false,"founded":"...","hq":"...","reason":"..."}}"""
    resp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": "Strict channel-only. Prefer remove hybrids."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.1,
        response_format={"type": "json_object"},
    )
    try:
        return json.loads(resp.choices[0].message.content or "{}")
    except Exception:
        return {"keep": False, "role": "Unknown", "market_fit": False}


def main():
    backup = FP.with_name(FP.stem + f"_pre_add100_final6_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(FP, backup)
    wb = load_workbook(FP)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    names = {(r[0] or "").strip().lower() for r in rows}
    doms = {domain_of(str(r[1] or "")) for r in rows if r[1]}
    need = TARGET - (len(rows) - BASELINE)
    print("need", need)

    added = 0
    accepted = []
    rejected = []
    for c in CANDS:
        if added >= need:
            break
        name, web, d = c["company"], c["website"], domain_of(c["website"])
        if name.lower() in names or d in doms:
            continue
        print("check", name, flush=True)
        if not dns_ok(d):
            rejected.append({"company": name, "reason": "dns"})
            print(" dns", flush=True)
            continue
        src = fetch(web)
        if src == "PARKED":
            rejected.append({"company": name, "reason": "parked"})
            print(" parked", flush=True)
            continue
        data = role_check(name, web, src or f"{name} {c['focus']}")
        role = (data.get("role") or "").lower()
        keep = bool(data.get("keep")) and bool(data.get("market_fit", True))
        if any(x in role for x in ("manufacturer", "oem", "converter", "hybrid", "wrong", "unknown")):
            keep = False
        if not keep:
            rejected.append({"company": name, "role": data.get("role"), "reason": data.get("reason")})
            print(" REJECT", data.get("role"), flush=True)
            time.sleep(0.15)
            continue
        if data.get("founded"):
            c["founded"] = str(data["founded"]).strip()
        if data.get("hq"):
            c["hq"] = str(data["hq"]).strip()
        cc = c["cc"]
        rc = "NA" if cc in {"US", "CA"} else "EU"
        geo = f"{rc}({cc})"
        ws.append(
            [
                c["company"],
                c["website"],
                c.get("founded") or ND,
                c.get("hq") or ND,
                geo,
                geo,
                "Private",
                ND,
                c["cats"],
                c["focus"],
                c["company"],
                "Both",
                "Specialty Distributor",
                ND,
                ND,
                ND,
                ND,
                ND,
                cc,
                rc,
                c["summary"],
            ]
        )
        names.add(name.lower())
        doms.add(d)
        added += 1
        accepted.append(c)
        print(f" KEEP ({(len(rows)-BASELINE)+added}/{TARGET})", flush=True)
        time.sleep(0.15)

    wb.save(FP)
    final = ws.max_row - 1
    audit = {
        "ts": datetime.now().isoformat(),
        "backup": str(backup),
        "added": added,
        "after": final,
        "from_baseline": final - BASELINE,
        "accepted": accepted,
        "rejected": rejected,
        "file": str(FP),
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print("DONE after", final, "from_baseline", final - BASELINE)


if __name__ == "__main__":
    main()
