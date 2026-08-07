#!/usr/bin/env python3
"""Strict second DeepSeek pass: ADC-channel distributors only.

Remove pure logistics/3PL, insurers, hospitals, manufacturers, and general
wholesalers with no oncology/ADC/specialty biologics channel signal.
Also ask DeepSeek to correct key filled columns.
"""
from __future__ import annotations

import json
import os
import re
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")

OUT_DIR = ROOT / "output/web_expand/global_antibody_drug_conjugates_market_global"
XLSX = OUT_DIR / "global_antibody_drug_conjugates_market_FINAL.xlsx"
SUMMARY = OUT_DIR / "adc_strict_deepseek_pass2_summary.json"
REMOVED = OUT_DIR / "adc_strict_deepseek_pass2_removed.json"
MODEL = os.getenv("OPENAI_MODEL", "deepseek-chat")
BATCH = 10

HEADERS = [
    "Company", "Founded", "Headquarters", "Cities / Regions", "Ownership",
    "Business Type", "Employees", "Turnover / Scale", "Contact Person", "Role",
    "Email", "Phone / WhatsApp", "LinkedIn", "Website", "Core Categories",
    "Specialty Focus", "Price Segment", "Key Brands Represented", "Exclusivity",
    "Partnership Duration", "Retail Chains", "Specialty Grocers", "Foodservice / HoReCa",
    "E-commerce", "Channel Strength", "Distribution Type", "Regions Served",
    "Regional Extensions", "Warehouse / Logistics", "Delivery / Storage",
    "Competitive Benchmark", "Data Confidence", "Quality Score", "Data Sources", "Summary",
]
WIDTHS = {
    "Company": 28, "Headquarters": 22, "Key Brands Represented": 32,
    "Specialty Focus": 28, "Core Categories": 28, "Summary": 36,
    "Website": 24, "Regions Served": 22,
}

NPD = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd", "null"}

# Hard remove: pure logistics / insurers / hospitals / OEMs that slipped through
HARD_REMOVE = {
    "cryoport systems",
    "world courier",
    "dhl supply chain life sciences",
    "blue shield california specialty",
    "khalifa medical / uae hospital supply partners",
    "biobridgex",
}


def empty(v: Any) -> bool:
    return str(v or "").strip().lower() in NPD


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def load_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    rows = []
    for r in range(3, ws.max_row + 1):
        d = {h: str(ws.cell(r, i).value or "").strip() for i, h in enumerate(hdr, 1) if h}
        if d.get("Company"):
            rows.append(d)
    return rows


def write_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ADC Distributors"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title = ws.cell(1, 1, "Global Antibody Drug Conjugates Market — Distributors / Channel Partners (strict audited)")
    title.font = Font(bold=True, size=13, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E79")
    header_fill = PatternFill("solid", fgColor="2E75B6")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="BDD7EE"),
        right=Side(style="thin", color="BDD7EE"),
        top=Side(style="thin", color="BDD7EE"),
        bottom=Side(style="thin", color="BDD7EE"),
    )
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(2, i, h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = thin
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 16)
    for ri, row in enumerate(rows, 3):
        for ci, h in enumerate(HEADERS, 1):
            cell = ws.cell(ri, ci, row.get(h, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
    cws = wb.create_sheet("Audit Checklist")
    cws["A1"] = "Strict ADC channel audit (DeepSeek deepseek-chat)"
    cws["A2"] = "Kept: distributors / specialty pharmacies / wholesalers in ADC/oncology/specialty biologics channel"
    cws["A3"] = "Removed: manufacturers, CDMOs, insurers, hospitals, pure 3PL logistics, unrelated"
    cws["A4"] = f"Rows: {len(rows)} | {datetime.now(timezone.utc).isoformat()}"
    wb.save(path)


SYSTEM = """You audit Global Antibody Drug Conjugates (ADC) market CHANNEL partners.

KEEP only if the company is primarily a DISTRIBUTOR / WHOLESALER / SPECIALTY PHARMACY / IMPORTER / AUTHORIZED CHANNEL PARTNER that participates in:
- ADC / oncology injectable / specialty biologics pharmaceutical distribution, OR
- life-science antibody / ADC-research reagent distribution into labs/hospitals.

REMOVE if ANY of these:
- manufacturer, CDMO, CMO, CRO, ADC drug developer, brand owner, reagent OEM that mainly manufactures
- health insurer / PBM only / hospital network (not a distributor)
- pure cold-chain logistics / 3PL / courier with no pharmaceutical wholesale or specialty pharmacy role
- unrelated to oncology/ADC/specialty biologics/antibody channel (e.g. only FMCG, only medical devices with no pharma/biologics)
- ambiguous company identity / wrong entity / not a real distributor

General national pharmacy wholesalers that routinely supply hospitals with oncology injectables MAY be kept as Distributor.

Correct fields only when you are reasonably confident. Use "Not publicly disclosed" if unknown.
Return STRICT JSON: {"results":[...]}"""


def call_batch(client: OpenAI, rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    payload = []
    for r in rows:
        payload.append({
            "company": r.get("Company"),
            "founded": r.get("Founded"),
            "hq": r.get("Headquarters"),
            "cities": r.get("Cities / Regions"),
            "ownership": r.get("Ownership"),
            "business_type": r.get("Business Type"),
            "employees": r.get("Employees"),
            "website": r.get("Website"),
            "core_categories": r.get("Core Categories"),
            "specialty_focus": r.get("Specialty Focus"),
            "brands": r.get("Key Brands Represented"),
            "regions": r.get("Regions Served"),
            "distribution_type": r.get("Distribution Type"),
            "summary": (r.get("Summary") or "")[:350],
        })
    user = {
        "instruction": "Decide keep/remove and optionally correct fields for each company.",
        "each_result_fields": {
            "company": "exact name",
            "keep": "boolean",
            "reason": "short",
            "role": "distributor|specialty_distributor|channel_partner|manufacturer|cdmo|logistics|insurer|hospital|unrelated|other",
            "corrections": {
                "Founded": "YYYY optional",
                "Headquarters": "city, country preferred optional",
                "Cities / Regions": "optional",
                "Ownership": "optional",
                "Business Type": "Distributor|Specialty Distributor optional",
                "Employees": "optional",
                "Core Categories": "optional",
                "Specialty Focus": "optional",
                "Key Brands Represented": "optional",
                "Regions Served": "optional",
                "Distribution Type": "optional",
                "Summary": "1-2 sentences optional",
            },
        },
        "companies": payload,
    }
    for attempt in range(3):
        try:
            resp = client.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": SYSTEM},
                    {"role": "user", "content": json.dumps(user, ensure_ascii=False)},
                ],
                temperature=0.05,
                max_tokens=5000,
                response_format={"type": "json_object"},
            )
            data = json.loads(resp.choices[0].message.content or "{}")
            if isinstance(data, dict):
                for k in ("results", "companies", "items", "data"):
                    if isinstance(data.get(k), list):
                        return data[k]
                if "company" in data:
                    return [data]
            if isinstance(data, list):
                return data
            return []
        except Exception as e:
            print(f"  batch error: {e}", flush=True)
            time.sleep(1.5 * (attempt + 1))
    return []


def apply(row: dict[str, str], d: dict[str, Any]) -> dict[str, str]:
    out = dict(row)
    corr = d.get("corrections") or {}
    if isinstance(corr, dict):
        for k, v in corr.items():
            if k in HEADERS and v is not None and str(v).strip() and not empty(v):
                out[k] = str(v).strip()
    role = str(d.get("role") or "").lower()
    if role in ("specialty_distributor", "specialty distributor"):
        out["Business Type"] = "Specialty Distributor"
    elif role in ("distributor", "channel_partner", "channel partner"):
        if "specialty" not in (out.get("Business Type") or "").lower():
            out["Business Type"] = "Distributor"
    # retail nonsense cleanup
    for col in ("Retail Chains", "Specialty Grocers", "Foodservice / HoReCa"):
        val = (out.get(col) or "").lower()
        if any(x in val for x in ("supermarket", "grocery", "restaurant", "horeca", "foodservice")):
            out[col] = "Not applicable"
    src = out.get("Data Sources") or ""
    if "adc_strict_deepseek_pass2" not in src:
        out["Data Sources"] = (src + "; adc_strict_deepseek_pass2").strip("; ")
    if empty(out.get("Summary")) and d.get("reason"):
        out["Summary"] = str(d.get("reason"))[:300]
    # confidence bump
    try:
        conf = int(re.sub(r"\D", "", out.get("Data Confidence") or "70") or "70")
    except Exception:
        conf = 70
    out["Data Confidence"] = str(min(95, conf + 3))
    out["Quality Score"] = out["Data Confidence"]
    return out


def main() -> int:
    rows = load_rows(XLSX)
    print(f"Loaded {len(rows)}; model={MODEL}")
    client = OpenAI(
        api_key=os.getenv("OPENAI_API_KEY"),
        base_url=os.getenv("OPENAI_BASE_URL", "https://api.deepseek.com/v1"),
    )
    smoke = client.chat.completions.create(
        model=MODEL, messages=[{"role": "user", "content": "OK"}], max_tokens=3
    )
    print(f"DeepSeek OK: {smoke.choices[0].message.content} ({smoke.model})")

    kept: list[dict[str, str]] = []
    removed: list[dict[str, Any]] = []

    # hard remove first
    work: list[dict[str, str]] = []
    for r in rows:
        if norm(r["Company"]) in HARD_REMOVE:
            removed.append({"company": r["Company"], "reason": "hard_remove_logistics_or_non_channel", "stage": "hard"})
            print(f"REMOVE hard: {r['Company']}")
        else:
            work.append(r)

    for i in range(0, len(work), BATCH):
        chunk = work[i : i + BATCH]
        print(f"\n[{i+1}-{i+len(chunk)}/{len(work)}] DeepSeek strict...", flush=True)
        decisions = call_batch(client, chunk)
        by = {norm(str(d.get("company") or "")): d for d in decisions if d.get("company")}
        for r in chunk:
            d = by.get(norm(r["Company"]))
            if not d:
                # keep with caution if DeepSeek miss
                d = {"company": r["Company"], "keep": True, "reason": "fallback_keep_missing_decision", "role": "distributor", "corrections": {}}
            keep = bool(d.get("keep"))
            role = str(d.get("role") or "").lower()
            if role in ("manufacturer", "cdmo", "logistics", "insurer", "hospital", "unrelated"):
                keep = False
            if keep:
                kept.append(apply(r, d))
                print(f"  KEEP  {r['Company']} — {d.get('reason')}", flush=True)
            else:
                removed.append({
                    "company": r["Company"],
                    "reason": d.get("reason"),
                    "role": d.get("role"),
                    "stage": "deepseek",
                })
                print(f"  REMOVE {r['Company']} — {d.get('reason')}", flush=True)
        time.sleep(0.2)

    seen = set()
    unique = []
    for r in kept:
        k = norm(r["Company"])
        if k in seen:
            continue
        seen.add(k)
        unique.append(r)
    unique.sort(key=lambda x: norm(x["Company"]))
    write_xlsx(XLSX, unique)
    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "model": MODEL,
        "input": len(rows),
        "kept": len(unique),
        "removed": len(removed),
        "business_types": dict(Counter(r.get("Business Type") for r in unique)),
        "removed_reasons": dict(Counter(str(x.get("reason", ""))[:120] for x in removed)),
        "removed_sample": removed[:60],
    }
    SUMMARY.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    REMOVED.write_text(json.dumps(removed, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nDONE kept={len(unique)} removed={len(removed)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
