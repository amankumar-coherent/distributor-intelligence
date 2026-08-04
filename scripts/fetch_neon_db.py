#!/usr/bin/env python3
"""Fetch data from a Neon (or any) PostgreSQL database.

Set DATABASE_URL in .env (never commit .env):

  DATABASE_URL=postgresql://user:pass@host/dbname?sslmode=require

Examples:
  python scripts/fetch_neon_db.py --list-tables
  python scripts/fetch_neon_db.py --table companies --limit 50
  python scripts/fetch_neon_db.py --table companies --export output/neon/companies.csv
  python scripts/fetch_neon_db.py --sql "SELECT id, name FROM companies LIMIT 10"
  python scripts/fetch_neon_db.py --export-all
  python scripts/fetch_neon_db.py --export-all --out-dir output/neon_export --formats csv,json,xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

_IDENT = re.compile(r"^[a-zA-Z_][a-zA-Z0-9_]*$")


def _database_url() -> str:
    url = (os.environ.get("DATABASE_URL") or "").strip()
    if not url:
        raise SystemExit(
            "DATABASE_URL is not set. Add it to .env (copy from .env.example).\n"
            "Example: DATABASE_URL=postgresql://user:pass@host/dbname?sslmode=require"
        )
    return url


def _connect():
    import psycopg

    return psycopg.connect(_database_url())


def list_tables(conn) -> list[tuple[str, str, int]]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT table_schema, table_name
            FROM information_schema.tables
            WHERE table_schema NOT IN ('pg_catalog', 'information_schema')
              AND table_type = 'BASE TABLE'
            ORDER BY table_schema, table_name
            """
        )
        tables = cur.fetchall()
        out: list[tuple[str, str, int]] = []
        for schema, name in tables:
            qualified = f'"{schema}"."{name}"' if schema != "public" else f'"{name}"'
            cur.execute(f"SELECT COUNT(*) FROM {qualified}")
            out.append((schema, name, int(cur.fetchone()[0])))
        return out


def _qualified_table(schema: str, name: str) -> str:
    if not _IDENT.match(schema) or not _IDENT.match(name):
        raise SystemExit(f"Invalid table identifier: {schema}.{name}")
    return f'"{schema}"."{name}"' if schema != "public" else f'"{name}"'


def fetch_table(conn, schema: str, table: str, *, limit: int | None) -> tuple[list[str], list[tuple]]:
    qualified = _qualified_table(schema, table)
    sql = f"SELECT * FROM {qualified}"
    if limit is not None:
        sql += f" LIMIT {int(limit)}"
    with conn.cursor() as cur:
        cur.execute(sql)
        columns = [d.name for d in cur.description] if cur.description else []
        rows = cur.fetchall()
    return columns, rows


def fetch_sql(conn, sql: str, *, limit: int | None) -> tuple[list[str], list[tuple]]:
    stripped = sql.strip().rstrip(";")
    upper = stripped.upper()
    if not upper.startswith("SELECT"):
        raise SystemExit("Only SELECT queries are allowed with --sql.")
    if limit is not None and " LIMIT " not in upper:
        stripped = f"{stripped} LIMIT {int(limit)}"
    with conn.cursor() as cur:
        cur.execute(stripped)
        columns = [d.name for d in cur.description] if cur.description else []
        rows = cur.fetchall()
    return columns, rows


def _serialize(value):
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def rows_to_dicts(columns: list[str], rows: list[tuple]) -> list[dict]:
    return [{col: _serialize(val) for col, val in zip(columns, row)} for row in rows]


def export_csv(path: Path, columns: list[str], rows: list[tuple]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([_serialize(v) for v in row])


def export_json(path: Path, columns: list[str], rows: list[tuple]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = rows_to_dicts(columns, rows)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def export_xlsx(path: Path, columns: list[str], rows: list[tuple], *, sheet_name: str) -> None:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(columns)
    for row in rows:
        ws.append([_serialize(v) for v in row])
    wb.save(path)


def export_all_tables(
    conn,
    out_dir: Path,
    *,
    formats: set[str],
) -> dict:
    tables = list_tables(conn)
    manifest: dict = {
        "exported_at": __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat(),
        "tables": [],
    }
    out_dir.mkdir(parents=True, exist_ok=True)

    for schema, name, _count in tables:
        columns, rows = fetch_table(conn, schema, name, limit=None)
        base = out_dir / name
        written: list[str] = []
        if "csv" in formats:
            export_csv(base.with_suffix(".csv"), columns, rows)
            written.append(f"{name}.csv")
        if "json" in formats:
            export_json(base.with_suffix(".json"), columns, rows)
            written.append(f"{name}.json")
        if "xlsx" in formats:
            export_xlsx(base.with_suffix(".xlsx"), columns, rows, sheet_name=name)
            written.append(f"{name}.xlsx")
        manifest["tables"].append(
            {
                "schema": schema,
                "table": name,
                "rows": len(rows),
                "columns": columns,
                "files": written,
            }
        )
        print(f"  {name}: {len(rows)} rows -> {', '.join(written)}")

    manifest_path = out_dir / "_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return manifest


def resolve_table(conn, table_arg: str) -> tuple[str, str]:
    if "." in table_arg:
        schema, name = table_arg.split(".", 1)
        return schema, name
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT table_schema
            FROM information_schema.tables
            WHERE table_schema NOT IN ('pg_catalog', 'information_schema')
              AND table_name = %s
            ORDER BY CASE WHEN table_schema = 'public' THEN 0 ELSE 1 END
            """,
            (table_arg,),
        )
        hits = [r[0] for r in cur.fetchall()]
    if not hits:
        raise SystemExit(f"Table not found: {table_arg}")
    if len(hits) > 1:
        opts = ", ".join(f"{s}.{table_arg}" for s in hits)
        raise SystemExit(f"Ambiguous table {table_arg!r}. Use schema prefix: {opts}")
    return hits[0], table_arg


def main() -> None:
    parser = argparse.ArgumentParser(description="Fetch rows from PostgreSQL (Neon).")
    parser.add_argument("--list-tables", action="store_true", help="List schemas, tables, row counts")
    parser.add_argument("--table", metavar="NAME", help="Table name (optionally schema.table)")
    parser.add_argument("--sql", metavar="SELECT", help="Run a read-only SELECT query")
    parser.add_argument("--limit", type=int, default=100, help="Max rows (default: 100; 0 = no limit)")
    parser.add_argument("--export", metavar="PATH", help="Write CSV or JSON (.csv / .json)")
    parser.add_argument("--preview", action="store_true", help="Print first rows to stdout (default if no --export)")
    parser.add_argument(
        "--export-all",
        action="store_true",
        help="Export every table to separate files (no merge with pipeline data)",
    )
    parser.add_argument(
        "--out-dir",
        default=str(ROOT / "output" / "neon_export"),
        help="Output directory for --export-all (default: output/neon_export)",
    )
    parser.add_argument(
        "--formats",
        default="csv,json,xlsx",
        help="Comma-separated: csv, json, xlsx (default: all three)",
    )
    args = parser.parse_args()

    if not args.list_tables and not args.table and not args.sql and not args.export_all:
        parser.print_help()
        raise SystemExit(1)

    limit = None if args.limit == 0 else args.limit
    formats = {f.strip().lower() for f in args.formats.split(",") if f.strip()}
    bad = formats - {"csv", "json", "xlsx"}
    if bad:
        raise SystemExit(f"Unknown format(s): {', '.join(sorted(bad))}")

    with _connect() as conn:
        if args.export_all:
            out_dir = Path(args.out_dir)
            print(f"Exporting all tables to {out_dir} ...")
            manifest = export_all_tables(conn, out_dir, formats=formats)
            print(f"\nDone. {len(manifest['tables'])} tables -> {out_dir}")
            print(f"Manifest: {out_dir / '_manifest.json'}")
            return

        if args.list_tables:
            tables = list_tables(conn)
            if not tables:
                print("No user tables found.")
                return
            print(f"{'Schema':<20} {'Table':<40} {'Rows':>8}")
            print("-" * 70)
            for schema, name, count in tables:
                print(f"{schema:<20} {name:<40} {count:>8}")
            return

        if args.sql:
            columns, rows = fetch_sql(conn, args.sql, limit=limit)
        else:
            schema, name = resolve_table(conn, args.table)
            columns, rows = fetch_table(conn, schema, name, limit=limit)
            print(f"Fetched {len(rows)} row(s) from {schema}.{name}")

        if args.export:
            path = Path(args.export)
            if path.suffix.lower() == ".json":
                export_json(path, columns, rows)
            else:
                export_csv(path, columns, rows)
            print(f"Wrote {path}")

        if args.preview or not args.export:
            records = rows_to_dicts(columns, rows)
            print(json.dumps(records[: min(10, len(records))], indent=2, ensure_ascii=False))
            if len(records) > 10:
                print(f"... ({len(records) - 10} more rows)")


if __name__ == "__main__":
    main()
