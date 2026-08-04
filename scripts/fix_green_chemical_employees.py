"""Fix green chemical FINAL Employees → verified LinkedIn-style ranges."""
from __future__ import annotations

import re
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

PATHS = [
    Path(
        r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL.xlsx"
    ),
    Path(
        r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL_320plus.xlsx"
    ),
]

# Exact-name overrides (agent web research / filings)
OVERRIDES: dict[str, str] = {
    # raw / garbage numbers
    "brenntag": "10001-50000",  # ~17,300 (FY2025 AR)
    "telko": "201-500",  # ~350 experts (company site); 19197 was bad scrape
    "airedale chemical": "1-10",  # sheet had "1"
    "k. k. chempro india": "51-200",  # "100 Employees"
    "green chemistry": "1-10",  # "0-1 employees"
    "ravago chemicals": "1001-5000",
    # NPD fills
    "ghcl trading division": "501-1000",  # GHCL Ltd ~518 LinkedIn
    "kolmar group": "51-200",  # ~82–130 LinkedIn/RocketReach
    "koda distribution group": "201-500",  # US specialty chem distribution scale
    # large / 10000+ / coarse bands
    "imcd group": "5001-10000",  # ~5,246 (AR 2025)
    "univar solutions": "10001-50000",  # ~11.4k
    "nagase & co": "5001-10000",
    "omya distribution": "5001-10000",
    "merck kgaa / milliporesigma distribution": "50001-100000",  # Merck KGaA ~62k
    "thermo fisher scientific chemicals channel": "100001-500000",  # ~125k
    "vwr / avantor distribution": "10001-50000",  # Avantor ~14.6k
    "sinochem plastics / sinochem international distribution": "100001-500000",
    "sojitz corporation (chemicals)": "10001-50000",
    "sumitomo corporation chemicals": "50001-100000",
    "toyota tsusho chemicals": "50001-100000",
    "wilmar international (oleochemicals)": "100001-500000",
    "aeci specialty chemicals / much asphalt chem channel": "5001-10000",
    "kowa company chemicals": "5001-10000",
    "braskem green pe channel partners": "5001-10000",
    "croda distribution partners / croda industrial specialties channel": "5001-10000",
}

NPD = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd"}

BUCKETS = [
    (10, "1-10"),
    (50, "11-50"),
    (200, "51-200"),
    (500, "201-500"),
    (1000, "501-1000"),
    (5000, "1001-5000"),
    (10000, "5001-10000"),
    (50000, "10001-50000"),
    (100000, "50001-100000"),
    (500000, "100001-500000"),
    (1000000, "500001-1000000"),
]


def norm_name(s: str) -> str:
    return re.sub(r"[^a-z0-9&]+", " ", (s or "").lower()).strip()


def parse_int(tok: str) -> int | None:
    tok = tok.replace(",", "").replace(" ", "")
    if not tok.isdigit():
        return None
    return int(tok)


def number_to_range(n: int) -> str:
    for hi, label in BUCKETS:
        if n <= hi:
            return label
    return "1000001+"


def looks_like_garbage(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if re.search(
        r"(?i)\b(trust|foundation|bank|llc|inc\.?|ltd|gmbh|corp|partners|holdings)\b",
        t,
    ):
        if not re.search(r"\d", t):
            return True
    if re.fullmatch(r"[A-Za-z][A-Za-z &.'-]{2,}", t) and t.lower() not in NPD:
        return True
    return False


def normalize_employees(raw: str, company: str) -> tuple[str, str]:
    key = norm_name(company)
    ov_map = {norm_name(k): (k, v) for k, v in OVERRIDES.items()}
    if key in ov_map:
        needle, rng = ov_map[key]
        return rng, f"verified_override:{needle}"

    v = (raw or "").strip()
    if v.lower() in NPD:
        return "Not publicly disclosed", "npd"

    if looks_like_garbage(v):
        return "Not publicly disclosed", f"cleared_garbage:{v}"

    v2 = re.sub(r"(?i)\s*(employees?|staff|team members?)\s*$", "", v).strip()
    v2 = v2.replace("–", "-").replace("—", "-").replace("−", "-")
    v2 = re.sub(r"\s+", "", v2)
    v2 = re.sub(r",(?=\d{3}(\D|$))", "", v2)

    # 0-1 → 1-10
    m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", v2)
    if m:
        a, b = parse_int(m.group(1)), parse_int(m.group(2))
        if a is not None and b is not None and a <= b:
            mid = max(1, (a + b) // 2)
            return number_to_range(mid), f"normalized_range:{a}-{b}"

    m = re.fullmatch(r"(\d+)\+", v2)
    if m:
        n = parse_int(m.group(1))
        if n is not None:
            return number_to_range(n + 1), f"open_ended:{n}+"

    m = re.fullmatch(r"\d+", v2)
    if m:
        n = parse_int(v2)
        if n is not None:
            return number_to_range(n), f"point_to_range:{n}"

    nums = re.findall(r"\d+", v2)
    if nums:
        n = parse_int(nums[0])
        if n is not None:
            return number_to_range(n), f"extracted:{n}"

    return "Not publicly disclosed", f"unparsed:{v}"


def fix_file(path: Path) -> dict:
    bak = path.with_name(
        path.stem + f"_pre_emp_fix_{datetime.now().strftime('%Y%m%d_%H%M%S')}" + path.suffix
    )
    shutil.copy2(path, bak)

    wb = openpyxl.load_workbook(path)
    ws = (
        wb["Green Chemical Companies"]
        if "Green Chemical Companies" in wb.sheetnames
        else wb.active
    )

    hdr_row = None
    hdr = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), 1):
        vals = [str(c or "") for c in row]
        if "Company" in vals and "Employees" in vals:
            hdr_row = i
            hdr = vals
            break
    if not hdr_row:
        raise RuntimeError(f"No header in {path}")

    emp_col = hdr.index("Employees") + 1
    co_col = hdr.index("Company") + 1

    changes = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        company = str(ws.cell(r, co_col).value or "").strip()
        if not company:
            continue
        old = str(ws.cell(r, emp_col).value or "").strip()
        new, reason = normalize_employees(old, company)
        if new != old:
            ws.cell(r, emp_col).value = new
            changes.append((company, old, new, reason))

    wb.save(path)

    dist = Counter()
    garbage = []
    raw = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        company = str(ws.cell(r, co_col).value or "").strip()
        if not company:
            continue
        e = str(ws.cell(r, emp_col).value or "").strip()
        dist[e] += 1
        if re.search(r"(?i)trust", e):
            garbage.append((company, e))
        if re.fullmatch(r"[\d,]+", e) and parse_int(e.replace(",", "") or "0") and parse_int(e.replace(",", "")) >= 1000:
            raw.append((company, e))

    return {
        "path": path.name,
        "backup": bak.name,
        "changed": len(changes),
        "samples": changes[:30],
        "dist": dist.most_common(),
        "garbage_left": garbage,
        "raw_left": raw,
    }


def main() -> None:
    for p in PATHS:
        if not p.exists():
            print("missing", p)
            continue
        rep = fix_file(p)
        print("FILE", rep["path"])
        print("  backup", rep["backup"])
        print("  changed", rep["changed"])
        print("  samples:")
        for co, old, new, reason in rep["samples"]:
            print(f"    {co}: {old!r} -> {new!r} ({reason})")
        print("  distribution:")
        for v, c in rep["dist"]:
            print(f"    {c:3d}  {v}")
        if rep["garbage_left"]:
            print("  WARNING garbage", rep["garbage_left"])
        if rep["raw_left"]:
            print("  WARNING raw", rep["raw_left"])
        print()


if __name__ == "__main__":
    main()
