# -*- coding: utf-8 -*-
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

FILES = {
    "electrical": Path(r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global\global_electrical_distribution_market_global_FINAL.xlsx"),
    "avocado": Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market.xlsx"),
    "smartphone": Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    "firewall": Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market.xlsx"),
}
NPD = {"", "nan", "none", "n/a", "na", "not publicly disclosed", "-", "—", "unknown", "tbd", "null"}

for mkt, path in FILES.items():
    wb = load_workbook(path, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            headers = vals
            hr = r
            break
    samples = []
    prefixes = Counter()
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip()) for i in range(len(headers))}
        if not d.get("Company"):
            continue
        person = (d.get("Contact Person") or "").strip()
        li = (d.get("LinkedIn") or "").strip()
        if not person or person.lower() in NPD:
            continue
        if not li or li.lower() in NPD:
            continue
        low = li.lower()
        if low.startswith("https://www.linkedin.com/in/"):
            prefixes["https://www.linkedin.com/in/"] += 1
        elif low.startswith("http://www.linkedin.com/in/"):
            prefixes["http://www.linkedin.com/in/"] += 1
        elif low.startswith("https://linkedin.com/in/"):
            prefixes["https://linkedin.com/in/"] += 1
        elif "linkedin.com/in/" in low:
            prefixes["other /in/"] += 1
            if len(samples) < 8:
                samples.append(li)
        elif "linkedin.com/company/" in low:
            prefixes["company"] += 1
            samples.append(li)
        else:
            prefixes["other"] += 1
            samples.append(li)
    print("====", mkt, "====")
    print(dict(prefixes))
    for s in samples[:8]:
        print(" SAMPLE:", s[:140])
    wb.close()
