#!/usr/bin/env python3
"""Fill Employees + Turnover/Scale for ADC FINAL via Growjo/Wikipedia/DeepSeek.

Owler blocks bots (403); use Growjo + Wikipedia + site + DeepSeek verify.
Label estimates clearly when from Growjo/estimate sources.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

import httpx
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")

OUT = ROOT / "output/web_expand/global_antibody_drug_conjugates_market_global"
XLSX = OUT / "global_antibody_drug_conjugates_market_FINAL.xlsx"
SUMMARY = OUT / "adc_fill_employees_turnover_summary.json"
PROGRESS = OUT / "adc_fill_emp_turnover_progress.jsonl"
MODEL = os.getenv("OPENAI_MODEL", "deepseek-chat")
BATCH = 8
CONCURRENCY = 6

HEADERS = [
    "Company", "Founded", "Headquarters", "Cities / Regions", "Ownership",
    "Business Type", "Employees", "Turnover / Scale", "Contact Person", "Role",
    "Email", "Phone / WhatsApp", "LinkedIn", "Website", "Core Categories",
    "Specialty Focus", "Price Segment", "Key Brands Represented", "Exclusivity",
    "Partnership Duration", "Retail Chains", "Specialty Grocers", "Foodservice / HoReCa",
    "E-commerce", "Channel Strength", "Distribution Type", "Regions Served",
    "Regional Extensions", "Warehouse / Logistics", "Delivery / Storage",
    "Competitive Benchmark", "Data Confidence", "Quality Score", "Data Sources", "Summary",
]
WIDTHS = {
    "Company": 28, "Headquarters": 28, "Cities / Regions": 36,
    "Employees": 18, "Turnover / Scale": 22, "Summary": 36, "Website": 24,
}

NPD = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd", "null"}

# Public / well-sourced figures (employees, turnover). Prefer entity-specific where known.
KNOWN: dict[str, tuple[str, str]] = {
    "McKesson Specialty Health": ("~150–650 (division est.)", "Parent McKesson ~$359B revenue (FY2025); division n/d"),
    "McKesson Specialty Care Distribution": ("Part of McKesson (~45,000+ company)", "Parent McKesson ~$359B (FY2025)"),
    "McKesson Plasma and Biologics": ("Part of McKesson (~45,000+ company)", "Parent McKesson ~$359B (FY2025)"),
    "McKesson Medical-Surgical": ("Part of McKesson (~45,000+ company)", "Parent McKesson ~$359B (FY2025)"),
    "McKesson Canada": ("Part of McKesson Canada ops", "Part of McKesson group revenue"),
    "Biologics by McKesson": ("Part of McKesson specialty", "Part of McKesson specialty channel"),
    "Cardinal Health Specialty Solutions": ("Part of Cardinal Health (~48,000 company)", "Parent Cardinal Health ~$227B revenue"),
    "Cardinal Health Specialty Pharmaceutical Distribution": ("Part of Cardinal Health (~48,000)", "Parent Cardinal Health ~$227B"),
    "Cencora (AmerisourceBergen) Specialty": ("Part of Cencora (~46,000 company)", "Parent Cencora ~$262B revenue"),
    "ASD Healthcare": ("Part of Cencora specialty", "Part of Cencora group"),
    "Oncology Supply": ("Part of Cencora specialty", "Part of Cencora group"),
    "CuraScript SD": ("Part of Cencora specialty", "Part of Cencora group"),
    "AAH Pharmaceuticals": ("~2,800", "~£2.9B revenue (FY2025 Companies House)"),
    "Suzuken Co., Ltd.": ("~13,000+ (group)", "~¥2.3T+ group revenue (public)"),
    "Alfresa Holdings": ("~12,000+ (group)", "Public TSE; multi-trillion JPY group scale"),
    "Medipal Holdings": ("~10,000+ (group)", "Public TSE; large Japan wholesale scale"),
    "Toho Holdings / TOHO Pharmaceutical": ("~8,000+ (group)", "Public TSE; large Japan wholesale scale"),
    "Phoenix Pharma SE": ("~40,000+ (PHOENIX group)", "€30B+ PHOENIX group revenue scale"),
    "PHOENIX Pharmahandel": ("Part of PHOENIX group", "Part of PHOENIX ~€30B+ group"),
    "Phoenix Healthcare Distribution": ("Part of PHOENIX UK", "Part of PHOENIX group"),
    "Alliance Healthcare UK": ("Part of Cencora / Alliance Healthcare", "Part of Cencora European wholesale"),
    "CVS Specialty": ("Part of CVS Health (~300,000 company)", "Part of CVS Health ~$360B+ revenue"),
    "Walgreens Specialty Pharmacy": ("Part of Walgreens Boots Alliance", "Part of WBA group revenue"),
    "AllianceRx Walgreens Pharmacy": ("Part of Walgreens Boots Alliance", "Part of WBA group revenue"),
    "Optum Specialty Pharmacy": ("Part of UnitedHealth / Optum", "Part of UnitedHealth Group revenue"),
    "Accredo Health Group": ("Part of Cigna / Evernorth", "Part of Cigna group"),
    "Cigna Specialty Pharmacy / Accredo": ("Part of Cigna / Evernorth", "Part of Cigna group"),
    "Express Scripts Accredo network": ("Part of Cigna / Evernorth", "Part of Cigna group"),
    "Aetna Specialty Pharmacy": ("Part of CVS Health / Aetna", "Part of CVS Health group"),
    "Henry Schein Medical": ("Part of Henry Schein (~25,000)", "Parent Henry Schein ~$12B+ revenue"),
    "Medline Industries": ("~40,000+", "~$20B+ estimated revenue (private est.)"),
    "Owens & Minor": ("~20,000+", "Public; multi-billion USD revenue"),
    "VWR International": ("Part of Avantor", "Part of Avantor public revenue"),
    "Avantor Inc.": ("~14,000+", "Public; ~$7B revenue scale"),
    "Fisher Scientific": ("Part of Thermo Fisher", "Part of Thermo Fisher ~$40B+ revenue"),
    "Zuellig Pharma": ("~10,000+", "~$15B+ estimated revenue (private est.)"),
    "DKSH Health": ("Part of DKSH (~33,000 company)", "Parent DKSH ~CHF 11B+ revenue"),
    "DiethelmKellerSiberHegner Healthcare": ("Part of DKSH (~33,000)", "Parent DKSH ~CHF 11B+"),
    "Cofares": ("~5,000+ cooperative network scale", "Multi-billion EUR Spanish pharmacy cooperative"),
    "NOWEDA Apothekergenossenschaft": ("Cooperative wholesale scale (thousands)", "Multi-billion EUR German pharmacy cooperative"),
    "Sanacorp eG": ("Cooperative wholesale scale", "Multi-billion EUR German pharmacy cooperative"),
    "Neuca": ("~4,000+", "Public WSE; multi-billion PLN revenue"),
    "Sigma Healthcare": ("~5,000+ (group est.)", "Public ASX; multi-billion AUD wholesale"),
    "Symbion": ("Part of EBOS / Sigma network AU", "Part of Australian wholesale group scale"),
    "Uniphar": ("~3,000+", "Public; multi-hundred million EUR+"),
    "Galenica / Alloga": ("Part of Galenica (~7,000+)", "Parent Galenica public CHF revenue"),
    "Tamro": ("Part of Phoenix / Tamro Nordics", "Part of PHOENIX group"),
    "Oriola": ("~800–1,000+", "Public; Nordic wholesale scale"),
    "Hungaropharma": ("~1,000+", "Major Hungary pharma wholesale scale"),
    "Nadro": ("~5,000+ (est.)", "Major Mexico wholesale scale"),
    "Marzam": ("Large Mexico wholesale (thousands est.)", "Major Mexico wholesale scale"),
    "Casa Saba": ("Large Mexico wholesale (thousands est.)", "Major Mexico wholesale scale"),
    "Sinopharm": ("100,000+ (group)", "Public; one of world's largest pharma distributors"),
    "Jointown Pharmaceutical Group": ("30,000+ (group est.)", "Public China; large wholesale scale"),
    "Shanghai Pharmaceuticals distribution": ("Part of Shanghai Pharma group", "Public China pharma group scale"),
    "China Resources Pharmaceutical": ("Part of CR Pharma group", "Public; large China pharma group"),
    "PT Enseval Putera Megatrading": ("~3,000+", "Public IDX; Indonesia wholesale scale"),
    "Profarma": ("~2,000+", "Public Brazil; multi-billion BRL wholesale"),
    "Viveo": ("Large Brazil specialty (thousands est.)", "Brazil specialty/hospital scale"),
    "Oncoprod": ("Specialty Brazil (hundreds–thousands est.)", "Brazil oncology specialty scale"),
    "BioCare Specialty Distribution (BioCareSD)": ("Hundreds (specialty distributor est.)", "US specialty distributor scale (private)"),
    "Anda Inc.": ("Part of Teva / Anda wholesale", "Large US wholesale scale"),
    "Anda Pharmaceuticals": ("Part of Teva / Anda wholesale", "Large US wholesale scale"),
    "Morris & Dickson Co., LLC": ("~1,000+", "Independent US wholesaler scale"),
    "Walmart Specialty Pharmacy": ("Part of Walmart (~2M+ company)", "Part of Walmart group"),
    "Kroger Specialty Pharmacy": ("Part of Kroger (~400,000 company)", "Part of Kroger group"),
    "CenterWell Pharmacy": ("Part of Humana", "Part of Humana group"),
    "Option Care Health specialty infusion": ("~6,000+", "Public; multi-billion USD infusion"),
    "PANTHERx Rare": ("~500–1,000 (est.)", "US rare specialty pharmacy scale"),
    "Orsini Specialty Pharmacy": ("~500+ (est.)", "US specialty pharmacy scale"),
    "Onco360": ("~500+ (est.)", "US oncology specialty pharmacy scale"),
    "US Bioservices": ("Part of Cencora / specialty", "Part of Cencora specialty"),
    "TheraCom Specialty Pharmacy": ("Part of Cencora", "Part of Cencora specialty services"),
    "Sciensus": ("~1,000+ (est.)", "UK specialty pharmacy / homecare scale"),
    "Mawdsleys": ("~500+ (est.)", "UK pharmaceutical wholesale scale"),
    "Mawdsley-Brooks": ("~500+ (est.)", "UK pharmaceutical wholesale scale"),
    "VECTOR Pharma FZCO": ("~50–200 (est.)", "MENA specialty commercialization scale"),
    "NewBridge Pharmaceuticals": ("~100–300 (est.)", "MENA specialty pharma scale"),
    "2BScientific Ltd": ("~11–50 (est.)", "SME life-science distributor scale"),
    "2BScientific GmbH": ("~11–50 (est.)", "SME life-science distributor scale"),
    "Thomas Scientific": ("~200–500 (est.)", "US lab distributor scale"),
    "Cedarlane": ("~50–200 (est.)", "Canada life-science distributor scale"),
    "Cedarlane Laboratories": ("~50–200 (est.)", "Canada life-science distributor scale"),
    "Cosmo Bio Co., Ltd.": ("~100–300 (est.)", "Japan life-science distributor scale"),
    "Funakoshi": ("~100–300 (est.)", "Japan life-science distributor scale"),
    "Abacus dx": ("~50–200 (est.)", "ANZ diagnostics distributor scale"),
}


def empty(v: Any) -> bool:
    return str(v or "").strip().lower() in NPD


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def growjo_slug(company: str) -> str:
    s = re.sub(r"[^\w\s\-]", " ", company)
    s = re.sub(r"\s+", "_", s.strip())
    s = re.sub(r"_+", "_", s)
    return s.strip("_")


def parse_employees(text: str) -> str:
    patterns = [
        r"has\s+(?:<!-- -->)?([\d,]+)(?:<!-- -->)?\s+Employees",
        r"([\d,]+)\s+Employees",
        r'"employees"\s*:\s*(\d+)',
        r"(?:employees?|headcount|staff)[^\d]{0,24}(\d[\d,]*(?:\s*[-–]\s*\d[\d,]*)?)",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if m:
            val = m.group(1).replace(",", "").strip()
            if val.isdigit() and 1 <= int(val) <= 5_000_000:
                return f"{int(val):,}"
            if re.search(r"\d", val):
                return val
    return ""


def parse_revenue(text: str) -> str:
    patterns = [
        r"estimated annual revenue is currently\s*(?:<!-- -->)?(\$?[\d\.,]+[MBKmbk]?)(?:<!-- -->)?",
        r"estimated annual revenue[^\d$]{0,40}(\$[\d\.,]+\s*[MBKmbk]?)",
        r"(?:revenue|turnover)[^\d$€£¥]{0,30}([\$€£¥]\s?[\d\.,]+\s*(?:million|billion|trillion|m|bn|b)?)",
        r"Revenue\s*(?:\(est\))?\s*[:\-]?\s*(\$[\d\.,]+[MBKmbk]?)",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if m:
            return normalize_money(m.group(1))
    return ""


def normalize_money(raw: str) -> str:
    s = re.sub(r"\s+", "", (raw or "").strip())
    if not s:
        return ""
    m = re.match(r"^([\$€£¥]?)([\d\.,]+)([MBKmbk]?)$", s)
    if not m:
        return raw.strip()
    cur, num, suf = m.groups()
    cur = cur or "$"
    suf = suf.upper()
    try:
        n = float(num.replace(",", ""))
    except ValueError:
        return raw.strip()
    if suf == "K":
        return f"{cur}{n:.0f}K (est.)"
    if suf == "M":
        return f"{cur}{n:.1f}M (est.)" if n < 100 else f"{cur}{n:.0f}M (est.)"
    if suf == "B":
        return f"{cur}{n:.1f}B (est.)" if n < 100 else f"{cur}{n:.0f}B (est.)"
    if n >= 1_000_000_000:
        return f"{cur}{n/1_000_000_000:.1f}B (est.)"
    if n >= 1_000_000:
        return f"{cur}{n/1_000_000:.1f}M (est.)"
    return f"{cur}{n:,.0f} (est.)"


def load_rows(path: Path) -> list[dict[str, str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    rows = []
    for r in range(3, ws.max_row + 1):
        d = {h: str(ws.cell(r, i).value or "").strip() for i, h in enumerate(hdr, 1) if h}
        if d.get("Company"):
            rows.append(d)
    return rows


def write_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ADC Distributors"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title = ws.cell(
        1, 1,
        "Global Antibody Drug Conjugates Market — Distributors (Employees / Turnover filled)",
    )
    title.font = Font(bold=True, size=13, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E79")
    hf = PatternFill("solid", fgColor="2E75B6")
    hfont = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="BDD7EE"),
        right=Side(style="thin", color="BDD7EE"),
        top=Side(style="thin", color="BDD7EE"),
        bottom=Side(style="thin", color="BDD7EE"),
    )
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(2, i, h)
        c.fill = hf
        c.font = hfont
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = thin
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 16)
    for ri, row in enumerate(rows, 3):
        for ci, h in enumerate(HEADERS, 1):
            cell = ws.cell(ri, ci, row.get(h, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
    cws = wb.create_sheet("Column Rules")
    cws["A1"] = "Employees / Turnover sourced from Growjo estimates, Wikipedia/public filings, DeepSeek verify"
    cws["A2"] = "Owler blocks automated access (403); Growjo used as Owler-class estimate source"
    cws["A3"] = "Values marked (est.) are third-party estimates — verify for financial decisions"
    cws["A4"] = f"Updated: {datetime.now(timezone.utc).isoformat()} | rows={len(rows)}"
    wb.save(path)


async def fetch(client: httpx.AsyncClient, url: str) -> str:
    try:
        r = await client.get(url, timeout=12.0, follow_redirects=True)
        if r.status_code >= 400:
            return ""
        return r.text or ""
    except Exception:
        return ""


async def growjo_lookup(client: httpx.AsyncClient, company: str) -> dict[str, str]:
    out: dict[str, str] = {}
    slug = growjo_slug(company)
    # try a few slug variants
    variants = [slug]
    base = re.sub(r"_(Inc|LLC|Ltd|GmbH|Co|SA|Sdn_Bhd|Pte_Ltd|Corporation|Corp|Group|Holdings).*$", "", slug, flags=re.I)
    if base and base != slug:
        variants.append(base)
    short = "_".join(slug.split("_")[:4])
    if short not in variants:
        variants.append(short)

    for v in variants[:3]:
        html = await fetch(client, f"https://growjo.com/company/{v}")
        if not html or len(html) < 500:
            continue
        # must mention company-ish token
        token = company.split()[0]
        if token.lower() not in html.lower() and company[:12].lower() not in html.lower():
            continue
        emp = parse_employees(html)
        rev = parse_revenue(html)
        if emp:
            out["employees"] = emp
            out["source"] = "growjo"
        if rev:
            out["turnover"] = rev if "(est.)" in rev else f"{rev} (est.)"
            out["source"] = "growjo"
        if out:
            out["url"] = f"https://growjo.com/company/{v}"
            break
    return out


async def wikipedia_lookup(client: httpx.AsyncClient, company: str) -> dict[str, str]:
    out: dict[str, str] = {}
    try:
        api = (
            "https://en.wikipedia.org/w/api.php?action=query&list=search"
            f"&srsearch={quote(company)}&utf8=1&format=json&srlimit=1"
        )
        r = await client.get(api, timeout=8.0)
        hits = ((r.json().get("query") or {}).get("search") or [])
        if not hits:
            return out
        title = hits[0].get("title") or ""
        r2 = await client.get(
            f"https://en.wikipedia.org/api/rest_v1/page/summary/{quote(title)}",
            timeout=8.0,
        )
        extract = r2.json().get("extract") or ""
        emp = parse_employees(extract)
        rev = parse_revenue(extract)
        # also try infobox via raw page HTML briefly
        html = await fetch(client, f"https://en.wikipedia.org/wiki/{quote(title)}")
        blob = extract + "\n" + html[:20000]
        if not emp:
            emp = parse_employees(blob)
        if not rev:
            # Wikipedia revenue tables often "US$xx billion"
            m = re.search(
                r"(?:Revenue|Turnover).*?(US\$|\$|€|£|¥)\s?([\d\.,]+)\s*(billion|million|trillion)?",
                blob,
                re.I | re.S,
            )
            if m:
                cur, num, scale = m.group(1), m.group(2), (m.group(3) or "")
                cur = "$" if "US" in cur or cur == "$" else cur
                rev = normalize_money(f"{cur}{num}{scale[:1].upper() if scale else ''}")
                if scale and not rev.endswith("(est.)"):
                    # public wiki figure — mark as reported
                    rev = rev.replace(" (est.)", "") + " (reported)"
        if emp:
            out["employees"] = emp
            out["source"] = "wikipedia"
        if rev:
            out["turnover"] = rev
            out["source"] = out.get("source") or "wikipedia"
        if extract:
            out["wiki"] = extract[:600]
    except Exception:
        pass
    return out


async def site_lookup(client: httpx.AsyncClient, website: str) -> dict[str, str]:
    out: dict[str, str] = {}
    if not website.startswith("http"):
        return out
    html = await fetch(client, website)
    if not html:
        return out
    text = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    text = re.sub(r"(?is)<style.*?>.*?</style>", " ", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text)[:8000]
    emp = parse_employees(text)
    rev = parse_revenue(text)
    if emp:
        out["employees"] = emp
        out["source"] = "website"
    if rev:
        out["turnover"] = rev
        out["source"] = out.get("source") or "website"
    return out


async def gather_scale(client: httpx.AsyncClient, row: dict[str, str]) -> dict[str, Any]:
    company = row["Company"]
    # try growjo + wiki in parallel; site optional
    g, w = await asyncio.gather(
        growjo_lookup(client, company),
        wikipedia_lookup(client, company),
    )
    site: dict[str, str] = {}
    # only hit site if still missing both
    if (not g.get("employees") and not w.get("employees")) or (
        not g.get("turnover") and not w.get("turnover")
    ):
        site = await site_lookup(client, row.get("Website") or "")

    emp = g.get("employees") or w.get("employees") or site.get("employees") or ""
    turn = g.get("turnover") or w.get("turnover") or site.get("turnover") or ""
    sources = []
    for d in (g, w, site):
        if d.get("source"):
            sources.append(d["source"])
    return {
        "company": company,
        "employees": emp,
        "turnover": turn,
        "sources": list(dict.fromkeys(sources)),
        "wiki": w.get("wiki", ""),
        "growjo_url": g.get("url", ""),
        "evidence": {
            "growjo": {k: g.get(k) for k in ("employees", "turnover", "url") if g.get(k)},
            "wikipedia": {k: w.get(k) for k in ("employees", "turnover") if w.get(k)},
            "website": {k: site.get(k) for k in ("employees", "turnover") if site.get(k)},
        },
    }


SYSTEM = """You verify/normalize Employees and Turnover/Scale for ADC market distributors.
Rules:
- Prefer verified public figures; otherwise keep third-party estimates and mark (est.).
- Employees: integer with commas, or a range like 50-200, or 'Not publicly disclosed'.
- Turnover / Scale: currency amount with M/B if known, or qualitative scale; mark (est.) if estimate.
- For subsidiaries, you may say 'Part of Parent (~X employees / $YB)' if division figures unavailable.
- Do NOT invent precise figures with fake precision. If unknown, use Not publicly disclosed.
Return JSON: {"results":[{"company":"...","employees":"...","turnover":"...","confidence":"high|medium|low","note":"..."}]}"""


def deepseek_verify(client: OpenAI, items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    payload = []
    for it in items:
        row = it["row"]
        ev = it["evidence"]
        payload.append({
            "company": row["Company"],
            "hq": row.get("Headquarters"),
            "website": row.get("Website"),
            "business_type": row.get("Business Type"),
            "current_employees": row.get("Employees"),
            "current_turnover": row.get("Turnover / Scale"),
            "web_evidence": ev,
        })
    user = {"task": "Normalize employees and turnover for each company", "companies": payload}
    for attempt in range(3):
        try:
            resp = client.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": SYSTEM},
                    {"role": "user", "content": json.dumps(user, ensure_ascii=False)},
                ],
                temperature=0.1,
                max_tokens=3500,
                response_format={"type": "json_object"},
            )
            data = json.loads(resp.choices[0].message.content or "{}")
            if isinstance(data, dict):
                for k in ("results", "companies", "items"):
                    if isinstance(data.get(k), list):
                        return data[k]
                if "company" in data:
                    return [data]
            if isinstance(data, list):
                return data
            return []
        except Exception as e:
            print(f"  DeepSeek error: {e}", flush=True)
            time.sleep(1.2 * (attempt + 1))
    return []


def apply_values(row: dict[str, str], emp: str, turn: str, tag: str) -> dict[str, str]:
    out = dict(row)
    if emp and not empty(emp):
        out["Employees"] = emp.strip()
    elif empty(out.get("Employees")):
        out["Employees"] = "Not publicly disclosed"
    if turn and not empty(turn):
        out["Turnover / Scale"] = turn.strip()
    elif empty(out.get("Turnover / Scale")):
        out["Turnover / Scale"] = "Not publicly disclosed"
    src = out.get("Data Sources") or ""
    if tag not in src:
        out["Data Sources"] = (src + f"; {tag}").strip("; ")
    return out


async def main() -> int:
    rows = load_rows(XLSX)
    print(f"Loaded {len(rows)}; model={MODEL}")

    # Pass A: known
    known_n = 0
    for i, row in enumerate(rows):
        if row["Company"] in KNOWN:
            e, t = KNOWN[row["Company"]]
            rows[i] = apply_values(row, e, t, "adc_emp_turnover_known_public")
            known_n += 1
    print(f"Pass A known applied: {known_n}")

    oa = OpenAI(
        api_key=os.getenv("OPENAI_API_KEY"),
        base_url=os.getenv("OPENAI_BASE_URL", "https://api.deepseek.com/v1"),
    )
    smoke = oa.chat.completions.create(
        model=MODEL, messages=[{"role": "user", "content": "OK"}], max_tokens=3
    )
    print(f"DeepSeek OK: {smoke.choices[0].message.content} ({smoke.model})")

    if PROGRESS.exists():
        PROGRESS.unlink()

    # Process all rows needing fill (or re-verify all for consistency)
    targets = list(range(len(rows)))
    filled_emp = filled_turn = 0

    limits = httpx.Limits(max_connections=CONCURRENCY + 2, max_keepalive_connections=CONCURRENCY)
    async with httpx.AsyncClient(
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
            )
        },
        limits=limits,
        timeout=httpx.Timeout(12.0, connect=5.0),
    ) as http:
        sem = asyncio.Semaphore(CONCURRENCY)

        async def one(idx: int) -> tuple[int, dict[str, Any]]:
            async with sem:
                ev = await gather_scale(http, rows[idx])
                await asyncio.sleep(0.08)
                return idx, ev

        for i in range(0, len(targets), BATCH):
            idxs = targets[i : i + BATCH]
            print(f"\n[{i+1}-{i+len(idxs)}/{len(targets)}] Growjo/Wiki + DeepSeek...", flush=True)
            try:
                results = await asyncio.wait_for(
                    asyncio.gather(*[one(j) for j in idxs]),
                    timeout=90.0,
                )
            except asyncio.TimeoutError:
                print("  batch timeout", flush=True)
                results = []
                for j in idxs:
                    results.append((j, {
                        "company": rows[j]["Company"], "employees": "", "turnover": "",
                        "sources": [], "wiki": "", "evidence": {},
                    }))

            items = []
            ev_by_idx = {}
            for idx, ev in results:
                ev_by_idx[idx] = ev
                items.append({"row": rows[idx], "evidence": ev})

            decisions = deepseek_verify(oa, items)
            by = {norm(str(d.get("company") or "")): d for d in decisions if d.get("company")}

            for idx in idxs:
                row = rows[idx]
                ev = ev_by_idx[idx]
                d = by.get(norm(row["Company"])) or {}

                # Prefer known overrides
                if row["Company"] in KNOWN:
                    emp, turn = KNOWN[row["Company"]]
                    tag = "adc_emp_turnover_known_public"
                else:
                    emp = str(d.get("employees") or ev.get("employees") or row.get("Employees") or "").strip()
                    turn = str(d.get("turnover") or ev.get("turnover") or row.get("Turnover / Scale") or "").strip()
                    # if deepseek says NPD but we have growjo, keep growjo with est.
                    if empty(emp) and ev.get("employees"):
                        emp = ev["employees"]
                    if empty(turn) and ev.get("turnover"):
                        turn = ev["turnover"]
                    tag = "adc_emp_turnover_growjo_wiki_deepseek"
                    if ev.get("sources"):
                        tag += "_" + "+".join(ev["sources"][:3])

                before = (row.get("Employees"), row.get("Turnover / Scale"))
                rows[idx] = apply_values(row, emp, turn, tag)
                after = (rows[idx].get("Employees"), rows[idx].get("Turnover / Scale"))
                if not empty(after[0]) and after[0].lower() != "not publicly disclosed":
                    filled_emp += 1
                if not empty(after[1]) and after[1].lower() != "not publicly disclosed":
                    filled_turn += 1
                rec = {
                    "company": row["Company"],
                    "before": before,
                    "after": after,
                    "sources": ev.get("sources"),
                    "note": d.get("note"),
                    "confidence": d.get("confidence"),
                }
                with PROGRESS.open("a", encoding="utf-8") as pf:
                    pf.write(json.dumps(rec, ensure_ascii=False) + "\n")
                print(f"  {row['Company'][:40]:40} | emp={after[0][:28]} | turn={after[1][:32]}", flush=True)
            await asyncio.sleep(0.2)

    rows.sort(key=lambda x: norm(x["Company"]))
    write_xlsx(XLSX, rows)

    emp_filled = sum(
        1 for r in rows
        if not empty(r.get("Employees")) and r.get("Employees", "").lower() != "not publicly disclosed"
    )
    turn_filled = sum(
        1 for r in rows
        if not empty(r.get("Turnover / Scale")) and r.get("Turnover / Scale", "").lower() != "not publicly disclosed"
    )
    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "model": MODEL,
        "rows": len(rows),
        "employees_filled": emp_filled,
        "turnover_filled": turn_filled,
        "employees_npd": len(rows) - emp_filled,
        "turnover_npd": len(rows) - turn_filled,
        "note": "Owler blocked (403); used Growjo estimates + Wikipedia/public + DeepSeek verify",
        "output": str(XLSX),
    }
    SUMMARY.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nDONE emp_filled={emp_filled} turnover_filled={turn_filled} / {len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
