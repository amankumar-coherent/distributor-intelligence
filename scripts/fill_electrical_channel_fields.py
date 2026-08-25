"""Fill / normalize channel columns in electrical FINAL Excel."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from vendor_intel.enrichment.known_channel_fields import sanitize_channel_fields

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)

COLS = [
    "Core Categories",
    "Specialty Focus",
    "Key Brands Represented",
    "Retail / E-commerce / Both",
    "Distribution Type",
]


def main() -> None:
    backup = XLSX.with_name(XLSX.stem + "_BEFORE_CHANNEL_FIELDS_FILL.xlsx")
    if not backup.exists():
        backup.write_bytes(XLSX.read_bytes())

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    col_idx = {c: headers.index(c) + 1 for c in COLS}

    changes: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        old_row = {c: str(ws.cell(row_idx, col_idx[c]).value or "").strip() for c in COLS}
        new_row = sanitize_channel_fields(name, old_row)
        for c in COLS:
            if new_row[c] != old_row[c]:
                ws.cell(row_idx, col_idx[c]).value = new_row[c]
                changes.append(
                    {
                        "company": name,
                        "field": c,
                        "old": old_row[c] or "(empty)",
                        "new": new_row[c],
                    }
                )

    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(XLSX)

    dump = XLSX.parent / "_channel_fields_fill_audit.json"
    dump.write_text(
        json.dumps(
            {
                "when": datetime.now().isoformat(timespec="seconds"),
                "changed_cells": len(changes),
                "changes": changes,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(f"changed_cells={len(changes)}")
    print(f"wrote {XLSX}")
    print(f"audit {dump}")


if __name__ == "__main__":
    main()
