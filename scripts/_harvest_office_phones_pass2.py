# -*- coding: utf-8 -*-
"""Second pass: remaining empty Office No. via imprint/legal/contact paths."""
from __future__ import annotations

import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urljoin, urlparse

import httpx
from openpyxl import load_workbook

from _harvest_office_phones import (
    PATHS,
    UA,
    abs_url,
    apply_excel,
    extract_phones,
    has_val,
    load_cache,
    save_cache,
)

EXTRA = (
    "/imprint",
    "/impressum",
    "/legal-notice",
    "/mentions-legales",
    "/en/contact",
    "/contacto",
    "/kontakt",
    "/pages/contact",
)


def harvest_extra(website: str) -> str:
    base = abs_url(website)
    if not base:
        return ""
    parsed = urlparse(base)
    origin = f"{parsed.scheme}://{parsed.netloc}"
    urls = [urljoin(origin + "/", p.lstrip("/")) for p in EXTRA]
    headers = {"User-Agent": UA, "Accept": "text/html,application/xhtml+xml"}
    try:
        with httpx.Client(
            follow_redirects=True,
            timeout=httpx.Timeout(8.0, connect=5.0),
            headers=headers,
        ) as client:
            for url in urls:
                try:
                    resp = client.get(url)
                except Exception:
                    continue
                if resp.status_code >= 400:
                    continue
                phone = extract_phones(resp.text or "")
                if phone:
                    return phone
    except Exception:
        return ""
    return ""


def main() -> None:
    cache = load_cache()
    company_web = {}
    for p in PATHS:
        wb = load_workbook(p, data_only=True)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        headers = None
        hr = 1
        for r in range(1, 6):
            vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
            if "Company" in vals:
                headers = vals
                hr = r
                break
        co_i = headers.index("Company")
        web_i = headers.index("Website")
        off_i = headers.index("Office No.")
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            co = str(row[co_i] or "").strip()
            if not co:
                continue
            off = str(row[off_i] or "").strip()
            web = str(row[web_i] or "").strip()
            if has_val(off):
                continue
            if abs_url(web):
                company_web[co] = web
        wb.close()

    sites = []
    seen = set()
    for web in company_web.values():
        k = abs_url(web).lower()
        if k in seen:
            continue
        if cache.get(k):
            continue
        seen.add(k)
        sites.append(web)
    print(f"PASS2 remaining websites={len(sites)} companies={len(company_web)}")
    t0 = time.time()
    done = 0
    with ThreadPoolExecutor(max_workers=28) as ex:
        futs = {ex.submit(harvest_extra, w): w for w in sites}
        for fut in as_completed(futs):
            w = futs[fut]
            key = abs_url(w).lower()
            try:
                phone = fut.result() or ""
            except Exception:
                phone = ""
            if phone:
                cache[key] = phone
            done += 1
            if done % 50 == 0 or done == len(sites):
                save_cache(cache)
                hit = sum(1 for v in cache.values() if v)
                print(f"  progress {done}/{len(sites)} cache_hits={hit} elapsed={time.time()-t0:.0f}s")
    save_cache(cache)
    fills = {}
    for co, web in company_web.items():
        phone = cache.get(abs_url(web).lower()) or ""
        if phone:
            fills[co] = phone
    print(f"PASS2 fills={len(fills)}")
    for p in PATHS:
        n = apply_excel(p, fills)
        print(f"WRITE {p.name} filled={n}")


if __name__ == "__main__":
    main()
