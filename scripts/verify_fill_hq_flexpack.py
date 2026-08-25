"""HQ fill: DDGS + optional site contact fetch + DeepSeek evidence gate.
City must appear in evidence. Never invent. Failed rows are retried (only filled skipped).
"""
from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
CKPT = FP.parent / "hq_verify_fill_ckpt.json"
AUDIT = FP.parent / "hq_verify_fill_audit.json"

API_KEY = (os.getenv("OPENAI_API_KEY") or "").strip()
BASE_URL = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()

# Agent web-verified seeds (city must be real; update Excel first, then skip in loop)
MANUAL_SEEDS: dict[str, str] = {
    "Associated Bag": "Milwaukee, Wisconsin, United States",
    "National Flexible": "Birkenshaw, Bradford, United Kingdom",
    "Bunzl": "London, United Kingdom",
    "Macfarlane Packaging": "Glasgow, United Kingdom",
    "Uline": "Pleasant Prairie, Wisconsin, United States",
    "Veritiv": "Atlanta, Georgia, United States",
    "RAJAPACK UK": "Brogborough, Bedford, United Kingdom",
    "Ahlsell": "Stockholm, Sweden",
    "Antalis": "Boulogne-Billancourt, France",
    "Officeworks": "Chadstone, Victoria, Australia",
    "Fastenal": "Winona, Minnesota, United States",
    "Lyreco": "Marly, France",
    "The Home Depot": "Atlanta, Georgia, United States",
    "W.B. Mason": "Brockton, Massachusetts, United States",
    "All American Containers": "Miami, Florida, United States",
    "Askul": "Tokyo, Japan",
    "Associated Packaging": "Gallatin, Tennessee, United States",
    "Acme Paper & Supply": "Savage, Maryland, United States",
    "Ball & Doggett": "Dandenong South, Victoria, Australia",
    "Blackwoods": "Macquarie Park, New South Wales, Australia",
    "BradyPLUS": "Las Vegas, Nevada, United States",
    "Cospak": "Minto, New South Wales, Australia",
    "Crawford Packaging": "London, Ontario, Canada",
    "Abco Kovex": "Swords, Dublin, Ireland",
    "Atlantic Packaging Products": "Scarborough, Ontario, Canada",
    "Atlantic Packaging Corp": "Wilmington, North Carolina, United States",
    "Davpack": "Derby, United Kingdom",
    "BigDug": "Gloucester, United Kingdom",
    "Boxmart": "Shenstone, Staffordshire, United Kingdom",
    "Bags Unlimited": "Rochester, New York, United States",
    "Burch Bottle & Packaging": "Queensbury, New York, United States",
    "Container and Packaging Supply": "Eagle, Idaho, United States",
    "Dayton Bag & Burlap": "Dayton, Ohio, United States",
    "IndustryBuying": "New Delhi, India",
    "SupplyOne": "Newtown Square, Pennsylvania, United States",
    "Allied Packaging": "Botany, New South Wales, Australia",
    "American Packaging Distributors": "West Chester, Pennsylvania, United States",
    "BB-Verpackungen": "Delmenhorst, Germany",
    "BVS Verpacken": "Urdorf, Switzerland",
    "Signet": "Wakerley, Queensland, Australia",
    "Allpack": "Cannock, United Kingdom",
    "Aropack": "Hythe, Kent, United Kingdom",
    "Bags Plus": "Malvern, Victoria, Australia",
    "Box Brothers": "Woodland Hills, California, United States",
    "Carolina Retail Packaging": "Lexington, South Carolina, United States",
    "Horna": "Grafenrheinfeld, Germany",
    "Chase Packaging": "Enfield, United Kingdom",
    "Complete Packaging Products": "Troy, Michigan, United States",
    "Delaware Valley Packaging": "Fairfield, California, United States",
    "Springpack": "Worcester, United Kingdom",
    "RAJA": "Tremblay-en-France, France",
    "Packaging Corporation of America": "Lake Forest, Illinois, United States",
    "Bags and Bows": "Trevose, Pennsylvania, United States",
    "Boxworks": "Bayswater North, Victoria, Australia",
    "Brown Packaging": "Anaheim, California, United States",
    "Braden Packaging & Supply": "Clinton, New York, United States",
    "Diversified Packaging Solutions": "St. Louis, Missouri, United States",
    "Nashville Wraps": "Hendersonville, Tennessee, United States",
    "The Packaging Company": "Brampton, Ontario, Canada",
    "Paper Mart": "Orange, California, United States",
}

COUNTRY_ONLY = {
    "united states", "usa", "us", "united kingdom", "uk", "great britain", "england",
    "germany", "deutschland", "australia", "canada", "france", "india", "mexico",
    "spain", "ireland", "netherlands", "sweden", "japan", "italy", "singapore",
    "new zealand", "south africa", "switzerland", "austria", "belgium", "china",
    "brazil", "poland", "portugal", "denmark", "norway", "finland",
}


def needs_hq_fill(hq: str) -> bool:
    h = (hq or "").strip()
    if not h or "not publicly" in h.lower():
        return True
    if "," in h:
        left = h.split(",")[0].strip()
        if left and left.lower() not in COUNTRY_ONLY and len(left) >= 3:
            return False
    if h.lower() in COUNTRY_ONLY:
        return True
    return "," not in h and len(h.split()) <= 3


def ddgs_text(query: str, max_results: int = 6) -> str:
    try:
        from ddgs import DDGS
    except Exception as e:
        return f"[ddgs_error] {e}"
    lines = []
    try:
        with DDGS() as ddg:
            for r in ddg.text(query, max_results=max_results):
                lines.append(f"{r.get('title','')} | {r.get('body','')} | {r.get('href','')}")
    except Exception as e:
        return f"[ddgs_error] {e}"
    return "\n".join(lines)


def fetch_site_snippet(website: str) -> str:
    if not website or website.upper() in ("NPD", "N/A", "NA"):
        return ""
    raw = website if "://" in website else "https://" + website
    try:
        host = urlparse(raw).scheme + "://" + urlparse(raw).netloc
    except Exception:
        return ""
    paths = ["/contact", "/contact-us", "/about", "/about-us", "/company", "/locations"]
    chunks = []
    headers = {"User-Agent": "Mozilla/5.0 (compatible; VendorIntelHQ/1.0)"}
    for p in paths[:3]:
        try:
            r = requests.get(host + p, headers=headers, timeout=8, allow_redirects=True)
            if r.status_code == 200 and len(r.text) > 200:
                text = re.sub(r"<script[\s\S]*?</script>", " ", r.text, flags=re.I)
                text = re.sub(r"<style[\s\S]*?</style>", " ", text, flags=re.I)
                text = re.sub(r"<[^>]+>", " ", text)
                text = re.sub(r"\s+", " ", text).strip()
                chunks.append(text[:1800])
                break
        except Exception:
            continue
    return "\n".join(chunks)[:3500]


def deepseek_hq(company: str, website: str, current: str, evidence: str) -> dict:
    if not API_KEY or not evidence.strip():
        return {"hq": None, "confidence": "none", "reason": "no_api_or_evidence"}
    from openai import OpenAI

    client = OpenAI(api_key=API_KEY, base_url=BASE_URL, timeout=45.0)
    prompt = f"""Extract HEADQUARTERS city for "{company}" ({website}). Current Excel: {current}

Rules: ONLY from EVIDENCE. Prefer HQ / head office / corporate headquarters / registered office.
Format "City, Country" or "City, State/Region, Country". City MUST appear verbatim in evidence.
If only country is known, return hq null. Do not invent.
JSON only: {{"hq":"..."|null,"city":"..."|null,"confidence":"high"|"medium"|"low","reason":"..."}}

EVIDENCE:
{evidence[:7500]}
"""
    try:
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": "Extract HQ city only from evidence. Never invent. JSON only."},
                {"role": "user", "content": prompt},
            ],
            temperature=0,
        )
        text = (resp.choices[0].message.content or "").strip()
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I | re.M).strip()
        return json.loads(text)
    except Exception as e:
        return {"hq": None, "confidence": "error", "reason": str(e)[:160]}


def hq_format_ok(hq: str) -> bool:
    """Require real city-level format; reject bare 'City, United States' without state."""
    h = (hq or "").strip()
    if not h or "," not in h:
        return False
    if needs_hq_fill(h):
        return False
    parts = [p.strip() for p in h.split(",") if p.strip()]
    if len(parts) < 2:
        return False
    # US/Canada/Australia usually need region; allow 2-part for small countries
    last = parts[-1].lower()
    if last in ("united states", "usa", "us", "canada", "australia") and len(parts) < 3:
        return False
    return True


def city_in_evidence(city: str, evidence: str) -> bool:
    if not city or len(city.strip()) < 3:
        return False
    return re.search(re.escape(city.strip()), evidence, re.I) is not None


def gather_evidence(company: str, website: str) -> str:
    host = ""
    try:
        host = urlparse(website if "://" in website else "https://" + website).netloc.replace("www.", "")
    except Exception:
        pass
    parts = [
        ddgs_text(f'"{company}" headquarters OR "head office" OR "corporate headquarters"'),
        ddgs_text(f'"{company}" "{host}" address OR contact OR location') if host else "",
        fetch_site_snippet(website),
    ]
    return "\n---\n".join(p for p in parts if p)[:9000]


def main() -> None:
    apply = os.environ.get("APPLY_HQ_FIXES", "1") == "1"
    wb = load_workbook(FP)
    ws = wb.active

    # Apply agent-verified seeds first
    seeded = 0
    for r in ws.iter_rows(min_row=2):
        name = str(r[0].value or "").strip()
        if name in MANUAL_SEEDS:
            r[3].value = MANUAL_SEEDS[name]
            seeded += 1
    if apply and seeded:
        wb.save(FP)
        print(f"seeded={seeded}", flush=True)

    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        rows.append(
            {
                "excel_row": i,
                "company": str(r[0].value or "").strip(),
                "website": str(r[1].value or "").strip(),
                "hq": str(r[3].value or "").strip(),
            }
        )
    targets = [r for r in rows if needs_hq_fill(r["hq"])]

    done = {}
    if CKPT.exists():
        raw = json.loads(CKPT.read_text(encoding="utf-8")).get("done") or {}
        # Only keep successful fills; retry all failures
        done = {k: v for k, v in raw.items() if v.get("status") == "filled" and v.get("hq")}

    filled = failed = 0
    for idx, row in enumerate(targets, 1):
        key = row["company"]
        if key in done:
            filled += 1
            if apply and done[key].get("hq"):
                ws.cell(row["excel_row"], 4).value = done[key]["hq"]
            continue

        print(f"[{idx}/{len(targets)}] {key}", flush=True)
        evidence = gather_evidence(key, row["website"])
        pick = deepseek_hq(key, row["website"], row["hq"], evidence)
        hq_val = pick.get("hq")
        city = pick.get("city") or (str(hq_val).split(",")[0].strip() if hq_val else None)
        conf = (pick.get("confidence") or "").lower()
        ok = (
            hq_val
            and conf in ("high", "medium")
            and city_in_evidence(str(city), evidence)
            and hq_format_ok(str(hq_val))
        )
        if ok:
            item = {**row, "status": "filled", "hq": str(hq_val).strip(), "old": row["hq"], "gate": pick}
            filled += 1
            if apply:
                ws.cell(row["excel_row"], 4).value = str(hq_val).strip()
            print(f"  OK {hq_val}", flush=True)
        else:
            item = {**row, "status": "failed", "hq": None, "gate": pick, "evidence_len": len(evidence)}
            failed += 1
            print(f"  FAIL {conf} reason={str(pick.get('reason',''))[:80]}", flush=True)
        done[key] = item
        if idx % 8 == 0:
            CKPT.write_text(
                json.dumps({"updated": datetime.now().isoformat(timespec="seconds"), "done": done}, indent=2),
                encoding="utf-8",
            )
            if apply:
                wb.save(FP)
            print(f"  save filled={filled} failed={failed}", flush=True)
        time.sleep(0.05)

    if apply:
        wb.save(FP)
    CKPT.write_text(
        json.dumps({"updated": datetime.now().isoformat(timespec="seconds"), "done": done}, indent=2),
        encoding="utf-8",
    )

    weak = cityish = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        h = str(r[3] or "").strip()
        if needs_hq_fill(h):
            weak += 1
        else:
            cityish += 1
    audit = {
        "updated": datetime.now().isoformat(timespec="seconds"),
        "targets": len(targets),
        "filled": filled,
        "failed": failed,
        "manual_seeds": seeded,
        "sheet_cityish_hq": cityish,
        "sheet_still_weak_hq": weak,
        "sample_filled": [
            {"company": v["company"], "old": v.get("old"), "hq": v.get("hq")}
            for v in done.values()
            if v.get("status") == "filled"
        ][:80],
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps({k: audit[k] for k in ["targets", "filled", "failed", "manual_seeds", "sheet_cityish_hq", "sheet_still_weak_hq"]}, indent=2))


if __name__ == "__main__":
    main()
