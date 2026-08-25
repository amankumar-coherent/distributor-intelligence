#!/usr/bin/env python3
"""Verify/fix avocado OFFICE_SLIM: legal names, websites, continents, phones, summaries."""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

X = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
AUDIT = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_verify_all_audit.json"
)
NPD = "Not publicly disclosed"

# Known legal-name / website / contact overrides (agent web-verified)
OVERRIDES: dict[str, dict] = {
    "Dot Foods, Inc.": {
        "Website": "https://www.dotfoods.com/",
    },
    "Dot Foods": {
        "Company": "Dot Foods, Inc.",
        "Website": "https://www.dotfoods.com/",
        "Headquarters": "1 Dot Way, Mt. Sterling, Illinois, United States",
        "Continent / Geography": "North America",
        "Operational Presence": "United States; Canada",
        "Office No.": "+1 217-773-4411",
    },
    "Eosta": {
        "Company": "Eosta B.V.",
        "Website": "https://www.eosta.com/en/",
        "Key Brands Represented": "Nature & More, others",
        "Office Email": "info@eosta.com",
        "Office No.": "+31 (0)180 63 55 00",
    },
    "Liberty Richter": {
        "Company": "Liberty Richter (World Finer Foods)",
        "Website": "https://libertyrichter.com/",
        "Key Brands Represented": "Reese, DaVinci, Ty Ling, others",
    },
    "Al Maya Group": {
        "Company": "Al Maya Group",
        "Website": "https://www.almaya.ae/",
        "Key Brands Represented": "Al Maya Supermarket, Borders, Paperchase, others",
        "Employees": "1001-5000",
    },
    "Londis Wholesale": {
        "Company": "Londis (Booker Group / Tesco)",
        "Website": "https://www.booker.co.uk/content/pages/retail/londis.html",
        "Key Brands Represented": "Londis, Jacks, Euro Shopper, others",
        "Office No.": "0808 178 8644",
    },
    "Grupo Herdez Distribution": {
        "Company": "Grupo Herdez, S.A.B. de C.V.",
        "Website": "https://grupoherdez.com.mx/",
        "Key Brands Represented": "Herdez, Doña María, Del Fuerte, McCormick, Barilla, Kikkoman, others",
        "Office No.": "+52 55 5201-5655",
        "Office Email": "invrel@herdez.com",
        "Ownership": "Public (BMV: HERDEZ)",
        "Employees": "10001-50000",
    },
    "Horeca Trade": {
        "Company": "Bidfood UAE (formerly Horeca Trade LLC)",
        "Website": "https://www.bidfoodme.com/presence/uae",
        "Key Brands Represented": "S.Pellegrino, Acqua Panna, Lamb Weston, Anchor, Fever-Tree, Nutella, Impossible, others",
        "Office No.": "+971 4 338 8772",
        "Ownership": "Public (Bidcorp / Bidfood)",
        "Employees": "201-500",
    },
    "Levarht": {
        "Company": "Levarht (J.M. Levarht & Zonen B.V.)",
        "Website": "https://www.levarht.nl/en",
        "Key Brands Represented": "Levarht, others",
        "Ownership": "Privately held (administration history)",
        "Summary": "Levarht — Dutch fresh produce trader (website active); historic J.M. Levarht & Zonen entered administration in 2022; verify trading status before outreach.",
    },
    "GreenGold USA": {
        "Company": "GreenGold USA",
        "Website": "https://greengoldusa.com/",
        "Headquarters": "Cleveland, Ohio, United States",
        "Continent / Geography": "North America",
        "Operational Presence": "United States; Colombia",
        "Key Brands Represented": "GreenGold, others",
        "Office Email": "wholesale@greengoldusa.com",
        "Ownership": "Privately held",
        "Employees": "1-10",
        "Founded": "2020",
        "Country Code": "US",
        "Region Code": "NA",
        "Core Categories": "Avocado oil; avocado spreads/guacamole; wholesale edible oils, others",
        "Specialty Focus": "Colombian avocado oil (extra virgin cold-pressed & refined) wholesale to US retail/foodservice/distributors, others",
    },
    "Fresh Del Monte Europe": {
        "Website": "https://www.delmontefresh.com/",
        "Company": "Fresh Del Monte Produce Europe",
    },
    "Fresh Produce Consortium Members Distribution": {
        "Website": "https://www.freshproduce.org.uk/",
        "Company": "Fresh Produce Consortium (members distribution)",
    },
    "All Organic Treasures GmbH": {
        "Website": "https://www.aot.de/",
    },
    "Fairway Foodservice": {
        "Website": "https://www.fairwayfoodservice.com/",
    },
    "I-D Foods": {
        "Company": "I-D Foods Corporation",
        "Website": "https://idfoods.com/",
    },
    "Core-Mark": {
        "Company": "Core-Mark International, Inc. (Performance Food Group)",
        "Website": "https://www.core-mark.com/",
        "Ownership": "Public (NYSE: PFGC)",
    },
    "Vistar": {
        "Company": "Vistar (Performance Food Group)",
        "Website": "https://vistar.com/",
        "Ownership": "Public (NYSE: PFGC)",
    },
    "Sysco": {
        "Company": "Sysco Corporation",
        "Website": "https://www.sysco.com/",
        "Ownership": "Public (NYSE: SYY)",
    },
    "US Foods": {
        "Company": "US Foods, Inc.",
        "Website": "https://www.usfoods.com/",
        "Ownership": "Public (NYSE: USFD)",
    },
    "Performance Food Group": {
        "Company": "Performance Food Group Company",
        "Website": "https://www.pfgc.com/",
        "Ownership": "Public (NYSE: PFGC)",
    },
    "UNFI (United Natural Foods)": {
        "Company": "United Natural Foods, Inc. (UNFI)",
        "Website": "https://www.unfi.com/",
        "Ownership": "Public (NYSE: UNFI)",
    },
    "Supervalu Wholesale": {
        "Company": "Supervalu Wholesale (UNFI)",
        "Website": "https://www.unfi.com/",
        "Ownership": "Public (NYSE: UNFI)",
    },
    "SuperValu": {
        "Company": "Supervalu (UNFI)",
        "Website": "https://www.unfi.com/",
        "Ownership": "Public (NYSE: UNFI)",
    },
    "Associated Wholesale Grocers": {
        "Company": "Associated Wholesale Grocers, Inc.",
        "Website": "https://www.awginc.com/",
    },
    "C&S Wholesale Grocers": {
        "Company": "C&S Wholesale Grocers, LLC",
        "Website": "https://www.cswg.com/",
    },
    "Gordon Food Service": {
        "Company": "Gordon Food Service, Inc.",
        "Website": "https://www.gfs.com/",
    },
    "KeHE Distributors": {
        "Company": "KeHE Distributors, LLC",
        "Website": "https://www.kehe.com/",
    },
    "KeHE Specialty": {
        "Company": "KeHE Specialty (KeHE Distributors, LLC)",
        "Website": "https://www.kehe.com/",
    },
    "Brenntag": {
        "Company": "Brenntag SE",
        "Website": "https://www.brenntag.com/",
        "Ownership": "Public (ETR: BNR)",
    },
    "Barentz": {
        "Company": "Barentz International B.V.",
        "Website": "https://www.barentz.com/",
    },
    "Metro AG": {
        "Company": "METRO AG",
        "Website": "https://www.metroag.de/",
        "Ownership": "Public (ETR: B4B)",
    },
    "Costco": {
        "Company": "Costco Wholesale Corporation",
        "Website": "https://www.costco.com/",
        "Ownership": "Public (NASDAQ: COST)",
    },
    "Walmart": {
        "Company": "Walmart Inc.",
        "Website": "https://www.walmart.com/",
        "Ownership": "Public (NYSE: WMT)",
    },
    "Martin-Brower": {
        "Company": "The Martin-Brower Company, L.L.C.",
        "Website": "https://www.martinbrower.com/",
    },
    "McLane Company": {
        "Company": "McLane Company, Inc.",
        "Website": "https://www.mclaneco.com/",
    },
    "Reyes Holdings": {
        "Company": "Reyes Holdings, L.L.C.",
        "Website": "https://www.reyesholdings.com/",
    },
    "Labatt Food Service": {
        "Company": "Labatt Food Service",
        "Website": "https://www.labattfood.com/",
    },
    "Scavuzzo's": {
        "Company": "Scavuzzo's Foodservice",
        "Website": "https://scavuzzos.com/",
    },
    "Patties Food Group Distribution": {
        "Company": "Patties Food Group",
        "Website": "https://pattiesfoodgroup.com/",
    },
    "Olam Food Ingredients Distribution": {
        "Company": "ofi (Olam Food Ingredients)",
        "Website": "https://www.ofi.com/",
    },
    "Nutra Food Ingredients": {
        "Company": "Nutra Food Ingredients, LLC",
        "Website": "https://nutrafoodingredients.com/",
    },
    "Gillco Ingredients": {
        "Company": "Gillco Ingredients (Azelis)",
        "Website": "https://gillco.com/",
        "Office No.": NPD,
    },
    "ITOCHU Food Distribution": {
        "Company": "ITOCHU-SHOKUHIN Co., Ltd.",
        "Website": "https://www.itochu-shokuhin.com/",
    },
    "Mitsubishi Corporation Food Distribution": {
        "Company": "Mitsubishi Shokuhin Co., Ltd.",
        "Website": "https://www.mitsubishi-shokuhin.com/",
    },
    "Jan Oskam": {
        "Company": "Jan Oskam B.V.",
        "Website": "https://www.janoskam.nl/",
    },
    "Food Supplies Company": {
        "Company": "The Food Supply Company Ltd. (FOOSCO)",
        "Website": "https://www.foosco.com/",
    },
    "Gourmet Gulf Distribution": {
        "Company": "Gourmet Gulf LLC",
        "Website": "https://www.gourmetgulf.com/",
    },
    "Quimidroga": {
        "Company": "Quimidroga, S.A.",
        "Website": "https://www.quimidroga.com/",
        "Ownership": "Privately held",
        "Employees": "201-500",
        "Country Code": "ES",
        "Region Code": "EU",
    },
    "Circle Ingredients": {
        "Company": "Circle Ingredients B.V.",
        "Website": "https://circleingredients.com/",
        "Ownership": "Privately held",
        "Employees": "11-50",
        "Country Code": "NL",
        "Region Code": "EU",
    },
    "So African Natural Oils": {
        "Company": "So African Natural Oils (Pty) Ltd",
        "Website": "https://www.soafrican.co.za/",
        "Ownership": "Privately held",
        "Employees": "11-50",
        "Country Code": "ZA",
        "Region Code": "AF",
        "Office No.": NPD,
    },
    "Cavin Oil Products": {
        "Company": "Cavin Oil Products / Cavin Resource",
        "Website": "https://cavin-resource.com/",
        "Ownership": "Privately held",
        "Employees": "11-50",
        "Office No.": NPD,
    },
    "AOT GmbH": {
        "Website": "https://www.aot.de/",
    },
    "Diconsa": {
        "Website": "https://www.gob.mx/seguridadalimentar",
        "Company": "DICONSA (Seguridad Alimentaria Mexicana)",
    },
    "Sime Darby Oils Distribution": {
        "Website": "https://www.simedarbyoils.com/",
    },
    "Adams Australia": {
        "Company": "Adams Australia Pty Ltd",
        "Website": "https://www.adamsaustralia.com.au/",
    },
    "American International Foods": {
        "Company": "American International Foods, Inc.",
        "Website": "https://www.americaninternationalfoods.com/",
    },
    "Pratts Wholesale": {
        "Company": "Pratts Wholesale Food Service Ltd.",
        "Website": "https://pratts.com/",
    },
    "Swire Foods Distribution": {
        "Company": "Swire Foods / Swire Coca-Cola",
        "Website": "https://www.swire.com/en/businesses/beverages_foodchain.php",
    },
}

CC_TO_CONT = {
    "US": "North America",
    "CA": "North America",
    "MX": "Latin America",
    "GB": "Europe",
    "DE": "Europe",
    "FR": "Europe",
    "NL": "Europe",
    "BE": "Europe",
    "ES": "Europe",
    "IT": "Europe",
    "CH": "Europe",
    "AT": "Europe",
    "PL": "Europe",
    "SE": "Europe",
    "IE": "Europe",
    "PT": "Europe",
    "DK": "Europe",
    "FI": "Europe",
    "NO": "Europe",
    "LT": "Europe",
    "UA": "Europe",
    "AE": "Middle East",
    "SA": "Middle East",
    "QA": "Middle East",
    "BH": "Middle East",
    "OM": "Middle East",
    "KW": "Middle East",
    "JO": "Middle East",
    "IL": "Middle East",
    "AU": "Asia-Pacific",
    "NZ": "Asia-Pacific",
    "JP": "Asia-Pacific",
    "CN": "Asia-Pacific",
    "HK": "Asia-Pacific",
    "TW": "Asia-Pacific",
    "SG": "Asia-Pacific",
    "MY": "Asia-Pacific",
    "TH": "Asia-Pacific",
    "ID": "Asia-Pacific",
    "KR": "Asia-Pacific",
    "IN": "Asia-Pacific",
    "PH": "Asia-Pacific",
    "VN": "Asia-Pacific",
    "ZA": "Africa",
    "NG": "Africa",
    "EG": "Africa",
    "KE": "Africa",
    "BR": "Latin America",
    "AR": "Latin America",
    "CL": "Latin America",
    "CO": "Latin America",
    "PE": "Latin America",
    "UY": "Latin America",
    "CR": "Latin America",
}


def norm_web(w: str) -> str:
    w = (w or "").strip()
    if not w or w == NPD:
        return w
    if w.startswith("www."):
        return "https://" + w
    if w.startswith("http://"):
        return "https://" + w[len("http://") :]
    return w


def junk_phone(p: str) -> bool:
    p = (p or "").strip()
    if not p or p == NPD:
        return False
    digits = re.sub(r"\D", "", p)
    if digits in ("4294967295", "2147483647"):
        return True
    if digits.startswith("42949"):
        return True
    # clearly too short garbage
    if len(digits) < 7:
        return True
    return False


def make_summary(co: str, d: dict) -> str:
    hq = d.get("hq") or NPD
    web = d.get("web") or NPD
    core = d.get("core") or NPD
    focus = d.get("focus") or NPD
    brands = d.get("brands") or NPD
    op = d.get("op") or NPD
    founded = d.get("founded") or ""
    bits = [co]
    if founded and founded != NPD:
        bits.append(f"founded {founded}")
    if hq and hq != NPD:
        bits.append(f"HQ {hq}")
    if op and op != NPD:
        bits.append(f"ops: {op}")
    head = " — ".join(bits[:1]) + (f" ({'; '.join(bits[1:])})" if len(bits) > 1 else "")
    body = f"{head}. Core: {core}. Focus: {focus}."
    if brands and brands != NPD:
        body += f" Brands: {brands}."
    if web and web != NPD:
        body += f" Web: {web}."
    return body[:500]


wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

stats = {
    "override_rows": 0,
    "continent_fixed": 0,
    "website_normalized": 0,
    "phone_cleared": 0,
    "summary_rewritten": 0,
    "field_writes": 0,
}

for r in range(3, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if not co:
        continue

    # Apply overrides by current name
    if co in OVERRIDES:
        for k, v in OVERRIDES[co].items():
            if k in ic:
                ws.cell(r, ic[k]).value = v
                stats["field_writes"] += 1
        stats["override_rows"] += 1
        co = str(ws.cell(r, ic["Company"]).value or "").strip()

    # Normalize website
    web = str(ws.cell(r, ic["Website"]).value or "").strip()
    nw = norm_web(web)
    if nw != web:
        ws.cell(r, ic["Website"]).value = nw
        stats["website_normalized"] += 1
        stats["field_writes"] += 1

    # Continent from country code if weak
    cont = str(ws.cell(r, ic["Continent / Geography"]).value or "").strip()
    cc = str(ws.cell(r, ic["Country Code"]).value or "").strip().upper()
    weak_cont = (
        not cont
        or cont == NPD
        or cont.startswith("NA(")
        or cont.startswith("EU(")
        or cont.startswith("APAC(")
        or "(" in cont[:4]
    )
    if weak_cont and cc in CC_TO_CONT:
        ws.cell(r, ic["Continent / Geography"]).value = CC_TO_CONT[cc]
        stats["continent_fixed"] += 1
        stats["field_writes"] += 1
    elif cont.startswith("NA("):
        ws.cell(r, ic["Continent / Geography"]).value = "North America"
        stats["continent_fixed"] += 1
        stats["field_writes"] += 1
    elif cont in ("NA", "North america"):
        ws.cell(r, ic["Continent / Geography"]).value = "North America"
        stats["continent_fixed"] += 1
        stats["field_writes"] += 1
    elif cont in ("EU", "Europe "):
        ws.cell(r, ic["Continent / Geography"]).value = "Europe"
        stats["continent_fixed"] += 1
        stats["field_writes"] += 1
    elif cont in ("APAC", "Asia Pacific", "Asia-Pacific "):
        ws.cell(r, ic["Continent / Geography"]).value = "Asia-Pacific"
        stats["continent_fixed"] += 1
        stats["field_writes"] += 1
    elif cont in ("ME", "Middle east"):
        ws.cell(r, ic["Continent / Geography"]).value = "Middle East"
        stats["continent_fixed"] += 1
        stats["field_writes"] += 1
    elif cont in ("LATAM", "LatAm"):
        ws.cell(r, ic["Continent / Geography"]).value = "Latin America"
        stats["continent_fixed"] += 1
        stats["field_writes"] += 1

    # Junk phones → NPD
    phone = str(ws.cell(r, ic["Office No."]).value or "").strip()
    if junk_phone(phone):
        ws.cell(r, ic["Office No."]).value = NPD
        stats["phone_cleared"] += 1
        stats["field_writes"] += 1

    # Ownership/Employees empty → NPD
    for col in ("Ownership", "Employees", "Office Email"):
        val = str(ws.cell(r, ic[col]).value or "").strip()
        if not val:
            ws.cell(r, ic[col]).value = NPD
            stats["field_writes"] += 1

    # Rewrite weak summaries from verified fields
    summ = str(ws.cell(r, ic["Summary"]).value or "").strip()
    if (
        not summ
        or summ == NPD
        or "discovered via web search" in summ.lower()
        or summ.startswith("Not publicly")
    ):
        d = {
            "hq": str(ws.cell(r, ic["Headquarters"]).value or "").strip(),
            "web": str(ws.cell(r, ic["Website"]).value or "").strip(),
            "core": str(ws.cell(r, ic["Core Categories"]).value or "").strip(),
            "focus": str(ws.cell(r, ic["Specialty Focus"]).value or "").strip(),
            "brands": str(ws.cell(r, ic["Key Brands Represented"]).value or "").strip(),
            "op": str(ws.cell(r, ic["Operational Presence"]).value or "").strip(),
            "founded": str(ws.cell(r, ic["Founded"]).value or "").strip(),
        }
        ws.cell(r, ic["Summary"]).value = make_summary(co, d)
        stats["summary_rewritten"] += 1
        stats["field_writes"] += 1

wb.save(X)
wb.close()

# Re-scan issues
wb = openpyxl.load_workbook(X, data_only=True)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i for i, h in enumerate(hdr) if h}
bad_web_re = re.compile(
    r"tripadvisor|foodnetwork|dictionary\.cambridge|salesjump|commentcamarche|greengold\.tv|askjan\.org|truenutra\.com|fresh\.com/us|timeanddate|instagram\.com|liberty\.edu|ftd\.com|al\.com/$",
    re.I,
)
remain = []
for row in ws.iter_rows(min_row=3, values_only=True):
    c = str(row[ic["Company"]] or "").strip()
    if not c:
        continue
    web = str(row[ic["Website"]] or "").strip()
    summ = str(row[ic["Summary"]] or "").strip()
    flags = []
    if not web or web == NPD:
        flags.append("no_web")
    elif bad_web_re.search(web):
        flags.append("bad_web")
    if "discovered via web search" in summ.lower():
        flags.append("weak_sum")
    if flags:
        remain.append({"company": c, "web": web, "flags": flags})
total = sum(1 for row in ws.iter_rows(min_row=3, values_only=True) if str(row[ic["Company"]] or "").strip())
wb.close()

audit = {
    "ts": datetime.now(timezone.utc).isoformat(),
    "stats": stats,
    "total": total,
    "remaining_hard_issues": remain,
}
AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
print(json.dumps({"stats": stats, "total": total, "remaining_hard": len(remain)}, indent=2))
for x in remain:
    print(x["company"], "|", x["web"][:70], "|", ",".join(x["flags"]))
