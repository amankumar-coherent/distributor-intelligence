# -*- coding: utf-8 -*-
"""One-shot agent web-search fills: fake websites + sales/procurement contacts."""
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook

NPD = "Not publicly disclosed"

FILES = {
    "electrical": Path(
        r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
        r"\global_electrical_distribution_market_global_FINAL.xlsx"
    ),
    "avocado": Path(
        r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
        r"\global_avocado_oil_market.xlsx"
    ),
    "smartphone": Path(
        r"D:\vendor-intel\output\web_expand\global_smartphone_market_global"
        r"\global_smartphone_market.xlsx"
    ),
    "firewall": Path(
        r"D:\vendor-intel\output\run4\global_firewall_market_global"
        r"\global_firewall_market.xlsx"
    ),
}

# Website-only corrections (verified official domains).
WEBSITES = {
    "avocado": {
        "Der kleine Gourmet GmbH": "https://der-kleine-gourmet.de/",
        "ESENCIA ANDALUSI GOURMET": "https://www.virgenextra.com/",
        "Grupo Éxito Wholesale": "https://grupoexito.com.co/",
        "MIXT SP ZOO": "https://mixtsp.com/",
        "Range Products": "https://www.rangeproducts.com.au/",
        "Sam's Club Mexico": "https://www.sams.com.mx/",
        "Sanabio": "https://sanabio.bio/",
        "Soapgoods": "https://soapgoods.com/",
        "The Bulk Cart": "https://thebulkcart.com/",
        "Windy Point Soap Making Supplies": "https://windypointsoap.com/",
        "Black Swan Distribution": NPD,
        "DIVERSE WORLDWIDE TRADING": NPD,
    },
    "smartphone": {
        "A&H TECH LTD": "https://anhtech.uk/",
    },
    "firewall": {
        "Datagate": "https://datagate-i.com/",
        "Nealar": NPD,
    },
}

# Contact fills: only write if Contact Person currently empty.
CONTACTS = {
    "avocado": {
        "Grupo Éxito Wholesale": {
            "person": "Luz María Ferrer Serna",
            "role": "Vice President Commercial and Supply",
            "li": "https://www.linkedin.com/in/luz-mar%C3%ADa-ferrer-serna-b30b91146",
        },
        "Sam's Club Mexico": {
            "person": "Rubén González López",
            "role": "Senior Commercial Director, Perishables",
            "li": "https://www.linkedin.com/in/rubengonlop",
        },
        "Bulk Apothecary": {
            "person": "Nick Stoll",
            "role": "Director of Global Sourcing",
            "li": "https://www.linkedin.com/in/nicholas-r-stoll",
        },
        "Nisshin OilliO Distribution": {
            "person": "Taiji Teraguchi",
            "role": "Director, National Account Sales",
            "li": NPD,
        },
    },
    "electrical": {
        "Home Depot Canada": {
            "person": "Peter Muench",
            "role": "Vice President Finance and Procurement",
            "li": "https://www.linkedin.com/in/peter-muench-1363901",
        },
        "Zahid Group": {
            "person": "Feher Alghalib Alshareif",
            "role": "Sales Director",
            "li": "https://www.linkedin.com/in/feher-alghalib-alshareif-417a8017",
        },
        "Womack Electric Supply": {
            "person": "Robert Herring",
            "role": "Area Sales Manager",
            "li": "https://www.linkedin.com/in/robert-herring-417a908b",
        },
    },
    "smartphone": {
        "Celltel Middle East L.L.C.": {
            "person": "Mohammed Hussain",
            "role": "Sales Manager",
            "li": "https://www.linkedin.com/in/mohammed-hussain-89673b35",
        },
    },
    "firewall": {
        "Liquid Intelligent Technologies": {
            "person": "Vinay Hiralall",
            "role": "Chief Commercial Officer",
            "li": "https://www.linkedin.com/in/vinayhiralall",
        },
        "Sonda": {
            "person": "Rodrigo Garcia Cortes",
            "role": "Sales Manager",
            "li": "https://www.linkedin.com/in/rodrigo-garcia-cortes-ventas-tecnologia-informacion-integracion",
        },
        "Exel del Norte": {
            "person": "Sergio Diaz Cervantes",
            "role": "Commercial Director",
            "li": "https://www.linkedin.com/in/sergio-diaz-cervantes-3398522a",
        },
        "Datagate": {
            "person": "Ross Tyler",
            "role": "VP Business Development",
            "li": "https://www.linkedin.com/in/rossntyler",
        },
        "Arabian Advanced Systems": {
            "person": "Maan Badahdah",
            "role": "Director of Business Development",
            "li": "https://www.linkedin.com/in/maanbadahdah",
        },
    },
}

EMPTY = {"", "nan", "none", "n/a", "na", "not publicly disclosed", "-", "—", "unknown", "tbd", "null"}


def _headers(ws):
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            return r, {h: i + 1 for i, h in enumerate(vals) if h}
    raise RuntimeError("headers not found")


def _empty(v):
    return (str(v or "").strip().lower() in EMPTY) or (v is None)


def save_atomic(wb, path: Path):
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def main():
    web_n = 0
    cp_n = 0
    details = []
    for mkt, path in FILES.items():
        wb = load_workbook(path)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        hr, col = _headers(ws)
        c_co = col["Company"]
        c_web = col["Website"]
        c_cp = col["Contact Person"]
        c_role = col["Role"]
        c_li = col["LinkedIn"]
        for row in range(hr + 1, ws.max_row + 1):
            co = str(ws.cell(row, c_co).value or "").strip()
            if not co:
                continue
            if mkt in WEBSITES and co in WEBSITES[mkt]:
                new_web = WEBSITES[mkt][co]
                old = str(ws.cell(row, c_web).value or "").strip()
                if old != new_web:
                    ws.cell(row, c_web).value = new_web
                    web_n += 1
                    details.append(f"WEB {mkt} | {co} -> {new_web}")
            if mkt in CONTACTS and co in CONTACTS[mkt] and _empty(ws.cell(row, c_cp).value):
                rec = CONTACTS[mkt][co]
                ws.cell(row, c_cp).value = rec["person"]
                ws.cell(row, c_role).value = rec["role"]
                ws.cell(row, c_li).value = rec["li"]
                cp_n += 1
                details.append(f"CP {mkt} | {co} | {rec['person']} | {rec['role']}")
        save_atomic(wb, path)
        wb.close()
        print("saved", mkt, path.name)
    print("website_updates", web_n)
    print("contact_fills", cp_n)
    for d in details:
        print(d)


if __name__ == "__main__":
    main()
