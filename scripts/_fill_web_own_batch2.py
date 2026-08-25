# -*- coding: utf-8 -*-
"""Pass 2: remaining empty Contact Person + fake website fixes."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

NPD = "Not publicly disclosed"

FILES = {
    "electrical": Path(
        r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
        r"\global_electrical_distribution_market_global_FINAL.xlsx"
    ),
    "avocado": Path(
        r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
        r"\global_avocado_oil_market.xlsx"
    ),
    "smartphone": Path(
        r"D:\vendor-intel\output\web_expand\global_smartphone_market_global"
        r"\global_smartphone_market.xlsx"
    ),
    "firewall": Path(
        r"D:\vendor-intel\output\run4\global_firewall_market_global"
        r"\global_firewall_market.xlsx"
    ),
}

WEBSITES = {
    "electrical": {
        "Haji Husein Alireza & Co.": "https://www.hha.com.sa",
    },
    "smartphone": {
        "The Source": "https://www.thesource.ca",
        "T-Gaia": "https://www.t-gaia.co.jp/",
        "TRSCO Group": NPD,
    },
    "firewall": {
        "Mindware": "https://mindware.net/",
    },
}

CONTACTS = {
    "avocado": {
        "Mitsubishi Corporation Food Distribution": {
            "person": "Hirohide Hosoda",
            "role": "Executive Vice President (COO, Sales)",
            "li": NPD,
        },
    },
    "smartphone": {
        "Westcoast": {
            "person": "Phil Bell",
            "role": "Chief Revenue Officer",
            "li": "https://www.linkedin.com/in/phil-bell-57547018",
        },
        "Zones": {
            "person": "Sean Hobday",
            "role": "SVP of Global Sales",
            "li": "https://www.linkedin.com/in/seanhobday",
        },
        "TelePart": {
            "person": "Kathrin Wasselowski",
            "role": "Head of Sales & Service",
            "li": "https://www.linkedin.com/in/kathrin-wasselowski-32b350149",
        },
        "TelePart Distribution GmbH": {
            "person": "Kathrin Wasselowski",
            "role": "Head of Sales & Service",
            "li": "https://www.linkedin.com/in/kathrin-wasselowski-32b350149",
        },
        "iByte Atacado": {
            "person": "Marcelo Igreja",
            "role": "Commercial Director",
            "li": "https://br.linkedin.com/in/marcelo-igreja-treinamentos",
        },
        "Visions Electronics": {
            "person": "Kyle Frizzley",
            "role": "National Sales Manager",
            "li": "https://www.linkedin.com/in/kyle-frizzley-1a8973165",
        },
    },
    "firewall": {
        "Distritech Maghreb": {
            "person": "Alejandro Mesa Correa",
            "role": "Commercial Director",
            "li": "https://www.linkedin.com/in/alejandro-mesa-correa",
        },
        "MicroGlobal": {
            "person": "Christian Font",
            "role": "Sales Manager",
            "li": "https://ar.linkedin.com/in/christian-font-51680625",
        },
    },
}

EMPTY = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}


def _headers(ws):
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            return r, {h: i + 1 for i, h in enumerate(vals) if h}
    raise RuntimeError("headers not found")


def _empty(v):
    return (str(v or "").strip().lower() in EMPTY) or (v is None)


def save_atomic(wb, path: Path):
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def main():
    web_n = 0
    cp_n = 0
    details = []
    for mkt, path in FILES.items():
        wb = load_workbook(path)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        hr, col = _headers(ws)
        c_co = col["Company"]
        c_web = col["Website"]
        c_cp = col["Contact Person"]
        c_role = col["Role"]
        c_li = col["LinkedIn"]
        for row in range(hr + 1, ws.max_row + 1):
            co = str(ws.cell(row, c_co).value or "").strip()
            if not co:
                continue
            if mkt in WEBSITES and co in WEBSITES[mkt]:
                new_web = WEBSITES[mkt][co]
                old = str(ws.cell(row, c_web).value or "").strip()
                if old != new_web:
                    ws.cell(row, c_web).value = new_web
                    web_n += 1
                    details.append(f"WEB {mkt} | {co} -> {new_web}")
            if mkt in CONTACTS and co in CONTACTS[mkt] and _empty(ws.cell(row, c_cp).value):
                rec = CONTACTS[mkt][co]
                ws.cell(row, c_cp).value = rec["person"]
                ws.cell(row, c_role).value = rec["role"]
                ws.cell(row, c_li).value = rec["li"]
                cp_n += 1
                details.append(f"CP {mkt} | {co} | {rec['person']} | {rec['role']}")
        save_atomic(wb, path)
        wb.close()
        print("saved", mkt, path.name)
    print("website_updates", web_n)
    print("contact_fills", cp_n)
    for d in details:
        print(d)


if __name__ == "__main__":
    main()
