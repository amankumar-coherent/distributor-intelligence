"""Remove remaining wrong-market / non-channel rows after role verify."""
from datetime import datetime
from pathlib import Path
import json
import shutil

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
REMOVE_COMPANIES = {
    "packaging outlet dist",
    "packaging outlet",
}

backup = FP.with_name(FP.stem + f"_pre_outlet_cleanup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
shutil.copy2(FP, backup)

wb = load_workbook(FP)
ws = wb.active
header = [c.value for c in ws[1]]
rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
kept = []
removed = []
for r in rows:
    name = str(r[0] or "").strip()
    if name.lower() in REMOVE_COMPANIES:
        removed.append({"company": name, "website": r[1], "reason": "wrong_market_equipment_parts_not_flex_pack_channel"})
        continue
    kept.append(r)

ws.delete_rows(1, ws.max_row)
ws.append(header)
for r in kept:
    ws.append(r)
wb.save(FP)

audit = {
    "ts": datetime.now().isoformat(),
    "backup": str(backup),
    "before": len(rows),
    "after": len(kept),
    "removed": removed,
    "file": str(FP),
}
out = FP.parent / "channel_role_verify_post_cleanup_audit.json"
out.write_text(json.dumps(audit, indent=2), encoding="utf-8")
print("before", len(rows), "after", len(kept), "removed", removed)
print("backup", backup)
