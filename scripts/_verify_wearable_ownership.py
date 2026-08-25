# -*- coding: utf-8 -*-
"""
Verify / fill Ownership for wearable FINAL sheet.
Only high/medium confidence from website text + DeepSeek, plus hand-verified overrides.
Never invent owners. Leave blank if unverified.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
AUDIT = ROOT / "_wearable_ownership_audit.json"
CKPT = ROOT / "_wearable_ownership_ckpt.json"

load_dotenv(r"D:\vendor-intel\.env")

# Hand-verified (web/SEC/Companies House/IR pages). Prefer precise parent/listing.
HAND: dict[str, str] = {
    # Corrections of stale/wrong values
    "AAH Pharmaceuticals": "Subsidiary of Admenta UK Limited / Hallo Healthcare Group (Aurelius-backed)",
    "Alliance Healthcare": "Subsidiary of Cencora",
    "Bunzl": "Public (LSE: BNZL)",
    # Blanks / known publics & subsidiaries
    "AddLife AB": "Public (Nasdaq Stockholm: ALIF B)",
    "Apotek Hjärtat": "Subsidiary of ICA Gruppen",
    "Anixter": "Subsidiary of Wesco International",
    "Wesco International": "Public (NYSE: WCC)",
    "Arrow Electronics": "Public (NYSE: ARW)",
    "CDW Corporation": "Public (NASDAQ: CDW)",
    "Coloplast": "Public (Nasdaq Copenhagen: COLO B)",
    "eBay Inc.": "Public (NASDAQ: EBAY)",
    "Fastenal": "Public (NASDAQ: FAST)",
    "Grainger": "Public (NYSE: GWW)",
    "Grifols": "Public (BME: GRF / NASDAQ: GRFS)",
    "HD Supply": "Subsidiary of Home Depot",
    "MSC Industrial Supply": "Public (NYSE: MSM)",
    "Mouser Electronics": "Subsidiary of TTI, Inc.",
    "Orion Pharma": "Public (Nasdaq Helsinki: ORNBV) — Orion Corporation",
    "Phoenix Group": "Private (PHOENIX group)",
    "Rakuten Group": "Public (TSE: 4755)",
    "Shoprite Holdings": "Public (JSE: SHP)",
    "Zoro": "Subsidiary of W.W. Grainger",
    "Moore Medical": "Subsidiary of McKesson Medical-Surgical",
    "CuraScript SD": "Subsidiary of Cencora",
    "BOC Healthcare": "Subsidiary of Linde plc (BOC)",
    "Bic Camera Inc.": "Public (TSE: 3048)",
    "As One Corporation": "Public (TSE: 7476)",
    "Almedio Inc.": "Public (TSE: 7859)",
    "CJ ENM": "Public (KRX: 035760)",
    "Hakuto Co., Ltd.": "Public (TSE: 7433)",
    "Hyundai Home Shopping": "Public (KRX: 057050)",
    "Japan Medical Supply Co., Ltd.": "Public (TSE: 7702) — JMS Co., Ltd.",
    "Kojima Co., Ltd.": "Public (TSE: 7513)",
    "LGU+": "Public (KRX: 032640) — LG Uplus",
    "MatsukiyoCocokara & Co.": "Public (TSE: 3088)",
    "Murakami Corporation": "Public (TSE: 7292)",
    "Sundrug Co., Ltd.": "Public (TSE: 9989)",
    "Tsuruha Holdings": "Public (TSE: 3391)",
    "Pague Menos": "Public (B3: PGMN3)",
    "Jean Coutu": "Subsidiary of Metro Inc.",
    "Rexall Pharmacy": "Subsidiary of McKesson Canada",
    "MediSystem Pharmacy": "Subsidiary of McKesson Canada",
    "Insight Enterprises": "Public (NASDAQ: NSIT)",
    "Newegg Business": "Public (NASDAQ: NEGG) — Newegg Commerce",
    "Global Industrial": "Public (NYSE: GIC)",
    "Movianto": "Subsidiary of Yusen Logistics (NYK Group)",
    "Terumo Australia": "Subsidiary of Terumo Corporation (TSE: 4543)",
    "Welch Allyn Australia": "Subsidiary of Baxter International (Hillrom / Welch Allyn)",
    "Kenya Medical Supplies Authority (KEMSA)": "Government / state corporation (Kenya)",
    "Extrafarma": "Subsidiary of Pague Menos",
    "Meditrade GmbH": "Private (Meditrade Group)",
    "Pharma Dynamics": "Subsidiary of Lupin Limited",
    "Reliance Digital": "Subsidiary of Reliance Retail (Reliance Industries)",
    "Imperial Health Sciences": "Subsidiary of Imperial (DP World)",
    "HD Supply": "Subsidiary of The Home Depot",
}

UA = {"User-Agent": "Mozilla/5.0 (compatible; VendorIntelOwn/1.1)"}


def strip_html(html: str) -> str:
    html = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", html)
    html = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", html)
    html = re.sub(r"(?is)<[^>]+>", " ", html)
    return re.sub(r"\s+", " ", html).strip()


def fetch(url: str) -> str:
    try:
        with httpx.Client(follow_redirects=True, timeout=18.0, headers=UA) as c:
            r = c.get(url)
            if r.status_code >= 400:
                return ""
            return strip_html(r.text)[:9000]
    except Exception:
        return ""


def candidate_urls(website: str) -> list[str]:
    base = (website or "").strip().rstrip("/")
    if not base:
        return []
    if not base.startswith("http"):
        base = "https://" + base
    paths = [
        "",
        "/about",
        "/about-us",
        "/company",
        "/investors",
        "/en/about",
        "/en/company",
        "/corporate",
        "/who-we-are",
    ]
    out = []
    for p in paths:
        u = base if p == "" else urljoin(base + "/", p.lstrip("/"))
        if u not in out:
            out.append(u)
    return out[:6]


def deepseek_ownership(company: str, website: str, hq: str, text: str) -> dict:
    key = (os.getenv("OPENAI_API_KEY") or os.getenv("DEEPSEEK_API_KEY") or "").strip()
    if not key:
        return {"ownership": "", "confidence": "none", "reason": "no_key"}
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    system = (
        "Extract company ownership from evidence ONLY. Never invent.\n"
        "Return JSON: {ownership:string, confidence:high|medium|low|none, reason:string}.\n"
        "ownership formats preferred:\n"
        "- Public (EXCHANGE: TICKER)\n"
        "- Private\n"
        "- Subsidiary of PARENT\n"
        "- Cooperative\n"
        "- Government / state-owned\n"
        "- Private (PE firm name) when evidenced\n"
        "If evidence only says privately held without parent → 'Private'.\n"
        "If unclear / no evidence → ownership empty string and confidence none/low.\n"
        "Do not guess PE sponsors or parents."
    )
    user = f"Company: {company}\nWebsite: {website}\nHQ: {hq}\n\nEvidence:\n{text[:8500]}"
    try:
        with httpx.Client(timeout=60.0) as c:
            r = c.post(
                f"{base}/chat/completions",
                headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
                json={
                    "model": model,
                    "temperature": 0,
                    "messages": [
                        {"role": "system", "content": system},
                        {"user": "user", "content": user},
                    ],
                    "response_format": {"type": "json_object"},
                },
            )
            # fix messages key
            raise RuntimeError("placeholder")
    except Exception:
        pass
    try:
        with httpx.Client(timeout=60.0) as c:
            r = c.post(
                f"{base}/chat/completions",
                headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
                json={
                    "model": model,
                    "temperature": 0,
                    "messages": [
                        {"role": "system", "content": system},
                        {"role": "user", "content": user},
                    ],
                    "response_format": {"type": "json_object"},
                },
            )
            r.raise_for_status()
            return json.loads(r.json()["choices"][0]["message"]["content"])
    except Exception as e:
        return {"ownership": "", "confidence": "none", "reason": str(e)}


def normalize_own(s: str) -> str:
    t = re.sub(r"\s+", " ", (s or "").strip())
    if not t:
        return ""
    # Drop obviously garbage
    low = t.lower()
    if low in {"unknown", "n/a", "na", "none", "not found", "unclear"}:
        return ""
    if "verify" in low or "placeholder" in low:
        return ""
    # Cap length
    if len(t) > 140:
        t = t[:137] + "..."
    return t


def apply_format(ws) -> None:
    headers = [c.value for c in ws[1]]
    widths = {
        "Company": 34,
        "Website": 36,
        "Founded": 10,
        "Headquarters": 28,
        "Continent / Geography": 52,
        "Operational Presence": 48,
        "Ownership": 42,
        "Summary": 48,
    }
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(str(h or ""), 16)
    ws.row_dimensions[1].height = 32
    for i in range(1, len(headers) + 1):
        cell = ws.cell(1, i)
        cell.fill = header_fill
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
    for r in range(2, ws.max_row + 1):
        for i in range(1, len(headers) + 1):
            cell = ws.cell(r, i)
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            cell.border = thin
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = ROOT / f"global_wearable_medical_devices_market_global_FINAL_BEFORE_OWNERSHIP_{stamp}.xlsx"
    shutil.copy2(XLSX, bak)

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers) if h}

    ckpt = {"done": {}}
    if CKPT.exists():
        try:
            ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
        except Exception:
            pass

    hand_applied = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co not in HAND:
            continue
        own = normalize_own(HAND[co])
        if not own:
            continue
        ws.cell(r, ic["Ownership"]).value = own
        hand_applied.append({"company": co, "ownership": own})

    # Targets: blank ownership OR hand-corrected already done; DeepSeek only blanks not in HAND
    targets = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co in HAND:
            continue
        cur = str(ws.cell(r, ic["Ownership"]).value or "").strip()
        if not cur:
            targets.append(r)

    MAX = int(os.getenv("OWN_EXPAND_MAX", "80"))
    print("hand_applied", len(hand_applied), "blank_targets", len(targets), "max", MAX, flush=True)

    filled_ds = []
    processed = 0
    try:
        for r in targets:
            if processed >= MAX:
                break
            co = str(ws.cell(r, ic["Company"]).value or "").strip()
            if co in ckpt.get("done", {}):
                print("skip ckpt", co, flush=True)
                continue
            processed += 1
            website = str(ws.cell(r, ic["Website"]).value or "").strip()
            hq = str(ws.cell(r, ic["Headquarters"]).value or "").strip()
            snippets = []
            for u in candidate_urls(website):
                t = fetch(u)
                if len(t) > 200:
                    snippets.append(f"URL:{u}\n{t[:3200]}")
                time.sleep(0.15)
            text = "\n\n---\n\n".join(snippets) if snippets else f"(no site text) HQ={hq}"
            ds = deepseek_ownership(co, website, hq, text)
            conf = str(ds.get("confidence") or "").lower()
            own = normalize_own(str(ds.get("ownership") or ""))
            applied = False
            if conf in {"high", "medium"} and own:
                ws.cell(r, ic["Ownership"]).value = own
                applied = True
                filled_ds.append({"company": co, "ownership": own, "conf": conf, "deepseek": ds})
                print("FILLED", co, "=>", own, flush=True)
            else:
                print("leave_blank", co, conf, own or "(empty)", flush=True)
            ckpt.setdefault("done", {})[co] = {
                "applied": applied,
                "ownership": own,
                "deepseek": ds,
            }
            if processed % 5 == 0:
                CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")
                wb.save(XLSX)
                print("checkpoint", processed, flush=True)
    finally:
        apply_format(ws)
        wb.save(XLSX)
        CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")

    blank = filled = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, ic["Ownership"]).value
        if v is None or str(v).strip() == "":
            blank += 1
        else:
            filled += 1

    audit = {
        "backup": bak.name,
        "hand_applied": hand_applied,
        "deepseek_filled": filled_ds,
        "counts_after": {"filled": filled, "blank": blank},
        "processed_new": processed,
        "remaining_blank_targets_est": max(0, len(targets) - processed),
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print("DONE filled", filled, "blank", blank, "ds_new", len(filled_ds))


if __name__ == "__main__":
    main()
