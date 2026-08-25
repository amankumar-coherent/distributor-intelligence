"""Analyze FINAL sheet for website quality and multi-location duplicates."""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)

wb = load_workbook(FP, data_only=True)
ws = wb.active
header = [c.value for c in ws[1]]
print("headers", header)
rows = list(ws.iter_rows(min_row=2, values_only=True))
print("count", len(rows))

ends = Counter()
for r in rows:
    name = (r[0] or "").strip()
    low = name.lower()
    for suf in [
        " inc",
        " inc.",
        " ltd",
        " ltd.",
        " llc",
        " co.",
        " corp",
        " corporation",
        " company",
        " gmbh",
        " sa",
        " pty",
        " plc",
    ]:
        if low.endswith(suf) or low.endswith(suf + "."):
            ends[suf] += 1
print("suffix_hits", dict(ends))

empty_web = []
domains = []
for r in rows:
    name, web = (r[0] or "").strip(), (r[1] or "").strip()
    if not web:
        empty_web.append(name)
        continue
    try:
        d = urlparse(web if "://" in web else "https://" + web).netloc.lower().replace("www.", "")
    except Exception:
        d = web
    domains.append((name, web, d))

print("empty_website", len(empty_web))
for n in empty_web[:30]:
    print("  EMPTY", n)

dom_map = defaultdict(list)
for name, web, d in domains:
    dom_map[d].append(name)
dups = {k: v for k, v in dom_map.items() if len(v) > 1}
print("duplicate_domains", len(dups))
for k, v in sorted(dups.items(), key=lambda x: -len(x[1])):
    print(f"  {k}: {len(v)} -> {v}")


def base(n: str) -> str:
    n = re.sub(r"\s+", " ", (n or "").strip())
    n = re.sub(r"\s*\([^)]*\)\s*$", "", n)
    for tok in [
        " USA",
        " US",
        " UK",
        " AU",
        " CA",
        " Canada",
        " Europe",
        " MX",
        " ZA",
        " France",
        " Germany",
        " Italy",
        " Spain",
        " Netherlands",
        " Belgium",
        " Switzerland",
        " Sweden",
        " Norway",
        " Denmark",
        " Ireland",
        " Portugal",
        " Austria",
        " Hungary",
        " Romania",
        " Slovakia",
        " Czech",
        " Luxembourg",
        " Poland",
        " Corporate",
        " Merchant",
        " Group Corporate",
    ]:
        if n.endswith(tok):
            n = n[: -len(tok)].strip()
    n = re.sub(
        r"\s+(Inc\.?|Ltd\.?|LLC|Corp\.?|Corporation|Company|GmbH|Pty\.?|PLC|SA|BV|AG|Limited)$",
        "",
        n,
        flags=re.I,
    )
    return n.lower().strip()


base_map = defaultdict(list)
for r in rows:
    name = (r[0] or "").strip()
    base_map[base(name)].append((name, r[1]))
multi = {k: v for k, v in base_map.items() if len(v) > 1}
print("multi_location_or_similar_base", len(multi))
for k, v in sorted(multi.items(), key=lambda x: -len(x[1]))[:80]:
    print(f"  BASE[{k}] ({len(v)})")
    for name, web in v:
        print(f"    - {name} | {web}")

# Print all company|website for later DNS probe
out = FP.parent / "all_companies_websites_tmp.txt"
with out.open("w", encoding="utf-8") as f:
    for r in rows:
        f.write(f"{r[0]}\t{r[1]}\n")
print("wrote", out)
wb.close()
