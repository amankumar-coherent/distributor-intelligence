from openpyxl import load_workbook
from pathlib import Path

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
out = FP.parent / "_founded_empty.txt"
wb = load_workbook(FP, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
empty = []
for r in rows[1:]:
    f = r[2]
    if f is None or str(f).strip() in ("", "None", "N/A", "Not publicly disclosed", "-"):
        empty.append(f"{r[0]}\t{r[1]}")
out.write_text("\n".join(empty), encoding="utf-8")
print(len(empty), "->", out)
wb.close()
