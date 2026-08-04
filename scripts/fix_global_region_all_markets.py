"""Fix misclassified 'Global' region rows across smartwatch / firewall / green chemical FINALs.

Root cause: Regions Served = 'Global' while HQ city/state lacked country keywords.
Reassign by verified HQ and rebuild region sheets (Asia Pacific first).
"""
from __future__ import annotations

import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Exact company name (normalized) → region + HQ/served fills
FIXES: dict[str, dict[str, str]] = {
    # Smartwatch
    "npuone": {
        "Region Segment": "North America",
        "Headquarters": "Alpharetta, Georgia, United States",
        "Regions Served": "United States; Global channel",
    },
    "rptechindia": {
        "Region Segment": "Asia Pacific",
        "Headquarters": "Gandhinagar, Gujarat, India",
        "Regions Served": "India; Asia Pacific",
    },
    "usmphones": {
        "Region Segment": "North America",
        "Headquarters": "New Brunswick, New Jersey, United States",
        "Regions Served": "United States; Global channel",
    },
    "arrow electronics": {
        "Region Segment": "North America",
        "Headquarters": "Centennial / Englewood, Colorado, United States",
        "Regions Served": "Global (HQ North America)",
    },
    "redington": {
        "Region Segment": "Asia Pacific",
        "Headquarters": "Chennai, Tamil Nadu, India",
        "Regions Served": "India; Middle East; Africa; Global channel",
    },
    # Firewall
    "cranel": {
        "Region Segment": "North America",
        "Headquarters": "Columbus, Ohio, United States",
        "Regions Served": "United States; Global channel",
    },
    "ebttikar technology company": {
        "Region Segment": "Middle East",
        "Headquarters": "Riyadh, Saudi Arabia",
        "Regions Served": "Saudi Arabia; Middle East; Global channel",
    },
    "eplus technology": {
        "Region Segment": "North America",
        "Headquarters": "Herndon, Virginia, United States",
        "Regions Served": "United States; Global channel",
        "Cities / Regions": "Herndon, VA; United States",
    },
    "nuvias uc": {
        "Region Segment": "Europe",
        "Headquarters": "Basingstoke, Hampshire, United Kingdom",
        "Regions Served": "United Kingdom; Europe; Global channel",
    },
    "red river technology": {
        "Region Segment": "North America",
        "Headquarters": "Chantilly, Virginia, United States",
        "Regions Served": "United States; Global channel",
    },
    "softcat": {
        "Region Segment": "Europe",
        "Headquarters": "Marlow, Buckinghamshire, United Kingdom",
        "Regions Served": "United Kingdom; Europe; Global channel",
    },
    # Green chemical
    "green chemical store": {
        "Region Segment": "North America",
        "Headquarters": "Dallas, Texas, United States",
        "Regions Served": "United States; Global e-commerce",
    },
    "palmer holland": {
        "Region Segment": "North America",
        "Headquarters": "Westlake, Ohio, United States",
        "Regions Served": "United States; North America",
    },
    "ricardo molina": {
        "Region Segment": "Europe",
        "Headquarters": "Sant Cugat del Vallès / Barcelona, Spain",
        "Regions Served": "Spain; Europe (Caldic Iberica)",
        "Cities / Regions": "Barcelona; Spain",
    },
    "tilley distribution": {
        "Region Segment": "North America",
        "Headquarters": "Baltimore, Maryland, United States",
        "Regions Served": "United States; North America",
    },
    "velox gmbh": {
        "Region Segment": "Europe",
        "Headquarters": "Hamburg, Germany",
        "Regions Served": "Europe; Global specialty chemicals (IMCD)",
        "Cities / Regions": "Hamburg, Germany",
    },
}

PATHS = [
    (
        Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL.xlsx"),
        "Smartwatch Companies",
        "Global Smartwatch Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL_320plus.xlsx"
        ),
        "Smartwatch Companies",
        "Global Smartwatch Market — Region-wise segmentation",
    ),
    (
        Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL.xlsx"),
        "Firewall Companies",
        "Global Firewall Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL_320plus.xlsx"
        ),
        "Firewall Companies",
        "Global Firewall Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL.xlsx"
        ),
        "Green Chemical Companies",
        "Global Green Chemical Market — Region-wise segmentation",
    ),
    (
        Path(
            r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL_320plus.xlsx"
        ),
        "Green Chemical Companies",
        "Global Green Chemical Market — Region-wise segmentation",
    ),
]

REGION_ORDER = [
    "Asia Pacific",
    "Europe",
    "North America",
    "Middle East",
    "Latin America",
    "Africa",
    "Eurasia / CIS",
    "Global",
    "Other - Unclear",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
CELL_FONT = Font(name="Calibri", size=9)
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP = Alignment(wrap_text=True, vertical="top")

NOTES = {
    "Asia Pacific": "India, China, Japan, SEA, ANZ, Korea",
    "Europe": "UK + EU / wider Europe",
    "North America": "USA, Canada, Mexico",
    "Middle East": "GCC, UAE, Saudi, Turkey, Israel",
    "Latin America": "Brazil, Chile, Colombia, Argentina, Peru, Uruguay",
    "Africa": "South Africa, Nigeria, Egypt, Kenya, etc.",
    "Eurasia / CIS": "Russia / CIS channel partners",
    "Global": "Truly multi-HQ / no single primary region",
    "Other - Unclear": "Insufficient HQ/region text to classify",
}


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9&]+", " ", (s or "").lower()).strip()


def _region_rank(region: str) -> int:
    try:
        return REGION_ORDER.index(region)
    except ValueError:
        return len(REGION_ORDER)


def _style_header(ws, headers: list[str]) -> None:
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = 18 if h != "Summary" else 36
    ws.freeze_panes = "A2"


def _write_rows(ws, headers: list[str], rows: list[dict[str, str]]) -> None:
    _style_header(ws, headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
        rr = ws.max_row
        for i in range(1, len(headers) + 1):
            cell = ws.cell(rr, i)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = WRAP
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last}{ws.max_row}"


def fix_file(path: Path, sheet_name: str, summary_title: str) -> dict:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_pre_fix_global_region_{stamp}{path.suffix}")
    shutil.copy2(path, bak)

    wb = openpyxl.load_workbook(path)
    ws = wb[sheet_name]
    hr = 3
    headers = [str(c.value or "") for c in ws[hr]]
    if "Company" not in headers or "Region Segment" not in headers:
        raise RuntimeError(f"bad header in {path}")

    co_col = headers.index("Company") + 1
    seg_col = headers.index("Region Segment") + 1
    col_idx = {h: i + 1 for i, h in enumerate(headers)}

    fixed = []
    for r in range(hr + 1, ws.max_row + 1):
        company = str(ws.cell(r, co_col).value or "").strip()
        if not company:
            continue
        seg = str(ws.cell(r, seg_col).value or "").strip()
        key = _norm(company)
        if seg != "Global" and key not in FIXES:
            continue
        patch = FIXES.get(key)
        if not patch:
            continue
        if seg == "Global" or key in FIXES:
            old_seg = seg
            for field, val in patch.items():
                if field in col_idx:
                    ws.cell(r, col_idx[field]).value = val
            fixed.append((company, old_seg, patch["Region Segment"]))

    # Load all rows, sort APAC-first, rewrite data
    companies: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        if not row or not row[co_col - 1]:
            continue
        d = {
            headers[i]: str(row[i] if i < len(row) and row[i] is not None else "")
            for i in range(len(headers))
        }
        companies.append(d)

    companies.sort(
        key=lambda d: (_region_rank(d.get("Region Segment", "")), _norm(d.get("Company", "")))
    )

    if ws.max_row > hr:
        ws.delete_rows(hr + 1, ws.max_row - hr)
    for d in companies:
        ws.append([d.get(h, "") for h in headers])
        rr = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(rr, c)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = WRAP

    # Rebuild region sheets
    for name in list(wb.sheetnames):
        if name.startswith("Region - ") or name in {"Region Summary", "By Region"}:
            del wb[name]

    counts = Counter(d.get("Region Segment", "") for d in companies)
    total = len(companies)
    sum_ws = wb.create_sheet("Region Summary", 1)
    sum_ws["A1"] = summary_title
    sum_ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    sum_ws.merge_cells("A1:D1")
    for i, h in enumerate(["Region Segment", "Company Count", "% of Total", "Notes"], 1):
        cell = sum_ws.cell(3, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    r = 4
    for region in REGION_ORDER:
        n = counts.get(region, 0)
        if n == 0:
            continue
        sum_ws.cell(r, 1, region).border = THIN
        sum_ws.cell(r, 2, n).border = THIN
        sum_ws.cell(r, 3, f"{100.0 * n / total:.1f}%").border = THIN
        sum_ws.cell(r, 4, NOTES.get(region, "")).border = THIN
        r += 1
    sum_ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    sum_ws.cell(r, 2, total).font = Font(bold=True)
    sum_ws.column_dimensions["A"].width = 22
    sum_ws.column_dimensions["B"].width = 16
    sum_ws.column_dimensions["C"].width = 12
    sum_ws.column_dimensions["D"].width = 55

    by_region: dict[str, list[dict[str, str]]] = defaultdict(list)
    for d in companies:
        by_region[d.get("Region Segment", "Other - Unclear")].append(d)

    for region in REGION_ORDER:
        rows = by_region.get(region) or []
        if not rows:
            continue
        title = f"Region - {region.replace('/', '-')}"[:31]
        rws = wb.create_sheet(title)
        _write_rows(rws, headers, rows)

    desired = ["Region Summary"] + [
        f"Region - {r.replace('/', '-')}"[:31] for r in REGION_ORDER if by_region.get(r)
    ]
    for offset, name in enumerate(desired, start=1):
        if name in wb.sheetnames:
            current = wb.sheetnames.index(name)
            if current != offset:
                wb.move_sheet(name, offset=offset - current)

    wb.save(path)
    return {
        "file": path.name,
        "backup": bak.name,
        "fixed": fixed,
        "counts": {k: counts[k] for k in REGION_ORDER if counts.get(k)},
        "global_left": counts.get("Global", 0),
    }


def main() -> None:
    for path, sheet, title in PATHS:
        if not path.exists():
            print("missing", path)
            continue
        rep = fix_file(path, sheet, title)
        print("FILE", rep["file"])
        print("  backup", rep["backup"])
        print("  fixed", len(rep["fixed"]))
        for co, old, new in rep["fixed"]:
            print(f"    {co}: {old} -> {new}")
        print("  Global left:", rep["global_left"])
        print("  counts:", " | ".join(f"{k}={v}" for k, v in rep["counts"].items()))
        print()


if __name__ == "__main__":
    main()
