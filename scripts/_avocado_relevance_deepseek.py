#!/usr/bin/env python3
"""DeepSeek relevance check: avocado-oil channel / distributor / partner vs off-market."""
from __future__ import annotations

import json
import os
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

XLSX = ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_OFFICE_SLIM.xlsx"
CACHE = ROOT / "output/web_expand/global_avocado_oil_market_global/_avocado_relevance_ds_cache.json"
AUDIT = ROOT / "output/web_expand/global_avocado_oil_market_global/_avocado_relevance_ds_audit.json"


def client() -> OpenAI:
    return OpenAI(
        api_key=os.getenv("OPENAI_API_KEY"),
        base_url=os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1",
    )


def load_cache() -> dict:
    if CACHE.exists():
        try:
            return json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def classify(cli: OpenAI, model: str, facts: dict) -> dict:
    prompt = f"""Classify this company for a Global Avocado Oil Market CHANNEL landscape (distributors, wholesalers, cash & carry, foodservice redistributors, specialty importers, grocery wholesale, edible-oil traders, produce/avocado supply chain partners that can carry avocado oil or edible oils).

Return STRICT JSON only:
{{
  "relevant": true/false,
  "role": "distributor|wholesaler|retailer_wholesale|foodservice|ingredient_trader|produce_channel|cash_carry|other_channel|oem_brand|off_market|unclear",
  "avocado_oil_link": "direct|adjacent_edible_oils|food_channel_likely|weak|none",
  "confidence": 0.0-1.0,
  "reason": "one short sentence"
}}

KEEP (relevant=true) if they are a channel partner that distributes food / edible oils / avocado products / broadline FS / grocery wholesale / cash & carry / produce wholesale that can include avocado oil SKUs.
DROP (relevant=false) if primarily: cosmetics-only with no food oils, pure restaurant franchise operator with no wholesale, dead/wrong company, non-food chemicals only, real-estate, directories, OEM consumer brand with no channel role, or clearly unrelated.

Facts:
{json.dumps(facts, ensure_ascii=False)}
"""
    resp = cli.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": "You are a strict market-map QA classifier. Output JSON only. Never invent companies.",
            },
            {"role": "user", "content": prompt},
        ],
        temperature=0.1,
        max_tokens=220,
    )
    text = (resp.choices[0].message.content or "").strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.startswith("json"):
            text = text[4:].strip()
    try:
        data = json.loads(text)
    except Exception:
        # try extract braces
        start, end = text.find("{"), text.rfind("}")
        data = json.loads(text[start : end + 1]) if start >= 0 and end > start else {}
    return {
        "relevant": bool(data.get("relevant", False)),
        "role": str(data.get("role") or "unclear"),
        "avocado_oil_link": str(data.get("avocado_oil_link") or "unclear"),
        "confidence": float(data.get("confidence") or 0),
        "reason": str(data.get("reason") or text[:180]),
    }


def main() -> None:
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    cli = client()
    cache = load_cache()

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    ic = {str(h): i for i, h in enumerate(hdr) if h}

    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        co = str(row[ic["Company"]] or "").strip()
        if not co:
            continue
        facts = {
            "company": co,
            "website": str(row[ic["Website"]] or "").strip(),
            "core": str(row[ic["Core Categories"]] or "").strip(),
            "specialty": str(row[ic["Specialty Focus"]] or "").strip(),
            "brands": str(row[ic["Key Brands Represented"]] or "").strip(),
            "summary": str(row[ic["Summary"]] or "").strip(),
            "hq": str(row[ic["Headquarters"]] or "").strip(),
            "ownership": str(row[ic["Ownership"]] or "").strip(),
        }
        rows.append(facts)
    wb.close()

    keep, drop, weak = [], [], []
    for i, facts in enumerate(rows, 1):
        co = facts["company"]
        fp = json.dumps(facts, sort_keys=True, ensure_ascii=False)
        cached = cache.get(co)
        if isinstance(cached, dict) and cached.get("fp") == fp and "relevant" in cached:
            result = {k: cached[k] for k in ("relevant", "role", "avocado_oil_link", "confidence", "reason")}
            src = "cache"
        else:
            try:
                result = classify(cli, model, facts)
                src = "deepseek"
                cache[co] = {"fp": fp, **result, "ts": datetime.now(timezone.utc).isoformat()}
                if i % 15 == 0:
                    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
                time.sleep(0.3)
            except Exception as e:
                result = {
                    "relevant": True,
                    "role": "unclear",
                    "avocado_oil_link": "unclear",
                    "confidence": 0,
                    "reason": f"classify_error: {e}",
                }
                src = "error"
                time.sleep(1.0)

        item = {"company": co, **result}
        link = result.get("avocado_oil_link", "")
        if not result["relevant"] or link == "none":
            drop.append(item)
        elif link == "weak" or result.get("confidence", 1) < 0.45:
            weak.append(item)
        else:
            keep.append(item)
        print(
            f"[{i}/{len(rows)}] {src} relevant={result['relevant']} link={result.get('avocado_oil_link')} | {co}",
            flush=True,
        )

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
    audit = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "model": model,
        "total": len(rows),
        "keep": len(keep),
        "drop": len(drop),
        "weak": len(weak),
        "drop_list": drop,
        "weak_list": weak,
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"total": audit["total"], "keep": audit["keep"], "drop": audit["drop"], "weak": audit["weak"]}, indent=2))
    print("\nDROP:")
    for x in drop:
        print(f"- {x['company']} | {x['role']} | {x['avocado_oil_link']} | {x['reason']}")
    print("\nWEAK:")
    for x in weak:
        print(f"- {x['company']} | {x['role']} | {x['avocado_oil_link']} | {x['reason']}")


if __name__ == "__main__":
    main()
