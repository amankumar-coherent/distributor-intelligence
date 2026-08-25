# -*- coding: utf-8 -*-
"""Retry fill remaining blank Ownership rows — hand + scrape + DeepSeek."""
from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import httpx
import openpyxl
from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")
OUT = ROOT / "output/web_expand/global_wearable_medical_devices_market_global"
XLSX = OUT / "global_wearable_medical_devices_market_global_FINAL.xlsx"

HAND = {
    "Himedica": "Subsidiary of Crawford Manufacturing Limited",
    "MediSupplies Ltd": "Subsidiary of Poole Bay Holdings Limited",
    "Gulf Medical Company": "Private (Naghi & Sons / Cigalah Group)",
    "Health Care Australia": "Private (Crescent Capital Partners)",
    "Spire Health": "Private (venture-backed; Gilde Healthcare)",
    "Medex Supply": "Private",
    "Empsephar": "Private (Spanish SL)",
    "Shenzhen XFT Medical": "Private",
    "Australian Healthcare Supplies": "Private",
    "MediDirect": "Private",
    "Brasil Medical Distribuidora": "Private",
    "Pro Medical Distribuidora": "Private",
    "Tecnomed Brasil": "Private",
}

EXTRA = {
    "Al-Etihad Medical Supplies": (
        "UAE medical supplies distributor (etihadmedical.com). No public listing found."
    ),
    "Al-Hayat Medical Company": (
        "Saudi medical company (alhayatmedical.com). No public listing found."
    ),
    "Al Noor Medical": (
        "Abu Dhabi medical distributor site alnoormedical.com — NOT confirmed as "
        "Mediclinic Al Noor Hospitals; do not assign Mediclinic without proof on this site."
    ),
    "Africure Pharmaceuticals": (
        "Row website africure.com may not match Africure Pharmaceuticals Ltd "
        "(SEM-listed africurepharma.com). Only assign Public if evidence clearly same entity."
    ),
    "MediMark Europe GmbH": "European medical distributor based Grenoble. GmbH form.",
    "ProCare Pharmacy": (
        "US pharmacy brand site. Many ProCare entities exist; do not invent parent."
    ),
}

UA = {"User-Agent": "Mozilla/5.0 (compatible; VendorIntelOwn/1.2)"}


def strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def scrape(website: str) -> str:
    url = (website or "").strip()
    if not url:
        return ""
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    texts: list[str] = []
    try:
        with httpx.Client(
            headers=UA,
            verify=False,
            timeout=httpx.Timeout(12.0, connect=4.0),
            follow_redirects=True,
        ) as c:
            for path in ["", "/about", "/about-us", "/company", "/en/about"]:
                try:
                    u = url if not path else urljoin(url.rstrip("/") + "/", path.lstrip("/"))
                    r = c.get(u)
                    if r.status_code >= 400:
                        continue
                    t = strip_html(r.text or "")
                    if len(t) > 150:
                        texts.append(f"URL:{u}\n{t[:2500]}")
                    if len(texts) >= 2:
                        break
                except Exception:
                    continue
    except Exception:
        return ""
    return "\n\n---\n\n".join(texts)[:7000]


def parse_json(text: str) -> dict:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    try:
        return json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


def ds_own(client: OpenAI, model: str, company: str, website: str, hq: str, evidence: str) -> dict:
    system = (
        "Extract ownership from EVIDENCE ONLY. Never invent.\n"
        "Return JSON: {ownership:string, confidence:high|medium|low|none, reason:string}.\n"
        "Formats: Public (EXCHANGE:TICKER) | Private | Private (detail) | Subsidiary of PARENT.\n"
        "If insufficient: ownership empty string and confidence none."
    )
    user = f"Company: {company}\nWebsite: {website}\nHQ: {hq}\n\nEVIDENCE:\n{evidence[:6500]}"
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=220,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        return parse_json(resp.choices[0].message.content or "{}")
    except Exception as e:  # noqa: BLE001
        return {"ownership": "", "confidence": "none", "reason": str(e)[:120]}


def main() -> None:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    client = OpenAI(api_key=key, base_url=base, timeout=60.0) if key else None
    print(f"DeepSeek={'yes' if client else 'no'} model={model}", flush=True)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_OWNERSHIP_RETRY_{stamp}.xlsx")
    wb0 = openpyxl.load_workbook(XLSX)
    wb0.save(bak)
    wb0.close()
    print("backup", bak.name, flush=True)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    gaps = [
        r
        for r in range(2, ws.max_row + 1)
        if not str(ws.cell(r, ic["Ownership"]).value or "").strip()
    ]
    print("gaps", len(gaps), flush=True)

    hand_n = ds_n = left = 0
    audit = []
    for r in gaps:
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        hq = str(ws.cell(r, ic["Headquarters"]).value or "").strip()

        if co in HAND:
            ws.cell(r, ic["Ownership"]).value = HAND[co]
            hand_n += 1
            audit.append({"company": co, "source": "hand", "ownership": HAND[co]})
            print("HAND", co, "=>", HAND[co], flush=True)
            continue

        if not client:
            left += 1
            continue

        evidence = scrape(web)
        if co in EXTRA:
            evidence = (evidence + "\n\nWEB_NOTES:\n" + EXTRA[co])[:7000]
        if not evidence.strip():
            evidence = f"(no site text) HQ={hq}"
            if co in EXTRA:
                evidence += "\n\nWEB_NOTES:\n" + EXTRA[co]

        data = ds_own(client, model, co, web, hq, evidence)
        conf = str(data.get("confidence") or "").lower()
        own = re.sub(r"\s+", " ", str(data.get("ownership") or "").strip())
        if own.lower() in {"", "unknown", "n/a", "none", "unclear", "not found"}:
            own = ""
        if conf in {"high", "medium"} and own:
            ws.cell(r, ic["Ownership"]).value = own[:140]
            ds_n += 1
            audit.append({"company": co, "source": "deepseek", "ownership": own, "conf": conf})
            print("DS", co, "=>", own, conf, flush=True)
        else:
            left += 1
            audit.append(
                {
                    "company": co,
                    "source": "unverified",
                    "conf": conf,
                    "reason": data.get("reason"),
                    "draft": own,
                }
            )
            print("LEAVE", co, conf, own or "(empty)", flush=True)
        time.sleep(0.3)

    wb.save(XLSX)
    wb.close()

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    oi = headers.index("Ownership")
    ci = headers.index("Company")
    filled = blank = 0
    blanks = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = str(row[oi] or "").strip()
        if v:
            filled += 1
        else:
            blank += 1
            blanks.append(row[ci])
    wb.close()

    stats = {
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "hand": hand_n,
        "deepseek": ds_n,
        "left": left,
        "filled": filled,
        "blank": blank,
        "blanks": blanks,
        "backup": str(bak),
    }
    (OUT / "wearable_ownership_retry_summary.json").write_text(
        json.dumps(stats, indent=2), encoding="utf-8"
    )
    (OUT / "wearable_ownership_retry_audit.json").write_text(
        json.dumps(audit, indent=2), encoding="utf-8"
    )
    print(json.dumps(stats, indent=2), flush=True)


if __name__ == "__main__":
    try:
        import urllib3

        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    except Exception:
        pass
    main()
