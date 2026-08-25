"""Fill/upgrade Headquarters for all rows in electrical FINAL Excel to city-level where known."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from vendor_intel.enrichment.known_hq import looks_city_level, resolve_hq

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)


def main() -> None:
    backup = XLSX.with_name(XLSX.stem + "_BEFORE_HQ_FILL.xlsx")
    if not backup.exists():
        backup.write_bytes(XLSX.read_bytes())

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    hi = headers.index("Headquarters") + 1

    changes: list[dict[str, str]] = []
    city_ok = 0
    weak = 0
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        old = str(ws.cell(row_idx, hi).value or "").strip()
        new = resolve_hq(name, old) or old
        if new != old:
            ws.cell(row_idx, hi).value = new
            changes.append(
                {
                    "company": name,
                    "old": old or "(empty)",
                    "new": new,
                    "city_level": str(looks_city_level(new) or ("," in new)),
                }
            )
        if looks_city_level(new) or ("," in str(new)):
            city_ok += 1
        else:
            weak += 1
            if not any(c["company"] == name for c in changes):
                changes.append(
                    {
                        "company": name,
                        "old": old or "(empty)",
                        "new": new,
                        "city_level": "False",
                        "note": "city not verified — left country-level",
                    }
                )

    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(XLSX)

    audit = {
        "when": datetime.now().isoformat(timespec="seconds"),
        "total_rows": city_ok + weak,
        "city_level": city_ok,
        "country_only_or_weak": weak,
        "changed": len([c for c in changes if c.get("old") != c.get("new")]),
        "changes": changes,
    }
    dump = XLSX.parent / "_hq_fill_audit.json"
    dump.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"city_level={city_ok} weak={weak} changed={audit['changed']}")
    print(f"wrote {XLSX}")
    print(f"audit {dump}")
    print(f"backup {backup}")
    for c in changes:
        if c.get("city_level") in (False, "False") or c.get("note"):
            print("WEAK", c["company"], "=>", c["new"])


if __name__ == "__main__":
    main()
