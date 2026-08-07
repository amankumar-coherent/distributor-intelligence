# -*- coding: utf-8 -*-
"""Merge multi-location company families into one unique parent per family.

Uses audit output from audit_multi_location_companies.py
(Alliance Healthcare Portugal/Polska → Alliance Healthcare, etc.).
"""
from __future__ import annotations

import json
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl

from audit_multi_location_companies import (
    company_key,
    nk,
    pick_slash_brand,
    strip_geo_suffix,
)

AUDIT_JSON = Path(
    r"D:\vendor-intel\output\run4\_multi_location_company_audit\multi_location_families.json"
)
OUT_DIR = Path(r"D:\vendor-intel\output\run4\_multi_location_company_audit")

# stem → nice display name
CANONICAL = {
    "exertis": "Exertis",
    "istore": "iStore",
    "bidfood": "Bidfood",
    "makro": "Makro",
    "sysco": "Sysco",
    "transgourmet": "Transgourmet",
    "azelis": "Azelis",
    "barentz": "Barentz",
    "brenntag": "Brenntag",
    "gordon food service": "Gordon Food Service",
    "alliance healthcare": "Alliance Healthcare",
    "nordic biosite": "Nordic BioSite",
    "2bscientific": "2BScientific",
    "all eights": "All Eights",
    "biologix pharma": "Biologix Pharma",
    "farmaenlace": "Farmaenlace",
    "genomax technologies": "Genomax Technologies",
    "zuellig pharma": "Zuellig Pharma",
}

# Prefer these HQ / website values when consolidating
PARENT_DEFAULTS: dict[str, dict] = {
    "Exertis": {
        "Headquarters": "United Kingdom",
        "Website": "https://www.exertis.com",
        "Region Segment": "Europe",
        "Regions Served": "Europe (UK & Ireland)",
    },
    "iStore": {
        "Headquarters": "Multi-country (Australia / South Africa)",
        "Region Segment": "Global",
        "Regions Served": "Australia; South Africa",
    },
    "Bidfood": {
        "Headquarters": "United Kingdom / international",
        "Website": "https://www.bidfood.co.uk",
        "Region Segment": "Global",
        "Regions Served": "Global / multi-region (consolidated branches)",
    },
    "Makro": {
        "Headquarters": "Amsterdam-Duivendrecht, Netherlands",
        "Website": "https://www.makro.nl",
        "Region Segment": "Global",
        "Regions Served": "Europe; Latin America",
    },
    "Sysco": {
        "Headquarters": "Houston, Texas, USA",
        "Website": "https://www.sysco.com",
        "Region Segment": "Global",
        "Regions Served": "Global",
    },
    "Transgourmet": {
        "Headquarters": "Basel, Switzerland",
        "Website": "https://www.transgourmet.ch",
        "Region Segment": "Europe",
        "Regions Served": "Europe",
    },
    "Azelis": {
        "Headquarters": "Antwerp, Belgium",
        "Website": "https://www.azelis.com",
        "Region Segment": "Global",
        "Regions Served": "Global",
    },
    "Barentz": {
        "Headquarters": "Hoofddorp, Netherlands",
        "Website": "https://www.barentz.com",
        "Region Segment": "Global",
        "Regions Served": "Global",
    },
    "Brenntag": {
        "Headquarters": "Essen, Germany",
        "Website": "https://www.brenntag.com",
        "Region Segment": "Global",
        "Regions Served": "Global",
    },
    "Gordon Food Service": {
        "Headquarters": "Wyoming, Michigan, USA",
        "Website": "https://www.gfs.com",
        "Region Segment": "North America",
        "Regions Served": "North America (USA & Canada)",
    },
    "Alliance Healthcare": {
        "Headquarters": "Chessington, United Kingdom",
        "Website": "https://www.alliance-healthcare.co.uk",
        "Region Segment": "Europe",
        "Regions Served": "Europe (multi-country consolidated)",
        "Ownership": "Subsidiary of Cencora / Walgreens Boots Alliance network",
        "Business Type": "Pharmaceutical Wholesaler / Distributor",
    },
    "Nordic BioSite": {
        "Headquarters": "Taby, Sweden",
        "Website": "https://www.nordicbiosite.com",
        "Region Segment": "Europe",
        "Regions Served": "Nordics",
    },
    "2BScientific": {
        "Headquarters": "Upper Heyford, United Kingdom",
        "Website": "https://www.2bscientific.com",
        "Region Segment": "Europe",
        "Regions Served": "Europe (UK & Germany)",
    },
    "All Eights": {
        "Headquarters": "Singapore / Malaysia",
        "Website": "https://www.alleights.com",
        "Region Segment": "Asia Pacific",
        "Regions Served": "Singapore; Malaysia",
    },
    "Biologix Pharma": {
        "Region Segment": "Middle East",
        "Regions Served": "MENA",
    },
    "Farmaenlace": {
        "Region Segment": "Latin America",
        "Regions Served": "Colombia; Ecuador",
    },
    "Genomax Technologies": {
        "Region Segment": "Asia Pacific",
        "Regions Served": "Malaysia; Singapore",
    },
    "Zuellig Pharma": {
        "Region Segment": "Asia Pacific",
        "Regions Served": "Asia Pacific",
    },
}

MARKET_FILES: dict[str, list[Path]] = {
    "Global Firewall Market": [
        Path(
            r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL.xlsx"
        ),
        Path(
            r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL_320plus.xlsx"
        ),
    ],
    "Global Smartwatch Market": [
        Path(
            r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL.xlsx"
        ),
        Path(
            r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL_320plus.xlsx"
        ),
    ],
    "Global Smartphone Market": [
        Path(
            r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market_global_FINAL.xlsx"
        ),
    ],
    "Global Avocado Oil Market": [
        Path(
            r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_global_FINAL.xlsx"
        ),
    ],
    "Global Antibody Drug Conjugate Market": [
        Path(
            r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market_FINAL.xlsx"
        ),
    ],
}


def find_header(ws) -> tuple[int, list]:
    for r in range(1, 6):
        vals = [c.value for c in ws[r]]
        labels = [str(v or "").strip() for v in vals]
        if "Company" in labels:
            return r, labels
    raise RuntimeError("Company header not found")


def score_row(ws, r: int, cols: dict) -> int:
    score = 0
    for f in (
        "Founded",
        "Headquarters",
        "Ownership",
        "Employees",
        "Turnover / Scale",
        "Contact Person",
        "Website",
        "Summary",
        "LinkedIn",
        "Cities / Regions",
    ):
        if f not in cols:
            continue
        v = str(ws.cell(r, cols[f]).value or "").strip().lower()
        if v and v not in {
            "n/a",
            "na",
            "not publicly disclosed",
            "-",
            "unknown",
            "seed_file",
        }:
            score += 1
    return score


def branch_label(company: str, canonical: str) -> str:
    raw = company.strip()
    # Prefer Alliance Healthcare side of slash names
    if "/" in raw:
        parts = re.split(r"\s*/\s*", raw)
        for p in parts:
            if nk(canonical) in nk(p):
                raw = p.strip()
                break
    pat = re.compile(re.escape(canonical), re.I)
    rest = pat.sub("", raw, count=1).strip(" -–—|/()")
    # legal leftovers
    rest = re.sub(
        r"\b(ltd|limited|llc|inc|gmbh|ag|sa|oy|ab|as|spa|srl|pte|sdn|bhd|group)\b",
        "",
        rest,
        flags=re.I,
    )
    rest = re.sub(r"\s+", " ", rest).strip(" -–—|/()")
    if not rest:
        # try last token(s)
        parts = raw.split()
        if len(parts) > len(canonical.split()):
            rest = " ".join(parts[len(canonical.split()) :]).strip(" -–—|/")
    return rest or raw


def family_matches(company: str, stem: str, variant_names: set[str]) -> bool:
    if company in variant_names:
        return True
    key = company_key(company)
    if key in {company_key(v) for v in variant_names}:
        return True
    # GEHE / Alliance Healthcare Deutschland
    if stem == "alliance healthcare" and "alliance healthcare" in key:
        return True
    stem_got = strip_geo_suffix(company)
    return stem_got == stem


def merge_family_in_sheet(
    ws,
    hdr_row: int,
    cols: dict,
    stem: str,
    canonical: str,
    variant_names: list[str],
) -> dict | None:
    variant_set = set(variant_names)
    members: list[tuple[int, str]] = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        company = str(ws.cell(r, cols["Company"]).value or "").strip()
        if not company:
            continue
        if family_matches(company, stem, variant_set):
            members.append((r, company))

    if len(members) <= 1:
        return None

    parent_key = nk(canonical)
    parent = None
    for r, company in members:
        if company_key(company) == parent_key or nk(company) == parent_key:
            parent = (r, company)
            break
    if parent is None:
        # prefer shortest / bare name
        bare = [
            (r, c)
            for r, c in members
            if strip_geo_suffix(c) == stem and company_key(c) in {stem, parent_key}
        ]
        if bare:
            parent = max(bare, key=lambda m: score_row(ws, m[0], cols))
        else:
            parent = max(members, key=lambda m: score_row(ws, m[0], cols))

    parent_row = parent[0]
    branch_rows = [r for r, _ in members if r != parent_row]

    cities: list[str] = []
    seen: set[str] = set()
    for r, company in members:
        label = branch_label(company, canonical)
        existing = (
            str(ws.cell(r, cols["Cities / Regions"]).value or "").strip()
            if "Cities / Regions" in cols
            else ""
        )
        hq = (
            str(ws.cell(r, cols["Headquarters"]).value or "").strip()
            if "Headquarters" in cols
            else ""
        )
        for cand in (label, existing, hq):
            if not cand:
                continue
            low = cand.lower()
            if low in seen or low in {
                "n/a",
                "not publicly disclosed",
                "global",
                "seed_file",
            }:
                continue
            if len(cand) > 90:
                continue
            # skip if cand is just the canonical again
            if nk(cand) == parent_key:
                continue
            cities.append(cand)
            seen.add(low)

    ws.cell(parent_row, cols["Company"]).value = canonical
    if "Cities / Regions" in cols and cities:
        ws.cell(parent_row, cols["Cities / Regions"]).value = "; ".join(cities)

    defaults = PARENT_DEFAULTS.get(canonical, {})
    for field, val in defaults.items():
        if field in cols:
            cur = str(ws.cell(parent_row, cols[field]).value or "").strip()
            if (
                not cur
                or cur.lower()
                in {"n/a", "na", "not publicly disclosed", "-", "unknown", "seed_file"}
                or field in {"Region Segment", "Regions Served", "Company"}
                or (field == "Headquarters" and len(members) >= 3)
                or (field == "Website" and "not publicly" in cur.lower())
            ):
                ws.cell(parent_row, cols[field]).value = val

    if "Region Segment" in cols and len(members) >= 3 and "Region Segment" not in defaults:
        ws.cell(parent_row, cols["Region Segment"]).value = "Global"
    if "Regions Served" in cols and len(members) >= 3 and "Regions Served" not in defaults:
        cur = str(ws.cell(parent_row, cols["Regions Served"]).value or "").strip()
        if not cur or cur.lower() in {"regional", "n/a", "not publicly disclosed"}:
            ws.cell(parent_row, cols["Regions Served"]).value = (
                "Global / multi-region (consolidated branches)"
            )

    # Prefer a non-junk website from any member
    if "Website" in cols:
        cur = str(ws.cell(parent_row, cols["Website"]).value or "").strip().lower()
        junk = {"", "n/a", "not publicly disclosed", "-", "seed_file"}
        if cur in junk or any(
            x in cur for x in ("loodgieters", "carehospitals", "nature.com")
        ):
            for r, _ in members:
                w = str(ws.cell(r, cols["Website"]).value or "").strip()
                wl = w.lower()
                if (
                    w
                    and wl not in junk
                    and not any(
                        x in wl for x in ("loodgieters", "carehospitals", "nature.com")
                    )
                ):
                    ws.cell(parent_row, cols["Website"]).value = w
                    break

    return {
        "canonical": canonical,
        "stem": stem,
        "kept_row": parent_row,
        "before_names": [c for _, c in members],
        "delete_rows": branch_rows,
        "cities": cities,
    }


def clean_region_sheets(wb, deleted_names: set[str]) -> int:
    removed = 0
    for sheet_name in list(wb.sheetnames):
        if not sheet_name.startswith("Region -"):
            continue
        rws = wb[sheet_name]
        try:
            hdr_row, hdr = find_header(rws)
        except RuntimeError:
            continue
        cols = {h: i + 1 for i, h in enumerate(hdr) if h}
        if "Company" not in cols:
            continue
        to_del = []
        for r in range(hdr_row + 1, rws.max_row + 1):
            co = str(rws.cell(r, cols["Company"]).value or "").strip()
            if co in deleted_names:
                to_del.append(r)
        for r in sorted(to_del, reverse=True):
            rws.delete_rows(r, 1)
            removed += 1
    return removed


def count_companies(ws, hdr_row: int, cols: dict) -> int:
    n = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        if str(ws.cell(r, cols["Company"]).value or "").strip():
            n += 1
    return n


def process_file(path: Path, families: list[dict]) -> dict:
    if not path.exists():
        return {"file": str(path), "missing": True}

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_pre_multi_loc_merge_{stamp}{path.suffix}")
    shutil.copy2(path, bak)

    wb = openpyxl.load_workbook(path)
    ws = None
    hdr_row = None
    cols = None
    for name in wb.sheetnames:
        try:
            hdr_row, hdr = find_header(wb[name])
            cols = {h: i + 1 for i, h in enumerate(hdr) if h}
            ws = wb[name]
            break
        except RuntimeError:
            continue
    if ws is None:
        raise RuntimeError(f"No company sheet in {path}")

    before = count_companies(ws, hdr_row, cols)
    merges = []
    all_delete_rows: list[int] = []
    deleted_names: set[str] = set()

    for fam in families:
        stem = fam["parent_stem"]
        # strip domain annotation if any
        stem = re.sub(r"\s*\[via domain.*?\]\s*$", "", stem).strip()
        canonical = CANONICAL.get(stem)
        if not canonical:
            # title-case fallback
            canonical = " ".join(w.capitalize() for w in stem.split())
        result = merge_family_in_sheet(
            ws, hdr_row, cols, stem, canonical, fam["variants"]
        )
        if not result:
            continue
        merges.append(result)
        for name in result["before_names"]:
            if name != canonical:
                deleted_names.add(name)
        all_delete_rows.extend(result["delete_rows"])

    for r in sorted(set(all_delete_rows), reverse=True):
        ws.delete_rows(r, 1)

    region_removed = clean_region_sheets(wb, deleted_names)
    after = count_companies(ws, hdr_row, cols)

    # Audit sheet
    audit_name = "Multi-Loc Merge Audit"
    if audit_name in wb.sheetnames:
        del wb[audit_name]
    aws = wb.create_sheet(audit_name)
    aws["A1"] = "Multi-location family merge"
    aws["A2"] = datetime.now().isoformat(timespec="seconds")
    aws["A3"] = f"Before: {before}"
    aws["A4"] = f"After: {after}"
    aws["A5"] = f"Removed: {before - after}"
    aws.append([])
    aws.append(["Canonical", "Variants before", "Cities rolled up"])
    for m in merges:
        aws.append(
            [
                m["canonical"],
                " | ".join(m["before_names"]),
                "; ".join(m["cities"][:20]),
            ]
        )

    wb.save(path)
    return {
        "file": path.name,
        "path": str(path),
        "backup": bak.name,
        "before": before,
        "after": after,
        "removed": before - after,
        "families_merged": len(merges),
        "merges": [
            {
                "canonical": m["canonical"],
                "before_names": m["before_names"],
                "cities": m["cities"],
            }
            for m in merges
        ],
        "region_rows_removed": region_removed,
    }


def sync_smartwatch_json(xlsx_path: Path) -> None:
    json_path = xlsx_path.with_suffix(".json")
    if not json_path.exists():
        return
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
    except Exception:
        return
    # only sync if list-of-companies shape
    if not isinstance(data, list):
        return
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = None
    hdr_row = cols = None
    for name in wb.sheetnames:
        try:
            hdr_row, hdr = find_header(wb[name])
            cols = {h: i + 1 for i, h in enumerate(hdr) if h}
            ws = wb[name]
            break
        except RuntimeError:
            continue
    if ws is None:
        wb.close()
        return
    names = set()
    for r in range(hdr_row + 1, ws.max_row + 1):
        co = str(ws.cell(r, cols["Company"]).value or "").strip()
        if co:
            names.add(co)
    wb.close()
    # drop deleted branch names from json if present as company field
    kept = []
    for row in data:
        if not isinstance(row, dict):
            kept.append(row)
            continue
        co = str(row.get("Company") or row.get("company") or "").strip()
        if not co or co in names:
            kept.append(row)
    json_path.write_text(json.dumps(kept, indent=2, ensure_ascii=False), encoding="utf-8")


def main() -> None:
    audit = json.loads(AUDIT_JSON.read_text(encoding="utf-8"))
    families_by_market = audit["families"]
    report = {"ran_at": datetime.now().isoformat(timespec="seconds"), "markets": []}

    for market, paths in MARKET_FILES.items():
        fams = families_by_market.get(market, [])
        print(f"\n======== {market} ({len(fams)} families) ========")
        if not fams:
            print("  nothing to merge")
            continue
        market_rep = {"market": market, "files": []}
        for path in paths:
            rep = process_file(path, fams)
            market_rep["files"].append(rep)
            if rep.get("missing"):
                print("  MISSING", path)
                continue
            print(
                f"  {rep['file']}: {rep['before']} -> {rep['after']} "
                f"(removed {rep['removed']}) families={rep['families_merged']} "
                f"region_cleaned={rep['region_rows_removed']}"
            )
            for m in rep["merges"]:
                print(
                    f"    {m['canonical']}: {len(m['before_names'])} -> 1 "
                    f"[{', '.join(m['before_names'][:6])}{'...' if len(m['before_names'])>6 else ''}]"
                )
            if "smartwatch" in path.name.lower() and path.name.endswith("_FINAL.xlsx"):
                sync_smartwatch_json(path)
        report["markets"].append(market_rep)

    out = OUT_DIR / "multi_location_merge_report.json"
    out.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nWrote {out}")


if __name__ == "__main__":
    main()
