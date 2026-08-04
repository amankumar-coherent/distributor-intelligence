"""Audit channel columns across run4 FINAL Excels."""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"C:\Users\aman.kumar\Desktop\coherent\output\run4")
COLS = [
    "Retail Chains",
    "Specialty Grocers",
    "Foodservice / HoReCa",
    "E-commerce",
    "Distribution Type",
    "Regions Served",
    "Regional Extensions",
    "Warehouse / Logistics",
    "Delivery / Storage",
    "Channel Strength",
]


def find_header(ws):
    for r in range(1, 8):
        vals = [c.value for c in ws[r]]
        if any(v and str(v).strip().lower() == "company" for v in vals):
            return r, vals
    return 1, [c.value for c in ws[1]]


def main() -> None:
    for mid in [
        "global_smartwatch_market_global",
        "global_green_chemical_market_global",
        "global_firewall_market_global",
    ]:
        short = re.sub(r"_global$", "", mid)
        xp = ROOT / mid / f"{short}_FINAL.xlsx"
        wb = load_workbook(xp, read_only=True, data_only=True)
        ws = wb.active
        hr, headers = find_header(ws)
        hmap = {str(h).strip(): i for i, h in enumerate(headers) if h}
        ci = hmap["Company"]
        print(f"\n=== {mid} ===")
        rows = [row for row in ws.iter_rows(min_row=hr + 1, values_only=True) if row and row[ci]]
        print(f"n={len(rows)}")
        for col in COLS:
            if col not in hmap:
                print(f"  {col}: MISSING")
                continue
            idx = hmap[col]
            blank = 0
            vals: Counter[str] = Counter()
            samples = []
            for row in rows:
                s = "" if row[idx] is None else str(row[idx]).strip()
                if not s:
                    blank += 1
                else:
                    vals[s[:90]] += 1
                    if len(samples) < 3:
                        samples.append((str(row[ci])[:28], s[:70]))
            print(f"  {col}: blank={blank}/{len(rows)} unique~{len(vals)}")
            for v, c in vals.most_common(6):
                print(f"    [{c}] {v}")
            for a, b in samples:
                print(f"    eg {a} | {b}")
        wb.close()


if __name__ == "__main__":
    main()
