import openpyxl

X = r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}
for r in range(3, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co in ("Dot Foods, Inc.", "Dot Foods"):
        ws.cell(r, ic["Key Brands Represented"]).value = "The Killer Brownie Co., others"
        ws.cell(r, ic["Summary"]).value = (
            "Dot Foods, Inc. — North America largest food redistributor (HQ Mt. Sterling, IL; founded 1960). "
            "Weekly multi-temp LTL consolidating 1,500+ manufacturer brands for distributors across US and Canada. "
            "No Dot proprietary private-label set; recognized distributed brands include The Killer Brownie Co. among specialty lines. "
            "Natural/specialty assortment and edible oils (mustard, sunflower, olive, rice bran, soybean)."
        )
        print("updated", co)
        print("brands:", ws.cell(r, ic["Key Brands Represented"]).value)
        break
wb.save(X)
wb.close()
print("ok")
