# -*- coding: utf-8 -*-
"""Verify/fill the "Online / Offline / Both" column for the Global Flexible
Packaging Market channel-partner workbook.

DeepSeek is currently out of balance, so this uses a pure HEURISTIC classifier
(no LLM call at all) over three evidence sources:
  1. Direct scan of the company's own website RAW HTML for e-commerce
     platform / cart / checkout signals (regex, deterministic).
  2. Google AI Overview scraper (dedicated instance, RAW markdown —
     llm_clean=False so it never touches DeepSeek).
  3. DDGS web search snippets.

Definitions:
  Online  — customers can complete a purchase directly on the company's own
            website (shopping cart / checkout / "Buy Now" / functioning store).
  Offline — no working online purchase mechanism; sales happen via phone,
            email, sales rep, account manager, or in-person/branch visits.
            A plain "Request a Quote" / "Contact Us" form with no cart still
            counts as Offline.
  Both    — real online ordering/checkout AND meaningful phone/rep/branch
            sales (common for B2B wholesale distributors).

No subagents, no DeepSeek — plain script, run directly:
  .venv\\Scripts\\python.exe scripts\\verify_flexpack_online_offline.py
  .venv\\Scripts\\python.exe scripts\\verify_flexpack_online_offline.py --limit 5   # smoke test

Resumable via checkpoint (skips already-decided companies on re-run).
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")

from vendor_intel.clients.ddgs_engines import normalize_text_backends  # noqa: E402
from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

WORKDIR = ROOT / "output" / "web_expand" / "global_flexible_packaging_market_global"
XLSX = WORKDIR / "global_flexible_packaging_market_global_FINAL.xlsx"
CKPT = WORKDIR / "_flexpack_online_offline_ckpt.json"
AUDIT = WORKDIR / "_flexpack_online_offline_audit.json"
COL = "Online / Offline / Both"

NPD = "Not publicly disclosed"
NPD_SET = {"", "n/a", "na", "-", "—", "not publicly disclosed", "unknown", "tbd", "none", "null"}
VALID = {"Online", "Offline", "Both"}
BACKEND_PARAM, BACKEND_NAMES, _DROPPED = normalize_text_backends()
GAI_BASE_URL = (os.getenv("GAI_BASE_URL") or "http://127.0.0.1:15551").rstrip("/")

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

ECOMMERCE_SIGNALS = [
    (re.compile(r"add[-_ ]?to[-_ ]?(cart|bag|basket)", re.I), "add_to_cart_button"),
    (re.compile(r"shopping[-_ ]?cart", re.I), "shopping_cart_widget"),
    (re.compile(r"/checkout|checkout[-_ ]?(page|process|flow)", re.I), "checkout_flow"),
    (re.compile(r"woocommerce", re.I), "woocommerce_platform"),
    (re.compile(r"cdn\.shopify|shopify\.com|/cart\.js", re.I), "shopify_platform"),
    (re.compile(r"bigcommerce", re.I), "bigcommerce_platform"),
    (re.compile(r"magento", re.I), "magento_platform"),
    (re.compile(r"buy[-_ ]?now", re.I), "buy_now_button"),
    (re.compile(r"your[-_ ]?cart|cart[-_ ]?(count|icon|total)", re.I), "cart_widget"),
    (re.compile(r"online[-_ ]?store|shop[-_ ]?online|order[-_ ]?online", re.I), "online_store_mention"),
]
OFFLINE_SIGNALS = [
    (re.compile(r"request\s+a\s+quote", re.I), "quote_request_only"),
    (re.compile(r"find\s+a\s+(branch|location|distributor)", re.I), "branch_locator"),
    (re.compile(r"call\s+(us\s+)?to\s+order", re.I), "call_to_order"),
    (re.compile(r"contact\s+(us|sales)\s+to\s+(order|purchase|buy)", re.I), "contact_to_order"),
    (re.compile(r"open\s+(a|an)\s+(trade|wholesale|business|credit)\s+account", re.I), "trade_account"),
    (re.compile(r"\btrade\s+login\b|\bwholesale\s+login\b|\bb2b\s+login\b", re.I), "b2b_login"),
    (re.compile(r"(our\s+)?sales\s+team|account\s+manager|sales\s+rep(resentative)?s?\b", re.I), "sales_rep_mention"),
    (re.compile(r"our\s+locations|branch\s+network|distributor\s+network|find\s+your\s+local", re.I), "location_network"),
    (re.compile(r'href=["\']tel:', re.I), "tel_link_present"),
    (re.compile(r"(call|phone|tel)[^.\n]{0,45}(order|sales|quote|wholesale|distributor)", re.I), "phone_near_sales_context"),
    (re.compile(r"(order|sales|quote|wholesale)[^.\n]{0,45}(call|phone|by\s+tel)", re.I), "sales_context_near_phone"),
    (re.compile(r"minimum\s+order\s+quantity|\bMOQ\b|bulk\s+(order|pricing)\s+(inquir|contact)", re.I), "bulk_order_moq"),
    (re.compile(r"physical\s+(branch(es)?|store(s)?|location(s)?)|brick[-\s]and[-\s]mortar", re.I), "physical_presence"),
    (re.compile(r"trade\s+centre(s)?|trade\s+center(s)?|showroom(s)?", re.I), "trade_centre_showroom"),
    (re.compile(r"walk[-\s]in|visit\s+(a|your\s+local|one\s+of\s+our)\s+(branch|store|location)", re.I), "walk_in_visit"),
    (re.compile(r"local\s+(branch|store|depot|warehouse)", re.I), "local_branch"),
]

ONLINE_TEXT_RE = re.compile(
    r"\b(shop\s+online|order(s|ing)?\s+online|online\s+store|add\s+to\s+cart|buy\s+now|"
    r"e-?commerce\s+(checkout|store)|purchase(s|d)?\s+online|online\s+ordering|"
    r"checkout\s+(page|process)|online\s+shop)\b",
    re.I,
)
OFFLINE_TEXT_RE = re.compile(
    r"\b(does\s+not\s+(offer|have|provide)\s+(an\s+)?online\s+(store|ordering|checkout)|"
    r"no\s+online\s+(store|ordering|checkout)|call\s+(us\s+)?to\s+order|"
    r"contact\s+(us|sales|a\s+rep)\s+to\s+(order|purchase|buy)|request\s+a\s+quote\s+to\s+(order|purchase)|"
    r"not\s+(currently\s+)?available\s+(for\s+purchase\s+)?online|"
    r"orders?\s+(are\s+)?(placed|taken)\s+(by|via)\s+phone|"
    r"sales\s+(team|rep(resentative)?)\s+will\s+contact\s+you|"
    r"(our\s+)?sales\s+team|account\s+manager|distributor\s+network|"
    r"open\s+(a|an)\s+(trade|wholesale|business|credit)\s+account|"
    r"minimum\s+order\s+quantity|\bMOQ\b|"
    r"physical\s+(branch(es)?|store(s)?|location(s)?)|brick[-\s]and[-\s]mortar|"
    r"(branch(es)?|trade\s+centre(s)?|trade\s+center(s)?|showroom(s)?)\s+(network|locations?)?|"
    r"walk[-\s]in|in[-\s]person\s+(purchase|shopping|visit)|"
    r"visit\s+(a|your\s+local|one\s+of\s+our)\s+(branch|store|location)|"
    r"local\s+(branch|store|depot|warehouse)|"
    r"\d+\s+(branches|stores|locations|depots))\b",
    re.I,
)


def emptyish(v: object) -> bool:
    return str(v or "").strip().lower() in NPD_SET


def host_of(website: str) -> str:
    w = (website or "").strip()
    if not w:
        return ""
    if not w.startswith("http"):
        w = "https://" + w
    try:
        h = urlparse(w).netloc
        return h[4:] if h.startswith("www.") else h
    except Exception:
        return ""


def strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html or "")
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def scan_website(website: str) -> tuple[list[str], list[str], int]:
    """Returns (ecommerce_signals_found, offline_signals_found, pages_fetched_ok).

    pages_fetched_ok distinguishes "scraped fine, found nothing" from "site
    blocked/timed out entirely" — the latter must NOT be treated as evidence
    of anything (large enterprise sites often block scrapers outright).
    """
    url = (website or "").strip()
    if not url:
        return [], [], 0
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    # Mixed order so a guaranteed early fetch covers both an e-commerce-likely
    # page AND a contact/about-likely page, instead of stopping on ecom alone.
    paths = [
        "", "/contact", "/shop", "/contact-us", "/about", "/store",
        "/about-us", "/products", "/wholesale", "/online-store", "/trade", "/cart",
    ]
    ecom_found: set[str] = set()
    off_found: set[str] = set()
    fetched = 0
    try:
        with httpx.Client(
            headers={"User-Agent": UA}, verify=False, timeout=httpx.Timeout(10.0, connect=4.0), follow_redirects=True
        ) as client:
            for path in paths:
                try:
                    u = url if not path else urljoin(url.rstrip("/") + "/", path.lstrip("/"))
                    r = client.get(u)
                    if r.status_code >= 400:
                        continue
                    raw = r.text
                    for pattern, label in ECOMMERCE_SIGNALS:
                        if pattern.search(raw):
                            ecom_found.add(label)
                    text = strip_html(raw)
                    for pattern, label in OFFLINE_SIGNALS:
                        if pattern.search(text):
                            off_found.add(label)
                    fetched += 1
                    if fetched >= 6:
                        break
                except Exception:
                    continue
    except Exception:
        return [], [], fetched
    return sorted(ecom_found), sorted(off_found), fetched


def ddg_search(query: str, max_results: int = 6) -> list[dict]:
    try:
        from ddgs import DDGS

        with DDGS() as d:
            return list(d.text(query, backend=BACKEND_PARAM, max_results=max_results)) or []
    except Exception as e:  # noqa: BLE001
        print(f"    ddgs error ({query[:60]}...): {e}", flush=True)
        return []


def gather_ddgs_evidence(company: str) -> str:
    queries = [
        f'"{company}" order online OR "shop online" OR e-commerce packaging',
        f'"{company}" "request a quote" OR "call to order" OR showroom',
    ]
    parts: list[str] = []
    for q in queries:
        hits = ddg_search(q, max_results=5)
        time.sleep(0.6)
        if not hits:
            continue
        lines = [
            f"- {h.get('title', '')}: {str(h.get('body', ''))[:250]} ({h.get('href', '')})" for h in hits
        ]
        parts.append(f"Query: {q}\n" + "\n".join(lines))
    return "\n\n".join(parts)[:6000]


_GAI_BAD_MARKERS = ("unusual traffic", "did not match", "ai overview is not available", "no ai overview")


async def gather_gai_evidence(company: str, website: str, gai_ok: bool) -> str:
    if not gai_ok:
        return ""
    host = host_of(website)
    q = f"{company} {host} does it sell online through its own website or only offline via phone quote sales rep"
    try:
        # llm_clean=False: raw scrape only, never touches DeepSeek.
        data = await google_ai_ask(q, base_url=GAI_BASE_URL, llm_clean=False)
    except Exception as e:  # noqa: BLE001
        print(f"    gai error: {e}", flush=True)
        return ""
    md = str(data.get("markdown") or "").strip()
    if not md or any(m in md.lower() for m in _GAI_BAD_MARKERS):
        return ""
    return md[:6000]


def classify_heuristic(
    existing: str, ecom: list[str], off: list[str], site_fetched: int, text_evidence: str
) -> tuple[str, str]:
    online_from_site = bool(ecom)
    offline_from_site = bool(off)
    online_from_text = bool(ONLINE_TEXT_RE.search(text_evidence))
    offline_from_text = bool(OFFLINE_TEXT_RE.search(text_evidence))

    online_signal = online_from_site or online_from_text
    offline_signal = offline_from_site or offline_from_text

    if online_signal and offline_signal:
        return "Both", (
            f"ecom={ecom or ('text-kw' if online_from_text else 'none')} + "
            f"offline={off or ('text-kw' if offline_from_text else 'none')}"
        )

    # The company's own site was completely unreachable (bot-blocked/timeout —
    # common for large enterprise sites), and we already have a verified
    # value. A single weak external-search-only signal (no direct site
    # evidence at all) is not enough to narrow Both -> Online/Offline, or
    # Online -> Offline, etc. Keep the existing value rather than guess.
    site_blocked = site_fetched == 0
    if site_blocked and existing in VALID and not (online_from_site or offline_from_site):
        if online_signal or offline_signal:
            return existing, (
                f"site unreachable (0 pages fetched); weak text-only signal "
                f"insufficient to override verified '{existing}'"
            )

    if online_signal:
        return "Online", f"ecom signals: {ecom or 'keyword match in evidence text'}"
    if offline_signal:
        return "Offline", f"offline signals: {off or 'keyword match in evidence text'}, no online ordering found"
    if existing in VALID:
        return existing, "no new evidence either way; kept existing verified value"
    return NPD, "no evidence found for online/offline classification"


def main() -> int:
    return asyncio.run(main_async())


async def main_async() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0, help="Process only first N unprocessed rows (smoke test)")
    args = ap.parse_args()

    print("DeepSeek: SKIPPED (insufficient balance) -- using pure heuristic classifier, no LLM calls", flush=True)

    gai_ok = False
    try:
        gai_ok = await google_ai_ping(base_url=GAI_BASE_URL)
    except Exception:
        gai_ok = False
    print(f"Google AI scraper @ {GAI_BASE_URL}: {'UP' if gai_ok else 'DOWN (site-scan + DDGS only)'}", flush=True)
    print(f"DDGS backends: {BACKEND_NAMES}", flush=True)

    if not XLSX.exists():
        print(f"ERROR: not found: {XLSX}", file=sys.stderr)
        return 1

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_ONLINE_OFFLINE_VERIFY_{stamp}.xlsx")
    if not bak.exists():
        bak.write_bytes(XLSX.read_bytes())
        print(f"backup -> {bak.name}", flush=True)

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    row_company: list[tuple[int, str]] = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co:
            row_company.append((r, co))
    print(f"total rows: {len(row_company)}", flush=True)

    ckpt: dict = {"done": {}}
    if CKPT.exists():
        try:
            ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
        except Exception:
            pass
    done = ckpt.setdefault("done", {})

    todo = [(r, co) for r, co in row_company if co not in done]
    if args.limit:
        todo = todo[: args.limit]
    print(f"to process this run: {len(todo)} (already done: {len(row_company) - len(todo)})", flush=True)

    def _scan(item: tuple[int, str]) -> tuple[int, str, list[str], list[str], int]:
        r, co = item
        website = str(ws.cell(r, ic["Website"]).value or "").strip()
        ecom, off, fetched = scan_website(website)
        return r, website, ecom, off, fetched

    scan_cache: dict[int, tuple[str, list[str], list[str], int]] = {}
    with ThreadPoolExecutor(max_workers=16) as ex:
        futs = [ex.submit(_scan, item) for item in todo]
        n = 0
        for fut in as_completed(futs):
            r, website, ecom, off, fetched = fut.result()
            scan_cache[r] = (website, ecom, off, fetched)
            n += 1
            if n % 40 == 0:
                print(f"  site-scan {n}/{len(todo)}", flush=True)

    # DDGS has been the slow/unreliable part all session (frequent timeouts).
    # Only fall back to it when the Google AI scraper didn't already return
    # substantial evidence, and run several rows' evidence-gathering
    # concurrently instead of one at a time.
    GAI_SUFFICIENT_LEN = 150
    CONCURRENCY = 6

    async def evidence_for_row(co: str, website: str) -> str:
        gai_ev = await gather_gai_evidence(co, website, gai_ok)
        if gai_ev and len(gai_ev) >= GAI_SUFFICIENT_LEN:
            return gai_ev
        ddgs_ev = await asyncio.to_thread(gather_ddgs_evidence, co)
        return "\n\n".join(p for p in (gai_ev, ddgs_ev) if p)

    changed = 0
    processed = 0
    for start in range(0, len(todo), CONCURRENCY):
        batch = todo[start : start + CONCURRENCY]
        evidences = await asyncio.gather(
            *(evidence_for_row(co, scan_cache.get(r, ("", [], [], 0))[0]) for r, co in batch)
        )
        for (r, co), text_evidence in zip(batch, evidences):
            website, ecom, off, site_fetched = scan_cache.get(r, ("", [], [], 0))
            existing = str(ws.cell(r, ic[COL]).value or "").strip()

            new_val, reason = classify_heuristic(existing, ecom, off, site_fetched, text_evidence)
            if new_val != existing:
                ws.cell(r, ic[COL]).value = new_val
                changed += 1

            done[co] = {"value": new_val, "prev": existing, "reason": reason}
            processed += 1
            print(
                f"[{processed}/{len(todo)}] {co} -> {new_val!r} (was {existing!r}) [{reason[:80]}]",
                flush=True,
            )

        wb.save(XLSX)
        CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"  checkpoint saved at {processed}/{len(todo)}", flush=True)

    wb.save(XLSX)
    CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")

    dist: dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(r, ic[COL]).value or "").strip()
        dist[v] = dist.get(v, 0) + 1

    audit = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "processed_this_run": len(todo),
        "changed_this_run": changed,
        "final_distribution": dist,
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(audit, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
