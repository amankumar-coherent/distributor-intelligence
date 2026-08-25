#!/usr/bin/env python3
"""Rebuild Specialty Focus as a semicolon-separated list of CLINICAL/APPLICATION
focus areas — a different dimension from Core Categories (which lists product
types), matching the requested format:

    Core Categories: Medical Wearables; Remote Monitoring Devices; Health Sensors
    Speciality Focus: Cardiac Monitoring; Diabetes Monitoring; Chronic Disease Management

Reads the ORIGINAL research prose from the checkpoint JSON files (untouched by
the earlier sentence-template reformat) plus the sheet's own Summary column.
Pure keyword rules, no LLM, no new web calls. Rows flagged "Needs manual
verification" are left untouched.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import openpyxl

XLSX = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/global_wearable_medical_devices_market_global_FINAL.xlsx"
SHEET = "Wearable Medical Companies"
CKPT_DIR = ROOT / "output/web_expand/global_wearable_medical_devices_market_global/verify_checkpoints"
FLAG_PREFIX = "Needs manual verification"

# (keyword pattern, application/clinical-focus label) -- checked in order,
# most specific first. This vocabulary is deliberately DIFFERENT from the
# Core Categories product-type vocabulary.
_FOCUS_RULES: list[tuple[str, str]] = [
    (r"continuous glucose|\bcgm\b|diabetes|glucose|insulin", "Diabetes Monitoring"),
    (r"cardiac|\becg\b|\bekg\b|heart|cardiovascular|holter", "Cardiac Monitoring"),
    (r"sleep (?:apnea|tracking|monitor)", "Sleep Monitoring"),
    (r"respiratory|pulmonary|oxygen therapy|ventilat|\bcpap\b", "Respiratory Care"),
    (r"neonatal|maternal|paediatric|pediatric", "Neonatal & Maternal Care"),
    (r"wound care|wound management", "Wound Care"),
    (r"surgical|surgery|operating room", "Surgical Support"),
    (r"rehabilitation|physiotherapy|physical therapy|mobility aid|wheelchair", "Rehabilitation & Mobility"),
    (r"infection (?:control|prevention)|sterili[sz]ation", "Infection Control"),
    (r"emergency|first aid|resuscitation|\baed\b|defibrillat", "Emergency Response"),
    (r"\bdental\b", "Dental Care"),
    (r"veterinar", "Veterinary Care"),
    (r"home\s?care|home\s?health", "Home Healthcare"),
    (r"hospice|palliative", "Palliative Care"),
    (r"diagnostic|laboratory|pathology", "Diagnostics"),
    (r"orthopedic|orthopaedic", "Orthopedic Care"),
    (r"mental health|behavioral health|psychiatric", "Behavioral Health"),
    (r"oncology|\bcancer\b", "Oncology Care"),
    (r"renal|dialysis|kidney", "Renal Care"),
    (r"vaccin|immuni[sz]ation", "Immunization"),
    (r"ostomy|urology|continence", "Urology & Continence Care"),
    (r"hearing|audiolog", "Hearing Care"),
    (r"vision|ophthalmic|optical", "Vision Care"),
    (r"fitness tracker|activity tracking|fitness monitoring", "Fitness & Wellness Monitoring"),
    (r"chronic (?:disease|condition)", "Chronic Disease Management"),
]

# Fallback when no specific clinical-focus keyword is found (e.g. a broadline
# wholesaler with no particular niche) -- still honest, not invented.
_FALLBACK_RULES: list[tuple[str, str]] = [
    (r"broadline|full.line|wholesale distribut|wholesaler", "General Pharmaceutical & Medical Distribution"),
    (r"retail pharmac|pharmacy chain|drugstore", "OTC & Retail Health Products"),
    (r"\bit\b|electronic|industrial", "General Industrial/IT Distribution"),
]
_DEFAULT_FALLBACK = "General Healthcare Distribution"


def derive_focus_areas(text: str, max_items: int = 4) -> str:
    low = text.lower()
    hits: list[str] = []
    for pattern, label in _FOCUS_RULES:
        if label in hits:
            continue
        if re.search(pattern, low):
            hits.append(label)
        if len(hits) >= max_items:
            break
    if hits:
        return "; ".join(hits)

    for pattern, label in _FALLBACK_RULES:
        if re.search(pattern, low):
            return label
    return _DEFAULT_FALLBACK


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}

    updated = 0
    skipped_flagged = 0
    for r in range(2, ws.max_row + 1):
        company = ws.cell(row=r, column=idx["Company"]).value
        if not company:
            continue
        core = str(ws.cell(row=r, column=idx["Core Categories"]).value or "").strip()
        if core.startswith(FLAG_PREFIX):
            skipped_flagged += 1
            continue

        ckpt_path = CKPT_DIR / f"{r:04d}.json"
        original_focus = ""
        if ckpt_path.exists():
            try:
                original_focus = json.loads(ckpt_path.read_text(encoding="utf-8")).get("specialty_focus", "")
            except Exception:
                original_focus = ""
        if original_focus.startswith(FLAG_PREFIX):
            skipped_flagged += 1
            continue

        summary = str(ws.cell(row=r, column=idx["Summary"]).value or "").strip()
        source_text = f"{core} {original_focus} {summary}"

        new_focus = derive_focus_areas(source_text)
        ws.cell(row=r, column=idx["Specialty Focus"], value=new_focus)
        updated += 1

    wb.save(XLSX)
    print(f"Updated {updated} rows, skipped {skipped_flagged} flagged rows.")


if __name__ == "__main__":
    main()
