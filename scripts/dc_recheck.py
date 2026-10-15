#!/usr/bin/env python3
"""Re-verify the 2026-10-09 Data Center System distributor list against the
Global Data Center Market definition (user, 2026-10-15) -- no discovery --
then run the usual final steps (status + ownership check, polished Excel +
JSON, row 2 = dc_markets.DEFINITIONS["dc_distributor"]).

Steps (each resumable -- re-running skips what is done):
  1. split the 421 companies into PARTS files (header row + rows);
  2. verify the parts in parallel, one verifier + browser profile each,
     market name "Global Data Center Market" so verify_distributors_ai_mode's
     data-center DISTRIBUTOR rule applies, VERIFY_CHUNK=8;
  3. merge the kept rows into output/final/dc/dc_verified/
     dc_distributor_verified_distributors.xlsx;
  4. status + ownership check and build -> output/final/dc/
     dc distributors (status checked)/.

Log: logs/dc_recheck.log (launch via logs/run_dc_recheck.ps1).
"""
from __future__ import annotations

import os
import subprocess
import sys
import time
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "output" / "market_research" / "global-data-center-system-market-distributors-only_verified_distributors.xlsx"
WORK = ROOT / "output" / "final" / "dc" / "recheck"
STAGE = ROOT / "output" / "final" / "dc" / "dc_verified"
OUT = ROOT / "output" / "final" / "dc" / "dc distributors (status checked)"
MARKET = "Global Data Center Market"
PARTS = 10  # user 2026-10-15: "open 10 browsers to verify, not only 3"
STATUS_WORKERS = int(os.environ.get("DC_STATUS_WORKERS", "20"))  # user: 20 browsers


def log(msg: str) -> None:
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def header_row(ws) -> int:
    for r in range(1, 15):
        if any(str(c.value).strip() == "Company" for c in ws[r] if c.value is not None):
            return r
    raise SystemExit(f"no Company header in {ws.title}")


def _old_verdicts_and_reset() -> list[dict]:
    """A split with a different PARTS count: keep its verdicts, move it aside."""
    import json
    import shutil

    verdicts = []
    for f in WORK.glob("*_verify_checkpoint.json"):
        try:
            verdicts += json.loads(f.read_text(encoding="utf-8")).get("verdicts", [])
        except (OSError, ValueError):
            pass
    shutil.move(str(WORK), str(WORK.with_name(f"recheck_old_{time.strftime('%Y%m%d_%H%M%S')}")))
    good = [v for v in verdicts if not v.get("backend_failure")]
    log(f"re-splitting into {PARTS} parts; carrying over {len(good)} verdicts")
    return good


def split() -> list[Path]:
    import json

    parts = [WORK / f"global-data-center-market-part{i + 1}.xlsx" for i in range(PARTS)]
    if all(p.is_file() for p in parts):
        return parts
    carried = _old_verdicts_and_reset() if WORK.is_dir() and any(WORK.glob("*part*.xlsx")) else []
    WORK.mkdir(parents=True, exist_ok=True)
    ws = openpyxl.load_workbook(SRC, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows[:15]) if r and "Company" in [str(c).strip() for c in r if c is not None])
    head, data = rows[hdr], [r for r in rows[hdr + 1:] if r and any(r)]
    size = -(-len(data) // PARTS)
    ci = [str(c).strip() if c is not None else "" for c in head].index("Company")
    for i, p in enumerate(parts):
        wb = openpyxl.Workbook()
        out = wb.active
        out.append(list(head))
        chunk = data[i * size:(i + 1) * size]
        for r in chunk:
            out.append(list(r))
        wb.save(p)
        names = {str(r[ci] or "").strip().lower() for r in chunk}
        seed = [v for v in carried if str(v.get("Company") or "").strip().lower() in names]
        if seed:
            ck = p.with_name(f"{p.stem}_verify_checkpoint.json")
            ck.write_text(json.dumps({"offset": 0, "verdicts": seed}, ensure_ascii=False, indent=2),
                          encoding="utf-8")
    log(f"split {len(data)} companies into {PARTS} parts of <= {size}")
    return parts


def verify(parts: list[Path]) -> list[Path]:
    procs = []
    for i, p in enumerate(parts):
        done = p.with_name(f"{p.stem}_verified_distributors.xlsx")
        env = dict(os.environ, VERIFY_CHUNK=os.environ.get("VERIFY_CHUNK", "8"),
                   GOOGLE_AI_MODE_MAX_RESETS="10", PYTHONIOENCODING="utf-8",
                   GOOGLE_AI_MODE_PROFILE_DIR=str(ROOT / "data" / f"ai_mode_dcverify_{i + 1}"))
        logf = (ROOT / "logs" / f"dc_recheck_verify_part{i + 1}.log").open("a", encoding="utf-8")
        procs.append(subprocess.Popen(
            [sys.executable, "-u", str(ROOT / "verify_distributors_ai_mode.py"),
             "--file", str(p), "--market-name", MARKET],
            cwd=str(ROOT), env=env, stdout=logf, stderr=subprocess.STDOUT))
        log(f"verifier part{i + 1} started (pid {procs[-1].pid}){' [resuming]' if done.exists() else ''}")
        time.sleep(15)  # stagger browser launches
    for p in procs:
        p.wait()
    outs = [p.with_name(f"{p.stem}_verified_distributors.xlsx") for p in parts]
    missing = [o.name for o in outs if not o.is_file()]
    if missing:
        raise SystemExit(f"verifier output missing: {missing} -- re-run to resume")
    return outs


def merge(outs: list[Path]) -> Path:
    STAGE.mkdir(parents=True, exist_ok=True)
    dest = STAGE / "dc_distributor_verified_distributors.xlsx"
    wb = openpyxl.load_workbook(outs[0])
    ws = wb.active
    h = header_row(ws)
    heads = [str(c.value).strip() if c.value is not None else "" for c in ws[h]]
    ci = heads.index("Company")
    seen = {str(ws.cell(r, ci + 1).value or "").strip().lower()
            for r in range(h + 1, ws.max_row + 1)} - {""}
    for o in outs[1:]:
        w2 = openpyxl.load_workbook(o, read_only=True).active
        rows = list(w2.iter_rows(values_only=True))
        h2 = next(i for i, r in enumerate(rows[:15]) if r and "Company" in [str(c).strip() for c in r if c is not None])
        heads2 = [str(c).strip() if c is not None else "" for c in rows[h2]]
        for r in rows[h2 + 1:]:
            row = dict(zip(heads2, r))
            name = str(row.get("Company") or "").strip()
            if not name or name.lower() in seen:
                continue
            seen.add(name.lower())
            ws.append([row.get(c) for c in heads])
    if "#" in heads:
        col, n = heads.index("#") + 1, 0
        for r in range(h + 1, ws.max_row + 1):
            if ws.cell(r, ci + 1).value:
                n += 1
                ws.cell(r, col).value = n
    wb.save(dest)
    log(f"merged {len(seen)} kept companies -> {dest}")
    return dest


def status_check() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    env = dict(os.environ, STATUS_WALL_WAITS="2", PYTHONIOENCODING="utf-8")
    log(f"status + ownership check ({STATUS_WORKERS} workers)")
    subprocess.run([sys.executable, "-u", str(ROOT / "scripts" / "status_check_folder.py"),
                    "--input", str(STAGE), "--output", str(OUT), "--workers", str(STATUS_WORKERS)],
                   cwd=str(ROOT), env=env)
    report = OUT / "status_check_report.xlsx"
    if report.is_file() and report.stat().st_mtime > time.time() - 900:
        log(f"DONE -> {OUT}")
    else:
        log("status check incomplete (blocked or unchecked companies) -- re-run to resume")


def main() -> None:
    log(f"source: {SRC.name}")
    parts = split()
    outs = verify(parts)
    merge(outs)
    status_check()


if __name__ == "__main__":
    main()
