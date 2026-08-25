# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pathlib import Path

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}
FILES = {
    "electrical": Path(
        r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
        r"\global_electrical_distribution_market_global_FINAL.xlsx"
    ),
    "avocado": Path(
        r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
        r"\global_avocado_oil_market.xlsx"
    ),
    "smartphone": Path(
        r"D:\vendor-intel\output\web_expand\global_smartphone_market_global"
        r"\global_smartphone_market.xlsx"
    ),
    "firewall": Path(
        r"D:\vendor-intel\output\run4\global_firewall_market_global"
        r"\global_firewall_market.xlsx"
    ),
}

print("market\trows\tcp_filled\tcp_pct\tli_url\tli_npd\tli_empty")
for key, path in FILES.items():
    wb = load_workbook(path, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            headers = vals
            hr = r
            break
    tot = cp = li_url = li_npd = li_empty = 0
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {
            headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
        }
        co = d.get("Company") or ""
        if not co:
            continue
        tot += 1
        person = (d.get("Contact Person") or "").strip()
        li = (d.get("LinkedIn") or "").strip()
        has_cp = bool(person) and person.lower() not in NPD
        if has_cp:
            cp += 1
        li_l = li.lower()
        if "linkedin.com" in li_l:
            li_url += 1
        elif not li or li_l in NPD:
            if has_cp:
                li_npd += 1
            else:
                li_empty += 1
        else:
            li_empty += 1
    pct = (100.0 * cp / tot) if tot else 0
    print(f"{key}\t{tot}\t{cp}\t{pct:.1f}%\t{li_url}\t{li_npd}\t{li_empty}")
    wb.close()
