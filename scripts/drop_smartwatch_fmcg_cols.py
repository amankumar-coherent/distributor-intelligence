"""Remove FMCG-only columns (Foodservice, Specialty Grocers) from smartwatch FINAL Excel."""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

DROP = {"Specialty Grocers", "Foodservice / HoReCa"}
PATHS = [
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL.xlsx"),
    Path(
        r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL_320plus.xlsx"
    ),
]


def drop_cols(path: Path) -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_pre_drop_fmcg_cols_{stamp}{path.suffix}")
    shutil.copy2(path, bak)
    wb = openpyxl.load_workbook(path)
    ws = wb["Smartwatch Companies"]
    hdr = [c.value for c in ws[3]]
    to_del = [i for i, h in enumerate(hdr, 1) if str(h or "") in DROP]
    for i in reversed(to_del):
        ws.delete_cols(i, 1)
    new_hdr = [c.value for c in ws[3]]
    wb.save(path)
    print(path.name)
    print("  backup", bak.name)
    print("  removed", [hdr[i - 1] for i in to_del])
    print("  cols now", len([h for h in new_hdr if h is not None]), ":", [str(h) for h in new_hdr if h])
    print()


if __name__ == "__main__":
    for p in PATHS:
        drop_cols(p)
