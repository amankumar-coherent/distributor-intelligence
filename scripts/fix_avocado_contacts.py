# -*- coding: utf-8 -*-
"""Clean junk phrases from avocado oil Contact Person column."""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_global_FINAL.xlsx"
)

NPD = "Not publicly disclosed"

# Exact old -> new (when regex alone is ambiguous)
EXACT = {
    "abdul sattar personal appointments": "Abdul Sattar",
    "odhran recently sat down": "Odhran",
    "jeff raikes and co-chair": NPD,  # wrong scrape / not a Diconsa contact
    "met with cesar c.": "Cesar C.",
    "phone number": NPD,
    "appointment of cathal geoghegan": "Cathal Geoghegan",
    "and key executive team": NPD,
    "of finance and administration": NPD,
    "spoke to tony reynolds": "Tony Reynolds",
    # leftovers after first clean pass
    "norkem. norkem's main": NPD,
    "roveg fruit b.v.": NPD,
    "employees listrichard scavuzzo": "Richard Scavuzzo",
    "swire hotels": NPD,
    "and darren fischer": "Darren Fischer",
}


def clean_contact(raw: str) -> tuple[str, str]:
    """Return (new_value, reason). Empty reason = unchanged."""
    s = (raw or "").strip()
    if not s:
        return s, ""
    low = re.sub(r"\s+", " ", s.lower()).strip()

    if low in {"not publicly disclosed", "n/a", "na", "unknown", "-", "none"}:
        return NPD if low != "not publicly disclosed" else s, (
            "normalize placeholder" if low != "not publicly disclosed" else ""
        )

    if low in EXACT:
        new = EXACT[low]
        return new, f"exact map -> {new}"

    original = s

    # Strip common scrape prefixes/suffixes
    patterns = [
        (r"(?i)^appointment of\s+", ""),
        (r"(?i)^appointments? of\s+", ""),
        (r"(?i)^met with\s+", ""),
        (r"(?i)^spoke (?:to|with)\s+", ""),
        (r"(?i)^interviewed\s+", ""),
        (r"(?i)\s+personal appointments?\s*$", ""),
        (r"(?i)\s+recently sat down\s*$", ""),
        (r"(?i)\s+sat down\s*$", ""),
        (r"(?i)^and\s+", ""),
        (r"(?i)^of\s+", ""),
        (r"(?i)\s+and co-?chair\s*$", ""),
        (r"(?i)\s+key executive team\s*$", ""),
        (r"(?i)^key executive team\s*$", NPD),
        (r"(?i)^phone number\s*$", NPD),
        (r"(?i)^email address\s*$", NPD),
        (r"(?i)^linkedin\s*$", NPD),
        (r"(?i)^ceo\s*$", NPD),
        (r"(?i)^president\s*$", NPD),
        (r"(?i)^managing director\s*$", NPD),
        (r"(?i)^finance and administration\s*$", NPD),
        (r"(?i)^of finance and administration\s*$", NPD),
    ]
    for pat, repl in patterns:
        if repl == NPD and re.fullmatch(pat, s):
            return NPD, f"placeholder via {pat}"
        news = re.sub(pat, repl if repl != NPD else "", s).strip(" ,;-")
        if news != s:
            s = news

    s = re.sub(r"\s+", " ", s).strip(" ,;-")

    # If leftover is clearly not a person name
    if not s or re.fullmatch(
        r"(?i)(team|executive team|management|leadership|board|phone|email|number|co-?chair|finance|administration)",
        s,
    ):
        return NPD, "not a person name"

    # Title-case obvious ALL CAPS first/last if short name
    if s.isupper() and 1 <= len(s.split()) <= 4:
        s = s.title()

    if s != original:
        return s, f"cleaned '{original}' -> '{s}'"
    return original, ""


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = PATH.with_name(f"{PATH.stem}_pre_contact_fix_{stamp}{PATH.suffix}")
    shutil.copy2(PATH, bak)

    wb = openpyxl.load_workbook(PATH)
    ws = wb["Avocado Oil Distributors"]
    hdr = 2
    cols = {str(c.value): i + 1 for i, c in enumerate(ws[hdr]) if c.value}
    cp_col = cols["Contact Person"]
    co_col = cols["Company"]

    fixes = []
    scanned = 0
    for r in range(hdr + 1, ws.max_row + 1):
        co = str(ws.cell(r, co_col).value or "").strip()
        if not co:
            continue
        scanned += 1
        old = str(ws.cell(r, cp_col).value or "").strip()
        new, reason = clean_contact(old)
        if reason and new != old:
            ws.cell(r, cp_col).value = new
            fixes.append((r, co, old, new, reason))

    # Pass 2: leftover scrape artifacts that look like company/role text, not a person
    ROW_FORCE = {
        # may already be cleaned once — force final good values
    }
    extra_bad_re = re.compile(
        r"(?i)employees\s*list|norkem'?s?\s+main|b\.?v\.?|swire hotels|"
        r"list[A-Z]|foodservice|distribution|trading ltd|"
        r"^[A-Z][a-z]+ [A-Z][a-z]+ (Inc|LLC|Ltd|GmbH|B\.?V\.?|Group|Foods|Fruit|Hotels)\b"
    )
    for r in range(hdr + 1, ws.max_row + 1):
        co = str(ws.cell(r, co_col).value or "").strip()
        if not co:
            continue
        old = str(ws.cell(r, cp_col).value or "").strip()
        if not old or old == NPD:
            continue
        new = None
        reason = ""
        # Glued 'Employees ListRichard Scavuzzo'
        m = re.search(r"(?i)(?:employees\s*)?list([A-Z][a-z]+(?:\s+[A-Z][a-z]+)+)", old)
        if m:
            new = m.group(1).strip()
            reason = "extract name after List"
        elif re.search(r"[a-z][A-Z]", old):  # CamelGlue
            m2 = re.search(r"([A-Z][a-z]+(?:\s+[A-Z][a-z]+)+)$", old)
            new = m2.group(1) if m2 else NPD
            reason = "unglue camel name"
        elif extra_bad_re.search(old) or old.lower() in {
            "roveg fruit b.v.",
            "swire hotels",
            "norkem. norkem's main",
        }:
            new = NPD
            reason = "company/role text not a person"
        elif re.fullmatch(
            r"(?i).*\b(inc|llc|ltd|gmbh|b\.?v\.?|company|group|hotels|fruit)\b.*",
            old,
        ) and not re.match(r"^[A-Z][a-z]+ [A-Z]", old):
            new = NPD
            reason = "looks like org name"
        if new and new != old:
            ws.cell(r, cp_col).value = new
            fixes.append((r, co, old, new, reason))

    still = []
    junk_re = re.compile(
        r"(?i)\b(personal|appointments?|recently|sat down|met with|spoke to|spoke with|"
        r"phone number|email address|key executive team|co-chair|appointment of|"
        r"employees list|norkem'?s main)\b"
    )
    for r in range(hdr + 1, ws.max_row + 1):
        co = str(ws.cell(r, co_col).value or "").strip()
        if not co:
            continue
        cp = str(ws.cell(r, cp_col).value or "").strip()
        if junk_re.search(cp) or re.search(r"[a-z][A-Z]", cp):
            still.append((r, co, cp))

    an = "Contact Fix Audit"
    if an in wb.sheetnames:
        del wb[an]
    aws = wb.create_sheet(an)
    aws.append(["Row", "Company", "Old Contact", "New Contact", "Reason"])
    for row in fixes:
        aws.append(list(row))

    wb.save(PATH)
    print(f"backup: {bak.name}")
    print(f"scanned: {scanned}")
    print(f"fixed: {len(fixes)}")
    for r, co, old, new, reason in fixes:
        print(f"  {r}: [{co}] {old!r} -> {new!r} ({reason})")
    print(f"still junk: {still or 'none'}")


if __name__ == "__main__":
    main()
