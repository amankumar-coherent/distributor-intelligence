import openpyxl

X = r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}
for r in range(3, ws.max_row + 1):
    if str(ws.cell(r, ic["Company"]).value or "").strip() == "I-D Foods":
        ws.cell(r, ic["Key Brands Represented"]).value = (
            "Haiku, M'Lord, St Regis, So Veggie, The Great Gentleman, others"
        )
        print("fixed I-D Foods")
        break
wb.save(X)
wb.close()
print("ok")
