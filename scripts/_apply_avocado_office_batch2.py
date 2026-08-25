#!/usr/bin/env python3
"""Apply web-verified office contacts batch 2."""
from __future__ import annotations

import openpyxl
from pathlib import Path

p = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
    r"\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
wb = openpyxl.load_workbook(p)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

FILLS = {
    "Brakes": {
        "website": "https://www.brake.co.uk",
        "office_no": "+44 345 606 9090",
        "office_email": "customer.services@brake.co.uk",
    },
    "C&S Wholesale Grocers": {
        "office_no": "(603) 354-7000",
        "office_email": "ir@cswg.com",
    },
    "Sysco": {
        "office_no": "(281) 584-1390",
        "office_email": "customer@sysco.com",
    },
    "Gordon Food Service": {
        "website": "https://www.gfs.com",
        "office_no": "(616) 530-7000",
    },
    "Brenntag": {
        "office_no": "+49 201 6496-0",
        "office_email": "ir@brenntag.de",
    },
    "Chefs' Warehouse": {
        "office_no": "(203) 894-1345",
        "office_email": "ir@chefswarehouse.com",
    },
}


def empty(v: object) -> bool:
    s = str(v or "").strip().lower()
    return s in ("", "not publicly disclosed", "n/a", "none")


n = 0
for row in range(3, ws.max_row + 1):
    co = str(ws.cell(row, ic["Company"]).value or "").strip()
    if co not in FILLS:
        continue
    hit = FILLS[co]
    if hit.get("website"):
        ws.cell(row, ic["Website"]).value = hit["website"]
    if hit.get("office_no"):
        ws.cell(row, ic["Office No."]).value = hit["office_no"]
    if hit.get("office_email"):
        ws.cell(row, ic["Office Email"]).value = hit["office_email"]
    n += 1
    print("filled", co)

wb.save(p)
wb.close()

wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c or "") for c in rows[1]]
ic2 = {h: i for i, h in enumerate(hdr) if h}
phone = sum(
    1
    for r in rows[2:]
    if r[ic2["Office No."]] and str(r[ic2["Office No."]]).strip() != "Not publicly disclosed"
)
email = sum(
    1
    for r in rows[2:]
    if r[ic2["Office Email"]] and str(r[ic2["Office Email"]]).strip() != "Not publicly disclosed"
)
print("TOTAL phone", phone, "email", email, "of", len(rows) - 2, "applied", n)
wb.close()
