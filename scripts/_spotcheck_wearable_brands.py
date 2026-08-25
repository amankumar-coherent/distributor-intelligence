# -*- coding: utf-8 -*-
from openpyxl import load_workbook

p = r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global\global_wearable_medical_devices_market_global_FINAL.xlsx"
wb = load_workbook(p)
ws = wb.active
h = [c.value for c in ws[1]]
ic = {x: i + 1 for i, x in enumerate(h) if x}
NPD = "Not publicly disclosed"

for r in range(2, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    v = str(ws.cell(r, ic["Key Brands Represented"]).value or "")
    if co == "Rexall Pharmacy" and ("M&M" in v or "Always" in v or "L'Oréal" in v):
        ws.cell(r, ic["Key Brands Represented"]).value = "Multi-brand pharmacy assortment, others"
        print("fixed Rexall")

check = [
    "AAH Pharmaceuticals",
    "Best Buy Health",
    "Zahrawi Group",
    "Bound Tree Medical",
    "Arrow Electronics",
    "Mediplus UK",
    "Medisave UK",
    "Performance Health",
    "LifeHealthcare",
    "Cardinal Health",
    "Medi-Dyne Healthcare Products",
    "Ecolab",
]
for r in range(2, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co in check:
        print(co, "=>", str(ws.cell(r, ic["Key Brands Represented"]).value)[:130])

wb.save(p)
filled = npd = 0
for r in range(2, ws.max_row + 1):
    v = str(ws.cell(r, ic["Key Brands Represented"]).value or "").strip()
    if v == NPD:
        npd += 1
    else:
        filled += 1
print("final filled", filled, "npd", npd, "rows", ws.max_row - 1)
