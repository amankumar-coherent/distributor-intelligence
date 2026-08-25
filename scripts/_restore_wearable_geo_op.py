# -*- coding: utf-8 -*-
"""
Restore Continent / Geography + Operational Presence from BEFORE_GEO_OP backup.
Keep Company / Website / Founded / Headquarters untouched.
- Convert REG(Country) -> Country names
- Restore prior multi-country commercial lists (cleaned)
- Merge verified multi-country expansions (union, never shrink searched data)
"""
from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
AUDIT = ROOT / "_wearable_geo_op_restore_audit.json"

ALIAS = {
    "usa": "United States",
    "us": "United States",
    "uk": "United Kingdom",
    "u.k.": "United Kingdom",
    "uae": "United Arab Emirates",
    "korea": "South Korea",
    "republic of korea": "South Korea",
    "hong kong sar": "Hong Kong",
    "mainland china": "China",
    "czech republic": "Czechia",
    "czechia": "Czechia",
    "türkiye": "Turkey",
    "turkiye": "Turkey",
    "the netherlands": "Netherlands",
    "holland": "Netherlands",
    "viet nam": "Vietnam",
}

DROP = {
    "europe",
    "asia",
    "asia pacific",
    "apac",
    "emea",
    "mena",
    "latam",
    "latin america",
    "middle east",
    "africa",
    "north america",
    "global",
    "international",
    "worldwide",
    "asean",
    "gcc",
    "pacific islands",
}

# Verified expansions to MERGE (union) — never remove prior searched countries
PHYSICAL_EXTRA: dict[str, list[str]] = {
    "Transmedic": [
        "Singapore",
        "Malaysia",
        "Indonesia",
        "Thailand",
        "Vietnam",
        "Philippines",
        "Hong Kong",
    ],
    "IDS Medical Systems": [
        "Singapore",
        "Malaysia",
        "Indonesia",
        "Philippines",
        "Thailand",
        "Vietnam",
    ],
    "Alliance Healthcare": [
        "United Kingdom",
        "France",
        "Germany",
        "Spain",
        "Netherlands",
        "Czechia",
        "Norway",
        "Romania",
        "Turkey",
    ],
    "Alloga": ["France", "Spain"],
    "Watsons": [
        "Hong Kong",
        "China",
        "Taiwan",
        "Macau",
        "Singapore",
        "Malaysia",
        "Thailand",
        "Philippines",
        "Indonesia",
        "Vietnam",
        "Turkey",
        "Qatar",
        "United Arab Emirates",
        "Saudi Arabia",
        "Bahrain",
    ],
    "EBOS Healthcare": ["Australia", "New Zealand"],
    "Device Technologies": ["Australia", "New Zealand"],
    "BioD Medica": ["Malaysia", "Singapore"],
    "Aster Pharmacy": [
        "United Arab Emirates",
        "India",
        "Saudi Arabia",
        "Oman",
        "Qatar",
        "Bahrain",
        "Jordan",
    ],
    "Zuellig Pharma": [
        "Singapore",
        "Malaysia",
        "Thailand",
        "Philippines",
        "Vietnam",
        "Indonesia",
        "Hong Kong",
        "Taiwan",
        "South Korea",
        "Brunei",
        "Cambodia",
        "Myanmar",
    ],
    "Walmart": ["United States", "Canada"],
    "Amazon": ["United States", "Canada"],
    "Henry Schein": [
        "United States",
        "Canada",
        "United Kingdom",
        "Germany",
        "France",
        "Spain",
        "Netherlands",
        "Australia",
        "New Zealand",
        "Brazil",
        "Singapore",
    ],
    "McKesson Corporation": ["United States", "Canada", "Germany"],
    "GrafiMedics BeLux": ["Belgium", "Luxembourg"],
    "Mediq": ["Netherlands", "Germany"],
    "DocMorris": ["Germany", "Switzerland", "Netherlands", "Spain", "France"],
    "Farmacias Cruz Verde": ["Chile", "Colombia"],
    "Invacare Europe": ["Switzerland"],
}

COMMERCIAL_EXTRA = {
    **PHYSICAL_EXTRA,
    "Zuellig Pharma": PHYSICAL_EXTRA["Zuellig Pharma"]
    + ["China", "India", "Australia", "Japan"],
    "GrafiMedics BeLux": ["Belgium", "Luxembourg"],
    "BinSina Pharmacy": ["United Arab Emirates", "Oman"],
}


def canon(token: str) -> str | None:
    t = (token or "").strip()
    if not t:
        return None
    low = t.lower().strip(" .")
    if low in DROP:
        return None
    if low in ALIAS:
        return ALIAS[low]
    m = re.fullmatch(r"[A-Za-z]{2,5}\((.+)\)", t)
    if m:
        return canon(m.group(1))
    if not re.search(r"[A-Za-z]", t):
        return None
    # normalize common full names already correct
    specials = {
        "united states": "United States",
        "united kingdom": "United Kingdom",
        "united arab emirates": "United Arab Emirates",
        "south africa": "South Africa",
        "south korea": "South Korea",
        "new zealand": "New Zealand",
        "hong kong": "Hong Kong",
        "saudi arabia": "Saudi Arabia",
        "dominican republic": "Dominican Republic",
        "el salvador": "El Salvador",
        "costa rica": "Costa Rica",
        "czech republic": "Czechia",
        "puerto rico": "Puerto Rico",
    }
    if low in specials:
        return specials[low]
    return t


def country_from_hq(hq: str) -> str | None:
    s = (hq or "").strip()
    if not s:
        return None
    if "," in s:
        parts = [p.strip() for p in s.split(",") if p.strip()]
        return canon(parts[-1])
    return canon(s)


def parse_list(raw: str) -> list[str]:
    s = (raw or "").strip()
    if not s or s.upper() in {"NPD", "N/A", "NA", "-"}:
        return []
    out: list[str] = []
    for ch in re.split(r"\s*;\s*", s):
        ch = ch.strip()
        if not ch:
            continue
        # region form whole chunk
        c = canon(ch)
        if c is None and "," in ch:
            # maybe City, Country inside OP by mistake — take last
            c = canon(ch.split(",")[-1].strip())
        if c:
            out.append(c)
    seen = set()
    uniq = []
    for c in out:
        if c not in seen:
            seen.add(c)
            uniq.append(c)
    return uniq


def union(a: list[str], b: list[str]) -> list[str]:
    seen = set()
    out = []
    for x in list(a) + list(b):
        c = canon(x) or x
        if c and c.lower() not in DROP and c not in seen:
            seen.add(c)
            out.append(c)
    return out


def fmt(countries: list[str]) -> str:
    return "; ".join(countries)


def apply_format(ws) -> None:
    headers = [c.value for c in ws[1]]
    widths = {
        "Company": 34,
        "Website": 36,
        "Founded": 10,
        "Headquarters": 28,
        "Continent / Geography": 42,
        "Operational Presence": 42,
        "Ownership": 18,
        "Employees": 12,
        "Core Categories": 28,
        "Specialty Focus": 36,
        "Key Brands Represented": 32,
        "Retail / E-commerce / Both": 16,
        "Distribution Type": 22,
        "Contact Person": 18,
        "Role": 16,
        "Email": 28,
        "LinkedIn": 28,
        "Office No.": 16,
        "Country Code": 10,
        "Region Code": 10,
        "Summary": 48,
    }
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="Calibri", size=10)
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="left")
    header_align = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(str(h or ""), 18)

    ws.row_dimensions[1].height = 32
    for i in range(1, len(headers) + 1):
        cell = ws.cell(1, i)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin

    geo_i = headers.index("Continent / Geography") + 1
    op_i = headers.index("Operational Presence") + 1
    for r in range(2, ws.max_row + 1):
        geo = str(ws.cell(r, geo_i).value or "")
        op = str(ws.cell(r, op_i).value or "")
        longest = max(len(geo), len(op), 20)
        lines = max(2, min(8, (longest // 42) + 1))
        ws.row_dimensions[r].height = 15 * lines
        for i in range(1, len(headers) + 1):
            cell = ws.cell(r, i)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def main() -> None:
    bak = sorted(ROOT.glob("*BEFORE_GEO_OP*.xlsx"))[-1]
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safety = ROOT / f"global_wearable_medical_devices_market_global_FINAL_BEFORE_RESTORE_GEO_{stamp}.xlsx"
    shutil.copy2(XLSX, safety)

    wb_cur = load_workbook(XLSX)
    ws = wb_cur.active
    wb_bak = load_workbook(bak)
    ws_b = wb_bak.active

    h = [c.value for c in ws[1]]
    hb = [c.value for c in ws_b[1]]
    ic = {x: i + 1 for i, x in enumerate(h) if x}
    icb = {x: i + 1 for i, x in enumerate(hb) if x}

    # index backup by company
    bak_rows = {}
    for r in range(2, ws_b.max_row + 1):
        co = str(ws_b.cell(r, icb["Company"]).value or "").strip()
        bak_rows[co] = {
            "geo": str(ws_b.cell(r, icb["Continent / Geography"]).value or "").strip(),
            "op": str(ws_b.cell(r, icb["Operational Presence"]).value or "").strip(),
            "website": str(ws_b.cell(r, icb["Website"]).value or "").strip(),
            "founded": str(ws_b.cell(r, icb["Founded"]).value or "").strip(),
            "hq": str(ws_b.cell(r, icb["Headquarters"]).value or "").strip(),
        }

    core_ok = 0
    core_bad = []
    restored = []

    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        website = str(ws.cell(r, ic["Website"]).value or "").strip()
        founded = str(ws.cell(r, ic["Founded"]).value or "").strip()
        hq = str(ws.cell(r, ic["Headquarters"]).value or "").strip()

        b = bak_rows.get(co)
        if not b:
            continue

        # verify core unchanged (and leave them alone)
        if website != b["website"] or founded != b["founded"] or hq != b["hq"]:
            core_bad.append(
                {
                    "company": co,
                    "cur": {"website": website, "founded": founded, "hq": hq},
                    "bak": {
                        "website": b["website"],
                        "founded": b["founded"],
                        "hq": b["hq"],
                    },
                }
            )
        else:
            core_ok += 1

        hq_country = country_from_hq(hq)
        old_geo = parse_list(b["geo"])
        old_op = parse_list(b["op"])

        # Continent / Geography: restored countries from old GEO + HQ country + verified physical extras
        geo = list(old_geo)
        if hq_country and hq_country not in geo:
            geo = [hq_country] + geo
        if co in PHYSICAL_EXTRA:
            geo = union(geo, PHYSICAL_EXTRA[co])
        # If GEO was only region form with one country, old_geo already parsed

        # Operational Presence: restore ALL prior OP countries + verified commercial extras
        op = list(old_op)
        if hq_country and not op:
            op = [hq_country]
        if hq_country and hq_country not in op:
            # keep prior list; ensure HQ country present
            op = [hq_country] + [c for c in op if c != hq_country]
        if co in COMMERCIAL_EXTRA:
            op = union(op, COMMERCIAL_EXTRA[co])

        geo = union(geo, [])
        op = union(op, [])
        new_geo = fmt(geo)
        new_op = fmt(op)

        ws.cell(r, ic["Continent / Geography"]).value = new_geo
        ws.cell(r, ic["Operational Presence"]).value = new_op
        restored.append(
            {
                "company": co,
                "bak_geo": b["geo"],
                "new_geo": new_geo,
                "bak_op": b["op"],
                "new_op": new_op,
            }
        )

    if ws.title.strip().lower() == "smartwatch companies":
        ws.title = "Wearable Medical Companies"

    apply_format(ws)
    wb_cur.save(XLSX)

    multi_op = sum(
        1
        for r in range(2, ws.max_row + 1)
        if ";" in str(ws.cell(r, ic["Operational Presence"]).value or "")
    )
    multi_geo = sum(
        1
        for r in range(2, ws.max_row + 1)
        if ";" in str(ws.cell(r, ic["Continent / Geography"]).value or "")
    )

    audit = {
        "backup_source": bak.name,
        "safety_copy": safety.name,
        "core_unchanged_rows": core_ok,
        "core_mismatch_rows": core_bad,
        "restored_rows": len(restored),
        "multi_country_geo": multi_geo,
        "multi_country_op": multi_op,
        "sample": restored[:30],
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print("source", bak.name)
    print("core_ok", core_ok, "core_bad", len(core_bad))
    print("restored", len(restored), "multi_geo", multi_geo, "multi_op", multi_op)
    for s in restored:
        if s["company"] in {
            "Al-Ahli Medical",
            "Alliance Healthcare",
            "Transmedic",
            "Watsons",
            "Giant Eagle",
            "Henry Schein",
            "Al-Bahar Medical",
            "Zuellig Pharma",
        }:
            print(s["company"])
            print("  GEO", s["bak_geo"], "=>", s["new_geo"])
            print("  OP ", s["bak_op"], "=>", s["new_op"])


if __name__ == "__main__":
    main()
