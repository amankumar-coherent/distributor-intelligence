"""Post-fix audit: blanks and remaining junk in 4 columns."""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"C:\Users\aman.kumar\Desktop\coherent\output\run4")
COLS = ["Core Categories", "Specialty Focus", "Price Segment", "Key Brands Represented"]


def find_header(ws):
    for r in range(1, 8):
        vals = [c.value for c in ws[r]]
        if any(v and str(v).strip().lower() == "company" for v in vals):
            return r, vals
    return 1, [c.value for c in ws[1]]


def main() -> None:
    for mid in [
        "global_smartwatch_market_global",
        "global_green_chemical_market_global",
        "global_firewall_market_global",
    ]:
        short = re.sub(r"_global$", "", mid)
        xp = ROOT / mid / f"{short}_FINAL.xlsx"
        wb = load_workbook(xp, read_only=True, data_only=True)
        ws = wb.active
        hr, headers = find_header(ws)
        hmap = {str(h).strip(): i for i, h in enumerate(headers) if h}
        ci = hmap["Company"]
        print(f"\n=== {mid} FINAL ===")
        for col in COLS:
            idx = hmap[col]
            blank = 0
            junk = []
            prices: Counter[str] = Counter()
            samples = []
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                if not row or not row[ci]:
                    continue
                s = "" if row[idx] is None else str(row[idx]).strip()
                if col == "Price Segment":
                    prices[s or "(blank)"] += 1
                if not s:
                    blank += 1
                    junk.append((str(row[ci])[:35], "(blank)"))
                elif col == "Specialty Focus" and s.lower() in ("service provider", "industry manufacturer"):
                    junk.append((str(row[ci])[:35], s))
                elif col == "Key Brands Represented" and (
                    "professional source" in s.lower()
                    or "private investors" in s.lower()
                    or s.lower().startswith("smartwatches;")
                    or s.lower().startswith("firewalls;")
                    or s.lower().startswith("specialty chemicals;")
                    or s.lower().startswith("bio-based")
                ):
                    junk.append((str(row[ci])[:35], s[:100]))
                if len(samples) < 4 and s:
                    samples.append((str(row[ci])[:30], s[:90]))
            print(f"  {col}: blank={blank} leftover_junk={len(junk)}")
            for a, b in junk[:8]:
                print(f"    {a} | {b}")
            for a, b in samples:
                print(f"    eg {a} | {b}")
            if col == "Price Segment":
                print(f"    dist={dict(prices)}")
        # B&H check
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if row and row[ci] and "b&h" in str(row[ci]).lower():
                print("  B&H brands:", row[hmap["Key Brands Represented"]])
                print("  B&H specialty:", row[hmap["Specialty Focus"]])
                print("  B&H price:", row[hmap["Price Segment"]])
        wb.close()


if __name__ == "__main__":
    main()
