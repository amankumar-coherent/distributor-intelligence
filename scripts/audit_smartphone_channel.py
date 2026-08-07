#!/usr/bin/env python3
"""Audit Global Smartphone FINAL: keep only smartphone distributors/channel partners.

Uses live web search + DeepSeek (OPENAI_* → api.deepseek.com, model deepseek-chat).
Removes OEMs, off-market firms, and non-channel companies. Spot-checks key columns.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from dotenv import dotenv_values, load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Prefer Cmi DeepSeek keys
load_dotenv(ROOT / ".env")
_cmi = Path(r"C:\Users\aman.kumar\Desktop\Cmi\agents\vendor-intel\.env")
if _cmi.exists():
    for k, v in (dotenv_values(_cmi) or {}).items():
        if v and k.startswith(("OPENAI_", "LLM_", "COMPILER_", "CLASSIFIER_", "MARKET_MAP_")):
            os.environ[k] = v

FINAL = (
    ROOT
    / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx"
)
SHEET = "Companies"
MARKET = "Global Smartphone Market"
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip() or "deepseek-chat"
BASE = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
CONCURRENT = int(os.getenv("SMARTPHONE_AUDIT_CONCURRENT") or "4")
BATCH_LIMIT = int(os.getenv("SMARTPHONE_AUDIT_LIMIT") or "0")  # 0 = all

OEM_NAMES = {
    "apple", "samsung", "huawei", "xiaomi", "oppo", "vivo", "oneplus", "google",
    "motorola", "nokia", "sony", "realme", "honor", "tecno", "infinix", "nothing",
    "lg electronics", "htc", "zte", "lenovo", "asus", "sharp", "meizu", "fairphone",
}
# Platforms / non-channel (not a distributor/reseller of phones)
HARD_DROP = {
    "alibaba", "alibaba group", "alibaba.com",
    "wikipedia", "statista", "grand view research",
    "allied electronics & automation", "allied electronics",
    "acce global software pvt ltd", "acce global software",
}

# Well-known smartphone / CE channel firms — keep even if web is noisy
KNOWN_KEEP = {
    "also holding", "also", "also nestle",
    "ingram micro", "td synnex", "synnex", "tech data",
    "d&h distributing", "d&h", "westcon", "westcon-comstor",
    "arrow electronics", "avnet",
    "elko group", "elko",
    "amazon", "best buy", "currys", "media markt", "saturn",
    "fnac", "darty", "reliance digital", "croma", "vijay sales",
}

DIST_RE = re.compile(
    r"\b(distributor|distribution|wholesaler|wholesale|reseller|vad|var|channel\s*partner|"
    r"authorized\s*(dealer|reseller|partner)|importer|retailer|retail|dealer|"
    r"integrator|carrier\s*partner|mobile\s*operator|telecom\s*retail)\b",
    re.I,
)
PHONE_RE = re.compile(
    r"\b(smartphone|smart\s*phone|mobile\s*phone|handset|cellular|android\s*phone|"
    r"iphone|feature\s*phone|mobile\s*device|cell\s*phone|telecom|mobile\s*retail|"
    r"consumer\s*electronics|mobile\s*phones?)\b",
    re.I,
)
OFF_RE = re.compile(
    r"\b(industrial\s*automation|plc\s*distributor|electronic\s*components?\s*only|"
    r"saas\s*only|crm\s*software|erp\s*software|avocado|grocery|firewall\s*appliance|"
    r"chemical\s*distributor|market\s*research\s*report)\b",
    re.I,
)

NPD = {
    "", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown",
    "tbd", "todo", "not applicable", "not apllicable",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _load_rows(path: Path) -> tuple[list[str], list[dict[str, str]], int]:
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers: list[str] = []
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            header_idx = i
            headers = vals
            break
    if header_idx is None:
        raise RuntimeError("No Company header")
    data: list[dict[str, str]] = []
    for r in rows[header_idx + 1 :]:
        d = {
            headers[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
            for j in range(len(headers))
            if headers[j]
        }
        if d.get("Company"):
            data.append(d)
    return [h for h in headers if h], data, header_idx


def _hard_rules(row: dict[str, str]) -> tuple[str, str] | None:
    """Return (action, reason) if deterministic; else None."""
    name = row.get("Company") or ""
    nl = _norm(name)
    for keep in KNOWN_KEEP:
        if nl == keep or nl.startswith(keep + " "):
            return "keep", f"known_channel:{keep}"
    for oem in OEM_NAMES:
        if nl == oem or nl.startswith(oem + " ") or nl == oem.replace(" ", ""):
            if not DIST_RE.search(name):
                return "remove", f"oem_brand:{oem}"
    for drop in HARD_DROP:
        if nl == drop or nl.startswith(drop):
            return "remove", f"hard_drop:{drop}"
    blob = " ".join(
        [
            name,
            row.get("Business Type") or "",
            row.get("Specialty Focus") or "",
            row.get("Core Categories") or "",
            row.get("Summary") or "",
            row.get("Key Brands Represented") or "",
        ]
    )
    if OFF_RE.search(blob) and not PHONE_RE.search(blob):
        return "remove", "off_market_text"
    return None


def _row_phone_channel_text(row: dict[str, str]) -> bool:
    blob = " ".join(
        [
            row.get("Company") or "",
            row.get("Business Type") or "",
            row.get("Specialty Focus") or "",
            row.get("Core Categories") or "",
            row.get("Summary") or "",
            row.get("Key Brands Represented") or "",
        ]
    )
    return bool(PHONE_RE.search(blob) and DIST_RE.search(blob))


async def _web_evidence(name: str, website: str) -> str:
    queries = [
        f'"{name}" smartphone OR "mobile phone" OR handset distributor OR wholesaler OR reseller',
        f'"{name}" authorized distributor mobile phones',
    ]
    if website:
        host = re.sub(r"^https?://", "", website.strip(), flags=re.I).split("/")[0]
        if host:
            queries.append(f"{host} smartphones OR mobile phones")

    parts: list[str] = []

    # 1) Prefer fast DDGS
    try:
        from ddgs import DDGS

        def _ddgs_one(q: str) -> list[str]:
            lines: list[str] = []
            try:
                hits = list(DDGS().text(q, max_results=5)) or []
            except Exception:
                hits = []
            for h in hits:
                lines.append(
                    f"- {h.get('title', '')}: {str(h.get('body', ''))[:220]} ({h.get('href', '')})"
                )
            return lines

        for q in queries[:2]:
            lines = await asyncio.wait_for(asyncio.to_thread(_ddgs_one, q), timeout=25)
            if lines:
                parts.append(f"Query: {q}\n" + "\n".join(lines))
    except Exception as err:
        parts.append(f"(ddgs note: {type(err).__name__}: {err})")

    if any(p.startswith("Query:") for p in parts):
        return "\n\n".join(parts)[:5000]

    # 2) FreeSearchRouter fallback
    try:
        from vendor_intel.clients.search_router import FreeSearchRouter
        from vendor_intel.config import Settings
        from vendor_intel.enrichment.gap_fill.sources import _search_snippets

        router = FreeSearchRouter(Settings.load())
        for q in queries[:2]:
            try:
                snip = await asyncio.wait_for(
                    _search_snippets(router, q, max_hits=5), timeout=30
                )
            except Exception:
                snip = ""
            if snip:
                parts.append(f"Query: {q}\n{snip}")
    except Exception as err:
        parts.append(f"(router note: {type(err).__name__}: {err})")

    return "\n\n".join(parts)[:5000] or "(no web hits)"


def _client():
    from openai import OpenAI

    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    if not key:
        raise RuntimeError("OPENAI_API_KEY missing")
    return OpenAI(api_key=key, base_url=BASE)


def _llm_verdict(client: Any, row: dict[str, str], evidence: str) -> dict[str, Any]:
    name = row.get("Company") or ""
    profile = (
        f"Company: {name}\n"
        f"Website: {row.get('Website') or ''}\n"
        f"Business Type: {row.get('Business Type') or ''}\n"
        f"HQ: {row.get('Headquarters') or ''}\n"
        f"Founded: {row.get('Founded') or ''}\n"
        f"Core Categories: {row.get('Core Categories') or ''}\n"
        f"Specialty Focus: {row.get('Specialty Focus') or ''}\n"
        f"Key Brands: {row.get('Key Brands Represented') or ''}\n"
        f"Summary: {(row.get('Summary') or '')[:400]}\n"
    )
    system = (
        "You audit vendor-intel landscape rows for the Global Smartphone Market. "
        "KEEP only companies that are distributors, wholesalers, retailers, dealers, "
        "resellers, authorized channel partners, carriers/retail arms, or importers "
        "of smartphones / mobile handsets (or broad CE distributors that clearly sell phones). "
        "REMOVE: OEMs/brands that only make phones (Apple/Samsung/etc), pure software/SaaS, "
        "industrial automation component distributors, marketplaces that are not a selling "
        "channel entity, research firms, and any firm with no smartphone/mobile phone channel role. "
        "Also flag field errors. Return JSON only."
    )
    user = f"""MARKET: {MARKET}

STORED ROW:
{profile}

WEB EVIDENCE:
{evidence}

Return JSON:
{{
  "action": "keep" | "remove",
  "is_smartphone_related": true/false,
  "is_channel_partner": true/false,
  "reason": "short reason",
  "field_fixes": {{
     "Founded": "value or empty if OK",
     "Headquarters": "",
     "Website": "",
     "Business Type": "",
     "Key Brands Represented": "",
     "Core Categories": ""
  }},
  "confidence": "high"|"medium"|"low"
}}
Rules for field_fixes: only suggest a replacement when evidence clearly contradicts the stored value.
Use empty string when the stored value is fine or unknown. Never invent emails/phones.
If action=remove, field_fixes may be {{}}.
"""
    resp = client.chat.completions.create(
        model=MODEL,
        temperature=0.1,
        max_tokens=700,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
    )
    raw = resp.choices[0].message.content or "{}"
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{[\s\S]*\}", raw)
        data = json.loads(m.group(0)) if m else {"action": "remove", "reason": "bad_json"}
    if not isinstance(data, dict):
        data = {"action": "remove", "reason": "bad_json_type"}
    data.setdefault("action", "remove")
    data.setdefault("reason", "")
    data.setdefault("field_fixes", {})
    data["usage"] = {
        "in": getattr(resp.usage, "prompt_tokens", 0) or 0,
        "out": getattr(resp.usage, "completion_tokens", 0) or 0,
    }
    return data


def _apply_fixes(row: dict[str, str], fixes: dict[str, Any]) -> dict[str, str]:
    out = dict(row)
    if not isinstance(fixes, dict):
        return out
    for k, v in fixes.items():
        if k not in out:
            continue
        val = str(v or "").strip()
        if not val or val.lower() in NPD:
            continue
        # never invent contacts here
        if k in ("Email", "Phone / WhatsApp", "Contact Person", "Role"):
            continue
        old = str(out.get(k) or "").strip()
        if _norm(old) != _norm(val):
            out[k] = val
    # smartphone landscape: grocery cols N/A
    out["Specialty Grocers"] = "Not Applicable"
    out["Foodservice / HoReCa"] = "Not Applicable"
    for c in ("Exclusivity", "Partnership Duration"):
        if _norm(out.get(c) or "") in NPD or not out.get(c):
            out[c] = "Not Publicly Disclosed"
    return out


def _write_xlsx(
    path: Path,
    headers: list[str],
    rows: list[dict[str, str]],
    *,
    title: str,
    audit_rows: list[tuple[str, str]],
) -> None:
    wb = load_workbook(path)
    # backup first handled by caller
    ws = wb[SHEET]
    # clear data rows (keep title + header)
    # find header row
    header_row = 2
    for r in range(1, 6):
        if ws.cell(r, 1).value == "Company":
            header_row = r
            break
    # delete old data
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    # ensure headers match
    for j, h in enumerate(headers, 1):
        ws.cell(header_row, j).value = h
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    for i, row in enumerate(rows):
        r = header_row + 1 + i
        for j, h in enumerate(headers, 1):
            cell = ws.cell(r, j, row.get(h, ""))
            cell.alignment = wrap
            cell.border = thin
    ws.cell(1, 1).value = title

    # Audit sheet
    name = "Channel Audit"
    if name in wb.sheetnames:
        del wb[name]
    aus = wb.create_sheet(name)
    aus.append(["field", "value"])
    for k, v in audit_rows:
        aus.append([k, v])

    wb.save(path)


async def _audit_one(
    sem: asyncio.Semaphore,
    client: Any,
    row: dict[str, str],
    idx: int,
    total: int,
) -> dict[str, Any]:
    name = row.get("Company") or ""
    try:
        hard = _hard_rules(row)
        if hard:
            action, reason = hard
            print(f"  [{idx}/{total}] HARD {action}: {name} — {reason}", flush=True)
            return {
                "company": name,
                "action": action,
                "reason": reason,
                "is_smartphone_related": False,
                "is_channel_partner": False,
                "field_fixes": {},
                "confidence": "high",
                "method": "hard_rule",
            }

        async with sem:
            print(f"  [{idx}/{total}] web+DeepSeek: {name}", flush=True)
            evidence = await _web_evidence(name, row.get("Website") or "")
            verdict = await asyncio.to_thread(_llm_verdict, client, row, evidence)

        action = str(verdict.get("action") or "remove").lower().strip()
        if action not in ("keep", "remove"):
            action = "remove"
        verdict["evidence_chars"] = len(evidence or "")
        if action == "keep":
            if (
                verdict.get("is_smartphone_related") is False
                or verdict.get("is_channel_partner") is False
            ):
                action = "remove"
                verdict["reason"] = (
                    str(verdict.get("reason") or "")
                    + " | forced_remove: not related or not channel"
                ).strip(" |")
        elif action == "remove" and _row_phone_channel_text(row):
            # Don't drop clear phone-channel rows when web evidence was thin/noisy
            conf = str(verdict.get("confidence") or "").lower()
            reason_l = str(verdict.get("reason") or "").lower()
            weak_web = (
                conf in ("low", "medium")
                or "no evidence" in reason_l
                or "irrelevant" in reason_l
                or int(verdict.get("evidence_chars") or 0) < 200
            )
            if weak_web:
                action = "keep"
                verdict["reason"] = (
                    str(verdict.get("reason") or "")
                    + " | override_keep: phone+channel text with weak/noisy web"
                ).strip(" |")
                verdict["confidence"] = "medium"

        print(
            f"      → {action}: {str(verdict.get('reason', ''))[:120]} "
            f"(conf={verdict.get('confidence')})",
            flush=True,
        )
        verdict["company"] = name
        verdict["action"] = action
        verdict["method"] = "web+deepseek"
        return verdict
    except Exception as err:  # noqa: BLE001
        print(f"  [{idx}/{total}] ERROR {name}: {type(err).__name__}: {err}", flush=True)
        return {
            "company": name,
            "action": "remove",
            "reason": f"audit_error:{type(err).__name__}:{err}",
            "is_smartphone_related": False,
            "is_channel_partner": False,
            "field_fixes": {},
            "confidence": "low",
            "method": "error",
        }


async def main() -> int:
    if not FINAL.exists():
        print(f"ERROR: missing {FINAL}", file=sys.stderr)
        return 1

    headers, rows, _ = _load_rows(FINAL)
    if BATCH_LIMIT > 0:
        rows = rows[:BATCH_LIMIT]
    print(f"=== Smartphone channel audit ===", flush=True)
    print(f"  file : {FINAL}", flush=True)
    print(f"  rows : {len(rows)}", flush=True)
    print(f"  model: {MODEL}", flush=True)
    print(f"  base : {BASE}", flush=True)
    print(f"  conc : {CONCURRENT}", flush=True)

    client = _client()
    ping = client.chat.completions.create(
        model=MODEL,
        max_tokens=20,
        messages=[{"role": "user", "content": 'Reply JSON {"ok":true}'}],
        response_format={"type": "json_object"},
    )
    print(f"  deepseek ping: {(ping.choices[0].message.content or '')[:80]}", flush=True)

    out_dir = FINAL.parent
    ckpt_path = out_dir / "smartphone_channel_audit_checkpoint.json"
    prior: dict[str, Any] = {}
    if ckpt_path.exists() and not (os.getenv("SMARTPHONE_AUDIT_FRESH") or "").strip():
        try:
            prior = json.loads(ckpt_path.read_text(encoding="utf-8"))
        except Exception:
            prior = {}
    done_map: dict[str, dict[str, Any]] = {
        _norm(str(v.get("company") or "")): v
        for v in (prior.get("verdicts") or [])
        if isinstance(v, dict) and v.get("company")
    }
    if done_map:
        print(f"  resume: {len(done_map)} prior verdicts loaded", flush=True)

    sem = asyncio.Semaphore(CONCURRENT)
    verdicts: list[dict[str, Any]] = []
    chunk = 12
    pending_rows = [r for r in rows if _norm(r.get("Company") or "") not in done_map]
    # keep prior for already done, in original order at end
    for i in range(0, len(pending_rows), chunk):
        batch = pending_rows[i : i + chunk]
        print(
            f"--- batch {i // chunk + 1}: {len(batch)} companies "
            f"(remaining ~{len(pending_rows) - i}) ---",
            flush=True,
        )
        tasks = []
        for j, row in enumerate(batch):
            # global index approx
            idx = len(done_map) + i + j + 1
            tasks.append(_audit_one(sem, client, row, idx, len(rows)))
        batch_verdicts = await asyncio.gather(*tasks)
        for v in batch_verdicts:
            done_map[_norm(str(v.get("company") or ""))] = v
        # checkpoint after each chunk
        ckpt_path.write_text(
            json.dumps(
                {
                    "verdicts": list(done_map.values()),
                    "done": len(done_map),
                    "total": len(rows),
                },
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        print(f"  checkpoint saved ({len(done_map)}/{len(rows)})", flush=True)

    # preserve input order
    verdicts = [
        done_map.get(_norm(r.get("Company") or ""), {
            "company": r.get("Company"),
            "action": "remove",
            "reason": "missing_verdict",
            "method": "error",
        })
        for r in rows
    ]

    by_name = {_norm(v["company"]): v for v in verdicts}
    kept: list[dict[str, str]] = []
    removed: list[dict[str, Any]] = []
    fixes_applied = 0
    for row in rows:
        v = by_name.get(_norm(row.get("Company") or ""), {})
        if str(v.get("action")) == "keep":
            fixed = _apply_fixes(row, v.get("field_fixes") or {})
            if fixed != row:
                fixes_applied += 1
            src = fixed.get("Data Sources") or ""
            if "deepseek_channel_audit" not in src:
                fixed["Data Sources"] = (src + "; deepseek_channel_audit").strip("; ")
            kept.append(fixed)
        else:
            removed.append(
                {
                    "company": row.get("Company"),
                    "reason": v.get("reason"),
                    "method": v.get("method"),
                    "confidence": v.get("confidence"),
                }
            )

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    dry = (os.getenv("SMARTPHONE_AUDIT_DRY") or "").strip().lower() in ("1", "true", "yes")
    if dry or BATCH_LIMIT > 0:
        print(
            f"  DRY/LIMIT mode — not writing FINAL "
            f"(kept={len(kept)} removed={len(removed)} fixes={fixes_applied})",
            flush=True,
        )
        audit_path = out_dir / f"smartphone_channel_audit_DRY_{ts}.json"
        audit_path.write_text(
            json.dumps(
                {
                    "ts": ts,
                    "dry": True,
                    "limit": BATCH_LIMIT,
                    "kept": len(kept),
                    "removed": len(removed),
                    "removed_companies": removed,
                    "verdicts": verdicts,
                },
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        print(f"  dry audit: {audit_path}", flush=True)
        for r in removed[:20]:
            print(f"    - {r['company']}: {r['reason']}", flush=True)
        return 0

    bak = FINAL.with_suffix(FINAL.suffix + f".bak_channel_audit_{ts}")
    shutil.copy2(FINAL, bak)
    print(f"  backup: {bak.name}", flush=True)

    audit_path = out_dir / f"smartphone_channel_audit_{ts}.json"
    audit_payload = {
        "ts": ts,
        "market": MARKET,
        "model": MODEL,
        "base_url": BASE,
        "input_rows": len(rows),
        "kept": len(kept),
        "removed": len(removed),
        "fixes_applied_rows": fixes_applied,
        "removed_companies": removed,
        "verdicts": verdicts,
    }
    audit_path.write_text(json.dumps(audit_payload, indent=2, ensure_ascii=False), encoding="utf-8")

    title = (
        f"Global Smartphone Market — channel partners only "
        f"(DeepSeek audit {ts}; kept {len(kept)}/{len(rows)})"
    )
    _write_xlsx(
        FINAL,
        headers,
        kept,
        title=title,
        audit_rows=[
            ("ts", ts),
            ("model", MODEL),
            ("base_url", BASE),
            ("input_rows", str(len(rows))),
            ("kept", str(len(kept))),
            ("removed", str(len(removed))),
            ("fixes_applied_rows", str(fixes_applied)),
            ("audit_json", str(audit_path)),
            ("backup", str(bak)),
        ],
    )

    print("=== DONE ===", flush=True)
    print(f"  kept   : {len(kept)}", flush=True)
    print(f"  removed: {len(removed)}", flush=True)
    print(f"  fixes  : {fixes_applied} rows updated", flush=True)
    print(f"  audit  : {audit_path}", flush=True)
    if removed:
        print("  sample removals:", flush=True)
        for r in removed[:25]:
            print(f"    - {r['company']}: {r['reason']}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
