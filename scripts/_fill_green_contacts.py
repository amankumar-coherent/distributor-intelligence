# -*- coding: utf-8 -*-
"""Agent web-search fills for green chemical empty Contact Person / Role / LinkedIn."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

NPD = "Not publicly disclosed"
PATH = Path(
    r"D:\vendor-intel\output\run4\global_green_chemical_market_global"
    r"\global_green_chemical_market.xlsx"
)

EMPTY = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}

# Only write if Contact Person currently empty. Sales/procurement titles only.
CONTACTS = {
    "Barentz": {
        "person": "Christian Barker",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/christian-barker",
    },
    "Brenntag": {
        "person": "Jon Birrell",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/jon-birrell-05378834",
    },
    "Biesterfeld": {
        "person": "Simon Cox",
        "role": "Head of Sales Northern Europe",
        "li": "https://www.linkedin.com/in/simon-cox-6693391b8",
    },
    "Caldic": {
        "person": "Terri-Anne Eustace",
        "role": "Commercial Director, Ingredients",
        "li": "https://www.linkedin.com/in/terri-anne-eustace-ab7b0336",
    },
    "Bax Chemicals": {
        "person": "Jeroen Knijn",
        "role": "Sales Manager",
        "li": "https://www.linkedin.com/in/jeroen-knijn-616a4597",
    },
    "Wilmar International (Oleochemicals)": {
        "person": "Tom Ozzella",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/tomozzella",
    },
    "Algol Chemicals": {
        "person": "Jari Kerbs",
        "role": "Sales Director Plastics",
        "li": "https://www.linkedin.com/in/jari-kerbs-0a634b6a",
    },
    "Bang & Bonsomer": {
        "person": "Antti Mehtomaa",
        "role": "Head of Sales Pro-Glue",
        "li": "https://www.linkedin.com/in/anttimehtomaa",
    },
    "Behn Meyer Group": {
        "person": "Mai Huong Nguyen",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/maihuongnguyenbehnmeyer",
    },
    "Aug. Hedinger GmbH & Co. KG": {
        "person": "Serkan Er",
        "role": "Sales Manager",
        "li": "https://www.linkedin.com/in/serkan-er-9a7671136",
    },
    "C.H. Erbslöh GmbH & Co. KG": {
        "person": "Stefan Limbach",
        "role": "Sales Manager",
        "li": "https://www.linkedin.com/in/stefan-limbach-48890962",
    },
    "Cornelius Group": {
        "person": "Martin Chambers",
        "role": "Head of UK Sales",
        "li": "https://www.linkedin.com/in/martin-chambers-04351a28",
    },
    "Connect Chemicals": {
        "person": "Elaine Zhu",
        "role": "Director Sales Market",
        "li": "https://www.linkedin.com/in/elaine-zhu-21a1836a",
    },
    "Azelis": {
        "person": "Marnik Tack",
        "role": "Chief Commercial Officer EMEA",
        "li": "https://www.linkedin.com/in/marnik-tack-620aa13",
    },
    "Univar Solutions": {
        "person": "Brian Jurcak",
        "role": "Chief Commercial and Operations Officer, Ingredients and Specialties",
        "li": "https://www.linkedin.com/in/brian-jurcak-58b61b19",
    },
    "Nordmann": {
        "person": "Paula Rodrigues",
        "role": "Global Sales Director Life Sciences",
        "li": "https://www.linkedin.com/in/dr-paula-rodrigues-225b69a1",
    },
    "Safic-Alcan": {
        "person": "Andreas Havsengen",
        "role": "Head of Sales and Business Development",
        "li": "https://www.linkedin.com/in/andreas-havsengen-56624533",
    },
    "Distrupol": {
        "person": "James Stanton",
        "role": "Commercial and Operations Director",
        "li": "https://www.linkedin.com/in/jmgstanton",
    },
    "Faravelli Group": {
        "person": "Kristyna Miratska",
        "role": "Sales Director Central Europe",
        "li": "https://www.linkedin.com/in/kristyna-miratska-b8433376",
    },
    "HELM AG": {
        "person": "Chris Ibekwe",
        "role": "Head of Sales and Product Management",
        "li": "https://www.linkedin.com/in/chrisibekwe",
    },
    "OQEMA Group": {
        "person": "Peter Dodd",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/peter-dodd-87b3a68b",
    },
    "Eigenmann & Veronelli": {
        "person": "Marco Faccincani",
        "role": "Sales Manager Human Nutrition",
        "li": "https://www.linkedin.com/in/marco-faccincani-76b86813",
    },
    "Stockmeier Holding GmbH": {
        "person": "Pierre Riviere",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/pierre-riviere-38b405210",
    },
    "Grolman Group": {
        "person": "Lars Torp",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/lars-torp-0020b4b",
    },
    "Quimidroga": {
        "person": "Emilio Dalmau",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/emilio-dalmau-65846887",
    },
    "Comercial Quimica Masso": {
        "person": "Josep Sues",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/josep-sues-sues-57920b48",
    },
    "Lehmann & Voss": {
        "person": "Harald Hunke",
        "role": "Commercial Director Europe",
        "li": "https://www.linkedin.com/in/dr-harald-hunke-42941916",
    },
    "Krahn Chemie": {
        "person": "Jan Urbas",
        "role": "Sales Director, Coatings and Construction Chemicals",
        "li": "https://www.linkedin.com/in/jan-urba%C5%9B-4ba921a3",
    },
    "Manuchar": {
        "person": "Sabri Tektas",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/sabri-tektas-26116747",
    },
    "Telko": {
        "person": "Ashish Maharishi",
        "role": "Sales Head, India",
        "li": "https://www.linkedin.com/in/ashish-maharishi-524a2717",
    },
    "Resinex (Ravago)": {
        "person": "Iulian Stancu",
        "role": "Sales Manager",
        "li": "https://www.linkedin.com/in/iulian-stancu-43084574",
    },
    "Palmer Holland": {
        "person": "Alexander Snyder",
        "role": "National Sales Director",
        "li": "https://www.linkedin.com/in/alexander-snyder-86a2b61b",
    },
    "Harcros Chemicals": {
        "person": "Jared Reigle",
        "role": "Vice President, Global Sales and Strategy",
        "li": "https://www.linkedin.com/in/jaredreigle",
    },
    "Hawkins, Inc.": {
        "person": "Jeremy Adler",
        "role": "Director of Sales and Business Development",
        "li": "https://www.linkedin.com/in/jeremyadler",
    },
    "EMCO Chemical Distributors": {
        "person": "Jake Knolmayer",
        "role": "Director of Sales and Marketing",
        "li": "https://www.linkedin.com/in/jake-knolmayer-7818674",
    },
    "Ellsworth Adhesives": {
        "person": "Jay Richardson",
        "role": "Vice President Sales, Global",
        "li": "https://www.linkedin.com/in/jay-richardson-6444413",
    },
    "Nexeo Plastics": {
        "person": "Jay Simmons",
        "role": "Regional Sales Director",
        "li": "https://www.linkedin.com/in/jaysimmons",
    },
    "BTC Europe GmbH": {
        "person": "Thiebaut Wenzler",
        "role": "Head of Essential Chemicals and Digital Sales",
        "li": "https://www.linkedin.com/in/thiebaut-wenzler-a8b153165",
    },
    "Bodo Moller Chemie": {
        "person": "Andreas Moser",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/andreas-moser-3ab70915",
    },
    "IMCD": {
        "person": "Andreas Igerl",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/andreas-christopher-igerl-91735931",
    },
    "Pochteca": {
        "person": "Alberto Romero",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/alberto-romero-a0574911b",
    },
    "Quimica Anastacio": {
        "person": "Henrique Rios Barbosa",
        "role": "Head of Sales Pharma LATAM",
        "li": "https://www.linkedin.com/in/henrique-rios-barbosa-2640a915",
    },
    "Sea-Land Chemical Company": {
        "person": "John Kivi",
        "role": "Senior Vice President of Commercial Excellence",
        "li": "https://www.linkedin.com/in/john-kivi-3074b78",
    },
    "Shrieve Chemical Company": {
        "person": "David Colquitt",
        "role": "Director of Sales",
        "li": "https://www.linkedin.com/in/david-colquitt-306b8831",
    },
    "Chemicals United B.V.": {
        "person": "Felix Adam",
        "role": "Director of Sales & Innovation",
        "li": "https://www.linkedin.com/in/felix-adam-62737339",
    },
    "Eurosyn": {
        "person": "Giovanni Bonanno",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/giovanni-bonanno-13891aa",
    },
    "PVS Chemicals": {
        "person": "Garrett Speck",
        "role": "Director of Sales and Marketing",
        "li": "https://www.linkedin.com/in/garrett-speck-657480aa",
    },
    "Globe Chemicals": {
        "person": "Osvaldo Otmar Redel",
        "role": "Sales Manager Brazil",
        "li": "https://www.linkedin.com/in/osvaldo-otmar-redel-2b397539",
    },
    "Seeler Industries": {
        "person": "Joe McCue",
        "role": "Vice President of Sales",
        "li": "https://www.linkedin.com/in/joe-mccue-aba9564b",
    },
    "Keyser & Mackay": {
        "person": "Charles Lynch",
        "role": "UK Sales Director",
        "li": "https://www.linkedin.com/in/charles-lynch-a80465234",
    },
    "K.D. Feddersen": {
        "person": "Guillaume Fromont",
        "role": "Sales Director France & Maghreb",
        "li": "https://www.linkedin.com/in/guillaume-fromont-848830b4",
    },
    "Interpolimeri S.p.A.": {
        "person": "Alessandro Castiglioni",
        "role": "EU Commercial Director",
        "li": "https://www.linkedin.com/in/alessandro-castiglioni-a6241818",
    },
    "Grupo ADI": {
        "person": "Israel Diaz Velasco",
        "role": "Director Comercial EMEA",
        "li": "https://www.linkedin.com/in/israeldvelasco",
    },
    "Grupo ADI / ADI Química": {
        "person": "Israel Diaz Velasco",
        "role": "Director Comercial EMEA",
        "li": "https://www.linkedin.com/in/israeldvelasco",
    },
    "Laborspirit": {
        "person": "Teresa Reis",
        "role": "Purchasing",
        "li": "https://www.linkedin.com/in/teresa-reis-93398814",
    },
    "Donauchem GmbH": {
        "person": "Michael Hoefer",
        "role": "Head of Sales",
        "li": "https://www.linkedin.com/in/michael-h%C3%B6fer-0675801bb",
    },
    "ELTON Group": {
        "person": "Markos Roussos",
        "role": "Group Head Sales & Strategy",
        "li": "https://www.linkedin.com/in/markosroussos",
    },
    "Dolder Group": {
        "person": "Antonella Ramaioli",
        "role": "Commercial Director Italy",
        "li": "https://www.linkedin.com/in/antonella-ramaioli-b340627",
    },
    "Comindex S.A.": {
        "person": "Sonia Gil Castillo",
        "role": "Responsable Comercial",
        "li": "https://www.linkedin.com/in/sonia-gil-castillo-0a456a55",
    },
    "Febo Group": {
        "person": "Matteo Canobbio",
        "role": "Vice Commercial Director",
        "li": "https://www.linkedin.com/in/matteo-canobbio-925430138",
    },
    "Norkem": {
        "person": "Mark Gregory",
        "role": "Group Business Development Director",
        "li": "https://www.linkedin.com/in/mark-gregory-b09b3825",
    },
    "CG Chemikalien": {
        "person": "Marco Regel",
        "role": "Head of Sales & Business Development",
        "li": "https://www.linkedin.com/in/marco-regel-0664968b",
    },
    "Acme-Hardesty": {
        "person": "John Hamilton",
        "role": "National Sales Director, Industrial Specialties & PU CASE",
        "li": "https://www.linkedin.com/in/john-hamilton-56173273",
    },
    "Actylis": {
        "person": "Eric Escobar",
        "role": "Vice President Global Pharma Sales",
        "li": "https://www.linkedin.com/in/eric-escobar-0a19b914",
    },
    "Amco Polymers": {
        "person": "Eddie Oropeza",
        "role": "Regional Commercial Director",
        "li": "https://www.linkedin.com/in/eddie-oropeza-085a238",
    },
    "Bamberger Polymers": {
        "person": "Chris Hessenius",
        "role": "North American Sales Director",
        "li": "https://www.linkedin.com/in/chris-hessenius-76aa9a13",
    },
    "M. Holland Company": {
        "person": "Monica Christler",
        "role": "Executive Vice President, Commercial",
        "li": "https://www.linkedin.com/in/monica-christler-a65b4b29",
    },
    "Maroon Group": {
        "person": "Kathleen Eatough",
        "role": "Senior Vice President of Sales, HI&I",
        "li": "https://www.linkedin.com/in/kathleen-eatough-a7084042",
    },
    "Ravago Chemicals": {
        "person": "Eric Kostyszyn",
        "role": "Vice President of Sales & Industrial Markets",
        "li": "https://www.linkedin.com/in/eric-kostyszyn-7b879820",
    },
    "Charkit Chemical Company": {
        "person": "Panos Yannopoulos",
        "role": "Executive Vice President, Sales",
        "li": "https://www.linkedin.com/in/panos-yannopoulos-18b79931",
    },
    "Independent Chemical Corporation": {
        "person": "Marc Ullrich",
        "role": "Director of Sales & Procurement",
        "li": "https://www.linkedin.com/in/marc-ullrich-5194794b",
    },
    "Coyne Chemical": {
        "person": "John Drzewicki",
        "role": "Director of Sales and Product Development",
        "li": "https://www.linkedin.com/in/john-drzewicki-a92085178",
    },
    "Obegi Chemicals": {
        "person": "Tawfik Aboul-Magd",
        "role": "Commercial Development Director",
        "li": "https://www.linkedin.com/in/tawfik-aboul-magd-8a315052",
    },
    "Sabará Químicos e Ingredientes": {
        "person": "Jefferson Texeira",
        "role": "Diretor Comercial",
        "li": "https://www.linkedin.com/in/jefferson-texeira-76309327",
    },
    "Usiquímica do Brasil": {
        "person": "Osvane Lazarone Cerqueira",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/osvane-lazarone-cerqueira-07a53a14",
    },
    "Arkem Kimya": {
        "person": "Dennis Raschke",
        "role": "Commercial Director North-West Europe",
        "li": "https://www.linkedin.com/in/dennis-raschke",
    },
    "Tilley Distribution": {
        "person": "Jim Zabelny",
        "role": "Sales Director, Ingredients",
        "li": "https://www.linkedin.com/in/jim-zabelny",
    },
    "Wego Chemical Group": {
        "person": "Ronald van Strien",
        "role": "Sales Director Europe",
        "li": "https://www.linkedin.com/in/ronaldvanstrien",
    },
    "Tricon Energy": {
        "person": "Eduardo Girote",
        "role": "Commercial Director, Engineering Plastics",
        "li": "https://www.linkedin.com/in/eduardo-girote-0b29076",
    },
    "Atlantic Chemicals Trading": {
        "person": "Richard Sheehan",
        "role": "Sales Director, UK & Ireland",
        "li": "https://www.linkedin.com/in/richard-sheehan-8957b736",
    },
    "IMPAG Group": {
        "person": "Fabiana Grisendi",
        "role": "Sales Director Europe, Performance Chemicals",
        "li": "https://www.linkedin.com/in/fabiana-grisendi-mba-5112b624",
    },
    "Julius Hoesch GmbH & Co. KG": {
        "person": "Christian Irmen",
        "role": "Head of Sales & Purchasing, Industrial Chemicals",
        "li": "https://www.linkedin.com/in/irmen-christian",
    },
    "GB-Chemie": {
        "person": "Manuel Poser",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/manuel-poser-3bb956289",
    },
    "HSH Chemie": {
        "person": "Aniko Petrik",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/aniko-petrik-ab610621",
    },
    "KH Chemicals": {
        "person": "Maikel Bolt",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/maikel-bolt-a1506b1a8",
    },
    "Jungbunzlauer": {
        "person": "Franck Ueberschlag",
        "role": "Executive Vice President of Sales",
        "li": "https://www.linkedin.com/in/franck-ueberschlag-3a7344",
    },
    "Haeffner Group": {
        "person": "Florian Grandel",
        "role": "Head of Purchasing & Sales",
        "li": "https://www.linkedin.com/in/florian-grandel-306506127",
    },
    "Inventec Performance Chemicals": {
        "person": "Jonathan Cetier",
        "role": "EMEA Sales Director",
        "li": "https://www.linkedin.com/in/jonathan-cetier-40831192",
    },
    "Borregaard": {
        "person": "Anders Sjoede",
        "role": "Vice President Sales, Industrial",
        "li": "https://www.linkedin.com/in/anders-sj%C3%B6de-21252820",
    },
    "Corbion": {
        "person": "Ruud Peerbooms",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/ruud-peerbooms-44b47b",
    },
    "Attilio Carmagnani": {
        "person": "Alberto Fasoli",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/alberto-fasoli-b6629143",
    },
    "Barcelonesa Group": {
        "person": "Mathieu Morin",
        "role": "Commercial Director France",
        "li": "https://www.linkedin.com/in/mathieumorin1",
    },
    "Kolmar Group AG": {
        "person": "Owen Price-James",
        "role": "Director Global Business Development",
        "li": "https://www.linkedin.com/in/owenpricejames",
    },
    "Indena distribution partners": {
        "person": "Stefano Togni",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/stefanotogni",
    },
    "GFBiochemicals": {
        "person": "Vincent Enaux",
        "role": "Chief Commercial Officer",
        "li": "https://www.linkedin.com/in/vincent-enaux-",
    },
    "Enamine": {
        "person": "Andrey Tarnovskiy",
        "role": "Sales Director, Europe & UK",
        "li": "https://www.linkedin.com/in/andrey-tarnovskiy",
    },
    "LGC Standards": {
        "person": "Jane Firth",
        "role": "Commercial Director, LGC AXIO Proficiency Testing",
        "li": "https://www.linkedin.com/in/jane-firth-41300910",
    },
    "Merck KGaA / MilliporeSigma distribution": {
        "person": "Rajeev Nair",
        "role": "Chief Commercial Officer, Discovery Solutions",
        "li": "https://www.linkedin.com/in/rajeev-nair-533bb121",
    },
    "Equilex Chemicals": {
        "person": "Patrick Elie",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/eliepatrick",
    },
    "Daymsa": {
        "person": "Jose Carlos Amor",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/jos%C3%A9-carlos-amor-346835b6",
    },
    "National Chemical Company": {
        "person": "James Bliss",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/james-bliss-43096a11",
    },
    "LUMAR QUIMICA": {
        "person": "Alex Andres",
        "role": "Sales Director",
        "li": "https://www.linkedin.com/in/alex-andres-61958a6",
    },
    "Meade-King, Robinson & Co": {
        "person": "Anthony Hughes",
        "role": "Sales Director, Mekiro BV",
        "li": "https://www.linkedin.com/in/anthony-hughes-b8a25944",
    },
    "Nedstar": {
        "person": "Victor Zilio",
        "role": "Head of Sales and Operations",
        "li": "https://www.linkedin.com/in/victor-zilio-9b884523",
    },
    "Ilario Ormezzano SAI": {
        "person": "Riccardo Valle",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/riccardo-valle-46a82b81",
    },
    "Kadion": {
        "person": "Carlos Gonzalez",
        "role": "International Sales Director",
        "li": "https://www.linkedin.com/in/carlos-gonzalez-08a8b728",
    },
    "COUNT Energy Trading / COUNT Energy Distribution": {
        "person": "Ibrahim Farooqi",
        "role": "Commercial Director",
        "li": "https://www.linkedin.com/in/ibrahimfarooqi",
    },
    "Plastribution Ltd": {
        "person": "Dionne Wardle",
        "role": "Head of Sales",
        "li": "https://www.linkedin.com/in/dionne-wardle-37ab2a76",
    },
    "RAHN Group": {
        "person": "Sjoerd Romkes",
        "role": "Commercial Director EnergyCuring",
        "li": "https://www.linkedin.com/in/sromkes",
    },
    "Welding GmbH & Co. KG": {
        "person": "Oliver Sternhagen",
        "role": "Head of Sales Pharma API Europe",
        "li": "https://www.linkedin.com/in/oliver-sternhagen-3080657b",
    },
    "Van Iperen International": {
        "person": "Taeyoung Kwak",
        "role": "Regional Sales Director, APAC",
        "li": "https://www.linkedin.com/in/taeyoung-kwak-b4943376",
    },
}


def _headers(ws):
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            return r, {h: i + 1 for i, h in enumerate(vals) if h}
    raise RuntimeError("headers not found")


def _empty(v):
    return (str(v or "").strip().lower() in EMPTY) or (v is None)


def save_atomic(wb, path: Path):
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def main():
    wb = load_workbook(PATH)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    hr, col = _headers(ws)
    c_co = col["Company"]
    c_cp = col["Contact Person"]
    c_role = col["Role"]
    c_li = col["LinkedIn"]
    n = 0
    unmatched = set(CONTACTS)
    for row in range(hr + 1, ws.max_row + 1):
        co = str(ws.cell(row, c_co).value or "").strip()
        if co not in CONTACTS:
            continue
        unmatched.discard(co)
        if not _empty(ws.cell(row, c_cp).value):
            print("SKIP already filled", co)
            continue
        rec = CONTACTS[co]
        ws.cell(row, c_cp).value = rec["person"]
        ws.cell(row, c_role).value = rec["role"]
        ws.cell(row, c_li).value = rec["li"]
        n += 1
        print(f"FILL {co} | {rec['person']} | {rec['role']}")
    save_atomic(wb, PATH)
    wb.close()
    print("filled", n)
    if unmatched:
        print("UNMATCHED", unmatched)


if __name__ == "__main__":
    main()
