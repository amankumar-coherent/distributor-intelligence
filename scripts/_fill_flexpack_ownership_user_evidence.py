"""Polish + fill remaining ownership gaps via DeepSeek."""
from __future__ import annotations

import json
import os
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook

load_dotenv(r"D:\vendor-intel\.env")
PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)

RAW = {
    "Box and Packaging Supply": (
        "Privately held company; part of The Logson Group; "
        "acquired by Logson Group in 2004; UK Boxes and Packaging established 1988."
    ),
    "Bulk Packaging": (
        "Bulk Packaging Liners Private Limited, Bangalore, India; private company; "
        "owned/managed by directors Aswinikumar Lakhotia and Mala Lakhotia; "
        "CIN U28129KA2003PTC031993; incorporated 26 May 2003."
    ),
    "Global Packaging": (
        "Global Pack Packaging Solutions Inc.; privately held; "
        "principals/executive leadership control the company; Canada."
    ),
    "Kiwi Packaging": (
        "Kiwi Packaging NZ Limited; fully owned by Simranjeet Singh who holds all "
        "issued shares (2 shares) and is company director; Hamilton, New Zealand; "
        "incorporated 12 August 2020."
    ),
    "Packaging Warehouse": (
        "Acquired by All4Labels Global Packaging Group on 16 November 2021; "
        "now part of All4Labels international packaging network."
    ),
}

FALLBACKS = {
    "Box and Packaging Supply": "Private; part of The Logson Group",
    "Bulk Packaging": "Private (Bulk Packaging Liners Pvt Ltd)",
    "Global Packaging": "Private",
    "Kiwi Packaging": "Private (owned by Simranjeet Singh)",
    "Packaging Warehouse": "Subsidiary of All4Labels Global Packaging Group",
}


def main() -> None:
    client = OpenAI(
        api_key=os.getenv("OPENAI_API_KEY", "").strip(),
        base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/"),
        timeout=60.0,
    )
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    system = (
        "Polish ownership for a B2B vendor spreadsheet. "
        'Return JSON {"ownership":"..."} only. '
        "Style like: Private; Private (family-owned); Private (SAS); "
        "Subsidiary of X; Private (Company Ltd). "
        "Keep concise (max ~12 words). Do not invent beyond FACTS. "
        "Prefer Subsidiary of / Private (owned by) when clear."
    )

    polished: dict[str, str] = {}
    for name, facts in RAW.items():
        try:
            resp = client.chat.completions.create(
                model=model,
                temperature=0,
                max_tokens=80,
                response_format={"type": "json_object"},
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": f"Company: {name}\nFACTS: {facts}"},
                ],
            )
            data = json.loads((resp.choices[0].message.content or "{}").strip())
            own = str(data.get("ownership") or "").strip()
        except Exception as ex:
            print("ERR", name, str(ex)[:120])
            own = ""
        if not own:
            own = FALLBACKS[name]
        polished[name] = own
        print(name, "=>", own)

    wb = load_workbook(PATH)
    ws = wb.active
    h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {n: h.index(n) + 1 for n in h}
    filled = 0
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, ic["Company"]).value or "").strip()
        if name in polished:
            ws.cell(r, ic["Ownership"]).value = polished[name]
            filled += 1
    wb.save(PATH)
    wb.close()

    wb = load_workbook(PATH, data_only=True)
    ws = wb.active
    oi = h.index("Ownership")
    ci = h.index("Company")
    named = blank = 0
    still: list[str] = []
    print("\n=== Applied ===")
    for row in ws.iter_rows(min_row=2, values_only=True):
        n = str(row[ci] or "").strip()
        v = str(row[oi] or "").strip()
        if n in polished:
            print(f"{n}: {v}")
        if not v or v.lower() == "not publicly disclosed":
            blank += 1
            still.append(n)
        else:
            named += 1
    print(
        json.dumps(
            {"filled_now": filled, "named": named, "npd": blank, "still": still},
            indent=2,
        )
    )
    wb.close()


if __name__ == "__main__":
    main()
