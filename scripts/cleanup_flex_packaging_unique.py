"""
Clean Global Flexible Packaging FINAL:
- drop fabricated geo/location clones + dead domains
- merge same company across locations into one name (no country in name)
- normalize Inc/Ltd noise and '(already listed)'
- blank placeholder / unverified fills (no invented replacements)
"""
from __future__ import annotations

import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Font

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
PROBE = FP.parent / "_website_probe.tsv"
OUT_TSV = FP.parent / "unique_companies_cleaned.tsv"
REPORT = FP.parent / "cleanup_unique_report.txt"

LOC_WORDS = {
    "india", "usa", "us", "u.s.", "u.s.a.", "uk", "u.k.", "uae", "china", "japan",
    "germany", "france", "italy", "spain", "brazil", "mexico", "canada", "australia",
    "korea", "south korea", "thailand", "vietnam", "indonesia", "malaysia", "singapore",
    "taiwan", "philippines", "turkey", "poland", "netherlands", "belgium", "switzerland",
    "sweden", "norway", "denmark", "finland", "austria", "ireland", "portugal", "greece",
    "russia", "saudi arabia", "south africa", "egypt", "argentina", "chile", "colombia",
    "peru", "pakistan", "bangladesh", "sri lanka", "new zealand", "hong kong",
    "europe", "asia", "apac", "emea", "latam", "north america", "south america",
    "middle east", "global", "worldwide", "international", "latin america", "scandinavia",
    "africa", "uk & ireland", "nz",
    # BR states / regions often used as fake company tails
    "amazonas", "bahia", "ceará", "ceara", "goiás", "goias", "minas", "maranhão",
    "maranhao", "mato grosso", "paraná", "parana", "pará", "para", "nordeste",
    "centro-oeste", "distrito federal", "espírito santo", "espirito santo",
    "rio", "rio preto", "paulista", "sp", "do brasil", "do sul", "do vale",
    "do nordeste", "nicarágua", "nicaragua", "paraguai", "uruguai", "méxico",
}

# Verified official sites only (do not invent beyond these)
KNOWN_SITES = {
    "bunzl": "https://www.bunzl.com",
    "airgas": "https://www.airgas.com",
    "duni": "https://www.duni.com",
    "detpak": "https://www.detpak.com",
    "bilcare": "https://www.bilcare.com",
    "ckspackaging": "https://www.ckspackaging.com",
    "hatzopoulos": "https://www.hatzopoulos.gr",
    "uflex": "https://www.uflexltd.com",
    "polyplex": "https://www.polyplex.com",
    "reynoldsconsumerproducts": "https://www.reynoldsconsumerproducts.com",
}

FABRICATED_NAME = re.compile(
    r"(?ix)^("
    r"embalagens\s+flex[ií]veis(\s|$)|"
    r"empaques\s+flexibles(\s|$)|"
    r"embalajes\s+flexibles(\s|$)|"
    r"pl[aá]sticos\s+flexibles(\s|$)|"
    r"pl[aá]sticos\s+y\s+empaques(\s|$)|"
    r"embalagens\s+(amazonas|bahia|cear[aá]|goi[aá]s|minas|maranh[aã]o|mato\s+grosso|"
    r"paran[aá]|par[aá]|nordeste|centro-oeste|distrito\s+federal|esp[ií]rito\s+santo|"
    r"rio(\s+preto)?|paulista|sp|do\s+brasil|do\s+sul|do\s+vale|nacional|modernas|"
    r"e\s+cia|flex\b)|"
    r"canadian\s+(film|flexible|flexibles|packaging|paper|poly)\b|"
    r"flexible\s+packaging\s+(australia|canada|group|inc\.?|products|solutions)\b|"
    r"packaging\s+(direct|distributors|plus|products|solutions|wholesale|world|supply)\b.*"
    r"\b(canada|usa|us|international|nz|australia)\b|"
    r"(vietnam|thai|singapore|philippine|turkish|egyptian|algerian|lebanese)\s+"
    r"(flexible\s+)?packaging\b|"
    r"(accra|dar\s+es\s+salaam|casablanca|cape|abu\s+dhabi|dubai|bahrain|cyprus|"
    r"east\s+africa)\s+packaging\b|"
    r"bunzl\s+(africa|asia|australia|canada|europe|latin\s+america|middle\s+east|"
    r"new\s+zealand|scandinavia|uk|distribution(\s+usa)?)\b|"
    r"flex-?pack\s+usa\b|"
    r"flexo\s+pack\s+india\b|"
    r"poly\s+(packaging|pak)\s+canada\b|"
    r"polytex\s+packaging\s+new\s+zealand\b|"
    r"pt\.?\s+mitra\s+packaging\s+indonesia\b|"
    r"african\s+packaging\s+solutions\b|"
    r"australian\s+packaging\s+solutions\b|"
    r"brasil\s+embalagens\b|"
    r"embalagens\s+e\s+embalagens\b|"
    r"empaques\s+del\s+(norte|sur)\b|"
    r"envases\s+andinos\b|"
    r"cairo\s+(flexible\s+)?packaging\b"
    r")"
)

LEGAL_TAIL = re.compile(
    r"[\s,]+(inc\.?|incorporated|llc|l\.l\.c\.|ltd\.?|limited|corp\.?|corporation|"
    r"co\.?|plc|gmbh|ag|sa|bv|nv|pty|pvt\.?|pte\.?|private|holdings?|group|"
    r"company|companies|s\.?a\.?|s\.?a\.?s\.?)\s*$",
    re.I,
)


def load_probe() -> dict[int, str]:
    status: dict[int, str] = {}
    if not PROBE.exists():
        return status
    for line in PROBE.read_text(encoding="utf-8").splitlines()[1:]:
        parts = line.split("\t")
        if len(parts) >= 4:
            status[int(parts[0])] = parts[3]
    return status


def domain_of(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    host = urlparse(u).netloc.lower().split("@")[-1]
    return host[4:] if host.startswith("www.") else host


def strip_already_listed(name: str) -> str:
    return re.sub(r"\s*\(already listed\)\s*", "", name, flags=re.I).strip()


def strip_location(name: str) -> str:
    n = strip_already_listed(name)
    # (Country) tails
    n = re.sub(r"\s*\(([^)]+)\)\s*$", lambda m: "" if m.group(1).strip().lower() in LOC_WORDS else m.group(0), n)
    # "… Packaging India" / "… Flexíveis Argentina"
    loc_alt = "|".join(re.escape(w) for w in sorted(LOC_WORDS, key=len, reverse=True))
    n = re.sub(rf"[\s,\-/|]+({loc_alt})\s*$", "", n, flags=re.I)
    n = re.sub(rf"\s+de\s+({loc_alt})\s*$", "", n, flags=re.I)
    n = re.sub(rf"\s+of\s+({loc_alt})\s*$", "", n, flags=re.I)
    # "Foo India Inc" → "Foo Inc"
    n = re.sub(
        rf"[\s,\-/]+({loc_alt})[\s,\-/]+(inc\.?|llc|ltd\.?|corp\.?|gmbh|ag|sa|plc)\s*$",
        r" \2",
        n,
        flags=re.I,
    )
    return re.sub(r"\s+", " ", n).strip(" ,-/|")


def normalize_legal_display(name: str) -> str:
    """Fix ', Inc.' / double spaces; keep one clean legal suffix if present."""
    n = strip_location(name)
    n = re.sub(r"\s+", " ", n).strip()
    # normalize ", Inc." / " Inc." → " Inc."
    n = re.sub(r",\s*(Inc\.?|LLC|Ltd\.?|Corp\.?|PLC|GmbH)\s*$", r" \1", n, flags=re.I)
    n = re.sub(r"\bInc\b\.?", "Inc.", n, flags=re.I) if re.search(r"\bInc\b", n, re.I) else n
    n = re.sub(r"\bLtd\b\.?", "Ltd.", n, flags=re.I) if re.search(r"\bLtd\b", n, re.I) else n
    # drop trailing legal only when name is otherwise empty after strip — keep for real cos
    return n.strip(" ,")


def canon_key(name: str) -> str:
    n = strip_location(name)
    # peel legal suffixes repeatedly
    for _ in range(3):
        n2 = LEGAL_TAIL.sub("", n).strip(" ,./")
        if n2 == n:
            break
        n = n2
    n = re.sub(r"[^\w\s]", " ", n.lower())
    n = re.sub(r"\s+", " ", n).strip()
    n = re.sub(r"\b(flexible\s+)?packaging\b$", "", n).strip()
    return n


def website_ok(status: str) -> bool:
    return status.startswith("http_2") or status.startswith("http_3")


def is_placeholder(val) -> bool:
    s = str(val or "").strip().lower()
    return s in {"", "not publicly disclosed", "n/a", "na", "none", "-"}


def is_generic_specialty(val) -> bool:
    s = str(val or "").strip().lower()
    return bool(
        re.match(
            r"^(distributes|imports and distributes)\s+flexible packaging",
            s,
        )
    )


def score_row(row: tuple, status: str) -> int:
    score = 0
    if website_ok(status):
        score += 50
    elif status == "http_fail":
        score += 10  # domain may exist (403)
    name = str(row[0] or "")
    if re.search(r"\bplc\b", name, re.I):
        score += 20
    if "(already listed)" in name.lower():
        score -= 5
    if strip_location(name) != name:
        score -= 15  # prefer non-location name
    # richer real-looking categories
    if row[8] and not is_placeholder(row[8]):
        score += 5
    if row[2] and not is_placeholder(row[2]):
        score += 3
    return score


def merge_fields(rows: list[tuple]) -> list:
    """Pick best non-empty real values across duplicates; never invent."""
    base = list(rows[0])
    for r in rows[1:]:
        for i, v in enumerate(r):
            if i >= len(base):
                break
            cur = base[i]
            if is_placeholder(cur) and not is_placeholder(v):
                if i == 9 and is_generic_specialty(v):
                    continue
                base[i] = v
    return base


def main() -> None:
    probe = load_probe()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = FP.with_name(FP.stem + f"_backup_before_unique_{stamp}.xlsx")
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    raw = list(ws.iter_rows(min_row=2, values_only=True))

    kept_pre: list[tuple[int, tuple, str]] = []
    dropped: list[tuple[int, str, str, str]] = []

    for excel_row, row in enumerate(raw, start=2):
        name = str(row[0] or "").strip()
        web = str(row[1] or "").strip()
        st = probe.get(excel_row, "")
        if not name:
            dropped.append((excel_row, name, web, "empty_name"))
            continue
        # Drop fabricated geo clones / census filler
        if FABRICATED_NAME.search(name):
            dropped.append((excel_row, name, web, "fabricated_geo_or_clone"))
            continue
        # Drop dead websites unless http_fail (possible real 403)
        if st == "dns_fail":
            dropped.append((excel_row, name, web, "dns_fail_unverified"))
            continue
        if st in {"empty", "bad_url", "unparseable"}:
            dropped.append((excel_row, name, web, "bad_website"))
            continue
        kept_pre.append((excel_row, row, st))

    # Force-keep Bunzl plc if present even if somehow dropped — scan raw
    bunzl_rows = []
    for excel_row, row in enumerate(raw, start=2):
        name = str(row[0] or "").strip()
        if re.match(r"(?i)^bunzl\b", name):
            bunzl_rows.append((excel_row, row, probe.get(excel_row, "")))

    # Group by canon
    groups: dict[str, list[tuple[int, tuple, str]]] = defaultdict(list)
    for item in kept_pre:
        key = canon_key(str(item[1][0]))
        if not key or key in {"flexible", "packaging", "global", "national", "alliance"}:
            # too-generic after strip — keep as own key using full cleaned name
            key = re.sub(r"\s+", " ", strip_location(str(item[1][0])).lower())
        groups[key].append(item)

    # Inject Bunzl merge from all Bunzl rows
    if bunzl_rows:
        groups["bunzl"] = bunzl_rows

    merged_rows: list[list] = []
    merge_log: list[str] = []
    for key, items in sorted(groups.items(), key=lambda kv: kv[0]):
        items_sorted = sorted(items, key=lambda it: score_row(it[1], it[2]), reverse=True)
        best_status = items_sorted[0][2]
        merged = merge_fields([it[1] for it in items_sorted])
        clean_name = normalize_legal_display(str(merged[0]))
        # Prefer parent-style name without location
        for it in items_sorted:
            cand = normalize_legal_display(str(it[1][0]))
            if strip_location(str(it[1][0])) == str(it[1][0]).strip() or re.search(r"\bplc\b", cand, re.I):
                clean_name = cand
                break
        # Bunzl special
        if key == "bunzl":
            clean_name = "Bunzl"
            merged[1] = KNOWN_SITES["bunzl"]
        merged[0] = clean_name

        # Website: keep only if probe ok / http_fail; else blank. Apply known map.
        web = str(merged[1] or "").strip()
        d = domain_of(web).replace(".", "")
        # known site by canon
        if key in KNOWN_SITES:
            merged[1] = KNOWN_SITES[key]
        elif not website_ok(best_status) and best_status != "http_fail":
            merged[1] = ""
        elif best_status == "dns_fail":
            merged[1] = ""

        # Blank placeholders / mass templates — real data only
        for idx in (2, 3, 4, 5, 6, 7, 8):  # founded..core categories
            if idx < len(merged) and is_placeholder(merged[idx]):
                merged[idx] = None
        if len(merged) > 9 and (is_placeholder(merged[9]) or is_generic_specialty(merged[9])):
            merged[9] = None
        # Key brands placeholder
        if len(merged) > 10 and is_placeholder(merged[10]):
            merged[10] = None
        # Mass template channels
        if len(merged) > 11 and str(merged[11] or "").strip().lower() == "no":
            merged[11] = None
        if len(merged) > 12 and str(merged[12] or "").strip().lower() == "wholesaler":
            merged[12] = None

        # Contacts: keep as-is (already sparse); do not invent. Drop email/phone if no website.
        if not merged[1]:
            for idx in (13, 14, 15, 16, 17):
                if idx < len(merged):
                    merged[idx] = None

        merged_rows.append(merged)
        if len(items_sorted) > 1:
            names = [str(it[1][0]) for it in items_sorted]
            merge_log.append(f"{clean_name} <= {names}")

    # Sort by company name
    merged_rows.sort(key=lambda r: str(r[0] or "").lower())

    # Rewrite sheet
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in merged_rows:
        ws.append(list(r))
    # bold header already
    for cell in ws[1]:
        cell.font = Font(bold=True)

    wb.save(FP)

    # TSV of unique company + website
    with OUT_TSV.open("w", encoding="utf-8") as f:
        f.write("company\twebsite\n")
        for r in merged_rows:
            f.write(f"{r[0]}\t{r[1] or ''}\n")

    with_web = sum(1 for r in merged_rows if r[1])
    report = [
        f"backup: {backup}",
        f"before_rows: {len(raw)}",
        f"dropped: {len(dropped)}",
        f"after_unique: {len(merged_rows)}",
        f"with_website: {with_web}",
        f"without_website: {len(merged_rows) - with_web}",
        f"merge_groups: {len(merge_log)}",
        "",
        "=== MERGES ===",
        *merge_log,
        "",
        "=== DROP REASONS ===",
    ]
    reason_c: dict[str, int] = defaultdict(int)
    for _, _, _, reason in dropped:
        reason_c[reason] += 1
    for k, v in sorted(reason_c.items(), key=lambda x: -x[1]):
        report.append(f"{k}: {v}")
    report.append("")
    report.append("=== UNIQUE COMPANY + WEBSITE ===")
    for r in merged_rows:
        report.append(f"{r[0]}\t{r[1] or ''}")
    REPORT.write_text("\n".join(report), encoding="utf-8")
    print("\n".join(report[:40]))
    print(f"... wrote {REPORT}")
    print(f"UNIQUE={len(merged_rows)} WITH_WEB={with_web}")


if __name__ == "__main__":
    main()
