# -*- coding: utf-8 -*-
import json
from pathlib import Path
from openpyxl import load_workbook

NPD = {"", "nan", "none", "n/a", "na", "not publicly disclosed", "-", "—", "unknown", "tbd", "null"}
FILES = {
    "electrical": Path(r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global\global_electrical_distribution_market_global_FINAL.xlsx"),
    "avocado": Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market.xlsx"),
    "smartphone": Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    "firewall": Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market.xlsx"),
}
rows = []
for mkt, path in FILES.items():
    wb = load_workbook(path, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            headers = vals
            hr = r
            break
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip()) for i in range(len(headers))}
        co = d.get("Company") or ""
        if not co:
            continue
        person = (d.get("Contact Person") or "").strip()
        li = (d.get("LinkedIn") or "").strip()
        if not person or person.lower() in NPD:
            continue
        if li and "linkedin.com/in/" in li.lower():
            continue
        rows.append({
            "m": mkt,
            "company": co,
            "person": person,
            "role": d.get("Role") or "",
            "hq": d.get("Headquarters") or "",
        })
    wb.close()
Path(r"D:\vendor-intel\output\_npd_linkedin.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
from collections import Counter
print("npd_with_name", len(rows), dict(Counter(x["m"] for x in rows)))
# print complete names first (no asterisks)
clean = [x for x in rows if "***" not in x["person"] and "*" not in x["person"]]
print("clean_names", len(clean))
for x in clean[:80]:
    print(f"{x['m'][:3]} | {x['person']} | {x['role'][:40]} | {x['company'][:36]}")
