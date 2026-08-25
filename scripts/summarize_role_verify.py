import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

audit_p = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\channel_role_verify_388_audit.json"
)
a = json.loads(audit_p.read_text(encoding="utf-8"))
print("before", a["before"], "after", a["after"], "removed", a["removed_count"], "errors", a["errors"])
print("file", a["file"])
print("backup", a["backup"])
print("---REMOVED---")
for x in a["removed"]:
    company = x.get("company")
    role = x.get("role")
    reason = str(x.get("reason", ""))[:140]
    print(f"{company} | {role} | {reason}")

FP = Path(a["file"])
wb = load_workbook(FP, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(min_row=2, values_only=True))
print("excel_rows", len(rows))
print("dist_types", Counter(r[12] for r in rows))
keys = ["manufactur", "converter", "oem", "hybrid manufacturer", "we manufacture"]
hits = []
for r in rows:
    summ = str(r[20]) if len(r) > 20 and r[20] else ""
    blob = ((r[0] or "") + " " + summ).lower()
    if any(k in blob for k in keys):
        hits.append((r[0], blob[:100]))
print("mfr_language_left", len(hits))
for h in hits:
    print(" ", h[0], "|", h[1][:90])
wb.close()
