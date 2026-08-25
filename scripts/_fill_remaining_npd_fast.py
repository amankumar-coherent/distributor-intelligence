#!/usr/bin/env python3
"""Fill remaining NPD Key Brands + contacts from agent web verify (fast)."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

X = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
AUDIT = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_remaining_npd_fill_audit.json"
)
NPD = "Not publicly disclosed"

# Match by current Company cell value
FILLS: dict[str, dict] = {
    "Connect The World S.A. de C.V.": {
        "Key Brands Represented": "Connect The World, others",
        "Core Categories": "Hass avocados; frozen avocado pulp/guacamole; extra virgin avocado oil; chia & superfoods, others",
        "Specialty Focus": "Mexican avocado oil & avocado derivatives export (OEM/white-label capable), others",
        "Headquarters": "Av. Patria 2303, Guadalajara, Jalisco, Mexico",
        "Summary": "Connect The World S.A. de C.V. — Guadalajara exporter of Hass avocados, frozen pulp/guacamole, EVO avocado oil and chia; OEM/white-label. Web: https://www.ctworldmx.com",
    },
    "DKSH Indonesia": {
        "Key Brands Represented": "Ovaltine, King's Fisher, Cap Kaki Tiga, Cap Panda, Ellips, others",
        "Website": "https://www.dksh.com/id-en/home",
        "Summary": "DKSH Indonesia — market expansion services; distributes FMCG brands incl. Ovaltine, King's Fisher canned fish, and Kino lines (Cap Kaki Tiga, Cap Panda, Ellips).",
    },
    "Foodservice Direct": {
        "Key Brands Represented": "Hellmann's, Knorr, Sir Kensington's, Ben & Jerry's, Lipton, others",
        "Summary": "FoodServiceDirect.com — largest online US foodservice distributor; sells national branded products only (no private label); 250,000+ SKUs incl. Unilever Food Solutions brands.",
    },
    "Fresh Produce Consortium (members distribution)": {
        "Key Brands Represented": NPD,
        "Summary": "Fresh Produce Consortium (UK) — trade association for fresh produce companies; member distributors carry their own brands. No single FPC house brand.",
    },
    "Labatt Food Service": {
        "Key Brands Represented": NPD,
        "Office No.": "+1 210-661-4216",
        "Office Email": "sanantoniocustomerservice@labattfood.com",
        "Headquarters": "4500 Industry Park, San Antonio, Texas, United States",
        "Summary": "Labatt Food Service — San Antonio family-owned broadliner; explicitly no private labels (national brands only). HQ phone +1 210-661-4216.",
    },
    "Maines Paper & Food Service": {
        "Company": "Maines Paper & Food Service (assets acquired by Lineage Logistics)",
        "Key Brands Represented": "Burger King, Tim Hortons, Olive Garden, LongHorn Steakhouse, Cheddar's, The Capital Grille, others",
        "Ownership": "Acquired (Lineage Logistics, 2020)",
        "Website": "https://www.maines.net/",
        "Summary": "Maines Paper & Food Service — historic QSR/casual dining distributor (Burger King, Tim Hortons, Darden brands). Distribution assets acquired by Lineage Logistics in 2020; treat as legacy/channel record.",
    },
    "The Martin-Brower Company, L.L.C.": {
        "Key Brands Represented": "McDonald's, Chipotle, Chick-fil-A, Panera Bread, others",
        "Summary": "The Martin-Brower Company, L.L.C. (Reyes Holdings) — largest global QSR supply-chain partner to McDonald's; also serves Chipotle, Chick-fil-A, Panera in select markets. No Martin-Brower retail private label.",
    },
    "Pro*Act": {
        "Key Brands Represented": NPD,
        "Summary": "PRO*ACT — produce sourcing/category management network (SSI family); connects operators to inspected DCs. No PRO*ACT consumer house-brand line publicly disclosed (contrast Markon First Crop at Markon Cooperative).",
    },
    "Reyes Holdings, L.L.C.": {
        "Key Brands Represented": "Coca-Cola, Reyes Coca-Cola Bottling, Reyes Beverage Group, Martin Brower, others",
        "Office No.": "+1 847-227-6500",
        "Summary": "Reyes Holdings, L.L.C. — family holding company for Martin Brower (McDonald's distribution), Reyes Coca-Cola Bottling, and Reyes Beverage Group (US beer). Brands via bottling/distribution units.",
    },
    "Sojitz": {
        "Company": "Sojitz Corporation / Sojitz Foods Corporation (SOFCO)",
        "Website": "https://www.sofco.sojitz.com/en/",
        "Key Brands Represented": "Sojitz Foods / SOFCO, others",
        "Office No.": "+81 3-5574-3300",
        "Headquarters": "Roppongi T-Cube, Minato-ku, Tokyo, Japan",
        "Summary": "Sojitz Foods Corporation (SOFCO) — Sojitz Group food trading house; sugar, flour, oils & fats, dairy, cocoa/chocolate imports. No major consumer edible-oil house brand disclosed.",
    },
    "Vistar (Performance Food Group)": {
        "Key Brands Represented": "Hostess, Snickers, M&M's, Hershey, others",
        "Office No.": "+1 800-880-9900",
        "Headquarters": "188 Inverness Drive West, Englewood, Colorado, United States",
        "Summary": "Vistar (PFG) — national candy/snack/beverage specialty distributor for vending, OCS, theatres and concessions; carries major national brands (e.g. Hostess, Snickers, Hershey). No distinct Vistar edible-oil private label.",
    },
    "DIVERSE WORLDWIDE TRADING LIMITED": {
        "Key Brands Represented": NPD,
        "Summary": "DIVERSE WORLDWIDE TRADING LIMITED — Kingston upon Thames; europages edible-oil/FMCG listings. No verified proprietary brand set; UK Cos House SICs include real estate — verify before outreach.",
    },
}

wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

applied = 0
for r in range(3, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co not in FILLS:
        continue
    for k, v in FILLS[co].items():
        if k in ic:
            ws.cell(r, ic[k]).value = v
    applied += 1

wb.save(X)
wb.close()

wb = openpyxl.load_workbook(X, data_only=True)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i for i, h in enumerate(hdr) if h}
real = npd = 0
still = []
phone_npd = email_npd = 0
for row in ws.iter_rows(min_row=3, values_only=True):
    c = str(row[ic["Company"]] or "").strip()
    if not c:
        continue
    b = str(row[ic["Key Brands Represented"]] or "").strip()
    if b in ("", NPD):
        npd += 1
        still.append(c)
    else:
        real += 1
    if str(row[ic["Office No."]] or "").strip() in ("", NPD):
        phone_npd += 1
    if str(row[ic["Office Email"]] or "").strip() in ("", NPD):
        email_npd += 1
wb.close()

audit = {
    "ts": datetime.now(timezone.utc).isoformat(),
    "applied_rows": applied,
    "total": real + npd,
    "real_brands": real,
    "npd_brands": npd,
    "still_npd": still,
    "phone_npd": phone_npd,
    "email_npd": email_npd,
}
AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
print(json.dumps(audit, indent=2))
