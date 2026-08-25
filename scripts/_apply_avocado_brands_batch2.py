#!/usr/bin/env python3
import openpyxl
from pathlib import Path

X = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
NPD = "Not publicly disclosed"
EXTRA = {
    "Bidvest Australia": "Caterers' Choice, Menu Maker, Jeffersons, others",
    "Turner Price": "Caterfood Select, others",
    "Nature's Best": "CADIA, MADE WITH, others",
    "Grupo Herdez Distribution": "Herdez, Del Fuerte, Dona Maria, Embasa, Bufalo, Aires de Campo, Yemina, Wholly Guacamole, others",
    "Lulu Hypermarket": "LuLu, others",
    "Agthia Distribution": "Al Ain, Yoplait, Grand Mills, Capri Sun, Al Bayan, Alpin, others",
}

wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}
n = 0
for r in range(3, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co in EXTRA:
        ws.cell(r, ic["Key Brands Represented"]).value = EXTRA[co]
        n += 1
        print("set", co)

# remove Mitsubishi if still present (pharma/nutrition, weak avocado-oil channel)
for r in range(ws.max_row, 2, -1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co == "Mitsubishi International Food Ingredients, Inc.":
        ws.delete_rows(r, 1)
        print("removed", co)

wb.save(X)
wb.close()

wb = openpyxl.load_workbook(X, data_only=True)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i for i, h in enumerate(hdr) if h}
real = npd = 0
for row in ws.iter_rows(min_row=3, values_only=True):
    c = str(row[ic["Company"]] or "").strip()
    if not c:
        continue
    b = str(row[ic["Key Brands Represented"]] or "").strip()
    if b in ("", NPD):
        npd += 1
    else:
        real += 1
print({"total": real + npd, "real": real, "npd": npd, "applied": n})
wb.close()
