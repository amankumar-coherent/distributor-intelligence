# -*- coding: utf-8 -*-
from datetime import datetime
from pathlib import Path

import openpyxl

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
    r"\global_wearable_medical_devices_market_global_FINAL.xlsx"
)

FILLS = {
    "Africure Pharmaceuticals": (
        "Public (Stock Exchange of Mauritius) — controlling stakeholder Sinhue Bosco Noronha"
    ),
    "ProCare Pharmacy": "Private (closely held)",
}


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = XLSX.with_name(XLSX.stem + f"_BEFORE_OWNERSHIP_LAST2_{stamp}.xlsx")
    wb0 = openpyxl.load_workbook(XLSX)
    wb0.save(bak)
    wb0.close()

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co not in FILLS:
            continue
        old = str(ws.cell(r, ic["Ownership"]).value or "").strip() or "(blank)"
        ws.cell(r, ic["Ownership"]).value = FILLS[co]
        print(f"{co}: {old} -> {FILLS[co]}")

    wb.save(XLSX)
    wb.close()

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    oi = headers.index("Ownership")
    filled = sum(
        1 for row in ws.iter_rows(min_row=2, values_only=True) if str(row[oi] or "").strip()
    )
    print(f"Ownership filled {filled}/{ws.max_row - 1}")
    print(f"backup {bak.name}")
    wb.close()


if __name__ == "__main__":
    main()
