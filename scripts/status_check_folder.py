#!/usr/bin/env python3
"""Active / closed / acquired check for a folder of *_verified_distributors.xlsx.

Successor to status_check_deliverables.py for any folder. That script sent
company NAMES only, and AI Mode matched same-named strangers ("Nexus
Engineering -> Scientific Atlanta", "Fire Fox -> Hero Cycles", Indian
distributors labelled with US acquirers). Here every company goes out with
its website and HQ, the website is the identity, and a "closed" verdict
needs stated evidence before a row is removed.

Never edits the input files. Output folder gets, per input workbook:
  <same name>.xlsx  closed/absorbed companies removed,
                    acquired ones renamed "ABC (acquired by XYZ)"
  <same name>.json  JSON copy of that workbook
plus status_check_report.xlsx (every closed / absorbed / acquired verdict
with its evidence) and, with --zip, the folder zipped.

Usage:
  python scripts/status_check_folder.py --input DIR --output DIR --workers 5
  python scripts/status_check_folder.py --input DIR --status
  python scripts/status_check_folder.py --input DIR --output DIR --build --zip

--workers N runs N browsers (one profile each, data/ai_mode_status_<i>),
then builds the output. Every answer is checkpointed per worker in
data/status_check_v2/<run>/, so re-running the same command resumes.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import subprocess
import sys
import time
import zipfile
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

from dotenv import load_dotenv  # noqa: E402

load_dotenv(ROOT / ".env")

import openpyxl  # noqa: E402

from status_check_deliverables import clean_acquirer, extract_json, internet_up  # noqa: E402
from scripts.excel_polish import (  # noqa: E402
    polish_distributor_sheet, polish_flag_sheet, polish_report_sheet)

HEADER_ROW = 5
FIRST_DATA_ROW = 6
# Companies per query. STATUS_CHUNK=1 isolates a company whose website makes
# AI Mode refuse the whole query ("no response available") -- geo 2026-10-14:
# one such site blocked its 9 chunk-mates through every retry and profile reset.
CHUNK = max(1, int(os.environ.get("STATUS_CHUNK", "10")))
MAX_CONSECUTIVE_EMPTY = 3
MAX_PASSES = 5  # companies an answer skipped are re-asked, up to this many passes

_PROMPT = (
    "Check the CURRENT status of each company below. The Website identifies "
    "the company; a different company with a similar name is NOT this one.\n"
    "status: \"closed\" only if this company is confirmed shut down, "
    "dissolved, liquidated, struck off or ceased trading; \"absorbed\" if it "
    "was acquired or merged and no longer trades under its own name; "
    "\"active\" if it still operates; \"unknown\" if you cannot tell.\n"
    "acquired_by: the company or investor that acquired a majority or all of "
    "THIS company, even if it still trades under its own name; \"\" if it "
    "was not acquired (a minority stake or a distribution partnership is not "
    "an acquisition).\n"
    "year: year of the closure or acquisition, else \"\".\n"
    "evidence: one short sentence stating the fact and its source, for any "
    "closed, absorbed or acquired answer; else \"\".\n"
    "Answer every id. Reply with ONLY this JSON:\n"
    '{"results":[{"id":1,"Company":"","status":"active","acquired_by":"",'
    '"year":"","evidence":""}]}'
)


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def _norm(name: object) -> str:
    return "".join(ch for ch in str(name or "").lower() if ch.isalnum())


def _domain(url: object) -> str:
    u = re.sub(r"^(https?://)?(www\.)?", "", str(url or "").strip().lower())
    return u.split("/")[0]


def company_key(name: object, website: object) -> str:
    return f"{_norm(name)}|{_domain(website)}"


def input_files(folder: Path) -> list[Path]:
    return sorted(
        f for f in folder.glob("*_verified_distributors.xlsx")
        if not f.name.startswith("~$")
    )


def _col(header: list, name: str) -> int | None:
    try:
        return header.index(name)
    except ValueError:
        return None


def read_rows(xlsx: Path) -> list[dict]:
    ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(min_row=HEADER_ROW, values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    ci, wi, hi = _col(header, "Company"), _col(header, "Website"), _col(header, "Headquarters")
    out = []
    for r in rows[1:]:
        name = r[ci] if ci is not None and ci < len(r) else None
        if not name or not str(name).strip():
            continue
        out.append({
            "Company": str(name).strip(),
            "Website": str(r[wi] or "").strip() if wi is not None else "",
            "HQ": str(r[hi] or "").strip() if hi is not None else "",
        })
    return out


def all_companies(folder: Path) -> dict[str, dict]:
    """Distinct companies across every file, keyed by name + website domain."""
    seen: dict[str, dict] = {}
    for f in input_files(folder):
        for row in read_rows(f):
            seen.setdefault(company_key(row["Company"], row["Website"]), row)
    return dict(sorted(seen.items()))


# -- checkpoints ---------------------------------------------------------------

def run_dir(folder: Path) -> Path:
    d = ROOT / "data" / "status_check_v2" / re.sub(r"[^A-Za-z0-9]+", "_", folder.name).strip("_")
    d.mkdir(parents=True, exist_ok=True)
    return d


def load_json(p: Path) -> dict:
    try:
        # utf-8-sig: an overrides.json saved by Windows PowerShell starts with
        # a BOM, which plain utf-8 rejects -> {} -> every correction silently
        # ignored (soft 2026-10-21).
        return json.loads(p.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return {}


def save_json(p: Path, data: dict) -> None:
    tmp = p.with_suffix(f".{os.getpid()}.tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, p)


def all_verdicts(folder: Path) -> dict[str, dict]:
    merged: dict[str, dict] = {}
    for p in sorted(run_dir(folder).glob("shard_*.json")):
        merged.update(load_json(p))
    # Reviewer corrections, keyed like the verdicts: {"key": {"status": ...,
    # "acquired_by": ..., "note": why}}. Data, not code, and they survive
    # re-runs because the shard answers are never edited.
    for key, fix in load_json(run_dir(folder) / "overrides.json").items():
        if key in merged and isinstance(fix, dict):
            orig = {k: merged[key].get(k) for k in ("status", "acquired_by", "year", "evidence")}
            merged[key] = {**merged[key], **fix, "_orig": orig}  # keeps what AI Mode said, for the flag note
    return merged


_AGREED = re.compile(r"agree|signed|plans? to|intend|in talks", re.I)
_DONE = re.compile(r"complet|finali[sz]ed|closed|became|joined|integrated|"
                   r"acquired (by|on|in)\b|was acquired|has acquired", re.I)


def flag_reason(name: str, v: dict | None, action: str, acquirer: str) -> str | None:
    """Why this company's ownership label is suspicious, or None.

    Flagged companies are kept OUT of the main Excel (verified companies only)
    and listed in flagged_for_review.xlsx instead. Flagged means:
      * a reviewer correction (AI Mode's claim was not backed by its evidence;
        an override with "flag": false is a correction that needs no review);
      * two different owners found (the name already said "(acquired by X)" from
        the earlier check, the status check says Y);
      * a labelled deal whose evidence only shows an agreement, not completion.
    """
    if not v:
        return None
    if v.get("note") and v.get("_orig") and v.get("flag") is not False:
        return f"AI Mode said: {_said(v)}. Review: {v['note']}."
    if action != "renamed" or not acquirer:
        return None
    m = re.search(r"\(acquired by (.+)\)\s*$", name, re.I)
    if m and _norm(acquirer) not in _norm(m.group(1)) and _norm(m.group(1)) not in _norm(acquirer):
        return (f"Two different owners found: '{m.group(1).strip()}' (earlier check) vs "
                f"'{acquirer}' (status check).")
    ev = str(v.get("evidence") or "")
    if _AGREED.search(ev) and not _DONE.search(ev):
        return "The evidence shows a deal agreed or signed; completion is not confirmed."
    return None


# -- one worker ----------------------------------------------------------------

def _parse(answer: str, batch: list[tuple[str, dict]]) -> dict[str, dict]:
    parsed = extract_json(answer or "")
    if not parsed:
        return {}
    by_id = {i + 1: key for i, (key, _) in enumerate(batch)}
    by_name = {_norm(row["Company"]): key for key, row in batch}
    got: dict[str, dict] = {}
    for r in parsed.get("results") or []:
        if not isinstance(r, dict):
            continue
        try:
            key = by_id.get(int(r.get("id")))
        except (TypeError, ValueError):
            key = None
        key = key or by_name.get(_norm(r.get("Company")))
        if not key or key in got:
            continue
        company = dict(batch)[key]["Company"]
        status = str(r.get("status") or "").strip().lower()
        if status not in ("active", "closed", "absorbed", "unknown"):
            status = "unknown"
        got[key] = {
            "Company": company,
            "status": status,
            "acquired_by": clean_acquirer(r.get("acquired_by"), company),
            "year": str(r.get("year") or "").strip(),
            "evidence": str(r.get("evidence") or "").strip(),
        }
    return got


def run_worker(folder: Path, shard: int, shards: int) -> None:
    from vendor_intel.scraping import google_ai_mode as gam

    companies = all_companies(folder)
    mine = [(k, v) for i, (k, v) in enumerate(companies.items()) if i % shards == shard]
    ck = run_dir(folder) / f"shard_{shard}.json"
    verdicts = load_json(ck)
    # Answers saved under another worker count (a run with a different
    # --workers number shards differently); they are never asked again.
    elsewhere = all_verdicts(folder)
    mine = [(k, v) for k, v in mine if k in verdicts or k not in elsewhere]
    log(f"[w{shard}] {len(mine)} companies, {sum(k in verdicts for k, _ in mine)} already checked")

    queries = empty = waits_used = 0
    # After MAX_CONSECUTIVE_EMPTY empty answers: pause and retry this many
    # times before the worker stops (0 = stop at once, the old behaviour).
    wall_waits = int(os.getenv("STATUS_WALL_WAITS") or 0)
    wall_wait_sec = int(os.getenv("STATUS_WALL_WAIT_SEC") or 600)
    try:
        for pass_no in range(1, MAX_PASSES + 1):
            todo = [(k, v) for k, v in mine if k not in verdicts]
            if not todo:
                break
            log(f"[w{shard}] pass {pass_no}: {len(todo)} to check")
            for i in range(0, len(todo), CHUNK):
                batch = todo[i:i + CHUNK]
                payload = [{"id": n + 1, "Company": row["Company"], "Website": row["Website"],
                            "HQ": row["HQ"]} for n, (_, row) in enumerate(batch)]
                prompt = f"{_PROMPT}\n\nCompanies:\n{json.dumps(payload, ensure_ascii=False)}"
                got = _parse(gam.ask(prompt, retries=1), batch)
                queries += 1
                if not got:
                    if not internet_up():
                        log(f"[w{shard}] no network, waiting 60s")
                        time.sleep(60)
                        continue
                    empty += 1
                    log(f"[w{shard}] empty answer ({empty}/{MAX_CONSECUTIVE_EMPTY})")
                    if empty >= MAX_CONSECUTIVE_EMPTY:
                        if waits_used < wall_waits:
                            # Google's block usually lifts after a pause; a
                            # fresh profile afterwards drops the flagged cookies.
                            waits_used += 1
                            log(f"[w{shard}] {empty} empty answers in a row — pausing "
                                f"{wall_wait_sec // 60} min ({waits_used}/{wall_waits}), then fresh profile")
                            gam.shutdown()
                            time.sleep(wall_wait_sec)
                            if not gam.reset_session("blocked: pause before resuming"):
                                gam.shutdown()
                            empty = 0
                            continue
                        log(f"[w{shard}] STOPPING: {empty} empty answers in a row "
                            "(AI Mode limit for this IP). Re-run to resume.")
                        return
                    continue
                empty = 0
                verdicts.update(got)
                save_json(ck, verdicts)
                done = sum(k in verdicts for k, _ in mine)
                log(f"[w{shard}] q{queries}: {len(got)}/{len(batch)} answered | {done}/{len(mine)} done")
        left = sum(k not in verdicts for k, _ in mine)
        log(f"[w{shard}] finished: {queries} queries, {left} unanswered")
    finally:
        gam.shutdown()


# -- build output --------------------------------------------------------------

def label(v: dict | None) -> tuple[str, str]:
    """(action, acquirer) for one verdict. A removal needs stated evidence."""
    if not v:
        return "unchecked", ""
    if v["status"] in ("closed", "absorbed") and v.get("evidence"):
        return "removed", v.get("acquired_by", "")
    if v.get("acquired_by"):
        return "renamed", v["acquired_by"]
    return "kept", ""


_LEGAL = re.compile(
    r"\b(ltda|eireli|sa|me|epp|inc|llc|gmbh|ltd|limited|pvt|private|srl|spa|bv|nv|ag|corp|corporation|co|"
    r"brasil|brazil|brasileira|do|da|de)\b")
_SECOND_LEVEL = {"com", "org", "net", "gov", "edu", "co"}


def registrable_domain(url: object) -> str:
    """'https://discover.univarsolutions.com/x' -> 'univarsolutions.com'; handles
    two-label country suffixes ('www.gabco.com.br' -> 'gabco.com.br')."""
    host = re.sub(r"^[a-z]*://", "", str(url or "").strip().lower()).split("/")[0].split(":")[0]
    host = re.sub(r"^www\.", "", host)
    parts = [p for p in host.split(".") if p]
    if len(parts) >= 3 and len(parts[-1]) == 2 and parts[-2] in _SECOND_LEVEL:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:]) if len(parts) >= 2 else host


def name_key(name: object) -> str:
    """Company name without accents, legal forms ('Ltda', 'Inc'), country words
    and the '(acquired by ...)' label, letters and digits only."""
    import unicodedata

    s = re.sub(r"\(acquired by.*\)\s*$", "", str(name or ""), flags=re.I).lower()
    # "S.r.l." / "S.A." / "Pvt." -> "srl" / "sa" / "pvt" so the legal-form
    # filter below sees them (dotted forms escaped it: "G&P Center S.r.l."
    # did not match "G&P Center Srl").
    s = re.sub(r"(?<=\b[a-z])\.(?=[a-z]\b|[a-z]\.)", "", s).replace(".", " ")
    s = "".join(ch for ch in unicodedata.normalize("NFKD", s) if not unicodedata.combining(ch))
    s = _LEGAL.sub(" ", re.sub(r"[^a-z0-9& ]", " ", s))
    return re.sub(r"[^a-z0-9]", "", s)


def find_duplicates(rows: list[dict]) -> dict[int, int]:
    """{row to drop: row to keep} for the same company listed twice.

    Same company = same website domain AND same HQ country, or the same name
    (legal forms ignored) AND same HQ city. A row already labelled
    "(acquired by ...)" is preferred as the keeper, then the fullest row.
    rows: [{"row", "name", "website", "city", "country", "filled"}]."""
    parent = {r["row"]: r["row"] for r in rows}

    def find(x: int) -> int:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a: int, b: int) -> None:
        parent[find(a)] = find(b)

    seen: dict[tuple, int] = {}
    for r in rows:
        keys = []
        dom = registrable_domain(r["website"])
        if dom and r["country"]:
            keys.append(("d", dom, r["country"]))
        nk = name_key(r["name"])
        if nk and r["city"]:
            keys.append(("n", nk, r["city"]))
        for k in keys:
            if k in seen:
                union(r["row"], seen[k])
            else:
                seen[k] = r["row"]
    # Same name (legal forms ignored) in the SAME country, or where one row's
    # HQ has no country at all ("Milan" vs "Milan, Italy"): one company with a
    # second website or branch city. The city rule above missed these -- the
    # 2026-10-21 audit found "G&P Center Srl (Milan, Italy)" / "G&P Center
    # S.r.l. (Milan)" and "Prescription Supply Inc" in Northwood and Toledo.
    # Same name in two DIFFERENT countries stays separate (country arms).
    by_name: dict[str, list[dict]] = {}
    for r in rows:
        nk = name_key(r["name"])
        if nk and len(nk) >= 4:
            by_name.setdefault(nk, []).append(r)
    def _ctry(c: str) -> str:
        return "hong kong" if "hong kong" in c or c == "kowloon" else c

    for members in by_name.values():
        for i, a in enumerate(members):
            for b in members[i + 1:]:
                ca, cb = _ctry(a["country"]), _ctry(b["country"])
                same_site = (registrable_domain(a["website"]) or "x") == (registrable_domain(b["website"]) or "y")
                # Same name AND same website is one company even when the two
                # rows list different offices (Healthcare 21 UK: UK and Cork).
                if not ca or not cb or ca == cb or same_site:
                    union(a["row"], b["row"])
    groups: dict[int, list[dict]] = {}
    for r in rows:
        groups.setdefault(find(r["row"]), []).append(r)
    out: dict[int, int] = {}
    for members in groups.values():
        if len(members) < 2:
            continue
        best = max(members, key=lambda m: ((1000 if "(acquired by" in m["name"].lower() else 0) + m["filled"],
                                           -m["row"]))
        out.update({m["row"]: best["row"] for m in members if m is not best})
    return out


def _said(v: dict) -> str:
    """What AI Mode originally claimed about the company (for the flag list)."""
    o = v.get("_orig") or v
    by, st = o.get("acquired_by"), o.get("status") or ""
    if st in ("closed", "absorbed"):
        return f"{st} (acquired by {by})" if by else st
    return f"acquired by {by}" if by else "no claim"


def set_market_definition(ws, stem: str) -> None:
    """Row 2 = this market's own definition, row 3 = who is included.

    The verified files carry only the generic rule sentence in row 2 (the
    verifier rewrites the banner and drops the market's scope), so every
    market of one rule type showed the identical line and no file said what
    ITS market was. The market is found from the file name
    (<slug>_verified_distributors.xlsx); a file whose slug is unknown keeps
    its banner unchanged."""
    import market_batches

    slug = stem.replace("_verified_distributors", "")
    try:
        definition = market_batches.definition_for(slug)
    except KeyError:
        return
    if not definition:
        return
    rule = str(ws["A2"].value or "").strip()
    note = str(ws["A3"].value or "").strip()
    # The verifier writes its banner without registering the market's kind,
    # so every file says the physical-goods rule even where the rows were
    # judged as service / hybrid / software (geo 2026-10-14). Rebuild the
    # rule text from the kind the market was actually judged under.
    try:
        from vendor_intel.pipeline.market_research import discovery as D
        from vendor_intel.pipeline.market_research.export_xlsx import _definition_banner

        _slug, market, _scope, kind = market_batches.by_slug(slug)
        D.set_distributor_definition(market, {
            "distributor": D.DEF_PHYSICAL, "service": D.DEF_SERVICE,
            "hybrid": D.DEF_HYBRID, "software": D.DEF_SOFTWARE}[kind])
        lines = [" ".join(str(t).split()) for t, is_title in _definition_banner(market) if not is_title]
        if lines:
            rule, note = lines[0], " ".join(lines[1:])
    except Exception:  # noqa: BLE001 - keep the verifier's banner rather than fail the build
        pass
    ws["A2"] = f"MARKET DEFINITION: {definition}"
    ws["A3"] = f"WHO IS INCLUDED: {rule} {note}".strip()


def _strip_label(text: object, label: str) -> str | None:
    s = str(text or "").strip()
    return (s[len(label):].strip() if s.startswith(label) else s) or None


def market_name(ws) -> str:
    """The market title from the sheet's first row (its banner)."""
    return str(ws["A1"].value or "").strip()


def _val(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def workbook_json(wb, name: str) -> dict:
    ws = wb.active
    rows = [tuple(_val(c) for c in r) for r in ws.iter_rows(values_only=True)]
    # Rows 1-3 only: row 4 is the styled summary bar, not part of the source text.
    pre = [r[0] for r in rows[:3] if r and r[0] is not None]
    header = [c for c in rows[HEADER_ROW - 1] if c is not None]
    companies = [
        {col: (r[i] if i < len(r) else None) for i, col in enumerate(header)}
        for r in rows[HEADER_ROW:] if any(c is not None for c in r)
    ]
    return {
        "market": pre[0] if pre else None,
        "market_definition": _strip_label(pre[1], "MARKET DEFINITION:") if len(pre) > 1 else None,
        "who_is_included": _strip_label(pre[2], "WHO IS INCLUDED:") if len(pre) > 2 else None,
        "source_file": name,
        "sheet": ws.title,
        "total_companies": len(companies),
        "columns": header,
        "companies": companies,
    }


def build(folder: Path, out: Path, make_zip: bool) -> None:
    verdicts = all_verdicts(folder)
    out.mkdir(parents=True, exist_ok=True)
    report, flag_rows = [], []
    totals = {"kept": 0, "renamed": 0, "removed": 0, "unchecked": 0, "flagged": 0, "dupes": 0}
    for f in input_files(folder):
        wb = openpyxl.load_workbook(f)
        ws = wb.active
        set_market_definition(ws, f.stem)
        header = [str(c.value or "").strip() for c in ws[HEADER_ROW]]
        ci, wi = header.index("Company") + 1, header.index("Website") + 1
        # Product-evidence gate (market_batches.EVIDENCE): a verified company
        # whose own profile never mentions the market's product goes to the
        # flagged list, not the main Excel (soft serve / h2 manual passes).
        import market_batches
        evidence = market_batches.evidence_for(f.stem.replace("_verified_distributors", ""))
        ev_cols = [header.index(c) + 1 for c in ("Core Categories", "Specialty Focus",
                                                  "Key Products / Technologies") if c in header]
        drop = []
        counts = dict.fromkeys(totals, 0)
        row_action: dict[int, str] = {}   # kept / renamed rows, for the duplicate check
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            cell = ws.cell(row=r, column=ci)
            name = str(cell.value or "").strip()
            if not name:
                continue
            v = verdicts.get(company_key(name, ws.cell(row=r, column=wi).value))
            action, acq = label(v)
            # A reviewer correction of the owner (overrides.json with
            # "acquired_by" and "flag": false) wins over an "(acquired by X)"
            # label the name carried in from discovery: replace the label
            # instead of flagging the two owners as a conflict (caps
            # 2026-10-15: Orkila Algerie named "Barentz", owner is Azelis).
            # ...also when the reviewer CONFIRMS the status check's owner
            # against an outdated label (ent 2026-10-15: "Integra (acquired by
            # Scienture)", which sold it to Tollo Health in 2025).
            m_old = re.search(r"\(acquired by (.+)\)\s*$", name, re.I)
            corrected = bool(v and v.get("_orig") and v.get("flag") is False and (
                (v.get("acquired_by") or "") != (v["_orig"].get("acquired_by") or "")
                or (m_old and v.get("acquired_by")
                    and _norm(m_old.group(1)) != _norm(v.get("acquired_by")))))
            if corrected and re.search(r"\(acquired by .+\)\s*$", name, re.I):
                name = re.sub(r"\s*\(acquired by .+\)\s*$", "", name, flags=re.I)
                cell.value = name
            reason = None if corrected else flag_reason(name, v, action, acq)
            if not reason and evidence is not None and action != "removed":
                profile = " | ".join([name] + [str(ws.cell(row=r, column=c).value or "") for c in ev_cols])
                if not evidence.search(profile):
                    reason = ("No evidence of this market's product in its profile "
                              "(Core Categories / Specialty Focus / Key Products).")
            if reason and action != "removed":
                # Not "verified": out of the main file, into the flagged list
                # (shown there as the name it would have had).
                shown = f"{name} (acquired by {acq})" if (
                    action == "renamed" and "acquired by" not in name.lower()) else name
                hq = ws.cell(row=r, column=header.index("Headquarters") + 1).value if "Headquarters" in header else ""
                flag_rows.append([market_name(ws), shown, ws.cell(row=r, column=wi).value,
                                  hq, _said(v), reason])
                drop.append(r)
                counts["flagged"] += 1
                continue
            counts[action] += 1
            if action == "removed":
                drop.append(r)
            else:
                row_action[r] = action
                if action == "renamed" and "acquired by" not in name.lower():
                    cell.value = f"{name} (acquired by {acq})"
            if action in ("removed", "renamed"):
                report.append([f.stem, name, ws.cell(row=r, column=wi).value, v["status"],
                               acq, v.get("year", ""), v.get("evidence", ""), action])
        # The same company listed twice (spelling / website variants): keep one row.
        hq_col = header.index("Headquarters") + 1 if "Headquarters" in header else None
        cand = []
        for r in row_action:
            hq = str(ws.cell(row=r, column=hq_col).value or "") if hq_col else ""
            parts = [p.strip().lower() for p in hq.split(",")]
            cand.append({"row": r, "name": str(ws.cell(row=r, column=ci).value or ""),
                         "website": ws.cell(row=r, column=wi).value,
                         "city": parts[0] if len(parts) > 1 else "", "country": parts[-1] if len(parts) > 1 else "",
                         "filled": sum(1 for c in ws[r] if c.value not in (None, ""))})
        for r, keep in find_duplicates(cand).items():
            counts[row_action[r]] -= 1
            counts["dupes"] += 1
            drop.append(r)
            report.append([f.stem, str(ws.cell(row=r, column=ci).value or ""), ws.cell(row=r, column=wi).value,
                           "duplicate", "", "", f"Same company as '{ws.cell(row=keep, column=ci).value}' "
                           "(same website and country, or same name and city) - listed once.", "duplicate"])
        for r in sorted(set(drop), reverse=True):
            ws.delete_rows(r, 1)
        polish_distributor_sheet(ws, HEADER_ROW)  # presentation only (filter, freeze, look)
        wb.save(out / f.name)
        (out / f"{f.stem}.json").write_text(
            json.dumps(workbook_json(wb, f.name), ensure_ascii=False, indent=2), encoding="utf-8")
        for k in totals:
            totals[k] += counts[k]
        log(f"{f.stem}: kept {counts['kept']}, acquired-renamed {counts['renamed']}, "
            f"removed {counts['removed']}, flagged-out {counts['flagged']}, duplicates merged {counts['dupes']}"
            + (f", UNCHECKED {counts['unchecked']}" if counts["unchecked"] else ""))

    rwb = openpyxl.Workbook()
    rws = rwb.active
    rws.title = "Status changes"
    rws.append(["File", "Company", "Website", "Status", "Acquired by", "Year", "Evidence", "Action"])
    for row in report:
        rws.append(row)
    polish_report_sheet(rws)
    rwb.save(out / "status_check_report.xlsx")

    # Companies kept OUT of the market files because their ownership claim was
    # suspicious: listed here with the reason, so none is lost.
    fwb = openpyxl.Workbook()
    fws = fwb.active
    fws.title = "Flagged for review"
    fws.append(["Market", "Company", "Website", "Headquarters", "AI Mode said", "Why flagged"])
    for row in flag_rows:
        fws.append(row)
    polish_flag_sheet(fws)
    fwb.save(out / "flagged_for_review.xlsx")
    log(f"TOTAL kept {totals['kept']}, acquired-renamed {totals['renamed']}, "
        f"removed {totals['removed']}, flagged-out {totals['flagged']}, duplicates merged {totals['dupes']}, "
        f"unchecked {totals['unchecked']} -> {out}")

    if make_zip:
        zpath = out.with_suffix(".zip")
        with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
            for p in sorted(out.iterdir()):
                z.write(p, f"{out.name}/{p.name}")
        log(f"zip: {zpath}")


def show_status(folder: Path) -> None:
    companies = all_companies(folder)
    verdicts = all_verdicts(folder)
    done = [verdicts[k] for k in companies if k in verdicts]
    by = {}
    for v in done:
        by[v["status"]] = by.get(v["status"], 0) + 1
    acq = sum(1 for v in done if v.get("acquired_by"))
    left = len(companies) - len(done)
    log(f"checked {len(done)}/{len(companies)} | {by} | acquired {acq} | "
        f"{left} left (~{-(-left // CHUNK)} queries)")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True, type=Path, help="folder of *_verified_distributors.xlsx")
    ap.add_argument("--output", type=Path, help="output folder (created)")
    ap.add_argument("--workers", type=int, default=0, help="run N browsers, then build")
    ap.add_argument("--shard", type=int, help=argparse.SUPPRESS)
    ap.add_argument("--shards", type=int, default=1, help=argparse.SUPPRESS)
    ap.add_argument("--status", action="store_true", help="progress only, no queries")
    ap.add_argument("--build", action="store_true", help="write output from saved answers only")
    ap.add_argument("--zip", action="store_true", help="also zip the output folder")
    args = ap.parse_args()
    folder = args.input.resolve()

    if args.status:
        show_status(folder)
        return
    if args.shard is not None:
        run_worker(folder, args.shard, args.shards)
        return
    if args.workers:
        procs = []
        for i in range(args.workers):
            env = dict(os.environ, GOOGLE_AI_MODE_PROFILE_DIR=str(ROOT / "data" / f"ai_mode_status_{i}"))
            procs.append(subprocess.Popen(
                [sys.executable, __file__, "--input", str(folder), "--shard", str(i),
                 "--shards", str(args.workers)], env=env))
            time.sleep(15)  # stagger browser launches
        for p in procs:
            p.wait()
        show_status(folder)
        verdicts = all_verdicts(folder)
        left = sum(k not in verdicts for k in all_companies(folder))
        if left and not args.build:
            # A partial build looks finished and gets sent; build only when
            # every company is checked (or when --build is asked for).
            log(f"{left} companies still unchecked — output NOT built. Re-run to resume.")
            return
    if args.output and (args.workers or args.build):
        build(folder, args.output.resolve(), args.zip)


if __name__ == "__main__":
    main()
