"""Add verified RAJA / Manutan country packaging merchants + other agent-verified distributors."""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "agent_manual_verified_add_audit3.json"
ND = "Not publicly disclosed"

# Evidence: RAJA Group is Europe's leading packaging distributor (raja-group.com / rajapack sites).
# Country sites are RAJA packaging merchants. Contacts not invented.
ROWS = []

def row(**kwargs):
    base = {
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Ownership": kwargs.pop("Ownership", "Subsidiary of RAJA Group"),
        "Employees": kwargs.pop("Employees", ND),
        "Operational Presence": kwargs.pop("Operational Presence", ND),
        "Key Brands Represented": kwargs.pop("Key Brands Represented", "RAJA"),
        "Core Categories": kwargs.pop(
            "Core Categories",
            "Cartons; protective packaging; tapes; stretch/shrink; mailers; warehouse equipment",
        ),
        "Specialty Focus": kwargs.pop(
            "Specialty Focus",
            "Catalog packaging merchant/distributor for business packaging supplies",
        ),
    }
    base.update(kwargs)
    return base


# RAJA country merchants not already in FINAL (existing: AT,CZ,FR,DE,IT,NL,PL,PT,ES,UK,DK + group)
for name, url, cc, geo, region, hq in [
    ("RAJA Belgium", "https://www.rajapack.be", "BE", "EU(Belgium)", "EU", "Belgium"),
    ("RAJA Switzerland", "https://www.rajapack.ch", "CH", "EU(Switzerland)", "EU", "Switzerland"),
    ("RAJA Hungary", "https://www.rajapack.hu", "HU", "EU(Hungary)", "EU", "Hungary"),
    ("RAJA Romania", "https://www.rajapack.ro", "RO", "EU(Romania)", "EU", "Romania"),
    ("RAJA Slovakia", "https://www.rajapack.sk", "SK", "EU(Slovakia)", "EU", "Slovakia"),
    ("RAJA Sweden", "https://www.rajapack.se", "SE", "EU(Sweden)", "EU", "Sweden"),
    ("RAJA Norway", "https://www.rajapack.no", "NO", "EU(Norway)", "EU", "Norway"),
    ("RAJA Ireland", "https://www.rajapack.ie", "IE", "EU(Ireland)", "EU", "Ireland"),
    ("RAJA Luxembourg", "https://www.rajapack.lu", "LU", "EU(Luxembourg)", "EU", "Luxembourg"),
]:
    ROWS.append(
        row(
            Company=name,
            Website=url,
            Founded="1954",  # RAJA Group founding year (public); country ops may differ
            Headquarters=hq,
            **{
                "Continent / Geography": geo,
                "Country Code": cc,
                "Region Code": region,
                "Summary": f"{name} is a RAJA Group packaging merchant/distributor serving businesses with catalog packaging supplies.",
                "_evidence": "raja-group.com — European packaging distributor founded 1954; country rajapack.* merchant sites",
            },
        )
    )

# Manutan country catalogs — industrial/office distributors including packaging (public Manutan group)
for name, url, cc, geo, region, hq in [
    ("Manutan Switzerland", "https://www.manutan.ch", "CH", "EU(Switzerland)", "EU", "Switzerland"),
    ("Manutan Portugal", "https://www.manutan.pt", "PT", "EU(Portugal)", "EU", "Portugal"),
    ("Manutan Hungary", "https://www.manutan.hu", "HU", "EU(Hungary)", "EU", "Hungary"),
    ("Manutan Austria", "https://www.manutan.at", "AT", "EU(Austria)", "EU", "Austria"),
    ("Manutan Czech Republic", "https://www.manutan.cz", "CZ", "EU(Czech Republic)", "EU", "Czech Republic"),
    ("Manutan Sweden", "https://www.manutan.se", "SE", "EU(Sweden)", "EU", "Sweden"),
    ("Manutan Slovakia", "https://www.manutan.sk", "SK", "EU(Slovakia)", "EU", "Slovakia"),
    ("Manutan Romania", "https://www.manutan.ro", "RO", "EU(Romania)", "EU", "Romania"),
    ("Manutan Luxembourg", "https://www.manutan.lu", "LU", "EU(Luxembourg)", "EU", "Luxembourg"),
    ("Manutan Denmark", "https://www.manutan.dk", "DK", "EU(Denmark)", "EU", "Denmark"),
    ("Manutan Norway", "https://www.manutan.no", "NO", "EU(Norway)", "EU", "Norway"),
    ("Manutan Finland", "https://www.manutan.fi", "FI", "EU(Finland)", "EU", "Finland"),
    ("Manutan Ireland", "https://www.manutan.ie", "IE", "EU(Ireland)", "EU", "Ireland"),
]:
    ROWS.append(
        row(
            Company=name,
            Website=url,
            Founded=ND,
            Headquarters=hq,
            Ownership="Subsidiary of Manutan Group",
            **{
                "Continent / Geography": geo,
                "Country Code": cc,
                "Region Code": region,
                "Key Brands Represented": "Manutan",
                "Core Categories": "Industrial supplies; packaging; warehouse equipment; workplace consumables",
                "Specialty Focus": "B2B catalog distributor of industrial and packaging supplies",
                "Summary": f"{name} is a Manutan Group B2B catalog distributor including packaging and warehouse supplies.",
                "_evidence": "manutan.* country sites; Manutan Group industrial/office catalog distributor",
            },
        )
    )

# Other agent-verified
ROWS.extend(
    [
        row(
            Company="Hughes Enterprises (United Hughes / BradyPLUS)",
            Website="https://hughesent.com",
            Founded=ND,
            Headquarters="2151 Green Lane, Levittown, Pennsylvania 19057, USA",
            Ownership="Subsidiary / BradyPLUS company",
            **{
                "Continent / Geography": "NA(United States)",
                "Operational Presence": "United States (historic NJ/PA packaging distribution)",
                "Country Code": "US",
                "Region Code": "NA",
                "Key Brands Represented": "Multi-supplier packaging equipment and consumables",
                "Core Categories": "Industrial packaging supplies; packaging equipment; stretch/shrink; strapping; void fill; tapes",
                "Specialty Focus": "Distributor of packaging supplies and equipment (BradyPLUS / United Hughes)",
                "Summary": "US packaging distributor (Hughes Enterprises / United Hughes) supplying packaging materials and end-of-line equipment; BradyPLUS company.",
                "_evidence": "hughesent.com industrial packaging supply; ProcessRegister distributor profile",
            },
        ),
        row(
            Company="Stockpak",
            Website="https://www.stockpak.co.uk",
            Founded=ND,
            Headquarters="United Kingdom",
            Ownership="Private",
            **{
                "Continent / Geography": "EU(United Kingdom)",
                "Country Code": "GB",
                "Region Code": "EU",
                "Key Brands Represented": "Stockpak",
                "Core Categories": "Jewellery packaging; gift packaging; retail packaging supplies",
                "Specialty Focus": "UK packaging supplier/wholesaler for jewellery and gift packaging",
                "Summary": "UK packaging merchant supplying stock and bespoke jewellery/gift packaging.",
                "_evidence": "stockpak.co.uk alive with packaging supply/wholesale signals",
            },
        ),
        row(
            Company="BigDug",
            Website="https://www.bigdug.co.uk",
            Founded=ND,
            Headquarters="United Kingdom",
            Ownership="Private",
            **{
                "Continent / Geography": "EU(United Kingdom)",
                "Country Code": "GB",
                "Region Code": "EU",
                "Key Brands Represented": "BigDug",
                "Core Categories": "Warehouse equipment; shelving; packaging and storage supplies",
                "Specialty Focus": "UK B2B supplier of warehouse and packaging-related equipment/supplies",
                "Summary": "UK B2B catalog supplier of warehouse equipment and related packaging/storage supplies.",
                "_evidence": "bigdug.co.uk alive with wholesale/catalog supply signals",
            },
        ),
        row(
            Company="MaxPak Australia",
            Website="https://www.maxpak.com.au",
            Founded=ND,
            Headquarters="Australia",
            Ownership="Private",
            **{
                "Continent / Geography": "APAC(Australia)",
                "Country Code": "AU",
                "Region Code": "APAC",
                "Key Brands Represented": "MaxPak",
                "Core Categories": "Packaging supplies; shipping materials",
                "Specialty Focus": "Australian packaging supplies wholesaler/distributor",
                "Summary": "Australian packaging supplies wholesaler/distributor (MaxPak).",
                "_evidence": "maxpak.com.au alive with packaging wholesale/distributor signals",
            },
        ),
        row(
            Company="Ozpack",
            Website="https://www.ozpack.com.au",
            Founded=ND,
            Headquarters="Australia",
            Ownership="Private",
            **{
                "Continent / Geography": "APAC(Australia)",
                "Country Code": "AU",
                "Region Code": "APAC",
                "Key Brands Represented": "Ozpack",
                "Core Categories": "Packaging supplies; shipping materials",
                "Specialty Focus": "Australian packaging supplies distributor",
                "Summary": "Australian packaging supplies distributor (Ozpack).",
                "_evidence": "ozpack.com.au alive with packaging supply signals",
            },
        ),
        row(
            Company="Winc Australia",
            Website="https://www.winc.com.au",
            Founded=ND,
            Headquarters="Australia",
            Ownership="Private",
            **{
                "Continent / Geography": "APAC(Australia)",
                "Country Code": "AU",
                "Region Code": "APAC",
                "Key Brands Represented": "Winc",
                "Core Categories": "Workplace supplies; packaging; office and facility consumables",
                "Specialty Focus": "Australian B2B workplace supplier including packaging",
                "Summary": "Australian B2B workplace distributor (Winc) including packaging supplies.",
                "_evidence": "winc.com.au alive with catalog/wholesale supply signals",
            },
        ),
        row(
            Company="Antalis Germany",
            Website="https://www.antalis.de",
            Founded=ND,
            Headquarters="Germany",
            Ownership="Subsidiary of Antalis",
            **{
                "Continent / Geography": "EU(Germany)",
                "Country Code": "DE",
                "Region Code": "EU",
                "Key Brands Represented": "Antalis",
                "Core Categories": "Paper; packaging; visual communication materials",
                "Specialty Focus": "German Antalis packaging and paper merchant/distributor",
                "Summary": "Antalis Germany — packaging and paper merchant/distributor.",
                "_evidence": "antalis.de; Antalis is a known packaging/paper distributor group",
            },
        ),
        row(
            Company="Ahlsell Denmark",
            Website="https://www.ahlsell.dk",
            Founded=ND,
            Headquarters="Denmark",
            Ownership="Subsidiary of Ahlsell",
            **{
                "Continent / Geography": "EU(Denmark)",
                "Country Code": "DK",
                "Region Code": "EU",
                "Key Brands Represented": "Ahlsell",
                "Core Categories": "Industrial supplies; installation products; related packaging/consumables",
                "Specialty Focus": "Nordic industrial distributor (Ahlsell Denmark)",
                "Summary": "Ahlsell Denmark — B2B industrial distributor including packaging-related supplies.",
                "_evidence": "ahlsell.dk; Ahlsell already present as Nordic industrial distributor family",
            },
        ),
        row(
            Company="Ratioform Spain",
            Website="https://www.ratioform.es",
            Founded=ND,
            Headquarters="Spain",
            Ownership="Subsidiary of Ratioform / RAJA ecosystem",
            **{
                "Continent / Geography": "EU(Spain)",
                "Country Code": "ES",
                "Region Code": "EU",
                "Key Brands Represented": "Ratioform",
                "Core Categories": "Shipping packaging; cartons; protective packaging; tapes",
                "Specialty Focus": "Spanish packaging merchant (Ratioform)",
                "Summary": "Ratioform Spain — packaging merchant/distributor for shipping and protective packaging.",
                "_evidence": "ratioform.es; Ratioform is known DE/EU packaging catalog merchant",
            },
        ),
        row(
            Company="Cactus Containers",
            Website="https://www.cactuscontainers.com",
            Founded=ND,
            Headquarters="United States",
            Ownership="Private",
            **{
                "Continent / Geography": "NA(United States)",
                "Country Code": "US",
                "Region Code": "NA",
                "Key Brands Represented": "Cactus Containers",
                "Core Categories": "Containers; bottles; packaging containers",
                "Specialty Focus": "US packaging containers distributor/supplier",
                "Summary": "US packaging containers distributor (Cactus Containers).",
                "_evidence": "cactuscontainers.com alive with packaging supply/distributor signals",
            },
        ),
    ]
)

HEADERS = [
    "Company",
    "Website",
    "Founded",
    "Headquarters",
    "Continent / Geography",
    "Operational Presence",
    "Ownership",
    "Employees",
    "Core Categories",
    "Specialty Focus",
    "Key Brands Represented",
    "Retail / E-commerce / Both",
    "Distribution Type",
    "Contact Person",
    "Role",
    "Email",
    "LinkedIn",
    "Office No.",
    "Country Code",
    "Region Code",
    "Summary",
]


def domain_of(url: str) -> str:
    try:
        return urlparse(url if "://" in url else "https://" + url).netloc.lower().replace("www.", "")
    except Exception:
        return ""


def main():
    # Fix: Founded 1954 for all RAJA countries is slightly stretchy — keep for group, but note in audit.
    backup = FP.with_name(FP.stem + f"_pre_agent_manual3_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(FP, backup)
    wb = load_workbook(FP)
    ws = wb.active
    names, domains = set(), set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        names.add((r[0] or "").strip().lower())
        domains.add(domain_of(str(r[1] or "")))
    added, skipped = [], []
    for rowd in ROWS:
        d = domain_of(rowd["Website"])
        if rowd["Company"].lower() in names or d in domains:
            skipped.append(rowd["Company"])
            continue
        # For RAJA country rows: prefer NPD founded if we only know group year (avoid over-claim)
        if rowd["Company"].startswith("RAJA ") and rowd.get("Founded") == "1954":
            rowd = dict(rowd)
            rowd["Founded"] = ND
            rowd["Summary"] = rowd["Summary"] + " (RAJA Group founded 1954.)"
        ws.append([rowd[h] for h in HEADERS])
        names.add(rowd["Company"].lower())
        domains.add(d)
        rec = {k: rowd[k] for k in HEADERS}
        rec["_evidence"] = rowd.get("_evidence")
        added.append(rec)
    for col in ws.columns:
        for cell in col:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[col[0].column_letter].width = 22
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(FP)
    audit = {
        "ts": datetime.now().isoformat(),
        "backup": str(backup),
        "added_count": len(added),
        "skipped": skipped,
        "final_rows": ws.max_row - 1,
        "added": added,
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print("added", len(added), "skipped", len(skipped), "final", audit["final_rows"])


if __name__ == "__main__":
    main()
