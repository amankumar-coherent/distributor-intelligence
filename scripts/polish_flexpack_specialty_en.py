"""Polish non-English Specialty Focus rows to English via DeepSeek (evidence-preserving)."""
from __future__ import annotations

import json
import os
import re
from pathlib import Path

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv(r"D:\vendor-intel\.env")

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global\global_flexible_packaging_market_global_FINAL.xlsx"
)
NON_EN = re.compile(
    r"\b(ist ein|Groothandel|Online-Shop|Verpackungsmaterialien|Kartonagen|Folien|Beutel|Etiketten)\b",
    re.I,
)


def main() -> None:
    key = os.getenv("OPENAI_API_KEY")
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    if not key:
        raise SystemExit("no OPENAI_API_KEY")

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers)}
    ci, cci, sfi = ic["Company"], ic["Core Categories"], ic["Specialty Focus"]

    updated = []
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, ci).value
        core = ws.cell(row, cci).value or ""
        focus = str(ws.cell(row, sfi).value or "")
        if not NON_EN.search(focus):
            continue
        prompt = (
            f'Rewrite Specialty Focus into clear ENGLISH (1-2 detailed sentences). '
            f"Keep ONLY facts already present. NO country/region at the end. Do not invent products.\n"
            f"Company: {name}\nCore Categories: {core}\nCurrent Specialty Focus: {focus}\n"
            'Return ONLY JSON: {"focus":"..."}'
        )
        r = httpx.post(
            f"{base}/chat/completions",
            headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
            json={
                "model": model,
                "temperature": 0,
                "messages": [{"role": "user", "content": prompt}],
            },
            timeout=60.0,
        )
        r.raise_for_status()
        content = r.json()["choices"][0]["message"]["content"]
        m = re.search(r"\{[\s\S]*\}", content)
        if not m:
            print("fail parse", name)
            continue
        data = json.loads(m.group(0))
        nf = re.sub(r"\s+", " ", str(data.get("focus") or "")).strip()
        if len(nf) < 40:
            print("too short", name)
            continue
        if nf[-1] not in ".!?":
            nf += "."
        ws.cell(row, sfi).value = nf
        updated.append(name)
        print(name, "->", nf[:140])

    wb.save(FP)
    print("updated", len(updated), "saved", FP)


if __name__ == "__main__":
    main()
