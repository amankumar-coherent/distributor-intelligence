# -*- coding: utf-8 -*-
"""Audit wearable FINAL: multi-location families + website DNS health."""
from __future__ import annotations

import json
import re
import socket
import ssl
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import urllib.request

X = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
    r"\global_wearable_medical_devices_market_global_FINAL.xlsx"
)
OUT = X.parent / "_wearable_website_unique_audit.json"

COUNTRY_TOKENS = {
    "portugal", "polska", "poland", "espana", "españa", "spain", "france",
    "greece", "ireland", "italia", "italy", "nederland", "netherlands", "holland",
    "norge", "norway", "romania", "românia", "uk", "united kingdom", "britain",
    "germany", "deutschland", "usa", "us", "united states", "america", "canada",
    "mexico", "brasil", "brazil", "argentina", "chile", "colombia", "peru",
    "australia", "new zealand", "nz", "japan", "china", "hong kong", "taiwan",
    "korea", "south korea", "india", "singapore", "malaysia", "indonesia",
    "thailand", "vietnam", "philippines", "uae", "saudi", "saudi arabia",
    "qatar", "kuwait", "bahrain", "oman", "egypt", "south africa", "nigeria",
    "kenya", "ghana", "morocco", "turkey", "türkiye", "israel", "russia",
    "ukraine", "sweden", "denmark", "finland", "belgium", "benelux", "austria",
    "switzerland", "czech", "czechia", "hungary", "iberia", "asia pacific",
    "apac", "emea", "latam", "latin america", "middle east", "gcc", "nordic",
    "nordics", "europe", "africa", "asia", "oceania", "international", "global",
    "deutschland gmbh", "france sas", "uk ltd", "australia pty",
}


def strip_geo_suffix(name: str) -> str:
    s = re.sub(r"\s+", " ", (name or "").strip())
    s = re.sub(r"\s*[-–—,/|(].*$", "", s).strip()
    for _ in range(4):
        low = s.lower()
        changed = False
        for tok in sorted(COUNTRY_TOKENS, key=len, reverse=True):
            if low.endswith(" " + tok) or low == tok:
                s = s[: len(s) - len(tok)].rstrip(" -–—,/")
                changed = True
                break
        if not changed:
            break
    return re.sub(r"\s+", " ", s).strip()


def legal_core(name: str) -> str:
    s = strip_geo_suffix(name)
    s = re.sub(
        r"(?i)\s+(Inc\.?|LLC|Ltd\.?|Limited|Corp\.?|Corporation|GmbH|S\.?A\.?|"
        r"S\.?p\.?A\.?|Pty\.?|Pvt\.?|PLC|plc|AG|AB|BV|B\.V\.|KK|Co\.?,?\s*Ltd\.?|"
        r"Company|Holdings|Group|International)\s*$",
        "",
        s,
    ).strip()
    return s


def host_of(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    try:
        return (urlparse(u).hostname or "").lower().lstrip("www.")
    except Exception:
        return ""


def probe(url: str, timeout: float = 8.0) -> dict:
    raw = (url or "").strip()
    if not raw:
        return {"ok": False, "reason": "empty", "final": "", "status": None}
    u = raw if re.match(r"^https?://", raw, re.I) else "https://" + raw
    host = host_of(u)
    dns_ok = False
    try:
        socket.getaddrinfo(host, 443)
        dns_ok = True
    except Exception as e:
        return {"ok": False, "reason": f"dns:{e.__class__.__name__}", "final": "", "status": None, "host": host}

    ctx = ssl.create_default_context()
    last_err = None
    for candidate in [u, u.replace("https://", "http://", 1)]:
        try:
            req = urllib.request.Request(
                candidate,
                headers={"User-Agent": "Mozilla/5.0 (compatible; VendorIntelAudit/1.0)"},
                method="GET",
            )
            with urllib.request.urlopen(req, timeout=timeout, context=ctx) as resp:
                final = resp.geturl()
                code = getattr(resp, "status", None) or resp.getcode()
                body = resp.read(4000).decode("utf-8", errors="ignore").lower()
                parked = any(
                    x in body
                    for x in (
                        "domain is for sale",
                        "buy this domain",
                        "parked free",
                        "sedoparking",
                        "godaddy parking",
                        "this domain may be for sale",
                        "related searches",
                    )
                )
                return {
                    "ok": True and not parked and int(code or 0) < 400,
                    "reason": "parked" if parked else "ok",
                    "final": final,
                    "status": code,
                    "host": host,
                    "dns_ok": dns_ok,
                }
        except Exception as e:
            last_err = f"{e.__class__.__name__}"
            continue
    return {"ok": False, "reason": last_err or "fail", "final": "", "status": None, "host": host, "dns_ok": dns_ok}


def main() -> None:
    wb = openpyxl.load_workbook(X, data_only=True)
    ws = wb.active
    hdr = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    ic = {h: i for i, h in enumerate(hdr) if h}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not str(r[0] or "").strip():
            continue
        rows.append({h: ("" if r[ic[h]] is None else str(r[ic[h]]).strip()) for h in ic})
    wb.close()

    families = defaultdict(list)
    for r in rows:
        key = legal_core(r["Company"]).lower()
        families[key].append(r["Company"])
    multi = {k: v for k, v in families.items() if len(v) > 1}

    # host collisions (same website, different company names)
    by_host = defaultdict(list)
    for r in rows:
        h = host_of(r.get("Website", ""))
        if h:
            by_host[h].append(r["Company"])
    host_multi = {h: v for h, v in by_host.items() if len(set(v)) > 1}

    # brand token families
    brand_re = re.compile(
        r"(?i)^(Alliance Healthcare|Cardinal Health|Medline|Owens\s*&\s*Minor|McKesson|"
        r"Henry Schein|AmerisourceBergen|Cencora|PHOENIX|Phoenix Pharmahandel|"
        r"Mediq|Patterson|Owens and Minor|Thermo Fisher|Fisher Scientific|VWR|"
        r"Suzuken|Alfresa|Toho Holdings|Medtronic|Abbott|ResMed|Masimo|Dexcom|"
        r"Philips|GE Healthcare|Siemens Healthineers|Stryker|Zimmer|Nipro|"
        r"Fresenius|Baxter|BD|Becton Dickinson|Johnson\s*&\s*Johnson|J&J|"
        r"Smith\s*&\s*Nephew|Coloplast|ConvaTec|Alcon|Essilor|Walgreens|CVS)"
    )
    brand_fams = defaultdict(list)
    for r in rows:
        m = brand_re.match(r["Company"].strip())
        if m:
            brand_fams[m.group(1).lower()].append(r["Company"])

    legal_suffix = [
        r["Company"]
        for r in rows
        if re.search(
            r"(?i)\b(Inc\.?|LLC|Ltd\.?|Limited|Corp\.?|Corporation|GmbH)\b",
            r["Company"],
        )
    ]

    print("companies", len(rows))
    print("legal_core multi families", len(multi), "rows", sum(len(v) for v in multi.values()))
    for k, v in sorted(multi.items(), key=lambda x: -len(x[1]))[:40]:
        print(f"  {len(v)} | {k} => {v}")
    print("same-host multi names", len(host_multi))
    for h, v in sorted(host_multi.items(), key=lambda x: -len(x[1]))[:30]:
        print(f"  {len(v)} | {h} => {v[:10]}")
    print("brand families with >1", sum(1 for v in brand_fams.values() if len(v) > 1))
    for k, v in sorted(brand_fams.items(), key=lambda x: -len(x[1])):
        if len(v) > 1:
            print(f"  {len(v)} | {k} => {v}")
    print("names with Inc/LLC/Ltd/Corp/GmbH", len(legal_suffix))

    # probe all websites
    probes = {}
    with ThreadPoolExecutor(max_workers=24) as ex:
        futs = {ex.submit(probe, r.get("Website", "")): r["Company"] for r in rows}
        done = 0
        for fut in as_completed(futs):
            company = futs[fut]
            probes[company] = fut.result()
            done += 1
            if done % 50 == 0:
                print("probed", done)

    bad = {c: p for c, p in probes.items() if not p.get("ok")}
    print("bad websites", len(bad))
    reasons = defaultdict(int)
    for p in bad.values():
        reasons[p.get("reason", "?")] += 1
    print("bad reasons", dict(reasons))

    payload = {
        "count": len(rows),
        "multi_legal_core": multi,
        "host_multi": host_multi,
        "brand_fams": {k: v for k, v in brand_fams.items() if len(v) > 1},
        "legal_suffix_names": legal_suffix,
        "bad_websites": {c: probes[c] for c in sorted(bad)},
        "all_probes": probes,
        "rows": [
            {
                "Company": r["Company"],
                "Website": r.get("Website", ""),
                "Headquarters": r.get("Headquarters", ""),
                "Continent / Geography": r.get("Continent / Geography", ""),
            }
            for r in rows
        ],
    }
    OUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print("wrote", OUT)


if __name__ == "__main__":
    main()
