"""Agent-verified packaging DISTRIBUTOR rows — public evidence only; no invented contacts."""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "agent_manual_verified_add_audit.json"
ND = "Not publicly disclosed"

ROWS = [
    {
        "Company": "BOX Partners",
        "Website": "https://www.boxpartners.com",
        "Founded": "1989",
        "Headquarters": "2650 Galvin Drive, Elgin, Illinois 60124, USA",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "Nationwide US distribution centers",
        "Ownership": "Private (Charlesbank Capital Partners-backed)",
        "Employees": "51–200",
        "Core Categories": "Packaging; shipping supplies; industrial supplies; poly bags; mailers; labels; tape; boxes",
        "Specialty Focus": "National wholesaler of packaging/shipping supplies for distributors and resellers",
        "Key Brands Represented": "Aviditi; BOX USA; multi-brand wholesale assortment",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": "https://www.linkedin.com/company/boxpartners",
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "US national packaging wholesaler founded 1989; HQ Elgin IL; supplies distributors/resellers.",
        "_evidence": "Mergr/LinkedIn/boxpartners.com — distributor/wholesaler founded 1989 Elgin",
    },
    {
        "Company": "BASCO",
        "Website": "https://www.bascousa.com",
        "Founded": "1946",
        "Headquarters": "University Park, Illinois, USA",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "USA nationwide",
        "Ownership": "Private",
        "Employees": ND,
        "Core Categories": "Industrial containers; drums; pails; IBC totes; packaging and shipping supplies",
        "Specialty Focus": "Wholesale distributor of industrial packaging containers and shipping supplies",
        "Key Brands Represented": "BASCO",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": "https://www.linkedin.com/company/bascousa",
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "US industrial packaging distributor founded 1946; HQ University Park IL.",
        "_evidence": "bascousa.com/about-us founded July 20 1946",
    },
    {
        "Company": "Stamar Packaging",
        "Website": "https://stamarpackaging.com",
        "Founded": "1981",
        "Headquarters": "United States",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "United States",
        "Ownership": "Private",
        "Employees": ND,
        "Core Categories": "Corrugated boxes; poly bags; shrink packaging; tapes; packaging equipment; industrial supplies",
        "Specialty Focus": "Distributor of packaging, shipping, and industrial supplies",
        "Key Brands Represented": "Multi-supplier packaging assortment",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "US packaging distributor since 1981 offering corrugated, films, bags, and packaging equipment.",
        "_evidence": "stamarpackaging.com FAQ since 1981; trusted distributor wording",
    },
    {
        "Company": "Excel Pac LLC",
        "Website": "https://excelpacllc.com",
        "Founded": ND,
        "Headquarters": "San Diego, California, USA",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "United States (San Diego based; ships nationwide)",
        "Ownership": "Private",
        "Employees": ND,
        "Core Categories": "Corrugated boxes; stretch film; tape; mailers; poly bags; foodservice; janitorial; cold-chain; PPE",
        "Specialty Focus": "San Diego industrial packaging and shipping supplies distributor",
        "Key Brands Represented": "Multi-brand packaging supplies",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "San Diego packaging distributor supplying industrial shipping materials nationwide.",
        "_evidence": "excelpacllc.com/about — San Diego distributor",
    },
    {
        "Company": "EmpaSoluciones",
        "Website": "https://empasoluciones.com.mx",
        "Founded": ND,
        "Headquarters": "Tijuana, Baja California, Mexico",
        "Continent / Geography": "LATAM(Mexico)",
        "Operational Presence": "Mexico nationwide",
        "Ownership": "Private",
        "Employees": "1–50",
        "Core Categories": "Industrial packaging; plastic; carton; paper; eco packaging",
        "Specialty Focus": "Mexican wholesale distributor of industrial packaging materials",
        "Key Brands Represented": "Multi-manufacturer sourced packaging",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "MX",
        "Region Code": "LATAM",
        "Summary": "Mexico packaging wholesaler/distributor based in Tijuana serving industrial customers nationwide.",
        "_evidence": "empasoluciones.com.mx — comercializacion y distribucion; wholesale envases",
    },
    {
        "Company": "Adelaide Packaging Supplies",
        "Website": "https://www.adpack.com.au",
        "Founded": "1991",
        "Headquarters": "Gillman / Adelaide, South Australia, Australia",
        "Continent / Geography": "APAC(Australia)",
        "Operational Presence": "South Australia",
        "Ownership": "Private (family-owned)",
        "Employees": ND,
        "Core Categories": "Packaging supplies; office supplies",
        "Specialty Focus": "Family-owned packaging and office supply wholesaler in South Australia",
        "Key Brands Represented": "Multi-brand packaging assortment",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "AU",
        "Region Code": "APAC",
        "Summary": "Adelaide packaging wholesaler founded 1991; packaging and office supplies for SA businesses.",
        "_evidence": "adpack.com.au/about-us — 30 years in 2021 => 1991",
    },
    {
        "Company": "Industrial & Shippers Supply",
        "Website": "https://issipackaging.com",
        "Founded": "1954",
        "Headquarters": "Houston, Texas, USA",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "Gulf Coast / Houston region",
        "Ownership": "Private",
        "Employees": ND,
        "Core Categories": "Packaging equipment; shipping supplies; labeling; stretch/shrink; industrial packaging",
        "Specialty Focus": "Houston packaging distributor of supplies, equipment, and service",
        "Key Brands Represented": "Multi-brand packaging equipment and consumables",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "Houston packaging distributor founded 1954 supplying industrial shipping materials and equipment.",
        "_evidence": "issipackaging.com — original packaging distributor Houston founded 1954",
    },
    {
        "Company": "Fillmore Container",
        "Website": "https://www.fillmorecontainer.com",
        "Founded": ND,
        "Headquarters": "United States",
        "Continent / Geography": "NA(United States)",
        "Operational Presence": "United States",
        "Ownership": "Private",
        "Employees": ND,
        "Core Categories": "Glass bottles; jars; closures; packaging containers",
        "Specialty Focus": "Distributor/supplier of bottles, jars, and packaging containers",
        "Key Brands Represented": "Multi-brand containers and closures",
        "Retail / E-commerce / Both": "Both",
        "Distribution Type": "Specialty Distributor",
        "Contact Person": ND,
        "Role": ND,
        "Email": ND,
        "LinkedIn": ND,
        "Office No.": ND,
        "Country Code": "US",
        "Region Code": "NA",
        "Summary": "US packaging container distributor supplying bottles, jars, and closures.",
        "_evidence": "fillmorecontainer.com — packaging container supplier/distributor; founded year not confirmed on About",
    },
]

HEADERS = [
    "Company",
    "Website",
    "Founded",
    "Headquarters",
    "Continent / Geography",
    "Operational Presence",
    "Ownership",
    "Employees",
    "Core Categories",
    "Specialty Focus",
    "Key Brands Represented",
    "Retail / E-commerce / Both",
    "Distribution Type",
    "Contact Person",
    "Role",
    "Email",
    "LinkedIn",
    "Office No.",
    "Country Code",
    "Region Code",
    "Summary",
]


def domain_of(url: str) -> str:
    try:
        return urlparse(url if "://" in url else "https://" + url).netloc.lower().replace("www.", "")
    except Exception:
        return ""


def main():
    backup = FP.with_name(FP.stem + f"_pre_agent_manual_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    shutil.copy2(FP, backup)
    wb = load_workbook(FP)
    ws = wb.active
    names, domains = set(), set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        names.add((r[0] or "").strip().lower())
        domains.add(domain_of(str(r[1] or "")))
    added, skipped = [], []
    for row in ROWS:
        d = domain_of(row["Website"])
        if row["Company"].lower() in names or d in domains:
            skipped.append(row["Company"])
            continue
        ws.append([row[h] for h in HEADERS])
        names.add(row["Company"].lower())
        domains.add(d)
        rec = {k: row[k] for k in HEADERS}
        rec["_evidence"] = row.get("_evidence")
        added.append(rec)
    for col in ws.columns:
        for cell in col:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[col[0].column_letter].width = 22
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(FP)
    audit = {
        "ts": datetime.now().isoformat(),
        "backup": str(backup),
        "added": added,
        "skipped_existing": skipped,
        "final_rows": ws.max_row - 1,
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print("added", len(added), "skipped", skipped, "final", audit["final_rows"])


if __name__ == "__main__":
    main()
