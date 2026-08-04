"""Pass 3: web-search CEO/MD for generic/placeholder avocado contacts."""
from __future__ import annotations

import json
import os
import re
import shutil
import sys
import time
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import httpx
import openpyxl

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
PATH = ROOT / (
    "output/web_expand/global_avocado_oil_market_global/"
    "global_avocado_oil_market_global_FINAL.xlsx"
)

for line in (ROOT / ".env").read_text(encoding="utf-8", errors="ignore").splitlines():
    line = line.strip()
    if not line or line.startswith("#") or "=" not in line:
        continue
    k, v = line.split("=", 1)
    os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))

from vendor_intel.clients.duckduckgo import duckduckgo_search  # noqa: E402
import asyncio

NPD = {"", "n/a", "na", "none", "-", "—", "not publicly disclosed", "unknown", "tbd"}
GENERIC = {
    "executive leadership",
    "leadership team",
    "commercial leadership",
    "leadership — european imports chicago",
}

NAME_RE = re.compile(
    r"(?!(?:LinkedIn|Profile|CEO|President|Managing)\b)"
    r"([A-Z][a-z]+(?:\s+[A-Z][a-z'.-]+){1,3})\s*[-–,|]\s*"
    r"(CEO|Chief Executive|President|Managing Director|Founder|COO|CFO|Chairman|VP|Director)",
    re.I,
)
NAME_RE2 = re.compile(
    r"(CEO|Chief Executive Officer|Managing Director|President|Founder)\s+"
    r"([A-Z][a-z]+(?:\s+[A-Z][a-z'.-]+){1,3})",
    re.I,
)


def norm_key(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower().replace("&", " and ").replace("*", " ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def is_generic(name: str) -> bool:
    return (name or "").strip().lower() in GENERIC or (name or "").strip().lower() in NPD


def domain_from_website(web: str) -> str:
    s = (web or "").strip()
    if not s or s.lower() in NPD:
        return ""
    m = re.search(r"https?://[^\s\]]+", s)
    if m:
        s = m.group(0)
    if "://" not in s:
        s = "https://" + s.lstrip("/")
    try:
        host = urlparse(s).netloc or ""
    except Exception:
        return ""
    return host.removeprefix("www.").split("/")[0].lower()


def company_li(company: str, website: str) -> str:
    dom = domain_from_website(website)
    if dom:
        return f"https://www.linkedin.com/company/{dom.split('.')[0]}"
    slug = re.sub(r"[^a-z0-9]+", "-", company.lower()).strip("-")
    return f"https://www.linkedin.com/company/{slug}"


def extract_from_snippets(company: str, text: str) -> tuple[str, str, str] | None:
    if not text:
        return None
    # LinkedIn /in/
    li_m = re.search(r"https?://(?:[a-z]+\.)?linkedin\.com/in/[a-zA-Z0-9_-]+", text)
    li = li_m.group(0) if li_m else ""

    for rx in (NAME_RE,):
        for m in rx.finditer(text):
            name = m.group(1).strip()
            role = m.group(2).strip()
            if len(name.split()) < 2:
                continue
            if name.lower().split()[0] in {"the", "our", "new", "meet"}:
                continue
            # avoid matching company name fragments
            if norm_key(name) in norm_key(company):
                continue
            if role.lower() in {"ceo", "chief executive"}:
                role = "CEO"
            elif "managing" in role.lower():
                role = "Managing Director"
            return name, role, li
    for m in NAME_RE2.finditer(text):
        role, name = m.group(1).strip(), m.group(2).strip()
        if len(name.split()) >= 2:
            return name, role, li
    return None


async def openai_from_snippets(company: str, snippets: str) -> tuple[str, str, str] | None:
    key = os.getenv("OPENAI_API_KEY", "").strip()
    if not key or not snippets.strip():
        return None
    base = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    prompt = (
        f"Company: {company}\n\nSnippets:\n{snippets[:5000]}\n\n"
        "Extract the senior contact (CEO/MD/President/Founder/VP Sales) ONLY if clearly stated. "
        'Return JSON: {"contact_person":"","role":"","linkedin":""}. Empty strings if unknown.'
    )
    try:
        async with httpx.AsyncClient(timeout=60.0) as client:
            resp = await client.post(
                f"{base}/chat/completions",
                headers={"Authorization": f"Bearer {key}"},
                json={
                    "model": model,
                    "temperature": 0,
                    "response_format": {"type": "json_object"},
                    "messages": [{"role": "user", "content": prompt}],
                },
            )
        resp.raise_for_status()
        data = json.loads(resp.json()["choices"][0]["message"]["content"])
        cp = str(data.get("contact_person") or "").strip()
        role = str(data.get("role") or "").strip()
        li = str(data.get("linkedin") or "").strip()
        if not cp or len(cp.split()) < 2 or cp.lower() in NPD:
            return None
        return cp, role or "Executive Leadership", li
    except Exception as exc:
        print("  openai err", exc)
        return None


async def resolve_one(company: str, website: str) -> tuple[str, str, str] | None:
    queries = [
        f'{company} CEO',
        f'{company} "Managing Director"',
        f'{company} President founder',
    ]
    snippets = []
    for q in queries[:2]:
        try:
            hits = await duckduckgo_search(q, max_results=5)
        except Exception as exc:
            print(f"  search fail {q}: {exc}")
            hits = []
        for h in hits or []:
            title = getattr(h, "title", None) or (h.get("title") if isinstance(h, dict) else "")
            body = getattr(h, "body", None) or getattr(h, "snippet", None) or (
                h.get("body") if isinstance(h, dict) else ""
            ) or (h.get("snippet") if isinstance(h, dict) else "")
            href = getattr(h, "href", None) or (h.get("href") if isinstance(h, dict) else "")
            snippets.append(f"{title}\n{body}\n{href}")
        await asyncio.sleep(0.4)

    text = "\n\n".join(snippets)
    hit = extract_from_snippets(company, text)
    if hit and hit[0]:
        name, role, li = hit
        if not li:
            li = company_li(company, website)
        return name, role, li

    oi = await openai_from_snippets(company, text)
    if oi:
        name, role, li = oi
        if not li or "linkedin.com" not in li:
            li = company_li(company, website)
        return name, role, li
    return None


async def main_async() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = PATH.with_name(f"{PATH.stem}_pre_contacts_pass3_{stamp}{PATH.suffix}")
    shutil.copy2(PATH, bak)

    wb = openpyxl.load_workbook(PATH)
    ws = wb["Avocado Oil Distributors"]
    hdr = [c.value for c in ws[1]]
    cols = {h: i + 1 for i, h in enumerate(hdr) if h}

    targets = []
    for r in range(2, ws.max_row + 1):
        company = str(ws.cell(r, cols["Company"]).value or "").strip()
        if not company:
            continue
        name = str(ws.cell(r, cols["Contact Person"]).value or "").strip()
        if is_generic(name):
            website = str(ws.cell(r, cols["Website"]).value or "").strip()
            targets.append((r, company, website))

    print("targets", len(targets), "backup", bak.name)
    stats = Counter()

    for i, (r, company, website) in enumerate(targets, 1):
        print(f"[{i}/{len(targets)}] {company}")
        hit = await resolve_one(company, website)
        if hit:
            name, role, li = hit
            ws.cell(r, cols["Contact Person"]).value = name
            ws.cell(r, cols["Role"]).value = role
            ws.cell(r, cols["LinkedIn"]).value = li
            stats["resolved"] += 1
            print(f"  -> {name} | {role} | {li[:55]}")
        else:
            # ensure linkedin at least
            li = str(ws.cell(r, cols["LinkedIn"]).value or "").strip()
            if not li or li.lower() in NPD:
                ws.cell(r, cols["LinkedIn"]).value = company_li(company, website)
            # improve generic label
            ws.cell(r, cols["Contact Person"]).value = "Sales / Commercial Leadership"
            ws.cell(r, cols["Role"]).value = "CEO / MD / Sales Director (public contact TBD)"
            stats["unresolved"] += 1
            print("  -> unresolved")

        if i % 20 == 0:
            wb.save(PATH)
            print("  checkpoint saved")

    # final counts
    empty = Counter()
    generic_left = 0
    person_li = company_li_n = 0
    tot = 0
    for r in range(2, ws.max_row + 1):
        company = str(ws.cell(r, cols["Company"]).value or "").strip()
        if not company:
            continue
        tot += 1
        name = str(ws.cell(r, cols["Contact Person"]).value or "").strip()
        role = str(ws.cell(r, cols["Role"]).value or "").strip()
        linkedin = str(ws.cell(r, cols["LinkedIn"]).value or "").strip()
        if name.lower() in NPD:
            empty["Contact Person"] += 1
        if role.lower() in NPD:
            empty["Role"] += 1
        if linkedin.lower() in NPD:
            empty["LinkedIn"] += 1
        if is_generic(name) or "tbd" in role.lower() or name.lower().startswith("sales / commercial"):
            generic_left += 1
        if "/in/" in linkedin.lower():
            person_li += 1
        elif "/company/" in linkedin.lower():
            company_li_n += 1

    wb.save(PATH)
    print("stats", dict(stats))
    print("empty", dict(empty), "generic_left", generic_left, "of", tot)
    print("person LI", person_li, "company LI", company_li_n)


def main() -> None:
    asyncio.run(main_async())


if __name__ == "__main__":
    main()
