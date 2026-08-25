"""Fill remaining Founded years via About + DDGS + DeepSeek evidence gate.

NO Google AI scraper. DeepSeek may only return a year that appears in the
provided evidence with founded/established/since context.
"""
from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
MISS = FP.parent / "founded_missing_remaining.json"
CKPT = FP.parent / "founded_fill_remaining61_ckpt.json"
AUDIT = FP.parent / "founded_fill_remaining61_audit.json"

API_KEY = (os.getenv("OPENAI_API_KEY") or "").strip()
BASE_URL = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 VendorIntelFounded/4.0"
YEAR_RE = re.compile(r"\b((?:18|19|20)\d{2})\b")
STRONG = re.compile(
    r"(?:founded|established|est\.?|since|incorporated|gegr[uü]ndet|seit|"
    r"opgericht|fond[eée]|cr[eée][eée]|in business since|family[- ]owned since|"
    r"serving.{0,40}since|commenced trading|registered on|incorporation|"
    r"over \d+ years)\D{0,40}\b((?:18|19|20)\d{2})\b|"
    r"\b((?:18|19|20)\d{2})\b\D{0,25}(?:founded|established|gegr[uü]ndet|seit)",
    re.I,
)

# Agent-confirmed from About / registry / Mergr / company press (this pass)
SEED: dict[str, str] = {
    "American Packaging Distributors": "1988",
    "Carolina Retail Packaging": "1973",
    "Global Trading": "1979",  # Pomanda: Global Trading U K Ltd (gtuk.co.uk / Preston)
    "Packaging Products Corp": "1968",  # Mergr / MSCP: packagingproductscorp.com
    "TransPack Krumbach": "1991",  # company press: gegründet 1991
    "The Packaging Company Canada": "2016",  # LinkedIn TPC Brampton / thepackagingcompany
    "Network Distribution": "1968",  # us.networkdistribution.com/about November 1968
    "Mr Box Online": "1977",  # A&M Tape & Packaging / MrBoxOnline
    "Packaging Direct B.V.": "2011",  # Oozo/KVK PackagingDirect.nl B.V.
    "The Packaging Company": "2016",  # LinkedIn thepackagingcompany.com
    "Packaging Direct New Zealand": "2017",  # NZ Companies Office
    "Packaging Plus": "1994",  # NZ Companies Office / packagingplus.co.nz
    "The Packaging Source": "1984",  # About since April 1984
    "Zynpak Packaging": "1985",  # Frasers year established
    "Package Design & Supply": "1995",  # Mergr / NY DOS Sep 1995
    "Diamond Packaging": "2019",  # Companies House NI666224 diamondpackaging.co.uk
    "The Packaging People": "2008",  # ABR Coffee Bags Pty Ltd 2008-02-01 / LinkedIn
}


def fetch_about(website: str) -> str:
    if not website:
        return ""
    base = website if website.startswith("http") else "https://" + website
    base = base.rstrip("/")
    chunks = []
    for p in ["", "/about", "/about-us", "/history", "/company", "/unternehmen", "/ueber-uns", "/a-propos"]:
        try:
            req = Request(base + p, headers={"User-Agent": UA})
            with urlopen(req, timeout=12) as resp:
                html = resp.read(120_000).decode("utf-8", errors="ignore")
            text = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>|<[^>]+>", " ", html, flags=re.I)
            text = re.sub(r"\s+", " ", text)
            chunks.append(text[:4500])
            if len(chunks) >= 3:
                break
        except Exception:
            continue
        time.sleep(0.1)
    return "\n".join(chunks)[:10000]


def ddgs_text(query: str) -> str:
    try:
        from ddgs import DDGS
    except Exception as e:
        return f"[ddgs_error] {e}"
    lines = []
    try:
        with DDGS() as ddg:
            for r in ddg.text(query, max_results=8):
                lines.append(f"{r.get('title','')} | {r.get('body','')} | {r.get('href','')}")
    except Exception as e:
        return f"[ddgs_error] {e}"
    return "\n".join(lines)[:6500]


def strong_hits(text: str) -> list[tuple[str, str]]:
    out, seen = [], set()
    for m in STRONG.finditer(text or ""):
        y = m.group(1) or m.group(2)
        if not y or y in seen:
            continue
        yi = int(y)
        if yi < 1800 or yi > 2026:
            continue
        seen.add(y)
        sn = text[max(0, m.start() - 50) : min(len(text), m.end() + 50)].strip()
        out.append((y, sn))
    return out


def deepseek_pick(company: str, website: str, evidence: str) -> dict:
    if not API_KEY or not evidence.strip():
        return {"year": None, "confidence": "low", "reason": "no_api_or_evidence"}
    prompt = f"""Extract the founding/establishment year for company "{company}" (website {website}).

Rules:
- Use ONLY the EVIDENCE below.
- Return a year ONLY if evidence clearly says this company was founded/established/incorporated/since YEAR.
- Ignore copyright years, phone numbers, unrelated companies, and "tracked since" analytics dates.
- If unclear or multiple conflicting years, return year=null.

Return JSON only: {{"year":"YYYY"|null,"confidence":"high"|"medium"|"low","snippet":"...","reason":"..."}}

EVIDENCE:
{evidence[:7500]}
"""
    body = json.dumps(
        {
            "model": MODEL,
            "messages": [
                {
                    "role": "system",
                    "content": "Extract founding years only from evidence. Never invent. JSON only.",
                },
                {"role": "user", "content": prompt},
            ],
            "temperature": 0,
        }
    ).encode()
    req = Request(
        BASE_URL + "/chat/completions",
        data=body,
        headers={"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urlopen(req, timeout=70) as resp:
            data = json.loads(resp.read().decode())
        content = data["choices"][0]["message"]["content"].strip()
        if content.startswith("```"):
            content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content)
        return json.loads(content)
    except Exception as e:
        return {"year": None, "confidence": "low", "reason": f"api_error:{e}"}


def resolve(company: str, website: str) -> dict:
    if company in SEED:
        return {"company": company, "year": SEED[company], "source": "seed", "status": "filled"}

    about = fetch_about(website)
    domain = urlparse(website if "://" in website else "https://" + website).netloc.replace("www.", "")
    search = ""
    for q in [
        f'"{company}" founded OR established OR since OR gegründet OR opgericht OR fondé',
        f'"{company}" {domain} founded OR established OR incorporated',
    ]:
        search += "\n" + ddgs_text(q)
        time.sleep(1.0)

    evidence = f"ABOUT:\n{about}\n\nSEARCH:\n{search}"
    about_hits = strong_hits(about)
    search_hits = strong_hits(search)

    if len(about_hits) == 1:
        y, sn = about_hits[0]
        return {"company": company, "year": y, "source": "about:" + sn[:140], "status": "filled"}

    agree = {y for y, _ in about_hits} & {y for y, _ in search_hits}
    if len(agree) == 1:
        y = next(iter(agree))
        return {"company": company, "year": y, "source": "about+search_agree", "status": "filled"}

    pick = deepseek_pick(company, website, evidence)
    year = pick.get("year")
    conf = (pick.get("confidence") or "").lower()
    if year and re.fullmatch(r"(?:18|19|20)\d{2}", str(year)) and str(year) in evidence and conf in ("high", "medium"):
        return {
            "company": company,
            "year": str(year),
            "source": f"deepseek:{conf}:" + str(pick.get("snippet") or "")[:140],
            "status": "filled",
            "pick": pick,
        }

    # Do NOT auto-accept lone DDGS hits (too many false positives, e.g. Bidvest 2002).
    # DeepSeek gate above already covers search evidence when confidence is high/medium.

    return {
        "company": company,
        "year": None,
        "source": "unverified",
        "status": "missing",
        "about_hits": about_hits[:4],
        "search_hits": search_hits[:4],
        "pick": pick,
    }


def main() -> None:
    missing = json.loads(MISS.read_text(encoding="utf-8"))
    ckpt = json.loads(CKPT.read_text(encoding="utf-8")) if CKPT.exists() else {"results": {}}
    results = ckpt.get("results") or {}

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    fi = headers.index("Founded") + 1
    row_by = {str(ws.cell(r, ci).value or "").strip(): r for r in range(2, ws.max_row + 1)}

    filled_now = 0
    for i, m in enumerate(missing, 1):
        company = str(m["company"])
        website = str(m.get("website") or "")
        row = row_by.get(company)
        if not row:
            continue
        cur = str(ws.cell(row, fi).value or "").strip()
        if cur.isdigit():
            results[company] = {"company": company, "year": cur, "status": "already"}
            continue
        if company in results and results[company].get("year"):
            # re-apply prior success
            ws.cell(row, fi).value = results[company]["year"]
            filled_now += 1
            continue

        print(f"[{i}/{len(missing)}] {company}", flush=True)
        rec = resolve(company, website)
        results[company] = rec
        if rec.get("year"):
            ws.cell(row, fi).value = rec["year"]
            filled_now += 1
            print(f"  -> {rec['year']} ({rec.get('source','')[:80]})", flush=True)
        else:
            print("  -> NONE", flush=True)

        if i % 5 == 0:
            wb.save(FP)
            CKPT.write_text(
                json.dumps({"results": results, "updated": datetime.now().isoformat(timespec="seconds")}, indent=2),
                encoding="utf-8",
            )

    wb.save(FP)
    CKPT.write_text(
        json.dumps({"results": results, "updated": datetime.now().isoformat(timespec="seconds")}, indent=2),
        encoding="utf-8",
    )

    total_filled = sum(
        1 for r in range(2, ws.max_row + 1) if str(ws.cell(r, fi).value or "").strip().isdigit()
    )
    still = [
        str(ws.cell(r, ci).value)
        for r in range(2, ws.max_row + 1)
        if not str(ws.cell(r, fi).value or "").strip().isdigit()
    ]
    audit = {
        "filled_now": filled_now,
        "total_filled": total_filled,
        "still_missing": len(still),
        "still_names": still,
        "results": results,
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(f"DONE filled_now={filled_now} total_filled={total_filled} still={len(still)}")


if __name__ == "__main__":
    main()
