"""Fill / normalize Employees column in electrical FINAL Excel to standard ranges."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from vendor_intel.enrichment.known_employees import NPD, resolve_employees, sanitize_employees

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)


def main() -> None:
    backup = XLSX.with_name(XLSX.stem + "_BEFORE_EMPLOYEES_FILL.xlsx")
    if not backup.exists():
        backup.write_bytes(XLSX.read_bytes())

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    ei = headers.index("Employees") + 1

    changes: list[dict[str, str]] = []
    ranged = 0
    still_npd = 0
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        old = str(ws.cell(row_idx, ei).value or "").strip()
        new = sanitize_employees(name, old)
        if new != old:
            ws.cell(row_idx, ei).value = new
            changes.append(
                {
                    "company": name,
                    "old": old or "(empty)",
                    "new": new,
                    "source": "registry" if resolve_employees(name) else "normalize",
                }
            )
        if new == NPD:
            still_npd += 1
        else:
            ranged += 1

    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(XLSX)

    dump = XLSX.parent / "_employees_fill_audit.json"
    dump.write_text(
        json.dumps(
            {
                "when": datetime.now().isoformat(timespec="seconds"),
                "changed": len(changes),
                "with_range": ranged,
                "still_npd": still_npd,
                "changes": changes,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(f"changed {len(changes)} employee values")
    print(f"with_range={ranged} still_npd={still_npd}")
    for c in changes[:25]:
        print(f"  {c['company']}: {c['old']!r} -> {c['new']!r} ({c['source']})")
    print(f"wrote {XLSX}")
    print(f"audit {dump}")


if __name__ == "__main__":
    main()
