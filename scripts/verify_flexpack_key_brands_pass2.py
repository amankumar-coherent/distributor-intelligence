# -*- coding: utf-8 -*-
"""Second pass: extract named brands from NPD rows when evidence contains brand signals."""
from __future__ import annotations

import json
import os
import re
import time
from pathlib import Path
from urllib.parse import urljoin

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv(r"D:\vendor-intel\.env")

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global\global_flexible_packaging_market_global_FINAL.xlsx"
)
CKPT = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global\key_brands_ckpt.json"
)
NPD = "Not publicly disclosed"

KNOWN = re.compile(
    r"\b(3M|Sealed Air|Pregis|Ranpak|STORO?pack|Intertape|IPG|Lantech|Fromm|Mosca|Clysar|"
    r"Dart|Solo|Huhtamaki|Pactiv|Georgia-Pacific|Kimberly-Clark|Dixie|Ecolab|Diversey|"
    r"Jiffy|Scotch|Tesa|Brady|Rubbermaid|Master Lock|Berry Global|Smurfit|Shurtape|"
    r"Duck Brand|Instapak|Cryovac|Autobag|CMPC|Sappi|Detpak|Robopac|Shanklin|GOJO|"
    r"Essity|Clorox|Spartan|Hoffmaster|Tennant)\b",
    re.I,
)

EXTRA_HAND = {
    "Gertex Solutions": "Sealed Air; Robopac; Shanklin; Heat Seal; SMIpack; Fox Packaging; others",
    "Imperial Bag & Paper": "3M; Dart; P&G Professional; Essity; GP Pro; Rubbermaid; GOJO; Hoffmaster; Spartan Chemical; Clorox; Tennant; others",
    "Ball & Doggett": "CMPC; Sappi; others",
}


def fetch(url: str) -> str:
    if not url or not str(url).startswith("http"):
        return ""
    headers = {"User-Agent": "Mozilla/5.0", "Accept": "text/html"}
    chunks = []
    try:
        with httpx.Client(follow_redirects=True, timeout=14.0, headers=headers) as client:
            for path in (
                "",
                "/brands",
                "/our-brands",
                "/partners",
                "/suppliers",
                "/vendors",
                "/about",
                "/about-us",
                "/products",
            ):
                try:
                    u = url if not path else urljoin(url if url.endswith("/") else url + "/", path.lstrip("/"))
                    r = client.get(u)
                    if r.status_code >= 400:
                        continue
                    html = r.text
                    alts = " ".join(
                        re.findall(r'(?:alt|title)=["\']([^"\']{2,80})["\']', html, flags=re.I)
                    )
                    html2 = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", html)
                    html2 = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", html2)
                    html2 = re.sub(r"(?is)<[^>]+>", " ", html2)
                    html2 = re.sub(r"\s+", " ", html2)
                    text = (alts + " " + html2)[:7000]
                    if len(text) > 400:
                        chunks.append(text)
                    if path and len(chunks) >= 2:
                        break
                except Exception:
                    continue
    except Exception:
        return ""
    return "\n---\n".join(chunks)[:12000]


def extract(company: str, website: str, evidence: str) -> str | None:
    key = os.getenv("OPENAI_API_KEY")
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    if not key or len(evidence) < 250:
        return None
    if not KNOWN.search(evidence) and not re.search(
        r"(?i)(our brands|brand partners|authori[sz]ed|we (?:carry|stock|distribute|represent)|featured manufacturers|supplier partners)",
        evidence,
    ):
        return None
    prompt = f"""Extract manufacturer brand names that packaging distributor "{company}" ({website}) sells/distributes/represents.
ONLY brands explicitly in Evidence. No invention. No product categories as brands.
Format: Brand1; Brand2; Brand3; others
If none: conf=low brands=
JSON only: {{"brands":"...","conf":"high|medium|low"}}
Evidence:
{evidence[:8500]}
"""
    try:
        with httpx.Client(timeout=60.0) as client:
            r = client.post(
                f"{base}/chat/completions",
                headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
                json={
                    "model": model,
                    "temperature": 0,
                    "messages": [{"role": "user", "content": prompt}],
                },
            )
            r.raise_for_status()
            content = r.json()["choices"][0]["message"]["content"]
            m = re.search(r"\{[\s\S]*\}", content)
            if not m:
                return None
            data = json.loads(m.group(0))
            if str(data.get("conf") or "").lower() == "low":
                return None
            brands = re.sub(r"\s+", " ", str(data.get("brands") or "")).strip(" ;")
            if not brands or len(brands) < 3:
                return None
            cn = re.sub(r"[^a-z0-9]", "", company.lower())
            bn = re.sub(r"[^a-z0-9]", "", brands.lower())
            if bn == cn or (";" not in brands and len(cn) >= 6 and bn.startswith(cn[:6])):
                return None
            # Prefer when at least one known brand appears, OR multiple semicolon brands
            if not KNOWN.search(brands) and brands.count(";") < 1:
                return None
            return brands
    except Exception:
        return None


def main() -> None:
    ckpt = {}
    if CKPT.exists():
        try:
            ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
        except Exception:
            ckpt = {}

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(headers)}
    ci, wi, bi = ic["Company"], ic["Website"], ic["Key Brands Represented"]

    # Apply extra HAND
    for row in range(2, ws.max_row + 1):
        name = str(ws.cell(row, ci).value or "").strip()
        if name in EXTRA_HAND:
            ws.cell(row, bi).value = EXTRA_HAND[name]

    rows = []
    for row in range(2, ws.max_row + 1):
        b = str(ws.cell(row, bi).value or "").strip()
        name = str(ws.cell(row, ci).value or "").strip()
        if name in EXTRA_HAND:
            continue
        if b == NPD or not b:
            rows.append(row)

    print("npd_candidates", len(rows), flush=True)
    updated = 0
    skipped = 0
    limit = 80
    for i, row in enumerate(rows[:limit]):
        name = str(ws.cell(row, ci).value or "").strip()
        web = str(ws.cell(row, wi).value or "")
        print(f"[{i+1}/{min(limit,len(rows))}] {name}", flush=True)
        ev = fetch(web)
        res = extract(name, web, ev)
        if res:
            ws.cell(row, bi).value = res
            ckpt[name] = {"brands": res}
            CKPT.write_text(json.dumps(ckpt, indent=2, ensure_ascii=False), encoding="utf-8")
            updated += 1
            print("  ->", res[:120], flush=True)
        else:
            skipped += 1
        time.sleep(0.3)

    wb.save(FP)

    # stats
    named = npd = 0
    for row in range(2, ws.max_row + 1):
        b = str(ws.cell(row, bi).value or "").strip()
        if b == NPD or not b:
            npd += 1
        else:
            named += 1
    print(f"updated={updated} skipped={skipped} named={named} npd={npd}")
    print("saved", FP)


if __name__ == "__main__":
    main()
