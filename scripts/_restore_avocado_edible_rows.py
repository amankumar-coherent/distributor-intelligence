#!/usr/bin/env python3
"""Restore edible avocado-oil-relevant rows with merged snapshot fields."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
GEO = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_geo_audit_snapshot.json"
)
CATS = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_cats_audit.json"
)
BRAND = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_brand_audit_snapshot.json"
)
OFFICE = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\_avocado_deepseek_office_cache.json"
)
NPD = "Not publicly disclosed"

RESTORE = {
    "So African Natural Oils": {
        "brands": "Seed Essential, SANO, others",
        "focus": "Cold-pressed avocado, grape seed, macadamia & botanical oils for food and personal care, others",
    },
    "Quimidroga": {},
    "Cavin Oil Products": {},
    "Mitsubishi International Food Ingredients, Inc.": {},
    "Circle Ingredients": {},
}

geo = {r["company"]: r for r in json.loads(GEO.read_text(encoding="utf-8"))}
cats = {r["company"]: r for r in json.loads(CATS.read_text(encoding="utf-8"))}
brand = {r["company"]: r for r in json.loads(BRAND.read_text(encoding="utf-8"))}
office = json.loads(OFFICE.read_text(encoding="utf-8")) if OFFICE.exists() else {}


def office_for(name: str, website: str) -> tuple[str, str]:
    key = f"ds::{name}::{website}"
    hit = office.get(key) or {}
    # only accept phone; skip emails marked not in text / invented
    phone = (hit.get("phone") or "").strip() or NPD
    email = NPD
    raw = hit.get("raw") or {}
    if raw.get("email_in_text") and (hit.get("email") or "").strip():
        email = hit["email"].strip()
    return phone, email


wb = openpyxl.load_workbook(XLSX)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}
existing = {
    str(ws.cell(r, ic["Company"]).value or "").strip() for r in range(3, ws.max_row + 1)
}

added = []
for name, ov in RESTORE.items():
    if name in existing:
        continue
    g = geo.get(name, {})
    c = cats.get(name, {})
    b = brand.get(name, {})
    website = c.get("website") or b.get("website") or NPD
    phone, email = office_for(name, website if website.startswith("http") else "")
    # try loose office key match
    if phone == NPD:
        for k, v in office.items():
            if name in k:
                phone = (v.get("phone") or "").strip() or NPD
                raw = v.get("raw") or {}
                if raw.get("email_in_text") and (v.get("email") or "").strip():
                    email = v["email"].strip()
                break
    brands = ov.get("brands") or b.get("brands") or c.get("brands") or NPD
    focus = ov.get("focus") or c.get("focus") or b.get("focus") or NPD
    core = c.get("core") or b.get("core") or NPD
    values = {
        "Company": name,
        "Website": website,
        "Founded": g.get("founded") or NPD,
        "Headquarters": g.get("hq") or NPD,
        "Continent / Geography": g.get("continent") or NPD,
        "Operational Presence": g.get("presence") or NPD,
        "Ownership": NPD,
        "Employees": NPD,
        "Core Categories": core,
        "Specialty Focus": focus,
        "Key Brands Represented": brands,
        "Office No.": phone,
        "Office Email": email,
        "Country Code": NPD,
        "Region Code": NPD,
        "Summary": NPD,
    }
    row_idx = ws.max_row + 1
    for col_name, val in values.items():
        if col_name in ic:
            ws.cell(row_idx, ic[col_name]).value = val
    added.append(name)

wb.save(XLSX)
wb.close()
print("restored", added)
