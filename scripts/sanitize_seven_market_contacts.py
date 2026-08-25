#!/usr/bin/env python3
"""Sanitize Contact Person / Role / LinkedIn on the 7 market Excels.

Keeps only sales / commercial / procurement leadership contacts.
Clears CEO/Founder/garbage names and company LinkedIn URLs.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

MARKETS = [
    Path(r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global\global_electrical_distribution_market_global_FINAL.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market.xlsx"),
]

NPD = "Not publicly disclosed"
OK_ROLE = re.compile(
    r"sales|commercial|revenue|\bcro\b|procurement|purchasing|"
    r"business development|\bbd\b|channel|partner|account (executive|director|manager)",
    re.I,
)
BAD_ROLE = re.compile(
    r"^(?!.*sales).*\b(ceo|cfo|cto|coo|ciso|founder|co-founder|chairman|president|"
    r"managing director|owner|board member)\b",
    re.I,
)


def _find_header(ws):
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            return vals, r
    raise RuntimeError("header not found")


def _valid(cp: str, role: str, li: str) -> tuple[bool, str, str, str]:
    cp = re.sub(r"\s+", " ", (cp or "").replace("\n", " ")).strip()
    role = re.sub(r"\s+", " ", (role or "").replace("\n", " ")).strip()
    li = (li or "").strip()

    if not cp or cp.lower() in {"n/a", "na", "unknown", npd_l := NPD.lower()}:
        return False, "", "", ""

    if "*" in cp:
        # keep obfuscated Apollo sales names if role OK
        if role and OK_ROLE.search(role) and not (BAD_ROLE.search(role) and "sales" not in role.lower()):
            if "linkedin.com/company/" in li.lower() or not li:
                li = NPD
            elif "linkedin.com/in/" not in li.lower():
                li = NPD
            if role.count("|") >= 2:
                role = role.split("|")[0].strip()
            return True, cp, role, li
        return False, "", "", ""

    if len(cp.split()) < 2:
        return False, "", "", ""
    if re.search(r"location|linkedin|http|www\.", cp, re.I):
        return False, "", "", ""
    if not role or not OK_ROLE.search(role):
        return False, "", "", ""
    if BAD_ROLE.search(role) and "sales" not in role.lower() and "commercial" not in role.lower():
        return False, "", "", ""
    if role.count("|") >= 2 or len(role) > 120:
        role = role.split("|")[0].strip()
        if not OK_ROLE.search(role):
            return False, "", "", ""

    if "linkedin.com/in/" in li.lower():
        pass
    else:
        li = NPD

    return True, cp, role, li


def main() -> None:
    for path in MARKETS:
        if not path.exists():
            print("missing", path)
            continue
        wb = load_workbook(path)
        ws = wb.active
        headers, hr = _find_header(ws)
        col = {h: i + 1 for i, h in enumerate(headers)}
        kept = cleared = 0
        for ridx in range(hr + 1, ws.max_row + 1):
            name = str(ws.cell(ridx, col["Company"]).value or "").strip()
            if not name:
                continue
            cp = str(ws.cell(ridx, col["Contact Person"]).value or "").strip()
            role = str(ws.cell(ridx, col["Role"]).value or "").strip()
            li = str(ws.cell(ridx, col["LinkedIn"]).value or "").strip()
            if not cp and not role and not li:
                continue
            ok, cp2, role2, li2 = _valid(cp, role, li)
            if ok:
                ws.cell(ridx, col["Contact Person"]).value = cp2
                ws.cell(ridx, col["Role"]).value = role2
                ws.cell(ridx, col["LinkedIn"]).value = li2
                kept += 1
            else:
                ws.cell(ridx, col["Contact Person"]).value = ""
                ws.cell(ridx, col["Role"]).value = ""
                ws.cell(ridx, col["LinkedIn"]).value = ""
                cleared += 1
        for cell in ws[hr]:
            cell.font = Font(bold=True)
        wb.save(path)
        print(f"{path.name}: kept={kept} cleared={cleared}")


if __name__ == "__main__":
    main()
