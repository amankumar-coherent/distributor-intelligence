"""Correct verified wrong Founded years in electrical distribution FINAL Excel."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)
NPD = "Not publicly disclosed"

# Exact-name corrections from web sources (Wikipedia / company history pages)
EXACT: dict[str, str] = {
    # Graybar US: founded 1869 (Western Electric roots); incorporated Graybar 1925
    "Graybar Electric": "1869",
    # Harris & Roome lineage 1920 is correct for Canada — leave as-is if already 1920
    # Edmundson: trading origins 1801 (Wikipedia); 1991 is UK Ltd incorporation only
    "Edmundson Electrical": "1801",
    # Blackwoods Australia: company story says 1878 (file had 1876)
    "Blackwoods": "1878",
    # Rexel India had WESCO's 1922 — Rexel group founded 1967 as CDME
    "Rexel India": "1967",
    "Rexel France": "1967",
    "Rexel Germany": "1967",
    # Al-Futtaim Group Wikipedia: 1939 (several rows wrongly used 1930)
    "Al-Futtaim Automation": "1939",
    "Al-Futtaim Import": "1939",
    "Al-Futtaim Power": "1939",
    "Al-Futtaim Trading": "1939",
    "Al-Futtaim Wholesale": "1939",
    # Acklands-Grainger roots / Acklands founded 1889
    "Acklands-Grainger": "1889",
    # Telecom entity mislabeled in electrical list — clear invented/wrong year
    "Cable & Wireless (Kenya)": NPD,
}

# Fill empty/NPD for well-known parents (same legal group founding year)
FILL_IF_EMPTY: dict[str, str] = {
    "Rexel Australia": "1967",
    "Rexel Austria": "1967",
    "Rexel Belgium": "1967",
    "Rexel Energy Solutions": "1967",
    "Sonepar Australia": "1969",
    "Sonepar Denmark": "1969",
    "WESCO UK": "1922",
    "Wesco Canada": "1922",  # already 1922 usually
    "Farnell (Asia Pacific)": "1939",
}


def main() -> None:
    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    fi = headers.index("Founded") + 1

    changes: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        cur = str(ws.cell(row_idx, fi).value or "").strip()
        new = None
        reason = ""
        if name in EXACT:
            target = EXACT[name]
            if cur != target:
                new, reason = target, "exact_verified_correction"
        elif name in FILL_IF_EMPTY:
            target = FILL_IF_EMPTY[name]
            if (not cur) or cur.lower() in {
                "not publicly disclosed",
                "n/a",
                "npd",
                "unknown",
            }:
                new, reason = target, "fill_empty_known_parent"
            elif cur != target and name.startswith("Rexel") and cur in {"1922", "1974", "1975"}:
                new, reason = target, "rexel_wrong_year_fix"

        if new is not None:
            ws.cell(row_idx, fi).value = new
            changes.append(
                {"company": name, "old": cur, "new": new, "reason": reason}
            )

    # Bold header
    for cell in ws[1]:
        cell.font = Font(bold=True)

    wb.save(XLSX)
    dump = XLSX.parent / "_founded_corrections.json"
    dump.write_text(json.dumps(changes, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"corrected {len(changes)} founded values")
    for c in changes:
        print(f"  {c['company']}: {c['old']!r} -> {c['new']!r} ({c['reason']})")
    print(f"wrote {XLSX}")
    print(f"log {dump}")


if __name__ == "__main__":
    main()
