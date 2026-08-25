"""Pass3: add remaining channel distributors to reach ~200 net new."""
from __future__ import annotations

import json
import os
import re
import shutil
import socket
import ssl
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Font
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")
FP = ROOT / (
    "output/web_expand/global_flexible_packaging_market_global"
    "/global_flexible_packaging_market_global_FINAL.xlsx"
)
OUT = FP.parent / "channel_add_200_pass3.json"
NEED = 100

OEM = re.compile(
    r"(?i)\b(amcor|mondi|huhtamaki|coveris|constantia|winpak|printpack|sonoco|proampac|"
    r"sealed\s*air|berry|suedpack|südpack|wipak|novolex|uflex|polyplex|jindal|orora|"
    r"smurfit|ahlstrom|intertape|rajoo|association|alibaba|indiamart|amazon|example\.invalid)\b"
)


def domain_of(url: str) -> str:
    u = (url or "").strip()
    if not re.match(r"^https?://", u or "", re.I):
        u = "https://" + u
    host = urlparse(u).netloc.lower().split("@")[-1]
    return host[4:] if host.startswith("www.") else host


def normalize_url(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    return u.rstrip("/")


def probe_ok(url: str) -> tuple[bool, str]:
    u = normalize_url(url)
    host = domain_of(u)
    if not host or "." not in host:
        return False, "bad"
    try:
        socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
    except Exception:
        try:
            socket.getaddrinfo(host, 80, type=socket.SOCK_STREAM)
        except Exception:
            return False, "dns"
    ctx = ssl.create_default_context()
    for attempt in (u, "https://" + host, "http://" + host):
        try:
            req = Request(attempt, headers={"User-Agent": "Mozilla/5.0"}, method="GET")
            kw = {"timeout": 10}
            if attempt.startswith("https"):
                kw["context"] = ctx
            with urlopen(req, **kw) as resp:
                code = getattr(resp, "status", 200)
                if code < 400 or code in (401, 403):
                    return True, resp.geturl()
        except Exception:
            continue
    return False, "fail"


def deepseek_batch(n: int, avoid: list[str]) -> list[dict]:
    client = OpenAI(
        api_key=os.environ["OPENAI_API_KEY"],
        base_url=os.environ.get("OPENAI_BASE_URL") or "https://api.deepseek.com/v1",
    )
    prompt = f"""JSON array only, {n} objects.
Format: [{{"company":"...","website":"https://...","country":"...","role":"distributor"}}]
ONLY real packaging DISTRIBUTORS/WHOLESALERS (bags, films, stretch, mailers, packaging supplies).
NO OEMs/manufacturers/converters/associations/marketplaces.
Prefer well-known regional packaging supply houses with real domains.
Avoid: {', '.join(avoid[:50])}
"""
    resp = client.chat.completions.create(
        model=os.getenv("OPENAI_MODEL") or "deepseek-chat",
        messages=[
            {"role": "system", "content": "JSON array only"},
            {"role": "user", "content": prompt},
        ],
        temperature=0.15,
        max_tokens=6000,
    )
    raw = (resp.choices[0].message.content or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I).strip()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        last = raw.rfind("}")
        data = json.loads(raw[raw.find("[") : last + 1] + "]")
    return data if isinstance(data, list) else []


COUNTRY_MAP = {
    "united states": ("US", "NA"),
    "usa": ("US", "NA"),
    "canada": ("CA", "NA"),
    "united kingdom": ("GB", "EU"),
    "uk": ("GB", "EU"),
    "germany": ("DE", "EU"),
    "france": ("FR", "EU"),
    "netherlands": ("NL", "EU"),
    "india": ("IN", "APAC"),
    "australia": ("AU", "APAC"),
    "new zealand": ("NZ", "APAC"),
    "uae": ("AE", "ME"),
    "mexico": ("MX", "LATAM"),
    "brazil": ("BR", "LATAM"),
    "south africa": ("ZA", "AF"),
    "singapore": ("SG", "APAC"),
    "japan": ("JP", "APAC"),
    "spain": ("ES", "EU"),
    "italy": ("IT", "EU"),
    "ireland": ("IE", "EU"),
    "belgium": ("BE", "EU"),
    "poland": ("PL", "EU"),
    "sweden": ("SE", "EU"),
}


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = FP.with_name(FP.stem + f"_backup_before_channel_pass3_{stamp}.xlsx")
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    hidx = {str(h): i for i, h in enumerate(headers)}

    existing_names = set()
    existing_domains = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        existing_names.add(str(r[0] or "").strip().lower())
        d = domain_of(str(r[1] or ""))
        if d:
            existing_domains.add(d)

    cands = []
    avoid = list(existing_names)[:60]
    for i in range(10):
        print(f"DeepSeek {i+1}/10...")
        try:
            batch = deepseek_batch(28, avoid)
            print(f"  {len(batch)}")
            for c in batch:
                name = str(c.get("company") or "").strip()
                web = normalize_url(str(c.get("website") or ""))
                if not name or not web or OEM.search(name) or OEM.search(web):
                    continue
                key = name.lower()
                dom = domain_of(web)
                if key in existing_names or key in {x["company"].lower() for x in cands} or dom in existing_domains:
                    continue
                cands.append(
                    {
                        "company": name,
                        "website": web,
                        "country": str(c.get("country") or "").strip(),
                        "role": "distributor",
                    }
                )
                avoid.append(key)
        except Exception as e:
            print(f"  err {e}")
        time.sleep(0.6)

    print(f"cands={len(cands)}")
    verified = []
    with ThreadPoolExecutor(max_workers=24) as ex:
        futs = {ex.submit(probe_ok, c["website"]): c for c in cands}
        for fut in as_completed(futs):
            c = futs[fut]
            ok, note = fut.result()
            if ok:
                c["verified_website"] = note if str(note).startswith("http") else c["website"]
                verified.append(c)
    print(f"verified={len(verified)}")
    selected = verified[:NEED]
    OUT.write_text(json.dumps(selected, indent=2, ensure_ascii=False), encoding="utf-8")

    added = 0
    for c in selected:
        row = [None] * len(headers)
        row[hidx["Company"]] = c["company"]
        row[hidx["Website"]] = c.get("verified_website") or c["website"]
        row[hidx["Distribution Type"]] = "Specialty Distributor"
        row[hidx["Retail / E-commerce / Both"]] = "No"
        row[hidx["Core Categories"]] = "Flexible packaging supplies; films; bags; stretch/shrink; packaging materials"
        row[hidx["Specialty Focus"]] = "Business distributor/wholesaler of packaging supplies (channel partner)"
        row[hidx["Summary"]] = "Channel/distributor verified via live website check for flexible packaging supplies market."
        country = c.get("country") or ""
        if country:
            row[hidx["Headquarters"]] = country
        low = country.lower()
        for k, (cc, rc) in COUNTRY_MAP.items():
            if k in low:
                row[hidx["Country Code"]] = cc
                row[hidx["Region Code"]] = rc
                row[hidx["Continent / Geography"]] = f"{rc}({country})"
                break
        ws.append(row)
        added += 1

    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(FP)
    print(f"ADDED={added} TOTAL={ws.max_row-1} backup={backup}")


if __name__ == "__main__":
    main()
