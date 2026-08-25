"""Round-2 ownership fill for remaining NPD channel rows (web + DeepSeek)."""
from __future__ import annotations

import json
import os
import re
import time
from pathlib import Path
from urllib.parse import urljoin

import httpx
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook

load_dotenv(r"D:\vendor-intel\.env")

PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
NPD = "Not publicly disclosed"

HAND = {
    "Luv2Pak": "Subsidiary of Progress Luv2Pak International Ltd.",
    "Rusco Packaging": "Private",
    "Packaging Price": "Private (Packaging Price LLC)",
    "Seal Packaging": "Private (Seal Packaging Ltd)",
    "Pack Supplies": "Private",
    "Emballage Direct": "Private (SAS)",
    "Plastic Bag Source": "Subsidiary of Magcrest Packaging Inc.",
}

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)


def emptyish(v: object) -> bool:
    return str(v or "").strip().lower() in {
        "",
        "n/a",
        "na",
        "-",
        "not publicly disclosed",
        "unknown",
        "tbd",
        "null",
    }


def strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def scrape(website: str) -> str:
    url = (website or "").strip()
    if not url:
        return ""
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    texts: list[str] = []
    try:
        with httpx.Client(
            headers={"User-Agent": UA},
            verify=False,
            timeout=httpx.Timeout(12.0, connect=4.0),
            follow_redirects=True,
        ) as client:
            for path in [
                "",
                "/about",
                "/about-us",
                "/pages/about-us",
                "/company",
                "/our-story",
            ]:
                try:
                    u = url if not path else urljoin(url.rstrip("/") + "/", path.lstrip("/"))
                    r = client.get(u)
                    if r.status_code >= 400:
                        continue
                    t = strip_html(r.text or "")
                    if len(t) > 150:
                        texts.append(t[:3000])
                    if len(texts) >= 2:
                        break
                except Exception:
                    continue
    except Exception:
        return ""
    return "\n---\n".join(texts)[:7000]


def main() -> None:
    key = os.getenv("OPENAI_API_KEY", "").strip()
    client = (
        OpenAI(
            api_key=key,
            base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/"),
            timeout=60.0,
        )
        if key
        else None
    )
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()

    wb = load_workbook(PATH)
    ws = wb.active
    h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {n: h.index(n) + 1 for n in h}

    gaps = [
        r
        for r in range(2, ws.max_row + 1)
        if emptyish(ws.cell(r, ic["Ownership"]).value)
    ]
    print("gaps", len(gaps))

    filled = 0
    left: list[str] = []
    system = (
        "Extract ownership from EVIDENCE ONLY. Never invent. "
        'JSON {"ownership":"","confidence":"high|medium|low|none","reason":""}. '
        "Formats: Private / Private (family-owned) / Subsidiary of X / Public (EX: TICKER). "
        'If insufficient: ownership="" confidence=none.'
    )

    for r in gaps:
        name = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        hq = (
            str(ws.cell(r, ic["Headquarters"]).value or "").strip()
            if "Headquarters" in ic
            else ""
        )

        if name in HAND:
            ws.cell(r, ic["Ownership"]).value = HAND[name]
            filled += 1
            print("HAND", name, "=>", HAND[name])
            continue

        evidence = scrape(web)
        if client and evidence:
            user = (
                f"Company: {name}\nWebsite: {web}\nHQ: {hq}\n\n"
                f"EVIDENCE:\n{evidence[:6500]}"
            )
            try:
                resp = client.chat.completions.create(
                    model=model,
                    temperature=0,
                    max_tokens=220,
                    response_format={"type": "json_object"},
                    messages=[
                        {"role": "system", "content": system},
                        {"role": "user", "content": user},
                    ],
                )
                raw = (resp.choices[0].message.content or "").strip()
                data = json.loads(raw)
                own = str(data.get("ownership") or "").strip()
                conf = str(data.get("confidence") or "").lower()
                if (
                    conf in {"high", "medium"}
                    and own
                    and own.lower()
                    not in {"unknown", "none", "n/a", "not publicly disclosed"}
                ):
                    ws.cell(r, ic["Ownership"]).value = own
                    filled += 1
                    print("DS", name, "=>", own, f"[{conf}]")
                else:
                    left.append(name)
                    print("LEAVE", name, conf, own or "(empty)")
            except Exception as ex:
                left.append(name)
                print("ERR", name, str(ex)[:80])
            time.sleep(0.25)
        else:
            if evidence and re.search(r"\bfamily[- ]owned\b", evidence, re.I):
                ws.cell(r, ic["Ownership"]).value = "Private (family-owned)"
                filled += 1
                print("HEUR", name, "=> Private (family-owned)")
            elif evidence and re.search(
                r"\b(privately held|private limited|ltd\.?|gmbh|sas|pty)\b",
                evidence,
                re.I,
            ):
                ws.cell(r, ic["Ownership"]).value = "Private"
                filled += 1
                print("HEUR", name, "=> Private")
            else:
                left.append(name)
                print("LEAVE", name, "(no evidence)")
                ws.cell(r, ic["Ownership"]).value = NPD

    for r in range(2, ws.max_row + 1):
        if emptyish(ws.cell(r, ic["Ownership"]).value):
            ws.cell(r, ic["Ownership"]).value = NPD

    wb.save(PATH)
    wb.close()

    wb = load_workbook(PATH, data_only=True)
    ws = wb.active
    oi = h.index("Ownership")
    blank = named = 0
    still: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[h.index("Company")] or "").strip()
        v = str(row[oi] or "").strip()
        if not v or v.lower() == NPD.lower():
            blank += 1
            still.append(name)
        else:
            named += 1
    print(
        json.dumps(
            {
                "filled_this_pass": filled,
                "still_npd": blank,
                "named": named,
                "still_names": still,
            },
            indent=2,
        )
    )
    wb.close()


if __name__ == "__main__":
    main()
