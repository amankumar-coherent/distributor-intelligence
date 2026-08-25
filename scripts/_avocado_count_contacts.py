# -*- coding: utf-8 -*-
from pathlib import Path
from collections import Counter
import openpyxl

X = Path(
    r"output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
NPD_SET = {
    "",
    "npd",
    "n/a",
    "na",
    "-",
    "—",
    "none",
    "null",
    "not publicly disclosed",
}


def is_empty(v: object) -> bool:
    s = str(v or "").strip()
    return s.lower() in NPD_SET


wb = openpyxl.load_workbook(X, data_only=True)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i for i, h in enumerate(hdr) if h}
rows = []
for row in ws.iter_rows(min_row=3, values_only=True):
    co = str(row[ic["Company"]] or "").strip()
    if not co:
        continue
    rows.append(
        {
            "co": co,
            "web": str(row[ic["Website"]] or "").strip(),
            "phone": str(row[ic["Office No."]] or "").strip(),
            "email": str(row[ic["Office Email"]] or "").strip(),
        }
    )
need_p = [r for r in rows if is_empty(r["phone"])]
need_e = [r for r in rows if is_empty(r["email"])]
need_either = [r for r in rows if is_empty(r["phone"]) or is_empty(r["email"])]
print("companies", len(rows))
print("need_phone", len(need_p))
print("need_email", len(need_e))
print("need_either", len(need_either))
print("phone_top", Counter((r["phone"] or "(empty)")[:40] for r in rows).most_common(8))
print("email_top", Counter((r["email"] or "(empty)")[:40] for r in rows).most_common(8))
for r in need_either[:20]:
    print(
        "MISS",
        r["co"][:45],
        "|",
        r["phone"][:25] or "EMPTY",
        "|",
        r["email"][:35] or "EMPTY",
        "|",
        r["web"][:55],
    )
wb.close()
