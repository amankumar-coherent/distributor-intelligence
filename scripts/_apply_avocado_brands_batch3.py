#!/usr/bin/env python3
import openpyxl
from pathlib import Path

X = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
EXTRA = {
    "Harbor Foodservice": "Bountiful Harvest, Brickfire Bakery, Katy's Kitchen, Pierport, Prairie Creek, EverFry, Villa Frizzoni, Vista Verde, others",
    "Merchants Foodservice": "Peak Fresh Produce, others",
    "Grupo Éxito Wholesale": "Exito, Carulla, Taeq, Arkitect, Bronzini, Surtimax, others",
    "Fresh Del Monte Europe": "Del Monte, others",
    "Fraser and Neave Distribution": "F&N, 100PLUS, SEASONS, NUTRISOY, ICE MOUNTAIN, others",
    "Golden State Foods": "GSF proprietary liquid products, others",
}

wb = openpyxl.load_workbook(X)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}
n = 0
for r in range(3, ws.max_row + 1):
    co = str(ws.cell(r, ic["Company"]).value or "").strip()
    if co in EXTRA:
        ws.cell(r, ic["Key Brands Represented"]).value = EXTRA[co]
        n += 1
        print("set", co)
wb.save(X)
wb.close()
print("applied", n)
