"""Fill Employees, Turnover / Scale, Contact Person for top-10 expanded Excels.

- Employees + Turnover: curated public figures + web-derived ranges
- Contact: Apollo people search (distributor titles) → people/match for full name
  → web search to resolve remaining partial/obfuscated names
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import time
from pathlib import Path
from urllib.parse import urlparse

import httpx
import openpyxl

ROOT = Path(r"D:\vendor-intel")
OUTPUT = ROOT / "output"
CACHE = ROOT / "data" / "apollo_top10_contact_cache.json"
EMP_CACHE = ROOT / "data" / "emp_turnover_top10.json"
APOLLO_BASE = "https://api.apollo.io/api/v1"

# load .env
for line in (ROOT / ".env").read_text(encoding="utf-8", errors="ignore").splitlines():
    line = line.strip()
    if not line or line.startswith("#") or "=" not in line:
        continue
    k, v = line.split("=", 1)
    os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))

API_KEY = (os.getenv("APOLLO_API_KEY") or "").strip()
DELAY = float(os.getenv("APOLLO_REQUEST_DELAY_SEC", "1.8"))

DIST_TITLES = [
    "CEO",
    "Chief Executive Officer",
    "Founder",
    "Co-Founder",
    "President",
    "Managing Director",
    "Owner",
    "VP Sales",
    "Vice President of Sales",
    "VP Business Development",
    "Director of Sales",
    "Head of Sales",
    "Chief Commercial Officer",
    "General Manager",
]
SENIORITIES = ["c_suite", "founder", "owner", "vp", "head", "director"]

TITLE_SCORE = [
    (r"\bceo\b|chief executive", 100),
    (r"\bfounder\b|co-founder|owner", 95),
    (r"\bpresident\b|managing director", 85),
    (r"chief commercial|\bcco\b", 80),
    (r"vp.*sales|vice president.*sales|director.*sales|head of sales", 75),
    (r"vp.*business|business development", 70),
    (r"general manager|\bgm\b", 60),
]

# Known public / large company employee + turnover (approximate, labeled)
KNOWN_SCALE: dict[str, tuple[str, str]] = {
    "mckesson": ("~51,000", "~$309B revenue (FY recent)"),
    "cardinal health": ("~48,000", "~$205B revenue (FY recent)"),
    "owens & minor": ("~20,000", "~$10B revenue (FY recent)"),
    "henry schein": ("~25,000", "~$12B revenue (FY recent)"),
    "medline": ("~40,000+", "Private; multi-billion USD revenue"),
    "amazon": ("~1,500,000+", "~$600B+ revenue (FY recent)"),
    "walmart": ("~2,100,000", "~$650B+ revenue (FY recent)"),
    "cvs": ("~300,000", "~$350B+ revenue (FY recent)"),
    "walgreens": ("~250,000", "~$140B+ revenue (FY recent)"),
    "costco": ("~300,000+", "~$250B+ revenue (FY recent)"),
    "target": ("~400,000", "~$100B+ revenue (FY recent)"),
    "sysco": ("~70,000+", "~$78B revenue (FY recent)"),
    "performance food": ("~35,000+", "~$58B revenue (FY recent)"),
    "unfi": ("~28,000", "~$30B revenue (FY recent)"),
    "c&s": ("~17,000", "Private; multi-billion USD wholesale"),
    "gordon food": ("~20,000+", "Private; multi-billion USD foodservice"),
    "mclane": ("~20,000+", "Berkshire Hathaway; multi-billion USD"),
    "kehe": ("~5,000+", "Employee-owned; multi-billion USD"),
    "bunzl": ("~20,000+", "~£11B+ revenue (group)"),
    "avnet": ("~15,000+", "~$25B revenue (FY recent)"),
    "digi-key": ("~4,000+", "Private; multi-billion USD distribution"),
    "mouser": ("~3,000+", "Private; multi-billion USD distribution"),
    "future electronics": ("~5,000+", "Private; multi-billion USD"),
    "rs components": ("~8,000+ (RS Group)", "~£3B+ revenue (RS Group)"),
    "rs group": ("~8,000+", "~£3B+ revenue"),
    "te connectivity": ("~85,000+", "~$16B revenue (FY recent)"),
    "sager": ("~500–1,000", "Private electronics distributor"),
    "nestlé": ("~270,000+", "~CHF 90B+ revenue (group)"),
    "diageo": ("~30,000+", "~£20B+ revenue"),
    "pernod": ("~19,000+", "~€12B+ revenue"),
    "l'oréal": ("~90,000+", "~€40B+ revenue"),
    "loreal": ("~90,000+", "~€40B+ revenue"),
    "amorepacific": ("~5,000+", "Public K-beauty group"),
    "ulta": ("~50,000+", "~$11B revenue (FY recent)"),
    "sally beauty": ("~25,000+", "~$3B+ revenue"),
    "thg": ("~5,000+", "Public e-commerce / beauty"),
    "beauty bay": ("~200–500", "Online beauty retailer"),
    "joann": ("~15,000 (pre-restructure)", "Private equity / restructured"),
    "hobby lobby": ("~40,000+", "Private retail"),
    "shimano": ("~10,000+", "Public bike components"),
    "coats": ("~15,000+", "Public industrial thread"),
    "wiggle": ("~200–500", "Online cycling (Frasers Group)"),
    "chain reaction": ("~200–500", "Online cycling (Frasers Group)"),
    "aeon": ("~150,000+", "Public retail group Japan"),
    "rakuten": ("~30,000+", "Public internet services"),
    "jd.com": ("~400,000+", "Public China e-commerce"),
    "alibaba": ("~200,000+", "Public China e-commerce"),
    "sephora": ("~40,000+ (LVMH beauty)", "LVMH group retail"),
    "boots": ("~50,000+", "Walgreens Boots Alliance"),
    "watsons": ("~100,000+ (AS Watson)", "CK Hutchison / AS Watson"),
    "douglas": ("~18,000+", "European beauty retail"),
    "grainger": ("~25,000+", "~$16B revenue (FY recent)"),
    "medtronic": ("~95,000+", "~$32B revenue (FY recent)"),
    "stryker": ("~50,000+", "~$20B revenue (FY recent)"),
    "thermo fisher": ("~120,000+", "~$40B revenue (FY recent)"),
    "avantor": ("~14,000+", "~$7B revenue (FY recent)"),
    "cencora": ("~40,000+", "~$260B+ revenue (FY recent)"),
    "amerisource": ("~40,000+", "~$260B+ revenue (FY recent)"),
    "concordance": ("~2,000+", "Private healthcare distributor"),
    "breakthru": ("~5,000+", "Private wine/spirits wholesale"),
    "rndc": ("~5,000+", "Private wine/spirits wholesale"),
    "republic national": ("~5,000+", "Private wine/spirits wholesale"),
    "winebow": ("~1,000+", "Private fine wine wholesale"),
    "essity": ("~45,000+", "Public hygiene/health"),
    "3m": ("~85,000+", "~$30B+ revenue"),
    "ottobock": ("~9,000+", "Private medical technology"),
    "össur": ("~4,000+", "Public orthopedics"),
    "ossur": ("~4,000+", "Public orthopedics"),
    "sunrise medical": ("~2,000+", "Private mobility"),
    "invacare": ("~3,000+", "Public / restructured mobility"),
    "herbalife": ("~10,000+", "Public nutrition"),
    "now foods": ("~2,000+", "Private supplements"),
    "iherb": ("~2,000+", "Private online supplements"),
    "fullscript": ("~500–1,000", "Private practitioner platform"),
    "demant": ("~20,000+", "Public hearing healthcare"),
    "interacoustics": ("~500–1,000", "Demant group audiology"),
    "prosemic": ("~200–500", "Private shunt resistor manufacturer"),
    "hethstia": ("~500–1,000", "Private konjac / noodle manufacturer"),
    "dear sundays": ("~50–200", "Private nail wellness"),
    "dearsundays": ("~50–200", "Private nail wellness"),
}


def load_dotenv_key() -> str:
    if not API_KEY:
        raise SystemExit("APOLLO_API_KEY missing")
    return API_KEY


def domain_from_website(website: str) -> str:
    s = (website or "").strip().lower()
    if not s or s.startswith("not "):
        return ""
    if "://" not in s:
        s = "https://" + s
    try:
        host = urlparse(s).netloc or ""
    except Exception:
        return ""
    host = host.removeprefix("www.").split("/")[0]
    if host.count(".") >= 1 and " " not in host:
        return host
    return ""


def score_title(title: str) -> int:
    t = (title or "").lower()
    best = 0
    for pat, pts in TITLE_SCORE:
        if re.search(pat, t, re.I):
            best = max(best, pts)
    return best


def person_name(p: dict) -> str:
    name = str(p.get("name") or "").strip()
    if name:
        return name
    first = str(p.get("first_name") or "").strip()
    last = str(p.get("last_name") or "").strip()
    if first and last and "*" not in last:
        return f"{first} {last}"
    ob = str(p.get("last_name_obfuscated") or "").strip()
    if first and ob:
        return f"{first} {ob}"
    return first


def is_partial(name: str) -> bool:
    return bool(name) and ("*" in name or re.search(r"[A-Za-z]{2}\*{2,}", name))


def lookup_scale(company: str) -> tuple[str, str] | None:
    low = company.lower()
    best = None
    best_len = 0
    for k, v in KNOWN_SCALE.items():
        if k in low and len(k) > best_len:
            best = v
            best_len = len(k)
    return best


def list_companies() -> list[dict]:
    rows = []
    for path in sorted(OUTPUT.rglob("*_distributors_expanded.xlsx")):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        data = list(ws.iter_rows(values_only=True))
        hdr_i = next(i for i, r in enumerate(data[:12]) if r and str(r[0] or "").strip() == "Company")
        header = [str(h or "").strip() for h in data[hdr_i]]
        for r in data[hdr_i + 1 :]:
            if not r or not r[0]:
                continue
            rows.append(
                {
                    "market": path.parent.name,
                    "file": str(path),
                    "company": str(r[0]).strip(),
                    "website": str(r[header.index("Website")] or "").strip()
                    if "Website" in header
                    else "",
                    "founded": str(r[header.index("Founded")] or "").strip()
                    if "Founded" in header
                    else "",
                    "hq": str(r[header.index("Headquarters")] or "").strip()
                    if "Headquarters" in header
                    else "",
                }
            )
        wb.close()
    return rows


class ApolloClient:
    def __init__(self) -> None:
        self.key = load_dotenv_key()
        self.client: httpx.AsyncClient | None = None
        self.last = 0.0
        self.credits_exhausted = False

    async def __aenter__(self) -> ApolloClient:
        self.client = httpx.AsyncClient(
            timeout=45,
            headers={
                "x-api-key": self.key,
                "Content-Type": "application/json",
                "Cache-Control": "no-cache",
            },
        )
        return self

    async def __aexit__(self, *a: object) -> None:
        if self.client:
            await self.client.aclose()

    async def _post(self, path: str, params: list | dict) -> httpx.Response:
        assert self.client
        loop = asyncio.get_running_loop()
        wait = DELAY - (loop.time() - self.last)
        if wait > 0:
            await asyncio.sleep(wait)
        resp = await self.client.post(f"{APOLLO_BASE}{path}", params=params)
        self.last = loop.time()
        return resp

    async def search(self, domain: str, company: str) -> list[dict]:
        params: list[tuple[str, str | int]] = [("page", 1), ("per_page", 10)]
        if domain:
            params.append(("q_organization_domains_list[]", domain))
        else:
            params.append(("q_organization_name", company))
        for t in DIST_TITLES:
            params.append(("person_titles[]", t))
        for s in SENIORITIES:
            params.append(("person_seniorities[]", s))
        resp = await self._post("/mixed_people/api_search", params)
        if resp.status_code != 200:
            return []
        people = list(resp.json().get("people") or [])
        if not people and domain and company:
            params2: list[tuple[str, str | int]] = [
                ("page", 1),
                ("per_page", 10),
                ("q_organization_name", company),
            ]
            for s in SENIORITIES:
                params2.append(("person_seniorities[]", s))
            resp2 = await self._post("/mixed_people/api_search", params2)
            if resp2.status_code == 200:
                people = list(resp2.json().get("people") or [])
        return [p for p in people if isinstance(p, dict)]

    async def enrich(self, person_id: str, domain: str) -> dict | None:
        if self.credits_exhausted:
            return None
        resp = await self._post("/people/match", {"id": person_id, "domain": domain})
        if resp.status_code == 422 and "insufficient credits" in resp.text.lower():
            self.credits_exhausted = True
            return None
        if resp.status_code != 200:
            return None
        person = resp.json().get("person")
        return person if isinstance(person, dict) else None

    async def find_contact(self, company: str, website: str) -> dict:
        domain = domain_from_website(website)
        people = await self.search(domain, company)
        if not people:
            return {"contact": "", "role": "", "partial": False, "source": "apollo_none"}

        ranked = sorted(
            ((score_title(str(p.get("title") or "")), p) for p in people),
            key=lambda x: x[0],
            reverse=True,
        )
        # Prefer scored; else top results
        candidates = [p for sc, p in ranked if sc >= 50] or [p for _, p in ranked[:5]]

        # Full name from search
        for p in candidates:
            name = person_name(p)
            if name and not is_partial(name):
                return {
                    "contact": name,
                    "role": str(p.get("title") or "").strip(),
                    "partial": False,
                    "source": "apollo_search",
                    "person_id": str(p.get("id") or ""),
                    "linkedin": str(p.get("linkedin_url") or ""),
                }

        # Enrich top candidates
        if domain:
            for p in candidates[:4]:
                pid = str(p.get("id") or "")
                if not pid:
                    continue
                enriched = await self.enrich(pid, domain)
                if not enriched:
                    continue
                name = person_name(enriched)
                if name and not is_partial(name):
                    return {
                        "contact": name,
                        "role": str(enriched.get("title") or p.get("title") or "").strip(),
                        "partial": False,
                        "source": "apollo_match",
                        "person_id": pid,
                        "linkedin": str(enriched.get("linkedin_url") or ""),
                    }

        # Keep partial for later web resolve
        for p in candidates:
            name = person_name(p)
            if name:
                return {
                    "contact": name,
                    "role": str(p.get("title") or "").strip(),
                    "partial": is_partial(name),
                    "source": "apollo_partial",
                    "person_id": str(p.get("id") or ""),
                    "linkedin": str(p.get("linkedin_url") or ""),
                    "first_name": str(p.get("first_name") or "").strip(),
                    "last_obf": str(p.get("last_name_obfuscated") or "").strip(),
                }
        return {"contact": "", "role": "", "partial": False, "source": "apollo_none"}


async def run_apollo(limit: int | None = None) -> dict:
    rows = list_companies()
    if limit:
        rows = rows[:limit]
    cache: dict = {}
    if CACHE.is_file():
        cache = json.loads(CACHE.read_text(encoding="utf-8"))

    async with ApolloClient() as apollo:
        for i, row in enumerate(rows, 1):
            key = f"{row['market']}::{row['company']}"
            if key in cache and cache[key].get("contact"):
                print(f"[{i}/{len(rows)}] SKIP cached {row['company']}")
                continue
            print(f"[{i}/{len(rows)}] Apollo {row['company']} ...", flush=True)
            try:
                result = await apollo.find_contact(row["company"], row["website"])
            except Exception as e:
                result = {"contact": "", "role": "", "error": str(e), "source": "error"}
            cache[key] = {**row, **result, "ts": time.time()}
            if i % 10 == 0:
                CACHE.write_text(json.dumps(cache, indent=2), encoding="utf-8")
                print(f"  cached {i}", flush=True)
            if apollo.credits_exhausted:
                print("  Apollo credits exhausted for match — continuing search-only", flush=True)

    CACHE.write_text(json.dumps(cache, indent=2), encoding="utf-8")
    return cache


def resolve_partial_via_web(cache: dict) -> dict:
    """Use DuckDuckGo to resolve obfuscated Apollo names to full names."""
    try:
        from ddgs import DDGS
    except ImportError:
        from duckduckgo_search import DDGS

    partials = [k for k, v in cache.items() if v.get("partial") and v.get("contact")]
    print(f"Web-resolving {len(partials)} partial names...")
    with DDGS() as ddgs:
        for key in partials:
            v = cache[key]
            company = v.get("company") or key.split("::")[-1]
            first = v.get("first_name") or str(v.get("contact") or "").split()[0]
            role = v.get("role") or ""
            obf = v.get("last_obf") or ""
            q = f'{first} {company} {role}'.strip()
            if obf and "*" not in obf:
                q = f'{first} {obf} {company}'
            try:
                hits = list(ddgs.text(q, max_results=6))
            except Exception as e:
                print(f"  web err {company}: {e}")
                continue
            blob = " ".join(f"{h.get('title','')} {h.get('body','')}" for h in hits)
            # Find First Last near company
            pat = re.compile(
                rf"\b({re.escape(first)}\s+[A-Z][a-z]+(?:\s+[A-Z][a-z]+)?)\b"
            )
            cands = []
            for m in pat.finditer(blob):
                full = m.group(1).strip()
                if "*" in full:
                    continue
                if full.lower().startswith(first.lower()) and len(full.split()) >= 2:
                    cands.append(full)
            if cands:
                # most common
                from collections import Counter

                best = Counter(cands).most_common(1)[0][0]
                print(f"  {v.get('contact')} -> {best} ({company})")
                v["contact"] = best
                v["partial"] = False
                v["source"] = (v.get("source") or "apollo") + "+web"
            else:
                print(f"  unresolved {v.get('contact')} ({company})")
            time.sleep(0.4)
    CACHE.write_text(json.dumps(cache, indent=2), encoding="utf-8")
    return cache


def apply_to_excels(cache: dict) -> None:
    # group by file
    by_file: dict[str, dict[str, dict]] = {}
    for key, v in cache.items():
        f = v.get("file")
        company = v.get("company") or key.split("::")[-1]
        if not f:
            continue
        by_file.setdefault(f, {})[company] = v

    for fpath, cmap in by_file.items():
        path = Path(fpath)
        if not path.is_file():
            continue
        wb = openpyxl.load_workbook(path)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=False))
        hdr_i = next(
            i
            for i, r in enumerate(rows[:12])
            if r and r[0].value and str(r[0].value).strip() == "Company"
        )
        header = [str(c.value or "").strip() for c in rows[hdr_i]]
        ci = header.index("Company")
        ei = header.index("Employees")
        ti = header.index("Turnover / Scale")
        pi = header.index("Contact Person")
        ri = header.index("Role") if "Role" in header else None
        li = header.index("LinkedIn") if "LinkedIn" in header else None

        updated = 0
        for r in rows[hdr_i + 1 :]:
            company = str(r[ci].value or "").strip()
            if not company:
                continue
            # employees / turnover
            scale = lookup_scale(company)
            if scale:
                emp, turn = scale
                cur_e = str(r[ei].value or "").strip().lower()
                cur_t = str(r[ti].value or "").strip().lower()
                if (not cur_e) or "not " in cur_e or "disclosed" in cur_e or "confirm" in cur_e:
                    r[ei].value = emp
                    updated += 1
                if (not cur_t) or "not " in cur_t or "disclosed" in cur_t or "confirm" in cur_t:
                    r[ti].value = turn
                    updated += 1
            else:
                # SME default band if still empty
                cur_e = str(r[ei].value or "").strip().lower()
                cur_t = str(r[ti].value or "").strip().lower()
                if (not cur_e) or "not " in cur_e or "disclosed" in cur_e or "confirm" in cur_e:
                    r[ei].value = "Not publicly disclosed (SME / private)"
                    updated += 1
                if (not cur_t) or "not " in cur_t or "disclosed" in cur_t or "confirm" in cur_t:
                    r[ti].value = "Not publicly disclosed (SME / private)"
                    updated += 1

            hit = cmap.get(company)
            if hit and hit.get("contact"):
                cur_p = str(r[pi].value or "").strip().lower()
                if (not cur_p) or "not " in cur_p or "confirm" in cur_p:
                    r[pi].value = hit["contact"]
                    updated += 1
                    if ri is not None and hit.get("role"):
                        r[ri].value = hit["role"]
                    if li is not None and hit.get("linkedin"):
                        cur_l = str(r[li].value or "").strip().lower()
                        if (not cur_l) or "not " in cur_l or "confirm" in cur_l:
                            r[li].value = hit["linkedin"]

        if "Run Summary" in wb.sheetnames:
            wb["Run Summary"]["A15"] = f"Employees/Turnover/Contact fill updated cells≈{updated}"
        wb.save(path)
        wb.close()
        print(f"Wrote {path.parent.name}: ~{updated} cell updates")


async def amain() -> None:
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--skip-apollo", action="store_true")
    ap.add_argument("--web-only", action="store_true")
    ap.add_argument("--apply-only", action="store_true")
    args = ap.parse_args()

    CACHE.parent.mkdir(parents=True, exist_ok=True)

    if args.apply_only:
        cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.is_file() else {}
        apply_to_excels(cache)
        return

    if args.web_only:
        cache = json.loads(CACHE.read_text(encoding="utf-8"))
        cache = resolve_partial_via_web(cache)
        apply_to_excels(cache)
        return

    if not args.skip_apollo:
        cache = await run_apollo(limit=args.limit)
    else:
        cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.is_file() else {}

    cache = resolve_partial_via_web(cache)
    apply_to_excels(cache)

    filled = sum(1 for v in cache.values() if v.get("contact"))
    partial = sum(1 for v in cache.values() if v.get("partial"))
    print(f"\nContacts with name: {filled}; still partial: {partial}; cache={CACHE}")


if __name__ == "__main__":
    asyncio.run(amain())
