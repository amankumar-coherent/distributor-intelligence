"""Add verified major flexible-packaging players (Company + Website only)."""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
OUT_TSV = FP.parent / "unique_companies_cleaned.tsv"

# Verified official sites only. Berry→Amcor and Scholle IPN→SIG redirects omitted as separate rows.
ADD = [
    ("Amcor", "https://www.amcor.com/"),
    ("Sealed Air", "https://www.sealedair.com/"),  # HTTP 403 to bots; official domain
    ("Mondi", "https://www.mondigroup.com/"),
    ("Huhtamaki", "https://www.huhtamaki.com/"),
    ("Coveris", "https://www.coveris.com/"),
    ("Constantia Flexibles", "https://www.cflex.com/"),
    ("Winpak", "https://www.winpak.com/"),
    ("Printpack", "https://www.printpack.com/"),
    ("Sonoco", "https://www.sonoco.com/"),
    ("Glenroy", "https://www.glenroy.com/"),
    ("Wipak", "https://wipak.com/"),
    ("Novolex", "https://novolex.com/"),
    ("CCL Industries", "https://www.cclind.com/"),
    ("Suedpack", "https://www.suedpack.com/"),
    ("Cosmo Films", "https://www.cosmofilms.com/"),
    ("Bischof+Klein", "https://www.bischof-klein.com/"),
    ("American Packaging Corporation", "https://americanpackaging.com/"),
    ("ePac Flexible Packaging", "https://epacflexibles.com/"),
    ("RKW Group", "https://rkw-group.com/"),
    ("Walki", "https://www.walki.com/"),
    ("Flair Flexible Packaging", "https://www.flairpackaging.com/"),
    ("Danaflex", "https://www.danaflex.com/"),
    ("Korozo", "https://www.korozogroup.com/"),
    ("SIG", "https://www.sig.biz/"),
    ("Smurfit Westrock", "https://www.smurfitwestrock.com/"),
    ("Ahlstrom", "https://www.ahlstrom.com/"),
    ("Perlen Packaging", "https://www.perlenpackaging.com/"),
    ("Gualapack", "https://gualapack.com/"),
    ("Jindal Poly Films", "https://www.jindalpoly.com/"),
]


def keyify(name: str) -> str:
    n = re.sub(r"[^\w\s+]", " ", name.lower())
    return re.sub(r"\s+", " ", n).strip()


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = FP.with_name(FP.stem + f"_backup_before_majors_{stamp}.xlsx")
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    ncols = ws.max_column
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    existing = {keyify(str(r[0] or "")) for r in rows}

    added = []
    skipped = []
    for name, web in ADD:
        k = keyify(name)
        if k in existing:
            skipped.append(name)
            continue
        blank = [None] * ncols
        blank[0] = name
        blank[1] = web
        rows.append(blank)
        existing.add(k)
        added.append(f"{name}\t{web}")

    rows.sort(key=lambda r: str(r[0] or "").lower())

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in rows:
        while len(r) < ncols:
            r.append(None)
        ws.append(r[:ncols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(FP)

    with OUT_TSV.open("w", encoding="utf-8") as f:
        f.write("company\twebsite\n")
        for r in rows:
            f.write(f"{r[0]}\t{r[1] or ''}\n")

    print(f"backup: {backup}")
    print(f"added: {len(added)}")
    print(f"skipped_existing: {skipped}")
    print(f"unique_total: {len(rows)}")
    print("--- ADDED ---")
    print("\n".join(added))


if __name__ == "__main__":
    main()
