# -*- coding: utf-8 -*-
"""Normalize LinkedIn person URLs to https://www.linkedin.com/in/{slug}."""
from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook

FILES = {
    "electrical": Path(
        r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
        r"\global_electrical_distribution_market_global_FINAL.xlsx"
    ),
    "avocado": Path(
        r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
        r"\global_avocado_oil_market.xlsx"
    ),
    "smartphone": Path(
        r"D:\vendor-intel\output\web_expand\global_smartphone_market_global"
        r"\global_smartphone_market.xlsx"
    ),
    "firewall": Path(
        r"D:\vendor-intel\output\run4\global_firewall_market_global"
        r"\global_firewall_market.xlsx"
    ),
}

IN_RE = re.compile(
    r"(?:https?://)?(?:[a-z]{2}\.)?(?:www\.)?linkedin\.com/in/([^/?#\s]+)",
    re.I,
)


def canonicalize(li: str) -> str | None:
    s = (li or "").strip()
    if not s:
        return None
    m = IN_RE.search(s)
    if not m:
        return None
    slug = unquote(m.group(1)).strip().strip("/")
    if not slug or slug.lower() in {"dir", "unavailable"}:
        return None
    return f"https://www.linkedin.com/in/{slug}"


def save_atomic(wb, path: Path):
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def main():
    n = 0
    for mkt, path in FILES.items():
        wb = load_workbook(path)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        headers = None
        hr = 1
        for r in range(1, 6):
            vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
            if "Company" in vals and "Contact Person" in vals:
                headers = {h: i + 1 for i, h in enumerate(vals) if h}
                hr = r
                break
        c_li = headers["LinkedIn"]
        for row in range(hr + 1, ws.max_row + 1):
            old = str(ws.cell(row, c_li).value or "").strip()
            new = canonicalize(old)
            if new and new != old:
                ws.cell(row, c_li).value = new
                n += 1
        save_atomic(wb, path)
        wb.close()
        print("saved", mkt)
    print("normalized", n)


if __name__ == "__main__":
    main()
