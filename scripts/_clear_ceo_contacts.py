# -*- coding: utf-8 -*-
"""Clear Contact Person / Role / LinkedIn when Role is CEO (not current sales/procurement)."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

PATHS = [
    Path(r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global\global_electrical_distribution_market_global_FINAL.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market.xlsx"),
]

CEO_RE = re.compile(r"\bceo\b|chief executive|vorstandsvorsitzender", re.I)
# Current sales/procurement title — not "former Commercial Director"
CURRENT_SALES_RE = re.compile(
    r"(head of sales|vp sales|vice president of sales|sales director|commercial director|"
    r"chief commercial|chief revenue|chief sales|procurement|purchasing|"
    r"director of sales|national sales|sales head)",
    re.I,
)
FORMER_RE = re.compile(r"\b(former|ex-|ex |retired|left)\b", re.I)


def is_ceo_not_sales(role: str) -> bool:
    if not role or not CEO_RE.search(role):
        return False
    if FORMER_RE.search(role) and CURRENT_SALES_RE.search(role):
        # e.g. "CEO (and former Commercial Director)" → still CEO now
        return True
    if CURRENT_SALES_RE.search(role) and not FORMER_RE.search(role):
        return False
    return True


def main() -> None:
    for p in PATHS:
        wb = load_workbook(p)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        headers = None
        hr = 1
        for r in range(1, 6):
            vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
            if "Company" in vals:
                headers = vals
                hr = r
                break
        co_i = headers.index("Company") + 1
        role_i = headers.index("Role") + 1
        cp_i = headers.index("Contact Person") + 1
        li_i = headers.index("LinkedIn") + 1 if "LinkedIn" in headers else None
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            co = str(ws.cell(r, co_i).value or "").strip()
            if not co:
                continue
            role = str(ws.cell(r, role_i).value or "").strip()
            if not is_ceo_not_sales(role):
                continue
            ws.cell(r, role_i).value = None
            ws.cell(r, cp_i).value = None
            if li_i:
                ws.cell(r, li_i).value = None
            n += 1
        tmp = p.with_suffix(".xlsx.tmp")
        wb.save(tmp)
        wb.close()
        tmp.replace(p)
        print(f"CLEARED {p.name} ceo_rows={n}")


if __name__ == "__main__":
    main()
