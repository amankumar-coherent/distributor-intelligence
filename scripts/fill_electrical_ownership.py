"""Fill / normalize Ownership column in electrical FINAL Excel via known_ownership + web-verified map."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from vendor_intel.enrichment.known_ownership import resolve_ownership, sanitize_ownership

XLSX = Path(
    r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
    r"\global_electrical_distribution_market_global_FINAL.xlsx"
)


def main() -> None:
    backup = XLSX.with_name(XLSX.stem + "_BEFORE_OWNERSHIP_FILL.xlsx")
    if not backup.exists():
        backup.write_bytes(XLSX.read_bytes())

    wb = load_workbook(XLSX)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = headers.index("Company") + 1
    oi = headers.index("Ownership") + 1

    changes: list[dict[str, str]] = []
    filled = 0
    still_npd = 0
    for row_idx in range(2, ws.max_row + 1):
        name = str(ws.cell(row_idx, ci).value or "").strip()
        if not name:
            continue
        old = str(ws.cell(row_idx, oi).value or "").strip()
        new = sanitize_ownership(name, old)
        if new != old:
            ws.cell(row_idx, oi).value = new
            changes.append({"company": name, "old": old or "(empty)", "new": new})
        if new.lower() in {"not publicly disclosed", "private", "public"} and not resolve_ownership(
            name, old
        ):
            still_npd += 1
        else:
            filled += 1

    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(XLSX)

    dump = XLSX.parent / "_ownership_fill_audit.json"
    dump.write_text(
        json.dumps(
            {
                "when": datetime.now().isoformat(timespec="seconds"),
                "changed": len(changes),
                "detailed_or_known": filled,
                "still_generic_or_npd": still_npd,
                "changes": changes,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(f"changed {len(changes)} ownership values")
    print(f"detailed/known~{filled} still_generic/npd~{still_npd}")
    for c in changes[:20]:
        print(f"  {c['company']}: {c['old']!r} -> {c['new']!r}")
    print(f"wrote {XLSX}")
    print(f"audit {dump}")


if __name__ == "__main__":
    main()
