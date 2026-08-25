#!/usr/bin/env python3
"""Fix 3G Mobile Operational Presence / regions (USA was incorrect)."""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

PRESENCE = "South Africa; Namibia; Botswana; Mauritius"
CITIES = "South Africa; Namibia; Botswana; Mauritius"
REGIONS_SERVED = "South Africa; Namibia; Botswana; Mauritius"
REGIONAL_EXT = "Southern Africa; Indian Ocean (Mauritius)"
CONTINENT = "AF(South Africa, Namibia, Botswana, Mauritius)"
HQ = "Johannesburg, South Africa"

FILES = [
    Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market_global_FINAL.xlsx"),
]


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for path in FILES:
        if not path.exists():
            print("MISSING", path)
            continue
        bak = path.with_name(path.stem + f"_bak_3gmobile_{stamp}" + path.suffix)
        shutil.copy2(path, bak)
        wb = openpyxl.load_workbook(path)
        fixed = False
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            hi = None
            hdr = None
            for i, r in enumerate(rows[:10]):
                vals = [str(c or "").strip() for c in r]
                if "Company" in vals:
                    hi = i
                    hdr = vals
                    break
            if hi is None:
                continue
            ic = {h: i for i, h in enumerate(hdr) if h}
            for excel_row in range(hi + 2, ws.max_row + 1):
                co = ws.cell(excel_row, ic["Company"] + 1).value
                if not co or str(co).strip().lower() != "3g mobile":
                    continue

                def setcol(name: str, val: str) -> None:
                    if name not in ic:
                        return
                    col = ic[name] + 1
                    old = ws.cell(excel_row, col).value
                    if str(old or "").strip() != val:
                        ws.cell(excel_row, col).value = val
                        print(f"  {name}: {old!r} => {val!r}")

                print(f"{path.name} / {ws.title} row {excel_row}")
                setcol("Operational Presence", PRESENCE)
                setcol("Continent / Geography", CONTINENT)
                setcol("Cities / Regions", CITIES)
                setcol("Regions Served", REGIONS_SERVED)
                setcol("Regional Extensions", REGIONAL_EXT)
                if "Headquarters" in ic:
                    hq = str(ws.cell(excel_row, ic["Headquarters"] + 1).value or "")
                    if "united states" in hq.lower() or not hq.strip():
                        setcol("Headquarters", HQ)
                fixed = True
            if fixed:
                break
        wb.save(path)
        wb.close()
        print("saved", path)
        print("backup", bak)


if __name__ == "__main__":
    main()
