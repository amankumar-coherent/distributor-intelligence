#!/usr/bin/env python3
"""Fill empty Continent / Geography in FINAL_SLIM Excels.

1) Google AI scraper (batches of 25)
2) DeepSeek (OPENAI_* / DeepSeek key) for residuals
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
sys.path.insert(0, str(ROOT / "src"))
load_dotenv(ROOT / ".env")
os.environ["GOOGLE_AI_SCRAPER_ENABLED"] = "true"
os.environ["GOOGLE_AI_SCRAPER_FORCE_OVERVIEW"] = "true"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRIES"] = "2"
os.environ["GOOGLE_AI_SCRAPER_OVERVIEW_RETRY_SLEEP_SEC"] = "1.0"

from vendor_intel.clients.google_ai_scraper import google_ai_ask, google_ai_ping  # noqa: E402

SLIM_FILES = [
    ROOT / "output/run4/global_smartwatch_market_global/global_smartwatch_market_FINAL_SLIM.xlsx",
    ROOT / "output/run4/global_green_chemical_market_global/global_green_chemical_market_FINAL_SLIM.xlsx",
    ROOT / "output/run4/global_firewall_market_global/global_firewall_market_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_global_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL_SLIM.xlsx",
    ROOT / "output/web_expand/global_antibody_drug_conjugates_market_global/global_antibody_drug_conjugates_market_FINAL_SLIM.xlsx",
]

GEO = "Continent / Geography"
BATCH = 25
PROGRESS = ROOT / "scripts/_geo_fill_empty_progress.jsonl"


def norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [norm(c.value) for c in ws[r] if c.value not in (None, "")]
        if "company" in vals:
            return r
    return 1


def collect_empty() -> list[dict]:
    rows: list[dict] = []
    for f in SLIM_FILES:
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        hr = header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        if GEO not in headers:
            wb.close()
            continue
        ci = headers.index("Company")
        gi = headers.index(GEO)
        hi = headers.index("Headquarters") if "Headquarters" in headers else None
        wi = headers.index("Website") if "Website" in headers else None
        rsi = headers.index("Regions Served") if "Regions Served" in headers else None
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if not row or not row[ci]:
                continue
            geo = str(row[gi] or "").strip()
            if geo:
                continue
            company = str(row[ci]).strip()
            rows.append(
                {
                    "file": str(f),
                    "company": company,
                    "hq": str(row[hi] or "").strip() if hi is not None else "",
                    "website": str(row[wi] or "").strip() if wi is not None else "",
                    "regions_served": str(row[rsi] or "").strip() if rsi is not None else "",
                    "key": f"{norm(f.name)}::{norm(company)}",
                }
            )
        wb.close()
    return rows


def load_done() -> dict[str, str]:
    """key -> geo string for successful fills."""
    done: dict[str, str] = {}
    if PROGRESS.exists():
        for line in PROGRESS.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
                if obj.get("key") and obj.get("geo"):
                    done[obj["key"]] = str(obj["geo"])
            except Exception:
                pass
    return done


def normalize_countries(vals: object) -> str:
    if vals is None:
        return ""
    if isinstance(vals, str):
        # already a joined string
        if ";" in vals or "," in vals:
            parts = re.split(r"[,;/|]", vals)
        else:
            parts = [vals]
        items = [p.strip() for p in parts if p.strip()]
    elif isinstance(vals, list):
        items = [str(x).strip() for x in vals if str(x).strip()]
    else:
        return ""
    seen: set[str] = set()
    out: list[str] = []
    for it in items:
        k = norm(it)
        if not k or k in {"n/a", "npd", "unknown", "none"}:
            continue
        if k in seen:
            continue
        seen.add(k)
        out.append(it)
    return "; ".join(out)


def deepseek_client() -> OpenAI:
    return OpenAI(
        api_key=os.getenv("OPENAI_API_KEY") or "",
        base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/"),
    )


def parse_geo_batch(companies: list[dict], markdown: str, source: str) -> dict[str, str]:
    """Parse Google/DeepSeek text into company -> countries string."""
    if not markdown.strip():
        return {}
    names = [c["company"] for c in companies]
    client = deepseek_client()
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    prompt = f"""Extract countries served for each company from the text.

Companies:
{json.dumps(names, ensure_ascii=False)}

Text ({source}):
```
{markdown[:14000]}
```

Return JSON array only:
[{{"company":"exact name","countries":["Country1","Country2",...]}}]
Rules:
- List EVERY country mentioned for that company (even 50+).
- If text says Global/Worldwide with no country list, use ["Global"].
- If HQ country only is known and no regions, use that country alone.
- Do not invent. Skip company if nothing useful.
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    text = (resp.choices[0].message.content or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I | re.M).strip()
    try:
        data = json.loads(text)
    except Exception:
        return {}
    if isinstance(data, dict):
        data = data.get("companies") or data.get("results") or []
    out: dict[str, str] = {}
    if not isinstance(data, list):
        return out
    for row in data:
        if not isinstance(row, dict):
            continue
        name = row.get("company")
        geo = normalize_countries(row.get("countries") or row.get("countries_served"))
        if name and geo:
            out[norm(name)] = geo
    return out


def build_google_query(batch: list[dict]) -> str:
    lines = [
        "For EACH company, list ALL countries / regions served (every country name, even if more than 50).",
        "Also note headquarters country if known.",
        'Query style: "<company> founded year, headquarter and all region served mention all country"',
        "",
        "Companies:",
    ]
    for i, c in enumerate(batch, 1):
        bits = [c["company"]]
        if c.get("hq"):
            bits.append(f"HQ={c['hq']}")
        if c.get("website"):
            bits.append(f"site={c['website']}")
        if c.get("regions_served"):
            bits.append(f"regions_hint={c['regions_served'][:120]}")
        lines.append(f"{i}. {' | '.join(bits)}")
    lines.append("Do not invent countries.")
    return "\n".join(lines)


async def google_fill_batch(batch: list[dict]) -> dict[str, str]:
    q = build_google_query(batch)
    data = await google_ai_ask(q, llm_clean=True)
    md = str(data.get("markdown") or data.get("markdown_raw") or "").strip()
    found = parse_geo_batch(batch, md, "google_ai") if md else {}

    missing = [c for c in batch if norm(c["company"]) not in found]
    if missing:
        # miss re-batch
        q2 = build_google_query(missing)
        data2 = await google_ai_ask(q2, llm_clean=True)
        md2 = str(data2.get("markdown") or data2.get("markdown_raw") or "").strip()
        if md2:
            found.update(parse_geo_batch(missing, md2, "google_ai_retry"))

    # single asks for still-missing (cap 8)
    still = [c for c in batch if norm(c["company"]) not in found]
    for c in still[:8]:
        q1 = f"{c['company']} founded year, headquarter and all region served mention all country"
        d1 = await google_ai_ask(q1, llm_clean=True)
        md1 = str(d1.get("markdown") or d1.get("markdown_raw") or "").strip()
        if md1:
            found.update(parse_geo_batch([c], md1, "google_ai_single"))
    return found


def ddg_snippets(company: str, hq: str = "", website: str = "") -> str:
    try:
        from ddgs import DDGS
    except ImportError:
        try:
            from duckduckgo_search import DDGS  # type: ignore
        except ImportError:
            return ""
    q = f"{company} countries served OR markets OR regions OR headquarters distribution"
    if hq:
        q += f" {hq}"
    if website:
        q += f" {website}"
    bits: list[str] = []
    try:
        with DDGS() as ddgs:
            for r in ddgs.text(q, max_results=5):
                bits.append(f"- {r.get('title','')}: {r.get('body') or r.get('description') or ''}")
    except Exception as e:
        bits.append(f"(search error: {e})")
    return "\n".join(bits)


def deepseek_fill_one(c: dict) -> str:
    """Residual fill with DeepSeek using DDGS + existing hints."""
    snippets = ddg_snippets(c["company"], c.get("hq", ""), c.get("website", ""))
    client = deepseek_client()
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    prompt = f"""List ALL countries this company serves / operates in / distributes to.

Company: {c['company']}
Headquarters: {c.get('hq') or 'unknown'}
Website: {c.get('website') or 'unknown'}
Existing Regions Served hint: {c.get('regions_served') or 'none'}

Web snippets:
{snippets or '(none)'}

Return JSON only:
{{"countries":["Country1","Country2",...]}}
Rules:
- Include every country you can support from snippets/hints.
- If only HQ country is known, return that one country.
- If clearly global with no list, return ["Global"].
- Do not invent. Empty list if unknown.
"""
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
    )
    text = (resp.choices[0].message.content or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I | re.M).strip()
    try:
        data = json.loads(text)
    except Exception:
        return ""
    return normalize_countries(data.get("countries") if isinstance(data, dict) else data)


def apply_geo(fills: dict[str, str], rows: list[dict]) -> None:
    by_file: dict[str, dict[str, str]] = {}
    key_to_company = {r["key"]: r for r in rows}
    for key, geo in fills.items():
        row = key_to_company.get(key)
        if not row or not geo:
            continue
        by_file.setdefault(row["file"], {})[norm(row["company"])] = geo

    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for fpath, mapping in by_file.items():
        wb = load_workbook(fpath)
        ws = wb.active
        hr = header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        gi = headers.index(GEO)
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            company = ws.cell(r, ci + 1).value
            if not company:
                continue
            key = norm(str(company))
            if key not in mapping:
                continue
            cur = str(ws.cell(r, gi + 1).value or "").strip()
            if cur:
                continue
            cell = ws.cell(r, gi + 1, mapping[key])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin
            n += 1
        wb.save(fpath)
        print(f"Filled {n} geo rows -> {Path(fpath).name}")


async def main() -> None:
    empty = collect_empty()
    already = load_done()
    # skip already filled in progress
    todo = [r for r in empty if r["key"] not in already]
    print(f"Empty Continent/Geography: {len(empty)} | progress-done: {len(already)} | todo: {len(todo)}")

    if not await google_ai_ping():
        print("WARN: Google AI scraper not up — DeepSeek-only for residuals")
        google_ok = False
    else:
        google_ok = True
        print("Google AI scraper: connected")

    fills: dict[str, str] = dict(already)
    PROGRESS.parent.mkdir(parents=True, exist_ok=True)

    # Phase 1: Google AI batches
    if google_ok and todo:
        with PROGRESS.open("a", encoding="utf-8") as prog:
            for i in range(0, len(todo), BATCH):
                batch = todo[i : i + BATCH]
                bnum = i // BATCH + 1
                btot = (len(todo) + BATCH - 1) // BATCH
                print(f"\n[Google AI] Batch {bnum}/{btot} ({len(batch)})", flush=True)
                try:
                    found = await google_fill_batch(batch)
                except Exception as e:
                    print(f"  batch error: {e}", flush=True)
                    found = {}
                for c in batch:
                    geo = found.get(norm(c["company"]))
                    if geo:
                        fills[c["key"]] = geo
                        prog.write(
                            json.dumps(
                                {
                                    "key": c["key"],
                                    "company": c["company"],
                                    "file": c["file"],
                                    "geo": geo,
                                    "source": "google_ai",
                                },
                                ensure_ascii=False,
                            )
                            + "\n"
                        )
                        prog.flush()
                        print(f"  OK {c['company']}: {geo[:90]}", flush=True)
                    else:
                        print(f"  MISS {c['company']}", flush=True)
                await asyncio.sleep(0.4)

    # Phase 2: DeepSeek residual
    residual = [r for r in empty if r["key"] not in fills or not fills.get(r["key"])]
    # also include ones still empty after google
    residual = [r for r in empty if not fills.get(r["key"])]
    print(f"\n[DeepSeek residual] {len(residual)} companies", flush=True)
    with PROGRESS.open("a", encoding="utf-8") as prog:
        for i, c in enumerate(residual, 1):
            print(f"[{i}/{len(residual)}] DeepSeek {c['company']}", flush=True)
            try:
                geo = deepseek_fill_one(c)
            except Exception as e:
                print(f"  error: {e}", flush=True)
                geo = ""
            if geo:
                fills[c["key"]] = geo
                prog.write(
                    json.dumps(
                        {
                            "key": c["key"],
                            "company": c["company"],
                            "file": c["file"],
                            "geo": geo,
                            "source": "deepseek",
                        },
                        ensure_ascii=False,
                    )
                    + "\n"
                )
                prog.flush()
                print(f"  OK: {geo[:90]}", flush=True)
            else:
                print("  still empty", flush=True)

    print("\nApplying to SLIM Excels...")
    apply_geo(fills, empty)

    # recount
    left = collect_empty()
    print(f"Remaining empty Continent/Geography: {len(left)}")


if __name__ == "__main__":
    asyncio.run(main())
