"""Fix Summary column: non-English, blank, wrong-identity, unreadable run-ons."""
from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from vendor_intel.pipeline.csv_fields import company_summary_cell, summarize_for_csv  # noqa: E402

ROOT4 = Path(r"C:\Users\aman.kumar\Desktop\coherent\output\run4")
MIRROR = Path(r"C:\Users\aman.kumar\Desktop\Cmi\agents\vendor-intel\output\run4")

# Curated English summaries for known bad / wrong / blank / non-English rows
SUMMARY_FIXES: dict[str, str] = {
    # smartwatch
    "galaxtecc": (
        "Galaxtecc is a China-based smartwatch OEM/ODM manufacturer (Galax Technology / Shenzhen Xingliu Technology). "
        "It provides custom design, production, and branding for global distributors and private-label brands, "
        "including kids GPS and fitness wearables."
    ),
    "cover company": (
        "Cover Company is a Uruguay-based electronics retailer and distributor of technology "
        "products including smartphones, laptops, consoles, smartwatches, and audio devices, "
        "with a focus on Xiaomi and Amazfit wearable brands."
    ),
    "dealer – ielov": (
        "Dealer – ielov supplies parental-control and GPS smartwatches and related wearable "
        "tracking devices for distributors and retail partners."
    ),
    "aaradhya enterprises": (
        "Aaradhya Enterprises is an India-based trader/distributor of consumer electronics "
        "and wearable devices including smartwatches and accessories."
    ),
    "cowaudio offical store": (
        "Cowaudio Official Store is a consumer electronics seller offering audio products "
        "and wearable accessories through online retail channels."
    ),
    "rmtibd": (
        "Rmtibd is a regional distributor/seller of mobile devices and smartwatches "
        "serving the Bangladesh and nearby South Asian markets."
    ),
    "shoppingvelleys": (
        "Shoppingvelleys is an online retail channel offering consumer electronics "
        "and wearable products including smartwatches."
    ),
    "smartwatchsolutionsltd": (
        "Smartwatchsolutionsltd is a specialist seller/distributor focused on smartwatches "
        "and wearable technology solutions for retail and wholesale buyers."
    ),
    # green chemical
    "ghcl trading division": (
        "GHCL Trading Division markets and trades chemical commodities such as soda ash "
        "and related industrial chemicals for domestic and export customers."
    ),
    "k. k. chempro india": (
        "K. K. Chempro India is an Indian specialty chemical distributor supplying "
        "industrial and performance chemicals to manufacturers and formulators."
    ),
    "solventis": (
        "Solventis is a European chemical solvents supplier and producer of automotive "
        "fluids, glycols, and related specialty solvents, distributing bulk chemicals globally."
    ),
    "velox gmbh": (
        "Velox GmbH is a specialty chemicals distributor (now part of IMCD) serving "
        "plastics, coatings, and related industrial markets across Europe with technical "
        "support and multi-principal product portfolios."
    ),
    # firewall
    "giacom": (
        "Giacom is a UK technology distributor providing cloud, connectivity, and IT "
        "solutions to channel partners, including security and networking offerings."
    ),
    "konekti systems": (
        "Konekti Systems is a cybersecurity and network security solutions partner that "
        "designs and implements next-generation firewalls and hybrid-cloud security "
        "architectures for enterprise clients."
    ),
    "indiannetworkco": (
        "Indiannetworkco is an India-based networking and wireless equipment distributor "
        "supplying switches, access points, and related network infrastructure to channel partners."
    ),
    "bluestar": (
        "BlueStar is a value-added technology distributor serving channel partners with IT, "
        "networking, and cybersecurity solutions across regional markets."
    ),
    "eplus technology": (
        "ePlus Technology is an IT solutions and services provider helping enterprises source, "
        "deploy, and manage networking, security, and cloud technologies."
    ),
    "newegg commerce": (
        "Newegg Commerce is a global e-commerce retailer specializing in PC hardware, consumer "
        "electronics, gaming peripherals, and related tech products for consumers and businesses."
    ),
    "no thread": (
        "No Thread provides cybersecurity and network security solutions for enterprise "
        "and MSP partners in the firewall and threat-protection channel."
    ),
    "bahtera": (
        "Bahtera is a technology distribution partner supplying IT and cybersecurity "
        "solutions, including network security products, to regional channel partners."
    ),
    "keenetic rmm": (
        "Keenetic RMM supports remote monitoring and management of Keenetic networking "
        "equipment for MSPs and IT service partners."
    ),
    "nayeem hussain visitors graph": (
        "Listing appears to be a web analytics/visitor-graph page rather than a verified "
        "firewall channel partner; retained for review."
    ),
}


def key(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "").strip().lower())


def is_non_english(text: str) -> bool:
    s = (text or "").strip()
    if not s:
        return False
    if re.search(r"(mais de|soluções|indústria|wir sind|gmbh ist|notre société|компания|для)", s, re.I):
        return True
    # mostly non-ascii letters
    letters = [c for c in s if c.isalpha()]
    if not letters:
        return False
    non_latin = sum(1 for c in letters if ord(c) > 127 and not (0x00C0 <= ord(c) <= 0x024F and c.lower() in "àáâãäåæçèéêëìíîïñòóôõöùúûüýÿ"))
    # Portuguese/French/German with Latin accents still OK if English words present
    if re.search(r"\b(the|and|company|distributor|retail|security|chemical|market)\b", s, re.I):
        return False
    # if almost no English stopwords and has accented Portuguese-like phrase
    if re.search(r"[ãõçáéíóúâêô]", s, re.I) and not re.search(r"\b(the|and|is|of|for|with)\b", s, re.I):
        return True
    return False


def is_wrong_or_junk(text: str, company: str) -> bool:
    low = (text or "").lower()
    if not low:
        return True
    if text.count("•") >= 3:
        return True
    if any(
        x in low
        for x in (
            "digital marketing agency",
            "financial consulting",
            "wealth management",
            "not-for-profit society",
            "hardware and lumber",
            "casa giacomo",
            "indian society for quality",
        )
    ):
        return True
    if "xiaomistores" in low and "cover" in company.lower():
        return True
    return False


def template_summary(company: str, role: str, market: str, core: str, specialty: str) -> str:
    role_l = (role or "channel partner").strip() or "channel partner"
    core_l = (core or "").strip()
    spec_l = (specialty or "").strip()
    market_label = {
        "smartwatch": "smartwatch and wearable technology",
        "green_chemical": "green and specialty chemicals",
        "firewall": "firewall and cybersecurity",
    }[market]
    bits = [f"{company} is a {role_l.lower()} in the {market_label} market."]
    if core_l:
        bits.append(f"Core focus areas include {core_l.rstrip('.')}.")
    elif spec_l and spec_l.lower() not in ("not applicable", "service provider", "industry manufacturer"):
        bits.append(f"Specialty focus: {spec_l.rstrip('.')}.")
    return " ".join(bits)


def polish_existing(text: str, company: str, role: str, core: str) -> str:
    cleaned = company_summary_cell(
        text,
        company=company,
        role=role,
        key_products=core,
        max_sentences=3,
        max_len=320,
    )
    if cleaned and len(cleaned) >= 40 and not is_non_english(cleaned):
        return cleaned
    cleaned2 = summarize_for_csv(text, max_sentences=3, max_len=320)
    if cleaned2 and len(cleaned2) >= 40 and not is_non_english(cleaned2):
        return cleaned2
    return ""


def resolve_summary(row: dict, market: str) -> tuple[str, str]:
    company = str(row.get("company") or row.get("brand") or "").strip()
    k = key(company)
    if k in SUMMARY_FIXES:
        return SUMMARY_FIXES[k], "curated"
    crawl = row.get("_crawl") if isinstance(row.get("_crawl"), dict) else {}
    current = str(
        row.get("polished_summary")
        or row.get("company_summary")
        or row.get("summary")
        or crawl.get("linkedin_about")
        or ""
    ).strip()
    role = str(row.get("role") or crawl.get("company_type") or "")
    core = str(crawl.get("core_categories") or "")
    specialty = str(crawl.get("specialty_focus") or "")

    if is_non_english(current) or is_wrong_or_junk(current, company):
        polished = polish_existing(current, company, role, core)
        if polished and not is_non_english(polished) and not is_wrong_or_junk(polished, company):
            return polished, "polished"
        return template_summary(company, role, market, core, specialty), "template"

    if not current:
        return template_summary(company, role, market, core, specialty), "template"

    # Too long / run-on → polish to 2-3 sentences
    if len(current) > 360 or (current.count(".") < 2 and len(current) > 220):
        polished = polish_existing(current, company, role, core)
        if polished:
            return polished, "trimmed"
    return current, "keep"


def apply_json(mid: str, market: str) -> dict:
    path = ROOT4 / mid / f"{mid}.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    stats = {"curated": 0, "polished": 0, "template": 0, "trimmed": 0, "keep": 0}
    for r in data.get("relevant_companies") or []:
        new_s, how = resolve_summary(r, market)
        stats[how] = stats.get(how, 0) + 1
        r["polished_summary"] = new_s
        r["company_summary"] = new_s
        r["summary"] = new_s
        crawl = dict(r.get("_crawl") or {})
        crawl["summary_fixed"] = how
        r["_crawl"] = crawl
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return stats


def find_header(ws):
    for r in range(1, 8):
        vals = [c.value for c in ws[r]]
        if any(v and str(v).strip().lower() == "company" for v in vals):
            return r, vals
    return 1, [c.value for c in ws[1]]


def apply_excel(path: Path, by_company: dict[str, str]) -> dict:
    if not path.exists():
        return {"missing": True}
    wb = load_workbook(path)
    ws = wb.active
    hr, headers = find_header(ws)
    hmap = {str(h).strip().lower(): i for i, h in enumerate(headers) if h}
    ci = hmap.get("company")
    si = hmap.get("summary")
    if ci is None or si is None:
        return {"error": "no summary col"}
    n = 0
    for ridx in range(hr + 1, ws.max_row + 1):
        company = ws.cell(ridx, ci + 1).value
        if not company:
            continue
        sm = by_company.get(key(str(company)))
        if sm:
            ws.cell(ridx, si + 1).value = sm
            n += 1
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    try:
        tmp.replace(path)
        out = str(path)
    except PermissionError:
        alt = path.with_name(path.stem + "_SUMMARY_FIXED.xlsx")
        tmp.replace(alt)
        out = str(alt)
    return {"file": out, "updated": n}


def main() -> None:
    markets = {
        "global_smartwatch_market_global": "smartwatch",
        "global_green_chemical_market_global": "green_chemical",
        "global_firewall_market_global": "firewall",
    }
    for mid, market in markets.items():
        print(f"\n=== {mid} ===")
        print(" JSON", apply_json(mid, market))
        data = json.loads((ROOT4 / mid / f"{mid}.json").read_text(encoding="utf-8"))
        by_co = {
            key(str(r.get("company") or "")): str(r.get("polished_summary") or "")
            for r in data.get("relevant_companies") or []
        }
        short = re.sub(r"_global$", "", mid)
        for suffix in ["_FINAL.xlsx", "_founded_and_hq_v2.xlsx", "_founded_and_hq_v2_cleaned.xlsx"]:
            print(" XLSX", apply_excel(ROOT4 / mid / f"{short}{suffix}", by_co))
        # show galaxtecc / last rows check for smartwatch
        if market == "smartwatch":
            for name in ("Galaxtecc", "Cover Company", "Wildberries", "Arrow Electronics"):
                print(f"  {name}: {by_co.get(key(name), '')[:160]}")
        dst = MIRROR / mid
        dst.mkdir(parents=True, exist_ok=True)
        for pat in [f"{mid}.json", "*_FINAL.xlsx", "*_founded_and_hq_v2.xlsx", "*_founded_and_hq_v2_cleaned.xlsx"]:
            for p in (ROOT4 / mid).glob(pat):
                try:
                    shutil.copy2(p, dst / p.name)
                except PermissionError:
                    print(f"  skip locked {p.name}")


if __name__ == "__main__":
    main()
