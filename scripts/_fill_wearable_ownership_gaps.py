# -*- coding: utf-8 -*-
"""Fill empty Ownership cells on wearable FINAL — hand-verified + DeepSeek (no invent)."""
from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import httpx
import openpyxl
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

OUT = ROOT / "output/web_expand/global_wearable_medical_devices_market_global"
XLSX = OUT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
CACHE = OUT / "wearable_ownership_gap_cache.json"
AUDIT = OUT / "wearable_ownership_gap_audit.json"
SUMMARY = OUT / "wearable_ownership_gap_summary.json"

# Web-verified only (SEC/IR/Companies House/press/About). Applied only when Ownership blank.
HAND: dict[str, str] = {
    "Allegro Medical": "Subsidiary of Scrip Companies (Centre Lane Partners)",
    "Carex Health Brands": "Subsidiary of Compass Health Brands",
    "SpinLife": "Subsidiary of Numotion (AEA Investors / Leavitt Equity Partners)",
    "Durbin PLC": "Subsidiary of Uniphar plc",
    "Dental Cremer": "Subsidiary of Henry Schein",
    "Medshop Australia": "Subsidiary of Bunzl plc",
    "Medhold": "Private (Sanlam Private Equity majority)",
    "Medisave UK": "Private",
    "Health Products For You": "Private",
    "Darby Group Companies": "Private (family-owned)",
    "Mediplus UK": "Private (Mediplus Holdings)",
    "Sanitätshaus Aktuell AG": "Member-owned Verbundgruppe (member Sanitätshäuser)",
    "Burrows Company": "Subsidiary of Owens & Minor",
    "Shenzhen Jumper Medical Equipment": "Private",
    "Mediquip": "Private (PE-backed)",
    "Dispensing Solutions": "Subsidiary of McKesson Medical-Surgical (via PSS World Medical)",
}

EMPTYISH = {
    "",
    "n/a",
    "na",
    "none",
    "-",
    "—",
    "not publicly disclosed",
    "unknown",
    "tbd",
    "null",
}

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)


def emptyish(v: object) -> bool:
    return str(v or "").strip().lower() in EMPTYISH


def strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    t = re.sub(r"&\w+;", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def scrape_evidence(website: str) -> str:
    url = (website or "").strip()
    if not url:
        return ""
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    paths = ["", "/about", "/about-us", "/company", "/investors", "/en/about", "/who-we-are"]
    texts: list[str] = []
    try:
        with httpx.Client(
            headers={"User-Agent": UA},
            verify=False,
            timeout=httpx.Timeout(12.0, connect=4.0),
            follow_redirects=True,
        ) as client:
            for path in paths:
                try:
                    u = url if not path else urljoin(url.rstrip("/") + "/", path.lstrip("/"))
                    r = client.get(u)
                    if r.status_code >= 400:
                        continue
                    t = strip_html(r.text or "")
                    if len(t) > 180:
                        texts.append(f"URL:{u}\n{t[:2800]}")
                    if len(texts) >= 2:
                        break
                except Exception:
                    continue
    except Exception:
        return ""
    return "\n\n---\n\n".join(texts)[:7500]


def ds_client() -> tuple[OpenAI | None, str]:
    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    base = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    if not key:
        return None, model
    return OpenAI(api_key=key, base_url=base, timeout=60.0), model


def parse_json(text: str) -> dict:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S).strip()
    try:
        return json.loads(raw)
    except Exception:
        m = re.search(r"\{.*\}", raw, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except Exception:
                return {}
        return {}


def deepseek_own(
    client: OpenAI, model: str, company: str, website: str, hq: str, evidence: str
) -> dict:
    system = (
        "Extract company ownership from EVIDENCE ONLY. Never invent parents, PE firms, or tickers.\n"
        "Return JSON: {ownership:string, confidence:high|medium|low|none, reason:string}.\n"
        "Preferred formats:\n"
        "- Public (EXCHANGE: TICKER)\n"
        "- Private\n"
        "- Private (PE / family detail) when evidenced\n"
        "- Subsidiary of PARENT\n"
        "- Cooperative / Member-owned\n"
        "- Government / state-owned\n"
        "If evidence is insufficient, ownership='' and confidence=none.\n"
        "Do not guess."
    )
    user = (
        f"Company: {company}\nWebsite: {website}\nHQ: {hq}\n\n"
        f"EVIDENCE:\n{(evidence or '(none)')[:7000]}"
    )
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=250,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        data = parse_json(resp.choices[0].message.content or "{}")
    except Exception as ex:  # noqa: BLE001
        return {"ownership": "", "confidence": "none", "reason": str(ex)[:160]}

    own = re.sub(r"\s+", " ", str(data.get("ownership") or "").strip())
    conf = str(data.get("confidence") or "").lower()
    if own.lower() in EMPTYISH | {"unclear", "not found", "unknown"}:
        own = ""
    if len(own) > 140:
        own = own[:137] + "..."
    return {
        "ownership": own,
        "confidence": conf,
        "reason": str(data.get("reason") or "")[:200],
    }


def main() -> None:
    client, model = ds_client()
    print(f"DeepSeek={'yes' if client else 'no'} model={model}", flush=True)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = XLSX.with_name(XLSX.stem + f"_BEFORE_OWNERSHIP_GAPS_{stamp}.xlsx")
    wb0 = openpyxl.load_workbook(XLSX)
    wb0.save(backup)
    wb0.close()
    print(f"backup {backup.name}", flush=True)

    cache: dict = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {str(h): i + 1 for i, h in enumerate(headers) if h}

    gaps = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if not co:
            continue
        cur = str(ws.cell(r, ic["Ownership"]).value or "").strip()
        if emptyish(cur):
            gaps.append(r)
    print(f"empty ownership rows={len(gaps)}", flush=True)

    hand_n = 0
    ds_n = 0
    left = 0
    audit = []

    for r in gaps:
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        hq = str(ws.cell(r, ic.get("Headquarters", 0)).value or "").strip() if "Headquarters" in ic else ""

        if co in HAND:
            own = HAND[co]
            ws.cell(r, ic["Ownership"]).value = own
            hand_n += 1
            audit.append({"company": co, "source": "hand", "ownership": own})
            print(f"HAND {co} => {own}", flush=True)
            continue

        # DeepSeek from site evidence
        if not client:
            left += 1
            audit.append({"company": co, "source": "skip", "ownership": ""})
            continue

        evidence = ""
        if co in cache and cache[co].get("evidence") is not None:
            evidence = cache[co].get("evidence") or ""
        else:
            evidence = scrape_evidence(web)
            cache[co] = {"evidence": evidence[:4000]}

        ds = deepseek_own(client, model, co, web, hq, evidence)
        conf = ds.get("confidence") or ""
        own = ds.get("ownership") or ""
        if conf in {"high", "medium"} and own:
            ws.cell(r, ic["Ownership"]).value = own
            ds_n += 1
            audit.append({"company": co, "source": "deepseek", "ownership": own, "conf": conf})
            print(f"DS {co} => {own} [{conf}]", flush=True)
        else:
            left += 1
            audit.append(
                {
                    "company": co,
                    "source": "unverified",
                    "ownership": "",
                    "conf": conf,
                    "reason": ds.get("reason"),
                }
            )
            print(f"LEAVE {co} conf={conf} own={own or '(empty)'}", flush=True)

        cache[co]["ds"] = ds
        time.sleep(0.25)
        if (hand_n + ds_n + left) % 10 == 0:
            CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
            wb.save(XLSX)

    ws.column_dimensions[get_column_letter(ic["Ownership"])].width = 42
    for r in range(2, ws.max_row + 1):
        ws.cell(r, ic["Ownership"]).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(XLSX)
    wb.close()
    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")

    # recount
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    oi = headers.index("Ownership")
    filled = blank = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = str(row[oi] or "").strip()
        if emptyish(v):
            blank += 1
        else:
            filled += 1
    wb.close()

    stats = {
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "file": str(XLSX),
        "backup": str(backup),
        "gaps_at_start": len(gaps),
        "hand_filled": hand_n,
        "deepseek_filled": ds_n,
        "left_blank_unverified": left,
        "ownership_filled_now": filled,
        "ownership_blank_now": blank,
        "total_rows": filled + blank,
    }
    SUMMARY.write_text(json.dumps(stats, indent=2), encoding="utf-8")
    print(json.dumps(stats, indent=2), flush=True)


if __name__ == "__main__":
    try:
        import urllib3

        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    except Exception:
        pass
    main()
