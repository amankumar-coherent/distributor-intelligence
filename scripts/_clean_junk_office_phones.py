# -*- coding: utf-8 -*-
"""Remove obviously fake/placeholder Office No. values harvested from websites."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from _harvest_office_phones import PATHS

JUNK_EXACT = {
    "1234567",
    "12345678",
    "123456789",
    "1234567890",
    "0123456789",
    "1111111111",
    "0000000000",
    "821234567",
    "27821234567",
}
JUNK_RE = re.compile(
    r"(123[-.\s]*456[-.\s]*7890|555[-.\s]*0?1[-.\s]*0100|555[-.\s]*1234|82\s*123\s*4567)",
    re.I,
)


def digits(s: str) -> str:
    return re.sub(r"\D", "", s or "")


def is_junk(s: str) -> bool:
    d = digits(s)
    if not d:
        return True
    if d in JUNK_EXACT or d.lstrip("0") in JUNK_EXACT:
        return True
    if JUNK_RE.search(s or ""):
        return True
    if len(set(d)) <= 2:
        return True
    if d.startswith("12345") or d.endswith("1234567"):
        return True
    return False


def main() -> None:
    for p in PATHS:
        wb = load_workbook(p)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        headers = None
        hr = 1
        for r in range(1, 6):
            vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
            if "Company" in vals:
                headers = vals
                hr = r
                break
        off_i = headers.index("Office No.") + 1
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            cur = str(ws.cell(r, off_i).value or "").strip()
            if cur and is_junk(cur):
                print(f"  DROP {p.name} r{r} {cur}")
                ws.cell(r, off_i).value = None
                n += 1
        tmp = p.with_suffix(".xlsx.tmp")
        wb.save(tmp)
        wb.close()
        tmp.replace(p)
        print(f"CLEAN {p.name} dropped={n}")


if __name__ == "__main__":
    main()
