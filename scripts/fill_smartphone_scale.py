#!/usr/bin/env python3
"""Fill Employees + Turnover/Scale as ranges for smartphone FINAL (web + DeepSeek)."""
from __future__ import annotations

import asyncio
import json
import os
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from dotenv import dotenv_values, load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

load_dotenv(ROOT / ".env")
_cmi = Path(r"C:\Users\aman.kumar\Desktop\Cmi\agents\vendor-intel\.env")
if _cmi.exists():
    for k, v in (dotenv_values(_cmi) or {}).items():
        if v and k.startswith(("OPENAI_", "LLM_", "COMPILER_", "CLASSIFIER_", "MARKET_MAP_")):
            os.environ[k] = v

FINAL = (
    ROOT
    / "output/web_expand/global_smartphone_market_global/global_smartphone_market_global_FINAL.xlsx"
)
SHEET = "Companies"
MODEL = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip() or "deepseek-chat"
BASE = (os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/")
CONCURRENT = int(os.getenv("SCALE_CONCURRENT") or "4")
LIMIT = int(os.getenv("SCALE_LIMIT") or "0")
NPD = "Not publicly disclosed"

EMP_RANGES = (
    "1–10",
    "11–50",
    "51–200",
    "201–500",
    "501–1,000",
    "1,001–5,000",
    "5,001–10,000",
    "10,001–50,000",
    "50,001+",
)
TO_RANGES = (
    "<$1M",
    "$1–10M",
    "$10–50M",
    "$50–100M",
    "$100–500M",
    "$500M–1B",
    "$1–5B",
    "$5–10B",
    "$10–50B",
    "$50B+",
)

# Curated ranges for well-known firms
KNOWN: dict[str, dict[str, str]] = {
    "amazon": {"employees": "50,001+", "turnover": "$50B+"},
    "best buy": {"employees": "50,001+", "turnover": "$10–50B"},
    "currys": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "cdw": {"employees": "10,001–50,000", "turnover": "$10–50B"},
    "also holding": {"employees": "1,001–5,000", "turnover": "$10–50B"},
    "ab sa": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "ingram micro": {"employees": "10,001–50,000", "turnover": "$50B+"},
    "td synnex": {"employees": "10,001–50,000", "turnover": "$50B+"},
    "d&h distributing": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "esprinet": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "computacenter": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "bechtle": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "dicker data": {"employees": "501–1,000", "turnover": "$1–5B"},
    "elko group": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "brightstar": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "exertis": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "amazon": {"employees": "50,001+", "turnover": "$50B+"},
    "lazada": {"employees": "5,001–10,000", "turnover": "$1–5B"},
    "erajaya": {"employees": "5,001–10,000", "turnover": "$1–5B"},
    "insight enterprises": {"employees": "10,001–50,000", "turnover": "$5–10B"},
    "softcat": {"employees": "1,001–5,000", "turnover": "$1–5B"},
    "shi international": {"employees": "5,001–10,000", "turnover": "$10–50B"},
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def _emptyish(v: str) -> bool:
    return _norm(v) in {
        "",
        "n/a",
        "na",
        "none",
        "-",
        "—",
        "not publicly disclosed",
        "unknown",
        "tbd",
        "todo",
        "seed_file",
    }


def _emp_to_range(raw: str) -> str:
    s = (raw or "").strip()
    if _emptyish(s):
        return ""
    # already a valid range label
    for r in EMP_RANGES:
        if _norm(s).replace(",", "") == _norm(r).replace(",", ""):
            return r
        if _norm(s).replace("-", "–") == _norm(r):
            return r
    # extract number
    s2 = s.replace(",", "").replace(" ", "")
    m = re.search(r"(\d+(?:\.\d+)?)\s*([kKmM])?\+?", s2)
    if not m:
        # range like 100-500
        m2 = re.search(r"(\d+)\s*[-–to]+\s*(\d+)", s.replace(",", ""), re.I)
        if m2:
            mid = (int(m2.group(1)) + int(m2.group(2))) // 2
            return _emp_num_to_range(mid)
        return ""
    n = float(m.group(1))
    suf = (m.group(2) or "").lower()
    if suf == "k":
        n *= 1_000
    elif suf == "m":
        n *= 1_000_000
    return _emp_num_to_range(int(n))


def _emp_num_to_range(n: int) -> str:
    if n <= 0:
        return ""
    if n <= 10:
        return "1–10"
    if n <= 50:
        return "11–50"
    if n <= 200:
        return "51–200"
    if n <= 500:
        return "201–500"
    if n <= 1000:
        return "501–1,000"
    if n <= 5000:
        return "1,001–5,000"
    if n <= 10000:
        return "5,001–10,000"
    if n <= 50000:
        return "10,001–50,000"
    return "50,001+"


def _money_to_usd(raw: str) -> float | None:
    s = (raw or "").strip().lower().replace(",", "").replace(" ", "")
    if not s or _emptyish(raw):
        return None
    # pure absurd integers as revenue without unit → ignore if huge year-like
    if re.fullmatch(r"\d{4}", s):  # year mistaken
        return None
    mult = 1.0
    if "billion" in s or s.endswith("b") or "bn" in s:
        mult = 1_000_000_000
    elif "million" in s or s.endswith("m") or "mn" in s:
        mult = 1_000_000
    elif "thousand" in s or s.endswith("k"):
        mult = 1_000
    m = re.search(r"(\d+(?:\.\d+)?)", s)
    if not m:
        return None
    n = float(m.group(1))
    if mult == 1.0:
        # bare number: if small like 6000 likely employees leaked; if huge maybe currency units
        if n < 100_000:
            return None
        return n  # assume USD
    return n * mult


def _to_to_range(raw: str) -> str:
    s = (raw or "").strip()
    if _emptyish(s):
        return ""
    for r in TO_RANGES:
        if _norm(s).replace(" ", "") == _norm(r).replace(" ", ""):
            return r
    # existing range mid
    if re.search(r"[-–]", s) and re.search(r"[\d$€£]", s):
        # take first number as approx
        parts = re.split(r"[-–]", s)
        usd = _money_to_usd(parts[0])
        if usd is None and len(parts) > 1:
            usd = _money_to_usd(parts[-1])
        if usd is not None:
            return _usd_to_range(usd)
    usd = _money_to_usd(s)
    if usd is None:
        return ""
    return _usd_to_range(usd)


def _usd_to_range(usd: float) -> str:
    if usd < 1_000_000:
        return "<$1M"
    if usd < 10_000_000:
        return "$1–10M"
    if usd < 50_000_000:
        return "$10–50M"
    if usd < 100_000_000:
        return "$50–100M"
    if usd < 500_000_000:
        return "$100–500M"
    if usd < 1_000_000_000:
        return "$500M–1B"
    if usd < 5_000_000_000:
        return "$1–5B"
    if usd < 10_000_000_000:
        return "$5–10B"
    if usd < 50_000_000_000:
        return "$10–50B"
    return "$50B+"


def _needs_fix(row: dict[str, str]) -> bool:
    e = row.get("Employees") or ""
    t = row.get("Turnover / Scale") or ""
    er = _emp_to_range(e)
    tr = _to_to_range(t)
    # need if empty/NPD or not already a standard range
    e_std = any(_norm(e).replace(",", "").replace("-", "–") == _norm(r).replace(",", "") for r in EMP_RANGES)
    t_std = any(_norm(t).replace(" ", "") == _norm(r).replace(" ", "") for r in TO_RANGES)
    if _emptyish(e) or not e_std:
        return True
    if _emptyish(t) or not t_std:
        return True
    # junk years as employees
    if re.fullmatch(r"20\d{2}", e.strip()):
        return True
    return False


def _load_rows(path: Path) -> tuple[list[str], list[dict[str, str]], int]:
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers: list[str] = []
    for i, r in enumerate(rows[:20]):
        vals = [str(c or "").strip() for c in r]
        if "Company" in vals:
            header_idx = i
            headers = vals
            break
    data: list[dict[str, str]] = []
    for r in rows[header_idx + 1 :]:
        d = {
            headers[j]: ("" if j >= len(r) or r[j] is None else str(r[j]).strip())
            for j in range(len(headers))
            if headers[j]
        }
        if d.get("Company"):
            data.append(d)
    return [h for h in headers if h], data, header_idx


async def _web_evidence(name: str, website: str) -> str:
    host = ""
    if website:
        host = re.sub(r"^https?://", "", website.strip(), flags=re.I).split("/")[0].replace("www.", "")
    queries = [
        f'"{name}" employees OR "number of employees" OR headcount OR staff',
        f'"{name}" revenue OR turnover OR "annual sales" OR "net sales" OR billion OR million',
    ]
    if host:
        queries.insert(0, f'"{name}" {host} employees OR revenue OR about')

    parts: list[str] = []
    try:
        from ddgs import DDGS

        def _one(q: str) -> list[str]:
            lines: list[str] = []
            try:
                hits = list(DDGS().text(q, max_results=5)) or []
            except Exception:
                hits = []
            for h in hits:
                lines.append(
                    f"- {h.get('title', '')}: {str(h.get('body', ''))[:230]} ({h.get('href', '')})"
                )
            return lines

        for q in queries[:3]:
            try:
                lines = await asyncio.wait_for(asyncio.to_thread(_one, q), timeout=22)
            except Exception:
                lines = []
            if lines:
                parts.append(f"Query: {q}\n" + "\n".join(lines))
    except Exception as err:
        parts.append(f"(ddgs: {err})")

    if not any(p.startswith("Query:") for p in parts):
        try:
            from vendor_intel.clients.search_router import FreeSearchRouter
            from vendor_intel.config import Settings
            from vendor_intel.enrichment.gap_fill.sources import _search_snippets

            router = FreeSearchRouter(Settings.load())
            for q in queries[:2]:
                try:
                    snip = await asyncio.wait_for(
                        _search_snippets(router, q, max_hits=5), timeout=25
                    )
                except Exception:
                    snip = ""
                if snip:
                    parts.append(f"Query: {q}\n{snip}")
        except Exception as err:
            parts.append(f"(router: {err})")
    return "\n\n".join(parts)[:6500] or "(no web hits)"


def _client():
    from openai import OpenAI

    key = (os.getenv("OPENAI_API_KEY") or "").strip()
    if not key:
        raise RuntimeError("OPENAI_API_KEY missing")
    return OpenAI(api_key=key, base_url=BASE)


def _llm(client: Any, row: dict[str, str], evidence: str) -> dict[str, Any]:
    name = row.get("Company") or ""
    system = (
        "Extract company size for a smartphone channel landscape. Return JSON only. "
        "Never invent. If unknown leave empty (caller maps to Not publicly disclosed).\n"
        f"employees MUST be one of: {', '.join(EMP_RANGES)} or empty.\n"
        f"turnover MUST be one of: {', '.join(TO_RANGES)} or empty.\n"
        "Map any exact headcount/revenue into the closest range. "
        "Turnover = annual revenue / sales scale in USD-equivalent ranges."
    )
    user = f"""Company: {name}
Website: {row.get('Website') or ''}
HQ: {row.get('Headquarters') or ''}
Current Employees: {row.get('Employees') or ''}
Current Turnover / Scale: {row.get('Turnover / Scale') or ''}
Business Type: {row.get('Business Type') or ''}

WEB EVIDENCE:
{evidence}

Return JSON:
{{
  "employees": "one allowed range or empty",
  "turnover": "one allowed range or empty",
  "employees_raw": "number or note from evidence or empty",
  "turnover_raw": "figure from evidence or empty",
  "confidence": "high|medium|low",
  "notes": "short"
}}
"""
    resp = client.chat.completions.create(
        model=MODEL,
        temperature=0.1,
        max_tokens=400,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
    )
    raw = resp.choices[0].message.content or "{}"
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{[\s\S]*\}", raw)
        data = json.loads(m.group(0)) if m else {}
    return data if isinstance(data, dict) else {}


def _apply(row: dict[str, str], data: dict[str, Any]) -> dict[str, str]:
    out = dict(row)
    known = KNOWN.get(_norm(row.get("Company") or ""), {})
    conf = str(data.get("confidence") or "").lower()

    emp = _emp_to_range(str(data.get("employees") or "")) or _emp_to_range(
        str(data.get("employees_raw") or "")
    )
    turn = _to_to_range(str(data.get("turnover") or "")) or _to_to_range(
        str(data.get("turnover_raw") or "")
    )

    # also try converting existing cell if LLM empty
    if not emp:
        emp = _emp_to_range(out.get("Employees") or "")
    if not turn:
        turn = _to_to_range(out.get("Turnover / Scale") or "")

    if known.get("employees") and (conf != "high" or not emp):
        emp = known["employees"]
    elif known.get("employees"):
        emp = known["employees"]

    if known.get("turnover") and (conf != "high" or not turn):
        turn = known["turnover"]
    elif known.get("turnover"):
        turn = known["turnover"]

    out["Employees"] = emp or NPD
    out["Turnover / Scale"] = turn or NPD

    src = out.get("Data Sources") or ""
    if "deepseek_scale_fill" not in src:
        out["Data Sources"] = (src + "; deepseek_scale_fill").strip("; ")
    return out


def _write(path: Path, headers: list[str], rows: list[dict[str, str]], title: str, meta: dict) -> None:
    wb = load_workbook(path)
    ws = wb[SHEET]
    header_row = 2
    for r in range(1, 6):
        if ws.cell(r, 1).value == "Company":
            header_row = r
            break
    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)
    for j, h in enumerate(headers, 1):
        ws.cell(header_row, j).value = h
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    for i, row in enumerate(rows):
        rr = header_row + 1 + i
        for j, h in enumerate(headers, 1):
            cell = ws.cell(rr, j, row.get(h, ""))
            cell.alignment = wrap
            cell.border = thin
    ws.cell(1, 1).value = title
    name = "Scale Fill Audit"
    if name in wb.sheetnames:
        del wb[name]
    aus = wb.create_sheet(name)
    aus.append(["field", "value"])
    for k, v in meta.items():
        aus.append([k, str(v)])
    wb.save(path)


async def _one(sem, client, row, idx, total):
    name = row.get("Company") or ""
    async with sem:
        print(f"  [{idx}/{total}] {name}", flush=True)
        try:
            evidence = await _web_evidence(name, row.get("Website") or "")
            data = await asyncio.to_thread(_llm, client, row, evidence)
            fixed = _apply(row, data)
            print(
                f"      → Emp={fixed.get('Employees')} | TO={fixed.get('Turnover / Scale')} "
                f"(conf={data.get('confidence')})",
                flush=True,
            )
            return fixed, {
                "company": name,
                "employees": fixed.get("Employees"),
                "turnover": fixed.get("Turnover / Scale"),
                "confidence": data.get("confidence"),
                "notes": data.get("notes"),
            }
        except Exception as err:  # noqa: BLE001
            print(f"      ERROR {type(err).__name__}: {err}", flush=True)
            fixed = _apply(row, {})
            return fixed, {"company": name, "error": str(err), "employees": fixed.get("Employees"), "turnover": fixed.get("Turnover / Scale")}


async def main() -> int:
    if not FINAL.exists():
        print(f"ERROR missing {FINAL}", file=sys.stderr)
        return 1
    headers, rows, _ = _load_rows(FINAL)
    todo = [i for i, r in enumerate(rows) if _needs_fix(r)]
    # always process all if almost all need fix
    if len(todo) > len(rows) * 0.7:
        todo = list(range(len(rows)))
    if LIMIT > 0:
        todo = todo[:LIMIT]

    print("=== Smartphone Employees + Turnover ranges ===", flush=True)
    print(f"  file : {FINAL}", flush=True)
    print(f"  queue: {len(todo)} / {len(rows)}", flush=True)
    print(f"  model: {MODEL} @ {BASE}", flush=True)
    print(f"  emp ranges: {EMP_RANGES}", flush=True)
    print(f"  to ranges : {TO_RANGES}", flush=True)

    client = _client()
    ping = client.chat.completions.create(
        model=MODEL,
        max_tokens=20,
        messages=[{"role": "user", "content": 'Reply JSON {"ok":true}'}],
        response_format={"type": "json_object"},
    )
    print(f"  ping: {(ping.choices[0].message.content or '')[:40]}", flush=True)

    sem = asyncio.Semaphore(CONCURRENT)
    out_rows = list(rows)
    ckpt_path = FINAL.parent / "smartphone_scale_checkpoint.json"
    done_map: dict[str, dict[str, Any]] = {}
    if ckpt_path.exists() and not (os.getenv("SCALE_FRESH") or "").strip():
        try:
            prior = json.loads(ckpt_path.read_text(encoding="utf-8"))
            for v in prior.get("audits") or []:
                if v.get("company") and v.get("employees") and v.get("turnover"):
                    done_map[_norm(v["company"])] = v
            print(f"  resume: {len(done_map)}", flush=True)
        except Exception:
            pass

    for i, row in enumerate(rows):
        a = done_map.get(_norm(row.get("Company") or ""))
        if a:
            fixed = dict(row)
            fixed["Employees"] = a["employees"]
            fixed["Turnover / Scale"] = a["turnover"]
            out_rows[i] = fixed

    pending = [i for i in todo if _norm(rows[i].get("Company") or "") not in done_map]
    chunk = 12
    for start in range(0, len(pending), chunk):
        batch = pending[start : start + chunk]
        print(
            f"--- batch {start // chunk + 1}: {len(batch)} (left ~{len(pending) - start}) ---",
            flush=True,
        )
        tasks = [
            _one(sem, client, rows[i], len(done_map) + start + j + 1, len(todo))
            for j, i in enumerate(batch)
        ]
        results = await asyncio.gather(*tasks)
        for i, (fixed, audit) in zip(batch, results):
            out_rows[i] = fixed
            done_map[_norm(fixed.get("Company") or "")] = audit
        ckpt_path.write_text(
            json.dumps(
                {"audits": list(done_map.values()), "done": len(done_map)},
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

    audits = [
        done_map.get(
            _norm(r.get("Company") or ""),
            {
                "company": r.get("Company"),
                "employees": out_rows[i].get("Employees"),
                "turnover": out_rows[i].get("Turnover / Scale"),
            },
        )
        for i, r in enumerate(out_rows)
    ]

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    from collections import Counter

    if LIMIT > 0:
        p = FINAL.parent / f"smartphone_scale_DRY_{ts}.json"
        p.write_text(json.dumps(audits[:LIMIT], indent=2, ensure_ascii=False), encoding="utf-8")
        print("DRY", p, flush=True)
        print("Emp", Counter(a.get("employees") for a in audits[:LIMIT]), flush=True)
        print("TO", Counter(a.get("turnover") for a in audits[:LIMIT]), flush=True)
        return 0

    bak = FINAL.with_suffix(FINAL.suffix + f".bak_scale_{ts}")
    shutil.copy2(FINAL, bak)
    emp_c = Counter(r.get("Employees") for r in out_rows)
    to_c = Counter(r.get("Turnover / Scale") for r in out_rows)
    audit_path = FINAL.parent / f"smartphone_scale_{ts}.json"
    audit_path.write_text(
        json.dumps(
            {
                "ts": ts,
                "model": MODEL,
                "emp_ranges": EMP_RANGES,
                "to_ranges": TO_RANGES,
                "employees_counts": dict(emp_c),
                "turnover_counts": dict(to_c),
                "audits": audits,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    title = f"Global Smartphone Market — channel partners (scale ranges {ts})"
    _write(
        FINAL,
        headers,
        out_rows,
        title,
        {
            "ts": ts,
            "model": MODEL,
            "rows": len(out_rows),
            "employees_counts": dict(emp_c),
            "turnover_counts": dict(to_c),
            "audit_json": str(audit_path),
            "backup": str(bak),
        },
    )
    print("=== DONE ===", flush=True)
    print(f"  backup: {bak.name}", flush=True)
    print(f"  audit : {audit_path}", flush=True)
    print("  Employees:", dict(emp_c.most_common(12)), flush=True)
    print("  Turnover:", dict(to_c.most_common(12)), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
