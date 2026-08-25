from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
wb = load_workbook(FP, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(min_row=2, values_only=True))
names = []
domains = []
for r in rows:
    name = (r[0] or "").strip()
    web = (r[1] or "").strip()
    names.append(name.lower())
    if web:
        u = web if "://" in web else "https://" + web
        d = urlparse(u).netloc.lower().replace("www.", "")
        domains.append(d)
(FP.parent / "existing_names.txt").write_text("\n".join(sorted(set(names))), encoding="utf-8")
(FP.parent / "existing_domains.txt").write_text("\n".join(sorted(set(domains))), encoding="utf-8")
print("count", len(rows), "unique_names", len(set(names)), "unique_domains", len(set(domains)))
wb.close()
