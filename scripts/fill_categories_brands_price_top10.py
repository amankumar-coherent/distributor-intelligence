"""Fill/verify Core Categories, Specialty Focus, Key Brands, Price Segment
for the top-10 *_distributors_expanded.xlsx markets.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

ROOT = Path(r"D:\vendor-intel\output")

# Market-level defaults: (core_categories, specialty_focus, price_segment, key_brands_default)
MARKET_DEFAULTS: dict[str, tuple[str, str, str, str]] = {
    "axillary_crutches_market_global": (
        "Axillary crutches; durable medical equipment (DME); mobility aids; rehabilitation supplies",
        "Hospital, clinic, and home-care mobility distribution",
        "Mid to Premium (clinical / institutional)",
        "Multi-brand DME assortment (crutches, walkers, related mobility)",
    ),
    "baijiu_market_global": (
        "Baijiu; Chinese spirits; sauce-aroma and strong-aroma liquor; duty-free / on-trade spirits",
        "Import, wholesale, and retail distribution of Chinese Baijiu",
        "Premium to Luxury (with mass-market SKUs)",
        "Moutai; Wuliangye; Luzhou Laojiao; other regional Baijiu labels",
    ),
    "benign_positional_vertigo_market_global": (
        "Vestibular diagnostic devices; BPPV treatment / repositioning aids; audiology & ENT clinic supplies",
        "Clinical audiology / vestibular channel distribution",
        "Premium (clinical / professional)",
        "Interacoustics; Natus/Otometrics; Amplivox; related vestibular brands",
    ),
    "elderberry_supplements_market_global": (
        "Elderberry supplements; immune-support syrups, gummies, capsules; natural CPG wholesale",
        "Natural / specialty supplement wholesale and retail channel",
        "Mass to Premium natural",
        "Garden of Life; Nature's Way; NOW Foods; Sambucol; private-label elderberry",
    ),
    "instant_noodles_market_global": (
        "Instant noodles; ramen / cup noodles; Asian grocery & foodservice noodles",
        "Grocery wholesale, foodservice, and e-grocery noodle distribution",
        "Value to Mid (mass grocery); some Premium Asian specialty",
        "Nissin; Nestlé Maggi; Indomie; Samyang; regional instant-noodle brands",
    ),
    "mountain_bike_suspension_market_global": (
        "Mountain bike suspension forks; rear shocks; dropper posts; suspension service parts",
        "Specialty bicycle component retail and wholesale distribution",
        "Mid to Premium / performance",
        "Fox; RockShox; Öhlins; DVO; Marzocchi; related MTB suspension brands",
    ),
    "muslin_products_market_global": (
        "Cotton muslin fabric; apparel & craft muslin; theatrical / scenic muslin; baby muslin textiles",
        "Fabric wholesale, retail, and industrial textile distribution",
        "Value to Mid (craft/apparel); Mid-Premium theatrical",
        "Mill / house brands; ROC-LON; Liberty; Mood; regional muslin mills",
    ),
    "nail_polish_market_global": (
        "Nail polish; gel polish; professional nail systems; salon nail care",
        "Professional salon wholesale and beauty retail nail channel",
        "Mass to Professional / Premium",
        "CND; Gelish; OPI; Essie; Kiara Sky; Young Nails; Sally Beauty private brands",
    ),
    "sheet_face_masks_market_global": (
        "Sheet face masks; hydrogel / biocellulose masks; K-beauty skincare masks",
        "K-beauty and drugstore sheet-mask retail / wholesale distribution",
        "Mass to Premium K-beauty",
        "Innisfree; Laneige; Dr. Jart+; Mediheal; Dr. Althea; private-label sheet masks",
    ),
    "shunt_resistor_market_global": (
        "Current-sense / shunt resistors; precision resistors; power electronics passives",
        "Electronic component distribution for industrial, auto, and power designs",
        "Mid to Premium industrial / OEM",
        "Vishay; Bourns; Isabellenhütte; PROSEMI; Yageo; related CSR brands",
    ),
}

# Company-specific overrides: core, specialty, price, brands
COMPANY_OVERRIDES: dict[str, tuple[str, str, str, str]] = {
    # --- Crutches / DME ---
    "McKesson": (
        "Pharmaceutical & medical-surgical wholesale; DME; OTC mobility",
        "National pharmaceutical and medical supply wholesale",
        "Mid to Premium institutional",
        "Multi-manufacturer medical / DME portfolio",
    ),
    "McKesson Medical-Surgical": (
        "Medical-surgical supplies; DME; clinic consumables including crutches",
        "Medical-surgical wholesale to providers",
        "Mid to Premium institutional",
        "Multi-manufacturer med-surg / mobility portfolio",
    ),
    "Cardinal Health": (
        "Medical products distribution; DME; pharmacy solutions",
        "Integrated healthcare distribution",
        "Mid to Premium institutional",
        "Multi-manufacturer medical portfolio",
    ),
    "Henry Schein": (
        "Medical & dental distribution; practice supplies; DME",
        "Physician / clinic supply distribution",
        "Mid to Premium professional",
        "Multi-brand medical and dental supplies",
    ),
    "Owens & Minor": (
        "Medical distribution & logistics; procedural packs; DME logistics",
        "Provider logistics and medical distribution",
        "Mid to Premium institutional",
        "Multi-manufacturer medical products",
    ),
    "Medline Industries,": (
        "Manufacturer/distributor of medical supplies; mobility; patient care",
        "Vertically integrated medical supply",
        "Value to Mid institutional",
        "Medline house brands; multi-source DME",
    ),
    "Medline Industries": (
        "Manufacturer/distributor of medical supplies; mobility; patient care",
        "Vertically integrated medical supply",
        "Value to Mid institutional",
        "Medline house brands; multi-source DME",
    ),
    "AliMed": (
        "Rehabilitation; orthopedics; DME; clinic equipment",
        "Rehab and clinical specialty distribution",
        "Mid to Premium clinical",
        "AliMed; multi-brand rehab/mobility",
    ),
    "Amazon Medical Supplies": (
        "Online retail of DME and home medical supplies",
        "Marketplace / e-commerce medical retail",
        "Value to Mid consumer",
        "Multi-brand third-party and Amazon Basics medical",
    ),
    "Walmart Mobility / OTC": (
        "Mass retail OTC and mobility aids",
        "Mass-market pharmacy / OTC retail",
        "Value to Mid mass",
        "Private label + national OTC/mobility brands",
    ),
    "CVS Health DME/OTC": (
        "Pharmacy retail; OTC; limited DME",
        "Drugstore / pharmacy retail channel",
        "Value to Mid mass",
        "CVS Health brands + national OTC",
    ),
    "Walgreens DME/OTC": (
        "Pharmacy retail; OTC; limited DME",
        "Drugstore / pharmacy retail channel",
        "Value to Mid mass",
        "Walgreens brands + national OTC",
    ),
    # --- Baijiu ---
    "JD.com": (
        "E-commerce supermarket including Baijiu and spirits",
        "China online retail / marketplace spirits",
        "Mass to Luxury",
        "Moutai; Wuliangye; multi-brand Baijiu on JD",
    ),
    "Breakthru Beverage Group": (
        "Wine, spirits, and beer wholesale distribution",
        "Premium beverage wholesale (USA/Canada)",
        "Premium to Luxury",
        "Multi-supplier wine & spirits portfolio (incl. Asian spirits where carried)",
    ),
    "Republic National Distributing Company (RNDC)": (
        "Wine and spirits wholesale distribution",
        "Multi-state USA spirits wholesale",
        "Premium to Luxury",
        "Multi-supplier spirits portfolio",
    ),
    "Young's Market Company": (
        "Wine and spirits wholesale (Western USA)",
        "Regional fine wine & spirits wholesale",
        "Premium to Luxury",
        "Multi-supplier wine & spirits",
    ),
    "Winebow Group": (
        "Fine wine and specialty spirits wholesale",
        "Specialty import wine & spirits",
        "Premium to Luxury",
        "Specialty import portfolio",
    ),
    "CNS Imports": (
        "Chinese spirits / Baijiu import and wholesale",
        "Asian spirits import distribution (USA)",
        "Premium to Luxury",
        "Selected Baijiu and Chinese spirit brands",
    ),
    "Cnsimports": (
        "Chinese spirits / Baijiu import and wholesale",
        "Asian spirits import distribution (USA)",
        "Premium to Luxury",
        "Selected Baijiu and Chinese spirit brands",
    ),
    "Shiwandistillery": (
        "Chinese rice wine, cooking wine, Baijiu manufacturing & export",
        "Guangdong traditional liquor producer/exporter",
        "Mass to Mid heritage",
        "Shi Wan Pai; Yubingshao; related Shiwan labels",
    ),
    "DrinkUpNY": (
        "Online retail of wine and specialty spirits",
        "Specialty e-commerce spirits retail",
        "Mid to Premium",
        "Multi-brand specialty spirits (incl. Asian spirits)",
    ),
    "DFS Group (travel retail)": (
        "Travel retail luxury goods including spirits",
        "Airport / travel retail spirits",
        "Premium to Luxury",
        "Global luxury spirits including Baijiu where listed",
    ),
    "Pernod Ricard China / Asia distribution": (
        "International spirits including China portfolio",
        "Global spirits brand owner distribution",
        "Premium to Luxury",
        "Pernod Ricard brands; China on-trade network",
    ),
    "Diageo Greater China distribution": (
        "International spirits; China market distribution",
        "Global spirits brand owner distribution",
        "Premium to Luxury",
        "Diageo brands; China spirits channel",
    ),
    # --- BPV ---
    "Interacoustics distributor network": (
        "Audiology & vestibular diagnostic equipment",
        "Clinical audiology / vestibular instrumentation",
        "Premium clinical",
        "Interacoustics",
    ),
    "Interacoustics USA / audiology dealers": (
        "Audiology & vestibular diagnostic equipment",
        "USA clinical dealer channel",
        "Premium clinical",
        "Interacoustics",
    ),
    "Natus Medical / Otometrics distribution": (
        "Neuro / audiology / Otometrics vestibular systems",
        "Clinical diagnostic distribution",
        "Premium clinical",
        "Natus; Otometrics",
    ),
    "Amplivox distributors": (
        "Hearing and audiology screening equipment",
        "Occupational / clinical audiology",
        "Mid to Premium clinical",
        "Amplivox",
    ),
    "e3 Diagnostics": (
        "Audiology equipment sales and service",
        "USA audiology dealer / distributor",
        "Premium clinical",
        "Multi-brand audiology (Interacoustics and peers)",
    ),
    "Medtronic ENT distribution": (
        "ENT surgical and related clinical products",
        "Hospital ENT channel",
        "Premium clinical / surgical",
        "Medtronic ENT",
    ),
    "Karl Storz distribution partners": (
        "Endoscopy / ENT visualization systems",
        "Hospital surgical channel",
        "Premium clinical / surgical",
        "KARL STORZ",
    ),
    "Olympus Medical distribution": (
        "Medical endoscopy and ENT imaging",
        "Hospital imaging / endoscopy channel",
        "Premium clinical / surgical",
        "Olympus Medical",
    ),
    # --- Elderberry ---
    "KeHE Distributors": (
        "Natural & specialty grocery wholesale including supplements",
        "Natural CPG wholesale to retailers",
        "Mid to Premium natural",
        "Multi-brand natural supplements incl. elderberry SKUs",
    ),
    "UNFI (United Natural Foods)": (
        "Natural/organic grocery & supplement wholesale",
        "National natural foods wholesale",
        "Mid to Premium natural",
        "Multi-brand natural / organic supplements",
    ),
    "Garden of Life": (
        "Organic vitamins and supplements including elderberry",
        "Brand manufacturer (Nestlé Health Science)",
        "Premium natural / organic",
        "Garden of Life",
    ),
    "NOW Foods": (
        "Vitamins, minerals, herbal supplements including elderberry",
        "Brand manufacturer / wholesale",
        "Value to Mid natural",
        "NOW Foods",
    ),
    "Herb Pharm": (
        "Herbal extracts and liquid botanicals",
        "Premium herbal manufacturer",
        "Premium natural",
        "Herb Pharm",
    ),
    "Sundown Naturals": (
        "Mass-market vitamins including immune/elderberry SKUs",
        "Mass retail vitamin brand",
        "Value to Mid mass",
        "Sundown Naturals",
    ),
    "Elder Farms": (
        "American elderberry farm products and supplements",
        "Farm-to-brand elderberry",
        "Mid to Premium natural",
        "Elder Farms",
    ),
    "Heartland Elderberry Farms": (
        "Farm-produced elderberry syrup and juice",
        "Farm-direct elderberry products",
        "Mid premium natural / craft",
        "Heartland Elderberry Farms",
    ),
    "iHerb Wholesale / B2B channel": (
        "Online vitamins and supplements including elderberry",
        "Global online supplement retail",
        "Mass to Premium",
        "Multi-brand supplements on iHerb",
    ),
    # --- Instant noodles ---
    "Gordon Food Service": (
        "Foodservice wholesale including Asian noodles / dry grocery",
        "North America foodservice distribution",
        "Value to Mid foodservice",
        "Multi-brand foodservice grocery incl. noodles",
    ),
    "McLane Company": (
        "Grocery and convenience wholesale including instant noodles",
        "Convenience / grocery wholesale",
        "Value to Mid mass",
        "Multi-brand CPG including noodle cups",
    ),
    "Performance Food Group": (
        "Foodservice and specialty distribution",
        "USA foodservice wholesale",
        "Value to Mid foodservice",
        "Multi-brand foodservice portfolio",
    ),
    "Performance Food Group (PFG)": (
        "Foodservice and specialty distribution",
        "USA foodservice wholesale",
        "Value to Mid foodservice",
        "Multi-brand foodservice portfolio",
    ),
    "Costco": (
        "Warehouse club retail including instant noodles / Asian foods",
        "Membership warehouse retail",
        "Value to Mid bulk",
        "Kirkland + national noodle brands",
    ),
    "Walmart": (
        "Mass grocery retail including instant noodles",
        "Mass-market grocery retail",
        "Value to Mid mass",
        "Great Value + national noodle brands",
    ),
    "Nestlé Maggi distribution": (
        "Maggi instant noodles and culinary products",
        "Brand owner global noodle distribution",
        "Value to Mid mass",
        "Maggi",
    ),
    "Hethstia": (
        "Konjac / shirataki noodles; high-protein ramen; healthy instant noodles",
        "Healthy noodle manufacturer / exporter",
        "Mid to Premium health / specialty",
        "Hethstia",
    ),
    "Nissin Foods sales companies": (
        "Instant noodles / Cup Noodles brand distribution",
        "Brand owner noodle sales companies",
        "Value to Mid mass",
        "Nissin; Cup Noodles",
    ),
    "Aeon Group wholesale/retail": (
        "Grocery retail including instant noodles (Asia)",
        "Asia supermarket / GMS retail",
        "Value to Mid mass",
        "Private label + national noodle brands",
    ),
    "JD.com supermarket / Jingdong": (
        "Online grocery including instant noodles",
        "China e-grocery",
        "Value to Mid mass",
        "Multi-brand instant noodles on JD",
    ),
    "Alibaba / Tmall supermarket": (
        "Online marketplace grocery including instant noodles",
        "China e-commerce grocery",
        "Value to Mid mass",
        "Multi-brand instant noodles on Tmall",
    ),
    # --- MTB ---
    "Chain Reaction Cycles": (
        "Online bicycle retail including MTB suspension",
        "Global online bike / components retail",
        "Mid to Premium performance",
        "Fox; RockShox; Öhlins; multi-brand MTB parts",
    ),
    "Wiggle": (
        "Online cycling retail including suspension components",
        "Online bike / multisport retail",
        "Mid to Premium performance",
        "Multi-brand MTB suspension and parts",
    ),
    "Bilstein": (
        "Automotive and specialty dampers / suspension",
        "Suspension manufacturer (thyssenkrupp)",
        "Premium performance",
        "Bilstein",
    ),
    "Koni": (
        "Shock absorbers and dampers",
        "Suspension manufacturer",
        "Premium performance",
        "KONI",
    ),
    "DVO Suspension distributors": (
        "MTB suspension forks and shocks",
        "Specialty MTB suspension brand channel",
        "Premium performance",
        "DVO Suspension",
    ),
    "Shimano distributor network": (
        "Bicycle components including related MTB groupsets (not forks primary)",
        "Global bike component brand distribution",
        "Mid to Premium",
        "Shimano",
    ),
    "Pon.Bike distribution (Cannondale, Santa Cruz parent ops)": (
        "Premium bicycle brands and related suspension OEM partners",
        "Premium bike brand group distribution",
        "Premium to Luxury performance",
        "Cannondale; Santa Cruz; related Pon.Bike brands",
    ),
    "Universal Cycles": (
        "Bicycle parts wholesale including suspension",
        "Bike parts wholesale",
        "Mid to Premium",
        "Multi-brand bike components",
    ),
    "Worldwide Cyclery": (
        "Online MTB retail and components",
        "Specialty MTB e-commerce",
        "Mid to Premium",
        "Fox; RockShox; multi-brand MTB",
    ),
    "BikeDiscount / Mantel": (
        "Online EU bike and parts retail",
        "European online bike retail",
        "Value to Premium",
        "Multi-brand MTB suspension",
    ),
    # --- Muslin ---
    "Mood Fabrics": (
        "Apparel fabrics including cotton muslin",
        "Fashion fabric wholesale/retail",
        "Mid to Premium fashion",
        "Mood house assortment; designer fabrics",
    ),
    "JOANN (Fabric & Craft)": (
        "Craft and apparel fabrics including muslin",
        "Mass craft fabric retail",
        "Value to Mid craft",
        "JOANN private label + mill brands",
    ),
    "Joann Stores": (
        "Craft and apparel fabrics including muslin",
        "Mass craft fabric retail",
        "Value to Mid craft",
        "JOANN private label + mill brands",
    ),
    "Hobby Lobby Fabric": (
        "Craft cottons and muslin",
        "Mass craft retail",
        "Value to Mid craft",
        "Hobby Lobby assortment / mill brands",
    ),
    "Fabric.com (Amazon)": (
        "Online fabric retail including muslin",
        "E-commerce fabric retail",
        "Value to Mid",
        "Multi-mill fabrics via Amazon Fabric.com",
    ),
    "ROC-LON Fabrics (Rockland Industries)": (
        "Linings, interlinings, muslin, and industrial textiles",
        "Textile manufacturer / wholesale (est. 1832)",
        "Mid industrial / apparel",
        "ROC-LON",
    ),
    "Whaleys (Bradford) Ltd": (
        "Theatrical muslin, calico, scenic fabrics",
        "UK theatrical textile wholesale",
        "Mid to Premium theatrical",
        "Whaleys",
    ),
    "Liberty Fabrics (Wholesale)": (
        "Printed cottons and specialty fabrics",
        "Heritage fashion fabric wholesale",
        "Premium fashion",
        "Liberty Fabrics",
    ),
    "Rose Brand": (
        "Theatrical muslin, scrim, scenic textiles",
        "Entertainment industry textile supply",
        "Mid to Premium theatrical",
        "Rose Brand",
    ),
    "Sewing Machine Plus": (
        "Sewing machines, notions, and fabrics",
        "Specialty sewing retail",
        "Mid consumer / hobby",
        "Bernina; Brother; multi-brand fabrics",
    ),
    "Casa Fashions (Karur)": (
        "Cotton muslin home textiles and baby muslin products",
        "India manufacturer / exporter",
        "Value to Mid export",
        "Casa Fashions",
    ),
    "Central India Fabrics": (
        "Flannel, muslin, interlining, and cotton fabrics",
        "India manufacturer / exporter",
        "Value to Mid export",
        "Central India Fabrics",
    ),
    "Coats Group Trade Distribution": (
        "Industrial thread and textile inputs",
        "Global industrial textile inputs",
        "Mid industrial",
        "Coats",
    ),
    "OnlineFabricStore": (
        "Online wholesale/retail fabrics including muslin",
        "USA online fabric retail",
        "Value to Mid",
        "Multi-mill fabric assortment",
    ),
    "SyFabrics": (
        "Wholesale apparel fabrics (LA garment district)",
        "Apparel fabric wholesale",
        "Value to Mid",
        "House / imported fabric lots",
    ),
    # --- Nail ---
    "Sally Beauty Supply": (
        "Professional and consumer nail / beauty retail",
        "Beauty supply retail chain",
        "Mass to Professional",
        "Sally Beauty brands; OPI; CND; multi-brand nail",
    ),
    "Armstrong McCall": (
        "Professional beauty wholesale (BSG / Sally)",
        "Salon professional wholesale",
        "Professional Mid to Premium",
        "Professional nail and beauty brands",
    ),
    "CND (Creative Nail Design) distribution": (
        "Professional gel polish and nail systems (SHELLAC)",
        "Professional nail brand distribution",
        "Professional Premium",
        "CND SHELLAC; CND Vinylux",
    ),
    "Gelish / Hand & Nail Harmony distributors": (
        "Professional gel polish systems",
        "Professional nail brand distribution",
        "Professional Premium",
        "Gelish; Hand & Nail Harmony",
    ),
    "Kiara Sky distribution": (
        "Professional gel and dip nail systems",
        "Professional nail brand / e-commerce",
        "Professional Mid to Premium",
        "Kiara Sky",
    ),
    "Young Nails distribution": (
        "Professional nail enhancement systems",
        "Professional nail brand distribution",
        "Professional Premium",
        "Young Nails",
    ),
    "Dearsundays": (
        "Non-toxic / wellness nail polish and salon services",
        "Clean nail brand + studios",
        "Premium clean beauty",
        "sundays / Dear Sundays",
    ),
    "Beauty Joint": (
        "Beauty and nail product wholesale / distribution",
        "Beauty distributor (LA)",
        "Mid to Professional",
        "Multi-brand beauty / nail assortment",
    ),
    "Essie / L'Oréal distribution": (
        "Mass and salon nail color",
        "Global L'Oréal nail brand distribution",
        "Mass to Mid Premium",
        "Essie",
    ),
    "Ulta Beauty": (
        "Beauty retail including nail color",
        "Specialty beauty retail",
        "Mass to Premium",
        "Essie; OPI; Ulta brands; multi-brand nail",
    ),
    "Beauty Bay": (
        "Online beauty retail including nail",
        "Online beauty e-commerce",
        "Mid to Premium",
        "Multi-brand nail and beauty",
    ),
    "Douglas (Europe)": (
        "European perfume and beauty retail including nail",
        "European beauty retail chain",
        "Mid to Premium",
        "Multi-brand European beauty / nail",
    ),
    "Kaganails": (
        "Professional gel polish wholesale (UK)",
        "UK exclusive gel nail distributor",
        "Professional Mid",
        "Kaga; related gel lines",
    ),
    "The Nailhub / pro e-wholesalers": (
        "Professional gel nail supplies e-commerce",
        "Pro nail education + e-wholesale",
        "Professional Mid to Premium",
        "Curated pro gel brands",
    ),
    "Burmax Beauty Distributors": (
        "Salon beauty and nail wholesale",
        "USA salon wholesale",
        "Professional Value to Mid",
        "Multi-brand salon nail / beauty",
    ),
    # --- Sheet masks ---
    "Innisfree": (
        "K-beauty skincare including sheet masks",
        "Amorepacific K-beauty brand retail/distribution",
        "Mass to Mid Premium K-beauty",
        "Innisfree",
    ),
    "Innisfree distribution (Amorepacific)": (
        "K-beauty skincare including sheet masks",
        "Brand distribution (Amorepacific)",
        "Mass to Mid Premium K-beauty",
        "Innisfree",
    ),
    "Laneige / Amorepacific distribution": (
        "K-beauty skincare including sheet masks",
        "Brand distribution (Amorepacific)",
        "Mid to Premium K-beauty",
        "Laneige",
    ),
    "Dr. Jart+ distribution (Estée Lauder)": (
        "K-beauty clinical skincare including sheet masks",
        "Estée Lauder K-beauty brand distribution",
        "Premium K-beauty",
        "Dr. Jart+",
    ),
    "Doctor Althea Global": (
        "Gentle / sensitive-skin K-beauty including sheet masks",
        "K-beauty brand DTC / retail",
        "Mid to Premium clean K-beauty",
        "Dr. Althea",
    ),
    "Peach & Lily": (
        "K-beauty retail and brand including sheet masks",
        "K-beauty specialty retail / brand",
        "Mid to Premium K-beauty",
        "Peach & Lily; curated K-beauty mask brands",
    ),
    "Peach & Lily distribution": (
        "K-beauty retail and brand including sheet masks",
        "K-beauty specialty retail / brand",
        "Mid to Premium K-beauty",
        "Peach & Lily; curated K-beauty mask brands",
    ),
    "Soko Glam": (
        "K-beauty e-commerce including sheet masks",
        "K-beauty specialty e-retail",
        "Mid to Premium K-beauty",
        "Curated K-beauty sheet mask brands",
    ),
    "Sokoglam": (
        "K-beauty e-commerce including sheet masks",
        "K-beauty specialty e-retail",
        "Mid to Premium K-beauty",
        "Curated K-beauty sheet mask brands",
    ),
    "YesStyle": (
        "Asian beauty and fashion e-commerce including sheet masks",
        "Global online K-beauty marketplace",
        "Value to Premium",
        "Multi-brand K-beauty sheet masks",
    ),
    "Sephora K-beauty assortment": (
        "Beauty retail including K-beauty sheet masks",
        "Global specialty beauty retail",
        "Mid to Premium",
        "Dr. Jart+; Laneige; Sephora Collection masks",
    ),
    "Watsons / AS Watson beauty retail": (
        "Drugstore beauty including sheet masks",
        "Asia drugstore retail",
        "Value to Mid mass",
        "Private label + K-beauty mask brands",
    ),
    "Boots UK beauty": (
        "Drugstore beauty including sheet masks",
        "UK drugstore retail",
        "Value to Mid mass",
        "Boots brands + international sheet masks",
    ),
    "Beauty Bay sheet masks": (
        "Online beauty retail sheet masks",
        "Online beauty e-commerce",
        "Mid to Premium",
        "Multi-brand sheet masks",
    ),
    "Douglas Beauty Europe": (
        "European beauty retail including sheet masks",
        "European beauty retail chain",
        "Mid to Premium",
        "Multi-brand European / K-beauty masks",
    ),
    # --- Shunt ---
    "Digi-Key Electronics": (
        "Electronic components including current-sense / shunt resistors",
        "Catalog electronics distribution",
        "Mid to Premium OEM / engineering",
        "Vishay; Bourns; Susumu; multi-brand CSR",
    ),
    "Mouser Electronics": (
        "Electronic components including shunt / current-sense resistors",
        "Catalog electronics distribution",
        "Mid to Premium OEM / engineering",
        "Vishay; Bourns; multi-brand CSR",
    ),
    "Avnet": (
        "Electronic components and technology solutions distribution",
        "Global electronics distribution",
        "Mid to Premium OEM",
        "Multi-supplier passives including shunt resistors",
    ),
    "Future Electronics": (
        "Electronic components distribution",
        "Global electronics distribution",
        "Mid to Premium OEM",
        "Multi-supplier passives including CSR",
    ),
    "RS Components": (
        "Industrial and electronic components including resistors",
        "Industrial / electronics catalog distribution",
        "Mid industrial",
        "Multi-brand resistors / passives",
    ),
    "RS Components / RS Group": (
        "Industrial and electronic components including resistors",
        "Industrial / electronics catalog distribution",
        "Mid industrial",
        "Multi-brand resistors / passives",
    ),
    "TE Connectivity": (
        "Connectors, sensors, and related passives",
        "Manufacturer of interconnect / sensor products",
        "Premium industrial / automotive",
        "TE Connectivity",
    ),
    "Prosemic": (
        "Precision current-sense / shunt resistors; circuit protection",
        "China CSR manufacturer",
        "Mid to Premium OEM",
        "PROSEMI",
    ),
    "Sager Electronics": (
        "Power, thermal, and passives distribution",
        "Specialty electronics distribution",
        "Mid to Premium OEM",
        "Multi-brand power / passives",
    ),
    "Vishay authorized distributor network": (
        "Passive components including shunt resistors",
        "Authorized brand distribution",
        "Mid to Premium OEM",
        "Vishay",
    ),
    # --- Extra verified brand / specialty fills ---
    "Sunrise Medical": (
        "Manual/power wheelchairs; walking aids; axillary & elbow crutches",
        "Mobility manufacturer with dealer distribution",
        "Mid to Premium clinical / consumer mobility",
        "QUICKIE; BREEZY; Coopers; Guardian / Red Dot crutches",
    ),
    "Sunrise Medical dealer network": (
        "Wheelchairs and walking aids including crutches",
        "Authorized mobility dealer channel",
        "Mid to Premium clinical",
        "QUICKIE; BREEZY; Coopers; Sunrise Medical walking aids",
    ),
    "ScripHessco": (
        "Chiropractic and rehab supplies including crutches",
        "Chiropractic / rehab wholesale",
        "Mid professional",
        "Multi-brand chiropractic / DME (incl. crutches)",
    ),
    "Gymedicare": (
        "Home medical equipment; walking aids; bathroom safety",
        "Home-care DME retail / wholesale",
        "Value to Mid consumer DME",
        "Multi-brand home medical / mobility aids",
    ),
    "Health Products For You": (
        "Online DME and medical supplies including crutches",
        "E-commerce medical / DME retail",
        "Value to Mid consumer",
        "Multi-brand DME / mobility assortment",
    ),
    "ProMed Products": (
        "Medical supplies and DME wholesale",
        "Medical supply wholesale",
        "Mid institutional / clinic",
        "Multi-brand medical / mobility",
    ),
    "DME Supply USA": (
        "Durable medical equipment wholesale / retail",
        "DME channel distribution",
        "Value to Mid",
        "Multi-brand DME including crutches",
    ),
    "Activeforever": (
        "Online mobility and rehab products",
        "Consumer DME e-commerce",
        "Value to Mid",
        "Multi-brand mobility / rehab",
    ),
    "Cardinal Health Medical": (
        "Medical products distribution; DME; pharmacy solutions",
        "Integrated healthcare distribution",
        "Mid to Premium institutional",
        "Multi-manufacturer medical portfolio",
    ),
    "Henry Schein Medical": (
        "Medical & dental distribution; practice supplies; DME",
        "Physician / clinic supply distribution",
        "Mid to Premium professional",
        "Multi-brand medical and dental supplies",
    ),
    "Independence Medical (Cardinal specialty)": (
        "Specialty medical distribution (Cardinal)",
        "Specialty healthcare distribution",
        "Mid to Premium institutional",
        "Multi-manufacturer specialty medical",
    ),
    "Drive DeVilbiss Healthcare distribution": (
        "Respiratory and mobility products including walking aids",
        "Brand manufacturer distribution",
        "Value to Mid consumer / clinical",
        "Drive DeVilbiss; Drive Medical mobility",
    ),
    "Graham-Field / Lumex distribution": (
        "Patient aids and mobility products",
        "Brand manufacturer distribution",
        "Value to Mid clinical / consumer",
        "Graham-Field; Lumex",
    ),
    "Nova Medical Products distribution": (
        "Mobility aids and patient support products",
        "Brand manufacturer distribution",
        "Value to Mid consumer",
        "Nova Medical",
    ),
    "Hugo Mobility distribution": (
        "Rollators and mobility aids",
        "Brand manufacturer distribution",
        "Value to Mid consumer",
        "Hugo Mobility",
    ),
    "Invacare dealer network": (
        "Homecare beds, wheelchairs, and related DME",
        "Authorized mobility dealer channel",
        "Mid clinical / consumer",
        "Invacare",
    ),
    "Ottobock dealer network": (
        "Prosthetics, orthotics, and mobility solutions",
        "Authorized clinical dealer channel",
        "Premium clinical",
        "Ottobock",
    ),
    "Össur distribution partners": (
        "Orthopedics, bracing, and mobility",
        "Authorized clinical distribution",
        "Premium clinical",
        "Össur",
    ),
    "Performance Health / TheraBand channel": (
        "Rehab, therapy, and clinical exercise products",
        "Rehab / therapy wholesale",
        "Mid to Premium clinical",
        "TheraBand; Performance Health portfolio",
    ),
    "Patterson Medical / performance health legacy": (
        "Rehab and therapy equipment distribution",
        "Clinical rehab wholesale",
        "Mid to Premium clinical",
        "TheraBand; legacy Patterson Medical portfolio",
    ),
    "Allied Medical": (
        "Clinical medical equipment and supplies",
        "Clinical equipment distribution",
        "Mid to Premium clinical",
        "Multi-brand clinical / audiology-related supplies",
    ),
    "AmerisourceBergen": (
        "Pharmaceutical and medical product wholesale",
        "National pharmaceutical wholesale (Cencora)",
        "Mid to Premium institutional",
        "Multi-manufacturer pharma / medical portfolio",
    ),
    "AmerisourceBergen / Cencora": (
        "Pharmaceutical and consumer-health wholesale",
        "National pharmaceutical wholesale",
        "Mid to Premium institutional",
        "Multi-manufacturer consumer-health / Rx portfolio",
    ),
    "Cedar Gate Technologies": (
        "Healthcare analytics and care-management software",
        "Healthcare IT / analytics (not device OEM)",
        "Enterprise / B2B SaaS",
        "Cedar Gate platform (software; not vestibular hardware brands)",
    ),
    "Nature's Way": (
        "Herbal and dietary supplements including elderberry",
        "Brand manufacturer (Schwabe / Nature's Way)",
        "Mass to Premium natural",
        "Nature's Way; Sambucus elderberry line",
    ),
    "Puritan's Pride": (
        "Vitamins and supplements including immune SKUs",
        "Direct / retail supplement brand (Nature's Bounty Co.)",
        "Value to Mid mass",
        "Puritan's Pride",
    ),
    "Puritan's Pride / The Nature's Bounty Co.": (
        "Vitamins and supplements including immune SKUs",
        "Brand manufacturer / retail",
        "Value to Mid mass",
        "Puritan's Pride; Nature's Bounty",
    ),
    "Swanson Health Products": (
        "Vitamins, herbs, and specialty supplements",
        "Direct-to-consumer supplement retail / brand",
        "Value to Mid",
        "Swanson",
    ),
    "Vitacost": (
        "Online natural products retail including supplements",
        "E-commerce natural retail (Kroger)",
        "Value to Premium natural",
        "Multi-brand supplements incl. elderberry SKUs",
    ),
    "Vitacost (Kroger)": (
        "Online natural products retail including supplements",
        "E-commerce natural retail (Kroger)",
        "Value to Premium natural",
        "Multi-brand supplements incl. elderberry SKUs",
    ),
    "Herbalife": (
        "Nutrition and wellness products",
        "Direct-selling nutrition brand",
        "Mid to Premium",
        "Herbalife Nutrition (elderberry SKUs where offered)",
    ),
    "Elderberry Farms": (
        "Elderberry syrups, extracts, and farm products",
        "Specialty elderberry producer / wholesale",
        "Mid to Premium natural",
        "Elderberry Farms house brand",
    ),
    "Elderberry Farms (Payson UT)": (
        "Elderberry syrups, extracts, and farm products",
        "Specialty elderberry producer / wholesale",
        "Mid to Premium natural",
        "Elderberry Farms house brand",
    ),
    "Elder Farms": (
        "Elderberry beverages and wellness products",
        "Specialty elderberry brand / producer",
        "Mid natural",
        "Elder Farms; Nobleberry",
    ),
    "Heartland Elderberry Farms": (
        "Elderberry farm products and extracts",
        "Specialty elderberry producer",
        "Mid natural",
        "Heartland Elderberry Farms",
    ),
    "Zahler": (
        "Kosher vitamins and supplements",
        "Specialty supplement brand",
        "Mid to Premium",
        "Zahler",
    ),
    "Zahler (Jewish Natural Foods) Wholesale": (
        "Kosher vitamins and supplements",
        "Specialty supplement wholesale",
        "Mid to Premium",
        "Zahler",
    ),
    "Sundown Naturals": (
        "Mass-market vitamins and herbal supplements",
        "Brand manufacturer (Nature's Bounty Co.)",
        "Value to Mid mass",
        "Sundown Naturals",
    ),
    "Sundown Naturals Distribution": (
        "Mass-market vitamins and herbal supplements",
        "Brand distribution",
        "Value to Mid mass",
        "Sundown Naturals",
    ),
    "C&S Wholesale Grocers": (
        "Grocery wholesale including Asian / convenience noodles",
        "Grocery wholesale to retailers",
        "Value to Mid mass grocery",
        "Multi-brand grocery noodles (Nissin, Maruchan, regional)",
    ),
    "Walmart": (
        "Mass retail groceries including instant noodles",
        "Mass-market grocery retail",
        "Value to Mid",
        "Nissin; Maruchan; Maggi; Great Value; Samyang (select)",
    ),
    "Walmart grocery distribution": (
        "Mass retail groceries including instant noodles",
        "Mass-market grocery retail / distribution",
        "Value to Mid",
        "Nissin; Maruchan; Maggi; Great Value; Samyang (select)",
    ),
    "Sysco": (
        "Foodservice distribution including Asian noodles",
        "Foodservice wholesale",
        "Value to Mid foodservice",
        "Multi-brand foodservice noodles / Asian grocery",
    ),
    "Sysco Corporation": (
        "Foodservice distribution including Asian noodles",
        "Foodservice wholesale",
        "Value to Mid foodservice",
        "Multi-brand foodservice noodles / Asian grocery",
    ),
    "US Foods": (
        "Foodservice distribution including Asian noodles",
        "Foodservice wholesale",
        "Value to Mid foodservice",
        "Multi-brand foodservice noodles",
    ),
    "UNFI": (
        "Natural/specialty grocery wholesale",
        "Natural foods wholesale",
        "Mid specialty grocery",
        "Multi-brand specialty / Asian noodles where carried",
    ),
    "Bicycle Warehouse": (
        "Bicycles, parts, and accessories including suspension",
        "Specialty bicycle retail",
        "Mid to Premium performance",
        "Fox; RockShox; multi-brand MTB suspension",
    ),
    "Chain Reaction Cycles": (
        "Online bicycle parts and accessories",
        "Global online bike / components retail",
        "Mid to Premium performance",
        "Fox; RockShox; Öhlins; multi-brand suspension",
    ),
    "Chain Reaction Cycles / Wiggle (SIGNA)": (
        "Online bicycle parts and accessories",
        "Global online bike / components retail",
        "Mid to Premium performance",
        "Fox; RockShox; Öhlins; multi-brand suspension",
    ),
    "Performance Bicycle": (
        "Bicycle retail parts and service",
        "Specialty bicycle retail (legacy / related channels)",
        "Mid to Premium",
        "Fox; RockShox; multi-brand MTB parts",
    ),
    "Bike24": (
        "Online bicycle components and apparel",
        "European online bike retail",
        "Mid to Premium",
        "Fox; RockShox; Öhlins; multi-brand suspension",
    ),
    "Wiggle": (
        "Online bicycle and triathlon retail",
        "Online bike / sports retail",
        "Mid to Premium",
        "Fox; RockShox; multi-brand suspension",
    ),
    "Ohlins": (
        "Premium motorcycle and bicycle suspension",
        "Brand manufacturer (Öhlins)",
        "Premium / race",
        "Öhlins",
    ),
    "Öhlins Bicycle distribution partners": (
        "Premium bicycle suspension",
        "Authorized brand distribution",
        "Premium / race",
        "Öhlins",
    ),
    "Fox Racing Shox dealer/distributor network": (
        "MTB and powersports suspension",
        "Authorized brand distribution",
        "Premium performance",
        "Fox",
    ),
    "SRAM / RockShox distribution": (
        "Bicycle drivetrain and suspension",
        "Authorized brand distribution",
        "Mid to Premium performance",
        "RockShox; SRAM",
    ),
    "Quality Bicycle Products (QBP)": (
        "Bicycle parts wholesale to dealers",
        "Specialty bicycle wholesale (USA)",
        "Mid to Premium",
        "Fox; RockShox; multi-brand dealer parts",
    ),
    "Blick Art Materials": (
        "Art supplies including muslin and canvas fabrics",
        "Art materials retail / wholesale",
        "Value to Mid educational / studio",
        "Blick; multi-brand art muslin / canvas",
    ),
    "Fabric Depot": (
        "Apparel and craft fabrics including muslin",
        "Fabric retail",
        "Value to Mid",
        "Multi-brand craft / apparel muslin",
    ),
    "Hancock Fabrics": (
        "Craft and apparel fabrics (legacy channel brands)",
        "Fabric retail channel",
        "Value to Mid",
        "Multi-brand craft muslin assortment",
    ),
    "Joann Stores": (
        "Craft and apparel fabrics including muslin",
        "Craft fabric retail",
        "Value to Mid",
        "JOANN private label; multi-brand muslin",
    ),
    "JOANN (Fabric & Craft)": (
        "Craft and apparel fabrics including muslin",
        "Craft fabric retail",
        "Value to Mid",
        "JOANN private label; multi-brand muslin",
    ),
    "Sewing Machine Plus": (
        "Sewing machines, parts, and related textiles",
        "Sewing retail / parts",
        "Value to Mid",
        "Multi-brand sewing / fabric accessories",
    ),
    "Sewing Parts Online": (
        "Sewing machine parts and accessories",
        "Sewing parts e-commerce",
        "Value to Mid",
        "Multi-brand sewing parts / accessories",
    ),
    "Mood Fabrics": (
        "Fashion fabrics including muslin and specialty textiles",
        "Fashion fabric retail / wholesale",
        "Mid to Premium fashion",
        "Mood; designer / mill muslin assortment",
    ),
    "ROC-LON Fabrics (Rockland Industries)": (
        "Theatrical and industrial muslin / cotton fabrics",
        "Mill / brand manufacturer",
        "Mid theatrical / industrial",
        "ROC-LON",
    ),
    "Liberty Fabrics (Wholesale)": (
        "Printed fabrics including cotton / muslin-related lines",
        "Heritage fabric brand wholesale",
        "Premium fashion / craft",
        "Liberty Fabrics",
    ),
    "Rose Brand": (
        "Theatrical fabrics, muslin, and stage textiles",
        "Theatrical / event textile supply",
        "Mid to Premium theatrical",
        "Rose Brand muslin / scenic fabrics",
    ),
    "Beauty Joint": (
        "Professional nail and beauty products wholesale",
        "Beauty distributor (LA)",
        "Mid to Professional",
        "Multi-brand pro nail (OPI, CND, Gelish peers)",
    ),
    "Nail Polish Warehouse": (
        "Professional nail polish and salon supplies",
        "Pro nail wholesale / e-commerce",
        "Mass to Professional",
        "Multi-brand salon nail polish / gel systems",
    ),
    "Kaganails": (
        "Gel nail polish and nail art wholesale",
        "Pro / online nail brand wholesale",
        "Value to Mid professional",
        "Kaga; QLZ; Love Easy; BMG gel lines",
    ),
    "CosmoProf (Beauty Systems Group)": (
        "Professional beauty wholesale including nails",
        "Salon professional wholesale",
        "Professional / Premium",
        "OPI; CND; Gelish; multi-brand BSG nail",
    ),
    "Sally Beauty Supply": (
        "Beauty retail including nail color systems",
        "Beauty retail / pro consumer",
        "Mass to Professional",
        "Sally Hansen; OPI; China Glaze; Sally Beauty brands",
    ),
    "SalonCentric (L'Oréal)": (
        "Professional salon distribution (L'Oréal)",
        "Salon professional wholesale",
        "Professional / Premium",
        "Essie; OPI (select); L'Oréal Professionnel nail",
    ),
    "OPI authorized distributors": (
        "Professional nail lacquer and gel systems",
        "Authorized brand distribution",
        "Professional / Premium",
        "OPI",
    ),
    "CND (Creative Nail Design) distribution": (
        "Professional nail systems (Shellac / vinylux)",
        "Authorized brand distribution",
        "Professional / Premium",
        "CND",
    ),
    "Gelish / Hand & Nail Harmony distributors": (
        "Professional gel polish systems",
        "Authorized brand distribution",
        "Professional",
        "Gelish; Hand & Nail Harmony",
    ),
    "Beauty Bay": (
        "Online beauty retail including K-beauty sheet masks",
        "Online beauty e-commerce",
        "Mid to Premium",
        "Innisfree; Laneige; Dr. Jart+; multi-brand sheet masks",
    ),
    "Beauty Bay sheet masks": (
        "Online beauty retail including K-beauty sheet masks",
        "Online beauty e-commerce",
        "Mid to Premium",
        "Innisfree; Laneige; Dr. Jart+; multi-brand sheet masks",
    ),
    "Beauty Bridge": (
        "Beauty retail / distribution including sheet masks",
        "Beauty retail / wholesale",
        "Mass to Premium",
        "Multi-brand sheet masks / K-beauty",
    ),
    "Soko Glam": (
        "K-beauty retail including sheet masks",
        "K-beauty specialty retail",
        "Mid to Premium K-beauty",
        "Mediheal; COSRX; Peach & Lily; multi-brand K-beauty masks",
    ),
    "Sokoglam": (
        "K-beauty retail including sheet masks",
        "K-beauty specialty retail",
        "Mid to Premium K-beauty",
        "Mediheal; COSRX; Peach & Lily; multi-brand K-beauty masks",
    ),
    "Soko Glam wholesale partners": (
        "K-beauty wholesale / retail partners",
        "K-beauty channel partners",
        "Mid to Premium K-beauty",
        "Mediheal; COSRX; multi-brand K-beauty masks",
    ),
    "SkinStore": (
        "Online prestige beauty including sheet masks",
        "Online beauty retail",
        "Mid to Premium",
        "Multi-brand skincare / sheet masks",
    ),
    "The Face Shop": (
        "Korean mass skincare including sheet masks",
        "Brand manufacturer (LG Beauty)",
        "Mass to Mid K-beauty",
        "THE FACE SHOP",
    ),
    "The Face Shop distribution": (
        "Korean mass skincare including sheet masks",
        "Brand distribution",
        "Mass to Mid K-beauty",
        "THE FACE SHOP",
    ),
    "Innisfree": (
        "Korean green-tea inspired skincare including sheet masks",
        "Brand manufacturer (Amorepacific)",
        "Mass to Mid K-beauty",
        "Innisfree",
    ),
    "Innisfree distribution (Amorepacific)": (
        "Korean skincare including sheet masks",
        "Brand distribution (Amorepacific)",
        "Mass to Mid K-beauty",
        "Innisfree",
    ),
    "Target": (
        "Mass retail beauty including sheet masks",
        "Mass-market beauty retail",
        "Value to Mid",
        "Up&Up; multi-brand drugstore / K-beauty masks",
    ),
    "Target Beauty": (
        "Mass retail beauty including sheet masks / nail",
        "Mass-market beauty retail",
        "Value to Mid",
        "Up&Up; multi-brand beauty / nail / masks",
    ),
    "Target Beauty masks": (
        "Mass retail beauty including sheet masks",
        "Mass-market beauty retail",
        "Value to Mid",
        "Up&Up; multi-brand drugstore / K-beauty masks",
    ),
    "Ulta Beauty": (
        "Beauty retail including sheet masks and nail color",
        "Specialty beauty retail",
        "Mass to Premium",
        "Multi-brand sheet masks / nail (Sephora peers at Ulta)",
    ),
    "Ulta Beauty sheet masks": (
        "Beauty retail including sheet masks",
        "Specialty beauty retail",
        "Mass to Premium",
        "Multi-brand sheet masks / K-beauty",
    ),
    "Peach & Lily": (
        "Clean K-beauty retail and brand assortment",
        "K-beauty specialty retail / brand",
        "Mid to Premium clean beauty",
        "Peach & Lily; curated K-beauty sheet masks",
    ),
    "YesStyle": (
        "Asian beauty e-commerce including sheet masks",
        "Online Asian beauty retail",
        "Mass to Premium",
        "Innisfree; Mediheal; Cosrx; multi-brand sheet masks",
    ),
    "Avnet": (
        "Electronic components distribution including passives",
        "Global electronics distribution",
        "Mid to Premium OEM",
        "Vishay; Bourns; Yageo; multi-brand CSR / passives",
    ),
    "Future Electronics": (
        "Electronic components distribution",
        "Global electronics distribution",
        "Mid to Premium OEM",
        "Vishay; Bourns; multi-brand current-sense resistors",
    ),
    "Mouser Electronics": (
        "Catalog electronics distribution",
        "Catalog electronics distribution",
        "Mid to Premium OEM / engineering",
        "Vishay; Bourns; Susumu; Isabellenhütte; multi-brand CSR",
    ),
    "Arrow Electronics": (
        "Electronic components and enterprise IT distribution",
        "Global electronics distribution",
        "Mid to Premium OEM",
        "Vishay; Bourns; multi-brand passives / CSR",
    ),
    "Allied Electronics & Automation": (
        "Industrial automation and electronic components",
        "Industrial / electronics distribution",
        "Mid industrial",
        "Multi-brand resistors / industrial passives",
    ),
    "TTI, Inc.": (
        "Specialty passive / connector / discrete distribution",
        "Specialty electronics distribution",
        "Mid to Premium OEM",
        "Vishay; Bourns; multi-brand passives",
    ),
    "Farnell / element14": (
        "Catalog electronics distribution",
        "Catalog electronics distribution",
        "Mid OEM / engineering",
        "Vishay; Bourns; multi-brand CSR",
    ),
    "Newark": (
        "Catalog electronics distribution (element14/Farnell)",
        "Catalog electronics distribution",
        "Mid OEM / engineering",
        "Vishay; Bourns; multi-brand CSR",
    ),
    "DFS Group (travel retail)": (
        "Travel retail luxury goods including spirits",
        "Airport / travel retail spirits",
        "Premium to Luxury",
        "Moutai; Wuliangye; global luxury spirits (travel retail)",
    ),
    "Shiwandistillery": (
        "Chinese rice wine, cooking wine, Baijiu manufacturing & export",
        "Guangdong traditional liquor producer/exporter",
        "Mass to Mid heritage",
        "Shi Wan Pai; Yubingshao; Shiwan labels",
    ),
}


# Tokens that indicate a real brand / portfolio answer (not a product-type list)
_BRAND_SIGNAL = re.compile(
    r"(?i)\b("
    r"multi[- ]?brand|multi[- ]?manufacturer|portfolio|assortment|private[- ]?label|"
    r"house brands?|authorized|distributor|network|channel|"
    r"vishay|bourns|yageo|susumu|isabellen|prosemi|fox|rockshox|öhlins|ohlins|dvo|"
    r"marzocchi|manitou|shimano|sram|moutai|wuliangye|luzhou|yanghe|nissin|maggi|"
    r"indomie|samyang|nongshim|maruchan|ottogi|opi|essie|cnd|gelish|kiara|orly|"
    r"innisfree|laneige|mediheal|cosrx|dr\.?\s*jart|althea|sambucol|garden of life|"
    r"nature'?s way|now foods|medline|mckesson|alimed|quickie|coopers|guardian|"
    r"interacoustics|natus|amplivox|maico|grason|kirkland|roc-?lon|liberty|"
    r"mood|theraband|lumex|invacare|drive devilbiss|hugo|nova|"
    r"digi-?key|mouser|arrow|avnet|tti|farnell|newark|"
    r"sally beauty|cosmo|saloncentric|ulta|sephora|"
    r"kehe|unfi|sysco|costco|walmart|target|"
    r"peach\s*&\s*lily|boots|up\s*&\s*up|mediheal|cosrx|"
    r"quickie|breezy|coopers|guardian|red\s*dot|"
    r"curated|international sheet"
    r")\b"
)

_PRODUCT_LIST_HINTS = re.compile(
    r"(?i)\b("
    r"crutches|mobility aids|medical supplies|healthcare solutions?|"
    r"sheet (?:face )?masks|hydrogel masks|cotton sheet|"
    r"nail polish|nail care|nail art|"
    r"instant noodles?|cup noodles?|ramen|noodle (?:cups|bowls|packs)|"
    r"suspension forks?|shock absorbers?|dropper posts?|"
    r"muslin (?:fabric|rolls|bags|sheets)|quilting fabric|"
    r"shunt resistors?|semiconductors?|capacitors?|connectors?|"
    r"elderberry (?:extract|syrup|gummies|capsules|powder)|"
    r"bpv treatment|vestibular therapy|diagnostic tools|"
    r"wholesale|sourcing|buying|requests|locator|logo|goldsupplier"
    r")\b"
)


def is_weak(v: object) -> bool:
    s = str(v or "").strip().lower()
    if not s:
        return True
    if s.startswith("not ") or "not confirm" in s or "not publicly" in s:
        return True
    if s in {"n/a", "na", "unknown", "-", "—", "service provider", "organic/premium"}:
        return True
    if s.startswith("b2b distribution for "):
        return True
    if s == "category product wholesale":
        return True
    if s == "multi-brand / assortment varies":
        return True
    # junk scraped tokens
    if s in {"locator; from; race; track", "logo; dark; best; professional"}:
        return True
    if "goldsupplier" in s or s.startswith("logo;"):
        return True
    return False


def looks_like_product_list(brands: object, core: object = None) -> bool:
    """True when Key Brands is really a product-category dump, not brand names."""
    b = str(brands or "").strip()
    if not b:
        return True
    if is_weak(b):
        return True
    if _BRAND_SIGNAL.search(b):
        return False
    c = str(core or "").strip()
    if c and b.lower().replace(";", ",").replace("  ", " ") == c.lower().replace(";", ",").replace("  ", " "):
        return True
    # Many semicolon/comma parts that are product phrases
    parts = [p.strip() for p in re.split(r"[;,]", b) if p.strip()]
    if len(parts) >= 2:
        productish = sum(1 for p in parts if _PRODUCT_LIST_HINTS.search(p) or p.islower() or p[:1].islower())
        if productish >= max(2, len(parts) - 1) and not _BRAND_SIGNAL.search(b):
            return True
    if _PRODUCT_LIST_HINTS.search(b) and not _BRAND_SIGNAL.search(b):
        return True
    return False


def needs_specialty_fix(v: object) -> bool:
    s = str(v or "").strip().lower()
    return is_weak(v) or s == "service provider" or s.startswith("b2b distribution for")


def needs_brands_fix(v: object, core: object = None) -> bool:
    s = str(v or "").strip().lower()
    if is_weak(v) or s == "multi-brand / assortment varies":
        return True
    return looks_like_product_list(v, core)


def _norm_key(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", " ", (name or "").lower()).strip()
    s = re.sub(r"\s+", " ", s)
    # drop common suffixes for matching
    s = re.sub(r"\b(inc|llc|ltd|co|corp|corporation|company|group|ag|plc)\b", "", s)
    return re.sub(r"\s+", " ", s).strip()


_OVERRIDE_BY_NORM: dict[str, tuple[str, str, str, str]] = {
    _norm_key(k): v for k, v in COMPANY_OVERRIDES.items()
}

# Market-specific overrides win for companies that appear in multiple markets
MARKET_COMPANY_OVERRIDES: dict[str, dict[str, tuple[str, str, str, str]]] = {
    "nail_polish_market_global": {
        "Beauty Bay": (
            "Online beauty retail including nail color",
            "Online beauty e-commerce",
            "Mid to Premium",
            "Essie; OPI; multi-brand nail / beauty",
        ),
        "Ulta Beauty": (
            "Beauty retail including nail color",
            "Specialty beauty retail",
            "Mass to Premium",
            "Essie; OPI; Ulta brands; multi-brand nail",
        ),
        "Target": (
            "Mass retail beauty including nail color",
            "Mass-market beauty retail",
            "Value to Mid",
            "Essie; OPI; Target private-label nail",
        ),
        "Target Beauty": (
            "Mass retail beauty including nail color",
            "Mass-market beauty retail",
            "Value to Mid",
            "Essie; OPI; Target private-label nail",
        ),
        "Walmart": (
            "Mass retail beauty including nail color",
            "Mass-market beauty retail",
            "Value to Mid",
            "Sally Hansen; Wet n Wild; multi-brand nail",
        ),
        "Walmart Beauty assortment": (
            "Mass retail beauty including nail color",
            "Mass-market beauty retail",
            "Value to Mid",
            "Sally Hansen; Wet n Wild; multi-brand nail",
        ),
        "Sephora": (
            "Prestige beauty retail including nail",
            "Prestige beauty retail",
            "Premium",
            "Essie; OPI; Sephora Collection nail",
        ),
        "Douglas (Europe)": (
            "European beauty retail including nail",
            "European beauty retail chain",
            "Mid to Premium",
            "Essie; OPI; multi-brand European nail",
        ),
        "CVS Beauty": (
            "Drugstore beauty including nail color",
            "Drugstore beauty retail",
            "Value to Mid",
            "Essie; Sally Hansen; CVS private-label nail",
        ),
        "Amazon Beauty wholesale brands": (
            "Online beauty marketplace including nail",
            "Marketplace beauty retail",
            "Value to Premium",
            "Multi-brand nail color / Amazon beauty",
        ),
    },
    "sheet_face_masks_market_global": {
        "Beauty Bay": (
            "Online beauty retail including K-beauty sheet masks",
            "Online beauty e-commerce",
            "Mid to Premium",
            "Innisfree; Laneige; Dr. Jart+; multi-brand sheet masks",
        ),
        "Beauty Bay sheet masks": (
            "Online beauty retail including K-beauty sheet masks",
            "Online beauty e-commerce",
            "Mid to Premium",
            "Innisfree; Laneige; Dr. Jart+; multi-brand sheet masks",
        ),
        "Ulta Beauty": (
            "Beauty retail including sheet masks",
            "Specialty beauty retail",
            "Mass to Premium",
            "Multi-brand sheet masks / K-beauty",
        ),
        "Ulta Beauty sheet masks": (
            "Beauty retail including sheet masks",
            "Specialty beauty retail",
            "Mass to Premium",
            "Multi-brand sheet masks / K-beauty",
        ),
        "Target": (
            "Mass retail beauty including sheet masks",
            "Mass-market beauty retail",
            "Value to Mid",
            "Up&Up; multi-brand drugstore / K-beauty masks",
        ),
        "Target Beauty": (
            "Mass retail beauty including sheet masks",
            "Mass-market beauty retail",
            "Value to Mid",
            "Up&Up; multi-brand drugstore / K-beauty masks",
        ),
        "Target Beauty masks": (
            "Mass retail beauty including sheet masks",
            "Mass-market beauty retail",
            "Value to Mid",
            "Up&Up; multi-brand drugstore / K-beauty masks",
        ),
        "Sephora K-beauty assortment": (
            "Prestige K-beauty including sheet masks",
            "Prestige beauty retail",
            "Premium",
            "Dr. Jart+; Laneige; Sephora Collection masks",
        ),
        "Douglas Beauty Europe": (
            "European beauty retail including sheet masks",
            "European beauty retail chain",
            "Mid to Premium",
            "Multi-brand European / K-beauty sheet masks",
        ),
        "CVS Beauty masks": (
            "Drugstore beauty including sheet masks",
            "Drugstore beauty retail",
            "Value to Mid",
            "Multi-brand drugstore sheet masks",
        ),
        "Walgreens Beauty": (
            "Drugstore beauty including sheet masks",
            "Drugstore beauty retail",
            "Value to Mid",
            "Multi-brand drugstore sheet masks",
        ),
        "Amazon Beauty sheet mask sellers": (
            "Online beauty marketplace including sheet masks",
            "Marketplace beauty retail",
            "Value to Premium",
            "Multi-brand sheet masks / K-beauty",
        ),
    },
    "instant_noodles_market_global": {
        "Walmart": (
            "Mass retail groceries including instant noodles",
            "Mass-market grocery retail",
            "Value to Mid",
            "Nissin; Maruchan; Maggi; Great Value; Samyang (select)",
        ),
        "Walmart grocery distribution": (
            "Mass retail groceries including instant noodles",
            "Mass-market grocery retail / distribution",
            "Value to Mid",
            "Nissin; Maruchan; Maggi; Great Value; Samyang (select)",
        ),
        "C&S Wholesale Grocers": (
            "Grocery wholesale including Asian / convenience noodles",
            "Grocery wholesale to retailers",
            "Value to Mid mass grocery",
            "Multi-brand grocery noodles (Nissin, Maruchan, regional)",
        ),
        "Costco Wholesale": (
            "Warehouse club retail including instant noodles",
            "Membership warehouse retail",
            "Value to Mid bulk",
            "Kirkland + Nissin / Maruchan / Asian noodle brands",
        ),
        "KeHE Distributors": (
            "Natural & specialty grocery wholesale including Asian noodles",
            "Natural / specialty grocery wholesale",
            "Mid specialty grocery",
            "Multi-brand specialty / Asian noodles",
        ),
        "UNFI": (
            "Natural/specialty grocery wholesale",
            "Natural foods wholesale",
            "Mid specialty grocery",
            "Multi-brand specialty / Asian noodles where carried",
        ),
    },
    "elderberry_supplements_market_global": {
        "C&S Wholesale Grocers": (
            "Grocery wholesale including OTC supplements",
            "Grocery wholesale to retailers",
            "Value to Mid mass",
            "Multi-brand elderberry / immune supplements",
        ),
        "KeHE Distributors": (
            "Natural & specialty grocery wholesale including supplements",
            "Natural CPG wholesale to retailers",
            "Mid to Premium natural",
            "Garden of Life; Nature's Way; Sambucol; multi-brand elderberry",
        ),
        "UNFI (United Natural Foods)": (
            "Natural/organic grocery & supplement wholesale",
            "National natural foods wholesale",
            "Mid to Premium natural",
            "Garden of Life; Nature's Way; NOW Foods; multi-brand elderberry",
        ),
        "Cardinal Health": (
            "Pharmaceutical and consumer-health wholesale",
            "National pharmaceutical wholesale",
            "Mid to Premium institutional / retail Rx",
            "Multi-manufacturer OTC / consumer-health (incl. elderberry SKUs)",
        ),
        "McKesson": (
            "Pharmaceutical and consumer-health wholesale",
            "National pharmaceutical wholesale",
            "Mid to Premium institutional / retail Rx",
            "Multi-manufacturer OTC / consumer-health",
        ),
        "McKesson (Consumer Health)": (
            "Pharmaceutical and consumer-health wholesale",
            "National pharmaceutical wholesale",
            "Mid to Premium institutional / retail Rx",
            "Multi-manufacturer OTC / consumer-health",
        ),
        "Costco Wholesale (Kirkland / brand aisle)": (
            "Warehouse club retail including supplements",
            "Membership warehouse retail",
            "Value to Mid",
            "Kirkland; national elderberry / immune brands",
        ),
    },
    "axillary_crutches_market_global": {
        "Cardinal Health": (
            "Medical products distribution; DME; pharmacy solutions",
            "Integrated healthcare distribution",
            "Mid to Premium institutional",
            "Multi-manufacturer medical / DME portfolio",
        ),
        "Henry Schein": (
            "Medical & dental distribution; practice supplies; DME",
            "Physician / clinic supply distribution",
            "Mid to Premium professional",
            "Multi-brand medical and dental / mobility supplies",
        ),
        "McKesson": (
            "Pharmaceutical & medical-surgical wholesale; DME; OTC mobility",
            "National pharmaceutical and medical supply wholesale",
            "Mid to Premium institutional",
            "Multi-manufacturer medical / DME portfolio",
        ),
    },
    "benign_positional_vertigo_market_global": {
        "Cardinal Health": (
            "Medical products distribution; clinic supplies",
            "Integrated healthcare distribution",
            "Mid to Premium institutional",
            "Multi-manufacturer medical portfolio (clinic / ENT adjacent)",
        ),
        "Henry Schein": (
            "Medical & dental distribution; ENT / clinic supplies",
            "Physician / clinic supply distribution",
            "Mid to Premium professional",
            "Multi-brand medical / ENT clinic supplies",
        ),
        "Henry Schein Medical (ENT/audiology)": (
            "Medical distribution with ENT / audiology focus",
            "Physician / clinic supply distribution",
            "Mid to Premium professional",
            "Multi-brand ENT / audiology clinic supplies",
        ),
        "McKesson": (
            "Pharmaceutical & medical-surgical wholesale",
            "National pharmaceutical and medical supply wholesale",
            "Mid to Premium institutional",
            "Multi-manufacturer medical portfolio",
        ),
        "Medline Industries": (
            "Manufacturer/distributor of medical supplies",
            "Vertically integrated medical supply",
            "Value to Mid institutional",
            "Medline house brands; multi-source clinic supplies",
        ),
        "Medline": (
            "Manufacturer/distributor of medical supplies",
            "Vertically integrated medical supply",
            "Value to Mid institutional",
            "Medline house brands; multi-source clinic supplies",
        ),
    },
    "shunt_resistor_market_global": {
        "Avnet": (
            "Electronic components distribution including passives",
            "Global electronics distribution",
            "Mid to Premium OEM",
            "Vishay; Bourns; Yageo; multi-brand CSR / passives",
        ),
        "Mouser Electronics": (
            "Catalog electronics distribution",
            "Catalog electronics distribution",
            "Mid to Premium OEM / engineering",
            "Vishay; Bourns; Susumu; Isabellenhütte; multi-brand CSR",
        ),
        "RS Components": (
            "Industrial and electronic components including resistors",
            "Industrial / electronics catalog distribution",
            "Mid industrial",
            "Vishay; Bourns; multi-brand resistors / passives",
        ),
        "RS Components / RS Group": (
            "Industrial and electronic components including resistors",
            "Industrial / electronics catalog distribution",
            "Mid industrial",
            "Vishay; Bourns; multi-brand resistors / passives",
        ),
    },
}


def _find_in_map(company: str, mapping: dict[str, tuple[str, str, str, str]]):
    if company in mapping:
        return mapping[company]
    nk = _norm_key(company)
    by_norm = {_norm_key(k): v for k, v in mapping.items()}
    if nk in by_norm:
        return by_norm[nk]
    for key, val in mapping.items():
        kn = _norm_key(key)
        if kn and (kn in nk or nk in kn) and min(len(kn), len(nk)) >= 8:
            return val
    return None


def lookup_override(market: str, company: str):
    mkt = MARKET_COMPANY_OVERRIDES.get(market) or {}
    hit = _find_in_map(company, mkt)
    if hit:
        return hit
    return _find_in_map(company, COMPANY_OVERRIDES)


def ensure_price_column(ws, header_row: int, header: list[str]) -> list[str]:
    """Insert Price Segment after Specialty Focus if missing. Return updated header list."""
    if "Price Segment" in header:
        return header
    # Find Specialty Focus col (1-based)
    if "Specialty Focus" in header:
        insert_at = header.index("Specialty Focus") + 2  # after specialty
    elif "Key Brands Represented" in header:
        insert_at = header.index("Key Brands Represented") + 1
    else:
        insert_at = len(header) + 1
    ws.insert_cols(insert_at)
    ws.cell(header_row, insert_at).value = "Price Segment"
    # rebuild header from sheet
    new_header = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(header_row, col).value
        if val is None and col > insert_at + 5:
            break
        new_header.append(str(val or "").strip())
    # trim trailing empties
    while new_header and not new_header[-1]:
        new_header.pop()
    return new_header


def resolve_fields(market: str, company: str, core, spec, brands, price):
    mdef = MARKET_DEFAULTS.get(market)
    ovr = lookup_override(market, company)

    if ovr:
        # Company-specific curated values always win (verify/replace weak scraped text)
        return ovr[0], ovr[1], ovr[2], ovr[3]

    if not mdef:
        return (
            str(core or "").strip(),
            str(spec or "").strip(),
            str(price or "").strip(),
            str(brands or "").strip(),
        )

    c_def, s_def, p_def, b_def = mdef
    out_c = str(core or "").strip() if not is_weak(core) else c_def
    out_s = str(spec or "").strip() if not needs_specialty_fix(spec) else s_def
    out_p = str(price or "").strip() if not is_weak(price) else p_def
    out_b = str(brands or "").strip() if not needs_brands_fix(brands, core) else b_def

    if not out_c:
        out_c = c_def
    if not out_s:
        out_s = s_def
    if not out_p:
        out_p = p_def
    if not out_b:
        out_b = b_def
    return out_c, out_s, out_p, out_b


def process(path: Path) -> dict:
    market = path.parent.name
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=False))
    hdr_i = next(
        i
        for i, r in enumerate(rows[:12])
        if r and r[0].value and str(r[0].value).strip() == "Company"
    )
    header_row = hdr_i + 1
    header = [str(c.value or "").strip() for c in rows[hdr_i]]
    header = ensure_price_column(ws, header_row, header)

    # Re-read cells after possible insert
    rows = list(ws.iter_rows(values_only=False))
    header = [str(c.value or "").strip() for c in rows[hdr_i]]
    # map indices
    idx = {h: header.index(h) for h in header if h}

    ci = idx["Company"]
    core_i = idx["Core Categories"]
    spec_i = idx["Specialty Focus"]
    brand_i = idx["Key Brands Represented"]
    price_i = idx["Price Segment"]

    updated = 0
    for r in rows[hdr_i + 1 :]:
        company = str(r[ci].value or "").strip()
        if not company:
            continue
        oc, os_, op, ob = resolve_fields(
            market,
            company,
            r[core_i].value,
            r[spec_i].value,
            r[brand_i].value,
            r[price_i].value,
        )
        changed = False
        for cell, new in (
            (r[core_i], oc),
            (r[spec_i], os_),
            (r[price_i], op),
            (r[brand_i], ob),
        ):
            if new and str(cell.value or "").strip() != new:
                cell.value = new
                changed = True
        if changed:
            updated += 1

    if "Run Summary" in wb.sheetnames:
        wb["Run Summary"]["A17"] = f"Core/Specialty/Price/Brands fill: {updated} rows"

    wb.save(path)
    wb.close()
    return {"updated": updated, "market": market}


def audit() -> None:
    print(f"{'Market':<36} {'weak_core':>9} {'weak_spec':>9} {'weak_price':>10} {'weak_brand':>10}")
    for path in sorted(ROOT.rglob("*_distributors_expanded.xlsx")):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        data = list(ws.iter_rows(values_only=True))
        hdr_i = next(i for i, r in enumerate(data[:12]) if r and str(r[0] or "").strip() == "Company")
        h = [str(x or "").strip() for x in data[hdr_i]]
        rows = [r for r in data[hdr_i + 1 :] if r and r[0]]
        def wcount(col):
            if col not in h:
                return -1
            i = h.index(col)
            return sum(1 for r in rows if is_weak(r[i]) or (col == "Specialty Focus" and needs_specialty_fix(r[i])) or (col == "Key Brands Represented" and needs_brands_fix(r[i])))
        # stricter: after fill, weak should be empty-only
        def emptyish(col):
            if col not in h:
                return -1
            i = h.index(col)
            core_i = h.index("Core Categories") if "Core Categories" in h else None
            n = 0
            for r in rows:
                v = r[i]
                if col == "Key Brands Represented":
                    core_v = r[core_i] if core_i is not None else None
                    if needs_brands_fix(v, core_v):
                        n += 1
                elif col == "Specialty Focus":
                    if needs_specialty_fix(v):
                        n += 1
                elif is_weak(v) or not str(v or "").strip():
                    n += 1
            return n
        print(
            f"{path.parent.name:<36} {emptyish('Core Categories'):>9} {emptyish('Specialty Focus'):>9} {emptyish('Price Segment'):>10} {emptyish('Key Brands Represented'):>10}"
        )
        wb.close()


def main() -> None:
    for path in sorted(ROOT.rglob("*_distributors_expanded.xlsx")):
        st = process(path)
        print(f"{st['market']}: updated {st['updated']}")
    print("\nPost-fill empty audit:")
    audit()


if __name__ == "__main__":
    main()
