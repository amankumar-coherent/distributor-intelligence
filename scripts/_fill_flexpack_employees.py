#!/usr/bin/env python3
"""Fill Employees as LinkedIn-style ranges — evidence only (site scrape + DeepSeek).

Never invent. If no evidence → Not publicly disclosed.
Normalize all values to standard ranges with en-dash.
"""
from __future__ import annotations

import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urljoin

import httpx
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook

load_dotenv(r"D:\vendor-intel\.env")

PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = PATH.parent / "flexpack_employees_fill_audit.json"
CACHE = PATH.parent / "flexpack_employees_cache.json"
NPD = "Not publicly disclosed"
WORKERS = int(os.getenv("EMP_WORKERS") or "8")
USE_DEEPSEEK = os.getenv("EMP_USE_DEEPSEEK", "1") != "0"

RANGES = (
    "1–10",
    "11–50",
    "51–200",
    "201–500",
    "501–1,000",
    "1,001–5,000",
    "5,001–10,000",
    "10,001–50,000",
    "50,001+",
)

# Publicly evidenced / well-known company sizes (LinkedIn or company filings style bands).
# Only include when confident for THIS legal/brand entity.
HAND: dict[str, str] = {
    # LinkedIn / filings / reputable directories — mapped to bands
    "Acme Paper & Supply": "51–200",  # LinkedIn ~73
    "Ahlsell": "5,001–10,000",  # ~5,200
    "Antalis": "1,001–5,000",
    "Askul": "1,001–5,000",
    "Associated Bag": "501–1,000",
    "Atlantic Packaging Corp": "201–500",  # LinkedIn ~432
    "Bags and Bows": "51–200",
    "Ball & Doggett": "51–200",  # ~189
    "Blackwoods": "501–1,000",  # ~686 reported
    "BOX Partners": "51–200",
    "BradyPLUS": "501–1,000",  # LinkedIn ~841
    "Bunzl": "10,001–50,000",
    "Cospak": "51–200",
    "Fastenal": "10,001–50,000",
    "Global Industrial": "1,001–5,000",
    "Grainger": "10,001–50,000",
    "HD Supply": "10,001–50,000",
    "Hisco": "501–1,000",
    "Imperial Dade": "5,001–10,000",
    "Imperial Supplies": "501–1,000",
    "Kite Packaging": "201–500",
    "Macfarlane Group": "1,001–5,000",  # ~1,000+
    "MAC Papers": "51–200",  # ~170
    "MSC Industrial Supply": "5,001–10,000",
    "Network Packaging": "51–200",
    "RAJA": "1,001–5,000",
    "R.J. Schinner Co.": "501–1,000",  # over 600
    "Shorr Packaging": "501–1,000",
    "Uline": "5,001–10,000",  # ~9,000
    "Veritiv": "1,001–5,000",
    "Weald Packaging Supplies": "1–10",  # 9
}

# Force re-check (likely wrong prior values)
FORCE_CLEAR = {
    "Packaging Warehouse",  # was 5,001–10,000 — wrong entity / invented
    "Industrial Packaging",  # was 50,001+ — implausible for channel row
    "Global Packaging",  # was 201–500 — thin evidence
    "Bulk Packaging",  # was 51–200 — dead site
}

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)


def emptyish(v: object) -> bool:
    s = str(v or "").strip().lower()
    return s in {
        "",
        "n/a",
        "na",
        "-",
        "—",
        "not publicly disclosed",
        "unknown",
        "tbd",
        "null",
        "none",
    }


def emp_num_to_range(n: int) -> str:
    if n <= 0:
        return ""
    if n <= 10:
        return "1–10"
    if n <= 50:
        return "11–50"
    if n <= 200:
        return "51–200"
    if n <= 500:
        return "201–500"
    if n <= 1000:
        return "501–1,000"
    if n <= 5000:
        return "1,001–5,000"
    if n <= 10000:
        return "5,001–10,000"
    if n <= 50000:
        return "10,001–50,000"
    return "50,001+"


def normalize_existing(v: str) -> str:
    """Map messy values to a standard range when possible; else ''."""
    s = str(v or "").strip()
    if emptyish(s):
        return ""
    # already a standard range (any dash variant)
    s2 = s.replace("�", "–").replace("-", "–").replace("—", "–").replace("–", "–")
    s2 = re.sub(r"\s+", "", s2)
    # try match known
    for r in RANGES:
        if s2.lower() == r.replace(" ", "").lower() or s2.lower() == r.lower().replace("–", "-"):
            return r
        if s2.replace(",", "") == r.replace(",", "").replace("–", "-"):
            return r
    # common patterns
    m = re.search(r"(\d[\d,]*)\s*[–\-�—]\s*(\d[\d,]*)", s)
    if m:
        try:
            lo = int(m.group(1).replace(",", ""))
            hi = int(m.group(2).replace(",", ""))
            mid = (lo + hi) // 2
            return emp_num_to_range(mid)
        except ValueError:
            pass
    m = re.search(r"(?:~|about|over|more than)?\s*(\d[\d,]*)\+?", s, re.I)
    if m and re.search(r"employee|staff|\+|over", s, re.I):
        try:
            n = int(m.group(1).replace(",", ""))
            if 1 <= n <= 2_000_000 and n not in range(1900, 2031):
                return emp_num_to_range(n)
        except ValueError:
            pass
    # bare small integer like "9"
    if re.fullmatch(r"\d{1,4}", s):
        n = int(s)
        if 1 <= n <= 5000:
            return emp_num_to_range(n)
    return ""


def strip_html(html: str) -> str:
    t = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    t = re.sub(r"(?is)<style.*?>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def scrape_evidence(website: str) -> str:
    url = (website or "").strip()
    if not url:
        return ""
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    texts: list[str] = []
    paths = [
        "",
        "/about",
        "/about-us",
        "/pages/about-us",
        "/company",
        "/our-story",
        "/careers",
        "/about/company",
    ]
    try:
        with httpx.Client(
            headers={"User-Agent": UA},
            verify=False,
            timeout=httpx.Timeout(12.0, connect=4.0),
            follow_redirects=True,
        ) as client:
            for path in paths:
                try:
                    u = url if not path else urljoin(url.rstrip("/") + "/", path.lstrip("/"))
                    r = client.get(u)
                    if r.status_code >= 400:
                        continue
                    t = strip_html(r.text or "")
                    if len(t) > 120:
                        texts.append(t[:3500])
                    if len(texts) >= 2:
                        break
                except Exception:
                    continue
    except Exception:
        return ""
    return "\n---\n".join(texts)[:8000]


def parse_employees_from_text(text: str) -> int | None:
    if not text:
        return None
    patterns = [
        r"company size[:\s]+([\d,]+)\s*[–\-]\s*([\d,]+)",
        r"([\d,]+)\s*[–\-]\s*([\d,]+)\s+employees",
        r"([\d,]+(?:\.\d+)?)\s*(?:thousand|k)\+?\s*(?:employees|staff|people|team members)",
        r"(?:over|more than|approximately|about|around|nearly)\s+([\d,]+)\s*(?:employees|staff|people|team members)",
        r"([\d,]+)\s*\+?\s*(?:employees|staff|people|team members|colleagues)",
        r"(?:employees|staff|headcount|workforce)[:\s]+([\d,]+)",
        r"employs\s+(?:about\s+|approximately\s+|over\s+|more than\s+)?([\d,]+)",
        r"workforce of\s+(?:about\s+|approximately\s+|over\s+)?([\d,]+)",
        r"team of\s+(?:over\s+|more than\s+)?([\d,]+)",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if not m:
            continue
        # range form
        if m.lastindex and m.lastindex >= 2 and m.group(2) and re.search(r"[–\-]", m.group(0)):
            try:
                lo = int(m.group(1).replace(",", ""))
                hi = int(m.group(2).replace(",", ""))
                if 1 <= lo <= hi <= 2_000_000:
                    return (lo + hi) // 2
            except ValueError:
                pass
        raw = m.group(1).replace(",", "")
        try:
            n = float(raw)
        except ValueError:
            continue
        g0 = m.group(0).lower()
        if "thousand" in g0 or re.search(r"\bk\b", g0):
            n *= 1000
        n = int(n)
        if 1 <= n <= 2_000_000 and n not in range(1900, 2031):
            return n
    # LinkedIn-style textual bands in page
    band_map = {
        r"1[-–]10\s+employees": "1–10",
        r"11[-–]50\s+employees": "11–50",
        r"51[-–]200\s+employees": "51–200",
        r"201[-–]500\s+employees": "201–500",
        r"501[-–]1,?000\s+employees": "501–1,000",
        r"1,?001[-–]5,?000\s+employees": "1,001–5,000",
        r"5,?001[-–]10,?000\s+employees": "5,001–10,000",
        r"10,?001[-–]50,?000\s+employees": "10,001–50,000",
        r"10,?001\+\s+employees": "10,001–50,000",
        r"50,?001\+\s+employees": "50,001+",
    }
    for pat, rng in band_map.items():
        if re.search(pat, text, re.I):
            # encode as sentinel via midpoint
            mid = {
                "1–10": 5,
                "11–50": 30,
                "51–200": 100,
                "201–500": 350,
                "501–1,000": 750,
                "1,001–5,000": 3000,
                "5,001–10,000": 7500,
                "10,001–50,000": 25000,
                "50,001+": 60000,
            }[rng]
            return mid
    return None


def deepseek_extract(client: OpenAI | None, model: str, name: str, web: str, evidence: str) -> str:
    if not client or not evidence or len(evidence) < 80:
        return ""
    system = (
        "Extract employee headcount from EVIDENCE only. Never invent or guess. "
        'JSON: {"range":"","confidence":"high|medium|low|none","reason":""}. '
        f"range MUST be one of: {', '.join(RANGES)} OR empty string. "
        "If evidence lacks a clear headcount/band, return range=\"\" confidence=none."
    )
    user = f"Company: {name}\nWebsite: {web}\n\nEVIDENCE:\n{evidence[:6500]}"
    try:
        resp = client.chat.completions.create(
            model=model,
            temperature=0,
            max_tokens=160,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
        )
        data = json.loads((resp.choices[0].message.content or "{}").strip())
        rng = str(data.get("range") or "").strip().replace("-", "–")
        conf = str(data.get("confidence") or "").lower()
        # normalize commas
        for r in RANGES:
            if rng.replace(",", "") == r.replace(",", "") or rng == r:
                rng = r
                break
        else:
            # try map mid
            n = parse_employees_from_text(rng)
            rng = emp_num_to_range(n) if n else ""
        if conf in {"high", "medium"} and rng in RANGES:
            return rng
    except Exception:
        return ""
    return ""


def load_cache() -> dict:
    if CACHE.exists():
        try:
            return json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_cache(cache: dict) -> None:
    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")


def process_one(
    name: str,
    web: str,
    client: OpenAI | None,
    model: str,
    cache: dict,
) -> tuple[str, str, str]:
    """Return (range, source, note)."""
    if name in HAND:
        return HAND[name], "hand", "known public band"

    key = f"{name}|{web}".lower()
    if key in cache and cache[key].get("range"):
        return cache[key]["range"], cache[key].get("source", "cache"), cache[key].get("note", "")

    evidence = scrape_evidence(web)
    n = parse_employees_from_text(evidence)
    if n:
        rng = emp_num_to_range(n)
        cache[key] = {"range": rng, "source": "site_regex", "note": f"n={n}"}
        return rng, "site_regex", f"n={n}"

    if USE_DEEPSEEK and evidence:
        rng = deepseek_extract(client, model, name, web, evidence)
        if rng:
            cache[key] = {"range": rng, "source": "deepseek", "note": "from site evidence"}
            return rng, "deepseek", "from site evidence"
        time.sleep(0.15)

    cache[key] = {"range": "", "source": "none", "note": "no evidence"}
    return "", "none", "no evidence"


def main() -> None:
    key = os.getenv("OPENAI_API_KEY", "").strip()
    client = (
        OpenAI(
            api_key=key,
            base_url=(os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1").rstrip("/"),
            timeout=60.0,
        )
        if key and USE_DEEPSEEK
        else None
    )
    model = (os.getenv("OPENAI_MODEL") or "deepseek-chat").strip()
    cache = load_cache()

    wb = load_workbook(PATH)
    ws = wb.active
    h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ic = {n: h.index(n) + 1 for n in h}

    # Normalize / clear first
    cleared = normalized = 0
    targets: list[tuple[int, str, str]] = []
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, ic["Company"]).value or "").strip()
        web = str(ws.cell(r, ic["Website"]).value or "").strip()
        cur = str(ws.cell(r, ic["Employees"]).value or "").strip()

        if name in FORCE_CLEAR:
            ws.cell(r, ic["Employees"]).value = None
            cur = ""
            cleared += 1

        if not emptyish(cur):
            norm = normalize_existing(cur)
            if name in HAND:
                ws.cell(r, ic["Employees"]).value = HAND[name]
                normalized += 1
            elif norm:
                if norm != cur:
                    ws.cell(r, ic["Employees"]).value = norm
                    normalized += 1
            else:
                # unparseable → clear for re-fill
                ws.cell(r, ic["Employees"]).value = None
                cur = ""
                cleared += 1

        cur2 = str(ws.cell(r, ic["Employees"]).value or "").strip()
        if emptyish(cur2) or name in HAND:
            # always apply HAND; empty need fill
            if name in HAND or emptyish(cur2):
                targets.append((r, name, web))

    print(f"normalize/clear done; targets={len(targets)} cleared={cleared} normalized={normalized}")

    # Deduplicate work by company
    results: dict[str, tuple[str, str, str]] = {}
    unique = {}
    for r, name, web in targets:
        unique[name] = web

    def _job(item):
        name, web = item
        return name, process_one(name, web, client, model, cache)

    done = 0
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = [ex.submit(_job, item) for item in unique.items()]
        for fut in as_completed(futs):
            name, triple = fut.result()
            results[name] = triple
            done += 1
            if done % 25 == 0:
                print(f"  progress {done}/{len(unique)}")
                save_cache(cache)

    save_cache(cache)

    filled = hand = site = ds = none = 0
    audit_rows = []
    for r, name, web in targets:
        rng, source, note = results.get(name, ("", "none", ""))
        if name in HAND:
            rng, source, note = HAND[name], "hand", "known public band"
        if rng:
            ws.cell(r, ic["Employees"]).value = rng
            filled += 1
            if source == "hand":
                hand += 1
            elif source == "site_regex":
                site += 1
            elif source == "deepseek":
                ds += 1
        else:
            ws.cell(r, ic["Employees"]).value = NPD
            none += 1
        audit_rows.append(
            {"company": name, "website": web, "employees": rng or NPD, "source": source, "note": note}
        )

    wb.save(PATH)
    wb.close()

    # Final stats
    wb = load_workbook(PATH, data_only=True)
    ws = wb.active
    ei = h.index("Employees")
    named = blank = 0
    dist = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = str(row[ei] or "").strip()
        if emptyish(v) or v == NPD:
            blank += 1
        else:
            named += 1
            dist[v] = dist.get(v, 0) + 1
    wb.close()

    summary = {
        "filled_attempted": filled,
        "hand": hand,
        "site_regex": site,
        "deepseek": ds,
        "npd": none,
        "final_named": named,
        "final_npd": blank,
        "distribution": dict(sorted(dist.items(), key=lambda x: -x[1])),
    }
    AUDIT.write_text(
        json.dumps({"summary": summary, "rows": audit_rows}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
