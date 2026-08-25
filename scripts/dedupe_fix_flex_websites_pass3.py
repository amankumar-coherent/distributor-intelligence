"""Pass3: merge Packaging Wholesalers into BOX Partners; final uniqueness check."""
from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "unique_website_cleanup_pass3_audit.json"


def main():
    backup = FP.with_name(
        FP.stem + f"_pre_unique_web_pass3_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    actions = []

    keep = []
    box = None
    for r in rows:
        name = (r[0] or "").strip()
        low = name.lower()
        if low == "packaging wholesalers":
            actions.append(
                {
                    "from": name,
                    "into": "BOX Partners",
                    "web": "https://www.boxpartners.com",
                    "reason": "acquired_by_BOX_Partners",
                }
            )
            continue
        if low == "box partners":
            box = r
            r[1] = "https://www.boxpartners.com"
        if low == "u.s. plastic corp.":
            r[1] = "https://usplastic.com"
        keep.append(r)

    if box is not None:
        s = str(box[20] or "")
        note = " Includes The Packaging Wholesalers (acquired by BOX Partners)."
        if "Packaging Wholesalers" not in s:
            box[20] = (s + note).strip()

    final = []
    seen = set()
    for r in keep:
        n = re.sub(r"\s+", " ", str(r[0] or "").strip())
        r[0] = n
        k = n.lower()
        if k in seen:
            actions.append({"dup_drop": n})
            continue
        seen.add(k)
        if len(r) > 12:
            r[12] = "Specialty Distributor"
        final.append(r)

    final.sort(key=lambda x: str(x[0]).lower())

    out = Workbook()
    o = out.active
    o.title = "Channel Landscape"
    o.append(header)
    for c in o[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in final:
        row = list(r) + [None] * (len(header) - len(r))
        o.append(row[: len(header)])
    for row in o.iter_rows(min_row=2, max_row=o.max_row, max_col=len(header)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in {1: 34, 2: 38, 6: 36, 13: 20, 21: 55}.items():
        o.column_dimensions[get_column_letter(i)].width = w
    o.freeze_panes = "A2"
    out.save(FP)

    names = [(r[0] or "") for r in final]
    webs = [(r[1] or "") for r in final]
    unique_names = len({n.lower() for n in names})
    doms = []
    for w in webs:
        try:
            d = urlparse(w if "://" in w else "https://" + w).netloc.lower().replace("www.", "")
        except Exception:
            d = w
        doms.append(d)
    dup_domains = [(k, v) for k, v in Counter(doms).items() if v > 1]
    suffix = [
        n
        for n in names
        if re.search(r"(Inc\.?|LLC|Ltd\.?|Corp\.?|Corporation|Company|Co\.)$", n, re.I)
    ]

    audit = {
        "ts": datetime.now().isoformat(),
        "backup": str(backup),
        "before": len(rows),
        "after": len(final),
        "unique_names": unique_names,
        "dup_domains": dup_domains,
        "remaining_suffix_names": suffix,
        "actions": actions,
        "file": str(FP),
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print("backup", backup)
    print("before", len(rows), "after", len(final))
    print("unique_names", unique_names, "of", len(final))
    print("dup_domains", dup_domains)
    print("remaining_suffix_names", suffix)
    print("actions", actions)


if __name__ == "__main__":
    main()
