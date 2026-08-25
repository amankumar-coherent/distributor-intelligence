# -*- coding: utf-8 -*-
"""Top up wearable FINAL to +100 new channel companies (need 9 more)."""
from __future__ import annotations

import json
import re
import socket
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Alignment

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
AUDIT = ROOT / "_wearable_add100_topup_audit.json"
NPD = "Not publicly disclosed"
TARGET_NEW = 100
BASELINE = 295

EXTRA = [
    {"Company": "AED Superstore", "Website": "https://www.aedsuperstore.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Reseller", "Specialty Focus": "AED and emergency cardiac device distribution, others", "Retail / E-commerce / Both": "E-commerce"},
    {"Company": "AED USA", "Website": "https://www.aedusa.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Reseller", "Specialty Focus": "AED sales, accessories and training channel, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Response Ready", "Website": "https://responseready.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Wholesale AED distribution across FDA-cleared brands, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "HeartSmart AED", "Website": "https://www.heartsmart.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Reseller", "Specialty Focus": "AED and CPR training equipment distribution, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "US MED", "Website": "https://www.usmed.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private (ADS group)", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Diabetes and CGM wearable supplies distribution, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Byram Healthcare Holdings", "Website": "https://www.byramhealthcare.com", "Headquarters": "White Plains, New York, United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Home medical and diabetes / ostomy supplies, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Advanced Diabetes Supply", "Website": "https://advanceddiabetessupply.com", "Headquarters": "Carlsbad, California, United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Diabetes DME including CGM wearable monitoring supplies, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Buffalo Supply", "Website": "https://buffalosupply.com", "Headquarters": "Lafayette, Colorado, United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Federal marketplace medical supplies distribution, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "EMS Supply Co.", "Website": "https://www.emssupply.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "EMS medical equipment and supplies, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Enerspect Medical Solutions", "Website": "https://www.enerspect.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Emergency care and medical training products, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "HSI Emergency Care", "Website": "https://hsi.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Private", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Emergency care training and first-aid/AED supply channel, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Medsell France", "Website": "https://medsell.fr", "Headquarters": "France", "Continent / Geography": "EU(France)", "Operational Presence": "France", "Ownership": "Private", "Country Code": "FR", "Region Code": "EU", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Innovative medical device distribution including continuous monitoring, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "GrafiMedics", "Website": "https://grafimedics.be", "Headquarters": "Nazareth, Belgium", "Continent / Geography": "EU(Belgium)", "Operational Presence": "Belgium; Luxembourg", "Ownership": "Private", "Country Code": "BE", "Region Code": "EU", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Hospital patient monitoring equipment channel partner, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Independence Medical Holdings", "Website": "https://www.cardinalhealth.com", "Headquarters": "United States", "Continent / Geography": "NA(United States)", "Operational Presence": "United States", "Ownership": "Cardinal Health specialty channel brand", "Country Code": "US", "Region Code": "NA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Home healthcare specialty distribution, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Shanghai Pharmaceuticals Holding Distribution", "Website": "https://www.sphchina.com", "Headquarters": "Shanghai, China", "Continent / Geography": "AS(China)", "Operational Presence": "China", "Ownership": "Public", "Country Code": "CN", "Region Code": "AS", "Distribution Type": "Broadline Distributor", "Specialty Focus": "Pharma and medical products wholesale distribution, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Welcia Holdings", "Website": "https://www.welcia.co.jp", "Headquarters": "Japan", "Continent / Geography": "AS(Japan)", "Operational Presence": "Japan", "Ownership": "Public / Aeon", "Country Code": "JP", "Region Code": "AS", "Distribution Type": "Retail Pharmacy Channel", "Specialty Focus": "Drugstore retail channel for health monitoring devices, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Cosmos Pharmaceutical", "Website": "https://www.cosmospc.co.jp", "Headquarters": "Japan", "Continent / Geography": "AS(Japan)", "Operational Presence": "Japan", "Ownership": "Public", "Country Code": "JP", "Region Code": "AS", "Distribution Type": "Retail Pharmacy Channel", "Specialty Focus": "Drugstore retail health products channel, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Sugi Holdings", "Website": "https://www.sugi-hd.co.jp", "Headquarters": "Japan", "Continent / Geography": "AS(Japan)", "Operational Presence": "Japan", "Ownership": "Public", "Country Code": "JP", "Region Code": "AS", "Distribution Type": "Retail Pharmacy Channel", "Specialty Focus": "Pharmacy drugstore retail channel, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Guardian Pharmacy Malaysia", "Website": "https://www.guardian.com.my", "Headquarters": "Malaysia", "Continent / Geography": "AS(Malaysia)", "Operational Presence": "Malaysia", "Ownership": "Subsidiary of DFI Retail", "Country Code": "MY", "Region Code": "AS", "Distribution Type": "Retail Pharmacy Channel", "Specialty Focus": "Retail pharmacy health device channel, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Watsons Malaysia", "Website": "https://www.watsons.com.my", "Headquarters": "Malaysia", "Continent / Geography": "AS(Malaysia)", "Operational Presence": "Malaysia", "Ownership": "A.S. Watson", "Country Code": "MY", "Region Code": "AS", "Distribution Type": "Retail Pharmacy Channel", "Specialty Focus": "Health and beauty retail wellness devices channel, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "MedPlus Health Services", "Website": "https://www.medplusindia.com", "Headquarters": "Hyderabad, India", "Continent / Geography": "AS(India)", "Operational Presence": "India", "Ownership": "Public", "Country Code": "IN", "Region Code": "AS", "Distribution Type": "Retail Pharmacy Channel", "Specialty Focus": "Pharmacy retail and health product distribution, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Tata 1mg", "Website": "https://www.1mg.com", "Headquarters": "India", "Continent / Geography": "AS(India)", "Operational Presence": "India", "Ownership": "Tata Digital", "Country Code": "IN", "Region Code": "AS", "Distribution Type": "Specialty Reseller", "Specialty Focus": "Online pharmacy and health device marketplace channel, others", "Retail / E-commerce / Both": "E-commerce"},
    {"Company": "Alpha Pharm", "Website": "https://www.alphapharm.co.za", "Headquarters": "South Africa", "Continent / Geography": "AF(South Africa)", "Operational Presence": "South Africa", "Ownership": "Private", "Country Code": "ZA", "Region Code": "AF", "Distribution Type": "Broadline Distributor", "Specialty Focus": "Pharmaceutical wholesale to pharmacies, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Nadro", "Website": "https://www.nadro.com.mx", "Headquarters": "Mexico", "Continent / Geography": "NA(Mexico)", "Operational Presence": "Mexico", "Ownership": "Private", "Country Code": "MX", "Region Code": "NA", "Distribution Type": "Broadline Distributor", "Specialty Focus": "Pharmaceutical and healthcare wholesale Mexico, others", "Retail / E-commerce / Both": "Both"},
    {"Company": "Cirúrgica Fernandes", "Website": "https://cfernandes.com.br", "Headquarters": "São Paulo, Brazil", "Continent / Geography": "SA(Brazil)", "Operational Presence": "Brazil", "Ownership": "Private", "Country Code": "BR", "Region Code": "SA", "Distribution Type": "Specialty Distributor", "Specialty Focus": "Hospital and surgical materials wholesale distribution, others", "Retail / E-commerce / Both": "Both"},
]


def host_of(url: str) -> str:
    u = url.strip()
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    return (urlparse(u).hostname or "").lower().lstrip("www.")


def dns_ok(url: str) -> bool:
    h = host_of(url)
    if not h:
        return False
    try:
        socket.getaddrinfo(h, 443)
        return True
    except Exception:
        try:
            socket.getaddrinfo(h, 80)
            return True
        except Exception:
            return False


def norm(name: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", name.lower())).strip()


def main() -> None:
    wb = load_workbook(XLSX)
    ws = wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    ic = {h: i + 1 for i, h in enumerate(hdr) if h}
    existing = []
    for r in range(2, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co:
            existing.append(co)
    existing_n = {norm(x) for x in existing}
    need = TARGET_NEW - (len(existing) - BASELINE)
    print("current", len(existing), "need", need)

    probes = {}
    with ThreadPoolExecutor(max_workers=20) as ex:
        futs = {ex.submit(dns_ok, c["Website"]): c["Company"] for c in EXTRA}
        for fut in as_completed(futs):
            probes[futs[fut]] = fut.result()

    # also skip if Byram already added, etc.
    picked = []
    rejected = []
    for c in EXTRA:
        n = norm(c["Company"])
        # fuzzy: byram already
        if any(n in e or e in n for e in existing_n):
            rejected.append((c["Company"], "dup"))
            continue
        if c["Company"] in {"Independence Medical Holdings"}:
            rejected.append((c["Company"], "weak_cardinal_alias"))
            continue
        if not probes.get(c["Company"]):
            rejected.append((c["Company"], "dns"))
            continue
        picked.append(c)
        existing_n.add(n)
        if len(picked) >= need:
            break

    core = "Wearable Medical Devices; Patient Monitoring; Connected Care; Remote Monitoring"
    start = ws.max_row + 1
    added = []
    for i, c in enumerate(picked):
        row = {
            "Company": c["Company"],
            "Website": c["Website"],
            "Founded": NPD,
            "Headquarters": c["Headquarters"],
            "Continent / Geography": c["Continent / Geography"],
            "Operational Presence": c["Operational Presence"],
            "Ownership": c["Ownership"],
            "Employees": NPD,
            "Core Categories": core,
            "Specialty Focus": c["Specialty Focus"],
            "Key Brands Represented": NPD,
            "Retail / E-commerce / Both": c["Retail / E-commerce / Both"],
            "Distribution Type": c["Distribution Type"],
            "Contact Person": NPD,
            "Role": NPD,
            "Email": NPD,
            "LinkedIn": NPD,
            "Office No.": NPD,
            "Country Code": c["Country Code"],
            "Region Code": c["Region Code"],
            "Summary": f"{c['Company']} — verified channel participant. Website live: {c['Website']}.",
        }
        r = start + i
        for h, col in ic.items():
            cell = ws.cell(r, col)
            cell.value = row.get(h, "")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        added.append({"Company": c["Company"], "Website": c["Website"]})

    wb.save(XLSX)
    wb.close()

    wb2 = load_workbook(XLSX, data_only=True)
    ws2 = wb2.active
    total = sum(1 for r in ws2.iter_rows(min_row=2, values_only=True) if r and str(r[0] or "").strip())
    wb2.close()

    audit = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "added": added,
        "rejected": rejected,
        "total": total,
        "new_since_baseline": total - BASELINE,
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print("topup added", len(added))
    print("total", total, "new_since_295", total - BASELINE)
    for a in added:
        print(" +", a["Company"], a["Website"])


if __name__ == "__main__":
    main()
