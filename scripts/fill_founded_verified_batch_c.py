"""Batch C: additional verified Founded years only."""
from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

FP = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
AUDIT = FP.parent / "founded_fill_audit_batch_c.json"

VERIFIED: dict[str, tuple[str, str]] = {
    "Walki": ("1930", "walki.com who-we-are / Paperituote Valkeakoski 1930"),
    "Danaflex": ("2001", "danaflex-group.com established 2001"),
    "Davpack": ("1968", "davpack.co.uk About; Davenport Paper Company Ltd 1968"),
    "Nashville Wraps": ("1976", "TNECD / LeadIQ / company history Robby Meadows 1976"),
    "Macfarlane Group": ("1949", "macfarlanegroup.com company history"),
    "Allpack Packaging": ("1989", "Allpack Packaging Supplies Ltd incorporation / Macfarlane acquisition notes"),
    "Amazon Business Packaging": ("2015", "Amazon Business launch 2015"),
    "MSC Industrial Supply Co.": ("1941", "Wikipedia MSC"),
    "Global Industrial": ("1949", "Global Industrial history"),
    "Paper Mart": ("1921", "papermart.com about"),
    "Cormack Packaging": ("1957", "Cormack AU history"),
    "Cospak Pty Ltd.": ("1975", "Cospak history"),
    "Detpak": ("1945", "Detpak history"),
    "Plastpack Australia": ("1928", "Plastpack 1928"),
    "Springpack": ("1994", "springpack history"),
    "Cleveland Containers": ("1984", "clevelandcontainers history"),
    "Polybags Limited": ("1986", "polybags.co.uk"),
    "Swiftpak": ("1984", "swiftpak history"),
    "The Cary Company": ("1895", "thecarycompany"),
    "Royal Bag": ("1976", "royalbag — only if confirmed; else skip"),
    "TapeCase": ("1969", "tapecase — only if confirmed; else skip"),
    "Wrap-Tite": ("1986", "wraptite history"),
    "Schermerhorn Bros. Co.": ("1893", "schermerhorn history"),
    "Richards Packaging": ("1912", "richards history"),
    "Crawford Packaging": ("1953", "crawford history"),
    "Freund Container & Supply": ("1938", "freund history"),
    "Dayton Bag & Burlap": ("1910", "dayton bag history"),
    "International Plastics": ("1968", "interplas history"),
    "Interplas": ("1968", "interplas"),
    "Ratioform": ("1979", "ratioform 1979"),
    "Ratioform Italy": ("1979", "ratioform"),
    "Network Packaging": ("1988", "network packaging"),
    "BradyPLUS": ("SKIP", "BradyPLUS is roll-up brand; Brady Corp 1914 is different entity"),
    "Signet AU": ("1968", "Signet AU"),
    "Waltons ZA": ("1942", "Waltons SA"),
    "Makro ZA": ("1971", "Makro SA"),
    "Sodimac Empresas": ("1952", "Sodimac"),
    "Falabella Empresas Chile": ("1889", "Falabella"),
}

# Only keep ones with strong confirmation from this session
KEEP = {
    "Walki",
    "Danaflex",
    "Davpack",
    "Nashville Wraps",
    "Macfarlane Group",
    "Amazon Business Packaging",
}


def main() -> None:
    clean = {
        k: v
        for k, v in VERIFIED.items()
        if k in KEEP and re.fullmatch(r"(18|19|20)\d{2}", v[0])
    }

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = FP.with_name(FP.stem + f"_backup_before_founded_c_{stamp}.xlsx")
    shutil.copy2(FP, backup)

    wb = load_workbook(FP)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = {str(h): i + 1 for i, h in enumerate(headers) if h}

    filled = []
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, col["Company"]).value or "").strip()
        cur = ws.cell(r, col["Founded"]).value
        cur_s = "" if cur is None else str(cur).strip()
        if name not in clean:
            continue
        if cur_s and cur_s not in ("None", "N/A", "Not publicly disclosed", "-"):
            continue
        year, src = clean[name]
        ws.cell(r, col["Founded"]).value = year
        filled.append({"company": name, "founded": year, "source": src})

    wb.save(FP)
    empty = has = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col["Founded"]).value
        if v is None or str(v).strip() in ("", "None", "N/A", "Not publicly disclosed", "-"):
            empty += 1
        else:
            has += 1

    audit = {
        "backup": str(backup),
        "filled_now": filled,
        "filled_count": len(filled),
        "founded_filled_total": has,
        "founded_empty_total": empty,
        "total_rows": ws.max_row - 1,
        "policy": "Left blank where no reliable public founding year found (no hallucination)",
    }
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(f"FILLED_NOW={len(filled)} TOTAL_FILLED={has} STILL_EMPTY={empty}")
    for x in filled:
        print(f"  {x['company']}: {x['founded']}")


if __name__ == "__main__":
    main()
