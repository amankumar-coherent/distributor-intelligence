"""Merge branch-style duplicate companies into unique parents across 3 markets.

Examples: Ingram Micro Brazil / Ingram Micro US → Ingram Micro
Same pattern for ALSO, TD SYNNEX, Exclusive Networks, etc.
"""
from __future__ import annotations

import re
import shutil
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

MARKETS = [
    {
        "name": "firewall",
        "paths": [
            Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL.xlsx"),
            Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market_FINAL_320plus.xlsx"),
        ],
        "sheet": "Firewall Companies",
    },
    {
        "name": "smartwatch",
        "paths": [
            Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL.xlsx"),
            Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market_FINAL_320plus.xlsx"),
        ],
        "sheet": "Smartwatch Companies",
    },
    {
        "name": "green_chemical",
        "paths": [
            Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL.xlsx"),
            Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL_320plus.xlsx"),
        ],
        "sheet": "Green Chemical Companies",
    },
]

# Explicit family rules: (canonical_name, matcher_fn returning bool on norm key)
# Keep conservative — only clear same-company regional branches.


def nk(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def starts(*prefixes):
    def fn(key: str) -> bool:
        return any(key == p or key.startswith(p + " ") for p in prefixes)

    return fn


def equals_or_starts(exact: set[str], prefixes: tuple[str, ...]):
    def fn(key: str) -> bool:
        if key in exact:
            return True
        return any(key.startswith(p + " ") for p in prefixes)

    return fn


FAMILIES: list[tuple[str, callable, str]] = [
    # (canonical, match_fn, parent_key preferred)
    ("Ingram Micro", starts("ingram micro"), "ingram micro"),
    (
        "ALSO Holding",
        lambda k: k == "also" or k == "also holding" or k.startswith("also "),
        "also holding",
    ),
    (
        "TD SYNNEX",
        # Only TD SYNNEX / Tech Data — NOT Synnex Technology International (Taiwan)
        lambda k: k == "td synnex" or k.startswith("td synnex ") or k.startswith("tech data"),
        "td synnex",
    ),
    (
        "Synnex Technology International",
        lambda k: (
            k == "synnex technology international"
            or k.startswith("synnex technology international ")
            or (k.startswith("synnex ") and "td synnex" not in k and not k.startswith("td "))
            or k == "synnex"
        ),
        "synnex technology international",
    ),
    (
        "Exclusive Networks",
        starts("exclusive networks"),
        "exclusive networks",
    ),
    (
        "Westcon-Comstor",
        lambda k: k.startswith("westcon") or "comstor" in k,
        "westcon comstor",
    ),
    ("Pax8", starts("pax8"), "pax8"),
    ("Redington", starts("redington"), "redington"),
    ("CDW", starts("cdw"), "cdw"),
    ("Infinigate", starts("infinigate"), "infinigate"),
    ("Logicalis", starts("logicalis"), "logicalis"),
    (
        "SiS Distribution",
        lambda k: k in {"sis", "sis distribution"}
        or k.startswith("sis distribution ")
        or k.startswith("sis distribution"),
        "sis distribution",
    ),
    ("Arrow ECS", starts("arrow ecs"), "arrow ecs"),
    ("Computacenter", starts("computacenter"), "computacenter"),
    ("Exertis", starts("exertis"), "exertis"),
    ("Logicom", starts("logicom"), "logicom"),
    (
        "StarLink",
        lambda k: k in {"starlink", "star link"} or k.startswith("starlink ") or k.startswith("star link "),
        "starlink",
    ),
    ("SHI International", lambda k: k == "shi" or k.startswith("shi ") or k.startswith("shi international"), "shi international"),
    ("Euronics", starts("euronics"), "euronics"),
    ("Best Buy", starts("best buy"), "best buy"),
    ("IMCD", lambda k: k == "imcd" or k.startswith("imcd "), "imcd"),
]


def branch_label(company: str, canonical: str) -> str:
    """Derive a short region/branch label from company name."""
    raw = company.strip()
    # strip canonical prefix loosely
    pat = re.compile(re.escape(canonical), re.I)
    rest = pat.sub("", raw, count=1).strip(" -–—|/()")
    if not rest or nk(rest) in {"holding", "group", "international", "corp", "corporation", "inc", "ltd"}:
        # try remove first two words
        parts = raw.split()
        if len(parts) > 2:
            rest = " ".join(parts[2:]).strip(" -–—|/")
        elif len(parts) > 1:
            rest = parts[-1]
    rest = rest.strip()
    return rest or raw


def find_header(ws) -> tuple[int, list]:
    for r in range(1, 6):
        vals = [c.value for c in ws[r]]
        if vals and str(vals[0] or "").strip() == "Company":
            return r, vals
        if any(str(v or "").strip() == "Company" for v in vals):
            return r, [str(v or "").strip() if v is not None else "" for v in vals]
    raise RuntimeError("Company header not found")


def score_row(ws, r: int, cols: dict) -> int:
    score = 0
    for f in (
        "Founded",
        "Headquarters",
        "Ownership",
        "Employees",
        "Turnover / Scale",
        "Contact Person",
        "Website",
        "Summary",
        "LinkedIn",
    ):
        if f in cols:
            v = str(ws.cell(r, cols[f]).value or "").strip().lower()
            if v and v not in {"n/a", "na", "not publicly disclosed", "-", "unknown"}:
                score += 1
    return score


def merge_family_in_sheet(ws, hdr_row: int, cols: dict, canonical: str, match_fn, parent_key: str) -> dict | None:
    members: list[tuple[int, str, str]] = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        company = str(ws.cell(r, cols["Company"]).value or "").strip()
        if not company:
            continue
        key = nk(company)
        if match_fn(key):
            members.append((r, company, key))

    if len(members) <= 1:
        return None

    # prefer exact parent_key, else highest score, else first
    parent = None
    for r, company, key in members:
        if key == parent_key:
            parent = (r, company, key)
            break
    if parent is None:
        parent = max(members, key=lambda m: score_row(ws, m[0], cols))

    parent_row = parent[0]
    branch_rows = [r for r, _, _ in members if r != parent_row]

    # aggregate cities
    cities: list[str] = []
    seen: set[str] = set()
    for r, company, key in members:
        label = branch_label(company, canonical) if key != parent_key else None
        existing = str(ws.cell(r, cols.get("Cities / Regions", 0)).value or "").strip() if "Cities / Regions" in cols else ""
        for cand in (label, existing):
            if not cand:
                continue
            low = cand.lower()
            if low in seen or low in {"n/a", "not publicly disclosed", "global"}:
                continue
            if len(cand) > 80:
                continue
            cities.append(cand)
            seen.add(low)

    # set parent
    ws.cell(parent_row, cols["Company"]).value = canonical
    if "Cities / Regions" in cols and cities:
        ws.cell(parent_row, cols["Cities / Regions"]).value = "; ".join(cities)
    if "Regions Served" in cols:
        # mark multi-region if many branches
        if len(members) >= 3:
            cur = str(ws.cell(parent_row, cols["Regions Served"]).value or "").strip()
            if not cur or cur.lower() in {"regional", "n/a", "not publicly disclosed"}:
                ws.cell(parent_row, cols["Regions Served"]).value = "Global / multi-region (consolidated branches)"
    if "Region Segment" in cols and len(members) >= 3:
        ws.cell(parent_row, cols["Region Segment"]).value = "Global"

    return {
        "canonical": canonical,
        "kept_row": parent_row,
        "before_names": [c for _, c, _ in members],
        "delete_rows": branch_rows,
        "cities": cities,
    }


def clean_region_sheets(wb, deleted_names: set[str], canonicals: set[str]) -> int:
    removed = 0
    for sheet_name in list(wb.sheetnames):
        if not sheet_name.startswith("Region -"):
            continue
        rws = wb[sheet_name]
        try:
            hdr_row, hdr = find_header(rws)
        except RuntimeError:
            continue
        cols = {h: i + 1 for i, h in enumerate(hdr) if h}
        if "Company" not in cols:
            continue
        to_del = []
        for r in range(hdr_row + 1, rws.max_row + 1):
            co = str(rws.cell(r, cols["Company"]).value or "").strip()
            if co in deleted_names:
                to_del.append(r)
        for r in sorted(to_del, reverse=True):
            rws.delete_rows(r, 1)
            removed += 1
    return removed


def count_companies(ws, hdr_row: int, cols: dict) -> int:
    n = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        if str(ws.cell(r, cols["Company"]).value or "").strip():
            n += 1
    return n


def process_file(path: Path, sheet_name: str) -> dict:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_pre_unique_merge_{stamp}{path.suffix}")
    shutil.copy2(path, bak)

    wb = openpyxl.load_workbook(path)
    if sheet_name not in wb.sheetnames:
        # fallback
        sheet_name = next(s for s in wb.sheetnames if "Compan" in s or "Distributor" in s)
    ws = wb[sheet_name]
    hdr_row, hdr = find_header(ws)
    cols = {str(h): i + 1 for i, h in enumerate(hdr) if h}

    before = count_companies(ws, hdr_row, cols)
    merges = []
    all_delete_rows: list[int] = []
    deleted_names: set[str] = set()
    canonicals: set[str] = set()

    for canonical, match_fn, parent_key in FAMILIES:
        result = merge_family_in_sheet(ws, hdr_row, cols, canonical, match_fn, parent_key)
        if not result:
            continue
        merges.append(result)
        canonicals.add(canonical)
        for name in result["before_names"]:
            if name != canonical:
                deleted_names.add(name)
        all_delete_rows.extend(result["delete_rows"])

    # delete branch rows bottom-up
    for r in sorted(set(all_delete_rows), reverse=True):
        ws.delete_rows(r, 1)

    region_removed = clean_region_sheets(wb, deleted_names, canonicals)

    # refresh counts after deletes — header still same row
    after = count_companies(ws, hdr_row, cols)
    wb.save(path)
    return {
        "file": path.name,
        "backup": bak.name,
        "before": before,
        "after": after,
        "removed": before - after,
        "families_merged": len(merges),
        "merges": merges,
        "region_rows_removed": region_removed,
    }


def main() -> None:
    for market in MARKETS:
        print(f"\n======== {market['name'].upper()} ========")
        for path in market["paths"]:
            if not path.exists():
                print("MISSING", path)
                continue
            rep = process_file(path, market["sheet"])
            print(f"{rep['file']}: {rep['before']} -> {rep['after']} (removed {rep['removed']}) families={rep['families_merged']} region_cleaned={rep['region_rows_removed']}")
            for m in rep["merges"]:
                print(f"  {m['canonical']}: {len(m['before_names'])} branches -> 1")
                print(f"    from: {', '.join(m['before_names'][:8])}{'...' if len(m['before_names'])>8 else ''}")


if __name__ == "__main__":
    main()
