"""Fill smartwatch FINAL.xlsx from founded_and_hq_v2_cleaned.xlsx (match by company).

Also sync JSON + Neon so Distributor Intelligence matches the sheet.
"""
from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

CMI = Path(r"C:\Users\aman.kumar\Desktop\Cmi\agents\vendor-intel\output\run4\global_smartwatch_market_global")
COH = Path(r"C:\Users\aman.kumar\Desktop\coherent\output\run4\global_smartwatch_market_global")

FINAL_NAME = "global_smartwatch_market_FINAL.xlsx"
CLEAN_NAME = "global_smartwatch_market_founded_and_hq_v2_cleaned.xlsx"

# Excel header → crawl field
COL_TO_CRAWL = {
    "Founded": "founded_year",
    "Headquarters": "headquarters",
    "Cities / Regions": "cities_regions",
    "Ownership": "ownership_parent",
    "Business Type": None,  # top-level role
    "Employees": "employee_count",
    "Turnover / Scale": "revenue",
    "Contact Person": "contact_person",
    "Role": "contact_role",
    "Email": "email",
    "Phone / WhatsApp": "phone",
    "LinkedIn": "linkedin",
    "Website": "website",
    "Core Categories": "core_categories",
    "Specialty Focus": "specialty_focus",
    "Price Segment": "price_segment",
    "Key Brands Represented": "key_brands_represented",
    "Retail Chains": "retail_chains",
    "Specialty Grocers": "specialty_grocers",
    "Foodservice / HoReCa": "foodservice",
    "E-commerce": "ecommerce",
    "Distribution Type": "distribution_type",
    "Regions Served": "regions_served",
    "Regional Extensions": "regional_extensions",
    "Warehouse / Logistics": "warehouse_logistics",
    "Delivery / Storage": "delivery_storage",
    "Summary": "summary",
}

# Prefer cleaned whenever it has a value (cleaned is the curated source of truth for these).
PREFER_CLEANED = {
    "Cities / Regions",
    "Core Categories",
    "Specialty Focus",
    "Price Segment",
    "Key Brands Represented",
    "Retail Chains",
    "Specialty Grocers",
    "Foodservice / HoReCa",
    "E-commerce",
    "Distribution Type",
    "Regions Served",
    "Regional Extensions",
    "Warehouse / Logistics",
    "Delivery / Storage",
    "Headquarters",
    "Founded",
    "Ownership",
    "Business Type",
    "Summary",
}


def emptyish(v) -> bool:
    if v is None:
        return True
    s = str(v).strip()
    return s in {"", "—", "-", "None", "N/A", "n/a", "null"}


def norm_name(v: str) -> str:
    return re.sub(r"\s+", " ", (v or "").strip().lower())


def load_rows(path: Path) -> tuple[list, dict[str, list], list[str]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    rows: dict[str, list] = {}
    order: list[str] = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        name = str(row[0]).strip()
        rows[norm_name(name)] = list(row)
        order.append(name)
    wb.close()
    return headers, rows, order


def apply_cleaned_into_final(final_path: Path, clean_path: Path) -> dict:
    f_headers, f_rows, f_order = load_rows(final_path)
    c_headers, c_rows, _ = load_rows(clean_path)

    c_hmap = {str(h).strip(): i for i, h in enumerate(c_headers) if h}
    f_hmap = {str(h).strip(): i for i, h in enumerate(f_headers) if h}

    fwb = load_workbook(final_path)
    fws = fwb[fwb.sheetnames[0]]
    f_idx: dict[str, int] = {}
    for ridx, row in enumerate(fws.iter_rows(min_row=4, values_only=False), start=4):
        vals = [c.value for c in row]
        if not vals or vals[0] is None or str(vals[0]).strip() == "":
            continue
        f_idx[norm_name(str(vals[0]).strip())] = ridx

    filled = 0
    overwritten = 0
    missing = []

    for name in f_order:
        nk = norm_name(name)
        if nk not in c_rows:
            missing.append(name)
            continue
        crow = c_rows[nk]
        ridx = f_idx[nk]
        for h, fi in f_hmap.items():
            if h == "Company":
                continue
            ci = c_hmap.get(h)
            if ci is None:
                continue
            fv = f_rows[nk][fi] if fi < len(f_rows[nk]) else None
            cv = crow[ci] if ci < len(crow) else None
            if emptyish(cv):
                continue
            use = False
            if emptyish(fv):
                use = True
                filled += 1
            elif h in PREFER_CLEANED and str(fv).strip() != str(cv).strip():
                use = True
                overwritten += 1
            if use:
                fws.cell(ridx, fi + 1).value = cv

    fwb.save(final_path)
    fwb.close()
    return {
        "filled_empty": filled,
        "overwritten": overwritten,
        "missing_in_cleaned": missing,
        "companies": len(f_order),
    }


def sync_json_from_cleaned(market_dir: Path, clean_path: Path) -> Path:
    """Update main JSON + neon filtered from cleaned Excel values."""
    headers, rows, order = load_rows(clean_path)
    hmap = {str(h).strip(): i for i, h in enumerate(headers) if h}
    cleaned_by_name: dict[str, dict] = {}
    for name in order:
        nk = norm_name(name)
        row = rows[nk]
        cleaned_by_name[nk] = {
            h: row[i] for h, i in hmap.items() if i < len(row)
        }

    json_path = market_dir / "global_smartwatch_market_global.json"
    data = json.loads(json_path.read_text(encoding="utf-8"))
    updated = 0
    for row in data.get("relevant_companies") or []:
        name = norm_name(str(row.get("company") or row.get("brand") or ""))
        src = cleaned_by_name.get(name)
        if not src:
            continue
        c = row.setdefault("_crawl", {})
        for excel_h, crawl_k in COL_TO_CRAWL.items():
            val = src.get(excel_h)
            if emptyish(val):
                continue
            sval = str(val).strip()
            if excel_h == "Business Type":
                row["role"] = sval
                c["role"] = sval
                continue
            if crawl_k is None:
                continue
            c[crawl_k] = sval
            if crawl_k == "summary":
                row["summary"] = sval
                row["company_summary"] = sval
            if crawl_k == "website":
                row["website"] = sval
        updated += 1

    json_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    # rebuild neon filtered
    from export_founded_hq_v2 import has_founded, has_headquarters

    rows_all = [r for r in (data.get("relevant_companies") or []) if r.get("is_relevant", True)]
    try:
        from vendor_intel.funnel.channel_partner_types import (
            channel_partner_mode_enabled,
            is_production_exportable,
        )

        channel = channel_partner_mode_enabled(data.get("query_context") or {})
    except Exception:
        channel = False
        is_production_exportable = None  # type: ignore
    keep = "is_analyst_seed" if channel else "is_seed"
    filtered = [r for r in rows_all if (has_founded(r) and has_headquarters(r)) or r.get(keep)]
    if channel and is_production_exportable:
        filtered = [r for r in filtered if is_production_exportable(r)]
    ctx = data.get("query_context") or {}
    market = str(ctx.get("industry") or data.get("query") or "Global Smartwatch Market")
    neon = {
        "query": market,
        "query_context": {"industry": market, "country": str(ctx.get("country") or "global")},
        "relevant_companies": filtered,
        "completed_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "filter_note": "Founded + HQ FINAL set synced from cleaned Excel",
    }
    neon_path = market_dir / "_neon_filtered_export.json"
    neon_path.write_text(json.dumps(neon, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"JSON companies updated from cleaned: {updated}; neon rows={len(filtered)}")
    return neon_path


def main() -> None:
    push = "--push" in sys.argv
    # Prefer user-specified CMI paths; also mirror to coherent
    clean = CMI / CLEAN_NAME
    final = CMI / FINAL_NAME
    if not clean.exists() or not final.exists():
        raise SystemExit(f"Missing files under {CMI}")

    stats = apply_cleaned_into_final(final, clean)
    print("CMI FINAL:", stats)

    # Mirror FINAL + sync coherent if present
    if COH.exists():
        coh_final = COH / FINAL_NAME
        coh_clean = COH / CLEAN_NAME if (COH / CLEAN_NAME).exists() else clean
        try:
            # copy filled FINAL bytes to coherent after writing CMI
            coh_final.write_bytes(final.read_bytes())
            print(f"Mirrored FINAL -> {coh_final}")
        except Exception as exc:
            # try fill coherent in place
            if coh_final.exists() and coh_clean.exists():
                print("COH FINAL:", apply_cleaned_into_final(coh_final, coh_clean))
            else:
                print("coherent mirror failed:", exc)

        neon_path = sync_json_from_cleaned(COH, clean)
        # also mirror json to CMI
        try:
            (CMI / "global_smartwatch_market_global.json").write_bytes(
                (COH / "global_smartwatch_market_global.json").read_bytes()
            )
            (CMI / "_neon_filtered_export.json").write_bytes(neon_path.read_bytes())
        except Exception as exc:
            print("CMI json mirror:", exc)
    else:
        neon_path = sync_json_from_cleaned(CMI, clean)

    # verify empties
    wb = load_workbook(final, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r and r[0]]
    print("Post-fill empty counts:")
    for i, h in enumerate(headers):
        if not h or h not in PREFER_CLEANED:
            continue
        n = sum(1 for r in rows if emptyish(r[i] if i < len(r) else None))
        print(f"  {n:3d} empty | {h}")
    wb.close()

    if push:
        from vendor_intel.neon_push import push_json_to_neon

        print(json.dumps(push_json_to_neon(neon_path, merge_companies=False), indent=2, default=str))


if __name__ == "__main__":
    main()
