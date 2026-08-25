# -*- coding: utf-8 -*-
"""Fill verified person LinkedIn URLs where Contact Person is already named."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}
FILES = {
    "electrical": Path(
        r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global"
        r"\global_electrical_distribution_market_global_FINAL.xlsx"
    ),
    "avocado": Path(
        r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
        r"\global_avocado_oil_market.xlsx"
    ),
    "smartphone": Path(
        r"D:\vendor-intel\output\web_expand\global_smartphone_market_global"
        r"\global_smartphone_market.xlsx"
    ),
    "firewall": Path(
        r"D:\vendor-intel\output\run4\global_firewall_market_global"
        r"\global_firewall_market.xlsx"
    ),
}

# company -> (person substring to confirm, url)
URLS = {
    "electrical": {
        "Graybar Electric": ("Ipson", "https://www.linkedin.com/in/andy-ipson-379692216"),
        "Graybar Canada": ("Horncastle", "https://www.linkedin.com/in/peter-horncastle-58103511"),
        "HD Supply": ("Schnellenberger", "https://www.linkedin.com/in/sheila-schnellenberger-a948743"),
        "Lapp Group": ("Hiller", "https://www.linkedin.com/in/dr-christoph-hiller-8874b8b8"),
        "MSC Industrial Direct": ("Nadi", "https://www.linkedin.com/in/jahida-nadi-12516015"),
        "McNaughton-McKay Electric": ("Guidry", "https://www.linkedin.com/in/steven-guidry-7aa450137"),
        "McNaughton-McKay Electric Company": ("Guidry", "https://www.linkedin.com/in/steven-guidry-7aa450137"),
        "North Coast Electric": ("Cox", "https://www.linkedin.com/in/tommy-cox-827471262"),
        "Middy's": ("Carrington", "https://www.linkedin.com/in/mark-carrington-127401366"),
        "Middy's Electrical (NZ)": ("Carrington", "https://www.linkedin.com/in/mark-carrington-127401366"),
        "Jewson": ("Bentham", "https://www.linkedin.com/in/laurence-bentham-683a13101"),
        "City Electric Supply (CES)": ("Cooper", "https://www.linkedin.com/in/matthew-cooper-593a808"),
        "Fega & Schmitt Elektrogroßhandel": ("Schelkens", "https://www.linkedin.com/in/richard-schelkens-42948114b"),
        "Franklin Empire": ("Brunelle", "https://www.linkedin.com/in/bernard-brunelle-808a833b"),
        "Hillmann & Ploog": ("Pree", "https://www.linkedin.com/in/simon-pree-514b14248"),
        "Ideal Supply": ("Van Kessel", "https://www.linkedin.com/in/stevevankessel"),
        "K/E Electric Supply": ("Kuchenmeister", "https://www.linkedin.com/in/randy-kuchenmeister-2768ba28"),
        "Motion Australia": ("Gray", "https://www.linkedin.com/in/grant-gray-2352a531"),
        "Rexel Austria": ("Jozic", "https://www.linkedin.com/in/stjepan-jozic-057709162"),
        "Rexel Energy Solutions": ("Michaud", "https://www.linkedin.com/in/paul-michaud-b8a71b5"),
        "Rexel France": ("Lain", "https://www.linkedin.com/in/arnaud-lain%C3%A9-235181115"),
        "Rexel Germany": ("Hutterer", "https://www.linkedin.com/in/martin-hutterer-920b84171"),
        "Rexel Sweden": ("Fasth", "https://www.linkedin.com/in/morgan-fasth-b098662b"),
        "Rumsey Electric Company": ("Melvin", "https://www.linkedin.com/in/patrick-melvin-2670aa12"),
        "Sonepar France": ("Saint Germain", "https://www.linkedin.com/in/francois-saint-germain-b662a692"),
        "Sonepar New Zealand": ("Franklin", "https://www.linkedin.com/in/richard-franklin-a7b14886"),
        "RS Components (now RS Group)": ("Thavamani", "https://www.linkedin.com/in/sridher-thavamani-76ab28259"),
        "Sonepar": ("Choubey", "https://www.linkedin.com/in/bl-choubey-62119328"),
        "Scott Electrical": ("Stout", "https://www.linkedin.com/in/cris-stout-59702999"),
        "Sesco (Southeast Electrical Supply)": ("Barker", "https://www.linkedin.com/in/kevin-barker-54867515"),
        "Sonepar Sweden": ("Dittlau", "https://www.linkedin.com/in/rickard-dittlau-63710812"),
        "Summit Electric Supply": ("Holdsworth", "https://www.linkedin.com/in/steven-holdsworth-64a48aa"),
        "TLC Electrical Supplies": ("O'Shea", "https://www.linkedin.com/in/roboshea"),
        "D&W Electrical": ("Philippe", "https://www.linkedin.com/in/jason-philippe-476a2219b"),
        "Ideal Electrical": ("Fann", "https://www.linkedin.com/in/michaeldfann"),
        "Electrical Distributors (NZ)": ("Fann", "https://www.linkedin.com/in/michaeldfann"),
        "Farnell (Asia Pacific)": ("Chan", "https://www.linkedin.com/in/cy-chan-29848a86"),
    },
    "avocado": {
        "C&S Wholesale Grocers": ("Saia", "https://www.linkedin.com/in/jim-saia-26b2b023"),
        "Bidfood": ("Browne", "https://www.linkedin.com/in/tim-browne-a036295a"),
        "Associated Wholesale Grocers": ("Neumann", "https://www.linkedin.com/in/neumannjames"),
        "Batory Foods": ("O'Leary", "https://www.linkedin.com/in/erin-o-leary-b769a843"),
        "Ben E. Keith Foods": ("Werner", "https://www.linkedin.com/in/david-werner-5b122426b"),
        "Bozzuto's": ("Capanna", "https://www.linkedin.com/in/steve-capanna-22a17123"),
        "Brake Bros": ("Addinall", "https://www.linkedin.com/in/debbie-addinall-400b7411"),
        "Bidvest Australia": ("Hopman", "https://www.linkedin.com/in/luke-hopman-6810895a"),
        "Blakemore Wholesale": ("Senior", "https://www.linkedin.com/in/sara-senior-81b64395"),
        "Bestway Wholesale": ("Lewis", "https://www.linkedin.com/in/dan-lewis-86426621"),
        "American International Foods": ("Silva", "https://www.linkedin.com/in/tony-silva-96b18a30"),
        "Charkit Chemical": ("Yannopoulos", "https://www.linkedin.com/in/panos-yannopoulos-18b79931"),
        "Core-Mark": ("Beck", "https://www.linkedin.com/in/chandler-beck-44908515"),
        "Creed Foodservice": ("Butcher", "https://www.linkedin.com/in/charlotte-butcher-121b21110"),
        "Doerle Food Services": ("Martin", "https://www.linkedin.com/in/charlie-martin-32187a1a"),
        "Dubord & Rainville": ("Aubin", "https://www.linkedin.com/in/g%C3%A9rald-aubin-85672780"),
        "DKSH": ("Leenman", "https://www.linkedin.com/in/sander-leenman-112369134"),
        "Circle Ingredients": ("Dieleman", "https://www.linkedin.com/in/wilko-dieleman-70720227"),
        "ChemCeed": ("Geissler", "https://www.linkedin.com/in/katrina-geissler-8b6a1941"),
        "FreshPoint": ("Arthur", "https://www.linkedin.com/in/eric-arthur-mba-22b7475"),
        "KeHE Specialty": ("Widtfeldt", "https://www.linkedin.com/in/widtfeldt"),
        "Imperial Dade": ("Beck", "https://www.linkedin.com/in/richard-beck-792b5654"),
        "Jetro Restaurant Depot": ("Carpenter", "https://www.linkedin.com/in/chris-carpenter-570006a2"),
        "H.T. Hackney": ("Thomas", "https://www.linkedin.com/in/tommy-thomas-0a2b5967"),
        "Londis Wholesale": ("Hayes", "https://www.linkedin.com/in/conor-hayes-ab3676105"),
        "Fyffes": ("Purcell", "https://www.linkedin.com/in/purcellsean"),
        "Grocers Supply": ("Morton", "https://www.linkedin.com/in/bob-morton-aab845104"),
        "Labatt Food Service": ("Monacelli", "https://www.linkedin.com/in/ken-monacelli-b5878b82"),
    },
    "smartphone": {
        "CDW": ("DeCoster", "https://www.linkedin.com/in/tomdecostertech"),
        "Arrow Electronics": ("Rampersad", "https://www.linkedin.com/in/salesh-rampersad-3a74b23"),
        "3G Mobile": ("Camara", "https://www.linkedin.com/in/john-camara-5908a916"),
        "Esprinet": ("Florio", "https://www.linkedin.com/in/claudioflorio"),
        "Exertis Ireland": ("Coogan", "https://www.linkedin.com/in/john-coogan-768a634"),
        "Elkjøp": ("Samsami", "https://www.linkedin.com/in/farshadsamsami"),
        "Elko Group": ("Drebeinika", "https://www.linkedin.com/in/laura-drebeinika-ex-klava"),
        "Freedom Mobile": ("Rafanelli", "https://www.linkedin.com/in/nino-rafanelli-12707520"),
        "Fonehouse": ("Connell", "https://www.linkedin.com/in/peteoconnell16"),
        "Ingram Micro": ("Nambudiri", "https://www.linkedin.com/in/vijaykumarnambudiri1980"),
    },
    "firewall": {
        "Dicker Data": ("Elgammal", "https://www.linkedin.com/in/yasserelgammal"),
        "Insight Enterprises": ("Kane", "https://www.linkedin.com/in/robertfkane"),
        "Optiv Security": ("Hurley", "https://www.linkedin.com/in/johnchurley"),
        "Computer Gross": ("Gaini", "https://www.linkedin.com/in/gainifrancesco"),
        "Jenne": ("LoSchiavo", "https://www.linkedin.com/in/sal-loschiavo-b446284"),
        "Layer 3 Communications": ("Faxon", "https://www.linkedin.com/in/scottfaxon"),
        "Kyndryl": ("Cirocco", "https://www.linkedin.com/in/angelocirocco"),
        "Iron Bow Technologies": ("Cook", "https://www.linkedin.com/in/donjcook"),
        "MA Labs": ("Carvalho", "https://www.linkedin.com/in/george-de-carvalho-5593948"),
        "NetSync Network Solutions": ("Noble", "https://www.linkedin.com/in/len-noble-264a64"),
        "Wortmann AG": ("Kaps", "https://www.linkedin.com/in/volker-kaps-253530306"),
        "Logicalis": ("Au", "https://www.linkedin.com/in/iau01"),
        "Esprinet Iberica": ("Pamplona", "https://www.linkedin.com/in/anapamplona"),
        "Aptec (Ingram Micro Middle East)": ("Mishra", "https://www.linkedin.com/in/gyaneshwar-mishra-6a92b113"),
        "Alviva Holdings": ("Parkin", "https://www.linkedin.com/in/vaughn-parkin-378b9415"),
        "Intcomex": ("Bouhadana", "https://www.linkedin.com/in/josephbouhadana"),
    },
}


def save_atomic(wb, path: Path):
    tmp = path.with_suffix(path.suffix + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def main():
    n = 0
    for mkt, path in FILES.items():
        if mkt not in URLS:
            continue
        wb = load_workbook(path)
        ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
        headers = None
        hr = 1
        for r in range(1, 6):
            vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
            if "Company" in vals and "Contact Person" in vals:
                headers = {h: i + 1 for i, h in enumerate(vals) if h}
                hr = r
                break
        c_co = headers["Company"]
        c_cp = headers["Contact Person"]
        c_li = headers["LinkedIn"]
        for row in range(hr + 1, ws.max_row + 1):
            co = str(ws.cell(row, c_co).value or "").strip()
            if co not in URLS[mkt]:
                continue
            person, url = URLS[mkt][co]
            cp = str(ws.cell(row, c_cp).value or "")
            if person.lower() not in cp.lower():
                print("SKIP name mismatch", mkt, co, cp)
                continue
            old = str(ws.cell(row, c_li).value or "").strip()
            if "linkedin.com/in/" in old.lower() and old == url:
                continue
            ws.cell(row, c_li).value = url
            n += 1
            print(f"LI {mkt} | {co} | {cp} | {url}")
        save_atomic(wb, path)
        wb.close()
    print("linkedin_fills", n)


if __name__ == "__main__":
    main()
