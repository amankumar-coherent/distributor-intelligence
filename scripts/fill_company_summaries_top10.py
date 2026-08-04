"""Write a proper per-company Summary for all top-10 expanded Excels."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEMO = Path(r"C:\Users\aman.kumar\Desktop\Demo Output")
SRC = Path(r"D:\vendor-intel\output")

MARKET_LABELS = {
    "axillary_crutches_market_global": "axillary crutches / DME mobility",
    "baijiu_market_global": "Baijiu / Chinese spirits",
    "benign_positional_vertigo_market_global": "BPV / vestibular clinical products",
    "elderberry_supplements_market_global": "elderberry / immune supplements",
    "instant_noodles_market_global": "instant noodles / ramen",
    "mountain_bike_suspension_market_global": "mountain bike suspension components",
    "muslin_products_market_global": "muslin / cotton textile products",
    "nail_polish_market_global": "nail polish / professional nail systems",
    "sheet_face_masks_market_global": "sheet face masks / K-beauty skincare",
    "shunt_resistor_market_global": "shunt / current-sense resistors",
}

WIDTH_HINTS = {
    "company": 28,
    "website": 32,
    "core categories": 36,
    "specialty focus": 32,
    "price segment": 24,
    "key brands": 36,
    "headquarters": 26,
    "cities": 24,
    "ownership": 20,
    "business type": 18,
    "employees": 12,
    "turnover": 18,
    "contact": 18,
    "founded": 10,
    "email": 24,
    "phone": 16,
    "summary": 58,
    "role": 18,
    "linkedin": 28,
    "regions": 18,
    "distribution": 16,
    "competitive": 12,
    "data confidence": 12,
    "quality": 10,
    "data sources": 28,
}

thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
body_font = Font(name="Calibri", size=10)
wrap = Alignment(wrap_text=True, vertical="top", horizontal="left")
header_align = Alignment(wrap_text=True, vertical="center", horizontal="center")


def clean(v: object) -> str:
    s = str(v or "").strip()
    if not s:
        return ""
    low = s.lower()
    if low.startswith("not ") or low in {"n/a", "na", "unknown", "-", "—"}:
        return ""
    return s


def short(v: object, n: int = 120) -> str:
    s = clean(v)
    if len(s) <= n:
        return s
    cut = s[:n]
    for sep in ("; ", ", ", " - "):
        i = cut.rfind(sep)
        if i > 40:
            return cut[:i].rstrip(" ;,") + "..."
    return cut.rstrip() + "..."


def build_summary(market: str, row: dict) -> str:
    company = clean(row.get("Company")) or "This company"
    label = MARKET_LABELS.get(market, market.replace("_", " "))
    biz = clean(row.get("Business Type")) or "channel participant"
    hq = clean(row.get("Headquarters")) or clean(row.get("Cities / Regions"))
    founded = clean(row.get("Founded"))
    ownership = clean(row.get("Ownership"))
    core = short(row.get("Core Categories"), 110)
    spec = short(row.get("Specialty Focus"), 90)
    price = short(row.get("Price Segment"), 60)
    brands = short(row.get("Key Brands Represented"), 100)
    emp = clean(row.get("Employees"))
    scale = clean(row.get("Turnover / Scale"))
    contact = clean(row.get("Contact Person"))
    role = clean(row.get("Role"))
    website = clean(row.get("Website"))

    biz_l = biz.lower()
    article = "an" if biz_l[:1] in "aeiou" else "a"
    lead = f"{company} is {article} {biz_l} in the {label} market"
    if hq:
        lead += f", headquartered in {short(hq, 70)}"
    if founded:
        lead += f" (founded {founded})"
    lead += "."

    detail = ""
    if core:
        detail += f" Core categories include {core}."
    if spec:
        detail += f" Specialty focus is {spec}."
    if price:
        detail += f" Price segment: {price}."
    if brands:
        detail += f" Key brands represented: {brands}."

    extras: list[str] = []
    if ownership:
        extras.append(f"Ownership: {short(ownership, 70)}")
    if emp:
        extras.append(f"Employees: {emp}")
    if scale:
        extras.append(f"Scale: {short(scale, 60)}")
    if contact:
        who = contact + (f" ({role})" if role else "")
        extras.append(f"Contact: {who}")
    if website:
        extras.append(f"Website: {website}")
    extra = (" " + ". ".join(extras) + ".") if extras else ""

    text = (lead + detail + extra).strip()
    return re.sub(r"\s+", " ", text)


def col_width(header: str) -> float:
    h = (header or "").strip().lower()
    for key, w in WIDTH_HINTS.items():
        if key in h:
            return float(w)
    return float(max(12, min(30, len(header or "") + 3)))


def format_sheet(ws, header_row: int, max_col: int, max_row: int) -> None:
    for c in range(1, max_col + 1):
        h = str(ws.cell(header_row, c).value or "").strip()
        ws.column_dimensions[get_column_letter(c)].width = col_width(h)
        cell = ws.cell(header_row, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin

    for r in range(header_row + 1, max_row + 1):
        if not any(ws.cell(r, c).value for c in range(1, max_col + 1)):
            continue
        max_lines = 1
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin
            val = cell.value
            if not val:
                continue
            text = str(val)
            width = ws.column_dimensions[get_column_letter(c)].width or 20
            cpl = max(8, int(width * 1.05))
            lines = 0
            for part in text.replace("\r", "").split("\n"):
                lines += max(1, (len(part) + cpl - 1) // cpl)
            hname = str(ws.cell(header_row, c).value or "")
            cap = 8 if "summary" in hname.lower() else 5
            max_lines = max(max_lines, min(cap, lines))
        ws.row_dimensions[r].height = min(120, max(40, 13 * max_lines + 8))

    ws.row_dimensions[header_row].height = 34
    ws.freeze_panes = ws.cell(header_row + 1, 2).coordinate
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{max_row}"


def process(path: Path) -> int:
    market = path.parent.name
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=False))
    hdr_i = next(
        i
        for i, r in enumerate(rows[:12])
        if r and r[0].value and str(r[0].value).strip() == "Company"
    )
    header_row = hdr_i + 1
    header = [str(c.value or "").strip() for c in rows[hdr_i]]

    if "Summary" not in header:
        max_existing = max((i for i, h in enumerate(header) if h), default=0) + 1
        ws.cell(header_row, max_existing + 1).value = "Summary"
        rows = list(ws.iter_rows(values_only=False))
        header = [str(c.value or "").strip() for c in rows[hdr_i]]

    idx = {h: i for i, h in enumerate(header) if h}
    si = idx["Summary"]
    updated = 0
    for r in rows[hdr_i + 1 :]:
        company = str(r[idx["Company"]].value or "").strip()
        if not company:
            continue
        rowd = {h: (r[i].value if i < len(r) else None) for h, i in idx.items()}
        summary = build_summary(market, rowd)
        if str(r[si].value or "").strip() != summary:
            r[si].value = summary
            updated += 1

    max_col = max(i for i, h in enumerate(header) if h) + 1
    format_sheet(ws, header_row, max_col, ws.max_row)

    if "Run Summary" in wb.sheetnames:
        wb["Run Summary"]["A18"] = f"Company summaries written: {updated}"

    wb.save(path)
    wb.close()
    DEMO.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, DEMO / path.name)
    return updated


def main() -> None:
    total = 0
    for path in sorted(SRC.rglob("*_distributors_expanded.xlsx")):
        n = process(path)
        print(f"{path.parent.name}: {n} summaries")
        total += n
    print(f"TOTAL {total}")
    print(f"Demo Output: {DEMO}")


if __name__ == "__main__":
    main()
