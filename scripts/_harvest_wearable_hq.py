# -*- coding: utf-8 -*-
"""Harvest HQ city hints from company About/Contact pages. Never invent."""
from __future__ import annotations

import json
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlparse

import httpx
from openpyxl import load_workbook

ROOT = Path(
    r"D:\vendor-intel\output\web_expand\global_wearable_medical_devices_market_global"
)
XLSX = ROOT / "global_wearable_medical_devices_market_global_FINAL.xlsx"
CACHE = ROOT / "_wearable_hq_harvest_cache.json"
NEED = ROOT / "_wearable_hq_country_only.json"

ABOUT_PATHS = [
    "",
    "/about",
    "/about-us",
    "/about/",
    "/contact",
    "/contact-us",
    "/contact/",
    "/company",
    "/locations",
    "/en/about",
    "/en/company",
    "/pages/about-us",
    "/pages/contact",
]

# Common city, country / city, state patterns from footer/contact
ADDR_PATTERNS = [
    re.compile(
        r"(?i)\b([A-Z][A-Za-z .'-]{1,40}),\s*([A-Z]{2})\s+\d{5}(?:-\d{4})?\b"  # City, ST ZIP
    ),
    re.compile(
        r"(?i)\b([A-Z][A-Za-z .'-]{1,40}),\s*(United States|USA|United Kingdom|UK|Canada|Australia|Germany|France|Japan|Singapore|India|Brazil|Sweden|Netherlands|Belgium|Spain|Italy|Switzerland|Denmark|Finland|Norway|Mexico|Malaysia|South Africa|Saudi Arabia|UAE|United Arab Emirates|Hong Kong|China|South Korea|New Zealand|Ireland|Austria|Poland)\b"
    ),
]


def base_url(url: str) -> str:
    u = (url or "").strip()
    if not u:
        return ""
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    p = urlparse(u)
    return f"{p.scheme}://{p.netloc}"


def plain(html: str) -> str:
    t = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", html or "")
    t = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", t)
    t = re.sub(r"(?is)<[^>]+>", " ", t)
    t = re.sub(r"&nbsp;", " ", t)
    t = re.sub(r"\s+", " ", t)
    return t


def extract_hq(text: str, country_hint: str) -> tuple[str, str]:
    if not text:
        return "", ""
    hits = []
    for pat in ADDR_PATTERNS:
        for m in pat.finditer(text):
            city = m.group(1).strip(" ,.")
            rest = m.group(2).strip()
            # filter noise
            bad = {"copyright", "all rights", "privacy", "cookie", "terms", "follow", "share", "email", "phone"}
            if any(b in city.lower() for b in bad):
                continue
            if len(city) < 3 or len(city) > 40:
                continue
            if rest.upper() in {"AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA", "HI", "ID", "IL", "IN", "IA", "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH", "NJ", "NM", "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV", "WI", "WY", "DC"}:
                hq = f"{city}, {rest}, United States"
            else:
                # normalize country names
                cmap = {
                    "USA": "United States",
                    "US": "United States",
                    "UK": "United Kingdom",
                    "UAE": "United Arab Emirates",
                }
                country = cmap.get(rest, rest)
                hq = f"{city}, {country}"
            ctx = text[max(0, m.start() - 30) : m.end() + 30]
            hits.append((hq, ctx))
    if not hits:
        # headquarters in X
        m = re.search(
            r"(?i)\b(?:headquarters|head office|based in|located in|hq[:\s]+)\s*([A-Z][A-Za-z .'-]{2,40}(?:,\s*[A-Z][A-Za-z .'-]{2,40})?)",
            text,
        )
        if m:
            frag = m.group(1).strip(" .,")
            if 3 < len(frag) < 60:
                return frag, m.group(0)[:120]
        return "", ""
    # prefer hit mentioning country_hint
    hint = (country_hint or "").lower()
    for hq, ctx in hits:
        if hint and hint[:4] in hq.lower():
            return hq, ctx[:160]
    return hits[0][0], hits[0][1][:160]


def harvest_one(company: str, website: str, country_hint: str) -> dict:
    root = base_url(website)
    if not root:
        return {"company": company, "website": website, "hq": "", "source": "", "evidence": "", "ok": False}
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml",
    }
    best = ("", "", "")
    with httpx.Client(timeout=10.0, headers=headers, verify=False) as client:
        for path in ABOUT_PATHS:
            url = root if path == "" else root.rstrip("/") + path
            try:
                r = client.get(url, follow_redirects=True)
                if r.status_code >= 400:
                    continue
                hq, ev = extract_hq(plain(r.text[:200000]), country_hint)
                if hq:
                    best = (hq, url, ev)
                    if path and any(k in path for k in ("about", "contact", "compan", "location")):
                        break
            except Exception:
                continue
    return {
        "company": company,
        "website": website,
        "hq": best[0],
        "source": best[1],
        "evidence": best[2],
        "ok": bool(best[0]),
    }


def main() -> None:
    import urllib3

    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    need = json.loads(NEED.read_text(encoding="utf-8")) if NEED.exists() else []
    # rebuild need from excel if empty
    if not need:
        wb = load_workbook(XLSX, data_only=True)
        ws = wb.active
        hdr = [str(c.value or "") for c in ws[1]]
        ic = {h: i + 1 for i, h in enumerate(hdr)}
        countries = {
            "united states",
            "usa",
            "us",
            "united kingdom",
            "uk",
            "germany",
            "france",
            "italy",
            "spain",
            "netherlands",
            "belgium",
            "sweden",
            "norway",
            "denmark",
            "finland",
            "switzerland",
            "austria",
            "poland",
            "ireland",
            "canada",
            "mexico",
            "brazil",
            "argentina",
            "chile",
            "australia",
            "new zealand",
            "japan",
            "china",
            "south korea",
            "korea",
            "india",
            "singapore",
            "malaysia",
            "uae",
            "united arab emirates",
            "saudi arabia",
            "qatar",
            "kuwait",
            "bahrain",
            "oman",
            "egypt",
            "south africa",
            "kenya",
            "hong kong",
            "spain",
            "italy",
        }
        for r in range(2, ws.max_row + 1):
            co = str(ws.cell(r, ic["Company"]).value or "").strip()
            web = str(ws.cell(r, ic["Website"]).value or "").strip()
            hq = str(ws.cell(r, ic["Headquarters"]).value or "").strip()
            if hq.lower() in countries or ("," not in hq and hq.lower() in countries):
                need.append({"row": r, "company": co, "website": web, "hq": hq})
        wb.close()

    cache = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    todo = [t for t in need if not cache.get(t["company"], {}).get("ok")]
    print(f"need={len(need)} harvest={len(todo)}", flush=True)

    done = 0
    with ThreadPoolExecutor(max_workers=20) as ex:
        futs = {
            ex.submit(harvest_one, t["company"], t["website"], t.get("hq", "")): t for t in todo
        }
        for fut in as_completed(futs):
            t = futs[fut]
            try:
                res = fut.result()
            except Exception as e:
                res = {
                    "company": t["company"],
                    "website": t["website"],
                    "hq": "",
                    "source": "",
                    "evidence": str(e),
                    "ok": False,
                }
            cache[t["company"]] = res
            done += 1
            if done % 40 == 0 or done == len(todo):
                CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
                ok = sum(1 for v in cache.values() if v.get("ok"))
                print(f"harvested {done}/{len(todo)} ok~{ok}", flush=True)

    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")
    ok = [(k, v["hq"], v.get("source", "")) for k, v in cache.items() if v.get("ok")]
    print("OK hits", len(ok))
    for k, hq, src in sorted(ok)[:80]:
        print(f"{hq} | {k} | {src[:50]}")


if __name__ == "__main__":
    main()
