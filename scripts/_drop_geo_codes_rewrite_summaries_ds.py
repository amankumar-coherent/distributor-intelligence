#!/usr/bin/env python3
"""Remove Country/Region Code cols; rewrite Summaries via DeepSeek from verified row facts."""
from __future__ import annotations

import json
import os
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(r"D:\vendor-intel")
load_dotenv(ROOT / ".env")

XLSX = ROOT / "output/web_expand/global_avocado_oil_market_global/global_avocado_oil_market_OFFICE_SLIM.xlsx"
CACHE = ROOT / "output/web_expand/global_avocado_oil_market_global/_avocado_summary_ds_cache.json"
AUDIT = ROOT / "output/web_expand/global_avocado_oil_market_global/_avocado_summary_ds_audit.json"
NPD = "Not publicly disclosed"
DROP_COLS = {"Country Code", "Region Code"}


def load_cache() -> dict:
    if CACHE.exists():
        try:
            return json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_cache(cache: dict) -> None:
    CACHE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")


def drop_columns(ws) -> list[str]:
    hdr = [c.value for c in ws[2]]
    # delete from right so indices stay valid
    for i in range(len(hdr), 0, -1):
        h = hdr[i - 1]
        if h in DROP_COLS:
            ws.delete_cols(i)
    return [c.value for c in ws[2] if c.value]


def client() -> OpenAI:
    key = os.getenv("OPENAI_API_KEY") or ""
    base = os.getenv("OPENAI_BASE_URL") or "https://api.deepseek.com/v1"
    if not key:
        raise SystemExit("OPENAI_API_KEY missing in .env")
    return OpenAI(api_key=key, base_url=base)


def summarize(cli: OpenAI, model: str, facts: dict) -> str:
    prompt = f"""Write ONE clean professional company summary for a vendor-intel edible oils / avocado-oil channel landscape spreadsheet.

Rules:
- 2 short sentences, max 320 characters total.
- Use ONLY the facts provided. Do not invent phones, emails, brands, founded year, HQ, or ownership.
- If a field is "{NPD}" or empty, do not invent a replacement.
- Mention role (distributor/wholesaler/etc), geography if given, core categories/specialty briefly, and key brands if given and not NPD.
- Plain text only. No bullets, no markdown, no quotes around the whole summary.
- Do not start with "Summary:".

Facts JSON:
{json.dumps(facts, ensure_ascii=False)}
"""
    resp = cli.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": "You write concise verified business summaries. Never invent facts.",
            },
            {"role": "user", "content": prompt},
        ],
        temperature=0.2,
        max_tokens=220,
    )
    text = (resp.choices[0].message.content or "").strip()
    text = text.strip('"').strip("'").replace("\n", " ").strip()
    if len(text) > 480:
        text = text[:477].rstrip() + "..."
    return text


def main() -> None:
    model = os.getenv("OPENAI_MODEL") or "deepseek-chat"
    cli = client()
    cache = load_cache()

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    cols = drop_columns(ws)
    ic = {str(h): i + 1 for i, h in enumerate(cols) if h}
    assert "Summary" in ic and "Company" in ic
    assert "Country Code" not in ic and "Region Code" not in ic

    companies = []
    for r in range(3, ws.max_row + 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        if co:
            companies.append(r)

    rewritten = 0
    errors = []
    for idx, r in enumerate(companies, 1):
        co = str(ws.cell(r, ic["Company"]).value or "").strip()
        facts = {
            "company": co,
            "website": str(ws.cell(r, ic["Website"]).value or "").strip(),
            "founded": str(ws.cell(r, ic["Founded"]).value or "").strip(),
            "headquarters": str(ws.cell(r, ic["Headquarters"]).value or "").strip(),
            "continent": str(ws.cell(r, ic["Continent / Geography"]).value or "").strip(),
            "operational_presence": str(ws.cell(r, ic["Operational Presence"]).value or "").strip(),
            "ownership": str(ws.cell(r, ic["Ownership"]).value or "").strip(),
            "employees": str(ws.cell(r, ic["Employees"]).value or "").strip(),
            "core_categories": str(ws.cell(r, ic["Core Categories"]).value or "").strip(),
            "specialty_focus": str(ws.cell(r, ic["Specialty Focus"]).value or "").strip(),
            "key_brands": str(ws.cell(r, ic["Key Brands Represented"]).value or "").strip(),
            "office_no": str(ws.cell(r, ic["Office No."]).value or "").strip() if "Office No." in ic else "",
            "office_email": str(ws.cell(r, ic["Office Email"]).value or "").strip() if "Office Email" in ic else "",
        }
        # fingerprint so cache invalidates if facts change
        fp = json.dumps(facts, sort_keys=True, ensure_ascii=False)
        cached = cache.get(co)
        if isinstance(cached, dict) and cached.get("fp") == fp and cached.get("summary"):
            summary = cached["summary"]
            source = "cache"
        else:
            try:
                summary = summarize(cli, model, facts)
                source = "deepseek"
                cache[co] = {"fp": fp, "summary": summary, "ts": datetime.now(timezone.utc).isoformat()}
                if idx % 10 == 0:
                    save_cache(cache)
                    wb.save(XLSX)
            except Exception as e:
                errors.append({"company": co, "error": str(e)})
                print(f"[{idx}/{len(companies)}] ERROR {co}: {e}", flush=True)
                time.sleep(1.5)
                continue

        ws.cell(r, ic["Summary"]).value = summary
        rewritten += 1
        print(f"[{idx}/{len(companies)}] {source}: {co}", flush=True)
        if source == "deepseek":
            time.sleep(0.35)

    save_cache(cache)
    wb.save(XLSX)
    wb.close()

    # recount empties
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[2] if c.value]
    ic2 = {str(h): i for i, h in enumerate([c.value for c in ws[2]]) if h}
    n = empty_sum = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        c = str(row[ic2["Company"]] or "").strip()
        if not c:
            continue
        n += 1
        s = str(row[ic2["Summary"]] or "").strip()
        if not s or s == NPD:
            empty_sum += 1
    wb.close()

    audit = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "model": model,
        "dropped_columns": sorted(DROP_COLS),
        "headers_now": hdr,
        "companies": n,
        "summaries_written": rewritten,
        "summary_empty_or_npd": empty_sum,
        "errors": errors,
    }
    AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({k: audit[k] for k in ("dropped_columns", "companies", "summaries_written", "summary_empty_or_npd", "headers_now")}, indent=2))
    if errors:
        print("errors:", len(errors))


if __name__ == "__main__":
    main()
