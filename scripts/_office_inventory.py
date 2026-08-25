# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}

PATHS = [
    Path(r"D:\vendor-intel\output\web_expand\global_electrical_distribution_market_global\global_electrical_distribution_market_global_FINAL.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_firewall_market_global\global_firewall_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_smartphone_market_global\global_smartphone_market.xlsx"),
    Path(r"D:\vendor-intel\output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market.xlsx"),
    Path(r"D:\vendor-intel\output\run4\global_smartwatch_market_global\global_smartwatch_market.xlsx"),
]


def load(p: Path):
    wb = load_workbook(p, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals:
            headers = vals
            hr = r
            break
    rows = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {
            headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
        }
        if d.get("Company"):
            rows.append(d)
    wb.close()
    return headers, rows


def has_val(s: str) -> bool:
    return bool(s) and s.strip().lower() not in NPD


for p in PATHS:
    if not p.exists():
        print("MISSING", p)
        continue
    headers, rows = load(p)
    office_keys = [h for h in headers if "office" in h.lower() or "phone" in h.lower() or "whatsapp" in h.lower()]
    print("=" * 80)
    print(p.name)
    print("office/phone headers:", office_keys)
    office = 0
    phone = 0
    empty_office = 0
    empty_both = 0
    can_copy = 0
    empties = []
    for d in rows:
        off = d.get("Office No.") or d.get("Office No") or d.get("Office Number") or ""
        ph = d.get("Phone / WhatsApp") or d.get("Phone") or ""
        ho = has_val(off)
        hp = has_val(ph)
        if ho:
            office += 1
        else:
            empty_office += 1
            if hp:
                can_copy += 1
            else:
                empty_both += 1
                empties.append((d.get("Company"), d.get("Website"), d.get("Headquarters") or d.get("HQ") or ""))
        if hp:
            phone += 1
    print(f"rows={len(rows)} office_filled={office} phone_filled={phone} empty_office={empty_office} can_copy_from_phone={can_copy} empty_both={empty_both}")
    for i, e in enumerate(empties[:25], 1):
        print(f"  {i}\t{e[0]}\t{e[1]}\t{e[2]}")
    if len(empties) > 25:
        print(f"  ... +{len(empties)-25} more empty both")
