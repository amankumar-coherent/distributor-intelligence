#!/usr/bin/env python3
"""Apply extra curated fixes for remaining NPD + Alliance year correction."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"D:\vendor-intel")

FIXES: dict[str, dict[str, dict[str, str]]] = {
    str(ROOT / r"output\web_expand\global_antibody_drug_conjugates_market_global\global_antibody_drug_conjugates_market_FINAL_SLIM.xlsx"): {
        "Alliance Healthcare Ireland": {
            "Website": "https://www.alliance-healthcare.com",
            "Founded": "1938",
            "Headquarters": "Dublin, Ireland",
        },
        "Alliance Healthcare Polska": {
            "Website": "https://www.alliance-healthcare.com",
            "Founded": "1938",
            "Headquarters": "Warsaw, Poland",
        },
        "Alliance Healthcare Portugal": {
            "Website": "https://www.alliance-healthcare.com",
            "Founded": "1938",
            "Headquarters": "Lisbon, Portugal",
        },
        "Pharmalog Austria": {
            "Website": "https://www.pharma-logistik.at",
            "Founded": "2006",
            "Headquarters": "Wels, Austria",
        },
        "Highmark Specialty Pharmacy": {
            "Website": "https://www.highmark.com",
            "Founded": "1996",
            "Headquarters": "Pittsburgh, Pennsylvania, United States",
        },
        "LeaderNET specialty pharmacy network": {
            "Website": "https://www.cardinalhealth.com",
            "Founded": "1971",
            "Headquarters": "Dublin, Ohio, United States",
        },
    },
}


def norm(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def header_row(ws) -> int:
    for r in range(1, min(6, ws.max_row + 1)):
        vals = [norm(c.value) for c in ws[r] if c.value not in (None, "")]
        if "company" in vals:
            return r
    return 1


def main() -> None:
    for fpath, mapping in FIXES.items():
        p = Path(fpath)
        if not p.exists():
            print("MISSING", p)
            continue
        wb = load_workbook(p)
        ws = wb.active
        hr = header_row(ws)
        headers = [str(c.value or "").strip() for c in ws[hr]]
        ci = headers.index("Company")
        want = {norm(k): v for k, v in mapping.items()}
        n = 0
        for r in range(hr + 1, ws.max_row + 1):
            company = ws.cell(r, ci + 1).value
            if not company:
                continue
            ch = want.get(norm(str(company)))
            if not ch:
                continue
            for col, val in ch.items():
                if col in headers:
                    ws.cell(r, headers.index(col) + 1).value = val
            n += 1
            print(f"  FIX {company}: {ch}")
        wb.save(p)
        print(f"Updated {n} -> {p.name}")


if __name__ == "__main__":
    main()
