# -*- coding: utf-8 -*-
"""Apply web-verified sales/procurement contacts to flexpack FINAL."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

P = Path(
    r"D:\vendor-intel\output\web_expand\global_flexible_packaging_market_global"
    r"\global_flexible_packaging_market_global_FINAL.xlsx"
)
NPD = "Not publicly disclosed"

HAND = {
    "Allpack": (
        "Marcus Clarke",
        "Sales Director",
        "https://www.linkedin.com/in/marcus-clarke-523240a0",
    ),
    "Antalis": (
        "Gary Hood",
        "Head of Sales Operations",
        "https://www.linkedin.com/in/gary-hood-63761955",
    ),
    "Manutan": (
        "Jose Palomares Villar",
        "Head of Sales",
        "https://www.linkedin.com/in/jose-palomares-sales-director",
    ),
    "Fastenal": (
        "Heath Hazard",
        "Vice President of Sales",
        "https://www.linkedin.com/in/heath-hazard-39727440",
    ),
    "Grainger": (
        "Ryan Adams",
        "Regional Vice President of Sales",
        "https://www.linkedin.com/in/ryan-adams-aa42624",
    ),
    "Lyreco": (
        "David Brace",
        "Corporate Accounts Sales Director",
        "https://www.linkedin.com/in/david-brace",
    ),
    "RAJAPACK UK": (
        "Matthew Whitaker",
        "Commercial Director",
        "https://www.linkedin.com/in/matthew-whitaker-38849636",
    ),
    "Allied Packaging": (
        "Chris Scoggins",
        "Regional Vice President of Sales",
        "https://www.linkedin.com/in/chrisjscoggins",
    ),
    "Bunzl": (
        "Lee Ann Faulkner",
        "Sales Director AGGORA Group",
        "https://www.linkedin.com/in/lee-ann-faulkner-a4010661",
    ),
    "MSC Industrial Supply": (
        "Jim Greene",
        "Vice President of Field Sales",
        "https://www.linkedin.com/in/jim-greene-14582351",
    ),
    "Imperial Supplies": (
        "Jon Daul",
        "Vice President of Sales",
        "https://www.linkedin.com/in/jon-daul-09a58b4b",
    ),
    "Cospak": (
        "Jonathan Evans",
        "Head of Procurement and Manufacturing",
        "https://www.linkedin.com/in/jonathan-evans-12503511",
    ),
    "GWP Group": ("David Mason", "Sales Director", NPD),
    "Swiftpak": (
        "Josh Bebb",
        "Head of Sales",
        "https://www.linkedin.com/in/josh-bebb",
    ),
    "Uline": (
        "Todd W McBride",
        "Vice President Sales, North America",
        "https://www.linkedin.com/in/todd-w-mcbride-15a6464",
    ),
    "Macfarlane Packaging": (
        "Paul Armstrong",
        "Sales Director Europe",
        "https://www.linkedin.com/in/paul-armstrong-31659451",
    ),
    "Associated Packaging": (
        "Jeff Hoeppner",
        "Corporate Vice President of Sales",
        "https://www.linkedin.com/in/jeffhoeppner-mba",
    ),
    "Veritiv": (
        "Chris Wheat",
        "Vice President of Sales",
        "https://www.linkedin.com/in/chris-wheat-070bb410",
    ),
    "Shorr Packaging": (
        "Steve Golding",
        "East Region Vice President of Sales",
        "https://www.linkedin.com/in/steve-golding-65aa81235",
    ),
    "Imperial Dade": (
        "Jason Olthoff",
        "VP of Sales, East Central Region",
        "https://www.linkedin.com/in/jason-olthoff-404a3732",
    ),
    "Staples": (
        "Kelly Cramer",
        "Area Vice President, Vertical Markets (Facility Solutions)",
        "https://www.linkedin.com/in/kelly-cramer-9b02034",
    ),
    "RAJA": (
        "Brian Hall",
        "International Commercial Director, RAJA Office",
        "https://www.linkedin.com/in/brian-hall-bb928728",
    ),
    "Crawford Packaging": (
        "Martin Schieck",
        "Vice President Sales & Marketing",
        "https://www.linkedin.com/in/martin-schieck-46579b1a",
    ),
    "Acme Paper & Supply": (
        "Andy Attman",
        "Vice President of Sales",
        "https://www.linkedin.com/in/andy-attman-0404b75",
    ),
    "Nashville Wraps": (
        "Greg Phairas",
        "Director of Sales",
        "https://www.linkedin.com/in/greg-phairas-5303b21b",
    ),
    "IPS Packaging": (
        "Todd Meadows",
        "Vice President of Sales",
        "https://www.linkedin.com/in/todd-meadows",
    ),
    "Ratioform": (
        "Marc Zipperlen",
        "Head of Sales Excellence, Industrial & Packaging (TAKKT)",
        "https://www.linkedin.com/in/marc-zipperlen-5966821a2",
    ),
    "Davpack": (
        "Marc Zipperlen",
        "Head of Sales Excellence, Industrial & Packaging (TAKKT)",
        "https://www.linkedin.com/in/marc-zipperlen-5966821a2",
    ),
}

CLEAR_TO_NPD = {"Kite Packaging"}  # Sales Manager — below Head/VP bar


def main() -> None:
    wb = load_workbook(P)
    ws = wb.active
    h = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: h.index(name) + 1 for name in h}

    applied: list[str] = []
    cleared: list[str] = []
    company_li_cleared = 0

    for r in range(2, ws.max_row + 1):
        name = (ws.cell(r, idx["Company"]).value or "").strip()

        li = (ws.cell(r, idx["LinkedIn"]).value or "").strip()
        if "linkedin.com/company/" in li.lower():
            ws.cell(r, idx["LinkedIn"]).value = NPD
            company_li_cleared += 1

        if name in CLEAR_TO_NPD:
            ws.cell(r, idx["Contact Person"]).value = NPD
            ws.cell(r, idx["Role"]).value = NPD
            cur_li = (ws.cell(r, idx["LinkedIn"]).value or "").strip()
            if "linkedin.com/in/" not in cur_li.lower():
                ws.cell(r, idx["LinkedIn"]).value = NPD
            cleared.append(name)
            continue

        if name in HAND:
            person, role, linkedin = HAND[name]
            ws.cell(r, idx["Contact Person"]).value = person
            ws.cell(r, idx["Role"]).value = role
            ws.cell(r, idx["LinkedIn"]).value = linkedin
            applied.append(name)

        for col in ("Contact Person", "Role"):
            v = (ws.cell(r, idx[col]).value or "").strip()
            if "@" in v or v.lower().startswith("http"):
                ws.cell(r, idx[col]).value = NPD

    wb.save(P)
    wb.close()
    print("applied", len(applied), sorted(set(applied)))
    print("cleared", cleared)
    print("company_li_cleared", company_li_cleared)


if __name__ == "__main__":
    main()
