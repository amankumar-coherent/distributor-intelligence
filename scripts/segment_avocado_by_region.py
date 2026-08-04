"""Segment avocado oil FINAL Excel by macro-region (sheets + Region column)."""
from __future__ import annotations

import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PATH = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global\global_avocado_oil_market_global_FINAL.xlsx"
)

# Display order of region sheets (Excel sheet titles cannot contain / \ ? * [ ])
REGIONS = [
    "North America",
    "Europe",
    "Middle East and Africa",
    "Asia-Pacific",
    "Latin America",
    "Global Multi-Region",
]

REGION_LABELS = {
    "North America": "North America",
    "Europe": "Europe",
    "Middle East & Africa": "Middle East and Africa",
    "Asia-Pacific": "Asia-Pacific",
    "Latin America": "Latin America",
    "Global / Multi-Region": "Global Multi-Region",
}

COL_WIDTHS = {
    "Region Segment": 18,
    "Company": 28,
    "Founded": 10,
    "Headquarters": 28,
    "Cities / Regions": 32,
    "Ownership": 20,
    "Business Type": 26,
    "Employees": 12,
    "Turnover / Scale": 14,
    "Contact Person": 20,
    "Role": 22,
    "Email": 26,
    "Phone / WhatsApp": 16,
    "LinkedIn": 32,
    "Website": 28,
    "Core Categories": 36,
    "Specialty Focus": 38,
    "Price Segment": 20,
    "Key Brands Represented": 36,
    "Retail Chains": 36,
    "Specialty Grocers": 36,
    "Foodservice / HoReCa": 36,
    "E-commerce": 28,
    "Channel Strength": 28,
    "Distribution Type": 28,
    "Regions Served": 28,
    "Regional Extensions": 28,
    "Warehouse / Logistics": 32,
    "Delivery / Storage": 28,
    "Competitive Benchmark": 32,
    "Data Confidence": 12,
    "Quality Score": 12,
    "Data Sources": 40,
    "Summary": 48,
}


def classify_region(regions_served: str, hq: str, cities: str, company: str) -> str:
    blob = f"{regions_served} {hq} {cities} {company}".lower()

    # Explicit global / multi-continent
    if any(
        x in blob
        for x in (
            "global",
            "worldwide",
            "70+ countries",
            "30+ countries",
            "31-country",
            "31 countries",
            "multi-region",
            "emea hq; americas",
        )
    ):
        # Country-named Bidfood entities stay local even if HQ says international
        if re.match(r"(?i)bidfood\s+\w+", company.strip()) and "asia" not in company.lower():
            pass  # fall through to country rules
        elif "bidfood asia" in company.lower():
            return "Asia-Pacific"
        else:
            return "Global Multi-Region"

    # Latin America first (before generic "america")
    if any(
        x in blob
        for x in (
            "brazil",
            "chile",
            "colombia",
            "argentina",
            "peru",
            "mexico",
            "latam",
            "latin america",
            "caribbean",
            "costa rica",
            "panama",
            "bahamas",
        )
    ):
        # Mexico often grouped NA — keep LatAm if only Mexico/LatAm signals and no US/Canada
        if "mexico" in blob and not any(x in blob for x in ("united states", "usa", "canada", "u.s.")):
            return "Latin America"
        if any(x in blob for x in ("brazil", "chile", "colombia", "argentina", "peru", "latam", "latin america")):
            return "Latin America"
        if "mexico" in blob and any(x in blob for x in ("united states", "usa", "canada")):
            return "North America"
        if "mexico" in blob:
            return "Latin America"

    # Middle East & Africa
    if any(
        x in blob
        for x in (
            "uae",
            "dubai",
            "abu dhabi",
            "saudi",
            "gcc",
            "middle east",
            "qatar",
            "bahrain",
            "kuwait",
            "oman",
            "pakistan",  # often served via ME trading hubs in this dataset; keep APAC? → South Asia → APAC
            "south africa",
            "nigeria",
            "kenya",
            "egypt",
            "africa",
            "mena",
        )
    ):
        if "pakistan" in blob or "india" in blob or "bangladesh" in blob:
            return "Asia-Pacific"
        return "Middle East and Africa"

    # Asia-Pacific
    if any(
        x in blob
        for x in (
            "asia",
            "australia",
            "new zealand",
            "anz",
            "japan",
            "china",
            "singapore",
            "malaysia",
            "indonesia",
            "thailand",
            "korea",
            "india",
            "hong kong",
            "taiwan",
            "philippines",
            "vietnam",
            "sarawak",
            "asia-pacific",
            "asia pacific",
        )
    ):
        return "Asia-Pacific"

    # Europe
    if any(
        x in blob
        for x in (
            "united kingdom",
            " uk",
            "uk ",
            "england",
            "scotland",
            "wales",
            "netherlands",
            "germany",
            "france",
            "spain",
            "belgium",
            "poland",
            "portugal",
            "italy",
            "sweden",
            "ireland",
            "switzerland",
            "czech",
            "turkey",
            "austria",
            "denmark",
            "norway",
            "finland",
            "greece",
            "hungary",
            "romania",
            "baltics",
            "latvia",
            "lithuania",
            "estonia",
            "europe",
            "benelux",
        )
    ):
        return "Europe"

    # North America
    if any(
        x in blob
        for x in (
            "united states",
            "usa",
            "u.s.",
            "canada",
            "québec",
            "quebec",
            "north america",
            "northeast united states",
            "midwest",
            "texas",
            "california",
        )
    ):
        return "North America"

    return "Global Multi-Region"


def style_header(ws, ncols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def style_body(ws, nrows: int, ncols: int, headers: list[str]) -> None:
    font = Font(name="Calibri", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt = PatternFill("solid", fgColor="F2F2F2")
    for r in range(2, nrows + 1):
        ws.row_dimensions[r].height = 48
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font = font
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r % 2 == 0:
                cell.fill = alt
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS.get(str(h), 18)


def write_sheet(wb, title: str, headers: list[str], rows: list[list]) -> None:
    # Excel sheet name max 31 chars
    name = title[:31]
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c).value = h
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(r, c).value = val
    style_header(ws, len(headers))
    if rows:
        style_body(ws, len(rows) + 1, len(headers), headers)
        ws.auto_filter.ref = ws.dimensions


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = PATH.with_name(f"{PATH.stem}_pre_region_segment_{stamp}{PATH.suffix}")
    shutil.copy2(PATH, bak)

    wb = openpyxl.load_workbook(PATH)
    # Use main data sheet
    src_name = "Avocado Oil Distributors" if "Avocado Oil Distributors" in wb.sheetnames else wb.sheetnames[0]
    src = wb[src_name]
    hdr = [c.value for c in src[1]]
    cols = {h: i + 1 for i, h in enumerate(hdr) if h}

    # Read all rows
    data_rows: list[dict] = []
    for r in range(2, src.max_row + 1):
        row = {h: src.cell(r, cols[h]).value for h in cols}
        company = str(row.get("Company") or "").strip()
        if not company:
            continue
        region = classify_region(
            str(row.get("Regions Served") or ""),
            str(row.get("Headquarters") or ""),
            str(row.get("Cities / Regions") or ""),
            company,
        )
        row["Region Segment"] = region
        data_rows.append(row)

    new_headers = ["Region Segment"] + [h for h in hdr if h]

    # Rebuild main All sheet
    all_rows = [[row.get(h) for h in new_headers] for row in data_rows]

    # Remove old region sheets if re-run
    for name in list(wb.sheetnames):
        if name.startswith("Region -") or name in {"By Region Summary", "All Distributors"}:
            del wb[name]

    # Summary sheet first
    counts = Counter(r["Region Segment"] for r in data_rows)
    if "By Region Summary" in wb.sheetnames:
        del wb["By Region Summary"]
    summary = wb.create_sheet("By Region Summary", 0)
    summary["A1"] = "Region Segment"
    summary["B1"] = "Company Count"
    summary["C1"] = "% of Total"
    summary["D1"] = "Sheet Name"
    tot = len(data_rows)
    for i, region in enumerate(REGIONS, start=2):
        n = counts.get(region, 0)
        summary.cell(i, 1).value = region
        summary.cell(i, 2).value = n
        summary.cell(i, 3).value = round(100.0 * n / tot, 1) if tot else 0
        summary.cell(i, 4).value = f"Region - {region}"[:31]
    summary.cell(len(REGIONS) + 2, 1).value = "TOTAL"
    summary.cell(len(REGIONS) + 2, 2).value = tot
    summary.cell(len(REGIONS) + 2, 3).value = 100.0
    style_header(summary, 4)
    for col, w in enumerate([22, 16, 12, 28], start=1):
        summary.column_dimensions[get_column_letter(col)].width = w
    for r in range(2, len(REGIONS) + 3):
        for c in range(1, 5):
            summary.cell(r, c).alignment = Alignment(vertical="center")
            summary.cell(r, c).font = Font(name="Calibri", size=11)

    # All Distributors
    write_sheet(wb, "All Distributors", new_headers, all_rows)

    # Per-region sheets
    by_region: dict[str, list[list]] = defaultdict(list)
    for row in data_rows:
        by_region[row["Region Segment"]].append([row.get(h) for h in new_headers])

    for region in REGIONS:
        write_sheet(wb, f"Region - {region}", new_headers, by_region.get(region, []))

    # Keep original sheet updated with Region Segment as first col
    # Replace Avocado Oil Distributors content
    # Clear and rewrite src with new headers
    if src.max_row > 1:
        src.delete_rows(2, src.max_row - 1)
    # If Region Segment already somehow present, rebuild columns
    # Delete all columns and rewrite
    if src.max_column:
        src.delete_cols(1, src.max_column)
    for c, h in enumerate(new_headers, start=1):
        src.cell(1, c).value = h
    for r, row in enumerate(all_rows, start=2):
        for c, val in enumerate(row, start=1):
            src.cell(r, c).value = val
    style_header(src, len(new_headers))
    style_body(src, len(all_rows) + 1, len(new_headers), new_headers)

    # Drop noisy audit sheet if present (optional keep)
    # Keep Web Expand Audit if exists

    wb.save(PATH)

    print("backup", bak.name)
    print("total", tot)
    for region in REGIONS:
        print(f"  {region}: {counts.get(region, 0)}")
    print("sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
