# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pathlib import Path

NPD = {
    "", "nan", "none", "n/a", "na", "not publicly disclosed",
    "-", "—", "unknown", "tbd", "null",
}

xlsx = Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market.xlsx")
final = Path(r"D:\vendor-intel\output\run4\global_green_chemical_market_global\global_green_chemical_market_FINAL.xlsx")

def load(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Companies"] if "Companies" in wb.sheetnames else wb.active
    headers = None
    hr = 1
    for r in range(1, 6):
        vals = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=r, max_row=r))]
        if "Company" in vals and "Contact Person" in vals:
            headers = vals
            hr = r
            break
    rows = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        d = {
            headers[i]: ("" if i >= len(row) or row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
        }
        if d.get("Company"):
            rows.append(d)
    wb.close()
    return rows

xrows = load(xlsx)
frows = {r["Company"].strip().lower(): r for r in load(final)}

print("FILLED_XLSX")
for d in xrows:
    p = (d.get("Contact Person") or "").strip()
    if p and p.lower() not in NPD:
        print(f"{d['Company']}\t{p}\t{d.get('Role','')}\t{d.get('LinkedIn','')}")

print("\nEMPTY_XLSX")
n = 0
for d in xrows:
    p = (d.get("Contact Person") or "").strip()
    if p and p.lower() not in NPD:
        continue
    n += 1
    hint = frows.get(d["Company"].strip().lower(), {})
    hp = (hint.get("Contact Person") or "").strip()
    hr = (hint.get("Role") or "").strip()
    hli = (hint.get("LinkedIn") or "").strip()
    print(
        f"{n}\t{d['Company']}\t{d.get('Website','')}\t{d.get('Headquarters','')}\t"
        f"FINAL:{hp}|{hr}|{hli}"
    )
