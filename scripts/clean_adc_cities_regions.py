#!/usr/bin/env python3
"""Clean ADC Cities / Regions to geographic places only.

Remove fluff like: EU distribution, specialty, speciality, wholesale network,
ops, medicines, pharmacy network, life-science distribution, etc.
"""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/web_expand/global_antibody_drug_conjugates_market_global"
XLSX = OUT / "global_antibody_drug_conjugates_market_FINAL.xlsx"
SUMMARY = OUT / "adc_clean_cities_regions_summary.json"

HEADERS = [
    "Company", "Founded", "Headquarters", "Cities / Regions", "Ownership",
    "Business Type", "Employees", "Turnover / Scale", "Contact Person", "Role",
    "Email", "Phone / WhatsApp", "LinkedIn", "Website", "Core Categories",
    "Specialty Focus", "Price Segment", "Key Brands Represented", "Exclusivity",
    "Partnership Duration", "Retail Chains", "Specialty Grocers", "Foodservice / HoReCa",
    "E-commerce", "Channel Strength", "Distribution Type", "Regions Served",
    "Regional Extensions", "Warehouse / Logistics", "Delivery / Storage",
    "Competitive Benchmark", "Data Confidence", "Quality Score", "Data Sources", "Summary",
]
WIDTHS = {
    "Company": 28, "Headquarters": 28, "Cities / Regions": 40,
    "Key Brands Represented": 32, "Specialty Focus": 28, "Core Categories": 28,
    "Summary": 36, "Website": 24, "Regions Served": 22,
}

# Non-place tokens / phrases to strip from a segment
FLUFF_RE = re.compile(
    r"(?i)\b("
    r"eu\s+distribution|europe\s+distribution|apac|asia[- ]pacific|"
    r"mena|latam|gcc|"
    r"distribution|distributor|distributing|"
    r"specialty|speciality|specialised|specialized|"
    r"wholesale|wholesaler|network|networks|"
    r"ops|operations|operating|"
    r"medicines?|pharma(?:ceuticals?)?|pharmacy|pharmacies|"
    r"depot|depots|hubs?|logistics|channel|channels|"
    r"orphan|forward\s+stocking|mail[- ]?order|mail\s+network|"
    r"equipment|life[- ]science|diagnostics?|biotech(?:\s+park)?|"
    r"expansion|cooperative|injectables?|hospital|hospitals|"
    r"nationwide|community\s+oncology|oncology|biologics?|"
    r"research|clinical|infusion|rare(?:[- ]disease)?|"
    r"alliance\s+healthcare|alloga|cencora|dksh|"
    r"parentheticals?|"
    r"hq\b|headquarters?"
    r")\b"
)

# Whole segment is only fluff / region-code with fluff
REGION_ONLY_OK = {
    "europe", "eu", "asia", "apac", "mena", "latam", "gcc", "anz",
    "nordics", "uk", "us", "usa", "uae", "emea",
}

STATE_EXPAND = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas",
    "CA": "California", "CO": "Colorado", "CT": "Connecticut", "DE": "Delaware",
    "FL": "Florida", "GA": "Georgia", "HI": "Hawaii", "ID": "Idaho",
    "IL": "Illinois", "IN": "Indiana", "IA": "Iowa", "KS": "Kansas",
    "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland",
    "MA": "Massachusetts", "MI": "Michigan", "MN": "Minnesota", "MS": "Mississippi",
    "MO": "Missouri", "MT": "Montana", "NE": "Nebraska", "NV": "Nevada",
    "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico", "NY": "New York",
    "NC": "North Carolina", "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma",
    "OR": "Oregon", "PA": "Pennsylvania", "RI": "Rhode Island", "SC": "South Carolina",
    "SD": "South Dakota", "TN": "Tennessee", "TX": "Texas", "UT": "Utah",
    "VT": "Vermont", "VA": "Virginia", "WA": "Washington", "WV": "West Virginia",
    "WI": "Wisconsin", "WY": "Wyoming", "DC": "District of Columbia",
}


def clean_segment(seg: str) -> str:
    s = (seg or "").strip(" .,;:-|/()[]")
    if not s:
        return ""
    # remove parenthetical notes
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"\s+", " ", s).strip(" .,;:-|/")
    # drop if mostly fluff words
    stripped = FLUFF_RE.sub(" ", s)
    stripped = re.sub(r"\s+", " ", stripped).strip(" .,;:-|/")
    # remove leftover connector junk
    stripped = re.sub(r"^(?:and|or|the|of|for|in|with|to)\s+", "", stripped, flags=re.I)
    stripped = re.sub(r"\s+(?:and|or|the|of|for|in|with|to)$", "", stripped, flags=re.I)
    stripped = stripped.strip(" .,;:-|/")

    if not stripped:
        return ""

    low = stripped.lower()
    # reject pure fluff leftovers like "US", "UK" alone? Keep UK/US/UAE/Europe as region labels if alone
    if low in REGION_ONLY_OK:
        # normalize display
        mapping = {
            "eu": "Europe", "europe": "Europe", "us": "United States", "usa": "United States",
            "uk": "United Kingdom", "uae": "United Arab Emirates", "apac": "Asia-Pacific",
            "asia": "Asia", "mena": "MENA", "latam": "Latin America", "gcc": "GCC",
            "anz": "Australia / New Zealand", "nordics": "Nordics", "emea": "EMEA",
        }
        return mapping.get(low, stripped)

    # if still contains strong fluff and little place signal, drop
    if FLUFF_RE.search(stripped) and len(stripped.split()) <= 3:
        # try harder strip
        stripped2 = FLUFF_RE.sub(" ", stripped)
        stripped2 = re.sub(r"\s+", " ", stripped2).strip(" .,;:-|/")
        if not stripped2 or stripped2.lower() in {"us", "uk", "eu"}:
            return mapping_region(stripped2) if stripped2 else ""
        stripped = stripped2

    # Expand "City, ST" US style
    m = re.match(r"^([A-Za-z .'\-]+),\s*([A-Z]{2})$", stripped)
    if m and m.group(2) in STATE_EXPAND:
        return f"{m.group(1).strip()}, {STATE_EXPAND[m.group(2)]}, United States"

    # "City, ST; ..." already handled per segment
    # Reject if starts with lowercase fluff remnant
    if re.match(r"^(network|ops|distribution|specialty|speciality)\b", stripped, re.I):
        return ""

    # Reject company-name-like leftovers
    if re.search(r"alliance healthcare|walgreens|cencora|alloga", stripped, re.I):
        return ""

    # Must look like a place: has a capital letter
    if not re.search(r"[A-ZÀ-ÖØ-Þ]", stripped):
        return ""

    # too long descriptive
    if len(stripped) > 80:
        return ""

    return stripped


def mapping_region(s: str) -> str:
    low = (s or "").lower()
    return {
        "eu": "Europe", "us": "United States", "uk": "United Kingdom",
    }.get(low, s)


def clean_cities(cities: str, hq: str) -> str:
    raw = (cities or "").strip()
    if not raw:
        return hq or ""

    # normalize separators
    raw = raw.replace("|", ";").replace(" / ", "; ").replace("/", "; ")
    # expand common short codes in raw first
    raw = re.sub(r"(?i)\bNZ\b", "New Zealand", raw)
    raw = re.sub(r"(?i)\bUAE\b", "United Arab Emirates", raw)
    parts = re.split(r"\s*;\s*", raw)
    cleaned: list[str] = []
    seen: set[str] = set()
    for p in parts:
        p_work = p
        p_work = re.sub(
            r"(?i)\s*[,;]?\s*(?:EU|Europe|US|UK|APAC|MENA|LATAM|GCC)?\s*"
            r"(?:nationwide\s+)?(?:depot\s+)?(?:specialty|speciality|wholesale|life[- ]science|"
            r"diagnostics?|pharmaceutical|pharmacy|mail)?\s*"
            r"(?:distribution|distributor|network|ops|operations|logistics|channel|medicines?)?\s*$",
            "",
            p_work,
        ).strip(" .,;:-|/")
        c = clean_segment(p_work)
        if not c:
            c = clean_segment(p)
        if not c:
            continue
        key = re.sub(r"\s+", " ", c.lower())
        if key in seen:
            continue
        seen.add(key)
        cleaned.append(c)

    if not cleaned:
        return (hq or "").strip()

    # Ensure HQ city appears if missing
    hq_clean = clean_segment(hq) or (hq or "").strip()
    if hq_clean:
        hq_key = re.sub(r"\s+", " ", hq_clean.lower())
        joined = " ; ".join(cleaned).lower()
        hq_city = hq_clean.split(",")[0].strip().lower()
        if hq_key not in seen and hq_city not in joined:
            cleaned.insert(0, hq_clean)
            seen.add(hq_key)

    # Drop redundant country-only segments if already present inside a longer place
    countries = {
        "united states", "united kingdom", "united arab emirates", "germany", "france",
        "india", "brazil", "japan", "china", "canada", "australia", "italy", "spain",
        "mexico", "singapore", "switzerland", "netherlands", "belgium", "sweden",
        "norway", "denmark", "poland", "portugal", "ireland", "greece", "romania",
        "thailand", "malaysia", "indonesia", "taiwan", "hong kong", "south korea",
        "saudi arabia", "turkey", "austria", "finland", "israel", "egypt", "new zealand",
    }
    final: list[str] = []
    lowered = [re.sub(r"\s+", " ", x.lower()) for x in cleaned]
    for i, c in enumerate(cleaned):
        key = lowered[i]
        if key in countries:
            # keep only if no other segment already contains this country name
            if any(key in lowered[j] and j != i for j in range(len(cleaned))):
                continue
        final.append(c)

    return "; ".join(final) if final else (hq or "").strip()


def write_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ADC Distributors"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title = ws.cell(
        1, 1,
        "Global Antibody Drug Conjugates Market — Distributors (Cities/Regions = places only)",
    )
    title.font = Font(bold=True, size=13, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E79")
    hf = PatternFill("solid", fgColor="2E75B6")
    hfont = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="BDD7EE"),
        right=Side(style="thin", color="BDD7EE"),
        top=Side(style="thin", color="BDD7EE"),
        bottom=Side(style="thin", color="BDD7EE"),
    )
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(2, i, h)
        c.fill = hf
        c.font = hfont
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = thin
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 16)
    for ri, row in enumerate(rows, 3):
        for ci, h in enumerate(HEADERS, 1):
            cell = ws.cell(ri, ci, row.get(h, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
    cws = wb.create_sheet("Column Rules")
    cws["A1"] = "Headquarters = ONE primary HQ location only"
    cws["A2"] = "Cities / Regions = geographic places only (cities/regions/countries)"
    cws["A3"] = "Do NOT put: EU distribution, specialty/speciality, wholesale network, ops, medicines, etc."
    cws["A4"] = f"Cleaned: {datetime.now(timezone.utc).isoformat()} | rows={len(rows)}"
    wb.save(path)


def main() -> int:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    rows: list[dict[str, str]] = []
    for r in range(3, ws.max_row + 1):
        d = {h: str(ws.cell(r, i).value or "").strip() for i, h in enumerate(hdr, 1) if h}
        if d.get("Company"):
            rows.append(d)

    changed = 0
    samples = []
    for row in rows:
        before = row.get("Cities / Regions") or ""
        after = clean_cities(before, row.get("Headquarters") or "")
        if after != before:
            changed += 1
            if len(samples) < 30:
                samples.append({"company": row["Company"], "before": before, "after": after})
        row["Cities / Regions"] = after
        src = row.get("Data Sources") or ""
        if "adc_clean_cities_geo_only" not in src:
            row["Data Sources"] = (src + "; adc_clean_cities_geo_only").strip("; ")

    rows.sort(key=lambda x: x["Company"].lower())
    write_xlsx(XLSX, rows)

    # residual fluff check
    residual = []
    for row in rows:
        v = row.get("Cities / Regions") or ""
        if re.search(r"(?i)\b(distribution|specialty|speciality|wholesale network|ops\b|medicines)\b", v):
            residual.append((row["Company"], v))

    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "rows": len(rows),
        "changed": changed,
        "residual_fluff": residual[:20],
        "residual_count": len(residual),
        "samples": samples,
    }
    SUMMARY.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"DONE changed={changed}/{len(rows)} residual_fluff={len(residual)}")
    for s in samples[:12]:
        print(f"- {s['company']}")
        print(f"  before: {s['before']}")
        print(f"  after:  {s['after']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
