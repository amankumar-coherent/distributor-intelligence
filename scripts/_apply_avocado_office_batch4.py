#!/usr/bin/env python3
"""Apply web-verified office contacts batch 4."""
from __future__ import annotations

import openpyxl
from pathlib import Path

p = Path(
    r"D:\vendor-intel\output\web_expand\global_avocado_oil_market_global"
    r"\global_avocado_oil_market_OFFICE_SLIM.xlsx"
)
wb = openpyxl.load_workbook(p)
ws = wb.active
hdr = [c.value for c in ws[2]]
ic = {str(h): i + 1 for i, h in enumerate(hdr) if h}

FILLS = {
    "McLane Company": {
        "office_no": "(254) 771-7500",
        "office_email": "contact@mclaneco.com",
    },
    "Wakefern Food Corp": {
        "office_no": "(908) 527-3300",
    },
    "SpartanNash": {
        "office_no": "(616) 878-2000",
        "office_email": "IR@spartannash.com",
    },
    "Shamrock Foods": {
        "office_no": "(602) 233-6400",
        "office_email": "azinfo@shamrockfoods.com",
    },
    "Reyes Holdings": {
        "office_no": "(847) 227-6500",
    },
    "Jetro Restaurant Depot": {
        "office_no": "(718) 762-8700",
    },
    "Restaurant Depot": {
        "office_no": "(718) 762-8700",
    },
    "Atlantic Chemicals Trading": {
        "office_no": "+49 40 33 44 19-0",
    },
}

n = 0
for row in range(3, ws.max_row + 1):
    co = str(ws.cell(row, ic["Company"]).value or "").strip()
    if co not in FILLS:
        continue
    hit = FILLS[co]
    if hit.get("office_no"):
        ws.cell(row, ic["Office No."]).value = hit["office_no"]
    if hit.get("office_email"):
        ws.cell(row, ic["Office Email"]).value = hit["office_email"]
    n += 1
    print("filled", co)

wb.save(p)
wb.close()

wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c or "") for c in rows[1]]
ic2 = {h: i for i, h in enumerate(hdr) if h}
phone = sum(
    1
    for r in rows[2:]
    if r[ic2["Office No."]] and str(r[ic2["Office No."]]).strip() != "Not publicly disclosed"
)
email = sum(
    1
    for r in rows[2:]
    if r[ic2["Office Email"]] and str(r[ic2["Office Email"]]).strip() != "Not publicly disclosed"
)
print("TOTAL phone", phone, "email", email, "of", len(rows) - 2, "applied", n)
wb.close()
