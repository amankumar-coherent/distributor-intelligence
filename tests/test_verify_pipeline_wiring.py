"""Verification runs as a second pass after discovery, not inside it.

Requested: "when all company come then apply this verify". Discovery and
verification answer different questions -- discovery asks who is in the
market, verification asks whether a specific company genuinely operates as
a distributor there -- so they stay separate passes over the same rows.

Two properties matter enough to pin:
  1. Verification must be able to run on ANY workbook. MARKETS is a
     hardcoded list, so before this a newly discovered market could not be
     verified without editing the script first.
  2. The market name fed to the criteria prompt must not carry the
     pipeline's "- Distributors Only" filter suffix, which describes the
     FILTER rather than the market.
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

_ROOT = Path(__file__).resolve().parents[1]
_SPEC = importlib.util.spec_from_file_location(
    "verify_distributors_ai_mode", _ROOT / "verify_distributors_ai_mode.py"
)
vd = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(vd)


# --- arbitrary-file support ----------------------------------------------


def test_market_entry_can_be_built_for_any_file():
    """MARKETS is hardcoded; the pipeline needs to describe a brand-new
    market's workbook at call time."""
    entry = vd.market_from_args(
        "output/market_research/global-silicon-carbide-market-distributors-only.xlsx"
    )
    assert entry["path"].endswith("global-silicon-carbide-market-distributors-only.xlsx")
    assert entry["evidence_cols"], "evidence columns must be populated"
    assert entry["key"]


@pytest.mark.parametrize(
    "filename,expected",
    [
        ("global-silicon-carbide-market-distributors-only.xlsx",
         "Global Silicon Carbide Market"),
        ("global-polyurethane-market-distributors-only.xlsx",
         "Global Polyurethane Market"),
        ("global-solar-rooftop-market-distributors.xlsx",
         "Global Solar Rooftop Market"),
    ],
)
def test_filter_suffix_is_stripped_from_the_market_name(filename, expected):
    """"- Distributors Only" describes the FILTER, not the market. Leaving
    it in made the criteria prompt read "...Market - Distributors Only
    market"."""
    assert vd.market_from_args(f"out/{filename}")["market_name"] == expected


def test_explicit_market_name_wins():
    """The pipeline passes the real market string, which preserves casing
    the filename cannot (CPaaS, not Cpaas)."""
    entry = vd.market_from_args("out/whatever.xlsx", "Global CPaaS Market")
    assert entry["market_name"] == "Global CPaaS Market"


def test_evidence_columns_match_the_export_schema():
    """The evidence columns must actually exist in the exported sheet, or
    the model is handed empty strings as its supporting evidence."""
    from vendor_intel.pipeline.market_research import schema

    for col in vd.DEFAULT_EVIDENCE_COLS:
        assert col in schema.COLUMNS, col


# --- fail-closed behaviour ------------------------------------------------


def test_keep_requires_all_three_signals_and_confidence():
    """A plausible-but-wrong KEEP costs more than a missed distributor on a
    run covering thousands of candidates, so every signal must agree."""
    base = {"in_market": True, "is_distributor": True, "fits_criteria": True,
            "confidence": 95}
    assert vd.should_keep(base)

    for flag in ("in_market", "is_distributor", "fits_criteria"):
        assert not vd.should_keep({**base, flag: False}), flag


def test_low_confidence_is_dropped_even_when_flags_agree():
    result = {"in_market": True, "is_distributor": True, "fits_criteria": True,
              "confidence": vd.MIN_CONFIDENCE - 1}
    assert not vd.should_keep(result)


def test_confidence_threshold_is_at_the_boundary():
    at = {"in_market": True, "is_distributor": True, "fits_criteria": True,
          "confidence": vd.MIN_CONFIDENCE}
    assert vd.should_keep(at), "the threshold itself must pass"


def test_criteria_prompt_states_the_required_type_and_market():
    prompt = vd.criteria_prompt("Global Polyurethane Market")
    assert "Global Polyurethane Market" in prompt
    low = prompt.lower()
    assert "distributor" in low
    # ambiguity must fail closed at the model level too
    assert "fits_criteria=false" in low or "fits_criteria: false" in low
    # the parent/subsidiary loophole must be closed
    assert "parent" in low


def test_chunk_size_keeps_the_url_under_googles_limit():
    """The whole prompt rides in ?q=; Google 400s past ~8 KB once
    percent-encoded, and a truncated prompt loses the JSON template."""
    from vendor_intel.scraping.google_ai_mode import MAX_URL_CHARS, build_ai_mode_url

    rows = [
        {
            "Company": f"Some Fairly Long Distributor Company Name {i} Limited",
            "Website": f"https://www.company-number-{i}.example.com",
            "Core Categories": "x" * vd.SUMMARY_CAP,
            "Specialty Focus": "y" * vd.SUMMARY_CAP,
            "Key Products / Technologies": "z" * vd.SUMMARY_CAP,
        }
        for i in range(vd.CHUNK_SIZE)
    ]
    import json

    prompt = (
        vd.criteria_prompt("Global CPaaS (Communication Platform as a Service) Market")
        + "\n"
        + json.dumps(rows, ensure_ascii=False)
    )
    url = build_ai_mode_url(prompt)
    assert len(url) <= MAX_URL_CHARS, f"{len(url)} chars at CHUNK_SIZE={vd.CHUNK_SIZE}"


# --- pipeline wiring ------------------------------------------------------


def test_runner_exposes_a_verify_flag():
    src = (_ROOT / "run_market_research.py").read_text(encoding="utf-8")
    assert '"--verify"' in src


def test_runner_invokes_the_verifier_as_a_subprocess():
    """Not an import: the verifier drives its own Patchright session and
    calls shutdown(). In-process it would contend with the discovery run's
    browser handle for the same profile lock."""
    src = (_ROOT / "run_market_research.py").read_text(encoding="utf-8")
    assert "verify_distributors_ai_mode.py" in src
    assert "subprocess.run" in src


def test_runner_passes_the_real_market_name_through():
    """The real market string is forwarded (so CPaaS keeps its casing
    rather than being derived from the filename as 'Cpaas'), with only the
    filter suffix stripped."""
    src = (_ROOT / "run_market_research.py").read_text(encoding="utf-8")
    assert "--market-name" in src
    assert "_market_name_for_criteria(market)" in src


def test_verification_runs_after_export_not_instead_of_it():
    """The discovery export must survive: verification writes a separate
    file so a bad pass cannot destroy the discovery output."""
    src = (_ROOT / "run_market_research.py").read_text(encoding="utf-8")
    export_at = src.index("export_xlsx(args.market)")
    verify_at = src.index("_run_verification(path, args.market)")
    assert export_at < verify_at, "verification must run after the export"


# --- the banner broke the verifier ---------------------------------------
#
# Live failure: the definition banner added above the table pushed the
# column header to row 5, but load_rows() hardcoded min_row=1. It read the
# banner as the header, found no "Company" column, and printed
# "0 companies ... DONE: 0 kept, 0 dropped" for a file holding 402
# companies. A silent no-op that looked like success.


def _sheet_with_banner(tmp_path, banner_rows: int, companies: int):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    for n in range(banner_rows):
        ws.append([f"banner line {n + 1}"])
    ws.append(["#", "Company", "Website", "Core Categories"])
    for i in range(companies):
        ws.append([i + 1, f"Company {i}", f"https://c{i}.example", "cat"])
    path = tmp_path / "sheet.xlsx"
    wb.save(path)
    return path


def test_header_is_found_below_a_banner(tmp_path):
    path = _sheet_with_banner(tmp_path, banner_rows=4, companies=7)
    rows, headers, _ = vd.load_rows(path)
    assert "Company" in headers
    assert len(rows) == 7, "every company below the banner must be read"
    assert rows[0]["Company"] == "Company 0"


def test_header_on_row_one_still_works(tmp_path):
    """A plain sheet with no banner must not regress."""
    path = _sheet_with_banner(tmp_path, banner_rows=0, companies=5)
    rows, headers, _ = vd.load_rows(path)
    assert "Company" in headers
    assert len(rows) == 5


@pytest.mark.parametrize("banner_rows", [1, 2, 3, 4, 5, 8])
def test_header_found_at_any_banner_depth(tmp_path, banner_rows):
    path = _sheet_with_banner(tmp_path, banner_rows, companies=3)
    rows, _, _ = vd.load_rows(path)
    assert len(rows) == 3, f"failed with {banner_rows} banner rows"


def test_find_header_row_reports_the_right_index(tmp_path):
    path = _sheet_with_banner(tmp_path, banner_rows=4, companies=2)
    import openpyxl

    ws = openpyxl.load_workbook(path)["Companies"]
    assert vd.find_header_row(ws) == 5


def test_find_header_row_falls_back_to_one(tmp_path):
    """A sheet with no Company column at all must not raise."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Foo", "Bar"])
    path = tmp_path / "nocompany.xlsx"
    wb.save(path)

    import openpyxl

    ws = openpyxl.load_workbook(path).active
    assert vd.find_header_row(ws) == 1


def test_the_real_pipeline_export_is_readable(tmp_path, monkeypatch):
    """End-to-end against the actual export writer, so the two cannot
    drift apart again."""
    monkeypatch.setattr(
        "vendor_intel.pipeline.market_research.state.STATE_DIR", tmp_path
    )
    # OUTPUT_DIR too, or export_xlsx() reads the temp state and writes the
    # workbook into the REAL output/market_research/. A test that leaves a
    # file next to live market data is a landmine: an export sitting there
    # is indistinguishable from a real one, and the verifier will happily
    # run against it.
    monkeypatch.setattr(
        "vendor_intel.pipeline.market_research.export_xlsx.OUTPUT_DIR", tmp_path
    )
    from vendor_intel.pipeline.market_research import schema
    from vendor_intel.pipeline.market_research.export_xlsx import export_xlsx
    from vendor_intel.pipeline.market_research.state import MarketResearchState

    market = "Global Reader Test Market - Distributors Only"
    state = MarketResearchState.load(market)
    row = {c: f"v-{c}" for c in (*schema.REQUIRED_FIELDS, *schema.CONTACT_FIELDS)}
    row["Company"] = "Acme Distribution"
    row["Website"] = "https://acme.example"
    state.add_batch([row])
    state.save()

    rows, headers, _ = vd.load_rows(export_xlsx(market))
    assert "Company" in headers
    assert [r["Company"] for r in rows] == ["Acme Distribution"]


# --- market name passed to the criteria prompt ---------------------------


def test_runner_strips_the_filter_suffix():
    """The suffix describes the FILTER, not the market."""
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "run_market_research", _ROOT / "run_market_research.py"
    )
    rmr = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rmr)

    assert (
        rmr._market_name_for_criteria("Global Silicon Carbide Market - Distributors Only")
        == "Global Silicon Carbide Market"
    )
    # casing the filename could not preserve
    assert (
        rmr._market_name_for_criteria(
            "Global CPaaS (Communication Platform as a Service) Market - Distributors Only"
        )
        == "Global CPaaS (Communication Platform as a Service) Market"
    )
    # a market with no suffix is left alone
    assert rmr._market_name_for_criteria("Global Polyurethane Market") == (
        "Global Polyurethane Market"
    )
