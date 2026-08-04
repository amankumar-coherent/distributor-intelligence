"""Tests for FMCG landscape Excel export."""
from vendor_intel.export.fmcg_landscape import (
    AVOCADO_OIL_COLUMNS,
    _EXPORT_COLUMNS,
    _export_columns,
    crawl_snapshot,
    export_fmcg_xlsx,
    row_from_verdict,
)


def test_avocado_column_count():
    assert len(AVOCADO_OIL_COLUMNS) == 30


def test_row_from_verdict_maps_core_fields():
    verdict = {
        "company": "Chosen Foods",
        "brand": "Chosen Foods",
        "website": "https://chosenfoods.com",
        "domain": "chosenfoods.com",
        "role": "Manufacturer",
        "company_function": "manufacturer",
        "key_products": "avocado oil, cooking spray",
        "parent": "Independent (USA)",
        "confidence": 0.82,
        "quality_score": 0.78,
        "is_relevant": True,
        "data_sources": "discovery | smart_crawl",
        "_crawl": {
            "founded_year": 2011,
            "headquarters": "San Diego, California, USA",
            "employee_count": "50-100",
            "email": "hello@chosenfoods.com",
            "products": "avocado oil",
            "crawl_text": "organic extra virgin avocado oil sold at Whole Foods and Costco foodservice bulk",
        },
    }
    row = row_from_verdict(verdict)
    assert row["Company"] == "Chosen Foods"
    assert row["Founded"] == "2011"
    assert row["Headquarters"] == "San Diego, California, USA"
    assert row["Email"] == "hello@chosenfoods.com"
    assert row["Website"] == "https://chosenfoods.com"
    assert row["Foodservice / HoReCa"] == "Yes"
    assert row["Retail Chains"] == "Yes"
    assert "Premium" in row["Price Segment"] or row["Price Segment"]


def test_crawl_snapshot_empty():
    assert crawl_snapshot(None) == {}
    assert crawl_snapshot({"error": "no_crawl"}) == {}


def test_fmcg_schema_matches_legacy_export_columns():
    """market_type="fmcg" must produce byte-identical output to the original
    hardcoded 30-column schema — every existing FMCG market run keeps
    producing the same Excel it always did."""
    assert _export_columns("fmcg") == _EXPORT_COLUMNS


def test_general_schema_is_smaller_and_drops_fmcg_only_columns():
    general = _export_columns("general")
    fmcg = _export_columns("fmcg")
    assert len(general) < len(fmcg)
    for fmcg_only in ("Retail Chains", "Foodservice / HoReCa", "E-commerce", "Channel Strength"):
        assert fmcg_only not in general
    # Universal firmographic fields still present.
    for universal in ("Company", "Founded", "Headquarters", "Website", "Employees"):
        assert universal in general


def test_unknown_market_type_falls_back_to_general():
    assert _export_columns("some_future_vertical") == _export_columns("general")


def test_export_fmcg_xlsx_defaults_to_fmcg_columns(tmp_path):
    import openpyxl

    result = {
        "query_context": {"industry": "Avocado Oil Market", "country": "global"},
        "relevant_companies": [
            {"company": "Acme", "is_relevant": True, "role": "Manufacturer", "confidence": 0.8}
        ],
    }
    path = export_fmcg_xlsx(result, tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[3]]
    assert "Retail Chains" in headers
    assert len(headers) == len(_EXPORT_COLUMNS)


def test_export_fmcg_xlsx_general_market_type_omits_fmcg_columns(tmp_path):
    import openpyxl

    result = {
        "query_context": {"industry": "High Energy Laser Market", "country": "global"},
        "relevant_companies": [
            {"company": "LaserCo", "is_relevant": True, "role": "Manufacturer", "confidence": 0.8}
        ],
    }
    path = export_fmcg_xlsx(result, tmp_path / "out.xlsx", market_type="general")
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[3]]
    assert "Retail Chains" not in headers
    assert "Company" in headers
    assert len(headers) == len(_export_columns("general"))
