#!/usr/bin/env python3
"""Fill blank emails by copying a company's own verified email pattern.

Only applied where the company has ONE consistent domain across its already
verified addresses. That restriction matters: in this data Subway's known
contacts sit on @subway.com, @dominus.technology and a @knights.ucf.edu
university address, and Nuberg's on @indianperoxide.com and @nubergepc.com
— constructing "firstname@subway.com" from those would invent an address
for a person who does not use it.

Every constructed value is written with Email Status = "CONSTRUCTED (not
verified)" and Email Source = "pattern", so a guess can always be told
apart from a verified address downstream.
"""
import argparse
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))


def local_part_style(email: str, name: str) -> str | None:
    """Which pattern produced this address, judged against its owner's name."""
    local = email.split("@")[0].lower()
    parts = [p for p in re.split(r"[^a-z]+", (name or "").lower()) if p]
    if not parts:
        return None
    first, last = parts[0], parts[-1]
    if local == first:
        return "first"
    if local == f"{first}.{last}":
        return "first.last"
    if local == f"{first}{last}":
        return "firstlast"
    if local == f"{first[:1]}{last}":
        return "flast"
    if local == f"{first}_{last}":
        return "first_last"
    if local == last:
        return "last"
    return None


def build(name: str, style: str, domain: str) -> str | None:
    # Scraped names can carry stray glyphs ("Vivan Kumawat ↗"); splitting on
    # non-letters drops them, but reject anything left too short to be a
    # real first+last rather than building a malformed address.
    parts = [p for p in re.split(r"[^a-z]+", (name or "").lower()) if len(p) > 1]
    if len(parts) < 2:
        return None
    first, last = parts[0], parts[-1]
    local = {
        "first": first,
        "first.last": f"{first}.{last}",
        "firstlast": f"{first}{last}",
        "flast": f"{first[:1]}{last}",
        "first_last": f"{first}_{last}",
        "last": last,
    }.get(style)
    return f"{local}@{domain}" if local else None


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", default=str(ROOT / "output/salesql/leadership_verified_emails.xlsx"))
    parser.add_argument("--output", default=None, help="defaults to overwriting --input")
    parser.add_argument("--apply", action="store_true", help="write the file (otherwise dry-run)")
    args = parser.parse_args()

    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    src = Path(args.input)
    df = pd.read_excel(src)
    if "Email Source" not in df.columns:
        df["Email Source"] = df["Email"].apply(lambda v: "salesql verified" if isinstance(v, str) and v else "")

    filled = 0
    skipped = []

    for company, group in df.groupby("Company"):
        known = group[group["Email"].notna()]
        blanks = group[group["Email"].isna()]
        if known.empty or blanks.empty:
            continue

        domains = {str(e).split("@")[-1].lower() for e in known["Email"]}
        if len(domains) != 1:
            skipped.append((company, f"{len(domains)} different domains: {', '.join(sorted(domains))}"))
            continue
        domain = domains.pop()

        styles = Counter()
        for _, r in known.iterrows():
            s = local_part_style(str(r["Email"]), str(r["Name"]))
            if s:
                styles[s] += 1
        if not styles:
            skipped.append((company, "cannot infer a pattern from the known email"))
            continue
        style = styles.most_common(1)[0][0]

        for idx, r in blanks.iterrows():
            guess = build(str(r["Name"]), style, domain)
            if not guess:
                skipped.append((f"{company} / {r['Name']}", "name has no usable first+last"))
                continue
            df.at[idx, "Email"] = guess
            df.at[idx, "Email Status"] = "CONSTRUCTED (not verified)"
            df.at[idx, "Email Source"] = f"pattern {style}@{domain}"
            filled += 1
            print(f"  FILL  {company[:34]:36} {str(r['Name'])[:22]:24} -> {guess}")

    print(f"\nconstructed: {filled}")
    if skipped:
        print("skipped (left blank on purpose):")
        for what, why in skipped:
            print(f"  - {what[:46]:48} {why}")

    if not args.apply:
        print("\nDry run. Re-run with --apply to write.")
        return

    out = Path(args.output or src)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Leadership Emails")
        ws = writer.sheets["Leadership Emails"]
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        body = Alignment(horizontal="left", vertical="top", wrap_text=False)
        amber = PatternFill("solid", fgColor="FFF2CC")
        status_col = list(df.columns).index("Email Status") + 1
        for idx, column in enumerate(df.columns, start=1):
            lengths = sorted(len(str(v)) for v in df[column].tolist() if v not in (None, ""))
            typical = lengths[min(int(len(lengths) * 0.9), len(lengths) - 1)] if lengths else 0
            ws.column_dimensions[get_column_letter(idx)].width = min(
                max(typical + 2, len(str(column)) + 2, 10), 40
            )
            for cell in ws[get_column_letter(idx)][1:]:
                cell.alignment = body
        # Shade constructed rows so a guess is visible at a glance.
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=status_col).value or "").startswith("CONSTRUCTED"):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = amber
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 24
    print(f"\nwritten: {out}")


if __name__ == "__main__":
    main()
