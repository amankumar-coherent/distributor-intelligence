#!/usr/bin/env python3
"""Split cards_unique_companies_with_role.xlsx into one workbook per role.

Reads the classified companies and writes output/company_roles/<Role>.xlsx
(Manufacturer.xlsx, Brand.xlsx, Distributor.xlsx, ...), each keeping the
same columns as the source and renumbering the serial column from 1 so
each file reads as a standalone list.

Rows are split on the Role column exactly as it stands -- no reclassifying
and no merging of roles. Unknown gets its own file rather than being
dropped, because those rows are mostly non-commercial bodies (a
university, a government research centre, a trade agency) that are
genuinely not Brand/Manufacturer/Distributor, and silently discarding them
would lose companies from the total.

Usage:
    python split_roles_excel.py
    python split_roles_excel.py --source other.xlsx --outdir some/dir
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DEFAULT_SRC = ROOT / "cards_unique_companies_with_role.xlsx"
DEFAULT_OUT = ROOT / "output" / "company_roles"


def safe_name(role: str) -> str:
    """Filesystem-safe file stem for a role label."""
    cleaned = re.sub(r"[^A-Za-z0-9 _-]", "", role).strip()
    return cleaned or "Unspecified"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", default=str(DEFAULT_SRC), help="Workbook to split")
    ap.add_argument("--outdir", default=str(DEFAULT_OUT), help="Directory for the per-role files")
    ap.add_argument("--sheet", default="Unique Companies", help="Sheet to read")
    args = ap.parse_args()

    import openpyxl

    src = Path(args.source)
    if not src.exists():
        raise SystemExit(f"Source not found: {src}")

    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb[args.sheet] if args.sheet in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    data = rows[1:]

    try:
        role_idx = header.index("Role")
    except ValueError:
        raise SystemExit(f"No 'Role' column in {src.name}. Columns: {header}")

    buckets: dict[str, list] = {}
    for r in data:
        role = (r[role_idx] or "").strip() if isinstance(r[role_idx], str) else ""
        buckets.setdefault(role or "Unknown", []).append(r)

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    widths = [8, 22, 26, 40, 20, 20, 30, 28, 55, 14]
    written = []
    for role, items in sorted(buckets.items(), key=lambda kv: -len(kv[1])):
        out = openpyxl.Workbook()
        sheet = out.active
        sheet.title = safe_name(role)[:31] or "Sheet1"
        sheet.append(header)
        for i, r in enumerate(items, 1):
            # Renumber so each file is a standalone 1..N list; the original
            # serial stays recoverable via the combined workbook.
            sheet.append([i] + list(r[1:]))
        for col, w in zip("ABCDEFGHIJKLMN", widths):
            sheet.column_dimensions[col].width = w
        sheet.freeze_panes = "A2"
        path = outdir / f"{safe_name(role)}.xlsx"
        out.save(path)
        written.append((role, len(items), path))

    total = sum(n for _, n, _ in written)
    print(f"Source: {src.name}  ({len(data)} rows)\n")
    for role, n, path in written:
        print(f"  {role:<14} {n:4d}  ->  {path}")
    print(f"\n{len(written)} files written to {outdir}")
    print(f"Row check: {total} split == {len(data)} source" if total == len(data)
          else f"WARNING: {total} split != {len(data)} source")


if __name__ == "__main__":
    main()
