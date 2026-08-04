#!/usr/bin/env python3
"""
Reorder Excel columns so Summary comes right after Quality Score,
and Data Sources moves to the last column.
Current:  ... Quality Score | Data Sources | Summary
Target:   ... Quality Score | Summary | Data Sources
"""
import openpyxl
import copy
from pathlib import Path
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent

MARKETS = [
    "avocado_oil",
    "technical_foam",
    "spectral_sensor",
    "gel_ice_pack",
    "cyber_physical_systems_security",
    "high_energy_laser",
]


def find_header_row(ws) -> int:
    for i, row in enumerate(ws.rows, 1):
        if row[0].value == "Company":
            return i
    raise ValueError("No 'Company' header row found")


def get_col_index(ws, header_row: int, col_name: str) -> int:
    """Return 1-indexed column number for a given header name."""
    for cell in ws[header_row]:
        if cell.value == col_name:
            return cell.column
    raise ValueError(f"Column '{col_name}' not found in header row {header_row}")


def swap_columns(ws, col_a: int, col_b: int):
    """Swap entire columns col_a and col_b (1-indexed) in the worksheet."""
    max_row = ws.max_row
    for row_num in range(1, max_row + 1):
        cell_a = ws.cell(row=row_num, column=col_a)
        cell_b = ws.cell(row=row_num, column=col_b)

        # Swap values
        val_a = cell_a.value
        val_b = cell_b.value
        cell_a.value = val_b
        cell_b.value = val_a

        # Swap styles (copy font, fill, alignment, border)
        def copy_style(src, dst):
            if src.has_style:
                dst.font = copy.copy(src.font)
                dst.fill = copy.copy(src.fill)
                dst.alignment = copy.copy(src.alignment)
                dst.border = copy.copy(src.border)
                dst.number_format = src.number_format

        style_a_font = copy.copy(cell_a.font) if cell_a.has_style else None
        style_a_fill = copy.copy(cell_a.fill) if cell_a.has_style else None
        style_a_align = copy.copy(cell_a.alignment) if cell_a.has_style else None
        style_a_border = copy.copy(cell_a.border) if cell_a.has_style else None
        style_a_fmt = cell_a.number_format

        if cell_b.has_style:
            cell_a.font = copy.copy(cell_b.font)
            cell_a.fill = copy.copy(cell_b.fill)
            cell_a.alignment = copy.copy(cell_b.alignment)
            cell_a.border = copy.copy(cell_b.border)
            cell_a.number_format = cell_b.number_format

        if style_a_font:
            cell_b.font = style_a_font
            cell_b.fill = style_a_fill
            cell_b.alignment = style_a_align
            cell_b.border = style_a_border
            cell_b.number_format = style_a_fmt

    # Swap column widths
    col_a_letter = get_column_letter(col_a)
    col_b_letter = get_column_letter(col_b)
    dim_a = ws.column_dimensions[col_a_letter].width
    dim_b = ws.column_dimensions[col_b_letter].width
    ws.column_dimensions[col_a_letter].width = dim_b
    ws.column_dimensions[col_b_letter].width = dim_a


def process_market(market: str):
    xlsx_path = ROOT / f"output/{market}/global_{market}_market_global.xlsx"
    if not xlsx_path.exists():
        print(f"  [SKIP] File not found: {xlsx_path.name}")
        return

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    header_row = find_header_row(ws)

    try:
        col_summary = get_col_index(ws, header_row, "Summary")
        col_datasrc = get_col_index(ws, header_row, "Data Sources")
    except ValueError as e:
        print(f"  [SKIP] {e}")
        return

    print(f"  Before: Data Sources=col{col_datasrc}, Summary=col{col_summary}")

    # If Summary is already before Data Sources, nothing to do
    if col_summary < col_datasrc:
        print(f"  Already in correct order — skipping")
        wb.close()
        return

    # Swap: Summary goes where Data Sources is, Data Sources goes to end
    swap_columns(ws, col_summary, col_datasrc)

    # Verify
    new_col_summary = get_col_index(ws, header_row, "Summary")
    new_col_datasrc = get_col_index(ws, header_row, "Data Sources")
    print(f"  After:  Summary=col{new_col_summary}, Data Sources=col{new_col_datasrc}")

    wb.save(xlsx_path)
    print(f"  Saved: {xlsx_path.name}")


def main():
    print("=" * 60)
    print("REORDERING COLUMNS: Summary -> before Data Sources")
    print("Target order: ...Quality Score | Summary | Data Sources")
    print("=" * 60)

    for market in MARKETS:
        print(f"\n[{market.upper()}]")
        process_market(market)

    print("\n" + "=" * 60)
    print("DONE — column order updated in all Excel files.")
    print("=" * 60)


if __name__ == "__main__":
    main()
