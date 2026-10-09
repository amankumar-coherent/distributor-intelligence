#!/usr/bin/env python3
"""Gap-fill empty cells in wearable_medical_devices_CORE5_COMPLETE.xlsx only.

Scope: this file is a 358-row filtered view (Website+Founded+HQ+Geography+
Operational Presence already present for every row) of the 2306-company
Wearable Medical Devices checkpoint. It is not wired into refill_accepted.py
(that operates on the full checkpoint and writes to a different FINAL file).
This script runs the existing gap-fill engine directly against these 358
rows and writes results back into this same file, touching only cells that
are currently empty/placeholder.
"""
from __future__ import annotations

import argparse
import asyncio
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

# Company names in this market include non-cp1252 characters (e.g. Č); when
# stdout is redirected to a file on Windows it defaults to cp1252 and a bare
# print() crashes the whole run. Force UTF-8 with replacement so no company
# name can kill the batch.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import openpyxl

DEFAULT_PATH = ROOT / "output" / "web_expand" / "wearable_medical_devices_global" / "wearable_medical_devices_CORE5_COMPLETE.xlsx"
SHEET_NAME = "Wearable Medical Devices"

# xlsx column -> gap-fill _crawl key (only the ones this file actually leaves empty)
COLUMN_TO_GAP_KEY = {
    "Ownership": "ownership",
    "Employees": "employee_count",
    "Specialty Focus": "specialty_focus",
    "Online / Offline / Both": "retail_ecommerce",
    "Contact Person": "contact_person",
    "Role": "contact_role",
    "Email": "email",
    "LinkedIn": "linkedin",
    "Office No.": "phone",
    "Office Email": "office_email",
}
# Seed these too so gap-fill treats them as already known and never re-asks.
SEED_ONLY_COLUMNS = {
    "Website": "website",
    "Founded": "founded_year",
    "Headquarters": "headquarters",
    "Continent / Geography": "regions_served",
    "Operational Presence": "offices",
    "Core Categories": "core_categories",
    "Key Brands Represented": "key_brands_represented",
    "Summary": "summary",
}

SAVE_EVERY = 25
SOURCES = ("contact_page", "leadership", "backfill", "wikidata")


def _load_env() -> None:
    import run_cursor_expand as rce
    rce._load_env()


async def _run(path: Path, limit: int | None, start_row: int) -> None:
    from vendor_intel.clients.search_router import FreeSearchRouter
    from vendor_intel.config import Settings
    from vendor_intel.enrichment.gap_fill.orchestrator import gap_fill_row

    os.environ.setdefault("GAP_FILL_SOURCE_TIMEOUT_SEC", "45")

    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {h: i + 1 for i, h in enumerate(headers)}

    total_rows = ws.max_row - 1
    last_row = min(start_row - 1 + limit, ws.max_row) if limit else ws.max_row
    n = total_rows

    router = FreeSearchRouter(Settings.load())
    per_company_budget = 120.0
    filled_any = 0

    for i in range(start_row, last_row + 1):
        company = ws.cell(row=i, column=col_idx["Company"]).value or ""
        website = ws.cell(row=i, column=col_idx["Website"]).value or ""

        crawl: dict[str, str] = {}
        for col, key in {**COLUMN_TO_GAP_KEY, **SEED_ONLY_COLUMNS}.items():
            val = str(ws.cell(row=i, column=col_idx[col]).value or "").strip()
            if val and val.lower() != "not publicly disclosed":
                crawl[key] = val

        row = {
            "company": company,
            "brand": company,
            "domain": website,
            "website": website,
            "country": "",
            "market": "Wearable Medical Devices",
            "_crawl": crawl,
        }

        idx = i - 1
        try:
            await asyncio.wait_for(
                gap_fill_row(row, router=router, owler_client=None, owler_model="",
                             sources=SOURCES, idx=idx, total=n),
                timeout=per_company_budget,
            )
        except asyncio.TimeoutError:
            print(f"[{idx}/{n}] {str(company)[:44]} - timed out, keeping partial")
        except Exception as exc:
            print(f"[{idx}/{n}] {str(company)[:44]} - gap-fill failed: {exc}")
            continue

        new_crawl = row.get("_crawl") or {}
        changed = []
        for col, key in COLUMN_TO_GAP_KEY.items():
            existing = str(ws.cell(row=i, column=col_idx[col]).value or "").strip()
            if existing and existing.lower() != "not publicly disclosed":
                continue
            value = str(new_crawl.get(key) or "").strip()
            if value:
                ws.cell(row=i, column=col_idx[col]).value = value
                changed.append(col)

        label = ", ".join(changed) if changed else "nothing new"
        print(f"[{idx}/{n}] {str(company)[:44]} - {label}")
        if changed:
            filled_any += 1

        if idx % SAVE_EVERY == 0:
            wb.save(path)
            print(f"    ...progress saved ({idx}/{n})")

    wb.save(path)
    print(f"\ngap-fill improved {filled_any}/{n} companies, saved -> {path}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--path", default=str(DEFAULT_PATH))
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--start-row", type=int, default=2, help="xlsx row to resume from (2 = first data row)")
    args = ap.parse_args()

    _load_env()
    asyncio.run(_run(Path(args.path), args.limit, args.start_row))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
