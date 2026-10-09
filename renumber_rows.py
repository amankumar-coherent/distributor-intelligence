#!/usr/bin/env python3
"""Resequence the "#" column so it reads 1..N with no gaps.

The exported "#" carries the DISCOVERY index, not the row position: a
company discovered 546th keeps 546 even when verification dropped 125 of
the companies before it. The delivered sheets therefore run 1..546 across
421 rows, and Solar Rooftop jumps 1, 2, 4 in its first three rows, which
reads as missing data.

Also stores the value as a NUMBER. The export writes text ("1"), so Excel
sorts 1, 10, 100, 2 and right-alignment is lost.

Edits in place, preserving the banner rows, widths, freeze pane and
autofilter.

Usage:
    python renumber_rows.py <file.xlsx> [more.xlsx ...]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def header_row(ws, limit: int = 12) -> int | None:
    """Row index of the real header, past the single-cell banner rows."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit,
                                         values_only=True), 1):
        filled = sum(1 for c in row if c is not None and str(c).strip())
        if filled > 3:
            return i
    return None


def renumber(path: Path) -> str:
    wb = load_workbook(path)
    ws = wb.active
    hdr = header_row(ws)
    if hdr is None:
        return "no header found - skipped"
    if str(ws.cell(row=hdr, column=1).value or "").strip() != "#":
        return "column A is not '#' - skipped"

    # Anchor on the Company column: a blank "#" must still be renumbered,
    # so row presence cannot be decided from the number itself.
    n = 0
    first, last = None, None
    for r in range(hdr + 1, ws.max_row + 1):
        company = ws.cell(row=r, column=2).value
        if company is None or not str(company).strip():
            continue
        n += 1
        cell = ws.cell(row=r, column=1)
        if first is None:
            first = cell.value
        last = cell.value
        cell.value = n
        cell.number_format = "0"

    wb.save(path)
    return f"{n} rows renumbered 1..{n} (was {first}..{last})"


def main() -> None:
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        raise SystemExit(2)
    for raw in args:
        p = Path(raw)
        if not p.is_file():
            print(f"  {p.name}: not found")
            continue
        print(f"  {p.name}: {renumber(p)}")


if __name__ == "__main__":
    main()
