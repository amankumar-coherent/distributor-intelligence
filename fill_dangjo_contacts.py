#!/usr/bin/env python3
"""Fill the blank Contact Person / Role / LinkedIn cells in a delivered workbook.

Works on the XLSX directly rather than on market state, because the
delivered file is a curated subset (140 double-verified companies) that no
longer matches any single state file.

ONLY blank cells are written. A company that already has a contact keeps
it untouched -- those 30 were checked in earlier passes and re-asking would
risk replacing a good name with a worse one.

Person priority, per the user: procurement/purchasing lead or sales
leadership first; a C-suite officer only when no commercial contact is
known. That ordering lives in contact_lookup's prompt.

Blank stays blank. contact_lookup rejects invented emails and
name-derived LinkedIn URLs, and a person without a title (or a title
without a person) is discarded as unusable -- so this pass leaves a cell
empty rather than filling it with something plausible. Expect well under
110 of the 110 to come back filled; that is the pass working, not failing.

Usage:
    python fill_dangjo_contacts.py --file "<path.xlsx>" --dry-run
    python fill_dangjo_contacts.py --file "<path.xlsx>" --limit 15
    python fill_dangjo_contacts.py --file "<path.xlsx>"
"""
from __future__ import annotations

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

from openpyxl import load_workbook  # noqa: E402

from vendor_intel.placeholders.load_keys import apply_env_overrides  # noqa: E402

apply_env_overrides()

from vendor_intel.pipeline.market_research import contact_lookup  # noqa: E402

BATCH = 12


def header_row(ws) -> int:
    for r in range(1, 15):
        if ws.cell(r, 1).value == "#":
            return r
    raise SystemExit("could not find the header row (no '#' cell)")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--file", required=True)
    ap.add_argument("--limit", type=int, default=None,
                    help="only the first N companies needing a contact")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    path = Path(args.file)
    wb = load_workbook(path)
    ws = wb.active
    hrow = header_row(ws)
    hdr = [ws.cell(hrow, c).value for c in range(1, ws.max_column + 1)]

    col = {name: hdr.index(name) + 1 for name in
           ("Company", "Website", "Contact Person", "Role", "LinkedIn")
           if name in hdr}
    for need in ("Company", "Website", "Contact Person", "Role"):
        if need not in col:
            raise SystemExit(f"column {need!r} not in sheet")

    todo = []
    for r in range(hrow + 1, ws.max_row + 1):
        company = ws.cell(r, col["Company"]).value
        if not company:
            continue
        if str(ws.cell(r, col["Contact Person"]).value or "").strip():
            continue                      # already has a contact - leave it
        todo.append({
            "row": r,
            "name": str(company).strip(),
            "website": str(ws.cell(r, col["Website"]).value or "").strip(),
        })

    if args.limit:
        todo = todo[:args.limit]

    print(f"{path.name}: {len(todo)} companies need a contact")
    if not todo:
        return

    filled = {"Contact Person": 0, "Role": 0, "LinkedIn": 0}
    for i in range(0, len(todo), BATCH):
        chunk = todo[i:i + BATCH]
        print(f"  batch {i // BATCH + 1}/{-(-len(todo) // BATCH)} "
              f"({len(chunk)} companies)...", flush=True)
        rows = [{"Company": c["name"], "Website": c["website"]} for c in chunk]
        try:
            answers = contact_lookup.lookup_contacts(rows)
        except Exception as exc:  # noqa: BLE001
            print(f"    ! {type(exc).__name__}: {exc}", flush=True)
            continue

        # lookup_contacts returns {company_name: {field: value}} and does
        # its own name-based matching precisely so an answer can never be
        # written onto the wrong company. Look up by name -- zipping by
        # position would throw that protection away.
        for c in chunk:
            ans = (answers or {}).get(c["name"]) or {}
            person = ans.get("contact_person", "")
            role = ans.get("role", "")
            link = ans.get("linkedin", "")
            # contact_lookup already drops a person without a title; keep
            # that invariant rather than writing half a contact.
            if not person or not role:
                continue
            print(f"    + {c['name'][:32]:34}{person[:22]:24}{role[:28]}")
            filled["Contact Person"] += 1
            filled["Role"] += 1
            if link:
                filled["LinkedIn"] += 1
            if not args.dry_run:
                ws.cell(c["row"], col["Contact Person"]).value = person
                ws.cell(c["row"], col["Role"]).value = role
                if link and "LinkedIn" in col:
                    ws.cell(c["row"], col["LinkedIn"]).value = link

    print("")
    print(f"  contacts found: {filled['Contact Person']} of {len(todo)}"
          f"   (LinkedIn for {filled['LinkedIn']})")

    if args.dry_run:
        print("  dry run - nothing written")
        return

    backup = path.with_suffix(
        f".{datetime.now():%Y%m%d-%H%M%S}.bak.xlsx")
    shutil.copy2(path, backup)
    wb.save(path)
    print(f"  saved (backup: {backup.name})")


if __name__ == "__main__":
    main()
