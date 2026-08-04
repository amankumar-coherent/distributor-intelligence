#!/usr/bin/env python3
import json
import csv
import io
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent

MARKETS = [
    "avocado_oil",
    "technical_foam",
    "spectral_sensor",
    "gel_ice_pack",
    "cyber_physical_systems_security",
    "high_energy_laser",
]

def load_companies_from_json(market: str) -> list[str]:
    json_path = ROOT / f"output/{market}/global_{market}_market_global.json"
    if not json_path.exists():
        return []
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    companies = data.get("relevant_companies") or []
    return [str(c.get("company", "")).strip() for c in companies if c.get("company")]

def load_companies_from_csv(market: str) -> list[str]:
    csv_path = ROOT / f"output/{market}/global_{market}_market_global.csv"
    if not csv_path.exists():
        return []
    with open(csv_path, encoding="utf-8", errors="replace") as f:
        lines = f.readlines()
    if not lines:
        return []
    header = lines[0]
    data_lines = [l for l in lines[1:] if not l.strip().startswith("===")]
    reader = csv.DictReader(io.StringIO(header + "".join(data_lines)))
    companies = []
    for r in reader:
        company = r.get("Company")
        if company:
            companies.append(str(company).strip())
    return companies

def load_companies_from_excel(market: str) -> list[str]:
    xlsx_path = ROOT / f"output/{market}/global_{market}_market_global.xlsx"
    if not xlsx_path.exists():
        return []
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb.active
    
    header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row and row[0] == "Company":
            header_row = i
            break
            
    if header_row is None:
        wb.close()
        return []
        
    companies = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row and row[0]:
            companies.append(str(row[0]).strip())
    wb.close()
    return companies

def main():
    report_lines = []
    report_lines.append("# Market Data Alignment Audit Report\n")
    report_lines.append("Comparing company lists and counts between JSON, CSV, and Excel files for all markets.\n\n")
    
    report_lines.append("| Market | JSON Count | CSV Count | Excel Count | Match? |\n")
    report_lines.append("|---|---|---|---|---|\n")
    
    detailed_sections = []
    
    for m in MARKETS:
        json_companies = load_companies_from_json(m)
        csv_companies = load_companies_from_csv(m)
        excel_companies = load_companies_from_excel(m)
        
        json_count = len(json_companies)
        csv_count = len(csv_companies)
        excel_count = len(excel_companies)
        
        # Lowercase mapping for matching comparison
        json_set = {c.lower(): c for c in json_companies}
        csv_set = {c.lower(): c for c in csv_companies}
        excel_set = {c.lower(): c for c in excel_companies}
        
        # 1. JSON vs CSV
        json_not_csv = [json_set[k] for k in json_set if k not in csv_set]
        csv_not_json = [csv_set[k] for k in csv_set if k not in json_set]
        
        # 2. JSON vs Excel
        json_not_excel = [json_set[k] for k in json_set if k not in excel_set]
        excel_not_json = [excel_set[k] for k in excel_set if k not in json_set]
        
        # 3. CSV vs Excel
        csv_not_excel = [csv_set[k] for k in csv_set if k not in excel_set]
        excel_not_csv = [excel_set[k] for k in excel_set if k not in csv_set]
        
        counts_match = (json_count == csv_count == excel_count)
        names_match = (len(json_not_csv) == 0 and len(csv_not_json) == 0 and 
                       len(json_not_excel) == 0 and len(excel_not_json) == 0 and 
                       len(csv_not_excel) == 0 and len(excel_not_csv) == 0)
        
        overall_match = counts_match and names_match
        match_status = "✅ YES" if overall_match else "❌ NO"
        
        report_lines.append(f"| {m} | {json_count} | {csv_count} | {excel_count} | {match_status} |\n")
        
        if not overall_match:
            sec = []
            sec.append(f"## Market: {m.upper()}\n")
            sec.append(f"- **JSON count**: {json_count}\n")
            sec.append(f"- **CSV count**: {csv_count}\n")
            sec.append(f"- **Excel count**: {excel_count}\n\n")
            
            if json_not_csv:
                sec.append(f"### In JSON but NOT in CSV ({len(json_not_csv)})\n")
                sec.append(", ".join([f"`{c}`" for c in json_not_csv]) + "\n\n")
            if csv_not_json:
                sec.append(f"### In CSV but NOT in JSON ({len(csv_not_json)})\n")
                sec.append(", ".join([f"`{c}`" for c in csv_not_json]) + "\n\n")
                
            if json_not_excel:
                sec.append(f"### In JSON but NOT in Excel ({len(json_not_excel)})\n")
                sec.append(", ".join([f"`{c}`" for c in json_not_excel]) + "\n\n")
            if excel_not_json:
                sec.append(f"### In Excel but NOT in JSON ({len(excel_not_json)})\n")
                sec.append(", ".join([f"`{c}`" for c in excel_not_json]) + "\n\n")
                
            if csv_not_excel:
                sec.append(f"### In CSV but NOT in Excel ({len(csv_not_excel)})\n")
                sec.append(", ".join([f"`{c}`" for c in csv_not_excel]) + "\n\n")
            if excel_not_csv:
                sec.append(f"### In Excel but NOT in CSV ({len(excel_not_csv)})\n")
                sec.append(", ".join([f"`{c}`" for c in excel_not_csv]) + "\n\n")
            
            sec.append("---\n")
            detailed_sections.append("".join(sec))
            
    report_lines.append("\n---\n\n")
    report_lines.extend(detailed_sections)
    
    with open("comparison_report.md", "w", encoding="utf-8") as f:
        f.writelines(report_lines)
    print("Report written to comparison_report.md successfully.")

if __name__ == "__main__":
    main()
