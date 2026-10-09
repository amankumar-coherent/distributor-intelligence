#!/usr/bin/env python3
"""Print how many companies in a market still have no real verdict.

The verifier exits 0 and writes an output file even when chunks failed for
backend reasons (rate limit, wedged browser), so its exit code cannot tell
a caller whether the market is actually finished. This compares the
discovery export against the checkpoint and reports the difference, which
is what a retry loop needs in order to know when to stop.

Prints a single integer to stdout so a shell can compare it directly.

    python unverified_count.py global-polyurethane-market-distributors-only
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name or "").lower())


def main() -> int:
    if len(sys.argv) < 2:
        print("0")
        return 0

    slug = sys.argv[1]
    out = ROOT / "output" / "market_research"
    export = out / f"{slug}.xlsx"
    checkpoint = out / f"{slug}_verify_checkpoint.json"

    if not export.is_file():
        print("0")
        return 0

    import openpyxl

    ws = openpyxl.load_workbook(export, read_only=True).active
    header = 1
    company_col = 1
    for row in range(1, 26):
        values = [str(ws.cell(row, c).value or "") for c in range(1, 30)]
        if "Company" in values:
            header = row
            company_col = values.index("Company") + 1
            break

    names = set()
    for row in range(header + 1, ws.max_row + 1):
        key = _norm(ws.cell(row, company_col).value)
        if key:
            names.add(key)

    judged: set[str] = set()
    if checkpoint.is_file():
        try:
            data = json.loads(checkpoint.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            data = {}
        for verdict in data.get("verdicts") or []:
            # A backend failure is not a verdict -- those companies are
            # exactly the ones a retry needs to pick up.
            if verdict.get("backend_failure"):
                continue
            key = _norm(verdict.get("Company"))
            if key:
                judged.add(key)

    print(len(names - judged))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
